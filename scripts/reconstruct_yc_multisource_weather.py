#!/usr/bin/env python3
"""Audit YC multisource weather and build a candidate only after all gates pass."""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import json
import math
import re
import shutil
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

try:
    import openpyxl
    from docx import Document
    from pypdf import PdfReader
except ImportError as exc:
    raise SystemExit(
        "Missing audit dependency. Install openpyxl, python-docx, and pypdf "
        "in the project runtime, then rerun."
    ) from exc


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "external" / "yc_chinaflux" / "raw"
RESULTS = ROOT / "results" / "yc_multisource_weather_reconstruction"
PREVIOUS_GAPS = ROOT / "results" / "yc_chinaflux_weather_reconciliation" / "unresolved_weather_gaps.csv"
PREVIOUS_LEGACY_INVENTORY = ROOT / "results" / "yc_chinaflux_weather_reconciliation" / "legacy_source_inventory.json"
PREVIOUS_LEGACY_OBSERVATIONS = ROOT / "results" / "yc_chinaflux_weather_reconciliation" / "legacy_source_observations_2004_2010.csv"
TRAIN_START = date(2004, 1, 1)
TRAIN_END = date(2013, 12, 31)
TRAIN_CUTOFF = 2013
MISSING_TEXT = {"", "-", "--", "/", "－", "−", "NA", "N/A", "NAN", "NONE"}
SENTINELS = {-99999.0, -9999.0, 99999.0, 9999.0}


def json_dump(path: Path, obj: Any) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def csv_dump(path: Path, fieldnames: list[str], rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: _csv_value(row.get(k)) for k in fieldnames})


def _csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    return value


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_zip_member(archive: zipfile.ZipFile, member: str) -> str:
    digest = hashlib.sha256()
    with archive.open(member) as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def parsed_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value.upper() in MISSING_TEXT:
            return None
        try:
            value = float(value.replace("－", "-").replace("−", "-"))
        except ValueError:
            return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    value = float(value)
    if not math.isfinite(value) or value in SENTINELS:
        return None
    return value


def require_inside_root(path: Path, label: str) -> Path:
    resolved = path.resolve()
    try:
        resolved.relative_to(ROOT)
    except ValueError as exc:
        raise SystemExit(f"{label} must remain inside the project: {resolved}") from exc
    return resolved


def find_xlsx_by_headers(paths: list[Path], required: set[str]) -> Path:
    matches: list[Path] = []
    for path in paths:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            header = {str(v).strip() for v in next(wb.active.iter_rows(min_row=1, max_row=1, values_only=True)) if v is not None}
        finally:
            wb.close()
        if required <= header:
            matches.append(path)
    if len(matches) != 1:
        raise RuntimeError(f"Expected one workbook with headers {sorted(required)}, found {[p.name for p in matches]}")
    return matches[0]


def pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts.extend(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells) for row in table.rows)
    return "\n".join(parts)


def list_inputs(raw: Path) -> dict[str, Any]:
    direct = sorted(p for p in raw.iterdir() if p.is_file())
    if not direct:
        raise RuntimeError(f"No source files found under {raw}")
    cf_pdfs = {p.stem.replace("YCA_M_", "").lower(): p for p in direct if p.suffix.lower() == ".pdf" and p.name.startswith("YCA_M_")}
    cf_zips = {p.stem.replace("YCA_M_", "").lower(): p for p in direct if p.suffix.lower() == ".zip" and p.name.startswith("YCA_M_")}
    if set(cf_pdfs) != {"30min", "daily", "monthly", "yearly"} or set(cf_zips) != {"30min", "daily", "monthly", "yearly"}:
        raise RuntimeError("ChinaFLUX products could not be identified as four documented PDF/ZIP pairs.")

    historic_pdfs = [p for p in direct if p.suffix.lower() == ".pdf" and "DP2011" in p.name]
    historic_zips = [p for p in direct if p.suffix.lower() == ".zip" and "DP2011" in p.name]
    if len(historic_pdfs) != 1 or len(historic_zips) != 1:
        raise RuntimeError("Expected the official 1998-2006 product PDF and associated ZIP.")

    xlsx = [p for p in direct if p.suffix.lower() == ".xlsx"]
    met_xlsx = find_xlsx_by_headers(xlsx, {"日最高空气温度", "日最低空气温度", "日降水量"})
    rad_xlsx = find_xlsx_by_headers(xlsx, {"总辐射", "净辐射", "光合有效辐射"})
    docx_files = [p for p in direct if p.suffix.lower() == ".docx"]
    if len(docx_files) != 1:
        raise RuntimeError("Expected one 2005-2022 product data dictionary DOCX.")
    return {
        "direct": direct,
        "chinaflux_pdfs": cf_pdfs,
        "chinaflux_zips": cf_zips,
        "historic_pdf": historic_pdfs[0],
        "historic_zip": historic_zips[0],
        "met_xlsx": met_xlsx,
        "rad_xlsx": rad_xlsx,
        "new_docx": docx_files[0],
    }


def inventory_sources(inputs: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    metadata: dict[str, Any] = {
        "scope": "YC/YCA only",
        "analysis_cutoff_year": TRAIN_CUTOFF,
        "validation_weather_values_read": False,
        "products": [],
    }

    cf_titles = {"30min": "ChinaFLUX 2003-2010 半小时气象", "daily": "ChinaFLUX 2003-2010 日值气象", "monthly": "ChinaFLUX 2003-2010 月值气象", "yearly": "ChinaFLUX 2003-2010 年值气象"}
    cf_resolutions = {"30min": "30 min", "daily": "daily", "monthly": "monthly", "yearly": "yearly"}
    cf_vars = "air temperature; precipitation; global solar radiation; humidity; wind; pressure; soil temperature/moisture; net radiation; PAR"
    for level in ("30min", "daily", "monthly", "yearly"):
        doc = inputs["chinaflux_pdfs"][level]
        text = pdf_text(doc)
        if "2003-2010" not in text and "2003 年-2010 年" not in text:
            raise RuntimeError(f"Unexpected ChinaFLUX coverage in {doc.name}")
        resolution = cf_resolutions[level]
        metadata["products"].append({
            "product": cf_titles[level], "time_coverage": "2003-01-01..2010-12-31",
            "resolution": resolution, "variables": cf_vars,
            "units": {"temperature": "degC", "precipitation": "mm", "solar_radiation": "W m-2 for 30min/daily mean; aggregate explicitly before daily MJ m-2"},
            "missing_codes": ["-99999"], "quality_status": "official product-level QA and gap-filled; row-level flag semantics not documented",
            "sensor_heights": {"temperature": "1.6 m and 2.9 m listed, field pairing unresolved", "precipitation": "2.0 m", "global_solar_radiation": "1.5 m"},
            "timestamp_definition": "calendar year/month/day/hour/minute; interval boundary convention not stated in available dictionary",
            "documentation": doc.name,
        })
        rows.append({
            "container_name": doc.name, "file_name": doc.name, "sha256": sha256_file(doc), "product": cf_titles[level],
            "time_coverage": "2003-2010", "resolution": resolution, "variables": cf_vars,
            "units": "temperature degC; precipitation mm; solar radiation W m-2; see product PDF for other fields",
            "missing_codes": "-99999", "sensor_heights": "temperature 1.6/2.9 m unpaired; precipitation 2.0 m; total radiation 1.5 m",
            "data_status": "official QA/infilled product documentation", "role": "metadata",
        })
        archive = inputs["chinaflux_zips"][level]
        rows.append({
            "container_name": archive.name, "file_name": archive.name, "sha256": sha256_file(archive), "product": cf_titles[level],
            "time_coverage": "2003-2010", "resolution": resolution, "variables": cf_vars,
            "units": "see associated product PDF", "missing_codes": "-99999", "sensor_heights": "see associated product PDF",
            "data_status": "official archive", "role": "data archive",
        })
        with zipfile.ZipFile(archive, metadata_encoding="gbk") as z:
            for entry in z.infolist():
                if entry.is_dir():
                    continue
                match = re.search(r"(20\d{2})年", entry.filename)
                year = match.group(1) if match else "2003-2010"
                rows.append({
                    "container_name": archive.name, "file_name": entry.filename,
                    "sha256": sha256_zip_member(z, entry.filename), "product": cf_titles[level],
                    "time_coverage": year, "resolution": resolution, "variables": cf_vars,
                    "units": "see associated product PDF", "missing_codes": "-99999", "sensor_heights": "see associated product PDF",
                    "data_status": "official archived workbook", "role": "archive member",
                })

    hist_pdf = inputs["historic_pdf"]
    hist_text = pdf_text(hist_pdf)
    if "1998-2006" not in hist_text and "1998 年至2006 年" not in hist_text:
        raise RuntimeError("The 1998-2006 product document did not confirm its time coverage.")
    hist_product = {
        "product": "1998-2006 年山东禹城站气象和太阳辐射监测数据",
        "time_coverage": "1998-01-01..2006-12-31", "resolution": "monthly",
        "variables": "air temperature; humidity; pressure; precipitation; wind; surface temperature; solar radiation; PAR; reflected/net radiation; soil heat flux",
        "units": "not specified in product PDF text; workbook unit rows require an XLS-capable parser before numerical use",
        "missing_codes": "not specified in product PDF text",
        "sensor_heights": "not specified in product PDF text; related station PDF contains site/method descriptions but no verified mapping to ChinaFLUX exported temperature columns",
        "quality_status": "digitized monthly statistical product; automated and manual observation summaries",
        "daily_use": "not eligible for daily gap filling or daily WGEN weather",
        "documentation": hist_pdf.name,
    }
    metadata["products"].append(hist_product)
    rows.append({
        "container_name": hist_pdf.name, "file_name": hist_pdf.name, "sha256": sha256_file(hist_pdf),
        "product": hist_product["product"], "time_coverage": "1998-2006", "resolution": "monthly",
        "variables": hist_product["variables"], "units": hist_product["units"], "missing_codes": hist_product["missing_codes"],
        "sensor_heights": hist_product["sensor_heights"], "data_status": hist_product["quality_status"], "role": "metadata",
    })
    hist_zip = inputs["historic_zip"]
    rows.append({
        "container_name": hist_zip.name, "file_name": hist_zip.name, "sha256": sha256_file(hist_zip),
        "product": hist_product["product"], "time_coverage": "1998-2006", "resolution": "monthly",
        "variables": hist_product["variables"], "units": hist_product["units"], "missing_codes": hist_product["missing_codes"],
        "sensor_heights": hist_product["sensor_heights"], "data_status": hist_product["quality_status"], "role": "data archive",
    })
    with zipfile.ZipFile(hist_zip, metadata_encoding="gbk") as z:
        for entry in z.infolist():
            if entry.is_dir():
                continue
            rows.append({
                "container_name": hist_zip.name, "file_name": entry.filename,
                "sha256": sha256_zip_member(z, entry.filename), "product": hist_product["product"],
                "time_coverage": "1998-2006" if entry.filename.lower().endswith(".xls") else "supporting documentation",
                "resolution": "monthly" if entry.filename.lower().endswith(".xls") else "not applicable",
                "variables": hist_product["variables"] if entry.filename.lower().endswith(".xls") else "station/method documentation",
                "units": hist_product["units"], "missing_codes": hist_product["missing_codes"],
                "sensor_heights": hist_product["sensor_heights"],
                "data_status": "official archived monthly workbook" if entry.filename.lower().endswith(".xls") else "associated station documentation",
                "role": "archive member",
            })

    new_text = docx_text(inputs["new_docx"])
    if "2005-2022" not in new_text or "缺测插补" not in new_text:
        raise RuntimeError("The 2005-2022 source document did not confirm coverage and processing method.")
    new_product = {
        "product": "禹城站 2005-2022 年大气环境要素观测数据集",
        "time_coverage": "2005-01-01..2022-12-31 (analysis values clipped to 2013)",
        "resolution": "daily", "temperature_fields": {"mean": "日平均空气温度, degC", "maximum": "日最高空气温度, degC", "minimum": "日最低空气温度, degC"},
        "precipitation": "日降水量, mm", "radiation": "总辐射日总量, MJ m-2; net radiation and PAR are separate fields",
        "missing_codes": ["fullwidth hyphen U+FF0D observed in cells; not described in data dictionary"],
        "sensor_heights": "temperature HMP45D listed, height not specified; total-radiation sensor CM11 listed, height not specified in this product table",
        "quality_status": "hourly monitoring summarized daily; EcoFlow V1 basic QC, missing-value interpolation, daily aggregation; no per-cell observed/imputed flag in workbook",
        "sensor_history": "Milos520 automatic weather station; updated to MAWS301 in May 2014",
        "analysis_cutoff_year": TRAIN_CUTOFF,
        "documentation": inputs["new_docx"].name,
    }
    metadata["products"].append(new_product)
    for path, role, variables, units in [
        (inputs["met_xlsx"], "daily met workbook", "TMEAN; TMAX; TMIN; RH; pressure; wind; RAIN", "temperature degC; precipitation mm; other units in data dictionary"),
        (inputs["rad_xlsx"], "daily radiation workbook", "global total radiation; net radiation; PAR", "total/net radiation MJ m-2; PAR mol m-2 s-1"),
        (inputs["new_docx"], "metadata", "field units; QA; sensors", "see field dictionary"),
    ]:
        rows.append({
            "container_name": path.name, "file_name": path.name, "sha256": sha256_file(path), "product": new_product["product"],
            "time_coverage": "2005-2022", "resolution": "daily", "variables": variables, "units": units,
            "missing_codes": "U+FF0D observed in workbook cells; dictionary does not state code",
            "sensor_heights": new_product["sensor_heights"], "data_status": new_product["quality_status"], "role": role,
        })

    legacy_manifest = json.loads(PREVIOUS_LEGACY_INVENTORY.read_text(encoding="utf-8"))
    for path, role in ((PREVIOUS_GAPS, "003_03 gap-list input"),
                       (PREVIOUS_LEGACY_INVENTORY, "003_03 legacy provenance manifest"),
                       (PREVIOUS_LEGACY_OBSERVATIONS, "003_03 extracted legacy values")):
        rows.append({
            "container_name": "003_03 prior-audit reference", "file_name": path.relative_to(ROOT).as_posix(),
            "sha256": sha256_file(path), "product": "prior-audit evidence used to identify gaps and traceable legacy values",
            "time_coverage": "2004-2010", "resolution": "daily audit extract", "variables": "RAIN; TMAX; TMIN; SRAD",
            "units": "degC; mm; solar radiation as in prior manifest", "missing_codes": "blank retained as unresolved",
            "sensor_heights": "see prior source metadata", "data_status": "inherited 003_03 evidence; not re-extracted in this run", "role": role,
        })
    for source_name, digest in legacy_manifest.get("source_files", {}).items():
        rows.append({
            "container_name": "003_03 legacy_source_inventory.json", "file_name": source_name, "sha256": digest,
            "product": "legacy raw weather workbook referenced by prior audit", "time_coverage": "2004-2010",
            "resolution": "daily records extracted in prior audit", "variables": "see legacy_source_inventory.json",
            "units": "see original workbook; not re-read in this task", "missing_codes": "prior audit preserved blanks",
            "sensor_heights": "not recorded in inherited hash manifest", "data_status": "original workbook hash inherited from 003_03; file not re-read in this task",
            "role": "inherited raw-source hash",
        })

    known = {Path(r["container_name"]).name for r in rows}
    for path in inputs["direct"]:
        if path.name in known:
            continue
        rows.append({
            "container_name": path.name, "file_name": path.name, "sha256": sha256_file(path), "product": "supporting/undetermined source documentation",
            "time_coverage": "see document", "resolution": "see document", "variables": "see document", "units": "see document",
            "missing_codes": "see document", "sensor_heights": "see document", "data_status": "supporting documentation", "role": "metadata",
        })
    metadata["input_rows_policy"] = "Data rows after 2013 are not read; only file-level hashes and metadata are recorded for the 2005-2022 source."
    metadata["products"].append({"product": "raw-source inventory", "files": [p.name for p in inputs["direct"]]})
    return rows, metadata


def daily_yucheng(path: Path, kind: str) -> dict[date, dict[str, float | None]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[date, dict[str, float | None]] = {}
    try:
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            year, month, day = (int(row[i]) for i in (3, 4, 5))
            if year > TRAIN_CUTOFF:
                break
            assert year <= TRAIN_CUTOFF, "2014+ validation weather entered source analysis"
            dt = date(year, month, day)
            if kind == "met":
                raw = {"TMEAN": row[6], "TMAX": row[7], "TMIN": row[9], "RAIN": row[22]}
            else:
                raw = {"SRAD": row[6]}
            out[dt] = {key: parsed_number(value) for key, value in raw.items()}
    finally:
        wb.close()
    return out


def chinaflux_daily(zip_path: Path, years: range) -> dict[date, dict[str, Any]]:
    grouped: dict[date, dict[str, Any]] = defaultdict(lambda: {"n": 0, "near": [], "above": [], "rain": [], "srad": []})
    with zipfile.ZipFile(zip_path, metadata_encoding="gbk") as archive:
        for entry in archive.infolist():
            if entry.is_dir() or not entry.filename.lower().endswith(".xlsx"):
                continue
            year_match = re.search(r"(20\d{2})年", entry.filename)
            if not year_match or int(year_match.group(1)) not in years:
                continue
            year = int(year_match.group(1))
            assert year <= TRAIN_CUTOFF, "2014+ validation weather entered ChinaFLUX analysis"
            with archive.open(entry) as binary:
                wb = openpyxl.load_workbook(binary, read_only=True, data_only=True)
                try:
                    ws = wb.active
                    headers = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
                    required = {"年", "月", "日", "近地面空气温度", "冠层上方空气温度", "太阳辐射", "降水量"}
                    if not required <= set(headers):
                        raise RuntimeError(f"Unexpected ChinaFLUX schema in {entry.filename}: {headers}")
                    indexes = {name: headers.index(name) for name in required}
                    for row in ws.iter_rows(min_row=3, values_only=True):
                        try:
                            dt = date(int(row[indexes["年"]]), int(row[indexes["月"]]), int(row[indexes["日"]]))
                        except (TypeError, ValueError):
                            continue
                        assert dt.year <= TRAIN_CUTOFF
                        item = grouped[dt]
                        item["n"] += 1
                        for label, field in (("near", "近地面空气温度"), ("above", "冠层上方空气温度"), ("rain", "降水量"), ("srad", "太阳辐射")):
                            value = parsed_number(row[indexes[field]])
                            if value is not None:
                                item[label].append(value)
                finally:
                    wb.close()
    result: dict[date, dict[str, Any]] = {}
    for dt, item in grouped.items():
        n = item["n"]
        record: dict[str, Any] = {"n_records": n}
        record["rain_partial_mm"] = sum(item["rain"]) if item["rain"] else 0.0
        record["rain_missing_intervals"] = max(0, n - len(item["rain"])) + max(0, 48 - n)
        record["rain_mm"] = record["rain_partial_mm"] if n == 48 and len(item["rain"]) == 48 else None
        record["rain_status"] = "COMPLETE_48" if record["rain_mm"] is not None else ("NO_ROWS" if n == 0 else "MISSING_INTERVALS")
        for layer in ("near", "above"):
            is_full = n == 48 and len(item[layer]) == 48
            record[f"{layer}_tmax"] = max(item[layer]) if is_full else None
            record[f"{layer}_tmin"] = min(item[layer]) if is_full else None
        is_full_radiation = n == 48 and len(item["srad"]) == 48
        record["srad_mj_m2"] = sum(item["srad"]) * 1800.0 / 1_000_000.0 if is_full_radiation else None
        record["srad_status"] = "COMPLETE_48" if is_full_radiation else "MISSING_INTERVALS"
        result[dt] = record
    return result


def pearson(xs: list[float], ys: list[float]) -> float | None:
    n = len(xs)
    if n < 2:
        return None
    mx, my = sum(xs) / n, sum(ys) / n
    sx = sum((x - mx) ** 2 for x in xs)
    sy = sum((y - my) ** 2 for y in ys)
    if sx == 0 or sy == 0:
        return None
    return sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / math.sqrt(sx * sy)


def paired_metrics(pairs: list[tuple[float | None, float | None]]) -> dict[str, Any]:
    valid = [(float(x), float(y)) for x, y in pairs if x is not None and y is not None]
    if not valid:
        return {"n": 0, "bias_new_minus_cf": None, "mae": None, "rmse": None, "pearson_r": None}
    xs = [x for x, _ in valid]
    ys = [y for _, y in valid]
    diffs = [y - x for x, y in valid]
    return {
        "n": len(valid), "bias_new_minus_cf": sum(diffs) / len(diffs),
        "mae": sum(abs(v) for v in diffs) / len(diffs),
        "rmse": math.sqrt(sum(v * v for v in diffs) / len(diffs)),
        "pearson_r": pearson(xs, ys),
    }


def overlap_analysis(cf: dict[date, dict[str, Any]], met: dict[date, dict[str, float | None]], rad: dict[date, dict[str, float | None]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    pairs: dict[str, list[tuple[date, float | None, float | None]]] = defaultdict(list)
    for dt, item in cf.items():
        if dt.year not in (2005, 2006):
            continue
        newmet = met.get(dt, {})
        newrad = rad.get(dt, {})
        pairs["RAIN"].append((dt, item["rain_mm"], newmet.get("RAIN")))
        pairs["TMAX_near"].append((dt, item["near_tmax"], newmet.get("TMAX")))
        pairs["TMIN_near"].append((dt, item["near_tmin"], newmet.get("TMIN")))
        pairs["TMAX_above"].append((dt, item["above_tmax"], newmet.get("TMAX")))
        pairs["TMIN_above"].append((dt, item["above_tmin"], newmet.get("TMIN")))
        pairs["SRAD"].append((dt, item["srad_mj_m2"], newrad.get("SRAD")))

    rows: list[dict[str, Any]] = []
    detail: dict[str, Any] = {}
    for variable, observations in pairs.items():
        metrics = paired_metrics([(cfv, newv) for _, cfv, newv in observations])
        valid = [(dt, cfv, newv) for dt, cfv, newv in observations if cfv is not None and newv is not None]
        events = None
        if variable == "RAIN":
            agree = sum((cfv > 0) == (newv > 0) for _, cfv, newv in valid)
            events = {"rain_event_threshold_mm": 0.0, "event_agreement_days": agree, "event_agreement_rate": agree / len(valid) if valid else None}
        annual: dict[str, Any] = {}
        monthly: list[dict[str, Any]] = []
        for year in (2005, 2006):
            year_values = [(cfv, newv) for dt, cfv, newv in observations if dt.year == year and cfv is not None and newv is not None]
            if variable in ("RAIN", "SRAD"):
                cf_total = sum(x for x, _ in year_values)
                new_total = sum(y for _, y in year_values)
                annual[str(year)] = {"paired_days": len(year_values), "cf_total": cf_total, "new_total": new_total, "difference_new_minus_cf": new_total - cf_total}
            else:
                annual[str(year)] = {"paired_days": len(year_values), "cf_mean": sum(x for x, _ in year_values) / len(year_values) if year_values else None,
                                     "new_mean": sum(y for _, y in year_values) / len(year_values) if year_values else None}
        for year in (2005, 2006):
            for month in range(1, 13):
                sub = [(dt, cfv, newv) for dt, cfv, newv in valid if dt.year == year and dt.month == month]
                if not sub:
                    continue
                period = f"{year}-{month:02d}"
                if variable in ("RAIN", "SRAD"):
                    cf_total = sum(cfv for _, cfv, _ in sub)
                    new_total = sum(newv for _, _, newv in sub)
                    monthly.append({"year": year, "month": month, "period": period, "paired_days": len(sub), "cf_total": cf_total, "new_total": new_total, "difference_new_minus_cf": new_total - cf_total})
                else:
                    cf_mean = sum(cfv for _, cfv, _ in sub) / len(sub)
                    new_mean = sum(newv for _, _, newv in sub) / len(sub)
                    monthly.append({"year": year, "month": month, "period": period, "paired_days": len(sub), "cf_mean": cf_mean, "new_mean": new_mean, "difference_new_minus_cf": new_mean - cf_mean})
        detail[variable] = {**metrics, "annual": annual, "monthly": monthly, **({"rain_events": events} if events else {})}
        for level, groups in (("annual", annual), ("monthly", {r["period"]: r for r in monthly} )):
            for key, summary in groups.items():
                rows.append({"variable_comparison": variable, "aggregation": level, "period": key, **summary})
    return rows, detail


def build_gap_resolution(cf: dict[date, dict[str, Any]], met: dict[date, dict[str, float | None]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    with PREVIOUS_GAPS.open("r", encoding="utf-8-sig", newline="") as f:
        previous = list(csv.DictReader(f))
    gaps = [r for r in previous if r.get("candidate_resolution") == "unresolved_no_independent_numeric_daily_observation"]
    if len(gaps) != 35:
        raise RuntimeError(f"Expected 35 prior unresolved precipitation dates, found {len(gaps)}")
    rows: list[dict[str, Any]] = []
    provisional = 0
    no_value = 0
    for prior in gaps:
        dt = date.fromisoformat(prior["date"])
        cf_item = cf.get(dt)
        alt = met.get(dt, {}).get("RAIN")
        legacy_value = parsed_number(prior.get("legacy_raw_rain_value_mm"))
        if legacy_value is not None:
            selected, source = legacy_value, "legacy_raw_observation"
            reason, qc = "Explicit numeric legacy raw value exists; retained as a traceable alternate for comparison.", "PROVISIONAL_LEGACY_SOURCE_NOT_FINAL"
            provisional += 1
        elif alt is not None:
            selected, source = alt, "yucheng_2005_2022"
            reason = "Official daily value is numeric, but product-level QC/infill has no day-level observed-versus-imputed flag and overlap totals differ; comparison-only, not accepted into final candidate."
            qc = "PROVISIONAL_PRODUCT_VALUE_SOURCE_CONSISTENCY_UNRESOLVED"
            provisional += 1
        else:
            selected, source = None, "unresolved"
            reason = "No daily numeric source covers this date; monthly products cannot be disaggregated."
            qc = "UNRESOLVED_NO_DAILY_SOURCE"
            no_value += 1
        rows.append({
            "date": dt.isoformat(), "year": dt.year, "doy": dt.timetuple().tm_yday,
            "chinaflux_30min_status": cf_item["rain_status"] if cf_item else "NO_ROWS",
            "chinaflux_30min_missing_records": cf_item["rain_missing_intervals"] if cf_item else None,
            "chinaflux_30min_partial_sum_mm": cf_item["rain_partial_mm"] if cf_item else None,
            "source_1998_2006": "monthly_only_not_eligible_for_daily_fill" if 1998 <= dt.year <= 2006 else "outside_coverage",
            "source_2005_2022": alt, "legacy_raw_source": prior.get("legacy_raw_rain_status", ""),
            "selected_value": selected, "selected_source": source, "decision_reason": reason,
            "qc_status": qc,
        })
    summary = {
        "status": "BLOCKED_PRECIPITATION_GAPS",
        "prior_unresolved_days": len(gaps), "numeric_alternatives_from_2005_2022": sum(r["source_2005_2022"] is not None for r in rows),
        "provisional_values_not_admitted_to_candidate": provisional,
        "no_daily_numeric_source_days": no_value,
        "certified_resolved_days": 0,
        "reason": "The 1998-2006 product is monthly only; the 2005-2022 daily product lacks per-day observed/imputed provenance, and overlap precipitation totals differ by year. Five target dates fall in 2004, before the 2005 daily product begins.",
    }
    return rows, summary


def weather_2011_2013(met: dict[date, dict[str, float | None]], rad: dict[date, dict[str, float | None]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    yearly: dict[str, Any] = {}
    for year in range(2011, 2014):
        expected = (date(year + 1, 1, 1) - date(year, 1, 1)).days
        year_dates = [date(year, 1, 1).fromordinal(date(year, 1, 1).toordinal() + offset) for offset in range(expected)]
        values = {"RAIN": [], "TMAX": [], "TMIN": [], "SRAD": []}
        missing: dict[str, list[str]] = {k: [] for k in values}
        invalid: dict[str, list[str]] = {k: [] for k in values}
        for dt in year_dates:
            m = met.get(dt, {})
            r = rad.get(dt, {})
            record: dict[str, Any] = {"DATE": dt.isoformat(), "YEAR": year, "DOY": dt.timetuple().tm_yday}
            for variable, source_value in (("RAIN", m.get("RAIN")), ("TMAX", m.get("TMAX")), ("TMIN", m.get("TMIN")), ("SRAD", r.get("SRAD"))):
                record[variable] = source_value
                record[f"{variable}_source"] = "yucheng_2005_2022" if source_value is not None else "unresolved"
                values[variable].append(source_value)
                if source_value is None:
                    missing[variable].append(dt.isoformat())
                if variable == "SRAD" and source_value is not None and source_value < 0:
                    invalid[variable].append(dt.isoformat())
            if m.get("TMAX") is not None and m.get("TMIN") is not None and m["TMAX"] < m["TMIN"]:
                invalid.setdefault("TMAX_LT_TMIN", []).append(dt.isoformat())
            record["row_qc"] = "SOURCE_MISSING" if any(record[v] is None for v in values) else ("NEGATIVE_SRAD_REVIEW" if record["SRAD"] < 0 else "SOURCE_VALUE_NOT_FINAL")
            rows.append(record)
        yearly[str(year)] = {
            "expected_days": expected,
            "source_daily_rows": sum(dt in met and dt in rad for dt in year_dates),
            "missing_dates_by_variable": missing,
            "missing_count_by_variable": {k: len(v) for k, v in missing.items()},
            "invalid_dates_by_variable": invalid,
            "value_counts_by_variable": {k: sum(v is not None for v in vals) for k, vals in values.items()},
        }
    blocked = any(any(v for v in details["missing_count_by_variable"].values()) or details["invalid_dates_by_variable"] for details in yearly.values())
    return rows, {
        "status": "BLOCKED_2011_2013_WEATHER" if blocked else "PASS_2011_2013_WEATHER",
        "years": yearly,
        "radiation_definition": "PASS: product dictionary identifies daily total radiation (MJ m-2), separate from net radiation/PAR; CM11 total-radiation sensor listed.",
        "unit_conversion": "No conversion applied to 2005-2022 daily radiation (already MJ m-2); ChinaFLUX 30-minute W m-2 integrated as sum(value * 1800 s) / 1e6.",
        "validation_leakage_check": "Only rows through 2013 were iterated; the loop stops at first 2014 row and asserts year <= 2013.",
    }


def build_provenance(cf: dict[date, dict[str, Any]], met: dict[date, dict[str, float | None]], rad: dict[date, dict[str, float | None]], gap_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    gap_map = {date.fromisoformat(r["date"]): r for r in gap_rows}
    with PREVIOUS_GAPS.open("r", encoding="utf-8-sig", newline="") as f:
        prior_gaps = list(csv.DictReader(f))
    legacy_rain = {
        date.fromisoformat(r["date"]): parsed_number(r.get("legacy_raw_rain_value_mm"))
        for r in prior_gaps
        if r.get("variable") == "RAIN" and parsed_number(r.get("legacy_raw_rain_value_mm")) is not None
    }
    rows: list[dict[str, Any]] = []
    total_days = (TRAIN_END - TRAIN_START).days + 1
    for offset in range(total_days):
        dt = date.fromordinal(TRAIN_START.toordinal() + offset)
        c = cf.get(dt, {})
        m = met.get(dt, {})
        r = rad.get(dt, {})
        selected_rain: float | None = c.get("rain_mm")
        rain_source = "chinaflux_30min" if selected_rain is not None else "unresolved"
        rain_qc = "COMPLETE_48_INTERVALS" if selected_rain is not None else "UNRESOLVED"
        gap = gap_map.get(dt)
        if selected_rain is None and dt.year >= 2011:
            selected_rain = m.get("RAIN")
            rain_source = "yucheng_2005_2022" if selected_rain is not None else "unresolved"
            rain_qc = "SOURCE_VALUE_NOT_FINAL" if selected_rain is not None else "SOURCE_MISSING"
        elif selected_rain is None and dt in legacy_rain:
            selected_rain = legacy_rain[dt]
            rain_source = "legacy_raw_observation"
            rain_qc = "TRACEABLE_LEGACY_RAW_VALUE_NOT_FINAL_SOURCE_PRIORITY"
        elif selected_rain is None and gap:
            selected_rain = gap.get("selected_value")
            rain_source = gap.get("selected_source", "unresolved")
            rain_qc = gap.get("qc_status", "UNRESOLVED")
        selected_srad = c.get("srad_mj_m2") if dt.year <= 2010 else r.get("SRAD")
        srad_source = "chinaflux_30min" if dt.year <= 2010 and selected_srad is not None else ("yucheng_2005_2022" if dt.year >= 2011 and selected_srad is not None else "unresolved")
        srad_qc = "COMPLETE_48_INTERVALS" if dt.year <= 2010 and selected_srad is not None else (
            "SOURCE_VALUE_NOT_FINAL" if dt.year >= 2011 and selected_srad is not None and selected_srad >= 0 else (
                "INVALID_NEGATIVE_SRAD" if dt.year >= 2011 and selected_srad is not None else "UNRESOLVED"))
        if dt.year <= 2010:
            tmax = tmin = None
            tmax_source = tmin_source = "unresolved"
            tmax_qc = tmin_qc = "BLOCKED_TEMPERATURE_MAPPING"
            near_tmax, near_tmin = c.get("near_tmax"), c.get("near_tmin")
            above_tmax, above_tmin = c.get("above_tmax"), c.get("above_tmin")
        else:
            tmax, tmin = m.get("TMAX"), m.get("TMIN")
            tmax_source = "yucheng_2005_2022" if tmax is not None else "unresolved"
            tmin_source = "yucheng_2005_2022" if tmin is not None else "unresolved"
            tmax_qc = "SOURCE_VALUE_HEIGHT_UNSPECIFIED" if tmax is not None else "SOURCE_MISSING"
            tmin_qc = "SOURCE_VALUE_HEIGHT_UNSPECIFIED" if tmin is not None else "SOURCE_MISSING"
            near_tmax = near_tmin = above_tmax = above_tmin = None
        overall = "BLOCKED" if any(v is None for v in (tmax, tmin, selected_srad, selected_rain)) or dt.year <= 2010 or (selected_srad is not None and selected_srad < 0) else "NOT_FINAL_SOURCE_QC"
        rows.append({
            "date": dt.isoformat(), "year": dt.year, "doy": dt.timetuple().tm_yday,
            "TMAX": tmax, "TMAX_source": tmax_source, "TMIN": tmin, "TMIN_source": tmin_source,
            "SRAD": selected_srad, "SRAD_source": srad_source, "RAIN": selected_rain, "RAIN_source": rain_source,
            "TMAX_qc": tmax_qc, "TMIN_qc": tmin_qc, "SRAD_qc": srad_qc, "RAIN_qc": rain_qc, "overall_qc": overall,
            "TMAX_candidate_near": near_tmax, "TMAX_candidate_above": above_tmax,
            "TMIN_candidate_near": near_tmin, "TMIN_candidate_above": above_tmin,
        })
    return rows


def make_report_summary(cf: dict[date, dict[str, Any]], met: dict[date, dict[str, float | None]], rad: dict[date, dict[str, float | None]], gap_summary: dict[str, Any], temp_metrics: dict[str, Any], qc: dict[str, Any]) -> dict[str, Any]:
    height_blocked = True
    gate_a = "BLOCKED_TEMPERATURE_MAPPING" if height_blocked else "PASS_TEMPERATURE_MAPPING"
    gate_b = gap_summary["status"]
    gate_c = qc["status"]
    annual_rain: dict[str, Any] = {}
    for year in (2005, 2006):
        cf_sum = sum(v["rain_mm"] for d, v in cf.items() if d.year == year and v.get("rain_mm") is not None)
        new_values = [r.get("RAIN") for d, r in met.items() if d.year == year and r.get("RAIN") is not None]
        new_sum = sum(new_values)
        annual_rain[str(year)] = {"chinaflux_30min_complete_sum_mm": cf_sum, "yucheng_2005_2022_daily_sum_mm": new_sum,
                                  "difference_mm": new_sum - cf_sum,
                                  "difference_percent_of_chinaflux": 100 * (new_sum - cf_sum) / cf_sum if cf_sum else None}
    secondary = [status for status in (gate_b, gate_c) if status.startswith("BLOCKED")]
    if any(abs(v["difference_percent_of_chinaflux"] or 0) > 0 for v in annual_rain.values()):
        secondary.append("BLOCKED_SOURCE_INCONSISTENCY")
    return {
        "task": "003_04_yc_multisource_weather_reconstruction",
        "final_status": gate_a,
        "secondary_blockers": list(dict.fromkeys(secondary)),
        "site_scope": "YC/YCA only",
        "gates": {"A_temperature": gate_a, "B_precipitation": gate_b, "C_2011_2013": gate_c},
        "source_years_used_for_weather_values": "2004-2013 only",
        "validation_weather_values_read": False,
        "source_inventory_rows": None,
        "temperature_overlap_metrics": temp_metrics,
        "rain_overlap_annual": annual_rain,
        "precipitation_gap_resolution": gap_summary,
        "weather_2011_2013_qc": qc,
        "candidate_built": False,
        "candidate_qc": "NOT_RUN_GATE_BLOCKED",
        "cli_created": False,
        "wgen_run": False,
        "dssat_run": False,
        "ppo_run": False,
    }


def prepare_results_dir(path: Path, overwrite: bool) -> None:
    path.mkdir(parents=True, exist_ok=True)
    existing = [p for p in path.iterdir() if p.is_file()]
    if existing and not overwrite:
        raise SystemExit(f"Results directory already contains files; refusing to overwrite: {path}")
    if existing and overwrite:
        backup = path / "backups" / datetime.now().strftime("%Y%m%d_%H%M%S")
        backup.mkdir(parents=True, exist_ok=True)
        for item in existing:
            shutil.copy2(item, backup / item.name)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--audit-only", action="store_true", help="Audit only; this is the default mode.")
    parser.add_argument("--build-candidate", action="store_true", help="Build the WGEN fitting CSV only if every gate passes.")
    parser.add_argument("--raw-dir", type=Path, default=RAW)
    parser.add_argument("--results-dir", type=Path, default=RESULTS)
    parser.add_argument("--overwrite", action="store_true", help="Back up existing task outputs before replacing them.")
    args = parser.parse_args()
    if args.audit_only and args.build_candidate:
        parser.error("Choose --audit-only or --build-candidate, not both.")
    raw = require_inside_root(args.raw_dir, "raw-dir")
    out = require_inside_root(args.results_dir, "results-dir")
    if not raw.is_dir():
        raise SystemExit(f"Source directory does not exist: {raw}")
    prepare_results_dir(out, args.overwrite)

    inputs = list_inputs(raw)
    inventory, source_metadata = inventory_sources(inputs)
    cf = chinaflux_daily(inputs["chinaflux_zips"]["30min"], range(2004, 2011))
    met = daily_yucheng(inputs["met_xlsx"], "met")
    rad = daily_yucheng(inputs["rad_xlsx"], "rad")
    training_weather_years = sorted({dt.year for dataset in (cf, met, rad) for dt in dataset})
    assert max(training_weather_years) <= TRAIN_CUTOFF, "training weather includes validation-year values"
    overlap_rows, overlap = overlap_analysis(cf, met, rad)

    near_metrics = {k: overlap[k] for k in ("TMAX_near", "TMIN_near")}
    above_metrics = {k: overlap[k] for k in ("TMAX_above", "TMIN_above")}
    temp_evidence: list[dict[str, Any]] = []
    for candidate, metrics in (("近地面空气温度", near_metrics), ("冠层上方空气温度", above_metrics)):
        for variable, result in metrics.items():
            temp_evidence.append({
                "overlap_years": "2005-2006", "n_paired_days": result["n"], "chinaflux_candidate_field": candidate,
                "dssat_variable": variable.split("_")[0], "comparison_source": "yucheng_2005_2022 daily TMAX/TMIN; sensor height unspecified",
                "bias_daily_minus_chinaflux": result["bias_new_minus_cf"], "MAE": result["mae"], "RMSE": result["rmse"],
                "pearson_r": result["pearson_r"], "mapping_decision": "NOT_SELECTED_HEIGHT_MAPPING_UNRESOLVED",
                "height_evidence": "ChinaFLUX data dictionary lists 1.6 m and 2.9 m together without pairing them to these exported columns; comparison product lists HMP45D but no temperature height.",
            })
    csv_dump(out / "temperature_mapping_evidence.csv", list(temp_evidence[0]), temp_evidence)

    gap_rows, gap_summary = build_gap_resolution(cf, met)
    csv_dump(out / "precipitation_gap_resolution.csv", list(gap_rows[0]), gap_rows)
    weather_rows, weather_qc = weather_2011_2013(met, rad)
    weather_fields = ["DATE", "YEAR", "DOY", "RAIN", "RAIN_source", "TMAX", "TMAX_source", "TMIN", "TMIN_source", "SRAD", "SRAD_source", "row_qc"]
    csv_dump(out / "weather_2011_2013_reconstruction.csv", weather_fields, weather_rows)

    mapping = {
        "status": "BLOCKED_TEMPERATURE_MAPPING", "selected_field": None,
        "candidate_fields": [
            {"field": "近地面空气温度", "physical_meaning": "below vegetation canopy", "height_m": None,
             "numeric_evidence": {"TMAX": overlap["TMAX_near"], "TMIN": overlap["TMIN_near"]}},
            {"field": "冠层上方空气温度", "physical_meaning": "above vegetation canopy", "height_m": None,
             "numeric_evidence": {"TMAX": overlap["TMAX_above"], "TMIN": overlap["TMIN_above"]}},
        ],
        "available_chinaflux_heights_m": [1.6, 2.9], "new_daily_product_temperature_height_m": None,
        "decision": "Near-ground has slightly lower MAE in the 2005-2006 numeric comparison, but the comparison series has no documented installation height; the evidence cannot pair either ChinaFLUX column to 1.6 m or 2.9 m. No canonical TMAX/TMIN source is selected.",
        "metadata_evidence": [
            "ChinaFLUX 30-minute product dictionary describes near-ground air temperature below canopy and canopy-above air temperature above canopy.",
            "The same dictionary lists air-temperature installations at 1.6 m and 2.9 m without assigning either height to either exported field.",
            "The 2005-2022 daily product names HMP45D but does not state its temperature sensor height.",
            "The 1998-2006 official product is monthly, so it cannot establish daily TMAX/TMIN mapping.",
        ],
        "overlap_period": "2005-2006", "numeric_comparison_is_not_independent_validation": True,
    }
    json_dump(out / "temperature_sensor_mapping.json", mapping)

    # Source preference is conditional because observed-versus-imputed flags are absent at cell level.
    rain_overlap = overlap["RAIN"]
    source_priority = {
        "status": "BLOCKED_SOURCE_INCONSISTENCY",
        "overlap_period": "2005-2006",
        "same_site_multi_product_language": "These are same-site products and are not treated as independent observations.",
        "variables": {
            "RAIN": {
                "comparison": rain_overlap,
                "annual_totals": {str(y): overlap["RAIN"]["annual"][str(y)] for y in (2005, 2006)},
                "provisional_preference": "chinaflux_30min on complete 48-interval days; 2005-2022 daily product as a gap alternative only",
                "decision": "No final priority: units are both mm/day, but annual totals differ and metadata does not resolve sensor/processing differences; daily product has product-level QC/infill without per-day flags.",
                "periodic_metrics_file": "source_overlap_metrics.csv",
            },
            "TMAX_TMIN": {
                "comparison": {**near_metrics, **above_metrics},
                "provisional_preference": "not selected",
                "decision": "Temperature field-to-height mapping is unresolved; numerical closeness cannot identify sensor height because the comparison product height is undocumented.",
            },
            "SRAD": {
                "comparison": overlap["SRAD"],
                "provisional_preference": "ChinaFLUX 30-minute incoming total solar radiation for complete days in 2004-2010; 2005-2022 total-radiation daily product is the only in-scope daily source for 2011-2013",
                "decision": "Definition is compatible with incoming global solar radiation and unit conversion is explicit, but the 2011-2013 table contains missing and negative daily values; no complete 2004-2013 source priority is approved.",
            },
        },
        "unresolved_items": ["rain gauge/processing differences", "per-day observed-versus-imputed flags", "temperature sensor height pairing"],
    }
    json_dump(out / "source_priority_decision.json", source_priority)

    provenance = build_provenance(cf, met, rad, gap_rows)
    csv_dump(out / "yc_weather_provenance_daily_2004_2013.csv", list(provenance[0]), provenance)
    overlap_fields = ["variable_comparison", "aggregation", "period", "year", "month", "paired_days", "cf_total", "new_total", "cf_mean", "new_mean", "difference_new_minus_cf"]
    csv_dump(out / "source_overlap_metrics.csv", overlap_fields, overlap_rows)
    csv_dump(out / "source_inventory.csv", list(inventory[0]), inventory)
    source_metadata["input_row_count_by_dataset"] = {
        "yucheng_met_through_2013": len(met), "yucheng_radiation_through_2013": len(rad),
        "chinaflux_30min_daily_groups_2004_2010": len(cf),
    }
    legacy_manifest = json.loads(PREVIOUS_LEGACY_INVENTORY.read_text(encoding="utf-8"))
    source_metadata["legacy_source_reference"] = {
        "prior_gap_file": PREVIOUS_GAPS.relative_to(ROOT).as_posix(), "prior_gap_file_sha256": sha256_file(PREVIOUS_GAPS),
        "prior_legacy_inventory": PREVIOUS_LEGACY_INVENTORY.relative_to(ROOT).as_posix(),
        "prior_legacy_inventory_sha256": sha256_file(PREVIOUS_LEGACY_INVENTORY),
        "prior_legacy_observations": PREVIOUS_LEGACY_OBSERVATIONS.relative_to(ROOT).as_posix(),
        "prior_legacy_observations_sha256": sha256_file(PREVIOUS_LEGACY_OBSERVATIONS),
        "inherited_original_source_hashes": legacy_manifest.get("source_files", {}),
        "reextracted_this_run": False,
    }
    source_metadata["training_weather_years"] = training_weather_years
    source_metadata["validation_cutoff_assertion"] = "PASS: all weather rows admitted to analysis have year <= 2013."
    json_dump(out / "source_metadata.json", source_metadata)
    json_dump(out / "weather_2011_2013_qc.json", weather_qc)

    summary = make_report_summary(cf, met, rad, gap_summary, overlap, weather_qc)
    summary["source_inventory_rows"] = len(inventory)
    summary["overlap_metrics"] = overlap
    summary["precipitation_gap_rows"] = gap_rows
    json_dump(out / "audit_summary.json", summary)

    candidate_qc = [
        {"check": "Gate A temperature mapping", "status": "NOT_RUN_CANDIDATE_NOT_BUILT", "reason": mapping["decision"]},
        {"check": "Gate B precipitation completeness", "status": "NOT_RUN_CANDIDATE_NOT_BUILT", "reason": gap_summary["reason"]},
        {"check": "Gate C 2011-2013 completeness", "status": "NOT_RUN_CANDIDATE_NOT_BUILT", "reason": "Daily TMAX/TMIN/SRAD contain unresolved missing dates and SRAD includes negative values."},
        {"check": "Final physical/date/provenance QC", "status": "NOT_RUN_CANDIDATE_NOT_BUILT", "reason": "A canonical candidate cannot be assembled until all gates pass."},
    ]
    csv_dump(out / "final_candidate_qc.csv", list(candidate_qc[0]), candidate_qc)
    csv_dump(out / "final_candidate_monthly_summary.csv", ["status", "month", "precipitation_total_mm", "rainy_days", "mean_TMAX", "mean_TMIN", "mean_SRAD"],
             [{"status": "NOT_RUN_CANDIDATE_NOT_BUILT", "month": "", "precipitation_total_mm": "", "rainy_days": "", "mean_TMAX": "", "mean_TMIN": "", "mean_SRAD": ""}])
    csv_dump(out / "final_candidate_annual_summary.csv", ["status", "year", "precipitation_total_mm", "rainy_days", "max_daily_rainfall_mm", "max_consecutive_dry_days", "mean_TMAX", "mean_TMIN", "mean_SRAD"],
             [{"status": "NOT_RUN_CANDIDATE_NOT_BUILT", "year": "", "precipitation_total_mm": "", "rainy_days": "", "max_daily_rainfall_mm": "", "max_consecutive_dry_days": "", "mean_TMAX": "", "mean_TMIN": "", "mean_SRAD": ""}])

    all_gates_pass = all(summary["gates"][key].startswith("PASS") for key in ("A_temperature", "B_precipitation", "C_2011_2013"))
    if args.build_candidate:
        if not all_gates_pass:
            raise SystemExit(f"Candidate blocked: {summary['final_status']}; see {out / 'audit_summary.json'}")
        candidate_rows = []
        for row in provenance:
            candidate_rows.append({
                "DATE": row["date"], "YEAR": row["year"], "DOY": row["doy"], "SRAD": row["SRAD"], "TMAX": row["TMAX"],
                "TMIN": row["TMIN"], "RAIN": row["RAIN"], "SRAD_source": row["SRAD_source"], "TMAX_source": row["TMAX_source"],
                "TMIN_source": row["TMIN_source"], "RAIN_source": row["RAIN_source"], "qc_status": row["overall_qc"],
            })
        csv_dump(out / "yc_wgen_fitting_weather_2004_2013.csv", list(candidate_rows[0]), candidate_rows)
        print(f"Candidate CSV created under {out}; no CLI/WGEN/DSSAT/PPO was run.")
    else:
        print(f"Audit complete: {summary['final_status']}; candidate was not created.")
    print(f"Results: {out}")


if __name__ == "__main__":
    main()
