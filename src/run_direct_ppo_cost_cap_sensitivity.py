from __future__ import annotations

import argparse
import shutil
import traceback
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import yaml

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None
    RGBColor = None
    PP_ALIGN = None
    Inches = None
    Pt = None

import run_all_year_direct_action_safe_ppo as direct
from ppo_safe_rendering import PROJECT_ROOT


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_direct_ppo_cost_cap_sensitivity.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "all_year_direct_ppo_cost_cap_sensitivity"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-11_direct_ppo_cost_and_cap_sensitivity_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-11_direct_ppo_cost_and_cap_sensitivity_report.pptx"

LITERATURE = [
    {
        "name": "Gautron et al. 2022 / gym-DSSAT",
        "url": "https://arxiv.org/abs/2207.03270",
        "method_support": "DSSAT can be wrapped as a Gym-style RL environment for crop management experiments.",
    },
    {
        "name": "Wu et al. 2022",
        "url": "https://arxiv.org/abs/2204.10394",
        "method_support": "Nitrogen management can be formulated as an RL problem using DSSAT crop simulations and input-cost tradeoffs.",
    },
    {
        "name": "Tao et al. 2022/2023",
        "url": "https://www.ijcai.org/proceedings/2023/691",
        "method_support": "Water and nitrogen management can be jointly optimized in DSSAT-based crop simulation settings.",
    },
    {
        "name": "Kallenberg et al. 2023",
        "url": "https://www.cambridge.org/core/journals/environmental-data-science/article/nitrogen-management-with-reinforcement-learning-and-crop-growth-models/358749FAFAA4990B1448DAB7F48D641C",
        "method_support": "CropGym demonstrates RL for crop-growth-model nitrogen management with yield and environmental tradeoffs.",
    },
    {
        "name": "Saikai et al. 2023",
        "url": "https://journals.plos.org/water/article?id=10.1371/journal.pwat.0000169",
        "method_support": "Deep RL can be used for irrigation scheduling with crop-model simulations and profit-oriented evaluation.",
    },
]


def load_yaml(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def write_yaml(data: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        yaml.safe_dump(data, f, sort_keys=False, allow_unicode=True)


def df_to_markdown(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df.empty:
        return ""
    work = df.copy()
    if max_rows is not None:
        work = work.head(max_rows)
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(4)
    for col in work.columns:
        work[col] = work[col].map(lambda x: "" if pd.isna(x) else str(x))
    header = "| " + " | ".join(map(str, work.columns)) + " |"
    sep = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, sep, *rows])


def ensure_dirs() -> None:
    for sub in ["configs", "models", "evaluation", "figures", "reports", "scenarios"]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)


def scenario_config(base: dict, scenario: dict, output_root_rel: str) -> dict:
    cfg = dict(base)
    cfg["paths"] = dict(base["paths"])
    cfg["runtime"] = dict(base["runtime"])
    cfg["action_safety"] = dict(base["action_safety"])
    cfg["economics"] = dict(base["economics"])
    cfg["reward"] = dict(base["reward"])
    cfg["paths"]["output_root"] = output_root_rel
    cfg["action_safety"]["season_irrigation_soft_limit"] = float(scenario["season_irrigation_cap"])
    cfg["action_safety"]["season_n_soft_limit"] = float(scenario["season_n_cap"])
    cfg["economics"]["water_cost"] = float(scenario["water_cost"])
    cfg["economics"]["nitrogen_cost"] = float(scenario["nitrogen_cost"])
    cfg["reward"]["water_cost"] = float(scenario["water_cost"])
    cfg["reward"]["nitrogen_cost"] = float(scenario["nitrogen_cost"])
    return cfg


def add_params(df: pd.DataFrame, scenario: dict) -> pd.DataFrame:
    out = df.copy()
    out["scenario"] = scenario["scenario"]
    out["water_cost"] = float(scenario["water_cost"])
    out["nitrogen_cost"] = float(scenario["nitrogen_cost"])
    out["season_irrigation_cap"] = float(scenario["season_irrigation_cap"])
    out["season_n_cap"] = float(scenario["season_n_cap"])
    return out


def summarize_for_sensitivity(eval_summary: pd.DataFrame, diagnosis: pd.DataFrame, scenario: dict) -> pd.DataFrame:
    diag_cols = [
        "station_code",
        "year",
        "irrigation_during_or_before_swfac_stress",
        "fertilization_during_or_before_nstres",
        "decision_reasonableness_label",
    ]
    diag = diagnosis[[c for c in diag_cols if c in diagnosis.columns]].drop_duplicates(["station_code", "year"])
    merged = eval_summary.merge(diag, on=["station_code", "year"], how="left", suffixes=("", "_diag"))
    if "decision_reasonableness_label_diag" in merged.columns:
        merged["decision_reasonableness_label"] = merged["decision_reasonableness_label_diag"].combine_first(
            merged.get("decision_reasonableness_label")
        )
    merged = add_params(merged, scenario)
    keep = [
        "scenario",
        "station_code",
        "year",
        "split",
        "water_cost",
        "nitrogen_cost",
        "season_irrigation_cap",
        "season_n_cap",
        "total_irrigation",
        "total_n",
        "final_topwt",
        "final_grnwt",
        "mean_swfac",
        "max_swfac",
        "swfac_stress_days_gt_0p05",
        "mean_nstres",
        "max_nstres",
        "nstres_days_gt_0p05",
        "profit_simple",
        "irrigation_event_count",
        "n_event_count",
        "irrigation_during_or_before_swfac_stress",
        "fertilization_during_or_before_nstres",
        "decision_reasonableness_label",
        "run_status",
        "daily_csv_path",
    ]
    return merged[[c for c in keep if c in merged.columns]]


def set_direct_output_root(path: Path) -> None:
    direct.OUTPUT_ROOT = path


def run_single_scenario(master_config: dict, base_direct_config: dict, selection: pd.DataFrame, scenario: dict, train: bool = True) -> pd.DataFrame:
    scenario_name = str(scenario["scenario"])
    scenario_root = OUTPUT_ROOT / "scenarios" / scenario_name
    scenario_rel = str(scenario_root.relative_to(PROJECT_ROOT)).replace("\\", "/")
    scenario_root.mkdir(parents=True, exist_ok=True)
    if bool(scenario.get("reuse_00617", False)):
        eval_summary = pd.read_csv(PROJECT_ROOT / master_config["base_eval_summary_csv"])
        diagnosis = pd.read_csv(PROJECT_ROOT / master_config["base_diagnosis_csv"])
        summary = summarize_for_sensitivity(eval_summary, diagnosis, scenario)
        scenario_eval = scenario_root / "evaluation"
        scenario_eval.mkdir(parents=True, exist_ok=True)
        summary.to_csv(scenario_eval / "ppo_cost_cap_sensitivity_summary.csv", index=False, encoding="utf-8-sig")
        return summary

    cfg = scenario_config(base_direct_config, scenario, scenario_rel)
    write_yaml(cfg, OUTPUT_ROOT / "configs" / f"config_{scenario_name}.yaml")
    set_direct_output_root(scenario_root)
    direct.ensure_dirs()
    env_config = direct.build_env_config(cfg, selection[selection["selected_for_train"] | selection["selected_for_eval"]])
    direct.write_yaml(env_config, scenario_root / "configs" / "resolved_env_config.yaml")
    if train:
        train_summary = direct.train_station_models(cfg, env_config, selection, debug=False)
    else:
        train_path = scenario_root / "evaluation" / "training_run_summary.csv"
        train_summary = pd.read_csv(train_path) if train_path.exists() else pd.DataFrame()
    eval_summary = direct.evaluate_models(cfg, env_config, selection, train_summary)
    diagnosis = direct.build_reasonableness(cfg, eval_summary)
    summary = summarize_for_sensitivity(eval_summary, diagnosis, scenario)
    summary.to_csv(scenario_root / "evaluation" / "ppo_cost_cap_sensitivity_summary.csv", index=False, encoding="utf-8-sig")
    return summary


def add_recommendation_flags(master_config: dict, summary: pd.DataFrame) -> pd.DataFrame:
    derived_cols = [
        "baseline_final_grnwt",
        "irrigation_cap_fraction",
        "n_cap_fraction",
        "yield_fraction_vs_baseline",
        "case_recommended",
    ]
    summary = summary.drop(columns=[col for col in derived_cols if col in summary.columns])
    baseline = summary[summary["scenario"].eq("A_baseline")][["station_code", "year", "final_grnwt"]].rename(
        columns={"final_grnwt": "baseline_final_grnwt"}
    )
    out = summary.merge(baseline, on=["station_code", "year"], how="left")
    rec_cfg = master_config["recommendation"]
    out["irrigation_cap_fraction"] = out["total_irrigation"] / out["season_irrigation_cap"]
    out["n_cap_fraction"] = out["total_n"] / out["season_n_cap"]
    out["yield_fraction_vs_baseline"] = out["final_grnwt"] / out["baseline_final_grnwt"].replace(0, np.nan)
    out["case_recommended"] = (
        out["run_status"].eq("ok")
        & ~out["decision_reasonableness_label"].eq("cap_saturated")
        & (out["irrigation_cap_fraction"] < float(rec_cfg["max_cap_fraction"]))
        & (out["n_cap_fraction"] < float(rec_cfg["max_cap_fraction"]))
        & (out["yield_fraction_vs_baseline"] >= float(rec_cfg["min_yield_fraction_vs_baseline"]))
        & (out["final_grnwt"] >= float(rec_cfg["min_positive_grnwt"]))
        & (out["final_topwt"] >= float(rec_cfg["min_positive_topwt"]))
    )
    return out


def scenario_rollup(summary: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for scenario, group in summary.groupby("scenario"):
        ok = group[group["run_status"].eq("ok")]
        rows.append(
            {
                "scenario": scenario,
                "water_cost": float(group["water_cost"].iloc[0]),
                "nitrogen_cost": float(group["nitrogen_cost"].iloc[0]),
                "season_irrigation_cap": float(group["season_irrigation_cap"].iloc[0]),
                "season_n_cap": float(group["season_n_cap"].iloc[0]),
                "n_cases": int(len(group)),
                "ok_cases": int(len(ok)),
                "cap_saturated_cases": int(group["decision_reasonableness_label"].eq("cap_saturated").sum()),
                "cap_saturation_rate": float(group["decision_reasonableness_label"].eq("cap_saturated").mean()),
                "recommended_cases": int(group["case_recommended"].sum()),
                "recommended_rate": float(group["case_recommended"].mean()),
                "mean_irrigation": float(ok["total_irrigation"].mean()) if len(ok) else np.nan,
                "mean_n": float(ok["total_n"].mean()) if len(ok) else np.nan,
                "mean_final_grnwt": float(ok["final_grnwt"].mean()) if len(ok) else np.nan,
                "mean_final_topwt": float(ok["final_topwt"].mean()) if len(ok) else np.nan,
                "mean_profit_simple": float(ok["profit_simple"].mean()) if len(ok) else np.nan,
                "min_yield_fraction_vs_baseline": float(ok["yield_fraction_vs_baseline"].min()) if len(ok) else np.nan,
                "mean_yield_fraction_vs_baseline": float(ok["yield_fraction_vs_baseline"].mean()) if len(ok) else np.nan,
            }
        )
    return pd.DataFrame(rows).sort_values(
        ["recommended_rate", "cap_saturation_rate", "mean_profit_simple", "mean_final_grnwt"],
        ascending=[False, True, False, False],
    )


def recommended_configuration(rollup: pd.DataFrame) -> pd.DataFrame:
    viable = rollup[(rollup["recommended_cases"] > 0) & (rollup["cap_saturation_rate"] < 1.0)].copy()
    if viable.empty:
        fallback = rollup[
            (~rollup["scenario"].eq("A_baseline"))
            & (rollup["mean_yield_fraction_vs_baseline"] >= 0.9)
            & (rollup["ok_cases"] == rollup["n_cases"])
        ].copy()
        if fallback.empty:
            best = rollup.sort_values(
                ["cap_saturation_rate", "mean_profit_simple", "mean_yield_fraction_vs_baseline"],
                ascending=[True, False, False],
            ).head(1)
        else:
            best = fallback.sort_values(
                ["mean_profit_simple", "mean_yield_fraction_vs_baseline", "season_irrigation_cap", "season_n_cap"],
                ascending=[False, False, True, True],
            ).head(1)
        status = "fallback_for_mentor_discussion_not_final"
        reason = "No scenario met the strict non-saturation criterion; this fallback keeps yield near baseline with lower input caps, but still saturates its cap."
    else:
        best = viable.head(1)
        status = "recommended_for_mentor_display"
        reason = "This scenario had the best balance of non-saturation, yield retention, and simple profit among tested cost/cap settings."
    out = best.copy()
    out["recommendation_status"] = status
    out["recommendation_reason"] = reason
    out["advantages"] = np.where(
        out["cap_saturation_rate"] < 1.0,
        "Reduced cap saturation relative to A_baseline; keeps the same direct PPO method and train/eval split.",
        "Keeps TOPWT/GRNWT workflow comparable to 006_17, but still saturates caps.",
    )
    out["limitations"] = np.where(
        out["recommended_cases"] > 0,
        "Still single-seed and only tests four cost/cap combinations.",
        "No tested parameter set fully solved cap saturation; not suitable as final thesis result yet.",
    )
    out["suitable_as_main_thesis_result"] = out["recommendation_status"].eq("recommended_for_mentor_display")
    return out


def plot_comparisons(summary: pd.DataFrame, rollup: pd.DataFrame) -> None:
    fig_dir = OUTPUT_ROOT / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    metrics = [
        ("total_irrigation", "total_irrigation_comparison.png", "Total irrigation (mm)"),
        ("total_n", "total_n_comparison.png", "Total N (kg/ha)"),
        ("final_grnwt", "final_grnwt_comparison.png", "Final GRNWT"),
        ("profit_simple", "profit_simple_comparison.png", "Simple profit"),
        ("swfac_stress_days_gt_0p05", "swfac_stress_days_comparison.png", "SWFAC stress days"),
        ("nstres_days_gt_0p05", "nstres_stress_days_comparison.png", "NSTRES stress days"),
    ]
    for col, filename, ylabel in metrics:
        if col not in summary.columns:
            continue
        data = summary.groupby("scenario")[col].mean().reindex(rollup["scenario"])
        ax = data.plot(kind="bar", figsize=(8.5, 4.2), color="#4F81BD")
        ax.set_ylabel(ylabel)
        ax.grid(axis="y", alpha=0.25)
        ax.get_figure().tight_layout()
        ax.get_figure().savefig(fig_dir / filename, dpi=180)
        plt.close(ax.get_figure())
    ax = rollup.set_index("scenario")["cap_saturation_rate"].plot(kind="bar", figsize=(8.5, 4.2), color="#C0504D")
    ax.set_ylabel("Cap saturation rate")
    ax.set_ylim(0, 1.05)
    ax.grid(axis="y", alpha=0.25)
    ax.get_figure().tight_layout()
    ax.get_figure().savefig(fig_dir / "cap_saturation_rate_comparison.png", dpi=180)
    plt.close(ax.get_figure())


def write_report(summary: pd.DataFrame, rollup: pd.DataFrame, rec: pd.DataFrame) -> None:
    references = "\n".join(
        f"- [{item['name']}]({item['url']}): {item['method_support']}" for item in LITERATURE
    )
    lines = [
        "# Direct PPO Cost and Cap Sensitivity Report",
        "",
        "## Scope",
        "",
        "This stage follows 006_17 direct action-safe PPO. It does not use RF, behavior cloning, imitation learning, offline schedule search, expert replay, constrained PPO fine-tuning, rainfall scaling, reward-structure redesign, episode-level profit reward, phenology-window action design, event action design, or multi-objective PPO.",
        "",
        "Only four quantities were changed across scenarios: `water_cost`, `nitrogen_cost`, `season_irrigation_cap`, and `season_n_cap`.",
        "",
        "## Tested Scenarios",
        "",
        df_to_markdown(rollup[["scenario", "water_cost", "nitrogen_cost", "season_irrigation_cap", "season_n_cap", "n_cases", "ok_cases"]]),
        "",
        "## Scenario Rollup",
        "",
        df_to_markdown(rollup),
        "",
        "## Recommended Configuration",
        "",
        df_to_markdown(rec),
        "",
        "## Sensitivity Summary Preview",
        "",
        df_to_markdown(
            summary[
                [
                    "scenario",
                    "station_code",
                    "year",
                    "total_irrigation",
                    "total_n",
                    "final_grnwt",
                    "profit_simple",
                    "decision_reasonableness_label",
                    "case_recommended",
                    "yield_fraction_vs_baseline",
                ]
            ],
            80,
        ),
        "",
        "## Method Sources and Literature Support",
        "",
        references,
        "",
        "The literature supports the general formulation of crop management as simulator-based reinforcement learning and the use of input-cost tradeoffs. The specific finding here, namely whether the four tested cost/cap settings avoid cap saturation in the five Chinese station-year pool, is this project's experimental result and should not be attributed to the cited papers.",
    ]
    DOC_MD.write_text("\n".join(lines), encoding="utf-8")
    reports = OUTPUT_ROOT / "reports"
    reports.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(DOC_MD, reports / DOC_MD.name)
    write_ppt(summary, rollup, rec)
    if DOC_PPT.exists():
        shutil.copyfile(DOC_PPT, reports / DOC_PPT.name)


def add_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.5))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = title
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_bullets(prs, title: str, bullets: list[str]) -> None:
    if Presentation is None:
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(12), Inches(5.8))
    tf = box.text_frame
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0, 0, 0)


def add_table(prs, title: str, df: pd.DataFrame, max_rows: int = 12) -> None:
    if Presentation is None:
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title)
    work = df.head(max_rows).copy()
    if work.empty:
        return
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(3)
    rows, cols = len(work), len(work.columns)
    table_shape = slide.shapes.add_table(rows + 1, cols, Inches(0.25), Inches(1.05), Inches(12.85), Inches(6.0))
    table = table_shape.table
    for j, col in enumerate(work.columns):
        cell = table.cell(0, j)
        cell.text = str(col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(68, 114, 196)
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(7)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(work.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = "" if pd.isna(value) else str(value)
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(226, 235, 247)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(6)
                    run.font.color.rgb = RGBColor(0, 0, 0)


def add_picture(prs, title: str, path: Path) -> None:
    if Presentation is None:
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title)
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(0.7), Inches(1.05), width=Inches(11.9))


def write_ppt(summary: pd.DataFrame, rollup: pd.DataFrame, rec: pd.DataFrame) -> None:
    if Presentation is None:
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_bullets(
        prs,
        "007_01 Direct PPO Cost and Cap Sensitivity",
        [
            "Goal: test whether changing only water_cost, nitrogen_cost, season irrigation cap, and season N cap can avoid cap saturation.",
            "No RF, imitation learning, offline schedule search, expert replay, rainfall scaling, reward redesign, or action-design changes were used.",
            "Train/eval years, seed, PPO hyperparameters, action scaling, and direct PPO structure are inherited from 006_17.",
        ],
    )
    add_table(prs, "Scenario Rollup", rollup, 8)
    add_table(prs, "Recommended Configuration", rec, 4)
    for fig in [
        "total_irrigation_comparison.png",
        "total_n_comparison.png",
        "final_grnwt_comparison.png",
        "profit_simple_comparison.png",
        "cap_saturation_rate_comparison.png",
    ]:
        add_picture(prs, fig.replace("_", " ").replace(".png", ""), OUTPUT_ROOT / "figures" / fig)
    add_bullets(
        prs,
        "Method Sources and Project Results",
        [
            "Literature supports crop-management RL with process-based simulators and input-cost tradeoffs.",
            "This project's tested result is the sensitivity of the five-station all-year direct PPO baseline to four cost/cap settings.",
            "A scenario should be shown to the mentor only if it avoids cap saturation while maintaining reasonable TOPWT/GRNWT.",
        ],
    )
    prs.save(DOC_PPT)


def run_pipeline(config_path: Path, scenarios: set[str] | None = None, report_only: bool = False, skip_train: bool = False) -> None:
    master = load_yaml(config_path)
    ensure_dirs()
    shutil.copyfile(config_path, OUTPUT_ROOT / "configs" / config_path.name)
    base_direct = load_yaml(PROJECT_ROOT / master["base_config"])
    selection = pd.read_csv(PROJECT_ROOT / master["base_selection_csv"])
    all_summaries = []
    if report_only:
        scenario_summaries = []
        for scenario in master["scenarios"]:
            scenario_path = (
                OUTPUT_ROOT
                / "scenarios"
                / scenario["scenario"]
                / "evaluation"
                / "ppo_cost_cap_sensitivity_summary.csv"
            )
            if scenario_path.exists():
                scenario_summaries.append(pd.read_csv(scenario_path))
        if scenario_summaries:
            summary = pd.concat(scenario_summaries, ignore_index=True)
        else:
            existing = OUTPUT_ROOT / "evaluation" / "ppo_cost_cap_sensitivity_summary.csv"
            summary = pd.read_csv(existing)
    else:
        for scenario in master["scenarios"]:
            if scenarios and scenario["scenario"] not in scenarios:
                continue
            print(f"[007_01] scenario={scenario['scenario']}", flush=True)
            try:
                all_summaries.append(run_single_scenario(master, base_direct, selection, scenario, train=not skip_train))
            except Exception:
                failure = pd.DataFrame(
                    [
                        {
                            "scenario": scenario["scenario"],
                            "run_status": "failed",
                            "notes": traceback.format_exc(),
                        }
                    ]
                )
                all_summaries.append(add_params(failure, scenario))
        previous_path = OUTPUT_ROOT / "evaluation" / "ppo_cost_cap_sensitivity_summary.csv"
        if previous_path.exists() and scenarios:
            previous = pd.read_csv(previous_path)
            previous = previous[~previous["scenario"].isin(scenarios)]
            all_summaries.insert(0, previous)
        summary = pd.concat(all_summaries, ignore_index=True) if all_summaries else pd.DataFrame()
    summary = add_recommendation_flags(master, summary)
    summary.to_csv(OUTPUT_ROOT / "evaluation" / "ppo_cost_cap_sensitivity_summary.csv", index=False, encoding="utf-8-sig")
    rollup = scenario_rollup(summary)
    rollup.to_csv(OUTPUT_ROOT / "evaluation" / "ppo_cost_cap_sensitivity_rollup.csv", index=False, encoding="utf-8-sig")
    rec = recommended_configuration(rollup)
    rec.to_csv(OUTPUT_ROOT / "evaluation" / "recommended_configuration.csv", index=False, encoding="utf-8-sig")
    plot_comparisons(summary, rollup)
    write_report(summary, rollup, rec)
    print((OUTPUT_ROOT / "evaluation" / "ppo_cost_cap_sensitivity_summary.csv").relative_to(PROJECT_ROOT))
    print((OUTPUT_ROOT / "evaluation" / "ppo_cost_cap_sensitivity_rollup.csv").relative_to(PROJECT_ROOT))
    print((OUTPUT_ROOT / "evaluation" / "recommended_configuration.csv").relative_to(PROJECT_ROOT))
    print(DOC_MD.relative_to(PROJECT_ROOT))
    print(DOC_PPT.relative_to(PROJECT_ROOT))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default=str(CONFIG_PATH))
    parser.add_argument("--scenario", action="append", help="Run only the named scenario. Can be repeated.")
    parser.add_argument("--report-only", action="store_true")
    parser.add_argument("--skip-train", action="store_true")
    args = parser.parse_args()
    run_pipeline(Path(args.config), scenarios=set(args.scenario or []), report_only=args.report_only, skip_train=args.skip_train)


if __name__ == "__main__":
    main()
