"""Build the concise Chinese stage-record deck for the HL/FQ source audit."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "docs" / "hl_fq_random_weather_8seed_results.pptx"
AUDIT = ROOT / "results" / "hl_fq_random_weather_8seed" / "source_audit"
NAVY = RGBColor(24, 42, 67)
TEAL = RGBColor(0, 133, 145)
RED = RGBColor(190, 57, 61)
AMBER = RGBColor(204, 129, 34)
INK = RGBColor(47, 59, 72)
MUTED = RGBColor(105, 119, 134)
PALE = RGBColor(242, 246, 249)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, x, y, w, h, text, *, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.02)
    p = frame.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_card(slide, x, y, w, h, title, body, *, accent=TEAL, body_size=17):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(220, 228, 234)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    band.fill.solid()
    band.fill.fore_color.rgb = accent
    band.line.fill.background()
    add_text(slide, x + 0.24, y + 0.18, w - 0.42, 0.38, title, size=17, color=NAVY, bold=True)
    add_text(slide, x + 0.24, y + 0.65, w - 0.42, h - 0.8, body, size=body_size, color=INK)


def add_header(slide, title, section, page):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    add_text(slide, 0.55, 0.28, 9.5, 0.5, title, size=25, color=NAVY, bold=True)
    add_text(slide, 10.0, 0.34, 2.7, 0.35, section, size=11, color=TEAL, bold=True, align=PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.91), Inches(12.2), Inches(0.025))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(205, 218, 226)
    line.line.fill.background()
    add_text(slide, 0.6, 7.15, 11.9, 0.18, "HL / FQ WGEN 8-seed 阶段审计  |  2026-10-04", size=9, color=MUTED)
    add_text(slide, 12.25, 7.12, 0.45, 0.2, str(page), size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def add_table(slide, x, y, widths, rows, row_h=0.52, font_size=14):
    height = row_h * len(rows)
    table = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(height)).table
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = str(value)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.035)
            cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if ri == 0 else (WHITE if ri % 2 else RGBColor(232, 239, 243))
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                paragraph.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(font_size if ri else font_size - 1)
                    run.font.bold = ri == 0
                    run.font.color.rgb = WHITE if ri == 0 else INK
    return table


def main() -> None:
    audit = json.loads((AUDIT / "weather_source_audit_final.json").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Decision
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_text(slide, 0.8, 0.65, 7.2, 0.45, "阶段审计  ·  2026-10-04", size=15, color=RGBColor(130, 220, 224), bold=True)
    add_text(slide, 0.8, 1.45, 11.4, 1.0, "HL / FQ 随机天气增强", size=34, color=WHITE, bold=True)
    add_text(slide, 0.82, 2.35, 11.5, 0.65, "8-seed PPO 训练前的数据源门禁", size=24, color=RGBColor(209, 221, 231))
    add_text(slide, 0.82, 3.35, 11.0, 0.62, audit["final_decision"], size=26, color=RGBColor(255, 160, 150), bold=True)
    add_text(slide, 0.85, 4.15, 10.8, 0.85,
             "2013 年降水仍有未核实日期，拟合天气不完整；按任务停止条件，未生成 WGEN 或启动 PPO。",
             size=21, color=WHITE)
    add_card(slide, 0.85, 5.45, 5.4, 1.15, "HLA  ·  2004–2013", "111 个未覆盖降水日（2013）", accent=RED, body_size=19)
    add_card(slide, 6.55, 5.45, 5.4, 1.15, "FQA  ·  2005–2013", "157 个未覆盖降水日（2013）", accent=RED, body_size=19)
    add_text(slide, 0.85, 7.12, 11.6, 0.18, "阶段说明：来源审计完成；WGEN / runtime / PPO 未运行。", size=9, color=RGBColor(177, 193, 204))

    # 2. Existing source quality
    slide = prs.slides.add_slide(blank)
    add_header(slide, "现有文件不能直接作为 WGEN 拟合输入", "01  数据源", 2)
    add_table(slide, 0.72, 1.25, [1.05, 1.1, 1.3, 1.55, 1.55, 1.75], [
        ["站点", "拟合日数", "RAIN 空白", "温度异常/缺失", "SRAD 缺失/异常", "现有 WTH 风险"],
        ["HLA", "3,653", "2,691", "TMAX 154.3°C；各变量缺 31 日", "缺 29 日", "2010 极端值仍在 WTH"],
        ["FQA", "3,287", "2,892", "TMAX/TMIN 各缺 83 日", "缺 93 日；306.055 异常", "旧清洗 provenance 不完整"],
    ], row_h=0.72, font_size=13)
    add_card(slide, 0.75, 3.55, 5.75, 1.35, "NASA POWER：TMAX / TMIN / SRAD", "6 个变量-站点筛选均通过训练期 MAE 与相关性门槛；候选值只用于缺失或物理异常日期。", accent=TEAL, body_size=15)
    add_card(slide, 6.75, 3.55, 5.75, 1.35, "NASA POWER：RAIN", "湿 / 干一致率 HLA 48.8%、FQA 68.1%；低于 80%，不用于雨量补缺。", accent=AMBER, body_size=15)
    add_text(slide, 0.82, 5.35, 11.5, 0.55, "所有查询和校正均限制在各自训练拟合期；没有读取 2014–2023。", size=18, color=NAVY, bold=True)
    add_text(slide, 0.82, 6.1, 11.5, 0.55, "原始 Excel、历史 .WTH、reward、action mask 与训练配置均未改写。", size=16, color=MUTED)

    # 3. Rain gap
    slide = prs.slides.add_slide(blank)
    add_header(slide, "官方邻站能补大部分雨量空白，但 2013 仍未覆盖", "02  降水门禁", 3)
    add_table(slide, 0.78, 1.28, [1.0, 2.0, 1.35, 1.3, 1.55, 1.4], [
        ["站点", "最近通过筛选站", "距站點", "湿/干一致", "多站候选补值", "仍缺"],
        ["HLA", "HAILUN / CHM00050756", "5.38 km", "89.3%", "2,580 天", "111 天"],
        ["FQA", "ZHENGZHOU / CHM00057083", "87.91 km", "86.9%", "2,735 天", "157 天"],
    ], row_h=0.7, font_size=13)
    add_card(slide, 0.82, 3.55, 5.6, 1.35, "本地记录日敏感性核对", "只在本地有记录的日期上，HLA 85.1%（935 对日），FQA 83.2%（388 对日）。", accent=TEAL, body_size=15)
    add_card(slide, 6.7, 3.55, 5.65, 1.35, "剩余缺口与 NASA 降水", "所有未覆盖日均在 2013；NASA 雨量未通过，不能代填。", accent=RED, body_size=15)
    add_text(slide, 0.88, 5.2, 11.5, 0.75,
             "CMA 官方日值产品可包含站点降水，但下载要求实名注册；本地没有 2013 CMA 原始记录。",
             size=17, color=INK)
    add_text(slide, 0.88, 6.2, 11.5, 0.4,
             "候选补值尚未晋升为 fitting weather；不把未知空白默认写成 0 mm。",
             size=16, color=RED, bold=True)

    # 4. YC route
    slide = prs.slides.add_slide(blank)
    add_header(slide, "YC WGEN-PPO 路线已核实，可作为方法参考", "03  路线复用", 4)
    add_card(slide, 0.82, 1.28, 5.7, 1.4, "运行时天气模式", "task-local FileX 副本设 METHODS.WTHER=W；每个 episode 将 weather seed 注入 PDI rseed1。", accent=TEAL, body_size=15)
    add_card(slide, 6.72, 1.28, 5.7, 1.4, "天气与 PPO 随机种子分离", "共同训练池 1001–1080；held-out 1081–1100；各 PPO seed 复用同一 weather pool。", accent=TEAL, body_size=15)
    add_table(slide, 1.1, 3.15, [2.7, 3.2, 3.3], [
        ["证据", "已验证内容", "边界"],
        ["004_03 smoke", "WTHER=W；4 个 daily weather hash 不同", "仅 YC"],
        ["004_03 formal", "3 个 PPO seeds；100,000 请求 / 100,080 实际步", "不是 8-seed"],
        ["004_02 气候 QC", "100 个随机 episodes、同种子可复现", "共同 106 天窗口；全年未验证"],
    ], row_h=0.62, font_size=14)
    add_text(slide, 1.15, 6.05, 10.7, 0.55, "本任务未生成 HL/FQ CLI 或 random weather；不会把 YC 结果当作 HL/FQ 证据。", size=16, color=NAVY, bold=True)

    # 5. Status and resume condition
    slide = prs.slides.add_slide(blank)
    add_header(slide, "当前完成度与恢复条件", "04  执行状态", 5)
    add_table(slide, 0.86, 1.25, [3.2, 3.6, 3.6], [
        ["阶段", "HLA", "FQA"],
        ["拟合天气 / CLI", "未生成", "未生成"],
        ["WGEN pool / runtime smoke", "0/80 + 0/20 / 未运行", "0/80 + 0/20 / 未运行"],
        ["PPO / 历史对照", "0/8 / 未生成", "0/8 / 未生成"],
    ], row_h=0.58, font_size=14)
    add_card(slide, 0.9, 4.05, 11.5, 1.25, "恢复所需输入", "项目目录内的 HLA/FQA 2013 官方逐日降水记录，或可追溯证据证明相关空白确实表示 0 mm。拿到后先重新跑来源与完整物理 QC。", accent=AMBER, body_size=16)
    add_text(slide, 0.95, 5.8, 11.3, 0.42, "reward / mask / IC / 训练年 / 验证年：未改。", size=17, color=TEAL, bold=True)
    add_text(slide, 0.95, 6.35, 11.3, 0.36, "GitHub push：未执行（项目指令要求 push 前获用户确认）。", size=14, color=MUTED)

    prs.save(OUT)
    # Reopen the package to catch a corrupt PPTX container before delivery.
    check = Presentation(OUT)
    if len(check.slides) != 5:
        raise RuntimeError(f"Expected 5 slides, got {len(check.slides)}")
    print(json.dumps({"path": str(OUT), "slides": len(check.slides), "bytes": OUT.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    main()
