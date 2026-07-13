from __future__ import annotations

from datetime import date
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
import yaml
from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "benchmark_results" / "021_01"
FIG = OUT / "figures"
DOC = ROOT / "docs" / "2026-07-12_021_01_final_dqn_strategy_determination.md"
PPT = ROOT / "docs" / "2026-07-12_021_01_final_dqn_strategy_determination.pptx"
CFG = ROOT / "configs" / "final_dqn_candidate.yaml"


def load_lc() -> pd.DataFrame:
    rows = []
    for seed in (0, 1, 2):
        run = OUT / f"021_01_lc2010_seed_stability_retry__lc_2010_seed{seed}"
        row = pd.read_csv(run / "evaluations" / "season_summary.csv").iloc[0].to_dict()
        rows.append(
            {
                "station_code": "LC",
                "year": 2010,
                "seed": seed,
                "selection": "maximum reward_total; earliest checkpoint on ties",
                "checkpoint_step": int(row["checkpoint"]),
                "yield_kg_ha": float(row["yield_kg_ha"]),
                "biomass_kg_ha": float(row["biomass_kg_ha"]),
                "irrigation_mm": float(row["irrigation_mm"]),
                "nitrogen_kg_ha": float(row["nitrogen_kg_ha"]),
                "total_reward": float(row["reward_total"]),
                "irrigation_events": int(row["number_of_irrigation_events"]),
                "nitrogen_events": int(row["number_of_nitrogen_events"]),
                "water_budget_use_ratio": float(row["water_budget_use_ratio"]),
                "nitrogen_budget_use_ratio": float(row["nitrogen_budget_use_ratio"]),
                "runtime_audit_passed": True,
                "status": "completed_50k",
                "source": str((run / "evaluations" / "season_summary.csv").relative_to(ROOT)),
            }
        )
    return pd.DataFrame(rows)


def plot_lc(df: pd.DataFrame) -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    fig, axes = plt.subplots(2, 2, figsize=(9, 6.5), constrained_layout=True)
    seeds = df["seed"].astype(str)
    panels = [
        ("yield_kg_ha", "Grain yield (kg/ha)", "#222222"),
        ("irrigation_mm", "Irrigation (mm)", "#1f77b4"),
        ("nitrogen_kg_ha", "Nitrogen (kg/ha)", "#2ca02c"),
        ("total_reward", "Selected reward", "#d62728"),
    ]
    for ax, (column, ylabel, color) in zip(axes.flat, panels):
        ax.bar(seeds, df[column], color=color, width=0.62)
        ax.set_xlabel("Seed")
        ax.set_ylabel(ylabel)
        ax.grid(axis="y", color="#dddddd", linewidth=0.7)
        for x, value in enumerate(df[column]):
            ax.text(x, value, f"{value:.1f}", ha="center", va="bottom", fontsize=9)
    fig.suptitle("LC2010 frozen DQN: formal 50K cross-seed result", fontsize=13)
    fig.savefig(FIG / "lc2010_50k_cross_seed_stability.png", dpi=300, bbox_inches="tight")
    fig.savefig(FIG / "lc2010_50k_cross_seed_stability.svg", bbox_inches="tight")
    plt.close(fig)


def update_tables(lc: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    lc.to_csv(OUT / "lc_seed_stability_summary.csv", index=False, encoding="utf-8-sig")
    station = pd.read_csv(OUT / "station_diagnosis.csv")
    mask = station["station_code"].eq("LC")
    station.loc[mask, "status"] = "resource_use_seed_sensitive"
    station.loc[mask, "evidence"] = (
        "After repairing the LC adapter source, seed0/1/2 formal 50K yields were "
        "8737/8728/8707 kg/ha while irrigation was 120/60/30 mm; nitrogen was 0 for all seeds."
    )
    station.loc[mask, "next_action"] = "Retain as diagnostic evidence; do not claim cross-seed resource stability."
    station.to_csv(OUT / "station_diagnosis.csv", index=False, encoding="utf-8-sig")

    evidence = pd.read_csv(OUT / "unified_parameter_evidence.csv")
    lc_mask = evidence["station_code"].eq("LC")
    evidence.loc[lc_mask, "status"] = "resource_use_seed_sensitive"
    evidence.loc[lc_mask, "evidence"] = (
        "Formal LC2010 50K seed0/1/2 had similar yields but irrigation differed by 90 mm."
    )
    unified = evidence["station_code"].eq("UNIFIED_CANDIDATE")
    evidence.loc[unified, "status"] = "partial_lc_sy"
    evidence.loc[unified, "evidence"] = (
        "Frozen configuration is supported by HLA/YC/FQ; LC remains resource-use seed-sensitive and SY input provenance is blocked."
    )
    evidence.to_csv(OUT / "unified_parameter_evidence.csv", index=False, encoding="utf-8-sig")
    return station, evidence


def update_candidate() -> dict:
    candidate = yaml.safe_load(CFG.read_text(encoding="utf-8"))
    candidate["meta"]["status"] = "partial_lc_sy"
    candidate["meta"]["last_updated"] = str(date.today())
    applicability = candidate["candidate"]["applicability"]
    applicability["confirmed_sites"] = ["HLA", "YC", "FQ"]
    applicability["blocked_sites"] = {
        "LC": "formal 50K yields are similar across seeds, but irrigation differs by 90 mm",
        "SY": "input provenance ambiguous (site config expects IC=2, prepared 2014 treatment row is IC=0)"
    }
    text = yaml.safe_dump(candidate, allow_unicode=True, sort_keys=False)
    CFG.write_text(text, encoding="utf-8")
    (OUT / "final_dqn_candidate.yaml").write_text(text, encoding="utf-8")
    return candidate


def md_table(df: pd.DataFrame) -> str:
    return df.to_markdown(index=False)


def write_doc(lc: pd.DataFrame, station: pd.DataFrame, evidence: pd.DataFrame, candidate: dict) -> None:
    yc = pd.read_csv(OUT / "yc_nitrogen_cost_summary.csv")
    fq = pd.read_csv(OUT / "fq_water_cost_summary.csv")
    sy = pd.read_csv(OUT / "sy_input_provenance.csv")
    text = f"""# 021_01 最终 DQN 策略判定

## 1. 任务状态

状态：`partial`。HLA、YC、FQ、LC 的必要诊断已经完成；LC 显示资源投入的 seed 敏感性，SY 仍因输入来源冲突而阻塞。本任务没有修改 reward、IC、DSSAT 输入值或动作空间，也没有扩大参数扫描。

## 2. 当前统一配置

- 算法：DQN；`n_steps=5`；50K steps；每 5K 保存 checkpoint。
- 动作：灌溉 `[0, 15, 30] mm`；施氮 `[0, 50, 100] kg/ha`。
- 季节预算：`I<=120 mm`，`N<=300 kg/ha`；决策间隔 7 DAP。
- 奖励：相对本地 null 的终端产量增益，减去 `1*I + 5*N`。
- checkpoint 规则：最大 `reward_total`；并列时选择更早 checkpoint。

## 3. 五站点判定

{md_table(station)}

## 4. YC nitrogen_cost 诊断

{md_table(yc)}

`nitrogen_cost=2/5/8` 没有形成足以支持修改统一系数的稳定分离，因此保留 5。

## 5. FQ water_cost 诊断

{md_table(fq)}

`water_cost=0.5/1/2` 基本落在同一策略族，没有证据支持修改统一系数，因此保留 1。

## 6. LC2010 正式 50K 跨 seed 复核

{md_table(lc)}

三 seed 的最佳 checkpoint 不同（15K、40K、50K）。产量分别为 8737、8728、8707 kg/ha，差异仅 30 kg/ha；灌溉分别为 120、60、30 mm，差异达 90 mm；施氮均为 0。DSSAT auto 为 8738 kg/ha、138.5 mm、0 kg/ha，official expert 为 8739 kg/ha。LC 因而属于“产量近乎持平但资源投入跨 seed 敏感”，不能写成跨 seed 稳定成功，也不是 DQN 显著增产案例。

第一次运行挂起的根因不是训练或端口：Benchmark adapter 错把已经 null 化的 LC 输入作为 DQN 源，形成 `FERTI=L` 但 `MF=0`，DSSAT/PDI 在 `FertType_mod.for` 触发 `fertfile(0)`。修复后先通过 5K smoke，再串行完成三组 50K；全部 runtime audit 通过。

## 7. SY 输入来源审计

{md_table(sy)}

SY 的 authoritative IC/MZX 尚未统一，因此不能启动正式训练，也不能算作统一配置已完成五站验证。

## 8. 统一参数证据

{md_table(evidence)}

## 9. 最终候选配置

```yaml
{yaml.safe_dump(candidate, allow_unicode=True, sort_keys=False)}```

## 10. 已解决与未解决问题

- 已解决：YC nitrogen_cost、FQ water_cost、LC adapter 和三 seed 50K 正式诊断。
- 未解决：LC 资源投入跨 seed 敏感；SY2014 的 IC/MZX 输入来源冲突。
- 下一步：先确认 SY authoritative input，再按同一冻结配置做 smoke 和正式 seed 复核；不新增敏感性扫描。

## 11. Methods Source

- `benchmark/environment_adapter.py`
- `benchmark/train_runner.py`
- `configs/experiments/021_01_lc2010_seed_stability_retry*.yaml`
- `src/finalize_021_01_after_lc_retry.py`

## 12. 输出与 Git

- LC 图：`benchmark_results/021_01/figures/lc2010_50k_cross_seed_stability.png/.svg`
- 汇总：`benchmark_results/021_01/lc_seed_stability_summary.csv`
- Git commit：待本次文件核验后提交。
- Git push：上次因 SSH 22 端口连接关闭而失败；本次提交后按用户既有授权再次尝试。
"""
    DOC.write_text(text, encoding="utf-8")


def add_slide(prs: Presentation, title: str, lines: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.6))
    frame = box.text_frame
    frame.clear()
    for index, line in enumerate(lines):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line
        p.font.size = Pt(20)


def write_ppt(lc: pd.DataFrame, station: pd.DataFrame) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_slide(prs, "021_01 最终 DQN 策略判定", ["HLA / YC / FQ / LC 已完成必要证据；SY 输入来源仍阻塞"])
    add_slide(prs, "统一配置", ["DQN，n_steps=5，50K", "I=[0,15,30]，N=[0,50,100]", "I<=120，N<=300", "Reward = local-null yield gain - 1I - 5N"])
    add_slide(prs, "站点状态", [f"{r.station_code}: {r.status}" for r in station.itertuples()])
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "LC2010 三 seed 正式结果"
    slide.shapes.add_picture(str(FIG / "lc2010_50k_cross_seed_stability.png"), Inches(1.5), Inches(1.25), width=Inches(10.3))
    add_slide(prs, "LC 结论", ["产量：8737 / 8728 / 8707 kg/ha", "灌溉：120 / 60 / 30 mm；施氮均为 0", "最佳 checkpoint：15K / 40K / 50K", "产量接近，但资源投入跨 seed 敏感"])
    add_slide(prs, "阻塞与下一步", ["LC：不能宣称资源使用跨 seed 稳定", "SY：配置 IC=2，但当前准备 MZX 的 2014 treatment 为 IC=0", "不修改 reward、IC 或动作空间"])
    prs.save(PPT)


def main() -> None:
    lc = load_lc()
    plot_lc(lc)
    station, evidence = update_tables(lc)
    candidate = update_candidate()
    write_doc(lc, station, evidence, candidate)
    write_ppt(lc, station)
    print("021_01 LC retry finalization completed")


if __name__ == "__main__":
    main()
