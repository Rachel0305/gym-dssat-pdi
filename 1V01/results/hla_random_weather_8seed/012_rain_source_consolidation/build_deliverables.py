from pathlib import Path
import pandas as pd,json
from pptx import Presentation
from pptx.util import Inches,Pt
R=Path(__file__).resolve().parents[3];O=Path(__file__).resolve().parent
g=json.loads((O/'final_gate.json').read_text());y=pd.read_csv(O/'coverage_summary_by_year.csv');m=pd.read_csv(O/'coverage_summary_by_month.csv');u=pd.read_csv(O/'unresolved_dates_final.csv');runs=pd.read_csv(O/'unresolved_runs.csv');h=json.loads((O/'source_hierarchy.json').read_text())
def table(df):
 return '| '+' | '.join(df.columns)+' |\n|'+ '|'.join(['---']*len(df.columns))+'|\n'+'\n'.join('| '+' | '.join(str(v) for v in row)+' |' for row in df.itertuples(index=False,name=None))
months=m[m.year==2013][['month','unresolved_days']];non=u.nearest_available_nonapproved_station.value_counts(dropna=False).to_dict()
report=f'''# 012 HLA 2004–2013 降雨来源整合

执行日期：2026-10-05。仅使用拟合期2004–2013。所有旧证据和原始天气保持不变。

## 最终结论

`HLA_GHCN50756_GENERAL_FALLBACK = ACCEPTED`

`HLA_RAIN_PROVENANCE = PARTIAL`

`READY_TO_RESUME_010 = FALSE`

3653天中3542天已确定来源，剩余111天（3.039%）全部在2013。未生成标记为完整通过的 rainfall candidate，也未执行五变量拼接、CLI、WGEN、DSSAT或PPO。

## 依据与来源优先级

011已确认50756官方原始缓存、目标244天完整、0日事件一致率85.13%、距离约5.38 km和高差6 m。012明确授权把已接受的50756补源扩展到整个拟合期的本地空白日，仍要求该日PRCP有效、QFLAG为空，并保留MFLAG/SFLAG。这是新的source-policy适用范围扩展，011旧记录未被改写。

本地D222/旧表的实际数值优先；按012将其记录为共享来源链，不互相作为独立验证。两文件实际数值已核对一致。全空白保持未知，不据日期行或WTH的零推断无雨。

其次为50756，再其次为008已通过筛选的50658（101.15 km）和50854（168.91 km），每天按距离最近且有效的顺序选择。未通过站仅提供调查线索，NASA仅作参考。

008次级站的原事件筛选曾使用本地空白转零口径。本次严格按012复用被冻结的PASS名单和门槛，不声称这是新增的011式独立事件验证；该方法限制保留在source_hierarchy.json中。

## 实际来源贡献

| 来源 | 天数 |
|---|---:|
| 本地D222/旧表真实数值 | 962 |
| GHCN 50756 | 2507 |
| GHCN 50658 | 46 |
| GHCN 50854 | 27 |
| 未解决 | 111 |

相较011中仅接受244天窗口的策略，本次新增闭合2336天；2580个原本地空白日已有有效GHCN来源。

## 逐年覆盖

{table(y)}

2013的1149.8 mm只是已选日值小计，不能称全年总量，湿日数和最大日雨量也未作为完整年指标输出。其他年份完整，但属于逐日来源拼接的序列，不能称单一站点观测。

## 最终2013缺口

最早2013-01-03，最晚2013-12-28；最长连续缺口9天。完整日期见unresolved_dates_final.csv，连续段见unresolved_runs.csv，不能只凭最早/最晚日期视为连续全年缺口。

{table(months)}

111天上本地均无明确数值；50756、50658、50854均为原始PRCP=-9999，原因是缺测而非质量标记剔除。各日QFLAG和状态原样保留。NASA参考35天为0，76天>0；这些值均未升级为来源。

已有缓存中最近的未批准可用站分布：{non}。这只提供下一步调查线索，原FAIL不被改为PASS，不新增下载。

## 最小恢复需求

只需逐日解决列出的111天：优先获取海伦同站原始日记录、50756/CMA这些日期的官方补档，或针对已缓存但未通过的邻站取得能解除原筛选失败原因的新证据并另行授权审计。禁止直接选用FAIL站、NASA、插值或零值。待111天都有可信来源后，才能完成3653天RAIN级QC并恢复010 Phase B。

## 复核与文件

已核对008每个HLA缓存站.dly哈希和011输入哈希；原始数据未变。3653天日期唯一、非缺失选值非负、本地真实值保留；2580个补值的站点选择及数值与008保存候选表逐日一致。全部输入的路径、SHA-256、大小、mtime见input_provenance.json。

正式输出包括daily_source_selection_2004_2013.csv、source_hierarchy.json、逐年/月统计、secondary_candidate_ranking.csv、unresolved_dates_final.csv、2013_unresolved_focus.csv、final_gate.json及verification.json。原始缓存包含其他站点时，仅HLA相关记录参与计算。

建议提交012 prompt、新审计/留档脚本、小型CSV/JSON、中文报告和PPT；不提交原始天气数据。本次未commit或push。建议commit message：`hla: consolidate rainfall provenance across local and GHCN sources`。
'''
(O/'findings.md').write_text(report,encoding='utf-8');(R/'docs/hla_rain_source_consolidation_012.md').write_text(report,encoding='utf-8')
slides=[('HLA降水来源整合：仍缺111天',['011已接受2004年前八个月244天补源','012授权50756扩展至整个拟合期本地空白日','此次已闭合3542/3653天，97.0%']),('逐日来源优先级',['保留本地D222/旧表真实数值','本地无值时选择50756有效PRCP','仍缺时按距离选择008已批准50658、50854','均无值则保留未知，NASA只作参考']),('2004–2013来源覆盖',['本地962天；50756补2507天','50658补46天；50854补27天','2004–2012各年完整，2013仍有111天缺口','2013已选1149.8 mm仅为小计']),('2013缺口分布',['1–6月：16、11、16、4、8、4天','7–12月：3、3、6、15、9、16天','最长连续9天，三家批准GHCN站均为-9999','NASA参考：35天=0，76天>0']),('最终门禁与恢复条件',['50756 general fallback：ACCEPTED','HLA_RAIN_PROVENANCE：PARTIAL','READY_TO_RESUME_010：FALSE','下一步只需闭合精确列出的111天日来源'])]
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
for i,(title,lines) in enumerate(slides):
 s=p.slides.add_slide(p.slide_layouts[6]);b=s.shapes.add_textbox(Inches(.6),Inches(.6),Inches(12.1),Inches(1));b.text_frame.text=title;pp=b.text_frame.paragraphs[0];pp.font.name='Microsoft YaHei';pp.font.size=Pt(32);pp.font.bold=True
 b=s.shapes.add_textbox(Inches(.8),Inches(2),Inches(11.7),Inches(4.6));b.text_frame.word_wrap=True
 for j,line in enumerate(lines):
  pp=b.text_frame.paragraphs[0] if j==0 else b.text_frame.add_paragraph();pp.text=line;pp.font.name='Microsoft YaHei';pp.font.size=Pt(24);pp.space_after=Pt(22)
 b=s.shapes.add_textbox(Inches(.8),Inches(7),Inches(11),Inches(.25));b.text_frame.text=f'012 实验留档  2026-10-05   {i+1}/5';b.text_frame.paragraphs[0].font.size=Pt(11)
p.save(R/'docs/hla_rain_source_consolidation_012.pptx');print('Saved report and 5-slide record.');print(non)
