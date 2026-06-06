from __future__ import annotations

import shutil
import sys
import traceback
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd

from ppo_safe_rendering import PROJECT_ROOT, load_yaml
from ppo_train import train_one_policy
from reward_candidates_v2 import (
    BATCHES,
    REWARD_CANDIDATES_V2,
    candidate_table,
    patch_reward_candidate_v2,
    unpatch_reward_candidate_v2,
)


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_reward_cost_tuning.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "reward_cost_tuning"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_reward_cost_coefficient_tuning_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_reward_cost_coefficient_tuning_report.pptx"


def ensure_dirs() -> None:
    for sub in [
        "configs",
        "reward_versions",
        "smoke_checks",
        "rendered_inputs",
        "models",
        "logs",
        "tensorboard",
        "daily_outputs",
        "evaluation",
        "figures/summary",
        "reports",
    ]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)


def df_to_markdown(df: pd.DataFrame) -> str:
    if df.empty:
        return ""
    work = df.copy()
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    separator = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, separator, *rows])


def write_design_files() -> None:
    rows = pd.DataFrame(candidate_table())
    rows.to_csv(OUTPUT_ROOT / "reward_versions" / "reward_candidate_v2_table.csv", index=False, encoding="utf-8-sig")
    shutil.copy2(PROJECT_ROOT / "src" / "reward_candidates_v2.py", OUTPUT_ROOT / "reward_versions" / "reward_candidates_v2.py")
    design = """# Reward candidate v2 design

Generated at: 2026-06-06

This stage does not overwrite the original gym-DSSAT reward. Candidate rewards are injected at runtime from `src/reward_candidates_v2.py`.

## Candidate families

- D: stronger linear water and nitrogen costs.
- E: linear costs plus cumulative quadratic costs over total seasonal water and nitrogen used so far.
- F: target interval penalty, with upper bounds 220 mm irrigation and 320 kg ha-1 N.
- G: normalized economic proxy using daily positive grain-yield increments minus water and nitrogen costs.

The G family is not a real RMB economic return yet. It is a normalized debug proxy and should later be replaced by interpretable grain, water, and fertilizer prices.

## Candidate table

""" + df_to_markdown(rows)
    (OUTPUT_ROOT / "reward_versions" / "reward_candidate_v2_design.md").write_text(design, encoding="utf-8")


def previous_failure_review() -> None:
    prev = PROJECT_ROOT / "Leave_One_experiments" / "reward_cost_debug" / "evaluation" / "reward_candidate_comparison.csv"
    df = pd.read_csv(prev)
    cols = [
        "reward_version",
        "reward_family",
        "mean_yield",
        "mean_reward",
        "mean_irrigation",
        "mean_n",
        "mean_irrigation_saturation_ratio",
        "mean_n_saturation_ratio",
        "recommendation_pass",
    ]
    review = """# Previous candidate failure review

Generated at: 2026-06-06

The 006_03 HLA pilot tested current_reward_baseline, A1-A3, B1-B3, and C1. All completed HLA 2011/2007/2009 evaluations but all candidates still used 300 mm irrigation and 450 kg ha-1 N. Therefore no candidate passed the unsaturated filter and SYA/LCA extension was skipped.

## Previous comparison

""" + df_to_markdown(df[cols]) + """

## Failure diagnosis

- Linear A costs were too small relative to growth/recovery reward and did not change the action allocation.
- B quadratic excess penalties changed reward scale strongly, but with only 5000 timesteps PPO still evaluated as cap-saturated.
- C1 lacked a true terminal done flag and used incremental grain-yield proxy, so it did not provide a clean season-level objective.
- The reward callback is step-wise, while the desired behavior is season-level water and nitrogen budgeting. This scale mismatch can make coefficient tuning brittle.
- This round strengthens costs and tests cumulative penalties that are visible before the cap is reached.
"""
    (OUTPUT_ROOT / "evaluation" / "previous_candidate_failure_review.md").write_text(review, encoding="utf-8")


def build_plan(config: dict) -> pd.DataFrame:
    rows = []
    order = 1
    for batch in BATCHES:
        for reward_version in batch:
            cand = REWARD_CANDIDATES_V2[reward_version]
            rows.append(
                {
                    "run_order": order,
                    "station": "HLA",
                    "train_year": 2011,
                    "eval_years": "2011,2007,2009",
                    "reward_version": reward_version,
                    "reward_family": cand.family,
                    "batch": cand.batch,
                    "cap_name": config["cap"]["cap_name"],
                    "season_irrigation_cap": config["cap"]["season_irrigation_cap"],
                    "season_n_cap": config["cap"]["season_n_cap"],
                    "seed": config["seed"],
                    "total_timesteps": config["training"]["total_timesteps"],
                }
            )
            order += 1
    plan = pd.DataFrame(rows)
    plan.to_csv(OUTPUT_ROOT / "configs" / "reward_cost_tuning_plan.csv", index=False, encoding="utf-8-sig")
    shutil.copy2(CONFIG_PATH, OUTPUT_ROOT / "configs" / CONFIG_PATH.name)
    return plan


def row_metadata(config: dict, reward_version: str, station: str) -> dict:
    cand = REWARD_CANDIDATES_V2[reward_version]
    return {
        "reward_version": reward_version,
        "reward_family": cand.family,
        "batch": cand.batch,
        "cap_name": config["cap"]["cap_name"],
        "season_irrigation_cap": float(config["cap"]["season_irrigation_cap"]),
        "season_n_cap": float(config["cap"]["season_n_cap"]),
        "cv_type": config["training"]["cv_types"].get(station, ""),
    }


def run_training(config: dict, reward_version: str, station: str, train_year: int) -> None:
    patch_reward_candidate_v2(reward_version)
    try:
        train_one_policy(
            station=station,
            train_year=train_year,
            seed=int(config["seed"]),
            total_timesteps=int(config["training"]["total_timesteps"]),
            config_path=CONFIG_PATH,
            debug=False,
            policy_tag=reward_version,
            row_metadata=row_metadata(config, reward_version, station),
        )
    finally:
        unpatch_reward_candidate_v2()


def enrich_summary() -> pd.DataFrame:
    src = OUTPUT_ROOT / "evaluation" / "ppo_evaluation_summary.csv"
    if not src.exists():
        return pd.DataFrame()
    df = pd.read_csv(src)
    for col in ["total_irrigation", "total_n_fertilizer", "season_irrigation_cap", "season_n_cap", "final_grnwt"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["irrigation_saturation_ratio"] = df["total_irrigation"] / df["season_irrigation_cap"]
    df["n_saturation_ratio"] = df["total_n_fertilizer"] / df["season_n_cap"]
    df["quality_gate_pass"] = (
        (df["run_status"] == "ok")
        & (df["episode_completed"].astype(str).str.lower().isin(["true", "1"]))
        & (df["irrigation_saturation_ratio"] <= 1.000001)
        & (df["n_saturation_ratio"] <= 1.000001)
    )
    out = OUTPUT_ROOT / "evaluation" / "reward_cost_tuning_evaluation_summary.csv"
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return df


def baseline_yield() -> float:
    prev = pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "reward_cost_debug" / "evaluation" / "reward_candidate_comparison.csv")
    row = prev[prev["reward_version"] == "current_reward_baseline"]
    return float(row.iloc[0]["mean_yield"])


def compare_candidates(df: pd.DataFrame) -> pd.DataFrame:
    hla = df[df["station"] == "HLA"].copy()
    grouped = (
        hla.groupby(["reward_version", "reward_family", "batch"], dropna=False)
        .agg(
            eval_count=("eval_year", "count"),
            ok_count=("quality_gate_pass", "sum"),
            mean_yield=("final_grnwt", "mean"),
            std_yield=("final_grnwt", "std"),
            mean_reward=("mean_reward", "mean"),
            mean_irrigation=("total_irrigation", "mean"),
            mean_n=("total_n_fertilizer", "mean"),
            mean_irrigation_saturation_ratio=("irrigation_saturation_ratio", "mean"),
            mean_n_saturation_ratio=("n_saturation_ratio", "mean"),
            mean_swfac=("mean_swfac", "mean"),
            mean_nstres=("mean_nstres", "mean"),
        )
        .reset_index()
    )
    base_yield = baseline_yield()
    grouped["yield_loss_vs_baseline"] = (base_yield - grouped["mean_yield"]) / base_yield
    grouped["input_reduction_vs_baseline"] = 1.0 - (grouped["mean_irrigation"] + grouped["mean_n"]) / (300.0 + 450.0)
    grouped["strict_pass"] = (
        (grouped["mean_irrigation_saturation_ratio"] < 0.95)
        & (grouped["mean_n_saturation_ratio"] < 0.95)
        & (grouped["yield_loss_vs_baseline"] <= 0.10)
        & (grouped["ok_count"] == grouped["eval_count"])
    )
    grouped["relaxed_pass"] = (
        (grouped["mean_irrigation_saturation_ratio"] < 0.95)
        & (grouped["mean_n_saturation_ratio"] < 0.95)
        & (grouped["yield_loss_vs_baseline"] <= 0.15)
        & (grouped["ok_count"] == grouped["eval_count"])
    )
    grouped["partial_pass"] = (
        (
            (grouped["mean_irrigation_saturation_ratio"] < 0.95)
            | (grouped["mean_n_saturation_ratio"] < 0.95)
        )
        & (grouped["ok_count"] == grouped["eval_count"])
    )
    out = OUTPUT_ROOT / "evaluation" / "reward_cost_tuning_candidate_comparison.csv"
    grouped.to_csv(out, index=False, encoding="utf-8-sig")
    return grouped


def select_candidate(comparison: pd.DataFrame) -> str | None:
    passed = comparison[comparison["strict_pass"]].copy()
    if passed.empty:
        passed = comparison[comparison["relaxed_pass"]].copy()
    if passed.empty:
        return None
    passed = passed.sort_values(["yield_loss_vs_baseline", "input_reduction_vs_baseline"], ascending=[True, False])
    return str(passed.iloc[0]["reward_version"])


def copy_candidate_figures(df: pd.DataFrame) -> None:
    for _, row in df.iterrows():
        src = PROJECT_ROOT / str(row.get("figure_dir", ""))
        reward_version = row.get("reward_version")
        eval_year = row.get("eval_year")
        station = row.get("station")
        if not src.exists() or not reward_version or pd.isna(eval_year):
            continue
        dst = OUTPUT_ROOT / "figures" / str(station) / str(reward_version) / f"eval_{int(eval_year)}"
        if dst.exists():
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copytree(src, dst)


def plot_summary(df: pd.DataFrame, comparison: pd.DataFrame) -> None:
    fig_root = OUTPUT_ROOT / "figures" / "summary"
    fig_root.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = ["DejaVu Sans"]
    x = comparison["reward_version"]

    plt.figure(figsize=(11, 6))
    plt.scatter(comparison["mean_irrigation"] + comparison["mean_n"], comparison["mean_yield"])
    for _, row in comparison.iterrows():
        plt.annotate(row["reward_version"], (row["mean_irrigation"] + row["mean_n"], row["mean_yield"]), fontsize=8)
    plt.xlabel("Mean water + nitrogen input")
    plt.ylabel("Mean grain yield")
    plt.tight_layout()
    plt.savefig(fig_root / "reward_tuning_yield_vs_input.png", dpi=180)
    plt.close()

    plt.figure(figsize=(12, 6))
    plt.plot(x, comparison["mean_irrigation_saturation_ratio"], marker="o", label="irrigation ratio")
    plt.plot(x, comparison["mean_n_saturation_ratio"], marker="o", label="nitrogen ratio")
    plt.axhline(0.95, color="black", linestyle="--", linewidth=1)
    plt.xticks(rotation=45, ha="right")
    plt.ylabel("Saturation ratio")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_root / "reward_tuning_saturation_ratio.png", dpi=180)
    plt.close()

    plt.figure(figsize=(11, 6))
    plt.scatter(comparison["input_reduction_vs_baseline"], comparison["yield_loss_vs_baseline"])
    for _, row in comparison.iterrows():
        plt.annotate(row["reward_version"], (row["input_reduction_vs_baseline"], row["yield_loss_vs_baseline"]), fontsize=8)
    plt.axhline(0.10, color="black", linestyle="--", linewidth=1)
    plt.axhline(0.15, color="gray", linestyle="--", linewidth=1)
    plt.xlabel("Input reduction vs baseline")
    plt.ylabel("Yield loss vs baseline")
    plt.tight_layout()
    plt.savefig(fig_root / "reward_tuning_yield_loss_vs_input_reduction.png", dpi=180)
    plt.close()

    plt.figure(figsize=(11, 6))
    plt.scatter(comparison["mean_yield"], comparison["mean_reward"])
    for _, row in comparison.iterrows():
        plt.annotate(row["reward_version"], (row["mean_yield"], row["mean_reward"]), fontsize=8)
    plt.xlabel("Mean yield")
    plt.ylabel("Mean reward")
    plt.tight_layout()
    plt.savefig(fig_root / "reward_tuning_reward_vs_yield.png", dpi=180)
    plt.close()

    plt.figure(figsize=(12, 6))
    plt.plot(x, comparison["mean_swfac"], marker="o", label="swfac")
    plt.plot(x, comparison["mean_nstres"], marker="o", label="nstres")
    plt.xticks(rotation=45, ha="right")
    plt.ylabel("Mean stress indicators")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_root / "reward_tuning_swfac_nstres.png", dpi=180)
    plt.close()


def run_hla_batches(config: dict) -> tuple[pd.DataFrame, pd.DataFrame, str | None, list[str]]:
    completed_batches: list[str] = []
    selected: str | None = None
    for batch in BATCHES:
        batch_name = REWARD_CANDIDATES_V2[batch[0]].batch
        completed_batches.append(batch_name)
        for reward_version in batch:
            print(f"[reward-cost-tuning] HLA pilot: {reward_version}", flush=True)
            run_training(config, reward_version, "HLA", 2011)
        df = enrich_summary()
        copy_candidate_figures(df)
        comparison = compare_candidates(df)
        plot_summary(df, comparison)
        selected = select_candidate(comparison)
        if selected and config["training"].get("stop_after_batch_if_strict_pass", True):
            break
    return enrich_summary(), compare_candidates(enrich_summary()), selected, completed_batches


def run_extension(config: dict, selected: str | None) -> pd.DataFrame:
    if not selected:
        out = OUTPUT_ROOT / "evaluation" / "cross_site_reward_cost_tuning_test.csv"
        empty = pd.DataFrame([{"selected_reward": "", "run_status": "skipped", "notes": "No HLA candidate passed strict or relaxed criteria."}])
        empty.to_csv(out, index=False, encoding="utf-8-sig")
        return empty
    for station in ["SYA", "LCA"]:
        train_year = int(config["training"]["best_train_years"][station])
        print(f"[reward-cost-tuning] Extension: {station} {train_year} {selected}", flush=True)
        run_training(config, selected, station, train_year)
    df = enrich_summary()
    copy_candidate_figures(df)
    cross = df[df["station"].isin(["SYA", "LCA"]) & (df["reward_version"] == selected)].copy()
    cross.to_csv(OUTPUT_ROOT / "evaluation" / "cross_site_reward_cost_tuning_test.csv", index=False, encoding="utf-8-sig")
    return cross


def write_report(comparison: pd.DataFrame, cross: pd.DataFrame, selected: str | None, completed_batches: list[str]) -> None:
    report_table = comparison.copy()
    for col in report_table.select_dtypes(include=["float", "int"]).columns:
        report_table[col] = report_table[col].round(4)
    cross_table = cross.copy()
    if len(cross_table):
        for col in cross_table.select_dtypes(include=["float", "int"]).columns:
            cross_table[col] = cross_table[col].round(4)
    strict = comparison[comparison["strict_pass"]]["reward_version"].tolist()
    relaxed = comparison[comparison["relaxed_pass"]]["reward_version"].tolist()
    partial = comparison[comparison["partial_pass"]]["reward_version"].tolist()
    md = f"""# Reward cost coefficient tuning report

Generated at: 2026-06-06

## Goal

This stage continues from 006_03 and tests stronger cost coefficients under HLA 2011 action-safe PPO with cap 300 mm irrigation / 450 kg ha-1 N. It does not modify the original site-packages reward and does not train PPO without action safety.

## Why 006_03 failed

All previous A/B/C candidates completed evaluation, but all used 300 mm irrigation and 450 kg ha-1 N. Stronger reward costs are needed, or the reward structure must be redesigned around a clearer season-level objective.

## Batches completed

{', '.join(completed_batches)}

## HLA comparison

{df_to_markdown(report_table)}

## Pass summary

- Strict pass candidates: {', '.join(strict) if strict else 'None'}
- Relaxed pass candidates: {', '.join(relaxed) if relaxed else 'None'}
- Partial pass candidates: {', '.join(partial) if partial else 'None'}
- Recommended candidate: `{selected or 'None'}`

## SYA/LCA extension

{df_to_markdown(cross_table) if len(cross_table) else 'Extension was skipped because HLA did not identify a strict or relaxed passing candidate.'}

## Interpretation

If no candidate becomes unsaturated, coefficient tuning alone is insufficient. The likely next step is to redesign reward as an episode-level objective, such as final yield or profit minus total water and nitrogen costs, and expose terminal season information explicitly instead of approximating terminal reward from step-wise callbacks.

## Next step

Do not enter multi-seed or rainfall-scaling budget scenario unless HLA and the small SYA/LCA extension are no longer dominated by cap saturation.
"""
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)
    DOC_MD.write_text(md, encoding="utf-8")
    shutil.copy2(DOC_MD, OUTPUT_ROOT / "reports" / DOC_MD.name)
    write_ppt(comparison, selected, completed_batches)


def write_ppt(comparison: pd.DataFrame, selected: str | None, completed_batches: list[str]) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        (OUTPUT_ROOT / "reports" / "ppt_generation_skipped_in_container.txt").write_text(
            "python-pptx is not installed in the DSSAT container. Generate the PPT from Windows Python after training.",
            encoding="utf-8",
        )
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_title(slide, title: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True

    def add_body(slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.0), Inches(5.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(16)

    def picture_slide(title: str, file_name: str) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, title)
        path = OUTPUT_ROOT / "figures" / "summary" / file_name
        if path.exists():
            slide.shapes.add_picture(str(path), Inches(0.8), Inches(1.0), width=Inches(11.8))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Reward Cost Coefficient Tuning")
    add_body(slide, "目标：继续调大 reward 成本系数，检查 HLA 是否不再打满 300/450 cap。\n约束：action-safe PPO；不做多 seed；不训练 FQA/YCA；不进入 rainfall-scaling budget scenario。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Why 006_03 Failed")
    add_body(slide, "上一轮 A/B/C candidates 全部完成评估，但全部 mean_irrigation=300、mean_n=450。\n说明成本项改变了 reward 数值，但没有让 PPO 主动节约水氮。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "New Candidate Batches")
    add_body(slide, f"已运行批次：{', '.join(completed_batches)}\nD：强线性成本。\nE：累计二次成本。\nF：目标区间上限惩罚。\nG：经济学近似 reward。")

    picture_slide("Yield vs Input", "reward_tuning_yield_vs_input.png")
    picture_slide("Saturation Ratio", "reward_tuning_saturation_ratio.png")
    picture_slide("Yield Loss vs Input Reduction", "reward_tuning_yield_loss_vs_input_reduction.png")
    picture_slide("Reward vs Yield", "reward_tuning_reward_vs_yield.png")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Recommendation")
    strict = comparison[comparison["strict_pass"]]["reward_version"].tolist()
    relaxed = comparison[comparison["relaxed_pass"]]["reward_version"].tolist()
    add_body(slide, f"Strict pass：{', '.join(strict) if strict else '无'}\nRelaxed pass：{', '.join(relaxed) if relaxed else '无'}\n推荐 candidate：{selected or '暂无'}\n若仍无通过项，下一步应重构 reward，而不是继续简单加大系数。")

    DOC_PPT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DOC_PPT)
    shutil.copy2(DOC_PPT, OUTPUT_ROOT / "reports" / DOC_PPT.name)


def main() -> None:
    ensure_dirs()
    config = load_yaml(CONFIG_PATH)
    write_design_files()
    previous_failure_review()
    build_plan(config)
    try:
        df, comparison, selected, completed_batches = run_hla_batches(config)
        cross = run_extension(config, selected)
        df = enrich_summary()
        comparison = compare_candidates(df)
        plot_summary(df, comparison)
        write_report(comparison, cross, selected, completed_batches)
    except Exception:
        err = OUTPUT_ROOT / "reports" / "reward_cost_tuning_error.log"
        err.write_text(traceback.format_exc(), encoding="utf-8")
        print(traceback.format_exc(), file=sys.stderr)
        raise


if __name__ == "__main__":
    main()
