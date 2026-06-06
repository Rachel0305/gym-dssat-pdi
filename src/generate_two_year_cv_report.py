from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))

from ppo_safe_rendering import PROJECT_ROOT, load_yaml


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_ppo_action_safe_two_year_cv.yaml"
REPORT_STEM = f"{date.today().isoformat()}_fqa_yca_two_year_action_safe_cv_report"


def fmt(value, digits: int = 2) -> str:
    try:
        return f"{float(value):.{digits}f}"
    except Exception:
        return ""


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path) if path.exists() else pd.DataFrame()


def output_paths(config: dict) -> dict[str, Path]:
    output_root = PROJECT_ROOT / config["paths"]["output_root"]
    return {
        "output_root": output_root,
        "plan": output_root / "configs" / "two_year_cv_training_plan.csv",
        "summary": output_root / "evaluation" / "two_year_cv_ppo_evaluation_summary.csv",
        "saturation": output_root / "evaluation" / "action_safety_saturation_check.csv",
        "best": output_root / "strategy_selection" / "best_policy_by_site.csv",
        "ranking": output_root / "strategy_selection" / "policy_ranking_by_site.csv",
        "smoke": output_root / "smoke_checks" / "pretrain_smoke_check_summary.csv",
        "all_site": PROJECT_ROOT / "Leave_One_experiments" / "ppo_action_safe_summary" / "all_site_best_policy_summary.csv",
    }


def write_markdown(config: dict) -> Path:
    paths = output_paths(config)
    summary = read_csv(paths["summary"])
    saturation = read_csv(paths["saturation"])
    best = read_csv(paths["best"])
    smoke = read_csv(paths["smoke"])
    all_site = read_csv(paths["all_site"])
    docs_path = PROJECT_ROOT / "docs" / f"{REPORT_STEM}.md"
    lines: list[str] = []
    lines.append("# FQA/YCA two-year action-safe PPO cross validation report")
    lines.append("")
    lines.append(f"Generated at: {date.today().isoformat()}")
    lines.append("")
    lines.append("## Goal")
    lines.append("")
    lines.append("This stage trains only action-safe PPO models for FQA and YCA. Each site has two observed years, so the result is explicitly marked as `limited_two_year_cross_validation`, not full dry/normal/wet stability validation.")
    lines.append("")
    lines.append("## Constraints followed")
    lines.append("")
    lines.append("- No PPO without action safety was trained.")
    lines.append("- Reward functions and my_data originals were not modified.")
    lines.append("- No multi-seed training was performed.")
    lines.append("- HLA/SYA/LCA previous results were not overwritten.")
    lines.append(f"- Timesteps per model: {config.get('training', {}).get('total_timesteps')}")
    lines.append("")
    lines.append("## Action safety parameters")
    lines.append("")
    for key, value in config.get("action_safety", {}).items():
        lines.append(f"- `{key}`: {value}")
    lines.append("")
    lines.append("## Pretrain smoke checks")
    lines.append("")
    if smoke.empty:
        lines.append("No pretrain smoke check summary found.")
    else:
        lines.append("| station | train_year | policy | status | completed |")
        lines.append("|---|---:|---|---|---|")
        for row in smoke.itertuples():
            lines.append(f"| {row.station} | {row.train_year} | {row.policy_name} | {row.run_status} | {row.episode_completed} |")
    lines.append("")
    lines.append("## Evaluation summary")
    lines.append("")
    if summary.empty:
        lines.append("No evaluation summary found.")
    else:
        cols = ["station", "policy_name", "train_year", "eval_year", "final_grnwt", "total_irrigation", "total_n_fertilizer", "mean_reward", "quality_gate_pass"]
        lines.append("| station | policy | train_year | eval_year | yield | irrigation | nitrogen | mean_reward | gate |")
        lines.append("|---|---|---:|---:|---:|---:|---:|---:|---|")
        for row in summary[cols].itertuples(index=False):
            lines.append(f"| {row.station} | {row.policy_name} | {row.train_year} | {row.eval_year} | {fmt(row.final_grnwt)} | {fmt(row.total_irrigation)} | {fmt(row.total_n_fertilizer)} | {fmt(row.mean_reward)} | {row.quality_gate_pass} |")
    lines.append("")
    lines.append("## Action safety saturation")
    lines.append("")
    if saturation.empty:
        lines.append("No action safety saturation check found.")
    else:
        at_cap = saturation["irrigation_at_season_limit"].astype(str).str.lower().eq("true").all() and saturation["n_at_season_limit"].astype(str).str.lower().eq("true").all()
        lines.append(f"- All FQA/YCA evaluations at 200 mm / 300 kg/ha cap: {at_cap}")
        lines.append("| station | policy | eval_year | irrigation | nitrogen | trigger_days | dominant_rule |")
        lines.append("|---|---|---:|---:|---:|---:|---|")
        for row in saturation.itertuples():
            lines.append(f"| {row.station} | {row.policy_name} | {row.eval_year} | {fmt(row.total_irrigation)} | {fmt(row.total_n_fertilizer)} | {row.num_safety_trigger_days} | {row.dominant_safety_rule} |")
        if at_cap:
            lines.append("")
            lines.append("当前 action-safe PPO 的策略比较更多反映了固定 season cap 下的时序分配差异，而不是自由水氮优化；下一阶段需要考虑 reward 成本项或 season cap 敏感性分析。")
    lines.append("")
    lines.append("## FQA/YCA best policy")
    lines.append("")
    if best.empty:
        lines.append("No best policy table found.")
    else:
        lines.append("| station | best_policy | train_year | validation_years | mean_yield | mean_irrigation | mean_n | stability_score |")
        lines.append("|---|---|---:|---:|---:|---:|---:|---:|")
        for row in best.itertuples():
            lines.append(f"| {row.station} | {row.best_policy_name} | {row.best_train_year} | {row.validation_years} | {fmt(row.mean_yield)} | {fmt(row.mean_irrigation)} | {fmt(row.mean_n_fertilizer)} | {fmt(row.stability_score, 3)} |")
    lines.append("")
    lines.append("## Five-site best policy summary")
    lines.append("")
    if all_site.empty:
        lines.append("No five-site best policy summary found.")
    else:
        lines.append(f"- Path: `{paths['all_site'].relative_to(PROJECT_ROOT)}`")
        lines.append("| station | best_policy | train_year | cv_type | mean_yield | irrigation | nitrogen | limited_data |")
        lines.append("|---|---|---:|---|---:|---:|---:|---|")
        for row in all_site.itertuples():
            lines.append(f"| {row.station} | {row.best_policy} | {row.train_year} | {row.cross_validation_type} | {fmt(row.mean_yield)} | {fmt(row.mean_irrigation)} | {fmt(row.mean_n)} | {row.limited_data_flag} |")
    lines.append("")
    lines.append("## Next step recommendation")
    lines.append("")
    lines.append("Because FQA/YCA are also expected to be checked for cap saturation, the next robust step is season cap sensitivity analysis and/or adding explicit water/nitrogen cost terms before treating these policies as final management recommendations. Multi-seed stability analysis should come after the action/cost design is less dominated by fixed caps.")
    docs_path.write_text("\n".join(lines), encoding="utf-8")
    report_copy = paths["output_root"] / "reports" / docs_path.name
    report_copy.parent.mkdir(parents=True, exist_ok=True)
    report_copy.write_text(docs_path.read_text(encoding="utf-8"), encoding="utf-8")
    return docs_path


def write_pptx(config: dict, markdown_path: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    paths = output_paths(config)
    summary = read_csv(paths["summary"])
    saturation = read_csv(paths["saturation"])
    best = read_csv(paths["best"])
    all_site = read_csv(paths["all_site"])
    ppt_path = PROJECT_ROOT / "docs" / f"{REPORT_STEM}.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def title(slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.2), Inches(0.55))
        p = box.text_frame.paragraphs[0]
        p.text = text
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0)

    def text(slide, body: str, top: float = 1.0) -> None:
        box = slide.shapes.add_textbox(Inches(0.65), Inches(top), Inches(12.0), Inches(5.9))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(17)
            p.font.color.rgb = RGBColor(0, 0, 0)

    def table(slide, df: pd.DataFrame, top: float = 1.0) -> None:
        if df.empty:
            text(slide, "No rows available.", top)
            return
        rows = min(len(df), 8) + 1
        cols = len(df.columns)
        shape = slide.shapes.add_table(rows, cols, Inches(0.35), Inches(top), Inches(12.6), Inches(5.3))
        t = shape.table
        for c, col in enumerate(df.columns):
            cell = t.cell(0, c)
            cell.text = str(col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(79, 129, 189)
        for r, (_, row) in enumerate(df.head(8).iterrows(), start=1):
            for c, col in enumerate(df.columns):
                cell = t.cell(r, c)
                cell.text = str(row[col])
                if r % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(230, 236, 247)
        for tr in t.rows:
            for cell in tr.cells:
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for run in p.runs:
                        run.font.name = "Microsoft YaHei"
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(0, 0, 0)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "FQA/YCA two-year action-safe PPO")
    text(slide, "本阶段只训练 FQA 和 YCA。\n两站均只有 2 个实测年份，因此标记为 limited_two_year_cross_validation。\n不训练无 safety PPO，不修改 reward，不做多 seed，不覆盖 HLA/SYA/LCA 结果。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Two-year cross validation design")
    text(slide, "FQA: 2008 -> 2010, 2010 -> 2008\nYCA: 2014 -> 2008, 2008 -> 2014\n每个模型训练前做 null_zero 和 fixed_low_input smoke check。\n每个模型评估训练年和另一个验证年。")

    if not summary.empty:
        show = summary[["station", "policy_name", "eval_year", "final_grnwt", "total_irrigation", "total_n_fertilizer", "quality_gate_pass"]].copy()
        for col in ["final_grnwt", "total_irrigation", "total_n_fertilizer"]:
            show[col] = show[col].map(lambda v: fmt(v))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title(slide, "Evaluation and quality gate")
        table(slide, show)

    if not saturation.empty:
        sat = saturation[["station", "policy_name", "eval_year", "total_irrigation", "total_n_fertilizer", "num_safety_trigger_days"]].copy()
        for col in ["total_irrigation", "total_n_fertilizer"]:
            sat[col] = sat[col].map(lambda v: fmt(v))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title(slide, "Action safety saturation")
        table(slide, sat)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "FQA/YCA best policies")
    best_show = best[["station", "best_policy_name", "best_train_year", "validation_years", "mean_yield", "mean_irrigation", "mean_n_fertilizer", "stability_score"]].copy() if not best.empty else best
    if not best_show.empty:
        for col in ["mean_yield", "mean_irrigation", "mean_n_fertilizer", "stability_score"]:
            best_show[col] = best_show[col].map(lambda v: fmt(v, 3 if col == "stability_score" else 2))
    table(slide, best_show)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Five-site best policy summary")
    all_show = all_site[["station", "best_policy", "train_year", "cross_validation_type", "mean_yield", "mean_irrigation", "mean_n", "limited_data_flag"]].copy() if not all_site.empty else all_site
    if not all_show.empty:
        for col in ["mean_yield", "mean_irrigation", "mean_n"]:
            all_show[col] = all_show[col].map(lambda v: fmt(v))
    table(slide, all_show)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Limitations and next step")
    text(slide, "FQA/YCA 的结果是 two-year 初步验证，不能等价于三类干湿年份稳定性分析。\n如果所有 evaluation 都打满 200/300 上限，说明策略比较主要反映固定 season cap 下的时序差异。\n下一步优先做 season cap 敏感性分析或 reward 成本项修正，然后再做多 seed 稳定性分析。")

    prs.save(ppt_path)
    report_copy = paths["output_root"] / "reports" / ppt_path.name
    report_copy.write_bytes(ppt_path.read_bytes())
    return ppt_path


def main() -> None:
    config = load_yaml(CONFIG_PATH)
    md = write_markdown(config)
    ppt = write_pptx(config, md)
    print(md.relative_to(PROJECT_ROOT))
    print(ppt.relative_to(PROJECT_ROOT))


if __name__ == "__main__":
    main()
