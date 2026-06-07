from __future__ import annotations

import argparse
import itertools
import json
import shutil
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None
    RGBColor = None
    PP_ALIGN = None
    Inches = None
    Pt = None

from offline_schedule_policy import DeterministicSchedulePolicy, ScheduleEvent
from ppo_action_safety import ActionSafetyState, apply_action_safety, normalize_action, update_action_safety_state
from ppo_evaluate import latest_observation_dict, make_env, scalar
from ppo_experiment_plan import find_year
from ppo_plot_results import plot_episode
from ppo_safe_rendering import PROJECT_ROOT, load_yaml


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_water_nitrogen_factorial_diagnosis.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "water_nitrogen_factorial_diagnosis"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_water_nitrogen_factorial_diagnosis_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_water_nitrogen_factorial_diagnosis_report.pptx"


@dataclass(frozen=True)
class Treatment:
    treatment_id: str
    treatment_name: str
    nitrogen_by_dap: dict[int, float]
    irrigation_by_dap: dict[int, float]


def ensure_dirs() -> None:
    for sub in [
        "configs",
        "fixed_factorial_schedules",
        "daily_outputs",
        "evaluation",
        "irrigation_responsive_search",
        "expert_policy",
        "imitation_dataset",
        "figures",
        "reports",
        "rendered_inputs",
        "logs",
    ]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)


def safe_copy(src: Path, dst: Path) -> None:
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, dst)


def df_to_markdown(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df.empty:
        return ""
    work = df.copy()
    if max_rows is not None:
        work = work.head(max_rows)
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(4)
    for col in work.columns:
        work[col] = work[col].map(lambda v: "" if pd.isna(v) else str(v))
    header = "| " + " | ".join(work.columns) + " |"
    sep = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, sep, *rows])


def truthy(value: Any) -> bool:
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if pd.isna(value):
        return False
    return str(value).strip().lower() in {"true", "1", "yes", "y"}


def key_to_int_dict(data: dict | None) -> dict[int, float]:
    if not data:
        return {}
    return {int(k): float(v) for k, v in data.items()}


def load_treatments(config: dict) -> list[Treatment]:
    rows = []
    for item in config["factorial_treatments"]:
        rows.append(
            Treatment(
                treatment_id=str(item["treatment_id"]),
                treatment_name=str(item["treatment_name"]),
                nitrogen_by_dap=key_to_int_dict(item.get("nitrogen_by_dap")),
                irrigation_by_dap=key_to_int_dict(item.get("irrigation_by_dap")),
            )
        )
    return rows


def treatment_policy(treatment: Treatment) -> DeterministicSchedulePolicy:
    events: list[ScheduleEvent] = []
    for dap, amount in treatment.nitrogen_by_dap.items():
        events.append(ScheduleEvent(name=f"N{dap}", dap=int(dap), nitrogen=float(amount)))
    for dap, amount in treatment.irrigation_by_dap.items():
        events.append(ScheduleEvent(name=f"I{dap}", dap=int(dap), irrigation=float(amount)))
    return DeterministicSchedulePolicy(events)


def clip_to_action_space(env, real_action: dict[str, float]) -> tuple[dict[str, float], str]:
    spaces = getattr(env.formator.action_space_dict, "spaces", env.formator.action_space_dict)
    clipped: dict[str, float] = {}
    notes: list[str] = []
    for name in env.formator.action_names:
        space = spaces[name]
        low = float(np.asarray(space.low).flatten()[0])
        high = float(np.asarray(space.high).flatten()[0])
        raw = float(real_action.get(name, 0.0))
        value = min(max(raw, low), high)
        clipped[name] = value
        if abs(raw - value) > 1e-9:
            notes.append(f"{name}_env_bound_clip_{raw:g}_to_{value:g}")
    return clipped, ";".join(notes)


def profit(final_grnwt: float, irrigation: float, nitrogen: float, config: dict, low_water: bool = False) -> float:
    econ = config["economics"]
    water_cost = float(econ["water_cost_low"] if low_water else econ["water_cost_default"])
    return (
        float(econ["grain_value_coef"]) * float(final_grnwt)
        - water_cost * float(irrigation)
        - float(econ["n_cost"]) * float(nitrogen)
    )


def daily_path(station: str, year: int, treatment_name: str) -> Path:
    return OUTPUT_ROOT / "daily_outputs" / station / f"{station}_{year}_{treatment_name}_daily.csv"


def summarize_daily(daily: pd.DataFrame, station: str, year: int, treatment: Treatment, status: str, error: str, config: dict) -> dict:
    total_i = float(sum(treatment.irrigation_by_dap.values()))
    total_n = float(sum(treatment.nitrogen_by_dap.values()))
    if len(daily):
        total_i = float(pd.to_numeric(daily["real_action_amir"], errors="coerce").fillna(0).sum())
        total_n = float(pd.to_numeric(daily["real_action_anfer"], errors="coerce").fillna(0).sum())
    final_grnwt = float(pd.to_numeric(daily["grnwt"], errors="coerce").iloc[-1]) if len(daily) else np.nan
    final_topwt = float(pd.to_numeric(daily["topwt"], errors="coerce").iloc[-1]) if len(daily) else np.nan
    final_xlai = float(pd.to_numeric(daily["xlai"], errors="coerce").iloc[-1]) if len(daily) else np.nan
    swfac = pd.to_numeric(daily["swfac"], errors="coerce") if len(daily) and "swfac" in daily else pd.Series(dtype=float)
    nstres = pd.to_numeric(daily["nstres"], errors="coerce") if len(daily) and "nstres" in daily else pd.Series(dtype=float)
    out = {
        "station": station,
        "year": int(year),
        "treatment_id": treatment.treatment_id,
        "treatment_name": treatment.treatment_name,
        "total_irrigation": total_i,
        "total_n": total_n,
        "run_status": status,
        "error_message": error,
        "episode_completed": bool(len(daily) and bool(daily["done"].iloc[-1])),
        "episode_length": int(len(daily)),
        "final_grnwt": final_grnwt,
        "final_topwt": final_topwt,
        "final_xlai": final_xlai,
        "profit_default": profit(final_grnwt, total_i, total_n, config, low_water=False) if np.isfinite(final_grnwt) else np.nan,
        "profit_low_water_cost": profit(final_grnwt, total_i, total_n, config, low_water=True) if np.isfinite(final_grnwt) else np.nan,
        "mean_swfac": float(swfac.mean()) if len(swfac) else np.nan,
        "max_swfac": float(swfac.max()) if len(swfac) else np.nan,
        "swfac_stress_days_gt_0p05": int((swfac > 0.05).sum()) if len(swfac) else 0,
        "swfac_stress_days_gt_0p10": int((swfac > 0.10).sum()) if len(swfac) else 0,
        "mean_nstres": float(nstres.mean()) if len(nstres) else np.nan,
        "max_nstres": float(nstres.max()) if len(nstres) else np.nan,
        "nstres_days_gt_0p05": int((nstres > 0.05).sum()) if len(nstres) else 0,
        "daily_csv_path": str(daily_path(station, year, treatment.treatment_name).relative_to(PROJECT_ROOT)),
        "notes": "",
    }
    return out


def run_fixed_treatment(config: dict, station: str, year: int, treatment: Treatment) -> dict:
    path = daily_path(station, year, treatment.treatment_name)
    if path.exists():
        daily = pd.read_csv(path)
        return summarize_daily(daily, station, year, treatment, "ok_cached", "", config)

    path.parent.mkdir(parents=True, exist_ok=True)
    run_tag = f"factorial_{treatment.treatment_name}"
    env = None
    records: list[dict[str, Any]] = []
    try:
        env = make_env(config, station, int(year), int(config.get("seed", 0)), run_tag=run_tag, evaluation=True, action_safety_enabled=False)
        policy = treatment_policy(treatment)
        safety_state = ActionSafetyState()
        safety_config = {**config.get("action_safety", {}), "enabled": bool(config.get("action_safety", {}).get("enabled", False))}
        year_info = find_year(config, station, int(year))
        obs, info = env.reset()
        done = False
        step_count = 0
        cumulative_i = 0.0
        cumulative_n = 0.0
        while not done and step_count < int(config.get("runtime", {}).get("max_steps", 260)):
            latest = latest_observation_dict(env, obs, info)
            dap_raw = scalar(latest.get("dap", np.nan))
            dap = int(round(dap_raw)) if not np.isnan(dap_raw) and dap_raw > 0 else step_count + 1
            planned = policy.action_for_dap(dap)
            bounded, clip_note = clip_to_action_space(env, planned)
            safety_result = apply_action_safety(bounded, dap, safety_state, safety_config)
            safe = safety_result.safe_real_action
            normalized = normalize_action(env.formator.action_names, env.formator.action_space_dict, safe)
            obs, reward, terminated, truncated, info = env.step(normalized)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            update_action_safety_state(safety_state, safe, dap)
            real_i = float(safe.get("amir", 0.0))
            real_n = float(safe.get("anfer", 0.0))
            cumulative_i += real_i
            cumulative_n += real_n
            date = pd.Timestamp(year_info["planting_date"]) + pd.Timedelta(days=max(dap - 1, 0))
            totir_raw = scalar(latest.get("totir"))
            tofer_raw = scalar(latest.get("tofer"))
            norm_by_name = {name: float(value) for name, value in zip(env.formator.action_names, np.asarray(normalized).flatten())}
            records.append(
                {
                    "station": station,
                    "year": int(year),
                    "treatment_id": treatment.treatment_id,
                    "treatment_name": treatment.treatment_name,
                    "date": date.strftime("%Y-%m-%d"),
                    "doy": int(date.dayofyear),
                    "dap": scalar(latest.get("dap", dap)),
                    "planned_event_name": policy.event_name_for_dap(dap),
                    "planned_irrigation": float(planned.get("amir", 0.0)),
                    "planned_n": float(planned.get("anfer", 0.0)),
                    "real_action_amir": real_i,
                    "real_action_anfer": real_n,
                    "normalized_action_amir": norm_by_name.get("amir", np.nan),
                    "normalized_action_anfer": norm_by_name.get("anfer", np.nan),
                    "topwt": scalar(latest.get("topwt")),
                    "grnwt": scalar(latest.get("grnwt")),
                    "xlai": scalar(latest.get("xlai")),
                    "totir_raw": totir_raw,
                    "totir": totir_raw if not np.isnan(totir_raw) else cumulative_i,
                    "tofer_raw": tofer_raw,
                    "tofer": tofer_raw if not np.isnan(tofer_raw) else cumulative_n,
                    "swfac": scalar(latest.get("swfac")),
                    "nstres": scalar(latest.get("nstres")),
                    "reward": float(reward),
                    "done": done,
                    "env_clip_note": clip_note,
                    "safety_rule_triggered": safety_result.safety_rule_triggered,
                    "info": json.dumps(info, ensure_ascii=False, default=str),
                }
            )
            step_count += 1
        daily = pd.DataFrame(records)
        daily.to_csv(path, index=False, encoding="utf-8-sig")
        fig_dir = OUTPUT_ROOT / "figures" / "factorial_daily" / station / f"{year}_{treatment.treatment_name}"
        plot_episode(daily, fig_dir)
        status = "ok" if len(daily) and bool(daily["done"].iloc[-1]) else "failed"
        return summarize_daily(daily, station, year, treatment, status, "" if status == "ok" else "episode_not_completed", config)
    except Exception as exc:
        if records:
            pd.DataFrame(records).to_csv(path, index=False, encoding="utf-8-sig")
        return summarize_daily(pd.DataFrame(records), station, year, treatment, "failed", f"{type(exc).__name__}: {exc}", config)
    finally:
        if env is not None:
            try:
                env.close()
            except Exception:
                pass


def add_factorial_metrics(summary: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (station, year), group in summary.groupby(["station", "year"]):
        by_id = {row.treatment_id: row for row in group.itertuples(index=False)}
        t0 = by_id.get("T0")
        t2 = by_id.get("T2")
        for row in group.to_dict("records"):
            y = float(row.get("final_grnwt", np.nan))
            row["yield_gain_vs_T0"] = y - float(getattr(t0, "final_grnwt", np.nan)) if t0 else np.nan
            row["yield_gain_vs_N_only_medium"] = y - float(getattr(t2, "final_grnwt", np.nan)) if t2 else np.nan
            row["yield_gain_from_irrigation_at_same_N"] = np.nan
            row["yield_gain_from_n_at_same_irrigation"] = np.nan
            row["profit_gain_from_irrigation_at_same_N_default"] = np.nan
            row["profit_gain_from_irrigation_at_same_N_low_water_cost"] = np.nan
            if row["treatment_id"] in ["T4", "T5"] and t2 is not None:
                row["yield_gain_from_irrigation_at_same_N"] = y - float(t2.final_grnwt)
                row["profit_gain_from_irrigation_at_same_N_default"] = float(row["profit_default"]) - float(t2.profit_default)
                row["profit_gain_from_irrigation_at_same_N_low_water_cost"] = float(row["profit_low_water_cost"]) - float(t2.profit_low_water_cost)
            if row["treatment_id"] == "T3" and t0 is not None:
                row["yield_gain_from_irrigation_at_same_N"] = y - float(t0.final_grnwt)
                row["profit_gain_from_irrigation_at_same_N_default"] = float(row["profit_default"]) - float(t0.profit_default)
                row["profit_gain_from_irrigation_at_same_N_low_water_cost"] = float(row["profit_low_water_cost"]) - float(t0.profit_low_water_cost)
            if row["treatment_id"] in ["T1", "T2"] and t0 is not None:
                row["yield_gain_from_n_at_same_irrigation"] = y - float(t0.final_grnwt)
            row["water_productivity"] = y / (float(row["total_irrigation"]) / 100.0) if float(row["total_irrigation"]) > 0 else np.nan
            row["n_productivity"] = y / (float(row["total_n"]) / 100.0) if float(row["total_n"]) > 0 else np.nan
            rows.append(row)
    return pd.DataFrame(rows)


def screen_irrigation_response(summary: pd.DataFrame, config: dict) -> pd.DataFrame:
    cfg = config["irrigation_response_screening"]
    rows = []
    for (station, year), group in summary.groupby(["station", "year"]):
        by_id = {row.treatment_id: row for row in group.itertuples(index=False)}
        t2 = by_id.get("T2")
        t4 = by_id.get("T4")
        t5 = by_id.get("T5")
        if t2 is None:
            rows.append({"station": station, "year": int(year), "irrigation_candidate": False, "notes": "T2 missing"})
            continue
        stress_days_t2 = float(getattr(t2, "swfac_stress_days_gt_0p05", 0))
        max_swfac_t2 = float(getattr(t2, "max_swfac", 0))
        has_water_stress = bool(stress_days_t2 > 0 or max_swfac_t2 > float(cfg["stress_threshold"]))
        candidates = [x for x in [t4, t5] if x is not None]
        irrigation_reduces_swfac = any(
            float(getattr(x, "swfac_stress_days_gt_0p05", np.inf)) < stress_days_t2
            or float(getattr(x, "max_swfac", np.inf)) < max_swfac_t2
            for x in candidates
        )
        yield_gains = [float(getattr(x, "final_grnwt", np.nan)) - float(t2.final_grnwt) for x in candidates]
        finite_yield_gains = [gain for gain in yield_gains if np.isfinite(gain)]
        max_yield_gain = max(finite_yield_gains) if finite_yield_gains else np.nan
        best_irrigation_treatment = ""
        if finite_yield_gains:
            best_idx = int(np.nanargmax(yield_gains))
            best_irrigation_treatment = candidates[best_idx].treatment_name
        irrigation_increases_yield = bool(
            np.isfinite(max_yield_gain)
            and (
                max_yield_gain >= float(cfg["yield_gain_abs_threshold"])
                or (float(t2.final_grnwt) > 0 and (float(t2.final_grnwt) + max_yield_gain) > float(t2.final_grnwt) * float(cfg["yield_gain_ratio_threshold"]))
            )
        )
        profit_default_gains = [float(getattr(x, "profit_default", np.nan)) - float(t2.profit_default) for x in candidates]
        profit_low_gains = [float(getattr(x, "profit_low_water_cost", np.nan)) - float(t2.profit_low_water_cost) for x in candidates]
        finite_default_gains = [gain for gain in profit_default_gains if np.isfinite(gain)]
        finite_low_gains = [gain for gain in profit_low_gains if np.isfinite(gain)]
        max_profit_default_gain = max(finite_default_gains) if finite_default_gains else np.nan
        max_profit_low_gain = max(finite_low_gains) if finite_low_gains else np.nan
        irrigation_increases_profit_default = bool(np.isfinite(max_profit_default_gain) and max_profit_default_gain > 0)
        irrigation_increases_profit_low = bool(np.isfinite(max_profit_low_gain) and max_profit_low_gain >= float(cfg["low_water_profit_tolerance"]))
        irrigation_candidate = bool(irrigation_reduces_swfac and irrigation_increases_yield and irrigation_increases_profit_low)
        rows.append(
            {
                "station": station,
                "year": int(year),
                "has_water_stress": has_water_stress,
                "t2_swfac_stress_days_gt_0p05": stress_days_t2,
                "t2_max_swfac": max_swfac_t2,
                "irrigation_reduces_swfac": irrigation_reduces_swfac,
                "irrigation_increases_yield_at_same_N": irrigation_increases_yield,
                "max_yield_gain_from_irrigation_at_same_N": max_yield_gain,
                "best_irrigation_treatment": best_irrigation_treatment,
                "irrigation_increases_profit_default": irrigation_increases_profit_default,
                "max_profit_gain_default": max_profit_default_gain,
                "irrigation_increases_profit_low_water_cost": irrigation_increases_profit_low,
                "max_profit_gain_low_water_cost": max_profit_low_gain,
                "irrigation_candidate": irrigation_candidate,
                "dominant_yield_driver": classify_driver(by_id),
                "notes": "",
            }
        )
    return pd.DataFrame(rows)


def classify_driver(by_id: dict[str, Any]) -> str:
    t0, t2, t5 = by_id.get("T0"), by_id.get("T2"), by_id.get("T5")
    if t0 is None or t2 is None:
        return "unknown"
    n_gain = float(t2.final_grnwt) - float(t0.final_grnwt)
    i_gain = float(t5.final_grnwt) - float(t2.final_grnwt) if t5 is not None else np.nan
    if np.isfinite(i_gain) and i_gain > max(200.0, 0.25 * max(n_gain, 0.0)):
        return "water_and_nitrogen"
    if n_gain > 200:
        return "nitrogen"
    if np.isfinite(i_gain) and i_gain > 200:
        return "water"
    return "weak_response"


def current_site_status(config: dict) -> pd.DataFrame:
    all_mode_path = PROJECT_ROOT / "output_hl" / "all_water_stress_diagnostics" / "all_water_stress_summary_final.csv"
    method_path = PROJECT_ROOT / "Leave_One_experiments" / "constrained_ppo_multiseed_augmented_prior" / "evaluation" / "multiseed_method_comparison.csv"
    augmented_path = PROJECT_ROOT / "Leave_One_experiments" / "expert_dataset_augmentation" / "evaluation" / "augmented_policy_gate_summary.csv"
    all_mode = pd.read_csv(all_mode_path) if all_mode_path.exists() else pd.DataFrame()
    method = pd.read_csv(method_path) if method_path.exists() else pd.DataFrame()
    augmented = pd.read_csv(augmented_path) if augmented_path.exists() else pd.DataFrame()
    rf = augmented[augmented.get("policy_name", "") == "BC_random_forest_regressor_augmented"] if not augmented.empty else pd.DataFrame()
    rows = []
    def series_value(row: pd.Series | None, column: str) -> Any:
        if row is None or column not in row.index:
            return np.nan
        return row[column]

    for station, years in config["observed_years"].items():
        observed = ",".join(str(item["year"]) for item in years)
        short = station.replace("A", "")
        site_rows = all_mode[all_mode["site"].astype(str).isin([short, station])] if not all_mode.empty and "site" in all_mode else pd.DataFrame()
        usable = site_rows[site_rows["trace_source"].astype(str) != "missing_or_stalled"] if not site_rows.empty and "trace_source" in site_rows else pd.DataFrame()
        null = usable[usable["agent"].astype(str) == "null"] if not usable.empty and "agent" in usable else pd.DataFrame()
        base = null.iloc[0] if len(null) else (usable.iloc[0] if len(usable) else None)
        rows.append(
            {
                "station": station,
                "observed_years": observed,
                "PRCP": series_value(base, "PRCP"),
                "ETCP": series_value(base, "ETCP"),
                "PRCP_minus_ETCP": series_value(base, "PRCP_minus_ETCP"),
                "all_mode_swfac_available": bool(base is not None and pd.notna(series_value(base, "max_swfac"))),
                "swfac_stress_days_gt_0p05": series_value(base, "swfac_stress_days_gt_0.05"),
                "max_swfac": series_value(base, "max_swfac"),
                "mean_swfac": series_value(base, "mean_swfac"),
                "nstres_days_gt_0p05": series_value(base, "nstres_days_gt_0.05"),
                "max_nstres": series_value(base, "max_nstres"),
                "mean_nstres": series_value(base, "mean_nstres"),
                "augmented_rf_mean_yield": float(rf["mean_yield"].iloc[0]) if len(rf) and "mean_yield" in rf else np.nan,
                "augmented_rf_mean_irrigation": float(rf["mean_irrigation"].iloc[0]) if len(rf) and "mean_irrigation" in rf else np.nan,
                "augmented_rf_mean_n": float(rf["mean_n"].iloc[0]) if len(rf) and "mean_n" in rf else np.nan,
                "augmented_rf_mean_profit": float(rf["mean_profit"].iloc[0]) if len(rf) and "mean_profit" in rf else np.nan,
                "diagnosis_note": "all_mode_missing_or_stalled" if base is None else "all_mode_available",
            }
        )
    status = pd.DataFrame(rows)
    status.to_csv(OUTPUT_ROOT / "evaluation" / "current_site_status_summary.csv", index=False, encoding="utf-8-sig")
    lines = [
        "# Current Site Status Review",
        "",
        "This stage does not continue PPO and does not enter rainfall-scaling. It first asks whether irrigation has real observed-year value.",
        "",
        "## Why not continue PPO",
        "",
        "- 006_13 showed no 300/450 regression, but fine-tuned PPO did not improve over the augmented RF prior.",
        "- MS0 augmented RF prior replay kept water at 0 and N near 83 kg/ha with best mean profit.",
        "- MS2 increased yield by using the strict 100/200 guardrail and produced negative profit, so it is not a low-input optimum.",
        "",
        "## Why not directly rainfall-scaling",
        "",
        "- Rainfall-scaling is a stress-test scenario, not observed-year evidence.",
        "- The augmented imitation dataset has very few non-zero irrigation rows, so a learned policy can easily become a no-irrigation policy.",
        "- Existing all-mode diagnostics show weak or missing swfac evidence in observed years.",
        "",
        "## Site status summary",
        "",
        df_to_markdown(status),
    ]
    (OUTPUT_ROOT / "evaluation" / "current_site_status_review.md").write_text("\n".join(lines), encoding="utf-8")
    return status


def plot_factorial(summary: pd.DataFrame, screening: pd.DataFrame) -> None:
    fig_dir = OUTPUT_ROOT / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    summary = summary.copy()
    summary["site_year"] = summary["station"] + "_" + summary["year"].astype(str)
    for metric, filename, ylabel in [
        ("final_grnwt", "factorial_yield_by_treatment_site_year.png", "Final grain weight"),
        ("profit_default", "factorial_profit_by_treatment_site_year.png", "Profit default water cost"),
        ("max_swfac", "factorial_swfac_by_treatment_site_year.png", "Max SWFAC"),
    ]:
        pivot = summary.pivot_table(index="site_year", columns="treatment_id", values=metric, aggfunc="mean")
        ax = pivot.plot(kind="bar", figsize=(14, 6))
        ax.set_ylabel(ylabel)
        ax.set_xlabel("")
        ax.grid(axis="y", alpha=0.25)
        fig = ax.get_figure()
        fig.tight_layout()
        fig.savefig(fig_dir / filename, dpi=180)
        plt.close(fig)

    irr = summary[summary["treatment_id"].isin(["T4", "T5"])].copy()
    if not irr.empty:
        pivot = irr.pivot_table(index="site_year", columns="treatment_id", values="yield_gain_from_irrigation_at_same_N", aggfunc="mean")
        ax = pivot.plot(kind="bar", figsize=(14, 5), color=["#4F81BD", "#9BBB59"])
        ax.axhline(0, color="black", linewidth=0.8)
        ax.set_ylabel("Yield gain vs T2")
        ax.set_xlabel("")
        ax.grid(axis="y", alpha=0.25)
        fig = ax.get_figure()
        fig.tight_layout()
        fig.savefig(fig_dir / "yield_gain_from_irrigation_at_same_N.png", dpi=180)
        plt.close(fig)
        pivot = irr.pivot_table(index="site_year", columns="treatment_id", values="profit_gain_from_irrigation_at_same_N_default", aggfunc="mean")
        ax = pivot.plot(kind="bar", figsize=(14, 5), color=["#4F81BD", "#9BBB59"])
        ax.axhline(0, color="black", linewidth=0.8)
        ax.set_ylabel("Profit gain vs T2, default water cost")
        ax.set_xlabel("")
        ax.grid(axis="y", alpha=0.25)
        fig = ax.get_figure()
        fig.tight_layout()
        fig.savefig(fig_dir / "profit_gain_from_irrigation_at_same_N.png", dpi=180)
        plt.close(fig)

    if not screening.empty:
        cols = [
            "has_water_stress",
            "irrigation_reduces_swfac",
            "irrigation_increases_yield_at_same_N",
            "irrigation_increases_profit_default",
            "irrigation_increases_profit_low_water_cost",
            "irrigation_candidate",
        ]
        heat = screening.set_index(screening["station"] + "_" + screening["year"].astype(str))[cols].astype(int)
        fig, ax = plt.subplots(figsize=(10, max(5, len(heat) * 0.35)))
        im = ax.imshow(heat.values, aspect="auto", cmap="Blues", vmin=0, vmax=1)
        ax.set_yticks(range(len(heat.index)))
        ax.set_yticklabels(heat.index)
        ax.set_xticks(range(len(cols)))
        ax.set_xticklabels(cols, rotation=35, ha="right")
        for i in range(len(heat.index)):
            for j in range(len(cols)):
                ax.text(j, i, str(int(heat.values[i, j])), ha="center", va="center", fontsize=8)
        ax.set_title("Irrigation response screening")
        fig.colorbar(im, ax=ax, fraction=0.025, pad=0.02)
        fig.tight_layout()
        fig.savefig(fig_dir / "irrigation_response_screening_heatmap.png", dpi=180)
        plt.close(fig)


def schedule_from_amounts(schedule_id: str, station: str, year: int, irrigation_daps: list[int], irrigation_amounts: tuple[float, ...], nitrogen_daps: list[int], nitrogen_amounts: tuple[float, ...]) -> dict:
    row: dict[str, Any] = {
        "schedule_id": schedule_id,
        "station": station,
        "year": int(year),
        "total_irrigation": float(sum(irrigation_amounts)),
        "total_n": float(sum(nitrogen_amounts)),
    }
    for dap, amount in zip(irrigation_daps, irrigation_amounts):
        row[f"I_DAP_{dap}"] = float(amount)
    for dap, amount in zip(nitrogen_daps, nitrogen_amounts):
        row[f"N_DAP_{dap}"] = float(amount)
    return row


def evaluate_search_schedule(config: dict, row: dict[str, Any], ref: dict[str, Any]) -> dict:
    treatment = Treatment(
        treatment_id=str(row["schedule_id"]),
        treatment_name=str(row["schedule_id"]),
        nitrogen_by_dap={int(k.split("_")[-1]): float(v) for k, v in row.items() if k.startswith("N_DAP_") and float(v) > 0},
        irrigation_by_dap={int(k.split("_")[-1]): float(v) for k, v in row.items() if k.startswith("I_DAP_") and float(v) > 0},
    )
    result = run_fixed_treatment(config, str(row["station"]), int(row["year"]), treatment)
    result["schedule_id"] = row["schedule_id"]
    result["yield_gain_vs_augmented_rf_prior"] = float(result["final_grnwt"]) - float(ref.get("augmented_rf_yield", np.nan))
    result["profit_gain_vs_augmented_rf_prior_default"] = float(result["profit_default"]) - float(ref.get("augmented_rf_profit", np.nan))
    result["profit_gain_vs_augmented_rf_prior_low_water_cost"] = float(result["profit_low_water_cost"]) - float(ref.get("augmented_rf_profit", np.nan))
    result["yield_per_100mm_irrigation"] = float(result["final_grnwt"]) / (float(result["total_irrigation"]) / 100.0) if float(result["total_irrigation"]) > 0 else np.nan
    result["recommendation_type"] = ""
    return result


def run_irrigation_responsive_search(config: dict, screening: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    candidates = screening[screening["irrigation_candidate"].map(truthy)].copy()
    if candidates.empty:
        conclusion = [
            "# No Observed-Year Irrigation Response Conclusion",
            "",
            "No station-year passed irrigation_response_screening. Therefore irrigation-responsive expert search was not executed.",
            "",
            "Current observed-year evidence does not support treating irrigation as the main optimization dimension. Rainfall-scaling should remain a separate artificial stress-test, not the main observed-year conclusion.",
        ]
        (OUTPUT_ROOT / "evaluation" / "no_observed_year_irrigation_response_conclusion.md").write_text("\n".join(conclusion), encoding="utf-8")
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame()

    cfg = config["irrigation_responsive_search"]
    augmented = pd.read_csv(OUTPUT_ROOT.parent / "expert_dataset_augmentation" / "evaluation" / "augmented_policy_gate_summary.csv")
    rf = augmented[augmented["policy_name"] == "BC_random_forest_regressor_augmented"].iloc[0].to_dict()
    all_rows = []
    for item in candidates.itertuples(index=False):
        n_daps = [int(x) for x in cfg["nitrogen_event_daps"]]
        i_daps = [int(x) for x in cfg["irrigation_event_daps"]]
        n_combos = [
            combo
            for combo in itertools.product([float(x) for x in cfg["nitrogen_amounts"]], repeat=len(n_daps))
            if sum(combo) <= float(cfg["total_n_cap"])
        ]
        i_combos = [
            combo
            for combo in itertools.product([float(x) for x in cfg["irrigation_amounts"]], repeat=len(i_daps))
            if 0 < sum(combo) <= float(cfg["total_irrigation_cap"])
        ]
        # Keep the search bounded; prioritize moderate N and irrigation totals.
        n_combos = sorted(n_combos, key=lambda x: (abs(sum(x) - 150), sum(x)))[: int(cfg["top_n_fixed_n_schedules"])]
        pairs = list(itertools.product(i_combos, n_combos))
        pairs = sorted(pairs, key=lambda p: (abs(sum(p[0]) - 80), abs(sum(p[1]) - 150)))[: int(cfg["max_candidates_per_site_year"])]
        for idx, (i_combo, n_combo) in enumerate(pairs, start=1):
            row = schedule_from_amounts(f"{item.station}{int(item.year)}_IRR{idx:04d}", item.station, int(item.year), i_daps, i_combo, n_daps, n_combo)
            all_rows.append(evaluate_search_schedule(config, row, {"augmented_rf_yield": rf["mean_yield"], "augmented_rf_profit": rf["mean_profit"]}))
    search = pd.DataFrame(all_rows)
    if search.empty:
        return search, pd.DataFrame(), pd.DataFrame()
    search["is_pareto"] = False
    for (station, year), group in search.groupby(["station", "year"]):
        idxs = []
        for idx, row in group.iterrows():
            dominated = group[
                (group["final_grnwt"] >= row["final_grnwt"])
                & (group["profit_low_water_cost"] >= row["profit_low_water_cost"])
                & (group["total_irrigation"] <= row["total_irrigation"])
                & (group["total_n"] <= row["total_n"])
                & (
                    (group["final_grnwt"] > row["final_grnwt"])
                    | (group["profit_low_water_cost"] > row["profit_low_water_cost"])
                    | (group["total_irrigation"] < row["total_irrigation"])
                    | (group["total_n"] < row["total_n"])
                )
            ]
            if dominated.empty:
                idxs.append(idx)
        search.loc[idxs, "is_pareto"] = True
    pareto = search[search["is_pareto"]].copy()
    ranking = search.sort_values(["profit_low_water_cost", "final_grnwt"], ascending=[False, False]).copy()
    search.to_csv(OUTPUT_ROOT / "evaluation" / "irrigation_responsive_search_summary.csv", index=False, encoding="utf-8-sig")
    pareto.to_csv(OUTPUT_ROOT / "evaluation" / "irrigation_responsive_pareto.csv", index=False, encoding="utf-8-sig")
    ranking.to_csv(OUTPUT_ROOT / "expert_policy" / "irrigation_responsive_expert_schedule_ranking.csv", index=False, encoding="utf-8-sig")
    valid = ranking[
        (ranking["total_irrigation"] > 0)
        & (ranking["run_status"].astype(str).str.startswith("ok"))
        & (ranking["profit_gain_vs_augmented_rf_prior_low_water_cost"] >= 0)
        & ((ranking["yield_gain_vs_augmented_rf_prior"] >= 0) | (ranking["yield_gain_vs_augmented_rf_prior"] / ranking["final_grnwt"].abs() >= -0.10))
    ].copy()
    if not valid.empty:
        valid.to_csv(OUTPUT_ROOT / "expert_policy" / "irrigation_responsive_expert_library.csv", index=False, encoding="utf-8-sig")
        build_irrigation_imitation_dataset(valid)
    return search, pareto, ranking


def build_irrigation_imitation_dataset(valid: pd.DataFrame) -> None:
    frames = []
    for path in valid["daily_csv_path"].dropna().unique():
        full = PROJECT_ROOT / path
        if full.exists():
            daily = pd.read_csv(full)
            daily["source_schedule_type"] = "irrigation_responsive_search"
            frames.append(daily)
    if frames:
        dataset = pd.concat(frames, ignore_index=True)
        dataset.to_csv(OUTPUT_ROOT / "imitation_dataset" / "imitation_dataset_irrigation_responsive.csv", index=False, encoding="utf-8-sig")


def screening_report(screening: pd.DataFrame) -> None:
    lines = [
        "# Irrigation Response Screening Report",
        "",
        "## Screening table",
        "",
        df_to_markdown(screening),
        "",
        "## Findings",
        "",
        f"- Station-years with water stress: {int(screening['has_water_stress'].sum())}.",
        f"- Station-years where irrigation reduced swfac: {int(screening['irrigation_reduces_swfac'].sum())}.",
        f"- Station-years where irrigation increased yield at same N: {int(screening['irrigation_increases_yield_at_same_N'].sum())}.",
        f"- Station-years where irrigation improved default-cost profit: {int(screening['irrigation_increases_profit_default'].sum())}.",
        f"- Station-years where irrigation improved low-water-cost profit: {int(screening['irrigation_increases_profit_low_water_cost'].sum())}.",
        f"- Irrigation candidates: {int(screening['irrigation_candidate'].sum())}.",
        "",
        "## HLA and FQA notes",
        "",
        "- HLA yield gap is interpreted as nitrogen/action-prior related unless HLA rows pass irrigation_candidate screening.",
        "- FQA 2008 is treated as a possible observed-year water candidate only if it passes stress, yield, and profit screening.",
    ]
    (OUTPUT_ROOT / "evaluation" / "irrigation_response_screening_report.md").write_text("\n".join(lines), encoding="utf-8")


def build_report(summary: pd.DataFrame, screening: pd.DataFrame, search: pd.DataFrame, pareto: pd.DataFrame, ranking: pd.DataFrame) -> None:
    completed = summary[summary["run_status"].astype(str).str.startswith("ok")]
    site_years = completed[["station", "year"]].drop_duplicates().sort_values(["station", "year"])
    driver = screening.groupby("dominant_yield_driver").size().reset_index(name="count") if not screening.empty else pd.DataFrame()
    search_executed = not search.empty
    library_path = OUTPUT_ROOT / "expert_policy" / "irrigation_responsive_expert_library.csv"
    dataset_path = OUTPUT_ROOT / "imitation_dataset" / "imitation_dataset_irrigation_responsive.csv"
    bc_retrain = False
    if dataset_path.exists():
        data = pd.read_csv(dataset_path)
        nonzero_i = int((pd.to_numeric(data.get("real_action_amir", pd.Series(dtype=float)), errors="coerce").fillna(0) > 0).sum())
        unique_i = int(pd.to_numeric(data.get("real_action_amir", pd.Series(dtype=float)), errors="coerce").dropna().round(6).nunique())
        bc_retrain = bool(len(data) and nonzero_i >= 5 and unique_i >= 3)
    lines = [
        "# Water x Nitrogen Factorial Diagnosis Report",
        "",
        "Generated by `src/run_water_nitrogen_factorial_diagnosis.py`.",
        "",
        "## Why this stage",
        "",
        "006_13 showed that constrained PPO did not regress to 300/450, but it also did not outperform the augmented RF prior. Because observed-year irrigation samples are sparse and swfac evidence is weak, this stage diagnoses whether irrigation has real observed-year value before any rainfall-scaling or further PPO work.",
        "",
        "## Completed station-years",
        "",
        df_to_markdown(site_years),
        "",
        "## Dominant yield driver",
        "",
        df_to_markdown(driver),
        "",
        "## Irrigation response screening",
        "",
        df_to_markdown(screening),
        "",
        "## Factorial treatment summary preview",
        "",
        df_to_markdown(summary[["station", "year", "treatment_id", "final_grnwt", "total_irrigation", "total_n", "profit_default", "profit_low_water_cost", "max_swfac", "mean_nstres", "yield_gain_from_irrigation_at_same_N"]], 40),
        "",
        "## Irrigation-responsive expert search",
        "",
        f"- Executed: {search_executed}.",
        f"- Search rows: {len(search)}.",
        f"- Pareto rows: {len(pareto)}.",
        f"- Expert library updated: {library_path.exists()}.",
        f"- Imitation dataset updated: {dataset_path.exists()}.",
        f"- Retrain irrigation-responsive BC: {bc_retrain}.",
        "",
        "## Rainfall-scaling recommendation",
        "",
    ]
    if int(screening["irrigation_candidate"].sum()) > 0:
        lines.append("Observed-year irrigation candidates were found. Rainfall-scaling can be considered later as a stress-test, but not as the immediate main line until irrigation-responsive expert schedules are reviewed.")
    else:
        lines.append("No observed-year station-year passed irrigation screening. Do not enter rainfall-scaling as the main line; keep observed-year work focused on nitrogen management / augmented RF prior and treat rainfall-scaling only as a separate artificial stress-test.")
    DOC_MD.write_text("\n".join(lines), encoding="utf-8")
    report_dir = OUTPUT_ROOT / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    safe_copy(DOC_MD, report_dir / DOC_MD.name)
    build_ppt(summary, screening, search, driver)
    if DOC_PPT.exists():
        safe_copy(DOC_PPT, report_dir / DOC_PPT.name)


def add_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.3), Inches(12.4), Inches(0.6))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(23)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.15), Inches(12.0), Inches(5.8))
    tf = body.text_frame
    tf.clear()
    for idx, item in enumerate(bullets):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        para.text = item
        para.space_after = Pt(7)
        for run in para.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0, 0, 0)


def add_table_slide(prs: Presentation, title: str, df: pd.DataFrame, max_rows: int = 12) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.3), Inches(12.4), Inches(0.6))
    r = title_box.text_frame.paragraphs[0].add_run()
    r.text = title
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(22)
    r.font.bold = True
    rows = min(max_rows, len(df))
    cols = len(df.columns)
    table_shape = slide.shapes.add_table(rows + 1, cols, Inches(0.35), Inches(1.15), Inches(12.65), Inches(5.8))
    table = table_shape.table
    work = df.head(rows).copy()
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(3)
    for j, col in enumerate(work.columns):
        cell = table.cell(0, j)
        cell.text = str(col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(68, 114, 196)
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(8)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(work.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = "" if pd.isna(value) else str(value)
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(226, 235, 247)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(7)
                    run.font.color.rgb = RGBColor(0, 0, 0)


def build_ppt(summary: pd.DataFrame, screening: pd.DataFrame, search: pd.DataFrame, driver: pd.DataFrame) -> None:
    if Presentation is None:
        print("python-pptx is not available; skipped PPT generation.", flush=True)
        return
    prs = Presentation()
    add_slide(
        prs,
        "006_14 Water x Nitrogen Factorial Diagnosis",
        [
            "Purpose: diagnose whether observed-year yield response comes from water, nitrogen, or both before any further PPO or rainfall-scaling.",
            f"Completed fixed-treatment rows: {len(summary)}.",
            f"Irrigation candidates: {int(screening['irrigation_candidate'].sum()) if not screening.empty else 0}.",
            "No ordinary PPO, no unconstrained PPO, and no rainfall-scaling were run in this stage.",
        ],
    )
    add_table_slide(prs, "Dominant Yield Driver", driver if not driver.empty else pd.DataFrame({"dominant_yield_driver": [], "count": []}), 8)
    add_table_slide(
        prs,
        "Irrigation Screening",
        screening[[
            "station",
            "year",
            "has_water_stress",
            "irrigation_reduces_swfac",
            "irrigation_increases_yield_at_same_N",
            "irrigation_increases_profit_default",
            "irrigation_increases_profit_low_water_cost",
            "irrigation_candidate",
            "dominant_yield_driver",
        ]],
        16,
    )
    preview = summary[summary["treatment_id"].isin(["T0", "T2", "T5"])][
        ["station", "year", "treatment_id", "final_grnwt", "total_irrigation", "total_n", "profit_default", "max_swfac", "mean_nstres"]
    ]
    add_table_slide(prs, "Factorial Summary Preview", preview, 18)
    add_slide(
        prs,
        "Conclusion",
        [
            f"Irrigation-responsive search executed: {not search.empty}.",
            "If no candidate passed, observed-year irrigation is not yet a main optimization dimension.",
            "Next step should be irrigation-specific expert search only for screened candidates, or nitrogen-management prior refinement if no candidate exists.",
        ],
    )
    prs.save(DOC_PPT)


def run_factorial(config: dict) -> pd.DataFrame:
    treatments = load_treatments(config)
    rows = []
    for station, years in config["observed_years"].items():
        for item in years:
            for treatment in treatments:
                print(f"[factorial] {station} {item['year']} {treatment.treatment_name}", flush=True)
                rows.append(run_fixed_treatment(config, station, int(item["year"]), treatment))
                pd.DataFrame(rows).to_csv(OUTPUT_ROOT / "evaluation" / "water_nitrogen_factorial_summary_partial.csv", index=False, encoding="utf-8-sig")
    summary = add_factorial_metrics(pd.DataFrame(rows))
    summary.to_csv(OUTPUT_ROOT / "evaluation" / "water_nitrogen_factorial_summary.csv", index=False, encoding="utf-8-sig")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default=str(CONFIG_PATH))
    parser.add_argument("--report-only", action="store_true")
    args = parser.parse_args()
    config = load_yaml(Path(args.config))
    ensure_dirs()
    safe_copy(Path(args.config), OUTPUT_ROOT / "configs" / Path(args.config).name)
    current_site_status(config)
    summary_path = OUTPUT_ROOT / "evaluation" / "water_nitrogen_factorial_summary.csv"
    if args.report_only and summary_path.exists():
        summary = pd.read_csv(summary_path)
    else:
        summary = run_factorial(config)
    screening = screen_irrigation_response(summary, config)
    screening.to_csv(OUTPUT_ROOT / "evaluation" / "irrigation_response_screening.csv", index=False, encoding="utf-8-sig")
    screening_report(screening)
    plot_factorial(summary, screening)
    search, pareto, ranking = run_irrigation_responsive_search(config, screening)
    build_report(summary, screening, search, pareto, ranking)
    print(OUTPUT_ROOT / "evaluation" / "water_nitrogen_factorial_summary.csv")
    print(OUTPUT_ROOT / "evaluation" / "irrigation_response_screening.csv")
    print(DOC_MD)
    print(DOC_PPT)


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        raise
