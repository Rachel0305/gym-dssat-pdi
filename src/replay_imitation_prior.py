from __future__ import annotations

import argparse
import json
import shutil
import traceback
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from constrained_ppo_wrappers import ImitationPenaltyRewardWrapper, PriorActionTable, ResidualPriorActionWrapper
from episode_profit_reward import EpisodeProfitRewardWrapper
from ppo_action_safety import ActionSafetyState, apply_action_safety, normalize_action, update_action_safety_state
from ppo_evaluate import latest_observation_dict, make_env, scalar
from ppo_experiment_plan import find_year
from ppo_plot_results import plot_episode
from ppo_safe_rendering import PROJECT_ROOT, load_yaml


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_prior_replay_debug.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "prior_replay_debug"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_prior_replay_debug_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_prior_replay_debug_report.pptx"


def ensure_dirs() -> None:
    for sub in ["configs", "daily_outputs", "evaluation", "figures", "reports", "models", "tensorboard"]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)
    for station in ["HLA", "SYA", "LCA"]:
        (OUTPUT_ROOT / "daily_outputs" / station).mkdir(parents=True, exist_ok=True)
        (OUTPUT_ROOT / "figures" / station).mkdir(parents=True, exist_ok=True)


def df_to_markdown(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df.empty:
        return ""
    work = df.copy()
    if max_rows is not None:
        work = work.head(max_rows)
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(4)
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    separator = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, separator, *rows])


def profit_score(final_grnwt: float, irrigation: float, nitrogen: float, config: dict) -> float:
    econ = config.get("economics", {})
    return (
        float(econ.get("grain_value_coef", 0.01)) * float(final_grnwt)
        - float(econ.get("water_cost", 0.5)) * float(irrigation)
        - float(econ.get("n_cost", 0.25)) * float(nitrogen)
    )


def safety_config(config: dict, method: dict | None = None) -> dict:
    out = dict(config.get("action_safety", {}))
    out["enabled"] = True
    if method:
        out["season_irrigation_soft_limit"] = float(method.get("guardrail_irrigation", out.get("season_irrigation_soft_limit", 300.0)))
        out["season_n_soft_limit"] = float(method.get("guardrail_n", out.get("season_n_soft_limit", 450.0)))
        out["daily_n_max"] = float(method.get("daily_n_max", out.get("daily_n_max", 150.0)))
    return out


class SimDaySafeActionWrapper:
    def __init__(self, env, config: dict, method: dict | None = None):
        self.env = env
        self.safety_config = safety_config(config, method)
        self.state = ActionSafetyState()
        self.step_count = 0
        self.last_safety_result = None
        self.action_space = env.action_space
        self.observation_space = env.observation_space
        self.metadata = getattr(env, "metadata", {})

    def reset(self, *args, **kwargs):
        self.state = ActionSafetyState()
        self.step_count = 0
        self.last_safety_result = None
        return self.env.reset(*args, **kwargs)

    def step(self, action):
        self.step_count += 1
        from ppo_action_safety import denormalize_action

        raw_real = denormalize_action(self.env.formator.action_names, self.env.formator.action_space_dict, action)
        result = apply_action_safety(raw_real, self.step_count, self.state, self.safety_config)
        safe_norm = normalize_action(self.env.formator.action_names, self.env.formator.action_space_dict, result.safe_real_action)
        obs, reward, terminated, truncated, info = self.env.step(safe_norm)
        update_action_safety_state(self.state, result.safe_real_action, self.step_count)
        self.last_safety_result = result
        return obs, reward, terminated, truncated, info

    def close(self):
        return self.env.close()

    @property
    def unwrapped(self):
        return self.env.unwrapped

    def __getattr__(self, name):
        return getattr(self.env, name)


def make_simday_safe_env(config: dict, station: str, year: int, run_tag: str, evaluation: bool, method: dict | None = None):
    base = make_env(config, station, int(year), int(config.get("seed", 0)), run_tag=run_tag, evaluation=evaluation, action_safety_enabled=False)
    return SimDaySafeActionWrapper(base, config, method)


def action_table(config: dict) -> PriorActionTable:
    return PriorActionTable(PROJECT_ROOT / config["primary_prior"]["action_table"])


def normalized_action_dict(env, normalized_action) -> dict[str, float]:
    arr = np.asarray(normalized_action).flatten()
    return {name: float(value) for name, value in zip(env.formator.action_names, arr)}


def evaluate_prior_replay(config: dict, station: str, eval_year: int, suffix: str = "replayed_bc_two_stage") -> dict[str, Any]:
    prior = action_table(config)
    env = make_simday_safe_env(config, station, eval_year, f"{suffix}_{station}_{eval_year}", evaluation=True)
    records: list[dict[str, Any]] = []
    cumulative_irrigation = 0.0
    cumulative_n = 0.0
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        year_info = find_year(config, station, int(eval_year))
        while not done and step_count < int(config["runtime"].get("max_steps", 260)):
            sim_day = step_count + 1
            prior_action = prior.action(station, sim_day)
            action = normalize_action(env.formator.action_names, env.formator.action_space_dict, prior_action)
            obs, reward, terminated, truncated, info = env.step(action)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            result = env.last_safety_result
            real_amir = float(result.safe_real_action.get("amir", 0.0))
            real_anfer = float(result.safe_real_action.get("anfer", 0.0))
            cumulative_irrigation += real_amir
            cumulative_n += real_anfer
            date = pd.Timestamp(year_info["planting_date"]) + pd.Timedelta(days=step_count)
            totir_raw = scalar(latest.get("totir"))
            tofer_raw = scalar(latest.get("tofer"))
            norm = normalized_action_dict(env, action)
            records.append(
                {
                    "station": station,
                    "policy_name": suffix,
                    "eval_year": int(eval_year),
                    "date": date.strftime("%Y-%m-%d"),
                    "year": int(date.year),
                    "doy": int(date.dayofyear),
                    "sim_day": sim_day,
                    "dap": scalar(latest.get("dap", sim_day)),
                    "topwt": scalar(latest.get("topwt")),
                    "grnwt": scalar(latest.get("grnwt")),
                    "xlai": scalar(latest.get("xlai")),
                    "totir_raw": totir_raw,
                    "totir": totir_raw if not np.isnan(totir_raw) else cumulative_irrigation,
                    "tofer_raw": tofer_raw,
                    "tofer": tofer_raw if not np.isnan(tofer_raw) else cumulative_n,
                    "swfac": scalar(latest.get("swfac")),
                    "nstres": scalar(latest.get("nstres")),
                    "reward": float(reward),
                    "real_action_amir": real_amir,
                    "real_action_anfer": real_anfer,
                    "normalized_action_amir": norm.get("amir", np.nan),
                    "normalized_action_anfer": norm.get("anfer", np.nan),
                    "bc_prior_amir": float(prior_action["amir"]),
                    "bc_prior_anfer": float(prior_action["anfer"]),
                    "raw_real_action_amir": float(result.raw_real_action.get("amir", 0.0)),
                    "raw_real_action_anfer": float(result.raw_real_action.get("anfer", 0.0)),
                    "safety_rule_triggered": result.safety_rule_triggered,
                    "done": done,
                    "info": json.dumps(info, ensure_ascii=False, default=str),
                }
            )
            step_count += 1
    finally:
        try:
            env.close()
        except Exception:
            pass
    daily = pd.DataFrame(records)
    daily_dir = OUTPUT_ROOT / "daily_outputs" / station
    daily_dir.mkdir(parents=True, exist_ok=True)
    daily_csv = daily_dir / f"{station}_{suffix}_eval{eval_year}_daily.csv"
    daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
    fig_dir = OUTPUT_ROOT / "figures" / station / suffix / f"eval_{eval_year}"
    plot_episode(daily, fig_dir)
    episode_completed = bool(records and records[-1]["done"])
    total_i = float(daily["real_action_amir"].sum()) if len(daily) else 0.0
    total_n = float(daily["real_action_anfer"].sum()) if len(daily) else 0.0
    final_y = float(daily["grnwt"].iloc[-1]) if len(daily) else np.nan
    return {
        "station": station,
        "eval_year": int(eval_year),
        "run_status": "ok" if episode_completed else "failed",
        "episode_completed": episode_completed,
        "episode_length": int(len(daily)),
        "final_grnwt": final_y,
        "total_irrigation": total_i,
        "total_n_fertilizer": total_n,
        "mean_swfac": float(pd.to_numeric(daily["swfac"], errors="coerce").mean()) if len(daily) else np.nan,
        "mean_nstres": float(pd.to_numeric(daily["nstres"], errors="coerce").mean()) if len(daily) else np.nan,
        "profit_score": profit_score(final_y, total_i, total_n, config) if not np.isnan(final_y) else np.nan,
        "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
        "notes": "pure_simday_safe_action_table_replay",
    }


def add_00609_deltas(summary: pd.DataFrame) -> pd.DataFrame:
    old = pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "evaluation" / "imitation_policy_dssat_evaluation_summary.csv")
    old = old[old["policy_name"].eq("BC_two_stage_classifier_regressor")].copy()
    rows = []
    for _, row in summary.iterrows():
        match = old[(old["station"].eq(row["station"])) & (old["eval_year"].astype(int).eq(int(row["eval_year"])))]
        item = dict(row)
        if len(match):
            ref = match.iloc[0]
            item["yield_loss_vs_00609_bc_two_stage"] = (float(ref["final_grnwt"]) - float(row["final_grnwt"])) / float(ref["final_grnwt"])
            item["n_diff_vs_00609_bc_two_stage"] = float(row["total_n_fertilizer"]) - float(ref["total_n_fertilizer"])
            item["irrigation_diff_vs_00609_bc_two_stage"] = float(row["total_irrigation"]) - float(ref["total_irrigation"])
        rows.append(item)
    return pd.DataFrame(rows)


def write_mismatch_review() -> None:
    text = """# Prior replay mismatch review

Generated at: 2026-06-06

Root cause summary:

1. 006_09 BC_two_stage evaluation used exported action tables from the learned two-stage model and a custom evaluator that applies action safety with `sim_day`.
2. 006_10 FT0 replay used `SafeActionWrapper`, which reads the environment `dap` before each step. In these DSSAT runs the first several rows have `dap = 0`, so day-1 fertilizer events were rejected by `anfer_dap_range`.
3. 006_10 also changed `daily_n_max` to 80 kg/ha. The BC_two_stage action table contains events above 80 kg/ha: HLA 132.75, SYA 138.50, LCA 93.75 and 86.75. These were clipped or removed.
4. Therefore FT0 was not a pure replay of 006_09. It mixed prior replay with a different safety clock and a stricter daily N cap.

Checklist:

| item | 006_09 BC_two_stage | 006_10 FT0 | mismatch |
| --- | --- | --- | --- |
| station/eval years | HLA/SYA/LCA, 10 evals | HLA/SYA/LCA, 10 evals | no |
| prior action source | exported two-stage action table | same action table | no |
| live sklearn model/scaler | trained on Windows, exported to action table for DSSAT container replay | not used | no practical mismatch |
| feature columns | used during action table export | action table replay | no during replay |
| action unit | real kg/ha and mm | real kg/ha and mm before safety | no |
| safety day variable | sim_day | environment dap | yes |
| daily N max | 150 kg/ha in 006_09 replay config | 80 kg/ha | yes |
| residual/guardrail affecting FT0 | none | guardrail/safety still active | yes, through safety cap |

The fix is to use a pure replay evaluator and a sim-day-based safety wrapper for the fixed constrained PPO interface.
"""
    (OUTPUT_ROOT / "evaluation" / "prior_replay_mismatch_review.md").write_text(text, encoding="utf-8")


def make_alignment() -> pd.DataFrame:
    rows = []
    for station in ["HLA", "SYA", "LCA"]:
        for old_rel in sorted((PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "daily_outputs" / station).glob(f"{station}_BC_two_stage_classifier_regressor_eval*_daily.csv")):
            year = int(old_rel.stem.split("_eval")[-1].split("_")[0])
            ft0_path = PROJECT_ROOT / "Leave_One_experiments" / "constrained_ppo_finetuning" / "daily_outputs" / station / f"FT0_BC_two_stage_replay_{station}_train{ {'HLA':2011,'SYA':2012,'LCA':2010}[station] }_seed0_eval{year}_daily.csv"
            if not ft0_path.exists():
                continue
            bc = pd.read_csv(old_rel)
            ft0 = pd.read_csv(ft0_path)
            merged = bc[["station", "eval_year", "date", "sim_day", "dap", "real_action_amir", "real_action_anfer"]].merge(
                ft0[["sim_day", "real_action_amir", "real_action_anfer", "bc_prior_amir", "bc_prior_anfer", "safety_rule_triggered"]],
                on="sim_day",
                how="outer",
                suffixes=("_00609", "_00610"),
            )
            for _, row in merged.iterrows():
                rows.append(
                    {
                        "station": station,
                        "eval_year": year,
                        "date": row.get("date"),
                        "dap": row.get("dap"),
                        "sim_day": row.get("sim_day"),
                        "bc_two_stage_irrigation_00609": row.get("real_action_amir_00609", 0.0),
                        "bc_two_stage_n_00609": row.get("real_action_anfer_00609", 0.0),
                        "ft0_irrigation_00610": row.get("real_action_amir_00610", 0.0),
                        "ft0_n_00610": row.get("real_action_anfer_00610", 0.0),
                        "expert_schedule_irrigation": row.get("bc_prior_amir", 0.0),
                        "expert_schedule_n": row.get("bc_prior_anfer", 0.0),
                        "irrigation_diff": row.get("real_action_amir_00610", 0.0) - row.get("real_action_amir_00609", 0.0),
                        "n_diff": row.get("real_action_anfer_00610", 0.0) - row.get("real_action_anfer_00609", 0.0),
                        "is_event_day": bool((row.get("real_action_anfer_00609", 0.0) or 0) > 0 or (row.get("bc_prior_anfer", 0.0) or 0) > 0),
                        "notes": row.get("safety_rule_triggered", ""),
                    }
                )
    align = pd.DataFrame(rows)
    align.to_csv(OUTPUT_ROOT / "evaluation" / "bc_two_stage_vs_ft0_daily_action_alignment.csv", index=False, encoding="utf-8-sig")
    return align


def ppo_kwargs(config: dict) -> dict:
    allowed = ["learning_rate", "gamma", "n_steps", "batch_size", "ent_coef", "clip_range"]
    return {k: v for k, v in config.get("ppo", {}).items() if k in allowed and v is not None}


def build_fixed_train_env(config: dict, station: str, year: int, method: dict):
    base = make_simday_safe_env(config, station, year, f"{method['name']}_{station}_{year}_train", evaluation=False, method=method)
    profit = EpisodeProfitRewardWrapper(base, "P2_terminal_profit_medium_cost")
    prior = action_table(config)
    if method["kind"] == "residual":
        return ResidualPriorActionWrapper(profit, prior, station, float(method.get("irrigation_residual_range", 20.0)), float(method.get("n_residual_range", 30.0)))
    if method["kind"] == "imitation_penalty":
        return ImitationPenaltyRewardWrapper(profit, prior, station, float(method.get("lambda_bc", 0.5)))
    return profit


def latest_wrapper_info(env) -> dict:
    out: dict[str, Any] = {}
    for attr in ["last_bc_info", "last_residual_info", "last_profit_info"]:
        value = getattr(env, attr, None)
        if isinstance(value, dict):
            out.update(value)
    safety_result = getattr(env, "last_safety_result", None)
    if safety_result is not None:
        out["safety_rule_triggered"] = safety_result.safety_rule_triggered
        out["real_action_amir"] = float(safety_result.safe_real_action.get("amir", 0.0))
        out["real_action_anfer"] = float(safety_result.safe_real_action.get("anfer", 0.0))
    return out


def evaluate_fixed_policy(config: dict, method: dict, station: str, train_year: int, eval_year: int, model, model_path: Path | None) -> dict:
    if method["kind"] == "prior_replay":
        row = evaluate_prior_replay(config, station, eval_year, suffix=method["name"])
        row.update({"finetune_method": method["name"], "train_year": train_year, "seed": int(config.get("seed", 0)), "model_path": ""})
        return row
    env = build_fixed_train_env(config, station, eval_year, method)
    prior = action_table(config)
    records = []
    cumulative_i = 0.0
    cumulative_n = 0.0
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        year_info = find_year(config, station, int(eval_year))
        while not done and step_count < int(config["runtime"].get("max_steps", 260)):
            sim_day = step_count + 1
            action, _ = model.predict(obs, deterministic=True)
            action = np.asarray(action, dtype=np.float32).flatten()
            obs, reward, terminated, truncated, info = env.step(action)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            wrapper = latest_wrapper_info(env)
            real_i = float(wrapper.get("real_action_amir", 0.0))
            real_n = float(wrapper.get("real_action_anfer", 0.0))
            cumulative_i += real_i
            cumulative_n += real_n
            prior_action = prior.action(station, sim_day)
            date = pd.Timestamp(year_info["planting_date"]) + pd.Timedelta(days=step_count)
            records.append(
                {
                    "station": station,
                    "finetune_method": method["name"],
                    "train_year": int(train_year),
                    "eval_year": int(eval_year),
                    "seed": int(config.get("seed", 0)),
                    "date": date.strftime("%Y-%m-%d"),
                    "sim_day": sim_day,
                    "dap": scalar(latest.get("dap", sim_day)),
                    "topwt": scalar(latest.get("topwt")),
                    "grnwt": scalar(latest.get("grnwt")),
                    "xlai": scalar(latest.get("xlai")),
                    "swfac": scalar(latest.get("swfac")),
                    "nstres": scalar(latest.get("nstres")),
                    "reward": float(reward),
                    "real_action_amir": real_i,
                    "real_action_anfer": real_n,
                    "bc_prior_amir": float(prior_action["amir"]),
                    "bc_prior_anfer": float(prior_action["anfer"]),
                    "bc_action_deviation": float((((real_i - prior_action["amir"]) / 100.0) ** 2 + ((real_n - prior_action["anfer"]) / 150.0) ** 2) ** 0.5),
                    "done": done,
                    "info": json.dumps(info, ensure_ascii=False, default=str),
                    **wrapper,
                }
            )
            step_count += 1
    finally:
        try:
            env.close()
        except Exception:
            pass
    daily = pd.DataFrame(records)
    daily_dir = OUTPUT_ROOT / "daily_outputs" / station
    daily_csv = daily_dir / f"{station}_{method['name']}_eval{eval_year}_daily.csv"
    daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
    fig_dir = OUTPUT_ROOT / "figures" / station / method["name"] / f"eval_{eval_year}"
    plot_episode(daily, fig_dir)
    done = bool(records and records[-1]["done"])
    total_i = float(daily["real_action_amir"].sum()) if len(daily) else 0.0
    total_n = float(daily["real_action_anfer"].sum()) if len(daily) else 0.0
    final_y = float(daily["grnwt"].iloc[-1]) if len(daily) else np.nan
    return {
        "station": station,
        "eval_year": int(eval_year),
        "finetune_method": method["name"],
        "train_year": int(train_year),
        "seed": int(config.get("seed", 0)),
        "model_path": str(model_path.relative_to(PROJECT_ROOT)) if model_path else "",
        "run_status": "ok" if done else "failed",
        "episode_completed": done,
        "episode_length": int(len(daily)),
        "final_grnwt": final_y,
        "total_irrigation": total_i,
        "total_n_fertilizer": total_n,
        "mean_swfac": float(pd.to_numeric(daily["swfac"], errors="coerce").mean()) if len(daily) else np.nan,
        "mean_nstres": float(pd.to_numeric(daily["nstres"], errors="coerce").mean()) if len(daily) else np.nan,
        "profit_score": profit_score(final_y, total_i, total_n, config) if not np.isnan(final_y) else np.nan,
        "bc_action_deviation_mean": float(pd.to_numeric(daily["bc_action_deviation"], errors="coerce").mean()) if len(daily) else np.nan,
        "bc_action_deviation_max": float(pd.to_numeric(daily["bc_action_deviation"], errors="coerce").max()) if len(daily) else np.nan,
        "irrigation_saturation_ratio_300_450": total_i / 300.0,
        "n_saturation_ratio_300_450": total_n / 450.0,
        "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
        "notes": "fixed_simday_safety_constrained_eval",
    }


def run_fixed_training(config: dict, method: dict, station: str, train_year: int):
    from stable_baselines3 import PPO

    env = build_fixed_train_env(config, station, train_year, method)
    model_dir = OUTPUT_ROOT / "models" / station
    model_dir.mkdir(parents=True, exist_ok=True)
    model_path = model_dir / f"{method['name']}_{station}_train{train_year}_seed{config.get('seed',0)}"
    try:
        model = PPO("MlpPolicy", env, verbose=0, seed=int(config.get("seed", 0)), tensorboard_log=str(OUTPUT_ROOT / "tensorboard" / station), **ppo_kwargs(config))
        model.learn(total_timesteps=int(config["runtime"].get("total_timesteps", 5000)), progress_bar=False)
        model.save(str(model_path))
    finally:
        try:
            env.close()
        except Exception:
            pass
    return model, model_path.with_suffix(".zip")


def add_fixed_deltas(config: dict, summary: pd.DataFrame) -> pd.DataFrame:
    ppo_y = float(config["ppo_reference"]["mean_yield"])
    bc = pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "evaluation" / "imitation_policy_dssat_evaluation_summary.csv")
    bc = bc[bc["policy_name"].eq("BC_two_stage_classifier_regressor")]
    rows = []
    for _, row in summary.iterrows():
        item = dict(row)
        match = bc[(bc["station"].eq(row["station"])) & (bc["eval_year"].astype(int).eq(int(row["eval_year"])))]
        if len(match):
            ref = match.iloc[0]
            item["yield_loss_vs_00609_bc_two_stage"] = (float(ref["final_grnwt"]) - float(row["final_grnwt"])) / float(ref["final_grnwt"])
            item["n_diff_vs_00609_bc_two_stage"] = float(row["total_n_fertilizer"]) - float(ref["total_n_fertilizer"])
        item["yield_loss_vs_ppo_baseline"] = (ppo_y - float(row["final_grnwt"])) / ppo_y if not np.isnan(row["final_grnwt"]) else np.nan
        item["input_reduction_vs_ppo_baseline"] = 1.0 - ((float(row["total_irrigation"]) / 300.0 + float(row["total_n_fertilizer"]) / 450.0) / 2.0)
        rows.append(item)
    return pd.DataFrame(rows)


def run_fixed_constrained(config: dict) -> pd.DataFrame:
    rows = []
    training = []
    for method in config["fixed_methods"]:
        for station, site in config["sites"].items():
            train_year = int(site["train_year"])
            model = None
            model_path = None
            if method.get("train", False):
                print(f"[00611] train fixed {method['name']} {station}", flush=True)
                try:
                    model, model_path = run_fixed_training(config, method, station, train_year)
                    training.append({"station": station, "train_year": train_year, "finetune_method": method["name"], "run_status": "ok", "model_path": str(model_path.relative_to(PROJECT_ROOT))})
                except Exception as exc:
                    training.append({"station": station, "train_year": train_year, "finetune_method": method["name"], "run_status": "failed", "error_message": f"{type(exc).__name__}: {exc}", "traceback": traceback.format_exc()[-2000:]})
                    continue
            for eval_year in site["eval_years"]:
                print(f"[00611] eval fixed {method['name']} {station} {eval_year}", flush=True)
                try:
                    rows.append(evaluate_fixed_policy(config, method, station, train_year, int(eval_year), model, model_path))
                except Exception as exc:
                    rows.append({"station": station, "train_year": train_year, "eval_year": int(eval_year), "finetune_method": method["name"], "run_status": "failed", "episode_completed": False, "error_message": f"{type(exc).__name__}: {exc}", "traceback": traceback.format_exc()[-2000:]})
    pd.DataFrame(training).to_csv(OUTPUT_ROOT / "evaluation" / "fixed_constrained_ppo_training_summary.csv", index=False, encoding="utf-8-sig")
    out = add_fixed_deltas(config, pd.DataFrame(rows))
    out.to_csv(OUTPUT_ROOT / "evaluation" / "constrained_ppo_fixed_interface_summary.csv", index=False, encoding="utf-8-sig")
    return out


def make_figures(replay: pd.DataFrame, fixed: pd.DataFrame, align: pd.DataFrame) -> None:
    fig_dir = OUTPUT_ROOT / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    old = pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "evaluation" / "imitation_policy_dssat_evaluation_summary.csv")
    old = old[old["policy_name"].eq("BC_two_stage_classifier_regressor")]
    old_ft0 = pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "constrained_ppo_finetuning" / "evaluation" / "constrained_ppo_evaluation_summary.csv")
    old_ft0 = old_ft0[old_ft0["finetune_method"].eq("FT0_BC_two_stage_replay")]
    comp = pd.concat(
        [
            old.assign(source="00609_BC_two_stage").rename(columns={"eval_year": "year"}),
            old_ft0.assign(source="00610_FT0").rename(columns={"eval_year": "year"}),
            replay.assign(source="00611_fixed_replay").rename(columns={"eval_year": "year"}),
        ],
        ignore_index=True,
        sort=False,
    )
    plt.figure(figsize=(9, 5))
    for source, part in comp.groupby("source"):
        plt.scatter(part["total_n_fertilizer"], part["final_grnwt"], label=source)
    plt.xlabel("Total N")
    plt.ylabel("Yield")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_dir / "prior_replay_00609_vs_00610_total_inputs.png", dpi=180)
    plt.close()

    events = align[align["is_event_day"]].copy()
    plt.figure(figsize=(9, 5))
    plt.bar(np.arange(len(events)), events["bc_two_stage_n_00609"], alpha=0.6, label="00609")
    plt.bar(np.arange(len(events)), events["ft0_n_00610"], alpha=0.6, label="00610 FT0")
    plt.xticks(np.arange(len(events)), events["station"].astype(str) + events["eval_year"].astype(str), rotation=45, ha="right")
    plt.ylabel("N action")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_dir / "prior_replay_daily_action_alignment.png", dpi=180)
    plt.close()

    method = fixed[fixed["run_status"].eq("ok")].groupby("finetune_method").agg(mean_yield=("final_grnwt", "mean"), mean_profit=("profit_score", "mean"), mean_i=("total_irrigation", "mean"), mean_n=("total_n_fertilizer", "mean"), mean_loss_ppo=("yield_loss_vs_ppo_baseline", "mean")).reset_index()
    plt.figure(figsize=(9, 5))
    plt.scatter(method["mean_n"], method["mean_yield"])
    for _, row in method.iterrows():
        plt.text(row["mean_n"], row["mean_yield"], row["finetune_method"], fontsize=8)
    plt.xlabel("Mean N")
    plt.ylabel("Mean yield")
    plt.tight_layout()
    plt.savefig(fig_dir / "fixed_constrained_ppo_yield_vs_input.png", dpi=180)
    plt.close()

    plt.figure(figsize=(9, 5))
    plt.barh(method["finetune_method"], method["mean_profit"], color="#4F7ECF")
    plt.xlabel("Mean profit")
    plt.tight_layout()
    plt.savefig(fig_dir / "fixed_constrained_ppo_profit_comparison.png", dpi=180)
    plt.close()

    plt.figure(figsize=(9, 5))
    if "bc_action_deviation_mean" in fixed:
        plt.hist(pd.to_numeric(fixed["bc_action_deviation_mean"], errors="coerce").dropna(), bins=20, color="#4F7ECF")
    plt.xlabel("BC deviation")
    plt.tight_layout()
    plt.savefig(fig_dir / "fixed_constrained_ppo_bc_deviation.png", dpi=180)
    plt.close()

    plt.figure(figsize=(9, 5))
    plt.scatter(replay["total_n_fertilizer"], replay["final_grnwt"], label="00611 replay")
    plt.scatter(old["total_n_fertilizer"], old["final_grnwt"], label="00609")
    plt.xlabel("Total N")
    plt.ylabel("Yield")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_dir / "prior_replay_yield_comparison.png", dpi=180)
    plt.close()


def write_report(config: dict, replay: pd.DataFrame, fixed: pd.DataFrame) -> None:
    replay_summary = replay.agg(mean_yield=("final_grnwt", "mean"), mean_n=("total_n_fertilizer", "mean"), mean_i=("total_irrigation", "mean"), mean_yield_loss=("yield_loss_vs_00609_bc_two_stage", "mean"), mean_n_diff=("n_diff_vs_00609_bc_two_stage", "mean")).reset_index()
    fixed_summary = fixed[fixed["run_status"].eq("ok")].groupby("finetune_method").agg(eval_count=("eval_year", "count"), mean_yield=("final_grnwt", "mean"), mean_profit=("profit_score", "mean"), mean_i=("total_irrigation", "mean"), mean_n=("total_n_fertilizer", "mean"), mean_loss_ppo=("yield_loss_vs_ppo_baseline", "mean"), mean_loss_bc=("yield_loss_vs_00609_bc_two_stage", "mean")).reset_index()
    bc_profit = float(pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "evaluation" / "imitation_policy_dssat_evaluation_summary.csv").query("policy_name == 'BC_two_stage_classifier_regressor'")["profit_score"].mean())
    fixed_summary["profit_gap_vs_bc_two_stage"] = fixed_summary["mean_profit"] - bc_profit
    passed = fixed_summary[(fixed_summary["eval_count"] >= 10) & (fixed_summary["mean_i"] <= 100) & (fixed_summary["mean_n"] <= 200) & (fixed_summary["mean_loss_ppo"] <= 0.15) & (fixed_summary["mean_profit"] >= 0.9 * bc_profit)]
    text = f"""# Prior replay debug report

Generated at: 2026-06-06

## Mismatch cause

006_10 FT0 did not reproduce 006_09 because it used a different action safety interface:

1. `SafeActionWrapper` read environment `dap`, which is 0 on early rows; first-day N events were rejected by `anfer_dap_range`.
2. 006_10 used `daily_n_max = 80`, clipping HLA/LCA BC prior N events above 80 kg/ha.
3. 006_09 replay used exported BC action tables with sim-day action safety and daily N max 150.

## Pure replay result

{df_to_markdown(replay_summary)}

## Fixed constrained PPO result

{df_to_markdown(fixed_summary)}

## Recommendation

{df_to_markdown(passed) if len(passed) else 'No fixed fine-tuned PPO method passed all gates. The corrected FT0 replay should be treated as the reliable prior baseline; next work should improve expert dataset augmentation or cost/reward calibration before constrained PPO multi-seed.'}
"""
    DOC_MD.write_text(text, encoding="utf-8")
    shutil.copy2(DOC_MD, OUTPUT_ROOT / "reports" / DOC_MD.name)

    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        return

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def title(slide, value):
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.45))
        p = box.text_frame.paragraphs[0]
        p.text = value
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0)

    def bullets(slide, lines):
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12), Inches(5.8))
        tf = box.text_frame
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(15)
            p.font.color.rgb = RGBColor(0, 0, 0)

    def table(slide, df):
        data = df.head(8).copy()
        shp = slide.shapes.add_table(len(data) + 1, len(data.columns), Inches(0.35), Inches(1.0), Inches(12.6), Inches(5.8))
        tbl = shp.table
        for j, col in enumerate(data.columns):
            cell = tbl.cell(0, j)
            cell.text = str(col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(79, 126, 207)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(8)
                p.font.bold = True
                p.font.color.rgb = RGBColor(255, 255, 255)
                p.alignment = PP_ALIGN.CENTER
        for i, (_, row) in enumerate(data.iterrows(), start=1):
            for j, col in enumerate(data.columns):
                value = row[col]
                if isinstance(value, float):
                    value = round(value, 4)
                cell = tbl.cell(i, j)
                cell.text = "" if pd.isna(value) else str(value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(235, 240, 250) if i % 2 == 0 else RGBColor(255, 255, 255)
                for p in cell.text_frame.paragraphs:
                    p.font.name = "Microsoft YaHei"
                    p.font.size = Pt(7)
                    p.font.color.rgb = RGBColor(0, 0, 0)

    slide = prs.slides.add_slide(blank)
    title(slide, "006_11 Prior replay debug")
    bullets(slide, ["问题：006_10 FT0 未复现 006_09 BC two-stage。", "原因：安全 wrapper 使用 env dap=0，并且 daily_n_max 从 150 改成 80。", "修复：新增 sim-day safety 的纯 replay evaluator。"])
    slide = prs.slides.add_slide(blank)
    title(slide, "Pure replay reproduction")
    table(slide, replay_summary.round(4))
    slide = prs.slides.add_slide(blank)
    title(slide, "Fixed constrained PPO")
    table(slide, fixed_summary.round(4))
    slide = prs.slides.add_slide(blank)
    title(slide, "Next step")
    bullets(slide, ["若 fixed FT2/FT3 不优于 fixed FT0，不进入 multi-seed。", "优先做 expert dataset augmentation 或重新校准成本/奖励。", "暂不进入 rainfall-scaling budget scenario。"])
    prs.save(DOC_PPT)
    shutil.copy2(DOC_PPT, OUTPUT_ROOT / "reports" / DOC_PPT.name)


def main() -> None:
    parser = argparse.ArgumentParser(description="Debug BC prior replay consistency before constrained PPO fine-tuning.")
    parser.add_argument(
        "--report-only",
        action="store_true",
        help="Rebuild figures, Markdown, and PPT from existing 006_11 CSV outputs without running DSSAT or PPO.",
    )
    args = parser.parse_args()
    ensure_dirs()
    config = load_yaml(CONFIG_PATH)
    shutil.copy2(CONFIG_PATH, OUTPUT_ROOT / "configs" / CONFIG_PATH.name)
    if args.report_only:
        replay = pd.read_csv(OUTPUT_ROOT / "evaluation" / "replayed_bc_two_stage_evaluation_summary.csv")
        fixed = pd.read_csv(OUTPUT_ROOT / "evaluation" / "constrained_ppo_fixed_interface_summary.csv")
        align = pd.read_csv(OUTPUT_ROOT / "evaluation" / "bc_two_stage_vs_ft0_daily_action_alignment.csv")
        make_figures(replay, fixed, align)
        write_report(config, replay, fixed)
        return
    write_mismatch_review()
    align = make_alignment()
    replay_rows = []
    for station, site in config["sites"].items():
        for year in site["eval_years"]:
            print(f"[00611] replay {station} {year}", flush=True)
            replay_rows.append(evaluate_prior_replay(config, station, int(year)))
    replay = add_00609_deltas(pd.DataFrame(replay_rows))
    replay.to_csv(OUTPUT_ROOT / "evaluation" / "replayed_bc_two_stage_evaluation_summary.csv", index=False, encoding="utf-8-sig")
    replay_pass = (
        (replay["run_status"].eq("ok").all())
        and abs(float(replay["total_irrigation"].mean())) <= 1e-6
        and abs(float(replay["n_diff_vs_00609_bc_two_stage"].mean())) / max(1.0, float(pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "evaluation" / "imitation_policy_dssat_evaluation_summary.csv").query("policy_name == 'BC_two_stage_classifier_regressor'")["total_n_fertilizer"].mean())) <= 0.05
        and abs(float(replay["yield_loss_vs_00609_bc_two_stage"].mean())) <= 0.05
    )
    if not replay_pass:
        blocker = "# Prior replay blocker report\n\nPure replay did not reproduce 006_09 within the 5% threshold. Do not run fixed PPO fine-tuning.\n"
        (OUTPUT_ROOT / "evaluation" / "prior_replay_blocker_report.md").write_text(blocker, encoding="utf-8")
        fixed = pd.DataFrame()
    else:
        fixed = run_fixed_constrained(config)
    make_figures(replay, fixed if len(fixed) else pd.DataFrame(), align)
    write_report(config, replay, fixed if len(fixed) else pd.DataFrame())


if __name__ == "__main__":
    main()
