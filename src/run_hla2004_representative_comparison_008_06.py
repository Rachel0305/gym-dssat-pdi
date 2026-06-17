from __future__ import annotations

from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None
    RGBColor = None
    Inches = None
    Pt = None

import run_all_year_direct_action_safe_ppo as all_year
import run_management_scenario_comparison as msc
from ppo_safe_rendering import PROJECT_ROOT


STATION = "HLA"
YEAR = 2004
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "representative_management_comparison_008_06"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-13_008_06_hla2004_representative_management_comparison_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "009_008_06_hla2004_representative_management_comparison.pptx"
ALL_YEAR_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "all_year_weather_calibration_validation"
SCENARIO_POOL = ALL_YEAR_ROOT / "scenario_pool" / "all_year_weather_scenario_pool.csv"
DIRECT_CONFIG = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_all_year_direct_action_safe_ppo.yaml"

REUSED_SCENARIOS = {
    "null_zero": "T0_null_zero",
    "n_only_medium": "T1_N_only_medium",
    "fixed_I60_N150": "T2_N_medium_I_low",
    "fixed_I120_N150": "T3_N_medium_I_mid",
}

RUN_SCENARIOS = {
    "expert_reference_recorded": "expert_reference",
    "ppo_replay_transfer": "ppo_replay_transfer",
    "dssat_auto_attempt": "dssat_automatic",
}

SCENARIO_NOTES = {
    "null_zero": "No nitrogen and no irrigation; strict lower-bound baseline.",
    "n_only_medium": "Rainfed adequate-N reference; used to reveal water stress under normal crop growth.",
    "fixed_I60_N150": "Fixed irrigation response benchmark, not an expert strategy.",
    "fixed_I120_N150": "Adequate irrigation response benchmark, not an expert strategy.",
    "expert_reference_recorded": "Recorded site management reference from a single-year observed record; not a year-specific optimum and may underperform in extreme drought years.",
    "ppo_replay_transfer": "FQA 007_04J schedule transfer diagnostic only; not HLA-trained PPO and not model-weight transfer.",
    "dssat_auto_attempt": "DSSAT automatic management attempt through gym-DSSAT; diagnostic unless automatic totals are confirmed from logs.",
}


def ensure_dirs() -> None:
    for sub in ["daily_outputs/HLA", "event_outputs", "evaluation", "figures", "reports", "logs", "rendered_inputs"]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)


def to_num(series: pd.Series, default: float = 0.0) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").fillna(default)


def add_weather(daily: pd.DataFrame) -> pd.DataFrame:
    out = daily.copy()
    out["scenario"] = out.get("scenario", "")
    if "rain" in out.columns:
        return out
    weather = msc.read_weather(STATION, YEAR)
    if weather.empty:
        for col in ["rain", "srad", "tmax", "tmin"]:
            out[col] = np.nan
        return out
    out["date_dt"] = pd.to_datetime(out["date"])
    out = out.merge(weather, left_on="date_dt", right_on="date", how="left", suffixes=("", "_weather"))
    out = out.drop(columns=["date_dt", "date_weather"], errors="ignore")
    return out


def event_table_from_daily(daily: pd.DataFrame, scenario: str) -> pd.DataFrame:
    if daily.empty:
        return pd.DataFrame()
    mask = (to_num(daily.get("real_action_amir", pd.Series(dtype=float))) > 0) | (
        to_num(daily.get("real_action_anfer", pd.Series(dtype=float))) > 0
    )
    rows = []
    for idx, row in enumerate(daily.loc[mask].itertuples(index=False), start=1):
        rows.append(
            {
                "station": STATION,
                "year": YEAR,
                "scenario": scenario,
                "event_id": idx,
                "decision_dap": float(getattr(row, "dap", np.nan)),
                "irrigation": float(getattr(row, "real_action_amir", 0.0)),
                "nitrogen": float(getattr(row, "real_action_anfer", 0.0)),
                "swfac_after": float(getattr(row, "swfac", np.nan)),
                "nstres_after": float(getattr(row, "nstres", np.nan)),
            }
        )
    return pd.DataFrame(rows)


def normalize_reused_daily(scenario: str, treatment_name: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    src = ALL_YEAR_ROOT / "daily_outputs" / STATION / f"{STATION}_{YEAR}_{treatment_name}_daily.csv"
    daily = pd.read_csv(src)
    daily["scenario"] = scenario
    daily["action_record_source"] = "reused_all_year_fixed_management"
    daily["diagnostic_reward"] = to_num(daily.get("reward", pd.Series(dtype=float)))
    daily = add_weather(daily)
    out = OUTPUT_ROOT / "daily_outputs" / STATION / f"{YEAR}_{scenario}_daily.csv"
    daily.to_csv(out, index=False, encoding="utf-8-sig")
    events = event_table_from_daily(daily, scenario)
    event_path = OUTPUT_ROOT / "event_outputs" / f"{STATION}_{YEAR}_{scenario}_events.csv"
    events.to_csv(event_path, index=False, encoding="utf-8-sig")
    return daily, events


def build_hla2004_env_config() -> dict[str, Any]:
    cfg = all_year.load_yaml(DIRECT_CONFIG)
    cfg["paths"]["output_root"] = str(OUTPUT_ROOT.relative_to(PROJECT_ROOT))
    pool = pd.read_csv(SCENARIO_POOL)
    row = pool[(pool["station_code"].eq(STATION)) & (pool["year"].eq(YEAR))]
    if row.empty:
        raise KeyError(f"{STATION} {YEAR} missing from scenario pool")
    env_config = all_year.build_env_config(cfg, row)
    return env_config


def run_single_scenario(env_config: dict[str, Any], display_scenario: str, runner_scenario: str) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    original_output_root = msc.OUTPUT_ROOT
    try:
        msc.OUTPUT_ROOT = OUTPUT_ROOT
        msc.ensure_dirs([STATION])
        summary, daily, events = msc.run_scenario(
            env_config,
            station=STATION,
            year=YEAR,
            scenario=runner_scenario,
            seed=0,
            max_steps=260,
        )
    finally:
        msc.OUTPUT_ROOT = original_output_root
    if daily.empty:
        return daily, events, summary
    daily["scenario"] = display_scenario
    if display_scenario == "expert_reference_recorded":
        daily["action_record_source"] = "single_year_site_observed_management_record"
    if display_scenario == "dssat_auto_attempt":
        daily["action_record_source"] = "dssat_internal_attempt_not_fully_exposed"
    daily_path = OUTPUT_ROOT / "daily_outputs" / STATION / f"{YEAR}_{display_scenario}_daily.csv"
    daily.to_csv(daily_path, index=False, encoding="utf-8-sig")
    if not events.empty:
        events["scenario"] = display_scenario
    event_path = OUTPUT_ROOT / "event_outputs" / f"{STATION}_{YEAR}_{display_scenario}_events.csv"
    events.to_csv(event_path, index=False, encoding="utf-8-sig")
    summary = dict(summary)
    summary["scenario"] = display_scenario
    summary["scenario_label"] = display_scenario
    summary["daily_csv_path"] = str(daily_path.relative_to(PROJECT_ROOT))
    summary["event_csv_path"] = str(event_path.relative_to(PROJECT_ROOT))
    summary["scenario_note"] = SCENARIO_NOTES.get(display_scenario, "")
    return daily, events, summary


def summarize_daily(daily: pd.DataFrame, events: pd.DataFrame, scenario: str) -> dict[str, Any]:
    sw = to_num(daily.get("swfac", pd.Series(dtype=float)))
    ns = to_num(daily.get("nstres", pd.Series(dtype=float)))
    grnwt = to_num(daily.get("grnwt", pd.Series(dtype=float)), np.nan).dropna()
    topwt = to_num(daily.get("topwt", pd.Series(dtype=float)), np.nan).dropna()
    irrig = to_num(events.get("irrigation", pd.Series(dtype=float)))
    n = to_num(events.get("nitrogen", pd.Series(dtype=float)))
    event_dap = to_num(events.get("decision_dap", pd.Series(dtype=float)), np.nan)
    daily_path = OUTPUT_ROOT / "daily_outputs" / STATION / f"{YEAR}_{scenario}_daily.csv"
    event_path = OUTPUT_ROOT / "event_outputs" / f"{STATION}_{YEAR}_{scenario}_events.csv"
    return {
        "station": STATION,
        "year": YEAR,
        "scenario": scenario,
        "run_status": "ok" if len(daily) else "missing",
        "daily_steps": int(len(daily)),
        "final_grnwt": float(grnwt.iloc[-1]) if len(grnwt) else np.nan,
        "final_topwt": float(topwt.iloc[-1]) if len(topwt) else np.nan,
        "total_irrigation": float(irrig.sum()),
        "total_n_or_tnup": float(n.sum()),
        "profit_low_water_cost": (0.01 * float(grnwt.iloc[-1]) - 0.1 * float(irrig.sum()) - 0.25 * float(n.sum())) if len(grnwt) else np.nan,
        "max_swfac": float(sw.max()) if len(sw) else np.nan,
        "swfac_days_gt_0p05": int((sw > 0.05).sum()) if len(sw) else 0,
        "max_nstres": float(ns.max()) if len(ns) else np.nan,
        "nstres_days_gt_0p05": int((ns > 0.05).sum()) if len(ns) else 0,
        "first_irrigation_dap": float(event_dap[irrig > 0].dropna().iloc[0]) if (irrig > 0).any() and len(event_dap[irrig > 0].dropna()) else np.nan,
        "first_n_dap": float(event_dap[n > 0].dropna().iloc[0]) if (n > 0).any() and len(event_dap[n > 0].dropna()) else np.nan,
        "daily_csv_path": str(daily_path.relative_to(PROJECT_ROOT)),
        "event_csv_path": str(event_path.relative_to(PROJECT_ROOT)),
        "scenario_note": SCENARIO_NOTES.get(scenario, ""),
    }


def process_diagnosis(summary: pd.DataFrame) -> pd.DataFrame:
    rows = []
    n_only = summary[summary["scenario"].eq("n_only_medium")]
    n_grnwt = float(n_only["final_grnwt"].iloc[0]) if len(n_only) else np.nan
    n_swdays = int(n_only["swfac_days_gt_0p05"].iloc[0]) if len(n_only) else 0
    for row in summary.itertuples(index=False):
        yield_gain_vs_n_only = float(row.final_grnwt) - n_grnwt if np.isfinite(n_grnwt) else np.nan
        rows.append(
            {
                "station": STATION,
                "year": YEAR,
                "scenario": row.scenario,
                "water_stress_reduction_vs_n_only_days": n_swdays - int(row.swfac_days_gt_0p05),
                "yield_gain_vs_n_only": yield_gain_vs_n_only,
                "is_year_selection_control": row.scenario in {"n_only_medium", "fixed_I60_N150", "fixed_I120_N150"},
                "is_recorded_expert_reference": row.scenario == "expert_reference_recorded",
                "is_diagnostic_not_final_ppo": row.scenario in {"ppo_replay_transfer", "dssat_auto_attempt"},
                "interpretation": SCENARIO_NOTES.get(row.scenario, ""),
            }
        )
    return pd.DataFrame(rows)


def read_all_daily(scenarios: list[str]) -> pd.DataFrame:
    frames = []
    for scenario in scenarios:
        path = OUTPUT_ROOT / "daily_outputs" / STATION / f"{YEAR}_{scenario}_daily.csv"
        if path.exists():
            frames.append(pd.read_csv(path))
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def plot_process(scenarios: list[str]) -> Path:
    colors = {
        "null_zero": "#808080",
        "n_only_medium": "#4F81BD",
        "fixed_I60_N150": "#9BBB59",
        "fixed_I120_N150": "#00A2E8",
        "expert_reference_recorded": "#8064A2",
        "ppo_replay_transfer": "#C0504D",
        "dssat_auto_attempt": "#F79646",
    }
    fig, axes = plt.subplots(4, 1, figsize=(11, 8.5), dpi=180, sharex=True)
    for scenario in scenarios:
        df = pd.read_csv(OUTPUT_ROOT / "daily_outputs" / STATION / f"{YEAR}_{scenario}_daily.csv")
        dap = to_num(df["dap"])
        if scenario == "null_zero":
            axes[0].bar(dap, to_num(df.get("rain", pd.Series(dtype=float))), color="#4F81BD", alpha=0.35, width=1.0)
        axes[1].plot(dap, to_num(df["swfac"]), label=scenario, color=colors.get(scenario), linewidth=1.6)
        axes[2].plot(dap, to_num(df["nstres"]), label=scenario, color=colors.get(scenario), linewidth=1.6)
        axes[3].plot(dap, to_num(df["grnwt"]), label=scenario, color=colors.get(scenario), linewidth=1.6)
        actions = (to_num(df.get("real_action_amir", pd.Series(dtype=float))) > 0) | (
            to_num(df.get("real_action_anfer", pd.Series(dtype=float))) > 0
        )
        for event in df.loc[actions].itertuples(index=False):
            axes[1].axvline(float(event.dap), color=colors.get(scenario), alpha=0.12, linewidth=1.0)
    axes[0].set_ylabel("RAIN")
    axes[1].set_ylabel("SWFAC")
    axes[2].set_ylabel("NSTRES")
    axes[3].set_ylabel("GRNWT")
    axes[3].set_xlabel("DAP")
    for ax in axes:
        ax.grid(alpha=0.2)
    axes[1].legend(loc="upper left", fontsize=7, ncol=2)
    fig.tight_layout()
    path = OUTPUT_ROOT / "figures" / "HLA_2004_rain_stress_actions_growth.png"
    fig.savefig(path)
    plt.close(fig)
    return path


def plot_events(scenarios: list[str]) -> Path:
    fig, axes = plt.subplots(2, 1, figsize=(11, 5.5), dpi=180, sharex=True)
    offsets = np.linspace(-3.0, 3.0, len(scenarios))
    for offset, scenario in zip(offsets, scenarios):
        df = pd.read_csv(OUTPUT_ROOT / "daily_outputs" / STATION / f"{YEAR}_{scenario}_daily.csv")
        dap = to_num(df["dap"]) + offset
        axes[0].bar(dap, to_num(df.get("real_action_amir", pd.Series(dtype=float))), width=1.2, label=scenario)
        axes[1].bar(dap, to_num(df.get("real_action_anfer", pd.Series(dtype=float))), width=1.2, label=scenario)
    axes[0].set_ylabel("Irrigation mm")
    axes[1].set_ylabel("N kg/ha")
    axes[1].set_xlabel("DAP")
    for ax in axes:
        ax.grid(axis="y", alpha=0.25)
    axes[0].legend(loc="upper left", fontsize=7, ncol=2)
    fig.tight_layout()
    path = OUTPUT_ROOT / "figures" / "HLA_2004_management_events.png"
    fig.savefig(path)
    plt.close(fig)
    return path


def plot_summary(summary: pd.DataFrame) -> Path:
    fig, axes = plt.subplots(1, 3, figsize=(12, 4), dpi=180)
    x = np.arange(len(summary))
    labels = summary["scenario"].astype(str).str.replace("_", "\n")
    axes[0].bar(x, to_num(summary["final_grnwt"]), color="#4F81BD")
    axes[0].set_title("GRNWT")
    axes[1].bar(x, to_num(summary["total_irrigation"]), color="#00A2E8")
    axes[1].set_title("Irrigation")
    axes[2].bar(x, to_num(summary["total_n_or_tnup"]), color="#9BBB59")
    axes[2].set_title("Nitrogen/TNUP")
    for ax in axes:
        ax.set_xticks(x)
        ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=7)
        ax.grid(axis="y", alpha=0.2)
    fig.tight_layout()
    path = OUTPUT_ROOT / "figures" / "HLA_2004_yield_water_n_summary.png"
    fig.savefig(path)
    plt.close(fig)
    return path


def md_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df.empty:
        return "_No rows._"
    work = df.head(max_rows).copy() if max_rows else df.copy()
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    sep = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, sep, *rows])


def write_report(summary: pd.DataFrame, diagnosis: pd.DataFrame, fig_paths: list[Path]) -> None:
    cols = [
        "scenario",
        "final_grnwt",
        "total_irrigation",
        "total_n_or_tnup",
        "max_swfac",
        "swfac_days_gt_0p05",
        "max_nstres",
        "nstres_days_gt_0p05",
        "profit_low_water_cost",
        "scenario_note",
    ]
    lines = [
        "# 008_06 HLA 2004 Representative Management Comparison",
        "",
        "## Purpose",
        "",
        "This task checks HLA 2004 as a water-stress representative year before any new PPO training.",
        "The fixed irrigation scenarios are year-selection controls, not expert strategies.",
        "",
        "## Key Interpretation",
        "",
        "- HLA 2004 is an extreme water-stress case under rainfed adequate-N management.",
        "- `expert_reference_recorded` is a single-year observed management reference, not a year-specific optimum.",
        "- `ppo_replay_transfer` is a diagnostic FQA schedule transfer, not HLA-trained PPO.",
        "- `dssat_auto_attempt` remains diagnostic unless native DSSAT automatic management is fully confirmed from logs.",
        "",
        "## Scenario Summary",
        "",
        md_table(summary[cols].round(3)),
        "",
        "## Process Diagnosis",
        "",
        md_table(diagnosis.round(3)),
        "",
        "## Recommendation",
        "",
        "Use HLA 2004 as the primary water-stress showcase. The next model step should train a stress-aware HLA 2004 PPO or build an irrigation gate using HLA 2004 before cross-site generalization.",
        "",
        "## Figures",
        "",
    ]
    for path in fig_paths:
        lines.append(f"- `{path.relative_to(PROJECT_ROOT)}`")
    DOC_MD.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")


def add_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.6), Inches(0.45))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 0, 0)


def add_text(slide, text: str, left: float, top: float, width: float, height: float, size: int = 15) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(0, 0, 0)


def write_ppt(summary: pd.DataFrame, fig_paths: list[Path]) -> None:
    if Presentation is None:
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "008_06 HLA 2004 水分限制代表情景")
    add_text(
        slide,
        "目的：在训练新 PPO 前，先验证 HLA 2004 是否适合做水分优化展示年份。\n"
        "固定 I60/I120 是水分响应对照，不是专家策略。\n"
        "expert_reference_recorded 是单一年份实测管理记录，不代表极端干旱年的最优策略。",
        0.6,
        1.1,
        12.0,
        2.0,
        16,
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "关键结果")
    cols = ["scenario", "final_grnwt", "total_irrigation", "total_n_or_tnup", "swfac_days_gt_0p05", "max_swfac"]
    df = summary[cols].round(2)
    table = slide.shapes.add_table(len(df) + 1, len(cols), Inches(0.25), Inches(0.9), Inches(12.85), Inches(3.8)).table
    for j, col in enumerate(cols):
        cell = table.cell(0, j)
        cell.text = col
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(68, 114, 196)
        for p in cell.text_frame.paragraphs:
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(9)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = "" if pd.isna(value) else str(value)
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(221, 235, 247)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(8)
                p.font.color.rgb = RGBColor(0, 0, 0)
    add_text(slide, "结论：HLA 2004 是强水分限制情景，适合下一步做 stress-aware PPO。", 0.6, 5.15, 12, 0.8, 14)

    for path in fig_paths:
        slide = prs.slides.add_slide(blank)
        add_title(slide, path.stem)
        slide.shapes.add_picture(str(path), Inches(0.55), Inches(0.9), width=Inches(12.1))

    prs.save(DOC_PPT)


def main() -> int:
    ensure_dirs()
    scenarios = list(REUSED_SCENARIOS.keys()) + list(RUN_SCENARIOS.keys())
    daily_tables: dict[str, pd.DataFrame] = {}
    event_tables: dict[str, pd.DataFrame] = {}
    summaries: list[dict[str, Any]] = []

    for scenario, treatment_name in REUSED_SCENARIOS.items():
        daily, events = normalize_reused_daily(scenario, treatment_name)
        daily_tables[scenario] = daily
        event_tables[scenario] = events
        summaries.append(summarize_daily(daily, events, scenario))

    env_config = build_hla2004_env_config()
    for display_scenario, runner_scenario in RUN_SCENARIOS.items():
        daily, events, summary = run_single_scenario(env_config, display_scenario, runner_scenario)
        daily_tables[display_scenario] = daily
        event_tables[display_scenario] = events
        if daily.empty:
            summary["scenario"] = display_scenario
            summary["scenario_note"] = SCENARIO_NOTES.get(display_scenario, "")
            summaries.append(summary)
        else:
            summaries.append(summarize_daily(daily, events, display_scenario) | {"scenario_note": SCENARIO_NOTES.get(display_scenario, "")})

    summary_df = pd.DataFrame(summaries)
    diagnosis = process_diagnosis(summary_df)
    summary_path = OUTPUT_ROOT / "evaluation" / "HLA_2004_representative_scenario_summary.csv"
    diagnosis_path = OUTPUT_ROOT / "evaluation" / "HLA_2004_process_diagnosis.csv"
    summary_df.to_csv(summary_path, index=False, encoding="utf-8-sig")
    diagnosis.to_csv(diagnosis_path, index=False, encoding="utf-8-sig")

    fig_paths = [plot_process(scenarios), plot_events(scenarios), plot_summary(summary_df)]
    write_report(summary_df, diagnosis, fig_paths)
    write_ppt(summary_df, fig_paths)

    print(summary_df[["scenario", "run_status", "final_grnwt", "total_irrigation", "total_n_or_tnup", "swfac_days_gt_0p05", "max_swfac"]].to_string(index=False))
    print(f"summary_csv={summary_path}")
    print(f"diagnosis_csv={diagnosis_path}")
    print(f"report_md={DOC_MD}")
    print(f"report_ppt={DOC_PPT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

