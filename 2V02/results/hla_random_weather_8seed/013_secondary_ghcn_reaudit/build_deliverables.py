"""Write the 013 Chinese audit report and a compact 5-slide record deck."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent


def md_table(frame: pd.DataFrame) -> str:
    headers = list(frame.columns)
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    for row in frame.itertuples(index=False, name=None):
        lines.append("| " + " | ".join(str(value) for value in row) + " |")
    return "\n".join(lines)


def main() -> None:
    gate = json.loads((OUT / "final_gate.json").read_text(encoding="utf-8"))
    event = pd.read_csv(OUT / "station_event_metrics.csv")
    inv = pd.read_csv(OUT / "candidate_station_inventory.csv")
    summary = pd.read_csv(OUT / "station_reaudit_summary.csv")
    unresolved = pd.read_csv(OUT / "unresolved_dates_after_013.csv")
    annual = pd.read_csv(OUT / "station_annual_qa.csv")
    sanity = json.loads((OUT / "2013_climate_sanity.json").read_text(encoding="utf-8"))
    month_counts = pd.Series(gate["unresolved_month_counts"], dtype=int)
    month_table = pd.DataFrame({"月份": [f"2013-{int(month):02d}" for month in month_counts.index],
                                "未解决天数": month_counts.values})
    view = event[["station_id", "zero_day_comparable_days", "zero_day_event_agreement", "zero_day_hit_rate", "date_alignment"]].copy()
    view["zero_day_event_agreement"] = (view["zero_day_event_agreement"] * 100).round(1).astype(str) + "%"
    view["zero_day_hit_rate"] = (view["zero_day_hit_rate"] * 100).round(1).astype(str) + "%"

    original_reasons = inv[["station_id", "008_original_fail_reason"]].drop_duplicates()
    report = f"""# HLA 2013 次级 GHCN 重新审计（013）

执行日期：2026-10-05。仅使用 HLA 2004–2013 fitting period。未处理 FQA。

## 结论

`HLA_GHCN50756_GENERAL_FALLBACK` 在 012 已接受。013 重新审计了 {len(event)} 个有缓存原始 `.dly` 且覆盖至少 1 个剩余缺口的旧 FAIL 站，没有候选达到冻结的 `wet/dry agreement >= 80%`。

`HLA_RAIN_PROVENANCE = PARTIAL`  
`READY_TO_RESUME_010 = FALSE`

013 新解决 0 天，剩余 111 天，全部仍在 2013。012 的 2013 已选降雨 subtotal 为 1,149.8 mm；该年未闭合，不报告全年总雨量、wet-day count 或最大日雨量。

## Truth 口径

只把 HLA/D222 实际存在数值的 962 天作为 truth，明确的数值 0 作为 dry，正值作为 wet；D222 三字段全空白不进入事件比较。D222 与项目旧雨量表属于同一观测数据链，二者一致不算两份独立证据。未使用 `.WTH` 作 truth。

008 原事件筛选在后续来源审计修正前，把本地空白按 0 参与比较。013 不把这称为代码错误，而是按修正后的 truth 口径重审。8 个候选均为 008 原 FAIL；真值口径改为仅显式观测日后，事件一致率仍低于 80%，所以没有 FAIL→PASS。

## 候选站的事件比较

日期 shift 以 GHCN 日期移动后再连接 HLA truth。各站 0 日的一致率最高，±1 日均未提供更合理的整体日期对齐；但事件一致率均未过门槛。

{md_table(view)}

最低一致率为 {event.zero_day_event_agreement.min()*100:.1f}%，最高为 {event.zero_day_event_agreement.max()*100:.1f}%。其中 RSM00031587 仅有 {int(event.loc[event.station_id=='RSM00031587','zero_day_comparable_days'].iloc[0])} 个可比 truth 日，虽然命中率较高，其一致率仍低于门槛，且站点距离较远，不能升级。完整 MAE/RMSE、相关系数、false wet/false dry 和累计差见 `station_date_shift_metrics.csv`。

008 原 FAIL 原因如下：

{md_table(original_reasons)}

013 新分类全部为 `FAIL_EVENT_AGREEMENT`。站点 metadata、原始 hash、PRCP 单位/缺测/flag、覆盖、极端值和距离见 inventory 与相关 QA CSV。所有入选日值均要求 QFLAG 为空；`-9999` 按缺测保留。未删除或裁剪极端值。

## 覆盖剩余111日

这些日期在 50756 及 012 已批准次级站均缺测。所有 013 候选即使有些日期有数值，也未通过正式一致率门槛，因此 013 新获批站为 0，对最终 source selection 不作升级。111 日原样保留 unresolved。

| 月份 | 未解决天数 |
|---|---:|
{md_table(month_table)}

最早日期 `{gate['earliest_unresolved']}`，最晚 `{gate['latest_unresolved']}`，最长连续段 {gate['longest_consecutive_unresolved_run']} 天。NASA 在这些日期有 35 天为 0、76 天大于 0，仅作参考。逐日的每站有效值/缺测/QFLAG状态见 `unresolved_dates_after_013.csv` 和 `new_candidate_ranking_on_111_days.csv`。

## 量级、空间和2013 sanity

月尺度只基于 2004–2012 的显式观测日，local subtotal 不当作完整月总量；年尺度排除2013，不用它评定候选站量级。各年候选可用降水小计、同日配对小计、最大日值见 `station_annual_qa.csv`。有些站离 HLA 约189–273 km，空间代表性为 QUESTIONABLE；坐标来自 GHCN inventory，HLA坐标来自既有天气站表头。没有正式地形/站点场址材料可证明空间等价。

`HLA_2013_CLIMATE_SANITY = {sanity['status']}`。1,149.8 mm 是当前选中日值 subtotal，已高于 2004–2012 完整年总量范围 {sanity['2004_2012_complete_year_total_min_mm']:.1f}–{sanity['2004_2012_complete_year_total_max_mm']:.1f} mm 和 NASA 2013 的 {sanity['2013_NASA_total_mm']:.2f} mm；但2013仍有111天未解决，不能称其为完整年总量或据此删改高雨量日。50756 与其他获准站的2013可用日数及小计见 `2013_climate_sanity.json`。

## 最小下一步

只需为准确列出的111天提供可追溯的逐日 HLA/邻站降水来源，或提供足以支持某个已缓存候选站通过冻结80%事件门槛及量级/空间审查的独立同站重叠资料。现有未批准邻站日值不能因为能填缺口而直接使用，NASA也不能升级为正式source。获得新证据后，对候选重新审计，再重建3653天来源表。

本次复核了012冻结日序列、D222/旧表/NESDC、011目标窗口证据、008 station audit 和各相关原始 `.dly` SHA-256。未修改源文件，未执行 fitting weather、CLI、WGEN、DSSAT 或 PPO。

建议提交：013 prompt、`run_reaudit.py`、`build_deliverables.py`、小型 CSV/JSON、中文报告和 PPT；不提交原始天气。建议 commit：`hla: reaudit secondary GHCN stations for 2013 rainfall gaps`。未 commit 或 push。
"""
    (OUT / "findings.md").write_text(report, encoding="utf-8")
    (ROOT / "docs/hla_secondary_ghcn_reaudit_013.md").write_text(report, encoding="utf-8")

    slides = [
        ("012之后仍有111天待查", ["缺口全部在2013", "50756与008获批次级站在这些日期均缺测", "仅复核已缓存、曾未通过筛选的站点"]),
        ("013修正本地truth口径", ["只用D222真实数值日，明确0才算dry", "全空白日不参与wet/dry比较", "D222与旧表属共享数据链，不重复计证据"]),
        ("8个候选均未通过事件门槛", [f"0日一致率范围 {event.zero_day_event_agreement.min()*100:.1f}%–{event.zero_day_event_agreement.max()*100:.1f}%", "冻结门槛为80%", "0日对齐最佳，但空间距离较远的站点仍有代表性疑问"]),
        ("111天没有减少", ["013新批准站：0", "1–6月：16、11、16、4、8、4天", "7–12月：3、3、6、15、9、16天", "35天NASA参考为0，76天大于0，仅作参考"]),
        ("RAIN门禁仍为PARTIAL", ["2013已选subtotal：1149.8 mm，非完整年总量", "剩余111天，最长连续9天", "READY_TO_RESUME_010 = FALSE", "下一步需要逐日可追溯降水或合格候选的独立验证证据"]),
    ]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for index, (heading, bullets) in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.65), Inches(0.6), Inches(12), Inches(0.9))
        box.text_frame.text = heading
        run = box.text_frame.paragraphs[0].runs[0]
        run.font.name, run.font.size, run.font.bold = "Microsoft YaHei", Pt(32), True
        if index == 2:
            table_data = event[["station_id", "zero_day_event_agreement", "zero_day_hit_rate"]].copy()
            table_data["zero_day_event_agreement"] = (table_data["zero_day_event_agreement"] * 100).map(lambda x: f"{x:.1f}%")
            table_data["zero_day_hit_rate"] = (table_data["zero_day_hit_rate"] * 100).map(lambda x: f"{x:.1f}%")
            table = slide.shapes.add_table(9, 3, Inches(1.0), Inches(1.75), Inches(11.3), Inches(4.9)).table
            table.columns[0].width, table.columns[1].width, table.columns[2].width = Inches(5.3), Inches(3.0), Inches(3.0)
            headers = ["GHCN站", "0日事件一致率", "0日命中率"]
            for col, value in enumerate(headers):
                cell = table.cell(0, col)
                cell.text = value
            for row, values in enumerate(table_data.itertuples(index=False, name=None), start=1):
                for col, value in enumerate(values):
                    table.cell(row, col).text = str(value)
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        para.font.name, para.font.size = "Microsoft YaHei", Pt(18)
        else:
            body = slide.shapes.add_textbox(Inches(0.85), Inches(2), Inches(11.6), Inches(4.6))
            body.text_frame.word_wrap = True
            for item_index, text in enumerate(bullets):
                para = body.text_frame.paragraphs[0] if item_index == 0 else body.text_frame.add_paragraph()
                para.text = text
                para.font.name, para.font.size = "Microsoft YaHei", Pt(22 if index == 3 else 24)
                para.space_after = Pt(20)
        footer = slide.shapes.add_textbox(Inches(0.85), Inches(7), Inches(11.5), Inches(0.25))
        footer.text_frame.text = f"HLA次级GHCN重新审计  |  2026-10-05  |  {index + 1}/5"
        footer.text_frame.paragraphs[0].font.size = Pt(11)
    prs.save(ROOT / "docs/hla_secondary_ghcn_reaudit_013.pptx")


if __name__ == "__main__":
    main()
