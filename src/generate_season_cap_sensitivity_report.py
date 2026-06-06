from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))

from ppo_safe_rendering import PROJECT_ROOT, load_yaml


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_season_cap_sensitivity.yaml"
REPORT_STEM = f"{date.today().isoformat()}_season_cap_sensitivity_analysis_report"


def fmt(value, digits: int = 2) -> str:
    try:
        return f"{float(value):.{digits}f}"
    except Exception:
        return ""


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path) if path.exists() else pd.DataFrame()


def paths(config: dict) -> dict[str, Path]:
    root = PROJECT_ROOT / config["paths"]["output_root"]
    return {
        "root": root,
        "plan": root / "configs" / "season_cap_sensitivity_plan.csv",
        "summary": root / "evaluation" / "season_cap_sensitivity_evaluation_summary.csv",
        "saturation": root / "evaluation" / "cap_saturation_summary.csv",
        "marginal": root / "evaluation" / "marginal_response_by_cap.csv",
        "smoke": root / "smoke_checks" / "pretrain_smoke_check_summary.csv",
    }


def derive_findings(saturation: pd.DataFrame, marginal: pd.DataFrame) -> dict:
    all_at_cap = False
    if not saturation.empty:
        all_at_cap = saturation["irrigation_at_cap"].astype(str).str.lower().eq("true").all() and saturation["n_at_cap"].astype(str).str.lower().eq("true").all()
    sustained_yield_increase = {}
    sustained_reward_increase = {}
    diminishing = {}
    if not marginal.empty:
        for station, group in marginal.groupby("station"):
            group = group.sort_values("season_irrigation_cap")
            dy = pd.to_numeric(group["delta_yield_from_previous_cap"], errors="coerce").dropna()
            dr = pd.to_numeric(group["delta_reward_from_previous_cap"], errors="coerce").dropna()
            sustained_yield_increase[station] = bool((dy >= -1e-6).all()) if len(dy) else False
            sustained_reward_increase[station] = bool((dr >= -1e-6).all()) if len(dr) else False
            gains = pd.to_numeric(group["yield_gain_per_100mm_irrigation"], errors="coerce").dropna().tolist()
            diminishing[station] = bool(len(gains) >= 2 and gains[-1] < gains[0])
    return {
        "all_at_cap": all_at_cap,
        "sustained_yield_increase": sustained_yield_increase,
        "sustained_reward_increase": sustained_reward_increase,
        "diminishing": diminishing,
    }


def write_markdown(config: dict) -> Path:
    p = paths(config)
    plan = read_csv(p["plan"])
    summary = read_csv(p["summary"])
    saturation = read_csv(p["saturation"])
    marginal = read_csv(p["marginal"])
    smoke = read_csv(p["smoke"])
    findings = derive_findings(saturation, marginal)
    docs_path = PROJECT_ROOT / "docs" / f"{REPORT_STEM}.md"
    lines: list[str] = []
    lines.append("# Season cap sensitivity analysis report")
    lines.append("")
    lines.append(f"Generated at: {date.today().isoformat()}")
    lines.append("")
    lines.append("## Goal")
    lines.append("")
    lines.append("This stage diagnoses whether action-safe PPO behavior is dominated by the fixed season cap. It does not modify reward, does not train PPO without action safety, and should not be interpreted as final optimal water-nitrogen management.")
    lines.append("")
    lines.append("## Cap vs budget scenario")
    lines.append("")
    lines.append("- Action safety cap is a technical safety constraint on maximum seasonal water and nitrogen actions.")
    lines.append("- Rainfall-scaling budget scenario is a weather or water-availability experiment under altered rainfall conditions.")
    lines.append("- These are different experiments. Cap sensitivity should be used before a formal rainfall-scaling budget scenario.")
    lines.append("")
    lines.append("## Design")
    lines.append("")
    lines.append(f"- Total models planned: {len(plan)}")
    lines.append(f"- Timesteps per model: {config['training']['total_timesteps']}")
    lines.append("- Cap levels: 100/150, 150/225, 200/300, 250/375, 300/450")
    lines.append("- Run order: HLA pilot, then SYA, LCA, then FQA/YCA.")
    lines.append("")
    lines.append("## Pretrain smoke checks")
    lines.append("")
    if smoke.empty:
        lines.append("No pretrain smoke check table found.")
    else:
        ok_count = int(((smoke["run_status"] == "ok") & (smoke["episode_completed"].astype(str).str.lower().eq("true"))).sum())
        lines.append(f"- Passed smoke checks: {ok_count}/{len(smoke)}")
    lines.append("")
    lines.append("## Quality gate")
    lines.append("")
    if summary.empty:
        lines.append("No evaluation summary found.")
    else:
        passed = int(summary["quality_gate_pass"].astype(str).str.lower().eq("true").sum())
        lines.append(f"- Passed evaluations: {passed}/{len(summary)}")
        model_count = summary["policy_name"].nunique()
        lines.append(f"- Completed cap models: {model_count}")
    lines.append("")
    lines.append("## Cap saturation")
    lines.append("")
    lines.append(f"- All evaluations at or near their cap: {findings['all_at_cap']}")
    if not saturation.empty:
        sat_agg = saturation.groupby(["station", "cap_name", "season_irrigation_cap", "season_n_cap"], as_index=False).agg(
            irrigation_ratio=("irrigation_saturation_ratio", "mean"),
            n_ratio=("n_saturation_ratio", "mean"),
            trigger_days=("num_safety_trigger_days", "mean"),
        )
        lines.append("| station | cap | irrigation_cap | n_cap | irrigation_ratio | n_ratio | trigger_days |")
        lines.append("|---|---|---:|---:|---:|---:|---:|")
        for row in sat_agg.itertuples():
            lines.append(f"| {row.station} | {row.cap_name} | {fmt(row.season_irrigation_cap)} | {fmt(row.season_n_cap)} | {fmt(row.irrigation_ratio, 3)} | {fmt(row.n_ratio, 3)} | {fmt(row.trigger_days)} |")
    lines.append("")
    lines.append("## Marginal response")
    lines.append("")
    if not marginal.empty:
        show = marginal[["station", "cap_name", "mean_yield", "mean_reward", "mean_irrigation", "mean_n", "delta_yield_from_previous_cap", "yield_gain_per_100mm_irrigation", "limited_data_flag"]]
        lines.append("| station | cap | yield | reward | irrigation | n | delta_yield | yield_gain_per_100mm | limited |")
        lines.append("|---|---|---:|---:|---:|---:|---:|---:|---|")
        for row in show.itertuples():
            lines.append(f"| {row.station} | {row.cap_name} | {fmt(row.mean_yield)} | {fmt(row.mean_reward)} | {fmt(row.mean_irrigation)} | {fmt(row.mean_n)} | {fmt(row.delta_yield_from_previous_cap)} | {fmt(row.yield_gain_per_100mm_irrigation)} | {row.limited_data_flag} |")
    lines.append("")
    lines.append("## Interpretation")
    lines.append("")
    if findings["all_at_cap"]:
        lines.append("PPO filled every tested season cap. This means the current reward still encourages using all water and nitrogen allowed by action safety.")
    else:
        lines.append("PPO did not fill every tested season cap, so some cap levels may reveal genuine lower-input behavior.")
    lines.append("The 200/300 cap should be treated as a diagnostic safety setting, not as final management guidance.")
    lines.append("If yield increases flatten while actions remain saturated, 200/300 may be a reasonable conservative experimental cap; if yield/reward continue rising strongly up to 300/450, reward cost terms need attention before multi-seed validation.")
    lines.append("")
    lines.append("## Next recommendation")
    lines.append("")
    lines.append("Recommended next step: revise reward cost terms or run a constrained rainfall-scaling budget scenario only after deciding whether the cap should be fixed at a conservative value. Multi-seed stability should come after the cap/cost design is less dominated by the safety ceiling.")
    docs_path.write_text("\n".join(lines), encoding="utf-8")
    report_copy = p["root"] / "reports" / docs_path.name
    report_copy.parent.mkdir(parents=True, exist_ok=True)
    report_copy.write_text(docs_path.read_text(encoding="utf-8"), encoding="utf-8")
    return docs_path


def write_pptx(config: dict, markdown_path: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    p = paths(config)
    summary = read_csv(p["summary"])
    saturation = read_csv(p["saturation"])
    marginal = read_csv(p["marginal"])
    findings = derive_findings(saturation, marginal)
    ppt_path = PROJECT_ROOT / "docs" / f"{REPORT_STEM}.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def title(slide, body: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.2), Inches(0.55))
        para = box.text_frame.paragraphs[0]
        para.text = body
        para.font.name = "Microsoft YaHei"
        para.font.size = Pt(24)
        para.font.bold = True
        para.font.color.rgb = RGBColor(0, 0, 0)

    def text(slide, body: str, top: float = 1.0) -> None:
        box = slide.shapes.add_textbox(Inches(0.65), Inches(top), Inches(12.0), Inches(5.9))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(body.split("\n")):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            para.font.name = "Microsoft YaHei"
            para.font.size = Pt(17)
            para.font.color.rgb = RGBColor(0, 0, 0)

    def table(slide, df: pd.DataFrame, top: float = 1.0) -> None:
        if df.empty:
            text(slide, "No rows available.", top)
            return
        rows = min(len(df), 8) + 1
        cols = len(df.columns)
        shape = slide.shapes.add_table(rows, cols, Inches(0.35), Inches(top), Inches(12.6), Inches(5.2))
        tbl = shape.table
        for c, col in enumerate(df.columns):
            cell = tbl.cell(0, c)
            cell.text = str(col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(79, 129, 189)
        for r, (_, row) in enumerate(df.head(8).iterrows(), start=1):
            for c, col in enumerate(df.columns):
                cell = tbl.cell(r, c)
                cell.text = str(row[col])
                if r % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(230, 236, 247)
        for tr in tbl.rows:
            for cell in tr.cells:
                for para in cell.text_frame.paragraphs:
                    para.alignment = PP_ALIGN.CENTER
                    for run in para.runs:
                        run.font.name = "Microsoft YaHei"
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(0, 0, 0)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Season cap sensitivity analysis")
    text(slide, "本阶段是 action safety cap 的技术诊断。\n不训练无 safety PPO，不修改 reward，不做多 seed，不覆盖前面结果。\n目标是判断 PPO 是否被 season cap 主导。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Cap vs budget")
    text(slide, "Action safety cap: 限制单季总灌水/总施肥，是训练安全约束。\nRainfall-scaling budget scenario: 改变降雨或水资源情景，是气象情景实验。\n两者不是同一个实验，cap 敏感性应先于正式 budget scenario。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Experimental design")
    text(slide, "五组 cap: 100/150, 150/225, 200/300, 250/375, 300/450。\n每站只使用前一阶段 best policy 的 train_year。\n运行顺序: HLA pilot -> SYA -> LCA -> FQA -> YCA。\nFQA/YCA 标记为 limited two-year cross validation。")

    if not summary.empty:
        status = summary.groupby("station", as_index=False).agg(models=("policy_name", "nunique"), evaluations=("eval_year", "count"), passed=("quality_gate_pass", lambda s: int(s.astype(str).str.lower().eq("true").sum())))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title(slide, "Run status")
        table(slide, status)

    if not saturation.empty:
        sat = saturation.groupby(["station", "cap_name"], as_index=False).agg(irrigation_ratio=("irrigation_saturation_ratio", "mean"), n_ratio=("n_saturation_ratio", "mean"))
        sat["irrigation_ratio"] = sat["irrigation_ratio"].map(lambda v: fmt(v, 3))
        sat["n_ratio"] = sat["n_ratio"].map(lambda v: fmt(v, 3))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title(slide, "Cap saturation")
        table(slide, sat)

    if not marginal.empty:
        current = marginal[marginal["cap_name"].eq("cap_current")][["station", "mean_yield", "mean_reward", "mean_irrigation", "mean_n", "limited_data_flag"]].copy()
        for col in ["mean_yield", "mean_reward", "mean_irrigation", "mean_n"]:
            current[col] = current[col].map(lambda v: fmt(v))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title(slide, "Current 200/300 cap")
        table(slide, current)

        gain = marginal[["station", "cap_name", "delta_yield_from_previous_cap", "yield_gain_per_100mm_irrigation"]].copy()
        for col in ["delta_yield_from_previous_cap", "yield_gain_per_100mm_irrigation"]:
            gain[col] = gain[col].map(lambda v: fmt(v))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title(slide, "Marginal yield response")
        table(slide, gain)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Interpretation")
    text(slide, f"All tested caps saturated: {findings['all_at_cap']}\n200/300 不能直接解释为最终最优管理。\n如果所有 cap 都被打满，应优先修正 reward 成本项或明确采用固定 cap 情景。\n多 seed 稳定性分析建议放在 cap/cost 设计稳定之后。")

    prs.save(ppt_path)
    report_copy = p["root"] / "reports" / ppt_path.name
    report_copy.write_bytes(ppt_path.read_bytes())
    return ppt_path


def main() -> None:
    config = load_yaml(CONFIG_PATH)
    md = write_markdown(config)
    ppt = write_pptx(config, md)
    print(md.relative_to(PROJECT_ROOT))
    print(ppt.relative_to(PROJECT_ROOT))


if __name__ == "__main__":
    main()
