#!/usr/bin/env python3
"""Reconcile the YC ChinaFLUX multi-scale weather products without editing inputs."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import statistics
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW_ROOT = PROJECT_ROOT / "data" / "external" / "yc_chinaflux" / "raw"
RESULT_ROOT = PROJECT_ROOT / "results" / "yc_chinaflux_weather_reconciliation"
WTH_ROOT = PROJECT_ROOT / "DSSAT_auto_validation" / "multisite_new_cultivar_inputs_013_lowIC_manual" / "YC"
LEGACY_CSV = RESULT_ROOT / "legacy_source_observations_2004_2010.csv"
DATASETS = {
    "30min": "YCA_M_30min.zip",
    "daily": "YCA_M_daily.zip",
    "monthly": "YCA_M_monthly.zip",
    "yearly": "YCA_M_yearly.zip",
}
YEARS = range(2003, 2011)
WGEN_YEARS = range(2004, 2011)
SENTINEL = -99999.0


def is_leap(year: int) -> bool:
    return year % 4 == 0 and (year % 100 != 0 or year % 400 == 0)


def number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(result) or result == SENTINEL:
        return None
    return result


def missing_kind(value: Any) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return "blank"
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return "unparsed"
    if not math.isfinite(numeric):
        return "nonfinite"
    if numeric == SENTINEL:
        return "sentinel_minus_99999"
    return None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def zip_display_name(raw_name: str, flag_bits: int) -> str:
    if flag_bits & 0x800:
        return raw_name
    try:
        return raw_name.encode("cp437").decode("gbk")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return raw_name


def open_year_sheet(archive: zipfile.ZipFile, year: int):
    candidates = [
        info for info in archive.infolist()
        if info.filename.lower().endswith(".xlsx")
        and re.search(rf"(?<!\d){year}(?!\d)", info.filename)
    ]
    if len(candidates) != 1:
        raise ValueError(f"Expected one XLSX member for {year}, found {len(candidates)}")
    workbook = load_workbook(BytesIO(archive.read(candidates[0])), read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    units = next(rows)
    return workbook, candidates[0], headers, units, rows


def column_index(headers: tuple[Any, ...], token: str) -> int:
    matches = [i for i, item in enumerate(headers) if token in str(item)]
    if len(matches) != 1:
        raise ValueError(f"Expected one column containing {token!r}, found {matches}")
    return matches[0]


def parse_date_row(row: tuple[Any, ...], year: int, daily: bool = False) -> date:
    row_year = int(row[0])
    month, day = int(row[1]), int(row[2])
    result = date(row_year, month, day)
    if row_year != year:
        raise ValueError(f"Row year {row_year} does not match workbook year {year}")
    return result


def load_legacy_observations() -> dict[date, dict[str, Any]]:
    if not LEGACY_CSV.exists():
        raise FileNotFoundError(
            f"Missing {LEGACY_CSV.name}. Run scripts/reconcile_yc_chinaflux_legacy_excel.ps1 first."
        )
    result: dict[date, dict[str, Any]] = {}
    with LEGACY_CSV.open("r", encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            day = date.fromisoformat(row["date"])
            values: dict[str, Any] = {"row_present_in_any_source": row.get("row_present_in_any_source", "")}
            for variable in ("SRAD", "TMAX", "TMIN", "RAIN"):
                values[f"status_{variable}"] = row.get(f"status_{variable}", "") or "no_row"
                values[f"raw_{variable}"] = number(row.get(f"raw_{variable}"))
                values[f"source_{variable}"] = row.get(f"source_{variable}", "")
            result[day] = values
    return result


def read_cleaned_yca() -> dict[date, dict[str, float | None]]:
    path = PROJECT_ROOT / "weather_clean" / "YCA_weather_cleaned.csv"
    result: dict[date, dict[str, float | None]] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            day = date.fromisoformat(row["date"])
            result[day] = {key: number(row.get(key)) for key in ("SRAD", "TMAX", "TMIN", "RAIN")}
    return result


def read_old_wth(year: int) -> dict[date, dict[str, float]]:
    path = WTH_ROOT / f"CNYC{year % 100:02d}01.WTH"
    if not path.exists():
        return {}
    rows: dict[date, dict[str, float]] = {}
    with path.open("r", encoding="ascii", errors="replace") as stream:
        for line in stream:
            fields = line.split()
            if len(fields) < 5 or not re.fullmatch(r"\d{7}", fields[0]):
                continue
            year_code, doy = int(fields[0][:4]), int(fields[0][4:])
            if year_code != year:
                continue
            day = date(year, 1, 1) + timedelta(days=doy - 1)
            try:
                srad, tmax, tmin, rain = map(float, fields[1:5])
            except ValueError:
                continue
            rows[day] = {"SRAD": srad, "TMAX": tmax, "TMIN": tmin, "RAIN": rain}
    return rows


def safe_sum(values: list[float]) -> float:
    return sum(values)


def pearson(xs: list[float], ys: list[float]) -> float | None:
    if len(xs) < 2 or len(xs) != len(ys):
        return None
    mx, my = statistics.mean(xs), statistics.mean(ys)
    dx, dy = [x - mx for x in xs], [y - my for y in ys]
    denominator = math.sqrt(sum(x * x for x in dx) * sum(y * y for y in dy))
    return sum(x * y for x, y in zip(dx, dy)) / denominator if denominator else None


def spell_metrics(values: dict[date, float | None], complete: dict[date, bool]) -> tuple[int, float | None, int | None]:
    wet_days = 0
    max_daily = None
    best_dry, run = 0, 0
    for day in sorted(values):
        value = values[day]
        if not complete.get(day, False) or value is None:
            run = 0
            continue
        if value > 0.1:
            wet_days += 1
            max_daily = value if max_daily is None else max(max_daily, value)
            run = 0
        else:
            run += 1
            best_dry = max(best_dry, run)
    return wet_days, max_daily, best_dry


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: "" if row.get(key) is None else row.get(key) for key in fields})


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> int:
    global PROJECT_ROOT, RAW_ROOT, RESULT_ROOT, WTH_ROOT, LEGACY_CSV
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=PROJECT_ROOT, help="Project root")
    args = parser.parse_args()
    PROJECT_ROOT = args.root.resolve()
    RAW_ROOT = PROJECT_ROOT / "data" / "external" / "yc_chinaflux" / "raw"
    RESULT_ROOT = PROJECT_ROOT / "results" / "yc_chinaflux_weather_reconciliation"
    WTH_ROOT = PROJECT_ROOT / "DSSAT_auto_validation" / "multisite_new_cultivar_inputs_013_lowIC_manual" / "YC"
    LEGACY_CSV = RESULT_ROOT / "legacy_source_observations_2004_2010.csv"
    RESULT_ROOT.mkdir(parents=True, exist_ok=True)

    legacy = load_legacy_observations()
    cleaned = read_cleaned_yca()
    archive_paths = {name: RAW_ROOT / filename for name, filename in DATASETS.items()}
    absent = [str(path) for path in archive_paths.values() if not path.exists()]
    if absent:
        raise FileNotFoundError("Missing input ZIPs: " + ", ".join(absent))
    archives = {
        name: zipfile.ZipFile(path, metadata_encoding="gbk")
        for name, path in archive_paths.items()
    }

    try:
        source_inventory: dict[str, Any] = {
            "source": "ChinaFLUX / National Ecosystem Science Data Center",
            "site": "YC/YCA Yucheng Station",
            "data_period": "2003-2010",
            "official_resource_page": "https://www.nesdc.org.cn/otherProject/index?menuId=station&projectId=1047",
            "archives": {},
        }
        for key, path in archive_paths.items():
            archive = archives[key]
            members = [info for info in archive.infolist() if info.filename.lower().endswith(".xlsx")]
            member_years = sorted({int(match.group(0)) for info in members for match in re.finditer(r"(?<!\d)(200\d|2010)(?!\d)", info.filename)})
            source_inventory["archives"][path.name] = {
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "xlsx_member_count": len(members),
                "year_coverage": member_years,
                "internal_excel_filenames": [zip_display_name(info.filename, info.flag_bits) for info in members],
                "parse_status": "pending",
            }

        temperature_map = {
            "status": "BLOCKED_TEMPERATURE_SENSOR_MAPPING",
            "selected_field": None,
            "candidate_fields": [
                {
                    "field": "近地面空气温度",
                    "physical_meaning": "air temperature below the vegetation canopy",
                    "height_m": None,
                    "height_evidence": "Product documentation lists 1.6 m and 2.9 m for two temperature sensors but does not pair a height to this exported column.",
                },
                {
                    "field": "冠层上方空气温度",
                    "physical_meaning": "air temperature above the vegetation canopy",
                    "height_m": None,
                    "height_evidence": "Product documentation lists 1.6 m and 2.9 m for two temperature sensors but does not pair a height to this exported column.",
                },
            ],
            "available_heights_m": [1.6, 2.9],
            "decision": "Both temperature series are retained as parallel daily candidates. No TMAX/TMIN field is selected until the metadata owner confirms the field-to-height pairing and DSSAT-representative level.",
            "evidence": [
                "The local YCA_M_30min.pdf data dictionary defines near-ground temperature as below-canopy and the second field as above-canopy.",
                "The same product document lists both sensor heights, 1.6 m and 2.9 m, without an explicit field-to-height assignment.",
                "The 2003-2005 ChinaFLUX network paper reports YC air-temperature sensor levels at 2 m and 1 m, but does not identify which current exported column matches either sensor or the later 1.6/2.9 m pair.",
            ],
        }

        half_hour_audit: list[dict[str, Any]] = []
        missing_intervals: list[dict[str, Any]] = []
        cross_scale: list[dict[str, Any]] = []
        daily_candidates: list[dict[str, Any]] = []
        sr_daily_check: list[dict[str, Any]] = []
        monthly_cross: dict[tuple[int, int], dict[str, Any]] = {}
        daily_metrics: dict[date, dict[str, Any]] = {}
        precipitation_scale_detail: dict[int, dict[str, Any]] = {}
        legacy_row_missing: list[dict[str, Any]] = []

        for year in YEARS:
            workbook, info, headers, units, rows = open_year_sheet(archives["30min"], year)
            try:
                indexes = {
                    "below": column_index(headers, "近地面空气温度"),
                    "above": column_index(headers, "冠层上方空气温度"),
                    "srad": column_index(headers, "太阳辐射"),
                    "rain": column_index(headers, "降水量"),
                }
                expected_records = (366 if is_leap(year) else 365) * 48
                year_start = datetime(year, 1, 1) + timedelta(minutes=30)
                expected_times = {year_start + timedelta(minutes=30 * n) for n in range(expected_records)}
                seen_times: set[datetime] = set()
                day_values: dict[date, dict[str, list[float]]] = defaultdict(lambda: {k: [] for k in indexes})
                day_missing_times: dict[date, dict[str, list[datetime]]] = defaultdict(lambda: {k: [] for k in indexes})
                missing_counts = {key: 0 for key in indexes}
                missing_by_kind = {key: defaultdict(int) for key in indexes}
                longest_run = {key: 0 for key in indexes}
                current_run = {key: 0 for key in indexes}
                max_time_gap = 0
                out_of_order_count = 0
                previous: datetime | None = None
                row_count = 0
                for row in rows:
                    if not row or row[0] is None:
                        continue
                    row_count += 1
                    row_year, month, day_num, hour, minute, second = (int(row[i]) for i in range(6))
                    data_day = date(row_year, month, day_num)
                    timestamp = datetime(row_year, month, day_num) + timedelta(hours=hour, minutes=minute, seconds=second)
                    if previous is not None:
                        delta = (timestamp - previous).total_seconds() / 1800
                        if delta <= 0:
                            out_of_order_count += 1
                        elif delta > 1:
                            max_time_gap = max(max_time_gap, int(delta - 1))
                    previous = timestamp
                    seen_times.add(timestamp)
                    for key, index in indexes.items():
                        value = row[index] if index < len(row) else None
                        kind = missing_kind(value)
                        if kind:
                            missing_counts[key] += 1
                            missing_by_kind[key][kind] += 1
                            current_run[key] += 1
                            longest_run[key] = max(longest_run[key], current_run[key])
                            day_missing_times[data_day][key].append(timestamp)
                        else:
                            current_run[key] = 0
                            day_values[data_day][key].append(float(value))

                missing_grid = len(expected_times - seen_times)
                unexpected_grid = len(seen_times - expected_times)
                for key in indexes:
                    half_hour_audit.append({
                        "year": year,
                        "variable": {"below": "近地面空气温度", "above": "冠层上方空气温度", "srad": "太阳辐射", "rain": "降水量"}[key],
                        "expected_records": expected_records,
                        "actual_records": row_count,
                        "unique_timestamps": len(seen_times),
                        "duplicate_timestamps": row_count - len(seen_times),
                        "missing_or_duplicate_grid_slots": missing_grid,
                        "unexpected_grid_timestamps": unexpected_grid,
                        "out_of_order_steps": out_of_order_count,
                        "max_missing_time_grid_slots_between_rows": max_time_gap,
                        "missing_total": missing_counts[key],
                        "missing_fraction": missing_counts[key] / expected_records,
                        "minus_99999_count": missing_by_kind[key]["sentinel_minus_99999"],
                        "blank_count": missing_by_kind[key]["blank"],
                        "nonfinite_count": missing_by_kind[key]["nonfinite"],
                        "unparsed_count": missing_by_kind[key]["unparsed"],
                        "longest_consecutive_missing_records": longest_run[key],
                    })

                year_daily: dict[date, dict[str, Any]] = {}
                for offset in range(366 if is_leap(year) else 365):
                    day = date(year, 1, 1) + timedelta(days=offset)
                    values = day_values[day]
                    counts = {key: len(values[key]) for key in indexes}
                    partial_rain = safe_sum(values["rain"])
                    if day_missing_times[day]["rain"]:
                        first_missing = min(day_missing_times[day]["rain"])
                        last_missing = max(day_missing_times[day]["rain"])
                        missing_intervals.append({
                            "date": day.isoformat(),
                            "year": year,
                            "doy": offset + 1,
                            "variable": "RAIN",
                            "expected_30min_records": 48,
                            "valid_30min_records": counts["rain"],
                            "missing_30min_records": len(day_missing_times[day]["rain"]),
                            "first_missing_timestamp": first_missing.isoformat(sep=" "),
                            "last_missing_timestamp": last_missing.isoformat(sep=" "),
                            "30min_valid_partial_sum_mm": round(partial_rain, 6),
                        })
                    for key, field in (("below", "T_BELOW"), ("above", "T_ABOVE"), ("srad", "SRAD")):
                        if day_missing_times[day][key]:
                            first_missing = min(day_missing_times[day][key])
                            last_missing = max(day_missing_times[day][key])
                            missing_intervals.append({
                                "date": day.isoformat(),
                                "year": year,
                                "doy": offset + 1,
                                "variable": field,
                                "expected_30min_records": 48,
                                "valid_30min_records": counts[key],
                                "missing_30min_records": len(day_missing_times[day][key]),
                                "first_missing_timestamp": first_missing.isoformat(sep=" "),
                                "last_missing_timestamp": last_missing.isoformat(sep=" "),
                                "30min_valid_partial_sum_mm": "",
                            })
                    legacy_day = legacy.get(day, {})
                    legacy_rain_status = legacy_day.get("status_RAIN", "no_row")
                    legacy_rain = legacy_day.get("raw_RAIN")
                    rain_incomplete = bool(day_missing_times[day]["rain"])
                    rain_candidate = partial_rain
                    rain_source = "ChinaFLUX_30min"
                    rain_resolved = not rain_incomplete
                    if rain_incomplete:
                        rain_candidate = None
                        rain_source = "unresolved"
                        if legacy_rain_status == "numeric" and legacy_rain is not None:
                            rain_candidate = legacy_rain
                            rain_source = "legacy_original_daily_fallback"
                            rain_resolved = True
                    below = values["below"]
                    above = values["above"]
                    srad_values = values["srad"]
                    srad_complete = counts["srad"] == 48
                    rain_valid_count = counts["rain"]
                    candidate = {
                        "date": day.isoformat(),
                        "year": year,
                        "doy": offset + 1,
                        "TMAX": "",
                        "TMIN": "",
                        "TMAX_below_canopy_C": round(max(below), 6) if below else "",
                        "TMIN_below_canopy_C": round(min(below), 6) if below else "",
                        "TMAX_above_canopy_C": round(max(above), 6) if above else "",
                        "TMIN_above_canopy_C": round(min(above), 6) if above else "",
                        "SRAD": round(safe_sum(srad_values) * 0.0018, 6) if srad_values else "",
                        "RAIN": round(rain_candidate, 6) if rain_candidate is not None else "",
                        "RAIN_30min_valid_partial_sum": round(partial_rain, 6),
                        "temp_valid_count": min(counts["below"], counts["above"]),
                        "temp_valid_count_below_canopy": counts["below"],
                        "temp_valid_count_above_canopy": counts["above"],
                        "srad_valid_count": counts["srad"],
                        "rain_valid_count": rain_valid_count,
                        "rain_incomplete": str(rain_incomplete).lower(),
                        "rain_candidate_resolved": str(rain_resolved).lower(),
                        "rain_source_used": rain_source,
                        "source_temperature_field": "",
                        "source_dataset": "ChinaFLUX YCA_M_30min" + (" + legacy original rainfall fallback" if rain_source == "legacy_original_daily_fallback" else ""),
                        "qc_status": "BLOCKED_TEMPERATURE_SENSOR_MAPPING" + (";rain_gap_replaced_with_traceable_legacy_observation" if rain_source == "legacy_original_daily_fallback" else ";BLOCKED_PRECIPITATION_GAPS" if rain_incomplete else ""),
                    }
                    daily_candidates.append(candidate)
                    year_daily[day] = {
                        "below_tmax": max(below) if below else None,
                        "below_tmin": min(below) if below else None,
                        "above_tmax": max(above) if above else None,
                        "above_tmin": min(above) if above else None,
                        "srad": safe_sum(srad_values) * 0.0018 if srad_complete else None,
                        "rain_30min": partial_rain if not rain_incomplete else None,
                        "rain_partial": partial_rain,
                        "rain_complete": not rain_incomplete,
                        "rain_candidate": rain_candidate,
                        "rain_resolved": rain_resolved,
                        "rain_source": rain_source,
                    }
                    daily_metrics[day] = year_daily[day]

                daily_book, _, daily_headers, _, daily_rows = open_year_sheet(archives["daily"], year)
                daily_rain: dict[date, float | None] = {}
                daily_srad_mean: dict[date, float | None] = {}
                try:
                    rain_i = column_index(daily_headers, "日降水量")
                    srad_i = column_index(daily_headers, "日平均太阳辐射")
                    for row in daily_rows:
                        day = parse_date_row(row, year, daily=True)
                        daily_rain[day] = number(row[rain_i])
                        daily_srad_mean[day] = number(row[srad_i])
                finally:
                    daily_book.close()

                srad_diffs: list[float] = []
                srad_pairs: list[tuple[float, float]] = []
                for day, metric in year_daily.items():
                    daily_avg = daily_srad_mean.get(day)
                    daily_srad_check = daily_avg * 0.0864 if daily_avg is not None else None
                    if metric["srad"] is not None and daily_srad_check is not None:
                        diff = metric["srad"] - daily_srad_check
                        srad_diffs.append(diff)
                        srad_pairs.append((metric["srad"], daily_srad_check))
                    sr_daily_check.append({
                        "date": day.isoformat(),
                        "year": year,
                        "SRAD_from_30min_MJ_m2_day": metric["srad"],
                        "daily_mean_W_m2": daily_avg,
                        "SRAD_check_from_daily_MJ_m2_day": daily_srad_check,
                        "difference_30min_minus_daily": round(metric["srad"] - daily_srad_check, 8) if metric["srad"] is not None and daily_srad_check is not None else "",
                        "valid_30min_srad_records": len(day_values[day]["srad"]),
                    })

                monthly_book, _, monthly_headers, _, monthly_rows = open_year_sheet(archives["monthly"], year)
                monthly_rain: dict[int, float | None] = {}
                try:
                    monthly_rain_i = column_index(monthly_headers, "月降水量")
                    for row in monthly_rows:
                        if int(row[0]) == year:
                            month = int(row[1])
                            monthly_rain[month] = number(row[monthly_rain_i])
                            monthly_cross[(year, month)] = {
                                "year": year,
                                "month": month,
                                "chinaflux_monthly_rain_mm": monthly_rain[month],
                            }
                finally:
                    monthly_book.close()

                yearly_book, _, yearly_headers, _, yearly_rows = open_year_sheet(archives["yearly"], year)
                yearly_rain_values: list[float] = []
                try:
                    yearly_rain_i = column_index(yearly_headers, "年降水量")
                    for row in yearly_rows:
                        if int(row[0]) == year:
                            value = number(row[yearly_rain_i])
                            if value is not None:
                                yearly_rain_values.append(value)
                finally:
                    yearly_book.close()

                valid_daily_rain = [value for value in daily_rain.values() if value is not None]
                valid_monthly_rain = [value for value in monthly_rain.values() if value is not None]
                half_hour_rain_sum = sum(metric["rain_partial"] for metric in year_daily.values())
                daily_sum = sum(valid_daily_rain)
                monthly_sum = sum(valid_monthly_rain)
                yearly_value = yearly_rain_values[0] if yearly_rain_values else None
                old_wth = read_old_wth(year)
                old_rain_sum = sum(row["RAIN"] for row in old_wth.values()) if old_wth else None
                rain_gap_days = sum(not metric["rain_complete"] for metric in year_daily.values())
                rain_fallback_days = sum(metric["rain_source"] == "legacy_original_daily_fallback" for metric in year_daily.values())
                rain_unresolved_days = sum(not metric["rain_resolved"] for metric in year_daily.values())
                cross_scale.append({
                    "year": year,
                    "expected_days": 366 if is_leap(year) else 365,
                    "half_hour_records": row_count,
                    "half_hour_rain_valid_sum_mm_lower_bound": round(half_hour_rain_sum, 3),
                    "half_hour_rain_missing_records": missing_counts["rain"],
                    "half_hour_rain_incomplete_days": rain_gap_days,
                    "daily_rain_valid_sum_mm": round(daily_sum, 3),
                    "daily_rain_valid_days": len(valid_daily_rain),
                    "daily_rain_missing_days": len(daily_rain) - len(valid_daily_rain),
                    "monthly_rain_valid_sum_mm": round(monthly_sum, 3),
                    "monthly_rain_valid_months": len(valid_monthly_rain),
                    "yearly_rain_mm": round(yearly_value, 3) if yearly_value is not None else "",
                    "daily_minus_half_hour_mm": round(daily_sum - half_hour_rain_sum, 3),
                    "monthly_minus_half_hour_mm": round(monthly_sum - half_hour_rain_sum, 3),
                    "yearly_minus_half_hour_mm": round(yearly_value - half_hour_rain_sum, 3) if yearly_value is not None else "",
                    "daily_minus_monthly_mm": round(daily_sum - monthly_sum, 3),
                    "srad_crosscheck_days": len(srad_diffs),
                    "srad_crosscheck_mae_MJ_m2_day": round(statistics.mean(abs(x) for x in srad_diffs), 8) if srad_diffs else "",
                    "srad_crosscheck_max_abs_MJ_m2_day": round(max(abs(x) for x in srad_diffs), 8) if srad_diffs else "",
                    "old_wth_rain_total_mm": round(old_rain_sum, 3) if old_rain_sum is not None else "",
                    "legacy_rain_fallback_days": rain_fallback_days,
                    "rain_days_still_unresolved_after_legacy_check": rain_unresolved_days,
                })
                precipitation_scale_detail[year] = {
                    "daily": daily_rain,
                    "monthly": monthly_rain,
                    "yearly": yearly_value,
                    "daily_valid_sum": daily_sum,
                    "monthly_valid_sum": monthly_sum,
                    "half_hour_partial_sum": half_hour_rain_sum,
                }
                for month in range(1, 13):
                    entry = monthly_cross.setdefault((year, month), {"year": year, "month": month})
                    month_days = [day for day in year_daily if day.month == month]
                    daily_month_values = [daily_rain.get(day) for day in month_days]
                    entry["daily_valid_rain_sum_mm"] = round(sum(v for v in daily_month_values if v is not None), 3)
                    entry["daily_valid_days"] = sum(v is not None for v in daily_month_values)
                    entry["daily_missing_days"] = len(month_days) - entry["daily_valid_days"]
                    entry["half_hour_valid_partial_sum_mm"] = round(sum(year_daily[day]["rain_partial"] for day in month_days), 3)
                    entry["half_hour_incomplete_days"] = sum(not year_daily[day]["rain_complete"] for day in month_days)
                    entry["legacy_wth_rain_sum_mm"] = round(sum(old_wth[day]["RAIN"] for day in month_days if day in old_wth), 3)
                    entry["chinaflux_monthly_rain_mm"] = monthly_rain.get(month)
                    entry["daily_minus_monthly_mm"] = round(entry["daily_valid_rain_sum_mm"] - monthly_rain[month], 3) if monthly_rain.get(month) is not None else ""

            finally:
                workbook.close()

        # Add legacy provenance and day-level comparison without treating imputations as observations.
        daily_comparison: list[dict[str, Any]] = []
        comparison_summary: list[dict[str, Any]] = []
        monthly_compare_rows: list[dict[str, Any]] = []
        unresolved_gaps: list[dict[str, Any]] = []
        old_wth_by_year = {year: read_old_wth(year) for year in WGEN_YEARS}
        candidate_by_date = {date.fromisoformat(row["date"]): row for row in daily_candidates}

        for candidate in daily_candidates:
            day = date.fromisoformat(candidate["date"])
            year = day.year
            if year not in WGEN_YEARS:
                continue
            old = old_wth_by_year[year].get(day, {})
            raw = legacy.get(day, {})
            clean = cleaned.get(day, {})
            provenance: dict[str, str] = {}
            for variable in ("SRAD", "TMAX", "TMIN", "RAIN"):
                raw_status = raw.get(f"status_{variable}", "no_row")
                raw_value = raw.get(f"raw_{variable}")
                old_value = old.get(variable)
                if raw_status == "numeric" and raw_value is not None and old_value is not None and abs(old_value - raw_value) <= 0.11:
                    status = "legacy_original_observation_matches_WTH"
                elif variable == "RAIN" and raw_status in ("blank", "sentinel", "unparsed"):
                    status = "legacy_rain_blank_or_invalid_to_zero"
                elif raw_status == "no_row":
                    status = "legacy_no_source_row_then_cleaning_fill"
                elif raw_status in ("blank", "sentinel", "unparsed"):
                    status = "legacy_monthly_mean_imputed"
                elif raw_status == "numeric":
                    status = "legacy_numeric_source_WTH_mismatch"
                else:
                    status = "cannot_determine"
                provenance[variable] = status
                candidate[f"old_{variable}_provenance"] = status

            for key in ("SRAD", "TMAX", "TMIN", "RAIN"):
                candidate[f"old_cleaned_{key}"] = clean.get(key, "")
                candidate[f"old_wth_{key}"] = old.get(key, "")
            candidate["chinaflux_rain_missing_from_30min"] = candidate["rain_incomplete"]
            candidate["legacy_raw_rain_status"] = raw.get("status_RAIN", "no_row")
            candidate["legacy_raw_rain_value"] = raw.get("raw_RAIN") if raw.get("raw_RAIN") is not None else ""

            metric = daily_metrics[day]
            row = {
                "date": day.isoformat(),
                "year": year,
                "month": day.month,
                "doy": day.timetuple().tm_yday,
                "old_wth_SRAD": old.get("SRAD", ""),
                "old_wth_TMAX": old.get("TMAX", ""),
                "old_wth_TMIN": old.get("TMIN", ""),
                "old_wth_RAIN": old.get("RAIN", ""),
                "ChinaFLUX_TMAX_below_canopy_C": metric["below_tmax"],
                "ChinaFLUX_TMIN_below_canopy_C": metric["below_tmin"],
                "ChinaFLUX_TMAX_above_canopy_C": metric["above_tmax"],
                "ChinaFLUX_TMIN_above_canopy_C": metric["above_tmin"],
                "ChinaFLUX_SRAD_MJ_m2_day": metric["srad"],
                "ChinaFLUX_RAIN_30min_valid_partial_mm": metric["rain_partial"],
                "ChinaFLUX_RAIN_complete_30min_mm": metric["rain_30min"] if metric["rain_30min"] is not None else "",
                "ChinaFLUX_RAIN_candidate_after_traceable_fallback_mm": metric["rain_candidate"] if metric["rain_candidate"] is not None else "",
                "rain_incomplete": candidate["rain_incomplete"],
                "rain_source_used": metric["rain_source"],
                "legacy_raw_rain_status": raw.get("status_RAIN", "no_row"),
                "legacy_raw_rain_value_mm": raw.get("raw_RAIN") if raw.get("raw_RAIN") is not None else "",
                "old_SRAD_provenance": provenance["SRAD"],
                "old_TMAX_provenance": provenance["TMAX"],
                "old_TMIN_provenance": provenance["TMIN"],
                "old_RAIN_provenance": provenance["RAIN"],
                "chinaflux_missing_variables": "RAIN" if candidate["rain_incomplete"] == "true" else "",
                "source_temperature_field": "unresolved; both branches retained",
                "qc_status": candidate["qc_status"],
            }
            daily_comparison.append(row)
            if candidate["rain_incomplete"] == "true":
                unresolved_gaps.append({
                    "date": day.isoformat(),
                    "year": year,
                    "doy": day.timetuple().tm_yday,
                    "variable": "RAIN",
                    "chinaflux_30min_missing_records": 48 - int(candidate["rain_valid_count"]),
                    "chinaflux_30min_valid_partial_sum_mm": candidate["RAIN_30min_valid_partial_sum"],
                    "daily_product_rain_mm_reference_only": precipitation_scale_detail[year]["daily"].get(day, ""),
                    "monthly_product_rain_mm_reference_only": precipitation_scale_detail[year]["monthly"].get(day.month, ""),
                    "yearly_product_rain_mm_reference_only": precipitation_scale_detail[year]["yearly"] or "",
                    "legacy_raw_rain_status": raw.get("status_RAIN", "no_row"),
                    "legacy_raw_rain_value_mm": raw.get("raw_RAIN") if raw.get("raw_RAIN") is not None else "",
                    "candidate_resolution": "replaced_from_traceable_legacy_numeric_observation" if metric["rain_source"] == "legacy_original_daily_fallback" else "unresolved_no_independent_numeric_daily_observation",
                    "legacy_no_source_row": str(raw.get("row_present_in_any_source") == "False").lower(),
                })

        # Old-versus-ChinaFLUX statistics. Rainfall comparisons exclude incomplete 30-minute days and any legacy fallback.
        for year in WGEN_YEARS:
            comparisons = [row for row in daily_comparison if row["year"] == year]
            for candidate_name, new_tmax, new_tmin in (
                ("near_surface_below_canopy", "ChinaFLUX_TMAX_below_canopy_C", "ChinaFLUX_TMIN_below_canopy_C"),
                ("above_canopy", "ChinaFLUX_TMAX_above_canopy_C", "ChinaFLUX_TMIN_above_canopy_C"),
            ):
                for variable, field in (("TMAX", new_tmax), ("TMIN", new_tmin)):
                    xs, ys = [], []
                    for row in comparisons:
                        new_value, old_value = number(row[field]), number(row[f"old_wth_{variable}"])
                        if new_value is not None and old_value is not None:
                            xs.append(new_value)
                            ys.append(old_value)
                    diffs = [x - y for x, y in zip(xs, ys)]
                    comparison_summary.append({
                        "year": year,
                        "variable": variable,
                        "candidate_temperature_field": candidate_name,
                        "n_paired_days": len(diffs),
                        "mean_bias_candidate_minus_old_C": round(statistics.mean(diffs), 4) if diffs else "",
                        "MAE_C": round(statistics.mean(abs(x) for x in diffs), 4) if diffs else "",
                        "RMSE_C": round(math.sqrt(statistics.mean(x * x for x in diffs)), 4) if diffs else "",
                        "Pearson_r": round(pearson(xs, ys), 5) if pearson(xs, ys) is not None else "",
                    })
            srad_pairs = [
                (number(row["ChinaFLUX_SRAD_MJ_m2_day"]), number(row["old_wth_SRAD"]))
                for row in comparisons
            ]
            srad_pairs = [(x, y) for x, y in srad_pairs if x is not None and y is not None]
            srad_diffs = [x - y for x, y in srad_pairs]
            comparison_summary.append({
                "year": year,
                "variable": "SRAD",
                "candidate_temperature_field": "",
                "n_paired_days": len(srad_pairs),
                "mean_bias_candidate_minus_old_MJ_m2_day": round(statistics.mean(srad_diffs), 4) if srad_diffs else "",
                "MAE_MJ_m2_day": round(statistics.mean(abs(x) for x in srad_diffs), 4) if srad_diffs else "",
                "RMSE_MJ_m2_day": round(math.sqrt(statistics.mean(x * x for x in srad_diffs)), 4) if srad_diffs else "",
                "Pearson_r": round(pearson([x for x, _ in srad_pairs], [y for _, y in srad_pairs]), 5) if pearson([x for x, _ in srad_pairs], [y for _, y in srad_pairs]) is not None else "",
            })
            rain_pairs = [
                row for row in comparisons
                if row["rain_incomplete"] == "false"
                and number(row["ChinaFLUX_RAIN_complete_30min_mm"]) is not None
                and number(row["old_wth_RAIN"]) is not None
            ]
            new_rain = [number(row["ChinaFLUX_RAIN_complete_30min_mm"]) for row in rain_pairs]
            old_rain = [number(row["old_wth_RAIN"]) for row in rain_pairs]
            wet, max_daily, max_dry = spell_metrics(
                {date.fromisoformat(row["date"]): number(row["ChinaFLUX_RAIN_complete_30min_mm"]) for row in comparisons},
                {date.fromisoformat(row["date"]): row["rain_incomplete"] == "false" for row in comparisons},
            )
            summary_row = next((row for row in cross_scale if row["year"] == year), {})
            comparison_summary.append({
                "year": year,
                "variable": "RAIN",
                "candidate_temperature_field": "",
                "n_paired_days": len(rain_pairs),
                "mean_bias_candidate_minus_old_mm": round(statistics.mean(a - b for a, b in zip(new_rain, old_rain)), 4) if rain_pairs else "",
                "MAE_mm": round(statistics.mean(abs(a - b) for a, b in zip(new_rain, old_rain)), 4) if rain_pairs else "",
                "RMSE_mm": round(math.sqrt(statistics.mean((a - b) ** 2 for a, b in zip(new_rain, old_rain))), 4) if rain_pairs else "",
                "Pearson_r": round(pearson(new_rain, old_rain), 5) if pearson(new_rain, old_rain) is not None else "",
                "ChinaFLUX_complete_day_rain_sum_mm": round(sum(new_rain), 3),
                "old_WTH_annual_rain_sum_mm": summary_row.get("old_wth_rain_total_mm", ""),
                "ChinaFLUX_complete_day_wet_days_gt_0p1mm": wet,
                "ChinaFLUX_complete_day_max_daily_rain_mm": round(max_daily, 3) if max_daily is not None else "",
                "ChinaFLUX_max_observed_dry_spell_on_complete_days": max_dry,
                "rain_metric_scope": "30-minute complete days only; incomplete and fallback days excluded",
            })
            for month in range(1, 13):
                date_rows = [row for row in comparisons if date.fromisoformat(row["date"]).month == month]
                valid_new = [number(row["ChinaFLUX_RAIN_complete_30min_mm"]) for row in date_rows if row["rain_incomplete"] == "false"]
                old_month = [number(row["old_wth_RAIN"]) for row in date_rows]
                old_month_valid = [value for value in old_month if value is not None]
                monthly_compare_rows.append({
                    "year": year,
                    "month": month,
                    "ChinaFLUX_complete_day_rain_sum_mm": round(sum(value for value in valid_new if value is not None), 3),
                    "ChinaFLUX_complete_days": len(valid_new),
                    "ChinaFLUX_incomplete_days": sum(row["rain_incomplete"] == "true" for row in date_rows),
                    "old_WTH_monthly_rain_sum_mm": round(sum(old_month_valid), 3),
                    "old_WTH_days": len(old_month_valid),
                    "ChinaFLUX_complete_minus_old_WTH_mm": round(sum(value for value in valid_new if value is not None) - sum(old_month_valid), 3),
                })

        # The 2009-12 source-table gap has no raw row; keep it separate from cell-level blanks.
        for day, raw in legacy.items():
            if 2004 <= day.year <= 2010 and raw.get("row_present_in_any_source") == "False":
                legacy_row_missing.append({
                    "date": day.isoformat(),
                    "year": day.year,
                    "variable": "ALL_LEGACY_SOURCE_VARIABLES",
                    "status": "no_source_row",
                    "cleaned_YCA_weather_exists": day in cleaned,
                    "cleaned_values_filled": True,
                    "interpretation": "All three raw Excel sources have no row for this date; the cleaned daily record exists after calendar completion and filling.",
                })
        unresolved_gaps.extend(legacy_row_missing)

        for year, inventory in source_inventory["archives"].items():
            inventory["parse_status"] = "parsed_all_8_annual_workbooks"
        source_inventory["legacy_raw_excel_extraction"] = {
            "csv": LEGACY_CSV.name,
            "manifest": "legacy_source_inventory.json",
            "method": "Excel COM read-only extraction of original T2.xls, D32.xls, and HLLCYCFQ rainfall workbook",
        }
        source_inventory["product_metadata_documents"] = [
            "YCA_M_30min.pdf",
            "YCA_M_DAILY.pdf",
            "YCA_M_MONTHLY.pdf",
            "YCA_M_YEARLY.pdf",
        ]

        option_a_rows = [row for row in daily_candidates if 2004 <= row["year"] <= 2010]
        option_a_days = len(option_a_rows)
        option_a_unresolved_rain = sum(row["rain_candidate_resolved"] == "false" for row in option_a_rows)
        option_a_fallback = sum(row["rain_source_used"] == "legacy_original_daily_fallback" for row in option_a_rows)
        prefill_path = PROJECT_ROOT / "weather_clean" / "data_check_by_year_before_fill.csv"
        prefill = {}
        with prefill_path.open("r", encoding="utf-8-sig", newline="") as stream:
            for row in csv.DictReader(stream):
                if row.get("station") == "YCA" and 2011 <= int(row["year"]) <= 2013:
                    prefill[int(row["year"])] = {key: int(row[key]) for key in ("SRAD", "TMAX", "TMIN", "RAIN", "inserted_missing_dates")}
        option_b_source_gaps = sum(
            values[key]
            for values in prefill.values()
            for key in ("SRAD", "TMAX", "TMIN", "inserted_missing_dates")
        )
        fit_options = {
            "dssat_weatherMan_guidance": {
                "nominal_input_years": "5-10 years of daily weather data",
                "source": "DSSAT.net FAQ: How to generate daily weather data",
                "url": "https://dssat.net/5165/",
                "meaning": "Nominal record-length guidance only; it does not waive variable completeness or provenance requirements.",
            },
            "option_A": {
                "years": "2004-2010",
                "nominal_year_count": 7,
                "all_years_inside_original_train_window": True,
                "validation_or_test_information_used": False,
                "candidate_days": option_a_days,
                "rain_days_still_unresolved_after_traceable_legacy_fallback": option_a_unresolved_rain,
                "days_using_traceable_legacy_rain_fallback": option_a_fallback,
                "source_consistency": "ChinaFLUX 30-minute series plus explicitly flagged legacy daily observations on missing-interval dates; not ChinaFLUX-only after fallback.",
                "nominal_WeatherMan_length_adequate": True,
                "eligible_for_CLI": False,
                "reason": "Temperature field-to-height mapping remains unresolved; any residual rain gaps also block parameter fitting. Seven nominal years cannot override these gates.",
            },
            "option_B": {
                "years": "2004-2013",
                "nominal_year_count": 10,
                "all_years_inside_original_train_window": True,
                "validation_or_test_information_used": False,
                "chinaflux_years": "2004-2010",
                "legacy_only_years": "2011-2013",
                "legacy_2011_2013_prefill_counts": prefill,
                "legacy_2011_2013_source_gap_values_excluding_rain": option_b_source_gaps,
                "source_consistency": "Not met; the last three years rely on legacy sources and contain pre-fill SRAD/temperature or date gaps.",
                "nominal_WeatherMan_length_adequate": True,
                "eligible_for_CLI": False,
                "reason": "2011-2013 lack a comparable, complete weather source and some variables were monthly-mean imputed; temperature mapping for ChinaFLUX remains unresolved.",
            },
            "final_decision": "No final fitting window approved and no CLI generated.",
        }

        rain_unresolved = sum(row.get("candidate_resolution", "").startswith("unresolved") for row in unresolved_gaps)
        summary = {
            "task": "003_03_yc_chinaflux_weather_reconciliation",
            "final_status": "BLOCKED_TEMPERATURE_SENSOR_MAPPING",
            "secondary_blockers": ["BLOCKED_PRECIPITATION_GAPS"] if rain_unresolved else [],
            "site_scope": "YC/YCA only",
            "parsed_archives": len(source_inventory["archives"]),
            "annual_workbooks_parsed": 32,
            "chinaflux_years": list(YEARS),
            "train_candidate_years": list(WGEN_YEARS),
            "temperature_sensor_mapping_confirmed": False,
            "rain_missing_days_in_chinaflux_30min": sum(row["variable"] == "RAIN" for row in missing_intervals),
            "rain_days_replaced_from_traceable_legacy_source": option_a_fallback,
            "rain_days_unresolved_after_legacy_source_check": option_a_unresolved_rain,
            "source_table_dates_without_any_row": [row["date"] for row in legacy_row_missing],
            "daily_temperature_candidate_branches": 2,
            "canonical_daily_TMAX_TMIN_selected": False,
            "formal_WTH_created": False,
            "CLI_created": False,
            "PPO_runs": 0,
            "DSSAT_smoke_runs": 0,
            "fit_options": fit_options,
        }

        write_json(RESULT_ROOT / "source_inventory.json", source_inventory)
        write_json(RESULT_ROOT / "temperature_sensor_mapping.json", temperature_map)
        write_json(RESULT_ROOT / "wgen_fit_window_options.json", fit_options)
        write_json(RESULT_ROOT / "audit_summary.json", summary)
        write_csv(RESULT_ROOT / "half_hour_integrity.csv", half_hour_audit)
        write_csv(RESULT_ROOT / "missing_daily_intervals.csv", missing_intervals)
        write_csv(RESULT_ROOT / "cross_scale_precipitation_check.csv", cross_scale)
        write_csv(RESULT_ROOT / "srad_daily_crosscheck.csv", sr_daily_check)
        write_csv(RESULT_ROOT / "yc_chinaflux_daily_candidate_2003_2010.csv", daily_candidates)
        write_csv(RESULT_ROOT / "yc_wgen_candidate_2004_2010.csv", option_a_rows)
        write_csv(RESULT_ROOT / "old_vs_chinaflux_daily_comparison.csv", daily_comparison)
        write_csv(RESULT_ROOT / "old_vs_chinaflux_summary.csv", comparison_summary)
        write_csv(RESULT_ROOT / "old_vs_chinaflux_monthly_rain.csv", monthly_compare_rows)
        write_csv(RESULT_ROOT / "unresolved_weather_gaps.csv", unresolved_gaps)

        print(json.dumps({
            "final_status": summary["final_status"],
            "secondary_blockers": summary["secondary_blockers"],
            "archive_count": summary["parsed_archives"],
            "cross_scale_years": len(cross_scale),
            "daily_candidate_rows": len(daily_candidates),
            "unresolved_rain_days": option_a_unresolved_rain,
            "legacy_rain_fallback_days": option_a_fallback,
        }, ensure_ascii=False))
        return 0
    finally:
        for archive in archives.values():
            archive.close()


if __name__ == "__main__":
    raise SystemExit(main())
