#!/usr/bin/env python3
"""Create the Chinese 003_04 audit deck from machine-readable audit outputs."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "yc_multisource_weather_reconstruction"
OUTPUT = ROOT / "docs" / "yc_multisource_weather_reconstruction.pptx"

INK = "202D35"
MUTED = "596A72"
PAPER = "F5F7F6"
WHITE = "FFFFFF"
TEAL = "187C78"
BLUE = "39799A"
GOLD = "D59A2A"
CORAL = "C95643"
LINE = "D6DEDD"
PALE_TEAL = "E1F0ED"
PALE_GOLD = "F7EEDB"
PALE_CORAL = "F5E4E0"
FONT = "Microsoft YaHei"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def fill_slide(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def rect(slide, x, y, w, h, color, line_color=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.color.rgb = rgb(line_color or color)
    shape.line.width = Pt(0.7)
    return shape


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         valign=MSO_ANCHOR.MIDDLE, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def title(slide, number, heading, kicker):
    text(slide, f"003_04   /   {number}", 0.6, 0.32, 3.0, 0.28, 9, TEAL, True)
    text(slide, heading, 0.6, 0.7, 11.8, 0.55, 25, INK, True)
    text(slide, kicker, 0.62, 1.27, 12.0, 0.32, 11, MUTED)
    rect(slide, 0.62, 1.72, 12.05, 0.012, LINE)


def footer(slide, page):
    rect(slide, 0.62, 7.12, 12.05, 0.012, LINE)
    text(slide, "YC/YCA  |  数据审计，不是 WGEN 输入候选", 0.65, 7.18, 7.8, 0.18, 8, MUTED)
    text(slide, f"{page} / 6", 11.6, 7.16, 1.0, 0.2, 8, MUTED, align=PP_ALIGN.RIGHT)


def add_table(slide, x, y, widths, row_h, headers, rows, font_size=11, highlight_rows=None):
    xx = x
    for i, (head, width) in enumerate(zip(headers, widths)):
        rect(slide, xx, y, width, row_h, INK)
        text(slide, head, xx + 0.06, y + 0.02, width - 0.12, row_h - 0.04, font_size, WHITE, True)
        xx += width
    for ri, row in enumerate(rows):
        yy = y + row_h * (ri + 1)
        xx = x
        bg = PALE_CORAL if highlight_rows and ri in highlight_rows else (WHITE if ri % 2 == 0 else "EDF1F0")
        for value, width in zip(row, widths):
            rect(slide, xx, yy, width, row_h, bg, LINE)
            text(slide, str(value), xx + 0.06, yy + 0.02, width - 0.12, row_h - 0.04, font_size, INK)
            xx += width


def load_data():
    summary = json.loads((RESULTS / "audit_summary.json").read_text(encoding="utf-8"))
    qc = json.loads((RESULTS / "weather_2011_2013_qc.json").read_text(encoding="utf-8"))
    with (RESULTS / "precipitation_gap_resolution.csv").open(encoding="utf-8-sig", newline="") as f:
        gaps = list(csv.DictReader(f))
    return summary, qc, gaps


def slide_one(prs, summary):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide(slide, INK)
    rect(slide, 0, 0, 0.18, 7.5, TEAL)
    text(slide, "YC / WEATHER DATA AUDIT", 0.72, 0.58, 5.8, 0.35, 11, "8FD0C2", True)
    text(slide, "多源官方天气\n重建审计", 0.7, 1.2, 7.7, 1.55, 34, WHITE, True, valign=MSO_ANCHOR.TOP)
    text(slide, "训练天气窗口 2004–2013  |  单站 YC/YCA  |  003_04", 0.76, 3.0, 8.2, 0.42, 15, "C7D5D6")
    rect(slide, 0.76, 4.02, 4.4, 0.68, CORAL, radius=True)
    text(slide, summary["final_status"], 0.92, 4.08, 4.1, 0.52, 14, WHITE, True)
    text(slide, "温度高度映射未证实；降水与 2011–2013 仍有阻塞。\n不生成 fitting CSV / WTH / CLI；不运行 WGEN、DSSAT、PPO。", 0.8, 5.0, 7.4, 0.95, 15, "E5ECEA", valign=MSO_ANCHOR.TOP)
    rect(slide, 8.85, 0.85, 3.75, 5.8, "2C3C43", line_color="53666D", radius=True)
    text(slide, "本轮三道门", 9.15, 1.14, 3.1, 0.42, 18, WHITE, True)
    gates = [("A", "温度字段/高度", CORAL), ("B", "35 天降水缺口", GOLD), ("C", "2011–2013 日值", BLUE)]
    for i, (letter, label, accent) in enumerate(gates):
        yy = 2.0 + i * 1.13
        rect(slide, 9.2, yy, 0.58, 0.58, accent, radius=True)
        text(slide, letter, 9.2, yy, 0.58, 0.58, 17, WHITE, True, align=PP_ALIGN.CENTER)
        text(slide, label, 9.95, yy - 0.01, 2.2, 0.28, 13, WHITE, True)
        status = summary["gates"][{"A": "A_temperature", "B": "B_precipitation", "C": "C_2011_2013"}[letter]]
        text(slide, status, 9.95, yy + 0.34, 2.45, 0.25, 9, "CCD8D8")
    text(slide, "2026-09-24", 0.78, 7.08, 2.5, 0.18, 8, "AFC0C2")
    return slide


def slide_sources(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide(slide)
    title(slide, "01", "三类产品，三种可用边界", "实际解析说明与表头；按产品分辨率决定能否支持逐日重建")
    left, right = 1.0, 12.25
    axis_y = 2.23
    rect(slide, left, axis_y, right - left, 0.02, LINE)
    for yr in (1998, 2003, 2005, 2010, 2013, 2022):
        x = left + (yr - 1998) / (2022 - 1998) * (right - left)
        rect(slide, x, axis_y - 0.07, 0.018, 0.16, MUTED)
        text(slide, str(yr), x - 0.28, axis_y - 0.42, 0.62, 0.23, 9, MUTED, align=PP_ALIGN.CENTER)
    lanes = [
        ("1998–2006 禹城监测产品", 1998, 2006, GOLD, "月值 | 10 个 XLS 成员", "不能拆成日值"),
        ("ChinaFLUX YCA", 2003, 2010, TEAL, "30 min / 日 / 月 / 年", "温度高度配对未说明"),
        ("禹城站大气环境要素", 2005, 2022, BLUE, "逐日 | 气象 + 辐射 XLSX", "QC 插补无逐日标志"),
    ]
    for i, (label, start, end, color, detail, limitation) in enumerate(lanes):
        yy = 2.85 + i * 1.05
        text(slide, label, 0.78, yy - 0.05, 2.55, 0.28, 12, INK, True)
        x1 = left + (start - 1998) / 24 * (right - left)
        x2 = left + (end - 1998) / 24 * (right - left)
        rect(slide, x1, yy + 0.34, max(0.16, x2 - x1), 0.22, color, radius=True)
        text(slide, detail, 0.78, yy + 0.61, 4.6, 0.24, 10, MUTED)
        text(slide, limitation, 8.2, yy + 0.56, 4.0, 0.28, 10, CORAL if i != 1 else GOLD, True, align=PP_ALIGN.RIGHT)
    rect(slide, 0.78, 6.3, 11.9, 0.48, PALE_GOLD, radius=True)
    text(slide, "同站点多产品交叉核对，不视为相互独立观测。原始 ZIP/XLSX 保持只读。", 0.98, 6.36, 11.5, 0.32, 12, INK, True)
    footer(slide, 2)
    return slide


def slide_temperature(prs, summary):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide(slide)
    title(slide, "02", "Gate A｜温度曲线接近，不等于高度映射", "2005–2006 重叠日对照：705 个配对日；日表测温高度未知")
    rows = []
    for key, label in (("TMAX_near", "近地面 → TMAX"), ("TMIN_near", "近地面 → TMIN"), ("TMAX_above", "冠层上方 → TMAX"), ("TMIN_above", "冠层上方 → TMIN")):
        m = summary["overlap_metrics"][key]
        rows.append([label, f"{m['bias_new_minus_cf']:+.3f}", f"{m['mae']:.3f}", f"{m['rmse']:.3f}", f"{m['pearson_r']:.5f}"])
    add_table(slide, 0.78, 2.06, [2.8, 1.35, 1.35, 1.35, 1.35], 0.53,
              ["候选映射", "Bias °C", "MAE °C", "RMSE °C", "Pearson r"], rows, 11, highlight_rows={0, 1})
    rect(slide, 9.25, 2.06, 3.35, 2.65, PALE_CORAL, radius=True)
    text(slide, "结论", 9.52, 2.32, 2.7, 0.35, 17, CORAL, True)
    text(slide, "近地面列误差较小，但无法证明它是 1.6 m。\n\n元数据只列出 1.6 m、2.9 m，未配对到导出字段。\n\n不选 TMAX/TMIN 主源。", 9.52, 2.8, 2.75, 1.6, 11, INK, valign=MSO_ANCHOR.TOP)
    rect(slide, 0.8, 5.0, 11.8, 1.08, PALE_TEAL, radius=True)
    text(slide, "Gate A  BLOCKED_TEMPERATURE_MAPPING", 1.02, 5.18, 5.5, 0.32, 16, TEAL, True)
    text(slide, "字段定义：冠层下 / 冠层上；参考日表：HMP45D，但无温度高度。数值相似仅作同站点产品一致性检查。", 1.03, 5.58, 11.2, 0.3, 11, INK)
    footer(slide, 3)
    return slide


def slide_rain(prs, summary, gaps):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide(slide)
    title(slide, "03", "Gate B｜30 个零值候选，仍不能认证为实测零", "35 个原缺口逐日追查；月/年累计量不拆分，不以缺测补零")
    rect(slide, 0.8, 2.05, 4.0, 3.98, WHITE, LINE, radius=True)
    text(slide, "35", 1.12, 2.34, 1.7, 0.85, 39, INK, True)
    text(slide, "原 unresolved 日", 2.75, 2.54, 1.7, 0.3, 12, MUTED, True)
    rect(slide, 1.15, 3.5, 3.18, 0.34, LINE, radius=True)
    rect(slide, 1.15, 3.5, 3.18 * 30 / 35, 0.34, GOLD, radius=True)
    text(slide, "30 天有官方日表数值候选（均为 0）", 1.14, 4.03, 3.28, 0.34, 11, INK, True)
    text(slide, "5 天无日源：2004-10-16 至 10-20", 1.14, 4.62, 3.25, 0.36, 11, CORAL, True)
    text(slide, "新日表说明含缺测插补，但无逐日实测/插补标志。", 1.14, 5.21, 3.27, 0.55, 10, MUTED, valign=MSO_ANCHOR.TOP)
    rain = summary["overlap_metrics"]["RAIN"]
    annual = rain["annual"]
    rows = []
    for year in ("2005", "2006"):
        a = annual[year]
        rows.append([year, f"{a['cf_total']:.1f}", f"{a['new_total']:.1f}", f"+{a['difference_new_minus_cf']:.1f}"])
    add_table(slide, 5.2, 2.1, [1.5, 2.0, 2.0, 1.8], 0.55,
              ["年份", "ChinaFLUX mm", "新日表 mm", "差值 mm"], rows, 10)
    text(slide, f"日配对：bias {rain['bias_new_minus_cf']:+.3f} mm/d  |  MAE {rain['mae']:.3f}  |  RMSE {rain['rmse']:.3f}  |  r={rain['pearson_r']:.4f}", 5.25, 3.66, 7.0, 0.45, 11, INK, True)
    text(slide, f"雨/无雨一致：{rain['rain_events']['event_agreement_days']}/730（{rain['rain_events']['event_agreement_rate']:.1%}）", 5.25, 4.18, 7.0, 0.32, 11, MUTED)
    rect(slide, 5.22, 4.85, 7.2, 0.98, PALE_GOLD, radius=True)
    text(slide, "年总量差异 2005 +8.2%  |  2006 +6.2%\n单位均为 mm，但量测/QC/插补差异原因未解。Gate B 认证通过 0/35。", 5.48, 5.03, 6.7, 0.62, 12, INK, True)
    footer(slide, 4)
    return slide


def slide_late_years(prs, qc):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide(slide)
    title(slide, "04", "Gate C｜2011–2013 有日表，不等于四变量完整", "新官方日表：RAIN 完整；温度/辐射的缺测与物理异常原样保留")
    rows = []
    for year, detail in qc["years"].items():
        missing = detail["missing_count_by_variable"]
        invalid = detail["invalid_dates_by_variable"].get("SRAD", [])
        rows.append([year, str(missing["RAIN"]), str(missing["TMAX"]), str(missing["TMIN"]), str(missing["SRAD"]), str(len(invalid))])
    add_table(slide, 0.82, 2.08, [1.5, 1.45, 1.6, 1.6, 1.45, 1.55], 0.62,
              ["年份", "RAIN 缺失", "TMAX 缺失", "TMIN 缺失", "SRAD 缺失", "SRAD 负值"], rows, 11,
              highlight_rows={0, 2})
    rect(slide, 0.84, 4.56, 11.7, 1.28, PALE_CORAL, radius=True)
    text(slide, "影响最大的连续缺口", 1.08, 4.75, 2.8, 0.32, 14, CORAL, True)
    text(slide, "2011-05-18 至 06-03、2011-08-17 至 09-04：TMAX/TMIN 与 SRAD 同时缺失。\nSRAD 负值：2011-03-09、2011-04-20、2013-01-02；未截为 0。", 3.75, 4.69, 8.3, 0.78, 12, INK)
    text(slide, "总辐射定义与单位可确认（MJ/m²，区别于净辐射/PAR）；完整性仍不通过。", 0.95, 6.22, 11.4, 0.35, 12, TEAL, True)
    footer(slide, 5)
    return slide


def slide_decision(prs, summary):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide(slide)
    title(slide, "05", "审计停在数据门槛，不进入 WGEN", "所有 3,653 天均保留变量级来源或 unresolved；验证期天气值未读取")
    counts = [("ChinaFLUX 完整日", 2514, TEAL), ("2005–2022 日表", 1126, GOLD), ("旧原始值", 8, BLUE), ("降水无值", 5, CORAL)]
    x0, width, y = 0.95, 11.35, 2.33
    total = sum(v for _, v, _ in counts)
    cursor = x0
    for label, value, color in counts:
        w = width * value / total
        rect(slide, cursor, y, max(w, 0.04), 0.45, color)
        if w > 0.6:
            text(slide, str(value), cursor, y + 0.03, w, 0.35, 10, WHITE, True, align=PP_ALIGN.CENTER)
        cursor += w
    for i, (label, value, color) in enumerate(counts):
        xx = 0.96 + (i % 2) * 5.75
        yy = 3.05 + (i // 2) * 0.52
        rect(slide, xx, yy + 0.03, 0.16, 0.16, color)
        text(slide, f"{label}：{value:,} 天", xx + 0.28, yy - 0.01, 4.9, 0.26, 11, INK, True)
    rect(slide, 0.9, 4.42, 11.55, 1.18, PALE_CORAL, radius=True)
    text(slide, "最终门槛", 1.15, 4.62, 1.7, 0.35, 15, CORAL, True)
    text(slide, "BLOCKED_TEMPERATURE_MAPPING", 3.0, 4.58, 5.6, 0.4, 17, INK, True)
    text(slide, "次级：降水缺口未认证  ·  2011–2013 不完整  ·  降水年总量不一致", 3.02, 5.03, 8.9, 0.3, 11, MUTED)
    text(slide, "最小后续：确认温度高度配对；索取逐日 QC/插补标志；补充 2004 缺口日源；修复 2011–2013 TMAX/TMIN/SRAD 缺测与负值。", 1.0, 6.05, 11.4, 0.55, 12, INK, True)
    footer(slide, 6)
    return slide


def main():
    summary, qc, gaps = load_data()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "YC 多源官方天气重建审计 003_04"
    prs.core_properties.subject = "YC/YCA 2004-2013 train-only weather source gates"
    slide_one(prs, summary)
    slide_sources(prs)
    slide_temperature(prs, summary)
    slide_rain(prs, summary, gaps)
    slide_late_years(prs, qc)
    slide_decision(prs, summary)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
