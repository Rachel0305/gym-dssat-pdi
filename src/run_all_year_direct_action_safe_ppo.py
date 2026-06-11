from __future__ import annotations

import argparse
import json
import shutil
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import gymnasium as gym
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import yaml

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None
    RGBColor = None
    PP_ALIGN = None
    Inches = None
    Pt = None

from ppo_action_safety import ActionSafetyState, apply_action_safety, normalize_action, update_action_safety_state
from ppo_evaluate import latest_observation_dict, scalar
from ppo_safe_rendering import PROJECT_ROOT, build_env_args, ensure_project_on_path


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_all_year_direct_action_safe_ppo.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "all_year_direct_action_safe_ppo"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_all_year_direct_action_safe_ppo_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_all_year_direct_action_safe_ppo_report.pptx"

STATION_NAMES = {
    "HLA": "Hailun",
    "SYA": "Shenyang",
    "LCA": "Luancheng",
    "YCA": "Yucheng",
    "FQA": "Fengqiu",
}


@dataclass
class EnvYear:
    station_code: str
    year: int
    planting_date: str
    scenario_type: str


def ensure_dirs() -> None:
    for sub in [
        "configs",
        "rendered_inputs",
        "models",
        "logs",
        "tensorboard",
        "daily_outputs",
        "evaluation",
        "figures/four_panel",
        "figures/weather",
        "figures/stress",
        "figures/actions",
        "figures/growth",
        "figures/summary",
        "reports",
    ]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)


def load_yaml(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def write_yaml(data: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        yaml.safe_dump(data, f, sort_keys=False, allow_unicode=True)


def df_to_markdown(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df.empty:
        return ""
    work = df.copy()
    if max_rows is not None:
        work = work.head(max_rows)
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(4)
    for col in work.columns:
        work[col] = work[col].map(lambda x: "" if pd.isna(x) else str(x))
    header = "| " + " | ".join(map(str, work.columns)) + " |"
    sep = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, sep, *rows])


def truthy(value: Any) -> bool:
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if pd.isna(value):
        return False
    return str(value).strip().lower() in {"true", "1", "yes", "y"}


def representative_phenology() -> dict[str, dict[str, int]]:
    path = PROJECT_ROOT / "data" / "observed_phenology_dates_standardized.csv"
    pheno = pd.read_csv(path, parse_dates=["planting_date", "harvest_date"])
    reps: dict[str, dict[str, int]] = {}
    for station, group in pheno.groupby("station_code"):
        planting_doy = int(round(group["planting_date"].dt.dayofyear.median()))
        harvest_days = int(round((group["harvest_date"] - group["planting_date"]).dt.days.median() + 1))
        reps[str(station)] = {"planting_doy": planting_doy, "harvest_days": harvest_days}
    return reps


def date_from_doy(year: int, doy: int) -> pd.Timestamp:
    return pd.Timestamp(year=int(year), month=1, day=1) + pd.Timedelta(days=int(doy) - 1)


def build_env_config(config: dict, selection: pd.DataFrame) -> dict:
    base = load_yaml(PROJECT_ROOT / config["paths"]["base_env_config"])
    base["paths"]["output_root"] = config["paths"]["output_root"]
    base["paths"]["wth_dir"] = config["paths"]["wth_dir"]
    base["paths"]["my_data_dir"] = config["paths"]["my_data_dir"]
    base["paths"]["cultivar_file"] = config["paths"]["cultivar_file"]
    base["runtime"]["mode"] = config["runtime"]["mode"]
    base["runtime"]["max_steps"] = int(config["runtime"]["max_steps"])
    base["seed"] = int(config["seed"])
    base["action_safety"] = dict(config["action_safety"])
    reps = representative_phenology()
    observed_years: dict[str, list[dict[str, Any]]] = {}
    for row in selection.drop_duplicates(["station_code", "year"]).itertuples(index=False):
        station = str(row.station_code)
        year = int(row.year)
        planting = date_from_doy(year, reps[station]["planting_doy"])
        observed_years.setdefault(station, []).append(
            {
                "year": year,
                "label": str(row.scenario_type),
                "planting_date": planting.strftime("%Y-%m-%d"),
                "weather_file": str(row.weather_file),
                "source": "006_17_all_year_direct_action_safe_ppo",
            }
        )
    base["observed_years"] = observed_years
    return base


def find_year(env_config: dict, station: str, year: int) -> dict:
    for item in env_config["observed_years"].get(station, []):
        if int(item["year"]) == int(year):
            return item
    raise KeyError(f"{station} {year} not found in env_config")


def score_selection_rows(group: pd.DataFrame) -> pd.DataFrame:
    out = group.copy()
    out["priority_score"] = (
        out["irrigation_responsive"].map(truthy).astype(int) * 100
        + out["has_water_stress"].map(truthy).astype(int) * 80
        + out["recommended_for_calibration"].map(truthy).astype(int) * 30
        + out["recommended_for_validation"].map(truthy).astype(int) * 25
        + out["recommended_for_ppo_train"].map(truthy).astype(int) * 10
        + out["recommended_for_ppo_eval"].map(truthy).astype(int) * 8
        + out["has_nitrogen_stress"].map(truthy).astype(int) * 5
        + out["scenario_type"].astype(str).str.contains("normal_year", na=False).astype(int) * 3
        + out["scenario_type"].astype(str).str.contains("dry_year|wet_year", regex=True, na=False).astype(int) * 2
    )
    return out


def select_train_eval_years(config: dict) -> pd.DataFrame:
    pool = pd.read_csv(PROJECT_ROOT / config["paths"]["scenario_pool_csv"])
    rows: list[dict[str, Any]] = []
    n_train = int(config["runtime"]["train_years_per_station"])
    n_eval = int(config["runtime"]["eval_years_per_station"])
    for station, group in pool.groupby("station_code"):
        group = score_selection_rows(group)
        train_candidates = group[group["recommended_for_ppo_train"].map(truthy)].sort_values(
            ["priority_score", "year"], ascending=[False, True]
        )
        train_idx = list(train_candidates.head(n_train).index)
        if len(train_idx) < min(n_train, len(group)):
            for idx in group.sort_values(["priority_score", "year"], ascending=[False, True]).index:
                if idx not in train_idx:
                    train_idx.append(idx)
                if len(train_idx) >= n_train:
                    break
        eval_candidates = group.drop(index=train_idx, errors="ignore")
        eval_candidates = eval_candidates[
            eval_candidates["recommended_for_ppo_eval"].map(truthy)
            | eval_candidates["irrigation_responsive"].map(truthy)
            | eval_candidates["has_water_stress"].map(truthy)
            | eval_candidates["recommended_for_validation"].map(truthy)
        ].sort_values(["priority_score", "year"], ascending=[False, True])
        eval_idx = list(eval_candidates.head(n_eval).index)
        if len(eval_idx) < min(n_eval, len(group) - len(set(train_idx))):
            for idx in group.drop(index=train_idx, errors="ignore").sort_values(["priority_score", "year"], ascending=[False, True]).index:
                if idx not in eval_idx:
                    eval_idx.append(idx)
                if len(eval_idx) >= n_eval:
                    break
        for idx, row in group.iterrows():
            reason = []
            if idx in train_idx:
                reason.append("selected_train_priority_all_year_pool")
            if idx in eval_idx:
                reason.append("selected_eval_priority_all_year_pool")
            if truthy(row.get("has_water_stress")):
                reason.append("water_stress_year")
            if truthy(row.get("irrigation_responsive")):
                reason.append("irrigation_responsive_year")
            if truthy(row.get("has_nitrogen_stress")):
                reason.append("nitrogen_stress_year")
            if int(row["year"]) in {2020, 2021, 2022, 2023}:
                reason.append("2020_2023_calibration_validation_candidate")
            data = row.to_dict()
            data["selected_for_train"] = idx in train_idx
            data["selected_for_eval"] = idx in eval_idx
            data["selection_reason"] = ";".join(dict.fromkeys(reason))
            rows.append(data)
    selection = pd.DataFrame(rows).sort_values(["station_code", "year"]).reset_index(drop=True)
    cols = [
        "station_code",
        "station_name",
        "year",
        "weather_file",
        "scenario_type",
        "has_water_stress",
        "has_nitrogen_stress",
        "irrigation_responsive",
        "recommended_for_ppo_train",
        "recommended_for_ppo_eval",
        "selected_for_train",
        "selected_for_eval",
        "selection_reason",
        "growing_season_rain",
        "max_swfac",
        "nstres_days_gt_0p05",
        "yield_gain_from_irrigation_at_same_N",
    ]
    selection = selection[[c for c in cols if c in selection.columns]]
    out_csv = OUTPUT_ROOT / "configs" / "ppo_train_eval_year_selection.csv"
    selection.to_csv(out_csv, index=False, encoding="utf-8-sig")
    selected = selection[selection["selected_for_train"] | selection["selected_for_eval"]].copy()
    lines = [
        "# PPO Train/Eval Year Selection",
        "",
        "Selection is based on the 006_15 all-year weather scenario pool.",
        "Priority: water-stress, irrigation-responsive, nitrogen-stress, 2020-2023 calibration/validation candidates, and dry/normal/wet mix.",
        "",
        df_to_markdown(
            selected[
                [
                    "station_code",
                    "year",
                    "scenario_type",
                    "selected_for_train",
                    "selected_for_eval",
                    "selection_reason",
                    "growing_season_rain",
                    "max_swfac",
                ]
            ],
            80,
        ),
    ]
    (OUTPUT_ROOT / "configs" / "ppo_train_eval_year_selection.md").write_text("\n".join(lines), encoding="utf-8")
    return selection


def ppo_kwargs(config: dict) -> dict:
    allowed = ["learning_rate", "gamma", "n_steps", "batch_size", "ent_coef", "clip_range"]
    return {k: config.get("ppo", {}).get(k) for k in allowed if config.get("ppo", {}).get(k) is not None}


def physical_from_policy_action(config: dict, action, action_names: list[str]) -> dict[str, float]:
    arr = np.asarray(action, dtype=float).flatten()
    out: dict[str, float] = {}
    scale = config["action_scale"]
    for name, value in zip(action_names, arr):
        value01 = float(np.clip((float(value) + 1.0) / 2.0, 0.0, 1.0))
        if name == "amir":
            out[name] = value01 * float(scale["daily_irrigation_max"])
        elif name == "anfer":
            out[name] = value01 * float(scale["daily_n_max"])
        else:
            out[name] = value01
    return out


class DirectActionSafeGrowthRewardWrapper(gym.Env):
    def __init__(self, env, config: dict):
        super().__init__()
        self.env = env
        self.config = config
        self.action_space = env.action_space
        self.observation_space = env.observation_space
        self.metadata = getattr(env, "metadata", {})
        self.safety_state = ActionSafetyState()
        self.last_action_info: dict[str, Any] = {}
        self.last_obs_dict: dict[str, Any] = {}

    def reset(self, *args, **kwargs):
        self.safety_state = ActionSafetyState()
        self.last_action_info = {}
        result = self.env.reset(*args, **kwargs)
        obs, info = result if isinstance(result, tuple) else (result, {})
        self.last_obs_dict = latest_observation_dict(self.env, obs, info)
        return obs, info

    def step(self, action):
        prev = dict(self.last_obs_dict)
        dap_raw = scalar(prev.get("dap", 0))
        dap = int(round(dap_raw)) if np.isfinite(dap_raw) and dap_raw > 0 else 1
        action_names = list(self.env.formator.action_names)
        scaled_real = physical_from_policy_action(self.config, action, action_names)
        safety = {**self.config["action_safety"], "enabled": True}
        safety_result = apply_action_safety(scaled_real, dap, self.safety_state, safety)
        safe_norm = normalize_action(action_names, self.env.formator.action_space_dict, safety_result.safe_real_action)
        obs, _env_reward, terminated, truncated, info = self.env.step(safe_norm)
        update_action_safety_state(self.safety_state, safety_result.safe_real_action, dap)
        latest = latest_observation_dict(self.env, obs, info)
        self.last_obs_dict = latest
        topwt_delta = max(0.0, scalar(latest.get("topwt"), 0.0) - scalar(prev.get("topwt"), 0.0))
        grnwt_delta = max(0.0, scalar(latest.get("grnwt"), 0.0) - scalar(prev.get("grnwt"), 0.0))
        reward_cfg = self.config["reward"]
        safe_i = float(safety_result.safe_real_action.get("amir", 0.0))
        safe_n = float(safety_result.safe_real_action.get("anfer", 0.0))
        growth_reward = float(reward_cfg["topwt_delta_coef"]) * topwt_delta + float(reward_cfg["grnwt_delta_coef"]) * grnwt_delta
        reward = growth_reward - float(reward_cfg["water_cost"]) * safe_i - float(reward_cfg["nitrogen_cost"]) * safe_n
        raw_arr = np.asarray(action, dtype=float).flatten()
        raw_norm = {name: float(value) for name, value in zip(action_names, raw_arr)}
        safe_norm_by_name = {name: float(value) for name, value in zip(action_names, np.asarray(safe_norm).flatten())}
        self.last_action_info = {
            "raw_action_amir": raw_norm.get("amir", np.nan),
            "raw_action_anfer": raw_norm.get("anfer", np.nan),
            "scaled_action_amir": float(scaled_real.get("amir", 0.0)),
            "scaled_action_anfer": float(scaled_real.get("anfer", 0.0)),
            "safe_action_amir": safe_i,
            "safe_action_anfer": safe_n,
            "safe_normalized_action_amir": safe_norm_by_name.get("amir", np.nan),
            "safe_normalized_action_anfer": safe_norm_by_name.get("anfer", np.nan),
            "season_cumulative_irrigation": float(self.safety_state.cumulative_irrigation),
            "season_cumulative_n": float(self.safety_state.cumulative_n),
            "action_safety_triggered": safety_result.safety_rule_triggered,
            "growth_reward": growth_reward,
            "topwt_delta": topwt_delta,
            "grnwt_delta": grnwt_delta,
        }
        return obs, float(reward), terminated, truncated, info

    def close(self):
        return self.env.close()

    @property
    def unwrapped(self):
        return self.env.unwrapped

    def __getattr__(self, name):
        return getattr(self.env, name)


def make_base_env(env_config: dict, station: str, year: int, seed: int, run_tag: str, evaluation: bool):
    ensure_project_on_path()
    import gym
    from sb3_wrapper import GymDssatWrapper

    year_info = find_year(env_config, station, year)
    env_args = build_env_args(
        station=station,
        year=year,
        planting_date=year_info["planting_date"],
        seed=seed,
        config=env_config,
        run_tag=run_tag,
        evaluation=evaluation,
        mode=env_config.get("runtime", {}).get("mode", "all"),
    )
    return GymDssatWrapper(gym.make("gym_dssat_pdi:GymDssatPdi-v0", **env_args).unwrapped)


def make_training_env(config: dict, env_config: dict, station: str, year: int, seed: int, run_tag: str):
    return DirectActionSafeGrowthRewardWrapper(make_base_env(env_config, station, year, seed, run_tag, evaluation=False), config)


def train_station_models(config: dict, env_config: dict, selection: pd.DataFrame, debug: bool = False) -> pd.DataFrame:
    from stable_baselines3 import PPO

    rows: list[dict[str, Any]] = []
    seed = int(config["seed"])
    total_timesteps = int(config["runtime"]["debug_timesteps"] if debug else config["total_timesteps"])
    stations = [str(x) for x in sorted(selection.loc[selection["selected_for_train"], "station_code"].unique())]
    if debug:
        stations = [str(config["runtime"].get("smoke_station", "HLA"))]
    for station in stations:
        train_years = [
            int(y)
            for y in selection[(selection["station_code"].eq(station)) & (selection["selected_for_train"])]["year"].tolist()
        ]
        model_path = OUTPUT_ROOT / "models" / station / "ppo_direct_action_safe_seed0.zip"
        model_path.parent.mkdir(parents=True, exist_ok=True)
        tensorboard_dir = OUTPUT_ROOT / "tensorboard" / station
        tensorboard_dir.mkdir(parents=True, exist_ok=True)
        env = None
        model = None
        status = "ok"
        notes = ""
        try:
            per_year_steps = max(1, int(np.ceil(total_timesteps / max(len(train_years), 1))))
            for idx, year in enumerate(train_years):
                env = make_training_env(config, env_config, station, year, seed, f"{station}_{year}_direct_train")
                if model is None:
                    model = PPO(
                        "MlpPolicy",
                        env,
                        verbose=1,
                        seed=seed,
                        tensorboard_log=str(tensorboard_dir),
                        **ppo_kwargs(config),
                    )
                else:
                    model.set_env(env)
                model.learn(
                    total_timesteps=per_year_steps,
                    reset_num_timesteps=(idx == 0),
                    progress_bar=False,
                )
                env.close()
                env = None
            if model is None:
                status = "failed"
                notes = "no_train_years"
            else:
                model.save(str(model_path.with_suffix("")))
        except Exception:
            status = "failed"
            notes = traceback.format_exc()
        finally:
            if env is not None:
                try:
                    env.close()
                except Exception:
                    pass
        rows.append(
            {
                "station_code": station,
                "train_years": ",".join(map(str, train_years)),
                "seed": seed,
                "total_timesteps": total_timesteps,
                "run_status": status,
                "model_path": str(model_path.relative_to(PROJECT_ROOT)) if model_path.exists() else "",
                "mean_train_reward": np.nan,
                "final_train_reward": np.nan,
                "notes": notes[-1500:] if notes else "",
            }
        )
    summary = pd.DataFrame(rows)
    summary.to_csv(OUTPUT_ROOT / "evaluation" / "training_run_summary.csv", index=False, encoding="utf-8-sig")
    return summary


def weather_for_daily(config: dict) -> pd.DataFrame:
    path = PROJECT_ROOT / config["paths"]["weather_clean_csv"]
    weather = pd.read_csv(path)
    weather["date"] = pd.to_datetime(weather["date"])
    return weather.rename(columns={"station": "station_code", "RAIN": "rain", "SRAD": "srad", "TMAX": "tmax", "TMIN": "tmin"})


def evaluate_one(config: dict, env_config: dict, model, station: str, year: int, split: str, model_path: Path, weather: pd.DataFrame) -> dict:
    seed = int(config["seed"])
    env = make_training_env(config, env_config, station, year, seed, f"{station}_{year}_direct_eval")
    records: list[dict[str, Any]] = []
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        year_info = find_year(env_config, station, year)
        planting = pd.Timestamp(year_info["planting_date"])
        while not done and step_count < int(config["runtime"]["max_steps"]):
            latest = latest_observation_dict(env, obs, info)
            dap_raw = scalar(latest.get("dap", step_count + 1))
            dap = int(round(dap_raw)) if np.isfinite(dap_raw) and dap_raw > 0 else step_count + 1
            action, _ = model.predict(obs, deterministic=True)
            obs, reward, terminated, truncated, info = env.step(action)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            action_info = dict(env.last_action_info)
            date = planting + pd.Timedelta(days=max(dap - 1, 0))
            w = weather[(weather["station_code"].eq(station)) & (weather["date"].eq(date))]
            wrow = w.iloc[0].to_dict() if len(w) else {}
            records.append(
                {
                    "station_code": station,
                    "year": int(year),
                    "date": date.strftime("%Y-%m-%d"),
                    "doy": int(date.dayofyear),
                    "dap": dap,
                    "rain": scalar(wrow.get("rain"), np.nan),
                    "srad": scalar(wrow.get("srad"), np.nan),
                    "tmax": scalar(wrow.get("tmax"), np.nan),
                    "tmin": scalar(wrow.get("tmin"), np.nan),
                    "swfac": scalar(latest.get("swfac")),
                    "nstres": scalar(latest.get("nstres")),
                    "topwt": scalar(latest.get("topwt")),
                    "grnwt": scalar(latest.get("grnwt")),
                    "xlai": scalar(latest.get("xlai")),
                    "reward": float(reward),
                    **action_info,
                    "done": done,
                    "info": json.dumps(info, ensure_ascii=False, default=str),
                }
            )
            step_count += 1
    finally:
        env.close()
    daily = pd.DataFrame(records)
    daily_dir = OUTPUT_ROOT / "daily_outputs" / station
    daily_dir.mkdir(parents=True, exist_ok=True)
    daily_path = daily_dir / f"{year}_ppo_daily.csv"
    if split == "train" and daily_path.exists():
        daily_path = daily_dir / f"{year}_train_ppo_daily.csv"
    daily.to_csv(daily_path, index=False, encoding="utf-8-sig")
    swfac = pd.to_numeric(daily.get("swfac", pd.Series(dtype=float)), errors="coerce")
    nstres = pd.to_numeric(daily.get("nstres", pd.Series(dtype=float)), errors="coerce")
    irr = pd.to_numeric(daily.get("safe_action_amir", pd.Series(dtype=float)), errors="coerce").fillna(0)
    n = pd.to_numeric(daily.get("safe_action_anfer", pd.Series(dtype=float)), errors="coerce").fillna(0)
    final_grnwt = float(pd.to_numeric(daily.get("grnwt", pd.Series(dtype=float)), errors="coerce").iloc[-1]) if len(daily) else np.nan
    final_topwt = float(pd.to_numeric(daily.get("topwt", pd.Series(dtype=float)), errors="coerce").iloc[-1]) if len(daily) else np.nan
    total_i = float(irr.sum())
    total_n = float(n.sum())
    econ = config["economics"]
    profit_simple = float(econ["grain_value_coef"]) * final_grnwt - float(econ["water_cost"]) * total_i - float(econ["nitrogen_cost"]) * total_n if np.isfinite(final_grnwt) else np.nan
    return {
        "station_code": station,
        "station_name": STATION_NAMES.get(station, station),
        "year": int(year),
        "split": split,
        "scenario_type": str(find_year(env_config, station, year).get("label", "")),
        "run_status": "ok" if len(daily) and bool(daily["done"].iloc[-1]) else "failed",
        "episode_completed": bool(len(daily) and bool(daily["done"].iloc[-1])),
        "episode_length": int(len(daily)),
        "final_grnwt": final_grnwt,
        "final_topwt": final_topwt,
        "max_xlai": float(pd.to_numeric(daily.get("xlai", pd.Series(dtype=float)), errors="coerce").max()) if len(daily) else np.nan,
        "total_irrigation": total_i,
        "total_n": total_n,
        "profit_simple": profit_simple,
        "mean_swfac": float(swfac.mean()) if len(swfac) else np.nan,
        "max_swfac": float(swfac.max()) if len(swfac) else np.nan,
        "swfac_stress_days_gt_0p05": int((swfac > 0.05).sum()) if len(swfac) else 0,
        "mean_nstres": float(nstres.mean()) if len(nstres) else np.nan,
        "max_nstres": float(nstres.max()) if len(nstres) else np.nan,
        "nstres_days_gt_0p05": int((nstres > 0.05).sum()) if len(nstres) else 0,
        "irrigation_event_count": int((irr > 0).sum()),
        "n_event_count": int((n > 0).sum()),
        "first_irrigation_dap": int(daily.loc[irr > 0, "dap"].iloc[0]) if (irr > 0).any() else np.nan,
        "first_n_dap": int(daily.loc[n > 0, "dap"].iloc[0]) if (n > 0).any() else np.nan,
        "model_path": str(model_path.relative_to(PROJECT_ROOT)),
        "daily_csv_path": str(daily_path.relative_to(PROJECT_ROOT)),
        "notes": "",
    }


def evaluate_models(config: dict, env_config: dict, selection: pd.DataFrame, training_summary: pd.DataFrame) -> pd.DataFrame:
    from stable_baselines3 import PPO

    weather = weather_for_daily(config)
    rows: list[dict[str, Any]] = []
    for train_row in training_summary.itertuples(index=False):
        if str(train_row.run_status) != "ok":
            continue
        station = str(train_row.station_code)
        model_path = PROJECT_ROOT / str(train_row.model_path)
        model = PPO.load(str(model_path))
        station_sel = selection[selection["station_code"].eq(station)]
        years = []
        for split_name in ["train", "eval"]:
            col = "selected_for_train" if split_name == "train" else "selected_for_eval"
            for year in station_sel[station_sel[col]]["year"].astype(int).tolist():
                years.append((split_name, year))
        for split, year in years:
            try:
                rows.append(evaluate_one(config, env_config, model, station, int(year), split, model_path, weather))
            except Exception:
                rows.append(
                    {
                        "station_code": station,
                        "station_name": STATION_NAMES.get(station, station),
                        "year": int(year),
                        "split": split,
                        "run_status": "failed",
                        "episode_completed": False,
                        "model_path": str(model_path.relative_to(PROJECT_ROOT)),
                        "daily_csv_path": "",
                        "notes": traceback.format_exc()[-1500:],
                    }
                )
    summary = pd.DataFrame(rows)
    summary.to_csv(OUTPUT_ROOT / "evaluation" / "ppo_direct_eval_summary.csv", index=False, encoding="utf-8-sig")
    return summary


def preplant_weather(config: dict, station: str, year: int, planting_date: str) -> pd.DataFrame:
    weather = weather_for_daily(config)
    planting = pd.Timestamp(planting_date)
    end = planting + pd.Timedelta(days=int(config["runtime"]["max_steps"]) - 1)
    start = planting - pd.Timedelta(days=30)
    out = weather[(weather["station_code"].eq(station)) & (weather["date"] >= start) & (weather["date"] <= end)].copy()
    out["dap"] = (out["date"] - planting).dt.days + 1
    return out


def plot_representative_outputs(config: dict, env_config: dict, eval_summary: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for row in eval_summary[eval_summary["run_status"].eq("ok")].itertuples(index=False):
        station = str(row.station_code)
        year = int(row.year)
        daily_path = PROJECT_ROOT / str(row.daily_csv_path)
        if not daily_path.exists():
            continue
        daily = pd.read_csv(daily_path)
        year_info = find_year(env_config, station, year)
        weather = preplant_weather(config, station, year, year_info["planting_date"])
        stem = f"{station}_{year}"
        if not weather.empty:
            fig, ax1 = plt.subplots(figsize=(9, 4))
            ax1.bar(weather["dap"], weather["rain"], color="#4F81BD", alpha=0.75, label="daily rain")
            ax1.set_xlabel("DAP")
            ax1.set_ylabel("Rain (mm/day)")
            ax2 = ax1.twinx()
            ax2.plot(weather["dap"], weather["rain"].cumsum(), color="#C0504D", label="cumulative rain")
            ax2.set_ylabel("Cumulative rain (mm)")
            ax1.axvline(1, color="black", linewidth=0.8)
            fig.tight_layout()
            fig.savefig(OUTPUT_ROOT / "figures" / "weather" / f"{stem}_weather_rainfall.png", dpi=180)
            plt.close(fig)
        fig, ax = plt.subplots(figsize=(9, 4))
        ax.plot(daily["dap"], daily["swfac"], label="SWFAC", color="#4F81BD")
        ax.plot(daily["dap"], daily["nstres"], label="NSTRES", color="#C0504D")
        ax.axhline(0.05, color="black", linewidth=0.8, linestyle="--")
        ax.set_xlabel("DAP")
        ax.set_ylabel("Stress index")
        ax.legend()
        ax.grid(axis="y", alpha=0.25)
        fig.tight_layout()
        fig.savefig(OUTPUT_ROOT / "figures" / "stress" / f"{stem}_swfac_nstres.png", dpi=180)
        plt.close(fig)
        fig, ax1 = plt.subplots(figsize=(9, 4))
        ax1.bar(daily["dap"], daily["safe_action_amir"], color="#4F81BD", alpha=0.75, label="irrigation")
        ax1.set_ylabel("Irrigation (mm)")
        ax2 = ax1.twinx()
        ax2.bar(daily["dap"], daily["safe_action_anfer"], color="#9BBB59", alpha=0.45, label="N")
        ax2.set_ylabel("N fertilizer (kg/ha)")
        ax1.set_xlabel("DAP")
        fig.tight_layout()
        fig.savefig(OUTPUT_ROOT / "figures" / "actions" / f"{stem}_ppo_actions.png", dpi=180)
        plt.close(fig)
        fig, ax = plt.subplots(figsize=(9, 4))
        ax.plot(daily["dap"], daily["topwt"], label="TOPWT", color="#4F81BD")
        ax.plot(daily["dap"], daily["grnwt"], label="GRNWT", color="#C0504D")
        ax.set_xlabel("DAP")
        ax.set_ylabel("Biomass / grain weight")
        ax.legend()
        ax.grid(axis="y", alpha=0.25)
        fig.tight_layout()
        fig.savefig(OUTPUT_ROOT / "figures" / "growth" / f"{stem}_topwt_grnwt.png", dpi=180)
        plt.close(fig)
        fig, axes = plt.subplots(4, 1, figsize=(10, 11), sharex=False)
        if not weather.empty:
            axes[0].bar(weather["dap"], weather["rain"], color="#4F81BD")
        axes[0].set_title("Actual rainfall from one month before planting")
        axes[0].set_ylabel("Rain")
        axes[1].plot(daily["dap"], daily["swfac"], label="SWFAC")
        axes[1].plot(daily["dap"], daily["nstres"], label="NSTRES")
        axes[1].axhline(0.05, color="black", linestyle="--", linewidth=0.8)
        axes[1].set_title("SWFAC and NSTRES dynamics")
        axes[1].legend()
        axes[2].bar(daily["dap"], daily["safe_action_amir"], label="Irrigation", color="#4F81BD", alpha=0.75)
        axes[2].bar(daily["dap"], daily["safe_action_anfer"], label="N fertilizer", color="#9BBB59", alpha=0.55)
        axes[2].set_title("PPO irrigation and fertilization actions")
        axes[2].legend()
        axes[3].plot(daily["dap"], daily["topwt"], label="TOPWT")
        axes[3].plot(daily["dap"], daily["grnwt"], label="GRNWT")
        axes[3].set_title("Crop growth response under PPO strategy")
        axes[3].set_xlabel("DAP")
        axes[3].legend()
        for ax in axes:
            ax.grid(axis="y", alpha=0.2)
        fig.tight_layout()
        fig.savefig(OUTPUT_ROOT / "figures" / "four_panel" / f"{stem}_four_panel.png", dpi=180)
        plt.close(fig)
        rows.append({"station_code": station, "year": year, "four_panel_path": str((OUTPUT_ROOT / "figures" / "four_panel" / f"{stem}_four_panel.png").relative_to(PROJECT_ROOT))})
    return pd.DataFrame(rows)


def build_reasonableness(config: dict, eval_summary: pd.DataFrame) -> pd.DataFrame:
    rows = []
    i_cap = float(config["action_safety"]["season_irrigation_soft_limit"])
    n_cap = float(config["action_safety"]["season_n_soft_limit"])
    for row in eval_summary.itertuples(index=False):
        if str(row.run_status) != "ok":
            continue
        daily_path = PROJECT_ROOT / str(row.daily_csv_path)
        daily = pd.read_csv(daily_path) if daily_path.exists() else pd.DataFrame()
        swfac = pd.to_numeric(daily.get("swfac", pd.Series(dtype=float)), errors="coerce")
        nstres = pd.to_numeric(daily.get("nstres", pd.Series(dtype=float)), errors="coerce")
        peak_swfac_dap = int(daily.loc[swfac.idxmax(), "dap"]) if len(daily) and swfac.notna().any() else np.nan
        peak_nstres_dap = int(daily.loc[nstres.idxmax(), "dap"]) if len(daily) and nstres.notna().any() else np.nan
        irr_days = set(pd.to_numeric(daily.loc[pd.to_numeric(daily.get("safe_action_amir", 0), errors="coerce").fillna(0) > 0, "dap"], errors="coerce").dropna().astype(int).tolist()) if len(daily) else set()
        n_days = set(pd.to_numeric(daily.loc[pd.to_numeric(daily.get("safe_action_anfer", 0), errors="coerce").fillna(0) > 0, "dap"], errors="coerce").dropna().astype(int).tolist()) if len(daily) else set()
        irr_near = bool(np.isfinite(peak_swfac_dap) and any(abs(d - peak_swfac_dap) <= 7 or d <= peak_swfac_dap for d in irr_days))
        n_near = bool(np.isfinite(peak_nstres_dap) and any(abs(d - peak_nstres_dap) <= 7 or d <= peak_nstres_dap for d in n_days))
        cap_sat = float(row.total_irrigation) >= 0.95 * i_cap or float(row.total_n) >= 0.95 * n_cap
        if cap_sat:
            label = "cap_saturated"
        elif (int(row.swfac_stress_days_gt_0p05) > 0 and not irr_near) or (int(row.nstres_days_gt_0p05) > 0 and not n_near):
            label = "questionable"
        else:
            label = "reasonable"
        rows.append(
            {
                "station_code": row.station_code,
                "year": int(row.year),
                "scenario_type": getattr(row, "scenario_type", ""),
                "total_irrigation": row.total_irrigation,
                "total_n": row.total_n,
                "irrigation_event_count": row.irrigation_event_count,
                "n_event_count": row.n_event_count,
                "swfac_stress_days_gt_0p05": row.swfac_stress_days_gt_0p05,
                "nstres_days_gt_0p05": row.nstres_days_gt_0p05,
                "irrigation_during_or_before_swfac_stress": irr_near,
                "fertilization_during_or_before_nstres": n_near,
                "first_irrigation_dap": row.first_irrigation_dap,
                "first_n_dap": row.first_n_dap,
                "peak_swfac_dap": peak_swfac_dap,
                "peak_nstres_dap": peak_nstres_dap,
                "final_grnwt": row.final_grnwt,
                "final_topwt": row.final_topwt,
                "profit_simple": row.profit_simple,
                "decision_reasonableness_label": label,
                "notes": "",
            }
        )
    diagnosis = pd.DataFrame(rows)
    diagnosis.to_csv(OUTPUT_ROOT / "evaluation" / "ppo_decision_reasonableness_diagnosis.csv", index=False, encoding="utf-8-sig")
    return diagnosis


def plot_summary(eval_summary: pd.DataFrame, diagnosis: pd.DataFrame) -> None:
    fig_dir = OUTPUT_ROOT / "figures" / "summary"
    if eval_summary.empty or "run_status" not in eval_summary.columns:
        return
    ok = eval_summary[eval_summary["run_status"].eq("ok")].copy()
    if ok.empty:
        return
    ok["site_year"] = ok["station_code"] + "_" + ok["year"].astype(str)
    for col, filename, ylabel in [
        ("total_irrigation", "ppo_total_irrigation_by_station_year.png", "Total irrigation (mm)"),
        ("total_n", "ppo_total_n_by_station_year.png", "Total N (kg/ha)"),
        ("final_grnwt", "ppo_final_grnwt_by_station_year.png", "Final GRNWT"),
        ("profit_simple", "ppo_profit_simple_by_station_year.png", "Simple profit score"),
    ]:
        ax = ok.set_index("site_year")[col].plot(kind="bar", figsize=(11, 4.5), color="#4F81BD")
        ax.set_ylabel(ylabel)
        ax.grid(axis="y", alpha=0.25)
        ax.get_figure().tight_layout()
        ax.get_figure().savefig(fig_dir / filename, dpi=180)
        plt.close(ax.get_figure())
    stress = ok.set_index("site_year")[["swfac_stress_days_gt_0p05", "nstres_days_gt_0p05"]]
    ax = stress.plot(kind="bar", figsize=(11, 4.5), color=["#4F81BD", "#C0504D"])
    ax.set_ylabel("Stress days")
    ax.grid(axis="y", alpha=0.25)
    ax.get_figure().tight_layout()
    ax.get_figure().savefig(fig_dir / "ppo_swfac_nstres_stress_days_by_station_year.png", dpi=180)
    plt.close(ax.get_figure())
    if not diagnosis.empty:
        counts = diagnosis["decision_reasonableness_label"].value_counts()
        ax = counts.plot(kind="bar", figsize=(7, 4), color="#4F81BD")
        ax.set_ylabel("Cases")
        ax.grid(axis="y", alpha=0.25)
        ax.get_figure().tight_layout()
        ax.get_figure().savefig(fig_dir / "ppo_decision_reasonableness_summary.png", dpi=180)
        plt.close(ax.get_figure())


def build_report(config: dict, selection: pd.DataFrame, train_summary: pd.DataFrame, eval_summary: pd.DataFrame, diagnosis: pd.DataFrame, fig_index: pd.DataFrame) -> None:
    selected = selection[selection["selected_for_train"] | selection["selected_for_eval"]]
    ok_train = train_summary[train_summary["run_status"].eq("ok")] if "run_status" in train_summary.columns else pd.DataFrame()
    ok_eval = eval_summary[eval_summary["run_status"].eq("ok")] if "run_status" in eval_summary.columns else pd.DataFrame()
    cap_sat = diagnosis[diagnosis["decision_reasonableness_label"].eq("cap_saturated")] if not diagnosis.empty else pd.DataFrame()
    lines = [
        "# All-Year Direct Action-Safe PPO Simple Baseline Report",
        "",
        "## Scope",
        "",
        "This stage trains direct action-safe PPO from gym-DSSAT states. It does not use RF, behavior cloning, expert replay, offline schedule search, constrained PPO from prior, rainfall scaling, reward grids, episode-level profit reward, low-frequency action design, seasonal budget action, or scheduled event action.",
        "",
        "## Fixed Settings",
        "",
        f"- Timesteps per station: {config['total_timesteps']}",
        f"- Daily action scale: irrigation <= {config['action_scale']['daily_irrigation_max']} mm/day; N <= {config['action_scale']['daily_n_max']} kg/ha/day.",
        f"- Season action safety cap: irrigation <= {config['action_safety']['season_irrigation_soft_limit']} mm; N <= {config['action_safety']['season_n_soft_limit']} kg/ha.",
        f"- Reward: `0.1 * delta_topwt + delta_grnwt - 0.1 * irrigation - 0.25 * nitrogen`.",
        "",
        "## Selected Train/Eval Years",
        "",
        df_to_markdown(selected[["station_code", "year", "scenario_type", "selected_for_train", "selected_for_eval", "selection_reason"]], 80),
        "",
        "## Training Summary",
        "",
        df_to_markdown(train_summary),
        "",
        "## Evaluation Summary",
        "",
        df_to_markdown(eval_summary[[c for c in ["station_code", "year", "split", "run_status", "final_grnwt", "total_irrigation", "total_n", "profit_simple", "swfac_stress_days_gt_0p05", "nstres_days_gt_0p05", "daily_csv_path"] if c in eval_summary.columns]], 120) if not eval_summary.empty else "",
        "",
        "## Decision Reasonableness",
        "",
        df_to_markdown(diagnosis, 120) if not diagnosis.empty else "",
        "",
        "## Main Finding",
        "",
        f"- Successful station models: {len(ok_train)} / {len(train_summary)}.",
        f"- Successful evaluations: {len(ok_eval)} / {len(eval_summary)}.",
        f"- Cap-saturated cases: {len(cap_sat)}.",
        "- If many cases are cap_saturated or decisions do not occur near SWFAC/NSTRES stress, this simple direct PPO baseline should be reported as a baseline attempt, not as a final optimized policy.",
        "",
        "## Figure Outputs",
        "",
        df_to_markdown(fig_index, 120) if not fig_index.empty else "No figure index generated.",
    ]
    DOC_MD.write_text("\n".join(lines), encoding="utf-8")
    shutil.copyfile(DOC_MD, OUTPUT_ROOT / "reports" / DOC_MD.name)
    build_ppt(config, selected, train_summary, eval_summary, diagnosis, fig_index)
    if DOC_PPT.exists():
        shutil.copyfile(DOC_PPT, OUTPUT_ROOT / "reports" / DOC_PPT.name)


def add_ppt_title(slide, title: str):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.55))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = title
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_ppt_bullets(prs, title: str, bullets: list[str]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_ppt_title(slide, title)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(12.0), Inches(5.9))
    tf = box.text_frame
    tf.clear()
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0, 0, 0)


def add_ppt_table(prs, title: str, df: pd.DataFrame, max_rows: int = 12):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_ppt_title(slide, title)
    if df.empty:
        return
    work = df.head(max_rows).copy()
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(3)
    rows, cols = len(work), len(work.columns)
    table_shape = slide.shapes.add_table(rows + 1, cols, Inches(0.35), Inches(1.05), Inches(12.65), Inches(6.0))
    table = table_shape.table
    for j, col in enumerate(work.columns):
        cell = table.cell(0, j)
        cell.text = str(col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(68, 114, 196)
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(7)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(work.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = "" if pd.isna(value) else str(value)
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(226, 235, 247)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(6)
                    run.font.color.rgb = RGBColor(0, 0, 0)


def add_ppt_picture(prs, title: str, path: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_ppt_title(slide, title)
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(0.65), Inches(1.0), width=Inches(12.0))


def build_ppt(config: dict, selected: pd.DataFrame, train_summary: pd.DataFrame, eval_summary: pd.DataFrame, diagnosis: pd.DataFrame, fig_index: pd.DataFrame) -> None:
    if Presentation is None:
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_ppt_bullets(
        prs,
        "006_17 All-Year Direct Action-Safe PPO",
        [
            "Goal: train a simple direct PPO baseline using the 006_15 all-year scenario pool.",
            "Excluded: RF / behavior cloning / expert replay / offline schedule search / rainfall-scaling / reward tuning.",
            "Reward: 0.1*delta_TOPWT + delta_GRNWT - 0.1*irrigation - 0.25*N.",
            f"Action safety: season irrigation cap {config['action_safety']['season_irrigation_soft_limit']} mm; season N cap {config['action_safety']['season_n_soft_limit']} kg/ha.",
        ],
    )
    add_ppt_table(prs, "Selected Train/Eval Years", selected[["station_code", "year", "scenario_type", "selected_for_train", "selected_for_eval"]], 18)
    add_ppt_table(prs, "Training Summary", train_summary, 8)
    if not eval_summary.empty:
        add_ppt_table(prs, "Evaluation Summary", eval_summary[["station_code", "year", "split", "run_status", "final_grnwt", "total_irrigation", "total_n", "profit_simple"]], 16)
    if not diagnosis.empty:
        add_ppt_table(prs, "Decision Reasonableness", diagnosis[["station_code", "year", "total_irrigation", "total_n", "irrigation_during_or_before_swfac_stress", "fertilization_during_or_before_nstres", "decision_reasonableness_label"]], 16)
    for fig in [
        "ppo_total_irrigation_by_station_year.png",
        "ppo_total_n_by_station_year.png",
        "ppo_final_grnwt_by_station_year.png",
        "ppo_decision_reasonableness_summary.png",
    ]:
        add_ppt_picture(prs, fig.replace("_", " ").replace(".png", ""), OUTPUT_ROOT / "figures" / "summary" / fig)
    for path in (OUTPUT_ROOT / "figures" / "four_panel").glob("*_four_panel.png"):
        add_ppt_picture(prs, path.stem, path)
    add_ppt_bullets(
        prs,
        "Conclusion",
        [
            "This file records whether the simple direct PPO baseline is presentable.",
            "If outputs are cap-saturated, the result should be shown as a baseline limitation rather than a final management recommendation.",
            "The four figure types allow mentor-facing inspection of weather, stress, actions, and crop growth.",
        ],
    )
    prs.save(DOC_PPT)


def run_pipeline(config_path: Path, selection_only: bool = False, train_only: bool = False, evaluate_only: bool = False, report_only: bool = False, debug: bool = False) -> None:
    config = load_yaml(config_path)
    ensure_dirs()
    shutil.copyfile(config_path, OUTPUT_ROOT / "configs" / config_path.name)
    selection_path = OUTPUT_ROOT / "configs" / "ppo_train_eval_year_selection.csv"
    if selection_path.exists() and (evaluate_only or report_only):
        selection = pd.read_csv(selection_path)
    else:
        selection = select_train_eval_years(config)
    selected = selection[selection["selected_for_train"] | selection["selected_for_eval"]]
    env_config = build_env_config(config, selected)
    env_config_path = OUTPUT_ROOT / "configs" / "resolved_env_config_all_year_direct_action_safe_ppo.yaml"
    write_yaml(env_config, env_config_path)
    if selection_only:
        print(selection_path.relative_to(PROJECT_ROOT))
        print(env_config_path.relative_to(PROJECT_ROOT))
        return
    train_summary_path = OUTPUT_ROOT / "evaluation" / "training_run_summary.csv"
    if report_only or evaluate_only:
        train_summary = pd.read_csv(train_summary_path) if train_summary_path.exists() else pd.DataFrame()
    else:
        train_summary = train_station_models(config, env_config, selection, debug=debug)
    if train_only:
        print(train_summary_path.relative_to(PROJECT_ROOT))
        return
    eval_summary_path = OUTPUT_ROOT / "evaluation" / "ppo_direct_eval_summary.csv"
    if report_only:
        eval_summary = pd.read_csv(eval_summary_path) if eval_summary_path.exists() else pd.DataFrame()
    else:
        eval_summary = evaluate_models(config, env_config, selection, train_summary)
    fig_index = plot_representative_outputs(config, env_config, eval_summary) if not eval_summary.empty else pd.DataFrame()
    diagnosis = build_reasonableness(config, eval_summary) if not eval_summary.empty else pd.DataFrame()
    plot_summary(eval_summary, diagnosis)
    build_report(config, selection, train_summary, eval_summary, diagnosis, fig_index)
    print(selection_path.relative_to(PROJECT_ROOT))
    print(train_summary_path.relative_to(PROJECT_ROOT))
    print(eval_summary_path.relative_to(PROJECT_ROOT))
    print((OUTPUT_ROOT / "evaluation" / "ppo_decision_reasonableness_diagnosis.csv").relative_to(PROJECT_ROOT))
    print(DOC_MD.relative_to(PROJECT_ROOT))
    print(DOC_PPT.relative_to(PROJECT_ROOT))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default=str(CONFIG_PATH))
    parser.add_argument("--selection-only", action="store_true")
    parser.add_argument("--train-only", action="store_true")
    parser.add_argument("--evaluate-only", action="store_true")
    parser.add_argument("--report-only", action="store_true")
    parser.add_argument("--debug", action="store_true")
    args = parser.parse_args()
    run_pipeline(
        Path(args.config),
        selection_only=args.selection_only,
        train_only=args.train_only,
        evaluate_only=args.evaluate_only,
        report_only=args.report_only,
        debug=args.debug,
    )


if __name__ == "__main__":
    main()
