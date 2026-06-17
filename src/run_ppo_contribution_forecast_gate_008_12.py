from __future__ import annotations

from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None
    RGBColor = None
    Inches = None
    Pt = None

import run_all_year_direct_action_safe_ppo as direct
from ppo_evaluate import latest_observation_dict, scalar
from ppo_safe_rendering import PROJECT_ROOT
from stress_aware_stage_action_wrapper import ForecastStressGateDapStageActionWrapper


OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "ppo_contribution_forecast_gate_008_12"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-14_008_12_ppo_contribution_under_forecast_gate_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "015_008_12_ppo_contribution_under_forecast_gate.pptx"
SCENARIO_POOL = PROJECT_ROOT / "Leave_One_experiments" / "all_year_weather_calibration_validation" / "scenario_pool" / "all_year_weather_scenario_pool.csv"
DIRECT_CONFIG = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_all_year_direct_action_safe_ppo.yaml"
RULE_SUMMARY = PROJECT_ROOT / "Leave_One_experiments" / "forecast_gate_rule_replay_008_11" / "evaluation" / "008_11_forecast_gate_rule_replay_summary.csv"

CASES = [
    {"station": "HLA", "year": 2004, "case_role": "strong_water_limited_showcase"},
    {"station": "FQA", "year": 2008, "case_role": "non_primary_water_showcase_check"},
    {"station": "FQA", "year": 2016, "case_role": "water_stress_candidate"},
]

STAGES = [
    {"stage_id": "S1", "stage_name": "establishment", "start_dap": 0, "end_dap": 20, "irrigation_base": 0.0, "irrigation_cap": 0.0, "nitrogen_base": 50.0, "nitrogen_cap": 50.0},
    {"stage_id": "S2", "stage_name": "early_vegetative", "start_dap": 21, "end_dap": 45, "irrigation_base": 0.0, "irrigation_cap": 40.0, "nitrogen_base": 50.0, "nitrogen_cap": 50.0},
    {"stage_id": "S3", "stage_name": "mid_season", "start_dap": 46, "end_dap": 75, "irrigation_base": 0.0, "irrigation_cap": 40.0, "nitrogen_base": 50.0, "nitrogen_cap": 50.0},
    {"stage_id": "S4", "stage_name": "late_water_risk", "start_dap": 76, "end_dap": 100, "irrigation_base": 0.0, "irrigation_cap": 40.0, "nitrogen_base": 0.0, "nitrogen_cap": 0.0},
    {"stage_id": "S5", "stage_name": "terminal_water_risk", "start_dap": 101, "end_dap": "end", "irrigation_base": 0.0, "irrigation_cap": 40.0, "nitrogen_base": 0.0, "nitrogen_cap": 0.0},
]


def ensure_dirs() -> None:
    for sub in ["configs", "models", "tensorboard", "daily_outputs", "evaluation", "figures", "rendered_inputs", "logs", "reports"]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)
    for case in CASES:
        (OUTPUT_ROOT / "models" / case["station"]).mkdir(parents=True, exist_ok=True)
        (OUTPUT_ROOT / "daily_outputs" / case["station"]).mkdir(parents=True, exist_ok=True)
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)


def base_config() -> dict[str, Any]:
    return {
        "seed": 0,
        "total_timesteps": 2000,
        "runtime": {"mode": "all", "max_steps": 260, "max_stage_steps": 12},
        "ppo": {"learning_rate": 0.0003, "gamma": 0.99, "n_steps": 5, "batch_size": 5, "n_epochs": 10, "ent_coef": 0.01, "clip_range": 0.2},
        "reward": {"topwt_delta_coef": 0.001, "grnwt_delta_coef": 0.020, "trnu_delta_coef": 0.0, "water_cost": 0.100, "nitrogen_cost": 0.250, "terminal_grnwt_coef": 0.010},
        "forecast_stress_gate": {
            "enabled": True,
            "lookahead_days": 7,
            "swfac_threshold": 0.05,
            "always_block_stage_ids": ["S1"],
            "future_rain_threshold_by_stage": {"S2": 10.0, "S3": 10.0, "S4": 20.0, "S5": 20.0},
            "min_irrigation_when_triggered_by_stage": {"S2": 30.0, "S3": 30.0, "S4": 30.0, "S5": 30.0},
        },
        "stage_action": {"stages": STAGES},
    }


def ppo_kwargs(config: dict[str, Any]) -> dict[str, Any]:
    return {k: config["ppo"][k] for k in ["learning_rate", "gamma", "n_steps", "batch_size", "n_epochs", "ent_coef", "clip_range"]}


def build_env_config() -> dict[str, Any]:
    cfg = direct.load_yaml(DIRECT_CONFIG)
    cfg["paths"]["output_root"] = str(OUTPUT_ROOT.relative_to(PROJECT_ROOT))
    cfg["action_safety"] = {"enabled": False}
    cfg["runtime"]["mode"] = "all"
    cfg["runtime"]["max_steps"] = 260
    pool = pd.read_csv(SCENARIO_POOL)
    rows = []
    for case in CASES:
        hit = pool[(pool["station_code"].astype(str).eq(case["station"])) & (pd.to_numeric(pool["year"], errors="coerce").eq(int(case["year"])))]
        if hit.empty:
            raise KeyError(f"Missing scenario pool row: {case['station']} {case['year']}")
        rows.append(hit.iloc[0])
    return direct.build_env_config(cfg, pd.DataFrame(rows))


def case_config(config: dict[str, Any], env_config: dict[str, Any], station: str, year: int) -> dict[str, Any]:
    out = {**config, "forecast_stress_gate": dict(config["forecast_stress_gate"])}
    year_info = direct.find_year(env_config, station, year)
    out["forecast_stress_gate"]["planting_date"] = year_info["planting_date"]
    out["forecast_stress_gate"]["weather_csv"] = str(PROJECT_ROOT / "weather_clean_qc" / f"{station}_weather_cleaned_qc.csv")
    return out


def make_env(config: dict[str, Any], env_config: dict[str, Any], station: str, year: int, tag: str, evaluation: bool):
    base = direct.make_base_env(env_config, station, year, seed=0, run_tag=tag, evaluation=evaluation)
    return ForecastStressGateDapStageActionWrapper(base, config)


def model_path(station: str, year: int) -> Path:
    return OUTPUT_ROOT / "models" / station / f"ppo_forecast_gate_{station}_{year}_seed0.zip"


def train_one(config: dict[str, Any], env_config: dict[str, Any], station: str, year: int) -> dict[str, Any]:
    from stable_baselines3 import PPO

    env = None
    out = model_path(station, year)
    try:
        env = make_env(config, env_config, station, year, f"{station}_{year}_00812_train", evaluation=False)
        model = PPO("MlpPolicy", env, verbose=0, seed=0, tensorboard_log=str(OUTPUT_ROOT / "tensorboard" / station), **ppo_kwargs(config))
        model.learn(total_timesteps=int(config["total_timesteps"]), progress_bar=False)
        model.save(str(out.with_suffix("")))
        return {"station": station, "year": year, "run_status": "ok", "model_path": str(out.relative_to(PROJECT_ROOT)), "notes": ""}
    except Exception as exc:
        return {"station": station, "year": year, "run_status": "failed", "model_path": "", "notes": repr(exc)}
    finally:
        if env is not None:
            env.close()


def evaluate_one(config: dict[str, Any], env_config: dict[str, Any], station: str, year: int, case_role: str) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    from stable_baselines3 import PPO

    model = PPO.load(str(model_path(station, year)))
    env = make_env(config, env_config, station, year, f"{station}_{year}_00812_eval", evaluation=True)
    daily_rows: list[dict[str, Any]] = []
    stage_rows: list[dict[str, Any]] = []
    try:
        obs, info = env.reset()
        done = False
        step_no = 0
        planting = pd.Timestamp(direct.find_year(env_config, station, year)["planting_date"])
        while not done and step_no < 12:
            before = latest_observation_dict(env, obs, info)
            action, _ = model.predict(obs, deterministic=True)
            obs, reward, terminated, truncated, info = env.step(action)
            done = bool(terminated or truncated)
            after = latest_observation_dict(env, obs, info)
            step_no += 1
            stage_info = dict(env.last_action_info)
            stage_rows.append({"station": station, "year": year, "case_role": case_role, "stage_step": step_no, "start_dap_observed": int(round(scalar(before.get("dap"), step_no))), "end_dap_observed": int(round(scalar(after.get("dap"), step_no))), "reward": float(reward), "done": done, **stage_info})
            for item in env.last_stage_records:
                dap = int(round(float(item["dap_before_step"])))
                date = planting + pd.Timedelta(days=max(dap - 1, 0))
                daily_rows.append({"station": station, "year": year, "case_role": case_role, "date": date.strftime("%Y-%m-%d"), "doy": int(date.dayofyear), "dap": dap, "stage_step": step_no, **item, "reward_stage": float(reward), "done": done})
    finally:
        env.close()
    daily = add_weather(pd.DataFrame(daily_rows), station, year)
    stages = pd.DataFrame(stage_rows)
    daily_path = OUTPUT_ROOT / "daily_outputs" / station / f"{station}_{year}_ppo_forecast_gate_daily.csv"
    stage_path = OUTPUT_ROOT / "daily_outputs" / station / f"{station}_{year}_ppo_forecast_gate_steps.csv"
    daily.to_csv(daily_path, index=False, encoding="utf-8-sig")
    stages.to_csv(stage_path, index=False, encoding="utf-8-sig")
    return daily, stages, summarize(station, year, case_role, daily, stages, daily_path, stage_path)


def add_weather(daily: pd.DataFrame, station: str, year: int) -> pd.DataFrame:
    if daily.empty:
        return daily
    weather = pd.read_csv(PROJECT_ROOT / "weather_clean_qc" / f"{station}_weather_cleaned_qc.csv", parse_dates=["date"])
    weather = weather[pd.to_numeric(weather["year"], errors="coerce").eq(int(year))].copy()
    weather = weather.rename(columns={"RAIN": "rain", "SRAD": "srad", "TMAX": "tmax", "TMIN": "tmin"})[["date", "rain", "srad", "tmax", "tmin"]]
    out = daily.copy()
    out["date_dt"] = pd.to_datetime(out["date"])
    out = out.merge(weather, left_on="date_dt", right_on="date", how="left", suffixes=("", "_weather"))
    return out.drop(columns=["date_dt", "date_weather"], errors="ignore")


def summarize(station: str, year: int, case_role: str, daily: pd.DataFrame, stages: pd.DataFrame, daily_path: Path, stage_path: Path) -> dict[str, Any]:
    irrigation = pd.to_numeric(stages.get("stage_action_amir", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
    nitrogen = pd.to_numeric(stages.get("stage_action_anfer", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
    min_i = pd.to_numeric(stages.get("gate_min_irrigation_applied", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
    extra = np.maximum(irrigation - min_i, 0.0)
    before = pd.to_numeric(stages.get("irrigation_before_gate", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
    grnwt = pd.to_numeric(daily.get("grnwt", pd.Series(dtype=float)), errors="coerce").dropna()
    swfac = pd.to_numeric(daily.get("swfac", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
    return {
        "station": station,
        "year": year,
        "case_role": case_role,
        "run_status": "ok" if len(stages) and bool(stages["done"].iloc[-1]) else "failed",
        "total_irrigation": float(irrigation.sum()),
        "total_n": float(nitrogen.sum()),
        "final_grnwt": float(grnwt.iloc[-1]) if len(grnwt) else np.nan,
        "swfac_days_gt_0p05": int((swfac > 0.05).sum()) if len(swfac) else 0,
        "gate_trigger_count": int((min_i > 0).sum()),
        "total_gate_min_irrigation": float(min_i.sum()),
        "total_ppo_extra_above_min": float(extra.sum()),
        "mean_irrigation_before_gate": float(before.mean()) if len(before) else np.nan,
        "stages_with_ppo_extra": int((extra > 1e-6).sum()),
        "daily_csv_path": str(daily_path.relative_to(PROJECT_ROOT)),
        "stage_csv_path": str(stage_path.relative_to(PROJECT_ROOT)),
    }


def plot_case(daily: pd.DataFrame, stages: pd.DataFrame, station: str, year: int) -> Path:
    fig_dir = OUTPUT_ROOT / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    fig, axes = plt.subplots(4, 1, figsize=(10, 8), dpi=180, sharex=True)
    axes[0].bar(daily["dap"], pd.to_numeric(daily.get("rain", 0), errors="coerce").fillna(0), color="#4F81BD", alpha=0.35)
    axes[0].set_ylabel("Rain mm")
    axes[1].plot(daily["dap"], pd.to_numeric(daily["swfac"], errors="coerce"), label="SWFAC", color="#4F81BD")
    axes[1].plot(daily["dap"], pd.to_numeric(daily["nstres"], errors="coerce"), label="NSTRES", color="#C0504D")
    axes[1].axhline(0.05, color="black", linestyle="--", linewidth=0.8)
    axes[1].legend(frameon=False)
    axes[1].set_ylabel("Stress")
    axes[2].vlines(stages["decision_dap"], 0, pd.to_numeric(stages["stage_action_amir"], errors="coerce"), color="#4F81BD", label="PPO irrigation")
    axes[2].vlines(stages["decision_dap"], 0, pd.to_numeric(stages["stage_action_anfer"], errors="coerce"), color="#9BBB59", linestyle="dashed", label="N")
    axes[2].legend(frameon=False)
    axes[2].set_ylabel("Actions")
    axes[3].plot(daily["dap"], pd.to_numeric(daily["topwt"], errors="coerce"), label="TOPWT", color="#9BBB59")
    axes[3].plot(daily["dap"], pd.to_numeric(daily["grnwt"], errors="coerce"), label="GRNWT", color="#8064A2")
    axes[3].legend(frameon=False)
    axes[3].set_ylabel("Growth")
    axes[3].set_xlabel("DAP")
    fig.suptitle(f"{station} {year} PPO under forecast gate", fontsize=12)
    fig.tight_layout()
    out = fig_dir / f"{station}_{year}_ppo_vs_rule_process.png"
    fig.savefig(out)
    plt.close(fig)
    return out


def compare_rule(ppo_summary: pd.DataFrame) -> pd.DataFrame:
    rule = pd.read_csv(RULE_SUMMARY)
    merged = ppo_summary.merge(rule, on=["station", "year"], suffixes=("_ppo", "_rule"))
    merged["yield_delta_ppo_minus_rule"] = merged["final_grnwt_ppo"] - merged["final_grnwt_rule"]
    merged["irrigation_delta_ppo_minus_rule"] = merged["total_irrigation_ppo"] - merged["total_irrigation_rule"]
    merged["ppo_has_measurable_contribution"] = (
        (merged["yield_delta_ppo_minus_rule"] > 100.0)
        | ((merged["yield_delta_ppo_minus_rule"] > -100.0) & (merged["irrigation_delta_ppo_minus_rule"] < -5.0))
        | (merged["total_ppo_extra_above_min"] > 1.0)
    )
    return merged


def df_md(df: pd.DataFrame, max_rows: int = 30) -> str:
    if df.empty:
        return ""
    out = df.head(max_rows).copy()
    for col in out.select_dtypes(include=["float", "int"]).columns:
        out[col] = pd.to_numeric(out[col], errors="coerce").round(4)
    header = "| " + " | ".join(out.columns) + " |"
    sep = "| " + " | ".join(["---"] * len(out.columns)) + " |"
    rows = ["| " + " | ".join("" if pd.isna(x) else str(x) for x in row) + " |" for row in out.values.tolist()]
    return "\n".join([header, sep, *rows])


def build_report(train: pd.DataFrame, ppo: pd.DataFrame, comp: pd.DataFrame, figures: list[Path]) -> None:
    lines = [
        "# 008_12 PPO Contribution Under Forecast Gate",
        "",
        "## Purpose",
        "",
        "This diagnostic compares PPO under the rain10_min30 forecast gate with the 008_11 pure rule replay baseline.",
        "",
        "## Training",
        "",
        df_md(train),
        "",
        "## PPO Summary",
        "",
        df_md(ppo),
        "",
        "## PPO vs Rule",
        "",
        df_md(comp[["station", "year", "final_grnwt_ppo", "final_grnwt_rule", "yield_delta_ppo_minus_rule", "total_irrigation_ppo", "total_irrigation_rule", "irrigation_delta_ppo_minus_rule", "total_ppo_extra_above_min", "stages_with_ppo_extra", "ppo_has_measurable_contribution"]]),
        "",
        "## Interpretation",
        "",
        "- If PPO mostly stays at the gate minimum and does not improve yield or save water, PPO contribution is limited under this gate.",
        "- If PPO adds water above the gate minimum or improves yield/water efficiency, PPO has measurable contribution.",
    ]
    DOC_MD.write_text("\n".join(lines), encoding="utf-8")
    build_ppt(comp, figures)


def build_ppt(comp: pd.DataFrame, figures: list[Path]) -> None:
    if Presentation is None:
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def title(slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.55), Inches(0.25), Inches(12.3), Inches(0.55))
        p = box.text_frame.paragraphs[0]
        p.text = text
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0)

    def body(slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.75), Inches(1.0), Inches(11.8), Inches(5.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(14)
            p.font.color.rgb = RGBColor(0, 0, 0)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "008_12 PPO Contribution Diagnostic")
    body(s, "Same rain10_min30 forecast gate.\n\nCompare pure rule replay vs PPO trained from scratch per station-year.\n\nQuestion: does PPO add value beyond the gate minimum?")
    keep = ["station", "year", "yield_delta_ppo_minus_rule", "irrigation_delta_ppo_minus_rule", "total_ppo_extra_above_min", "ppo_has_measurable_contribution"]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "PPO vs Rule Summary")
    body(s, comp[keep].to_string(index=False))
    for fig in figures[:3]:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        title(s, fig.stem.replace("_", " "))
        s.shapes.add_picture(str(fig), Inches(0.85), Inches(0.95), width=Inches(11.8))
    prs.save(DOC_PPT)


def main() -> None:
    ensure_dirs()
    env_config = build_env_config()
    direct.write_yaml(env_config, OUTPUT_ROOT / "configs" / "rendered_env_config.yaml")
    config = base_config()
    train_rows = []
    ppo_rows = []
    figures = []
    for case in CASES:
        station, year = case["station"], int(case["year"])
        cfg = case_config(config, env_config, station, year)
        train = train_one(cfg, env_config, station, year)
        train_rows.append(train)
        if train["run_status"] == "ok":
            daily, stages, summary = evaluate_one(cfg, env_config, station, year, case["case_role"])
            ppo_rows.append(summary)
            figures.append(plot_case(daily, stages, station, year))
    train_df = pd.DataFrame(train_rows)
    ppo_df = pd.DataFrame(ppo_rows)
    comp = compare_rule(ppo_df) if not ppo_df.empty else pd.DataFrame()
    train_df.to_csv(OUTPUT_ROOT / "evaluation" / "008_12_ppo_training_summary.csv", index=False, encoding="utf-8-sig")
    ppo_df.to_csv(OUTPUT_ROOT / "evaluation" / "008_12_ppo_contribution_summary.csv", index=False, encoding="utf-8-sig")
    comp.to_csv(OUTPUT_ROOT / "evaluation" / "008_12_ppo_vs_rule_comparison.csv", index=False, encoding="utf-8-sig")
    build_report(train_df, ppo_df, comp, figures)
    print((OUTPUT_ROOT / "evaluation" / "008_12_ppo_contribution_summary.csv").relative_to(PROJECT_ROOT))
    print((OUTPUT_ROOT / "evaluation" / "008_12_ppo_vs_rule_comparison.csv").relative_to(PROJECT_ROOT))
    print(DOC_MD.relative_to(PROJECT_ROOT))
    print(DOC_PPT.relative_to(PROJECT_ROOT))


if __name__ == "__main__":
    main()
