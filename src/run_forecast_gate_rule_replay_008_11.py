from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None
    RGBColor = None
    Inches = None
    Pt = None

import run_all_year_direct_action_safe_ppo as direct
from ppo_action_safety import normalize_action
from ppo_evaluate import latest_observation_dict, scalar
from ppo_safe_rendering import PROJECT_ROOT


OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "forecast_gate_rule_replay_008_11"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-14_008_11_forecast_gate_rule_replay_validation_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "014_008_11_forecast_gate_rule_replay_validation.pptx"
SCENARIO_POOL = PROJECT_ROOT / "Leave_One_experiments" / "all_year_weather_calibration_validation" / "scenario_pool" / "all_year_weather_scenario_pool.csv"
DIRECT_CONFIG = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_all_year_direct_action_safe_ppo.yaml"
REFERENCE_SUMMARY = PROJECT_ROOT / "Leave_One_experiments" / "representative_management_comparison_008_06" / "evaluation" / "HLA_2004_representative_scenario_summary.csv"


@dataclass(frozen=True)
class RuleStage:
    stage_id: str
    stage_name: str
    start_dap: int
    end_dap: int | None
    nitrogen: float
    irrigation: float
    future_rain_threshold: float
    swfac_threshold: float = 0.05
    irrigation_allowed: bool = True


STAGES = [
    RuleStage("S1", "establishment", 0, 20, 50.0, 0.0, 0.0, irrigation_allowed=False),
    RuleStage("S2", "early_vegetative", 21, 45, 50.0, 30.0, 10.0),
    RuleStage("S3", "mid_season", 46, 75, 50.0, 30.0, 10.0),
    RuleStage("S4", "late_water_risk", 76, 100, 0.0, 30.0, 20.0),
    RuleStage("S5", "terminal_water_risk", 101, None, 0.0, 30.0, 20.0),
]

CASES = [
    {"station": "HLA", "year": 2004, "case_role": "strong_water_limited_showcase"},
    {"station": "FQA", "year": 2008, "case_role": "non_primary_water_showcase_check"},
    {"station": "FQA", "year": 2016, "case_role": "water_stress_candidate"},
]


def ensure_dirs() -> None:
    for sub in ["configs", "daily_outputs", "event_outputs", "evaluation", "figures", "rendered_inputs", "logs", "reports"]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)
    for case in CASES:
        (OUTPUT_ROOT / "daily_outputs" / case["station"]).mkdir(parents=True, exist_ok=True)
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)


def build_env_config() -> dict[str, Any]:
    cfg = direct.load_yaml(DIRECT_CONFIG)
    cfg["paths"]["output_root"] = str(OUTPUT_ROOT.relative_to(PROJECT_ROOT))
    cfg["action_safety"] = {"enabled": False}
    cfg["runtime"]["mode"] = "all"
    cfg["runtime"]["max_steps"] = 260
    pool = pd.read_csv(SCENARIO_POOL)
    selections = []
    for case in CASES:
        hit = pool[(pool["station_code"].astype(str).eq(case["station"])) & (pd.to_numeric(pool["year"], errors="coerce").eq(case["year"]))]
        if hit.empty:
            raise KeyError(f"Missing scenario pool row: {case['station']} {case['year']}")
        selections.append(hit.iloc[0])
    selection = pd.DataFrame(selections)
    return direct.build_env_config(cfg, selection)


def weather_path(station: str) -> Path:
    return PROJECT_ROOT / "weather_clean_qc" / f"{station}_weather_cleaned_qc.csv"


def read_weather(station: str, year: int) -> pd.DataFrame:
    path = weather_path(station)
    df = pd.read_csv(path, parse_dates=["date"])
    df = df[pd.to_numeric(df["year"], errors="coerce").eq(int(year))].copy()
    rain_col = "RAIN" if "RAIN" in df.columns else "rain"
    return df.rename(columns={rain_col: "rain", "SRAD": "srad", "TMAX": "tmax", "TMIN": "tmin"})[
        ["date", "rain", "srad", "tmax", "tmin"]
    ]


def stage_for_dap(dap: int) -> RuleStage:
    for stage in STAGES:
        if dap >= stage.start_dap and (stage.end_dap is None or dap <= stage.end_dap):
            return stage
    return STAGES[-1]


def next_stage(dap: int, stage: RuleStage) -> bool:
    return stage.end_dap is not None and dap > stage.end_dap


def future_rain(weather: pd.DataFrame, planting: pd.Timestamp, dap: int, lookahead_days: int = 7) -> float:
    date = planting + pd.Timedelta(days=max(int(dap) - 1, 0))
    end = date + pd.Timedelta(days=max(lookahead_days - 1, 0))
    mask = (weather["date"] >= date) & (weather["date"] <= end)
    return float(pd.to_numeric(weather.loc[mask, "rain"], errors="coerce").fillna(0.0).sum())


def make_env(env_config: dict[str, Any], station: str, year: int):
    return direct.make_base_env(env_config, station, year, seed=0, run_tag=f"{station}_{year}_00811_rule_replay", evaluation=True)


def step_real_action(env, real_action: dict[str, float]):
    names = list(env.formator.action_names)
    full = {name: 0.0 for name in names}
    full.update({k: max(0.0, float(v)) for k, v in real_action.items() if k in full})
    norm = normalize_action(names, env.formator.action_space_dict, full)
    obs, env_reward, terminated, truncated, info = env.step(norm)
    return obs, env_reward, terminated, truncated, info, full


def run_case(env_config: dict[str, Any], station: str, year: int, case_role: str) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    env = make_env(env_config, station, year)
    weather = read_weather(station, year)
    year_info = direct.find_year(env_config, station, year)
    planting = pd.Timestamp(year_info["planting_date"])
    daily_rows: list[dict[str, Any]] = []
    event_rows: list[dict[str, Any]] = []
    try:
        obs, info = env.reset()
        latest = latest_observation_dict(env, obs, info)
        done = False
        stage_step = 0
        while not done and stage_step < 12:
            prev = dict(latest)
            dap = int(round(scalar(prev.get("dap"), 1.0)))
            if dap <= 0:
                dap = 1
            stage = stage_for_dap(dap)
            swfac_at_decision = scalar(prev.get("swfac"), 0.0)
            rain7 = future_rain(weather, planting, dap)
            swfac_trigger = bool(swfac_at_decision > stage.swfac_threshold)
            forecast_trigger = bool(rain7 < stage.future_rain_threshold)
            allowed = bool(stage.irrigation_allowed and (swfac_trigger or forecast_trigger))
            irrigation = stage.irrigation if allowed else 0.0
            nitrogen = stage.nitrogen
            reason = "blocked_stage" if not stage.irrigation_allowed else (
                "allowed_by_swfac_and_forecast" if swfac_trigger and forecast_trigger else
                "allowed_by_swfac" if swfac_trigger else
                "allowed_by_forecast" if forecast_trigger else
                "blocked_no_stress_or_dry_forecast"
            )
            stage_step += 1
            event_rows.append(
                {
                    "station": station,
                    "year": year,
                    "case_role": case_role,
                    "stage_step": stage_step,
                    "stage_id": stage.stage_id,
                    "stage_name": stage.stage_name,
                    "decision_dap": dap,
                    "swfac_at_decision": swfac_at_decision,
                    "future_7d_rain": rain7,
                    "future_rain_threshold": stage.future_rain_threshold,
                    "swfac_trigger": swfac_trigger,
                    "forecast_trigger": forecast_trigger,
                    "irrigation_allowed": allowed,
                    "gate_reason": reason,
                    "irrigation": irrigation,
                    "nitrogen": nitrogen,
                }
            )
            obs, _reward, terminated, truncated, info, full = step_real_action(env, {"amir": irrigation, "anfer": nitrogen})
            latest = latest_observation_dict(env, obs, info)
            done = bool(terminated or truncated)
            append_daily(daily_rows, station, year, case_role, planting, weather, stage, stage_step, dap, "stage_rule_action", full, latest, done)
            while not done:
                latest = latest_observation_dict(env, obs, info)
                current_dap = int(round(scalar(latest.get("dap"), dap)))
                if current_dap <= 0:
                    current_dap = dap
                if next_stage(current_dap, stage):
                    break
                obs, _reward, terminated, truncated, info, full = step_real_action(env, {"amir": 0.0, "anfer": 0.0})
                latest = latest_observation_dict(env, obs, info)
                done = bool(terminated or truncated)
                append_daily(daily_rows, station, year, case_role, planting, weather, stage, stage_step, current_dap, "internal_zero", full, latest, done)
    finally:
        env.close()
    daily = pd.DataFrame(daily_rows)
    events = pd.DataFrame(event_rows)
    daily_path = OUTPUT_ROOT / "daily_outputs" / station / f"{station}_{year}_forecast_gate_rule_daily.csv"
    event_path = OUTPUT_ROOT / "event_outputs" / f"{station}_{year}_forecast_gate_rule_events.csv"
    daily.to_csv(daily_path, index=False, encoding="utf-8-sig")
    events.to_csv(event_path, index=False, encoding="utf-8-sig")
    summary = summarize_case(station, year, case_role, daily, events, daily_path, event_path)
    return daily, events, summary


def append_daily(
    rows: list[dict[str, Any]],
    station: str,
    year: int,
    case_role: str,
    planting: pd.Timestamp,
    weather: pd.DataFrame,
    stage: RuleStage,
    stage_step: int,
    dap: int,
    action_role: str,
    action: dict[str, float],
    obs: dict[str, Any],
    done: bool,
) -> None:
    date = planting + pd.Timedelta(days=max(int(dap) - 1, 0))
    wrow = weather[weather["date"].eq(date)]
    w = wrow.iloc[0].to_dict() if len(wrow) else {}
    rows.append(
        {
            "station": station,
            "year": year,
            "case_role": case_role,
            "date": date.strftime("%Y-%m-%d"),
            "doy": int(date.dayofyear),
            "dap": int(dap),
            "stage_step": stage_step,
            "stage_id": stage.stage_id,
            "stage_name": stage.stage_name,
            "action_role": action_role,
            "real_action_amir": float(action.get("amir", 0.0)),
            "real_action_anfer": float(action.get("anfer", 0.0)),
            "rain": scalar(w.get("rain"), np.nan),
            "srad": scalar(w.get("srad"), np.nan),
            "tmax": scalar(w.get("tmax"), np.nan),
            "tmin": scalar(w.get("tmin"), np.nan),
            "swfac": scalar(obs.get("swfac")),
            "nstres": scalar(obs.get("nstres")),
            "topwt": scalar(obs.get("topwt")),
            "grnwt": scalar(obs.get("grnwt")),
            "xlai": scalar(obs.get("xlai")),
            "done": done,
        }
    )


def summarize_case(station: str, year: int, case_role: str, daily: pd.DataFrame, events: pd.DataFrame, daily_path: Path, event_path: Path) -> dict[str, Any]:
    swfac = pd.to_numeric(daily.get("swfac", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
    nstres = pd.to_numeric(daily.get("nstres", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
    grnwt = pd.to_numeric(daily.get("grnwt", pd.Series(dtype=float)), errors="coerce").dropna()
    topwt = pd.to_numeric(daily.get("topwt", pd.Series(dtype=float)), errors="coerce").dropna()
    irrigation = pd.to_numeric(events.get("irrigation", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
    nitrogen = pd.to_numeric(events.get("nitrogen", pd.Series(dtype=float)), errors="coerce").fillna(0.0)
    triggers = events[pd.to_numeric(events.get("irrigation", 0), errors="coerce").fillna(0.0) > 0]
    return {
        "station": station,
        "year": year,
        "case_role": case_role,
        "run_status": "ok" if len(daily) else "failed",
        "daily_steps": int(len(daily)),
        "total_irrigation": float(irrigation.sum()),
        "total_n": float(nitrogen.sum()),
        "final_grnwt": float(grnwt.iloc[-1]) if len(grnwt) else np.nan,
        "final_topwt": float(topwt.iloc[-1]) if len(topwt) else np.nan,
        "max_swfac": float(swfac.max()) if len(swfac) else np.nan,
        "swfac_days_gt_0p05": int((swfac > 0.05).sum()) if len(swfac) else 0,
        "max_nstres": float(nstres.max()) if len(nstres) else np.nan,
        "nstres_days_gt_0p05": int((nstres > 0.05).sum()) if len(nstres) else 0,
        "gate_trigger_count": int(len(triggers)),
        "trigger_count_with_swfac_gt_0p05": int(pd.to_numeric(triggers.get("swfac_at_decision", pd.Series(dtype=float)), errors="coerce").gt(0.05).sum()) if len(triggers) else 0,
        "trigger_count_by_forecast_only": int((triggers.get("forecast_trigger", pd.Series(dtype=bool)).astype(str).str.lower().isin(["true", "1"]) & ~triggers.get("swfac_trigger", pd.Series(dtype=bool)).astype(str).str.lower().isin(["true", "1"])).sum()) if len(triggers) else 0,
        "first_irrigation_dap": float(triggers["decision_dap"].iloc[0]) if len(triggers) else np.nan,
        "daily_csv_path": str(daily_path.relative_to(PROJECT_ROOT)),
        "event_csv_path": str(event_path.relative_to(PROJECT_ROOT)),
    }


def plot_case(daily: pd.DataFrame, events: pd.DataFrame, station: str, year: int) -> Path:
    fig_dir = OUTPUT_ROOT / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    fig, axes = plt.subplots(4, 1, figsize=(10, 8), dpi=180, sharex=True)
    axes[0].bar(daily["dap"], pd.to_numeric(daily["rain"], errors="coerce").fillna(0.0), color="#4F81BD", alpha=0.35)
    axes[0].set_ylabel("Rain mm")
    axes[1].plot(daily["dap"], pd.to_numeric(daily["swfac"], errors="coerce"), label="SWFAC", color="#4F81BD")
    axes[1].plot(daily["dap"], pd.to_numeric(daily["nstres"], errors="coerce"), label="NSTRES", color="#C0504D")
    axes[1].axhline(0.05, color="black", linestyle="--", linewidth=0.8)
    axes[1].legend(frameon=False)
    axes[1].set_ylabel("Stress")
    axes[2].vlines(events["decision_dap"], 0, pd.to_numeric(events["irrigation"], errors="coerce"), color="#4F81BD", label="Irrigation")
    axes[2].vlines(events["decision_dap"], 0, pd.to_numeric(events["nitrogen"], errors="coerce"), color="#9BBB59", linestyle="dashed", label="N")
    axes[2].legend(frameon=False)
    axes[2].set_ylabel("Actions")
    axes[3].plot(daily["dap"], pd.to_numeric(daily["topwt"], errors="coerce"), label="TOPWT", color="#9BBB59")
    axes[3].plot(daily["dap"], pd.to_numeric(daily["grnwt"], errors="coerce"), label="GRNWT", color="#8064A2")
    axes[3].legend(frameon=False)
    axes[3].set_ylabel("Growth")
    axes[3].set_xlabel("DAP")
    fig.suptitle(f"{station} {year} forecast-gate rule replay", fontsize=12)
    fig.tight_layout()
    out = fig_dir / f"{station}_{year}_forecast_gate_rule_process.png"
    fig.savefig(out)
    plt.close(fig)
    return out


def build_report(summary: pd.DataFrame, audit: pd.DataFrame, figures: list[Path]) -> None:
    lines = [
        "# 008_11 Forecast Gate Rule Replay Validation",
        "",
        "## Purpose",
        "",
        "This test runs the forecast gate as a deterministic rule without PPO. It checks whether the rule boundary is agronomically reasonable before using it as an RL action constraint.",
        "",
        "## Summary",
        "",
        df_md(summary),
        "",
        "## Gate Trigger Audit",
        "",
        df_md(audit),
        "",
        "## Interpretation",
        "",
        "- HLA 2004 should irrigate because it is a strong water-limited case.",
        "- FQA 2008 is used to check whether the rule over-irrigates a case that was not suitable as the main water-optimization showcase.",
        "- FQA 2016 is a water-stress candidate and should behave as an intermediate/response case.",
        "- This result does not prove PPO contribution; it only validates the forecast-gate boundary.",
    ]
    DOC_MD.write_text("\n".join(lines), encoding="utf-8")
    build_ppt(summary, figures)


def df_md(df: pd.DataFrame) -> str:
    if df.empty:
        return ""
    out = df.copy()
    for col in out.select_dtypes(include=["float", "int"]).columns:
        out[col] = pd.to_numeric(out[col], errors="coerce").round(4)
    header = "| " + " | ".join(out.columns) + " |"
    sep = "| " + " | ".join(["---"] * len(out.columns)) + " |"
    rows = ["| " + " | ".join("" if pd.isna(x) else str(x) for x in row) + " |" for row in out.values.tolist()]
    return "\n".join([header, sep, *rows])


def build_ppt(summary: pd.DataFrame, figures: list[Path]) -> None:
    if Presentation is None:
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def title(slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.55), Inches(0.25), Inches(12.3), Inches(0.55))
        p = box.text_frame.paragraphs[0]
        p.text = text
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0)

    def body(slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.75), Inches(1.0), Inches(11.8), Inches(5.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(15)
            p.font.color.rgb = RGBColor(0, 0, 0)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "008_11 Forecast Gate Rule Replay")
    body(s, "Pure rule replay, no PPO training and no model weights.\n\nPurpose: test whether the forecast-stress gate is an agronomically reasonable action boundary.")
    keep = ["station", "year", "case_role", "total_irrigation", "total_n", "final_grnwt", "swfac_days_gt_0p05", "gate_trigger_count", "first_irrigation_dap"]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Summary")
    body(s, summary[keep].to_string(index=False))
    for fig in figures:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        title(s, fig.stem.replace("_", " "))
        s.shapes.add_picture(str(fig), Inches(0.85), Inches(0.95), width=Inches(11.8))
    prs.save(DOC_PPT)


def main() -> None:
    ensure_dirs()
    env_config = build_env_config()
    direct.write_yaml(env_config, OUTPUT_ROOT / "configs" / "rendered_env_config.yaml")
    summaries = []
    audits = []
    figures = []
    for case in CASES:
        daily, events, summary = run_case(env_config, case["station"], int(case["year"]), case["case_role"])
        summaries.append(summary)
        audits.append(events)
        figures.append(plot_case(daily, events, case["station"], int(case["year"])))
    summary_df = pd.DataFrame(summaries)
    audit_df = pd.concat(audits, ignore_index=True) if audits else pd.DataFrame()
    summary_path = OUTPUT_ROOT / "evaluation" / "008_11_forecast_gate_rule_replay_summary.csv"
    audit_path = OUTPUT_ROOT / "evaluation" / "008_11_forecast_gate_rule_trigger_audit.csv"
    summary_df.to_csv(summary_path, index=False, encoding="utf-8-sig")
    audit_df.to_csv(audit_path, index=False, encoding="utf-8-sig")
    build_report(summary_df, audit_df, figures)
    if DOC_MD.exists():
        (OUTPUT_ROOT / "reports").mkdir(parents=True, exist_ok=True)
        (OUTPUT_ROOT / "reports" / DOC_MD.name).write_text(DOC_MD.read_text(encoding="utf-8"), encoding="utf-8")
    print(summary_path.relative_to(PROJECT_ROOT))
    print(audit_path.relative_to(PROJECT_ROOT))
    print(DOC_MD.relative_to(PROJECT_ROOT))
    print(DOC_PPT.relative_to(PROJECT_ROOT))


if __name__ == "__main__":
    main()
