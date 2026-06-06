from __future__ import annotations

from pathlib import Path
import shutil
import textwrap

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import yaml
try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None
    RGBColor = None
    PP_ALIGN = None
    Inches = None
    Pt = None

from ppo_evaluate import make_env
from ppo_safe_rendering import PROJECT_ROOT, load_yaml


CURRENT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "ppo_observed_years"
DEBUG_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "ppo_action_debug"
CONFIG_SAFE = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_ppo_action_safe_debug.yaml"
REPORT_MD = PROJECT_ROOT / "docs" / "2026-06-05_ppo_action_scale_and_reward_debug_report.md"
REPORT_PPTX = PROJECT_ROOT / "docs" / "2026-06-05_ppo_action_scale_and_reward_debug_report.pptx"
BLUE = RGBColor(68, 114, 196) if RGBColor else None
LIGHT_BLUE = RGBColor(232, 238, 249) if RGBColor else None


def ensure_dirs() -> None:
    for sub in ["action_diagnostics", "reward_diagnostics", "config_versions", "models", "daily_outputs", "evaluation", "figures", "reports"]:
        (DEBUG_ROOT / sub).mkdir(parents=True, exist_ok=True)


def read_csv(path: Path) -> pd.DataFrame:
    if path.exists():
        return pd.read_csv(path)
    return pd.DataFrame()


def rel(path: Path) -> str:
    return str(path.relative_to(PROJECT_ROOT)).replace("\\", "/")


def action_summary_for_file(path: Path) -> dict:
    df = pd.read_csv(path)
    eval_year = int(df["eval_year"].iloc[0]) if "eval_year" in df else int(path.stem.split("_eval")[1].split("_")[0])
    amir = pd.to_numeric(df["real_action_amir"], errors="coerce").fillna(0)
    anfer = pd.to_numeric(df["real_action_anfer"], errors="coerce").fillna(0)
    norm_amir = pd.to_numeric(df["normalized_action_amir"], errors="coerce")
    norm_anfer = pd.to_numeric(df["normalized_action_anfer"], errors="coerce")
    return {
        "eval_year": eval_year,
        "daily_csv_path": rel(path),
        "days": len(df),
        "nonzero_irrigation_days": int((amir > 1e-9).sum()),
        "nonzero_fertilization_days": int((anfer > 1e-9).sum()),
        "max_daily_irrigation": float(amir.max()),
        "max_daily_fertilization": float(anfer.max()),
        "mean_daily_irrigation": float(amir.mean()),
        "mean_daily_fertilization": float(anfer.mean()),
        "sum_irrigation": float(amir.sum()),
        "sum_fertilization": float(anfer.sum()),
        "median_normalized_action_amir": float(norm_amir.median()),
        "median_normalized_action_anfer": float(norm_anfer.median()),
        "max_normalized_action_amir": float(norm_amir.max()),
        "max_normalized_action_anfer": float(norm_anfer.max()),
        "min_normalized_action_amir": float(norm_amir.min()),
        "min_normalized_action_anfer": float(norm_anfer.min()),
        "totir_last": float(pd.to_numeric(df.get("totir", pd.Series([0])), errors="coerce").iloc[-1]),
        "tofer_last": float(pd.to_numeric(df.get("tofer", pd.Series([0])), errors="coerce").iloc[-1]),
    }


def plot_daily_actions(path: Path, out: Path, title: str) -> None:
    df = pd.read_csv(path)
    fig, axes = plt.subplots(2, 1, figsize=(9, 6), sharex=True)
    axes[0].plot(df["dap"], df["real_action_amir"], label="real_action_amir")
    axes[0].plot(df["dap"], df["normalized_action_amir"], label="normalized_action_amir")
    axes[0].set_ylabel("irrigation")
    axes[0].legend(frameon=False)
    axes[0].grid(alpha=0.25)
    axes[1].plot(df["dap"], df["real_action_anfer"], label="real_action_anfer")
    axes[1].plot(df["dap"], df["normalized_action_anfer"], label="normalized_action_anfer")
    axes[1].set_ylabel("fertilization")
    axes[1].set_xlabel("DAP")
    axes[1].legend(frameon=False)
    axes[1].grid(alpha=0.25)
    fig.suptitle(title)
    fig.tight_layout()
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=150)
    plt.close(fig)


def current_action_diagnostics() -> Path:
    ensure_dirs()
    rows = []
    for path in sorted((CURRENT_ROOT / "daily_outputs" / "HLA").glob("HLA_train2007_eval*_seed0_daily.csv")):
        rows.append(action_summary_for_file(path))
        eval_year = rows[-1]["eval_year"]
        plot_daily_actions(path, DEBUG_ROOT / "figures" / f"current_ppo_daily_actions_eval{eval_year}.png", f"Current PPO daily actions eval {eval_year}")
    out = DEBUG_ROOT / "action_diagnostics" / "current_ppo_action_summary.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out


def total_recalculation_check() -> Path:
    summary = read_csv(CURRENT_ROOT / "evaluation" / "ppo_evaluation_summary.csv")
    rows = []
    for path in sorted((CURRENT_ROOT / "daily_outputs" / "HLA").glob("HLA_train2007_eval*_seed0_daily.csv")):
        df = pd.read_csv(path)
        eval_year = int(df["eval_year"].iloc[0])
        row = summary[summary["eval_year"].astype(int) == eval_year].iloc[0]
        manual_irrig = float(pd.to_numeric(df["real_action_amir"], errors="coerce").fillna(0).sum())
        manual_n = float(pd.to_numeric(df["real_action_anfer"], errors="coerce").fillna(0).sum())
        rows.append(
            {
                "eval_year": eval_year,
                "manual_sum_irrigation": manual_irrig,
                "summary_total_irrigation": float(row["total_irrigation"]),
                "irrigation_match": abs(manual_irrig - float(row["total_irrigation"])) < 1e-6,
                "manual_sum_fertilization": manual_n,
                "summary_total_n_fertilizer": float(row["total_n_fertilizer"]),
                "fertilization_match": abs(manual_n - float(row["total_n_fertilizer"])) < 1e-6,
                "totir_last": float(pd.to_numeric(df["totir"], errors="coerce").iloc[-1]),
                "tofer_last": float(pd.to_numeric(df["tofer"], errors="coerce").iloc[-1]),
            }
        )
    out = DEBUG_ROOT / "action_diagnostics" / "total_input_recalculation_check.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out


def action_space_check() -> Path:
    config = load_yaml(CONFIG_SAFE)
    env = make_env(config, "HLA", 2007, 0, run_tag="action_space_check", evaluation=True, action_safety_enabled=False)
    rows = []
    try:
        spaces = getattr(env.formator.action_space_dict, "spaces", env.formator.action_space_dict)
        for name in env.formator.action_names:
            space = spaces[name]
            low = float(space.low.flatten()[0])
            high = float(space.high.flatten()[0])
            rows.append(
                {
                    "station": "HLA",
                    "year": 2007,
                    "action_name": name,
                    "action_space_low": low,
                    "action_space_high": high,
                    "normalization_formula": "normalized = 2 * ((real - low) / (high - low)) - 1",
                    "denormalization_formula": "real = low + 0.5 * (normalized + 1) * (high - low)",
                    "example_normalized_minus1": low,
                    "example_normalized_0": low + 0.5 * (high - low),
                    "example_normalized_plus1": high,
                    "notes": "PPO outputs normalized action in [-1, 1].",
                }
            )
    finally:
        env.close()
    out = DEBUG_ROOT / "action_diagnostics" / "ppo_action_space_check.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out


def reward_review() -> tuple[Path, Path]:
    review = DEBUG_ROOT / "reward_diagnostics" / "current_reward_function_review.md"
    candidates = DEBUG_ROOT / "reward_diagnostics" / "reward_revision_candidates.md"
    reward_path = Path("/opt/gym_dssat_pdi/lib/python3.10/site-packages/gym_dssat_pdi/envs/configs/rewards.py")
    snippet = ""
    source_note = "Not directly readable from this host path; see docs/all_mode_reward_and_observation_notes.md."
    if reward_path.exists():
        text = reward_path.read_text(encoding="utf-8", errors="replace")
        start = text.find("def all_reward")
        end = text.find("\ndef ", start + 1) if start >= 0 else -1
        snippet = text[start:end if end > start else None]
        source_note = str(reward_path)
    if not snippet:
        notes = PROJECT_ROOT / "docs" / "all_mode_reward_and_observation_notes.md"
        snippet = notes.read_text(encoding="utf-8", errors="replace") if notes.exists() else ""
    review.write_text(
        "# Current reward function review\n\n"
        f"Source reviewed: `{source_note}`\n\n"
        "## Code or project note excerpt\n\n"
        "```python\n"
        + snippet[:6000]
        + "\n```\n\n"
        "## Diagnosis\n\n"
        "- 当前 all-mode reward 主要来自产量/生长信号与水氮动作成本的组合。\n"
        "- 从 HLA 2007 debug PPO 的 daily output 看，策略几乎每天输出正的 amir/anfer，说明当前训练仅靠 reward 还不足以约束动作频率。\n"
        "- 如果 action cost 或 season excess penalty 不够强，PPO 在短步数随机探索阶段会频繁尝试高水高氮动作。\n"
        "- 当前结果与 fixed_high_input 的 120 mm / 250 kg ha-1 对照不一致，说明不能直接批量训练。\n"
        "- 本阶段不覆盖 reward；先用 action safety 单独隔离动作尺度问题。\n",
        encoding="utf-8",
    )
    candidates.write_text(
        "# Reward revision candidates\n\n"
        "本文件只给候选方案，不覆盖当前 reward。\n\n"
        "## Candidate A: strong input cost reward\n\n"
        "```text\n"
        "reward = crop_growth_reward - irrigation_cost * amir - nitrogen_cost * anfer - excessive_input_penalty\n"
        "excessive_input_penalty = max(0, total_irrigation - irrigation_limit) * excess_irrigation_cost\n"
        "                        + max(0, total_n - nitrogen_limit) * excess_n_cost\n"
        "```\n\n"
        "优点：过程成本清晰，可以直接抑制每天大水大肥。缺点：成本系数需要参数扫描。\n\n"
        "## Candidate B: terminal yield plus process cost\n\n"
        "```text\n"
        "daily_reward = - daily_irrigation_cost - daily_nitrogen_cost\n"
        "terminal_reward = final_yield_value - total_irrigation_cost - total_nitrogen_cost\n"
        "```\n\n"
        "优点：经济解释更强。缺点：terminal reward 稀疏，训练可能更慢，需要更稳定的 PPO 设置。\n",
        encoding="utf-8",
    )
    return review, candidates


def compare_current_action_safe() -> Path:
    current = read_csv(CURRENT_ROOT / "evaluation" / "ppo_evaluation_summary.csv")
    safe = read_csv(DEBUG_ROOT / "evaluation" / "action_safe_debug_evaluation_summary.csv")
    if safe.empty:
        safe = read_csv(DEBUG_ROOT / "evaluation" / "ppo_evaluation_summary.csv")
    if not safe.empty:
        safe_out = DEBUG_ROOT / "evaluation" / "action_safe_debug_evaluation_summary.csv"
        safe.to_csv(safe_out, index=False, encoding="utf-8-sig")
    rows = []
    for name, df in [("current_ppo", current), ("action_safe_ppo", safe)]:
        if df.empty:
            continue
        for _, row in df.iterrows():
            rows.append(
                {
                    "model_type": name,
                    "eval_year": int(row["eval_year"]),
                    "final_grnwt": float(row["final_grnwt"]),
                    "total_irrigation": float(row["total_irrigation"]),
                    "total_n_fertilizer": float(row["total_n_fertilizer"]),
                    "mean_reward": float(row["mean_reward"]),
                    "sum_reward": float(row["sum_reward"]),
                    "run_status": row["run_status"],
                    "daily_csv_path": row["daily_csv_path"],
                    "figure_dir": row["figure_dir"],
                }
            )
    out = DEBUG_ROOT / "evaluation" / "current_vs_action_safe_comparison.csv"
    comp = pd.DataFrame(rows)
    comp.to_csv(out, index=False, encoding="utf-8-sig")
    if not comp.empty:
        metrics = [
            ("total_irrigation", "current_vs_action_safe_total_irrigation.png"),
            ("total_n_fertilizer", "current_vs_action_safe_total_n.png"),
            ("final_grnwt", "current_vs_action_safe_final_grnwt.png"),
            ("mean_reward", "current_vs_action_safe_reward.png"),
        ]
        for metric, filename in metrics:
            fig, ax = plt.subplots(figsize=(8, 4.5))
            comp.pivot(index="eval_year", columns="model_type", values=metric).plot(kind="bar", ax=ax)
            ax.set_ylabel(metric)
            ax.grid(axis="y", alpha=0.25)
            fig.tight_layout()
            fig.savefig(DEBUG_ROOT / "figures" / filename, dpi=150)
            plt.close(fig)
        for year in sorted(comp["eval_year"].unique()):
            cur = CURRENT_ROOT / "daily_outputs" / "HLA" / f"HLA_train2007_eval{year}_seed0_daily.csv"
            safe_daily = DEBUG_ROOT / "daily_outputs" / "HLA" / f"HLA_train2007_eval{year}_seed0_daily.csv"
            if cur.exists() and safe_daily.exists():
                cur_df = pd.read_csv(cur)
                safe_df = pd.read_csv(safe_daily)
                fig, axes = plt.subplots(2, 1, figsize=(9, 6), sharex=True)
                axes[0].plot(cur_df["dap"], cur_df["real_action_amir"], label="current amir", alpha=0.8)
                axes[0].plot(safe_df["dap"], safe_df["real_action_amir"], label="safe amir", alpha=0.8)
                axes[0].legend(frameon=False)
                axes[0].grid(alpha=0.25)
                axes[1].plot(cur_df["dap"], cur_df["real_action_anfer"], label="current anfer", alpha=0.8)
                axes[1].plot(safe_df["dap"], safe_df["real_action_anfer"], label="safe anfer", alpha=0.8)
                axes[1].legend(frameon=False)
                axes[1].grid(alpha=0.25)
                axes[1].set_xlabel("DAP")
                fig.tight_layout()
                fig.savefig(DEBUG_ROOT / "figures" / f"current_vs_action_safe_daily_actions_eval{year}.png", dpi=150)
                plt.close(fig)
    return out


def md_table(df: pd.DataFrame, max_rows: int = 20) -> str:
    if df.empty:
        return "_No data._"
    view = df.head(max_rows).copy()
    cols = list(view.columns)
    lines = ["| " + " | ".join(cols) + " |", "| " + " | ".join(["---"] * len(cols)) + " |"]
    for _, row in view.iterrows():
        values = []
        for col in cols:
            value = row[col]
            if pd.isna(value):
                values.append("")
            elif isinstance(value, float):
                values.append(f"{value:.4g}")
            else:
                values.append(str(value))
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def write_report_and_ppt() -> None:
    action_summary = read_csv(DEBUG_ROOT / "action_diagnostics" / "current_ppo_action_summary.csv")
    action_space = read_csv(DEBUG_ROOT / "action_diagnostics" / "ppo_action_space_check.csv")
    recalculation = read_csv(DEBUG_ROOT / "action_diagnostics" / "total_input_recalculation_check.csv")
    comparison = read_csv(compare_current_action_safe())
    safe = comparison[comparison["model_type"] == "action_safe_ppo"].copy() if not comparison.empty else pd.DataFrame()
    within_range = bool((safe["total_irrigation"] <= 300).all() and (safe["total_n_fertilizer"] <= 400).all()) if not safe.empty else False
    REPORT_MD.parent.mkdir(parents=True, exist_ok=True)
    (DEBUG_ROOT / "reports").mkdir(parents=True, exist_ok=True)
    REPORT_MD.write_text(
        "# PPO action scale and reward debug report\n\n"
        "## 1. 为什么不能直接批量 PPO\n\n"
        "上一阶段 HLA 2007 debug PPO 能跑通，但 total_irrigation 超过 3000 mm，total_n_fertilizer 超过 14000 kg/ha，远高于 fixed_high_input 的 120 mm / 250 kg/ha，因此不能直接批量训练。\n\n"
        "## 2. 当前 PPO 每日动作诊断\n\n"
        + md_table(action_summary)
        + "\n\n## 3. action_space 和动作转换\n\n"
        + md_table(action_space)
        + "\n\n## 4. total input 重算\n\n"
        + md_table(recalculation)
        + "\n\n## 5. reward 诊断\n\n"
        f"- 当前 reward review: `{rel(DEBUG_ROOT / 'reward_diagnostics' / 'current_reward_function_review.md')}`\n"
        f"- reward candidates: `{rel(DEBUG_ROOT / 'reward_diagnostics' / 'reward_revision_candidates.md')}`\n"
        "- 诊断结论：当前 reward/成本不足以在 debug PPO 中阻止高频水氮动作；本阶段先不改 reward，只隔离测试 action safety。\n\n"
        "## 6. action safety 方案\n\n"
        "参数来自 `experiments/ppo_observed_years/config_ppo_action_safe_debug.yaml`：daily irrigation <= 40 mm，daily N <= 80 kg/ha，season irrigation <= 200 mm，season N <= 300 kg/ha，并限制最小间隔和 DAP 范围。\n\n"
        "## 7. action-safe debug PPO 新旧对比\n\n"
        + md_table(comparison)
        + "\n\n## 8. 是否可以进入批量训练\n\n"
        + ("可以进入启用 action safety 的逐模型小批量训练准备，但不能使用无 safety 的旧 PPO 直接批量训练；仍建议先跑 HLA/SYA/LCA 每个模型的 pretrain smoke check。" if within_range else "暂不建议进入批量训练；需要继续修正 action 约束或 reward。")
        + "\n",
        encoding="utf-8",
    )
    shutil.copy2(REPORT_MD, DEBUG_ROOT / "reports" / REPORT_MD.name)

    if Presentation is None:
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def title(slide, text):
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text = text
        for r in p.runs:
            r.font.name = "Microsoft YaHei"
            r.font.size = Pt(23)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)

    def bullets(slide, items):
        box = slide.shapes.add_textbox(Inches(0.65), Inches(1.05), Inches(12), Inches(5.8))
        tf = box.text_frame
        tf.clear()
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = item
            for r in p.runs:
                r.font.name = "Microsoft YaHei"
                r.font.size = Pt(16)
                r.font.color.rgb = RGBColor(0, 0, 0)

    def table(slide, df, max_rows=8):
        if df.empty:
            bullets(slide, ["No data."])
            return
        data = df.head(max_rows)
        shape = slide.shapes.add_table(len(data) + 1, len(data.columns), Inches(0.45), Inches(1.05), Inches(12.4), Inches(5.8))
        tbl = shape.table
        for j, col in enumerate(data.columns):
            cell = tbl.cell(0, j)
            cell.text = str(col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.name = "Microsoft YaHei"
                    r.font.size = Pt(8)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
        for i, (_, row) in enumerate(data.iterrows(), start=1):
            for j, value in enumerate(row):
                cell = tbl.cell(i, j)
                cell.text = "" if pd.isna(value) else str(value)
                if i % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = LIGHT_BLUE
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for r in p.runs:
                        r.font.name = "Microsoft YaHei"
                        r.font.size = Pt(7)

    slides = [
        ("问题背景", ["当前 HLA 2007 debug PPO 可训练可评估，但水氮投入极端偏高。", "这说明流程跑通不等于策略有效，批量训练前必须先诊断动作尺度和 reward。"]),
        ("action safety 设计", ["每日灌水上限 40 mm，每日施氮上限 80 kg/ha。", "季节灌水软上限 200 mm，季节施氮软上限 300 kg/ha。", "限制最小间隔和 DAP 范围，所有裁剪写入 daily CSV。"]),
        ("reward 函数诊断", ["当前 reward/成本不足以阻止 debug PPO 高频高量动作。", "本阶段不覆盖原 reward，只生成 reward revision candidates。", "后续如仍需提高策略质量，再做 reward 参数扫描。"]),
    ]
    for t, b in slides:
        s = prs.slides.add_slide(blank)
        title(s, t)
        bullets(s, b)
    s = prs.slides.add_slide(blank)
    title(s, "当前 PPO 动作诊断")
    table(s, action_summary[["eval_year", "days", "nonzero_irrigation_days", "nonzero_fertilization_days", "sum_irrigation", "sum_fertilization"]])
    s = prs.slides.add_slide(blank)
    title(s, "action_space 检查")
    table(s, action_space[["action_name", "action_space_low", "action_space_high", "example_normalized_minus1", "example_normalized_0", "example_normalized_plus1"]])
    s = prs.slides.add_slide(blank)
    title(s, "新旧 PPO 对比")
    table(s, comparison[["model_type", "eval_year", "final_grnwt", "total_irrigation", "total_n_fertilizer", "mean_reward"]], max_rows=8)
    s = prs.slides.add_slide(blank)
    title(s, "下一步建议")
    bullets(s, ["action safety 已用于隔离动作尺度问题。", "若水氮进入合理范围，可进入逐模型训练准备。", "若产量或 reward 仍差，再调整 reward，而不是直接扩大批量训练。"])
    prs.save(REPORT_PPTX)
    shutil.copy2(REPORT_PPTX, DEBUG_ROOT / "reports" / REPORT_PPTX.name)


def main() -> None:
    ensure_dirs()
    shutil.copy2(CONFIG_SAFE, DEBUG_ROOT / "config_versions" / CONFIG_SAFE.name)
    current_action_diagnostics()
    if not (DEBUG_ROOT / "action_diagnostics" / "ppo_action_space_check.csv").exists():
        action_space_check()
    total_recalculation_check()
    reward_review()
    compare_current_action_safe()
    write_report_and_ppt()
    print(rel(DEBUG_ROOT / "action_diagnostics" / "current_ppo_action_summary.csv"))
    print(rel(REPORT_MD))
    print(rel(REPORT_PPTX))


if __name__ == "__main__":
    main()
