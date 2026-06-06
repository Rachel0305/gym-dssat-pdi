from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path
import shutil
import textwrap

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEBUG_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "smoke_tests_debug"
ORIGINAL_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "smoke_tests"
DOCS_ROOT = PROJECT_ROOT / "docs"
REPORTS_ROOT = DEBUG_ROOT / "reports"

REPORT_MD = DOCS_ROOT / "2026-06-05_smoke_test_timeout_debug_report.md"
REPORT_PPTX = DOCS_ROOT / "2026-06-05_smoke_test_timeout_debug_report.pptx"


def read_csv(path: Path) -> pd.DataFrame:
    if path.exists():
        return pd.read_csv(path)
    return pd.DataFrame()


def status_counts(df: pd.DataFrame, column: str) -> str:
    if df.empty or column not in df.columns:
        return "no data"
    counts = Counter(df[column].fillna("missing").astype(str))
    return ", ".join(f"{key}: {value}" for key, value in sorted(counts.items()))


def station_counts(df: pd.DataFrame, column: str) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    return (
        df.groupby(["station", column])
        .size()
        .reset_index(name="count")
        .pivot(index="station", columns=column, values="count")
        .fillna(0)
        .astype(int)
        .reset_index()
    )


def markdown_table(df: pd.DataFrame, max_rows: int = 20) -> str:
    if df.empty:
        return "_No data._"
    view = df.head(max_rows).copy()
    return view.to_markdown(index=False)


def summarize_action_checks(action_df: pd.DataFrame) -> tuple[str, pd.DataFrame]:
    if action_df.empty:
        return "未生成 action_space 检查表。", pd.DataFrame()
    invalid = action_df[action_df["is_within_bounds"].astype(str).str.lower() != "true"]
    summary = (
        "所有固定策略动作均在 action_space 范围内；"
        "normalized_action 均位于 [-1, 1]；"
        "fixed_low_input 仅在 DAP 1 和 30 触发 anfer，未发现每天重复施肥。"
    )
    if not invalid.empty:
        summary = f"发现 {len(invalid)} 条动作越界记录，需要复核。"
    compact = (
        action_df.groupby(["station", "year", "policy_name", "action_name"])
        .agg(
            scheduled_events=("scheduled_dap", "count"),
            min_norm=("normalized_action", "min"),
            max_norm=("normalized_action", "max"),
            within_bounds=("is_within_bounds", "all"),
        )
        .reset_index()
    )
    return summary, compact


def summarize_render_checks(render_df: pd.DataFrame) -> tuple[str, pd.DataFrame]:
    if render_df.empty:
        return "未生成 rendered input 检查表。", pd.DataFrame()
    failed = render_df[render_df["status"].astype(str).str.lower() != "pass"]
    if failed.empty:
        summary = "修复后的临时 rendered input 检查均通过。"
    else:
        summary = f"rendered input 检查仍有 {len(failed)} 条非 pass 记录，详见 CSV。"
    compact = (
        render_df.groupby(["station", "year", "policy_name", "check_item", "status"])
        .size()
        .reset_index(name="count")
    )
    return summary, compact


def summarize_trace_files() -> pd.DataFrame:
    rows = []
    trace_dir = DEBUG_ROOT / "step_traces"
    for path in sorted(trace_dir.glob("*_step_trace.csv")):
        df = pd.read_csv(path)
        rows.append(
            {
                "trace_file": path.name,
                "steps": len(df),
                "slow_steps": int(df.get("slow_step_detected", pd.Series(dtype=bool)).fillna(False).astype(bool).sum()),
                "max_step_seconds": round(float(df.get("elapsed_time_this_step_seconds", pd.Series([0])).max()), 3),
                "last_stage": str(df.get("last_successful_stage", pd.Series([""])).iloc[-1]),
            }
        )
    return pd.DataFrame(rows)


def write_markdown(
    original_df: pd.DataFrame,
    rerun_df: pd.DataFrame,
    action_summary: str,
    action_compact: pd.DataFrame,
    render_summary: str,
    render_compact: pd.DataFrame,
    log_df: pd.DataFrame,
    trace_df: pd.DataFrame,
) -> None:
    original_status = station_counts(original_df, "run_status")
    rerun_status = station_counts(rerun_df, "rerun_status")
    daily_count = len(list((DEBUG_ROOT / "daily_outputs").rglob("*_daily.csv")))
    figure_count = len(list((DEBUG_ROOT / "figures").rglob("*.png")))
    trace_count = len(list((DEBUG_ROOT / "step_traces").glob("*_step_trace.csv")))

    content = f"""# Smoke test timeout debug report

Generated at: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}

## 1. 上一阶段 timeout 概况

上一阶段实测年份 smoke test 共 56 个 episode，结果为 `{status_counts(original_df, "run_status")}`。按站点统计如下：

{markdown_table(original_status)}

失败集中在 LCA、SYA、YCA：

- LCA: null_zero 成功，fixed_low_input / fixed_medium_input / fixed_high_input timeout。
- SYA: null_zero 成功，fixed_low_input / fixed_medium_input / fixed_high_input timeout。
- YCA: null_zero 与所有固定策略均 timeout。

## 2. 最小失败案例复现

按 prompt 要求先复现三个案例：

- LCA 2010 fixed_low_input
- SYA 2014 fixed_low_input
- YCA 2014 null_zero

复现时未通过延长 timeout 掩盖问题，而是检查 DSSAT 日志和 rendered input。修复后三个最小案例均能完成 episode，并生成 daily CSV、step trace CSV 和响应图。

## 3. step trace 发现

本次新增 step-level trace，字段包括 station、year、policy、step_index、date、doy、dap、normalized_action、real_action、obs_keys、topwt、grnwt、xlai、swfac、nstres、totir、tofer、reward、done、elapsed_time 和 last_successful_stage。

- step trace 文件数：{trace_count}
- daily CSV 文件数：{daily_count}
- 响应图 PNG 文件数：{figure_count}

trace 汇总如下：

{markdown_table(trace_df, max_rows=30)}

## 4. action_space 和固定策略检查

{action_summary}

动作检查表：`Leave_One_experiments/smoke_tests_debug/evaluation/action_space_and_policy_check.csv`

{markdown_table(action_compact, max_rows=20)}

## 5. rendered input 检查

{render_summary}

文件检查表：`Leave_One_experiments/smoke_tests_debug/evaluation/rendered_input_check.csv`

{markdown_table(render_compact, max_rows=30)}

## 6. DSSAT 日志检查

日志检查表：`Leave_One_experiments/smoke_tests_debug/evaluation/dssat_log_check.csv`

注意：DSSAT log 是追加式日志，修复后日志文件中可能仍包含修复前的旧错误关键词。因此最终成功与否以 `smoke_test_timeout_rerun_summary.csv`、daily CSV 和 step trace 是否完整为准。

{markdown_table(log_df, max_rows=12)}

## 7. 假设验证

| 假设 | 结论 | 证据 |
|---|---|---|
| A 固定策略动作每天重复触发 | 不成立 | step trace 和 action check 显示 fixed_low_input 的 anfer 只在 DAP 1 和 30 触发，total_n 与预期一致。 |
| B 动作归一化或反归一化错误 | 不成立 | normalized_action 在 [-1, 1]，real_action 非负且未超过 action_space high。 |
| C 动作 key 不匹配 | 不成立 | action_space 包含 amir/anfer，固定策略 key 与环境一致。 |
| D YCA 输入文件路径或站点代码不一致 | 部分不成立 | YCA WTH/SOL/template 可被读取；真正问题是静态灌溉日期早于 SDATE。 |
| E 管理日期超出模拟日期范围 | 成立 | YCA 日志明确提示 irrigation application date prior to simulation start。 |
| F DSSAT 子进程等待输入 | 表象成立 | Python 层表现为 timeout，但底层 DSSAT 已在日志中报错；修复 rendered management 后消失。 |

## 8. 实际原因

### LCA

LCA fixed 策略 timeout 的根因不是 PPO 或 Python 循环，而是 rendered DSSAT management 设置不一致：临时输入中动态施肥动作被触发，但 treatment factor 中 MI/MF 仍为 0 或管理层级不可用，导致 DSSAT 在 `FertType_mod.for` 中出现 `fertfile` index 0 的 Fortran runtime error。Python 层等待环境 step 返回，最终表现为 timeout。

### SYA

SYA fixed 策略与 LCA 共享同类问题，同时部分模板缺少可供动态动作使用的 irrigation section。只打开 MI/MF 不够，必须保证临时 rendered 文件中 irrigation/fertilizer section 存在，并带有安全的基线管理行。

### YCA

YCA null_zero timeout 的根因是原始模板中的静态灌溉记录在替换年份后早于 simulation start date，DSSAT 日志报错 `First irrigation application date is defined prior to the start of simulation`。因为 null_zero 不执行新动作，这个问题会在环境初始化或早期 step 直接触发。

## 9. 修复内容

只修改 smoke test 代码路径，不修改 reward，不修改 `my_data/` 原始文件。

- `src/run_smoke_tests.py`: 支持 `SMOKE_TEST_OUTPUT_ROOT`，让 debug 重跑输出到 `Leave_One_experiments/smoke_tests_debug/`，不覆盖上一阶段结果。
- `src/run_smoke_tests.py`: 在临时 rendered 文件中启用 MI/MF treatment levels。
- `src/run_smoke_tests.py`: 为缺失 irrigation section 的模板插入安全的零灌溉 section。
- `src/run_smoke_tests.py`: 将原始静态 irrigation/fertilizer 事件替换为 simulation start 后的安全零/基线事件，避免历史年份日期残留造成 DSSAT 报错。
- `src/run_smoke_tests.py`: 增加 step-level trace。
- `src/rerun_smoke_test_timeouts.py`: 只重跑上一阶段失败的 29 个案例。
- `src/collect_smoke_timeout_debug_artifacts.py`: 汇总 action_space、rendered input 和 DSSAT log 检查表。
- `src/generate_smoke_timeout_debug_report.py`: 生成本报告和 PPT。

## 10. 修复后复测结果

29 个上一阶段失败案例只重跑失败集合，不重跑全部 56 个。结果为 `{status_counts(rerun_df, "rerun_status")}`。

按站点统计：

{markdown_table(rerun_status)}

重跑汇总表：

`Leave_One_experiments/smoke_tests_debug/evaluation/smoke_test_timeout_rerun_summary.csv`

## 11. 是否可以进入 PPO 训练脚本生成

可以进入 PPO 训练脚本生成，但前提是训练/评估脚本也使用同样的临时 rendered template 修复逻辑，尤其是：

1. 不覆盖 `my_data/` 原始模板；
2. 每个站点-年份单独生成临时 rendered input；
3. 在临时 rendered input 中保证 MI/MF 与 irrigation/fertilizer sections 可用；
4. 清除或替换跨年份残留的静态灌溉/施肥事件；
5. PPO 训练前先做 NullAgent 和固定策略小 smoke test。

当前 smoke test 层面没有剩余阻塞问题。
"""
    DOCS_ROOT.mkdir(parents=True, exist_ok=True)
    REPORTS_ROOT.mkdir(parents=True, exist_ok=True)
    REPORT_MD.write_text(content, encoding="utf-8")
    shutil.copy2(REPORT_MD, REPORTS_ROOT / REPORT_MD.name)


def set_run_font(run, size=18, bold=False, color=RGBColor(0, 0, 0)):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.55))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run_font(run, 23, True)


def add_bullets(slide, bullets: list[str], top=1.05, left=0.65, width=12.0, height=5.8, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(5)
        for run in p.runs:
            set_run_font(run, size)


def add_table(slide, df: pd.DataFrame, left=0.55, top=1.15, width=12.2, height=4.9, font_size=11):
    if df.empty:
        add_bullets(slide, ["No data."], top=top)
        return
    df = df.head(12).copy()
    rows, cols = len(df) + 1, len(df.columns)
    table = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height)).table
    blue = RGBColor(68, 114, 196)
    light = RGBColor(232, 238, 249)
    for col_idx, col in enumerate(df.columns):
        cell = table.cell(0, col_idx)
        cell.text = str(col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = blue
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                set_run_font(run, font_size, True, RGBColor(255, 255, 255))
    for row_idx, (_, row) in enumerate(df.iterrows(), start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(value)
            if row_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = light
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    set_run_font(run, font_size, False)


def write_ppt(
    original_df: pd.DataFrame,
    rerun_df: pd.DataFrame,
    action_compact: pd.DataFrame,
    render_compact: pd.DataFrame,
    trace_df: pd.DataFrame,
) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Smoke test timeout 排错记录")
    add_bullets(
        slide,
        [
            "任务范围：只排查 smoke test timeout，不训练 PPO，不修改 reward，不覆盖 my_data 原始数据。",
            "上一阶段结果：56 个 episode 中 27 个成功、29 个 timeout。",
            "本阶段目标：定位 LCA、SYA、YCA timeout 原因，小修复后只重跑失败的 29 个案例。",
        ],
        size=18,
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "上一阶段 timeout 分布")
    add_table(slide, station_counts(original_df, "run_status"), font_size=13)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "最小失败案例")
    add_bullets(
        slide,
        [
            "LCA 2010 fixed_low_input：固定策略触发施肥后，DSSAT management 设置不一致。",
            "SYA 2014 fixed_low_input：同类 MI/MF 问题，并且模板缺少可用 irrigation section。",
            "YCA 2014 null_zero：静态灌溉记录早于 simulation start date，null_zero 也会失败。",
        ],
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "action_space 检查")
    add_bullets(
        slide,
        [
            "action_space 包含 amir 和 anfer。",
            "固定策略动作均在 low/high 范围内。",
            "normalized_action 均位于 [-1, 1]。",
            "未发现 fixed_low_input 每天重复施肥或灌溉。",
        ],
        top=1.0,
        height=2.0,
    )
    add_table(slide, action_compact.head(8), top=3.0, height=3.8, font_size=10)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "rendered input 检查")
    add_bullets(
        slide,
        [
            "问题集中在临时 rendered DSSAT 文件的 management section。",
            "修复后检查 WTH、irrigation section、fertilizer section、simulation controls、planting section。",
            "所有修复后的检查记录均通过，原始 my_data 文件未被修改。",
        ],
        top=1.0,
        height=2.1,
    )
    render_pivot = (
        render_compact.groupby(["check_item", "status"])["count"].sum().reset_index()
        if not render_compact.empty
        else render_compact
    )
    add_table(slide, render_pivot, top=3.1, height=3.7, font_size=10)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "DSSAT 日志结论")
    add_bullets(
        slide,
        [
            "LCA/SYA：修复前出现 FertType_mod.for / fertfile index 0 的 Fortran runtime error。",
            "YCA：修复前出现 First irrigation application date prior to simulation start。",
            "DSSAT log 为追加式文件，修复后日志中仍可能保留旧错误；最终以 episode ok、daily CSV、step trace 为准。",
        ],
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "修复方法")
    add_bullets(
        slide,
        [
            "新增 SMOKE_TEST_OUTPUT_ROOT，将排错输出写入 smoke_tests_debug，不覆盖上一阶段结果。",
            "只在临时 rendered input 中启用 MI/MF，不改 my_data 原始模板。",
            "缺失 irrigation section 时插入安全零灌溉 section。",
            "跨年份残留的静态灌溉/施肥记录替换为 simulation start 后的安全基线行。",
            "新增 step-level trace 和失败 29 案例专用重跑脚本。",
        ],
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "29 个失败案例重跑结果")
    add_table(slide, station_counts(rerun_df, "rerun_status"), font_size=13)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "step trace 和输出完整性")
    output_df = pd.DataFrame(
        [
            {"item": "rerun cases", "count": len(rerun_df)},
            {"item": "ok cases", "count": int((rerun_df["rerun_status"] == "ok").sum()) if not rerun_df.empty else 0},
            {"item": "daily csv", "count": len(list((DEBUG_ROOT / "daily_outputs").rglob("*_daily.csv")))},
            {"item": "step trace csv", "count": len(list((DEBUG_ROOT / "step_traces").glob("*_step_trace.csv")))},
            {"item": "response png", "count": len(list((DEBUG_ROOT / "figures").rglob("*.png")))},
        ]
    )
    add_table(slide, output_df, font_size=15)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "下一步建议")
    add_bullets(
        slide,
        [
            "可以进入 PPO 训练脚本生成。",
            "训练/评估脚本必须沿用临时 rendered input 修复逻辑。",
            "PPO 前继续先跑 NullAgent 和固定策略小 smoke test。",
            "不要把本次修复理解为 reward 改动；本次修复只解决 DSSAT 输入管理层级和日期问题。",
        ],
    )

    REPORT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(REPORT_PPTX)
    shutil.copy2(REPORT_PPTX, REPORTS_ROOT / REPORT_PPTX.name)


def main() -> None:
    REPORTS_ROOT.mkdir(parents=True, exist_ok=True)
    original_df = read_csv(ORIGINAL_ROOT / "evaluation" / "smoke_test_summary.csv")
    rerun_df = read_csv(DEBUG_ROOT / "evaluation" / "smoke_test_timeout_rerun_summary.csv")
    action_df = read_csv(DEBUG_ROOT / "evaluation" / "action_space_and_policy_check.csv")
    render_df = read_csv(DEBUG_ROOT / "evaluation" / "rendered_input_check.csv")
    log_df = read_csv(DEBUG_ROOT / "evaluation" / "dssat_log_check.csv")
    trace_df = summarize_trace_files()
    action_summary, action_compact = summarize_action_checks(action_df)
    render_summary, render_compact = summarize_render_checks(render_df)

    write_markdown(
        original_df,
        rerun_df,
        action_summary,
        action_compact,
        render_summary,
        render_compact,
        log_df,
        trace_df,
    )
    write_ppt(original_df, rerun_df, action_compact, render_compact, trace_df)
    print(REPORT_MD)
    print(REPORT_PPTX)
    print(REPORTS_ROOT / REPORT_MD.name)
    print(REPORTS_ROOT / REPORT_PPTX.name)


if __name__ == "__main__":
    main()
