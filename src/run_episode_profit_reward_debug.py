from __future__ import annotations

import shutil
import traceback
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from episode_profit_reward import (
    PROFIT_REWARD_CANDIDATES,
    EpisodeProfitRewardWrapper,
    candidate_table,
)
from ppo_evaluate import latest_observation_dict, make_env, normalized_action_dict
from ppo_experiment_plan import find_year
from ppo_plot_results import plot_episode
from ppo_safe_rendering import PROJECT_ROOT, load_yaml
from ppo_train import ppo_kwargs, run_pretrain_smoke_check


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_episode_profit_reward_debug.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "episode_profit_reward_debug"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_episode_profit_reward_debug_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_episode_profit_reward_debug_report.pptx"


def ensure_dirs() -> None:
    for sub in [
        "configs",
        "reward_versions",
        "smoke_checks",
        "rendered_inputs",
        "models/HLA",
        "logs/HLA",
        "tensorboard/HLA",
        "daily_outputs/HLA",
        "evaluation",
        "figures/summary",
        "reports",
    ]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)


def df_to_markdown(df: pd.DataFrame) -> str:
    if df.empty:
        return ""
    work = df.copy()
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    separator = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, separator, *rows])


def write_interface_check() -> Path:
    text = """# Terminal reward interface check

Generated at: 2026-06-06

## Answers required by prompt 006_05

1. The original gym-DSSAT reward function can access `_next_state`.
2. The original reward function cannot reliably judge episode done because the callback signature is `_previous_state, _next_state, _history, _cultivar` and has no done flag.
3. The original reward function can access `_history`.
4. `_history` includes actions in previous training/evaluation traces.
5. `_history` includes daily observations through the environment history used by the wrapper/evaluation utilities.
6. `final_grnwt` can be read from the final observation / history after the environment step returns done.
7. Daily `amir` and `anfer` can be accumulated reliably from `SafeActionWrapper.last_safety_result.safe_real_action`.
8. Terminal reward can be added on the last step at the gymnasium wrapper layer, after `terminated` or `truncated` is known.
9. Because the original reward callback has no done flag, terminal reward should be added in a new project-level wrapper, not in site-packages.
10. Minimal implementation: keep `site-packages` reward unchanged, keep `src/ppo_train.py` unchanged, add `src/episode_profit_reward.py` and a dedicated runner that wraps the action-safe env.

## Feasibility conclusion

Terminal reward is feasible through `EpisodeProfitRewardWrapper`. This is the least invasive design: action safety still controls the physical water/N limits, while the profit wrapper replaces the training reward with daily input cost plus terminal grain-yield profit.
"""
    out = OUTPUT_ROOT / "evaluation" / "reward_interface_check.md"
    out.write_text(text, encoding="utf-8")
    return out


def write_design_files(config: dict) -> None:
    table = pd.DataFrame(candidate_table())
    table.to_csv(OUTPUT_ROOT / "reward_versions" / "episode_profit_reward_candidate_table.csv", index=False, encoding="utf-8-sig")
    shutil.copy2(PROJECT_ROOT / "src" / "episode_profit_reward.py", OUTPUT_ROOT / "reward_versions" / "episode_profit_reward.py")
    shutil.copy2(CONFIG_PATH, OUTPUT_ROOT / "configs" / CONFIG_PATH.name)
    plan = []
    for order, version in enumerate(config["reward_candidates"]["versions"], start=1):
        cand = PROFIT_REWARD_CANDIDATES[version]
        plan.append(
            {
                "run_order": order,
                "station": config["training"]["hla_pilot_station"],
                "train_year": config["training"]["hla_train_year"],
                "eval_years": ",".join(str(y) for y in config["training"]["hla_eval_years"]),
                "reward_version": version,
                "reward_family": cand.family,
                "season_irrigation_cap": config["cap"]["season_irrigation_cap"],
                "season_n_cap": config["cap"]["season_n_cap"],
                "seed": config["seed"],
                "total_timesteps": config["training"]["total_timesteps"],
            }
        )
    pd.DataFrame(plan).to_csv(OUTPUT_ROOT / "configs" / "episode_profit_reward_plan.csv", index=False, encoding="utf-8-sig")
    design = """# Episode-level profit reward design

This stage does not overwrite the original gym-DSSAT reward file. The new reward is injected by `EpisodeProfitRewardWrapper`.

For P1-P6:

```text
daily_reward = - daily_water_cost * daily_irrigation - daily_n_cost * daily_n
terminal_reward = grain_value_coef * final_grnwt
                  - season_water_cost * total_irrigation
                  - season_n_cost * total_n
```

P0 is a pass-through baseline using the current environment reward under the same action-safety cap.

The coefficients are normalized scores, not real RMB prices.

## Candidate table

""" + df_to_markdown(table)
    (OUTPUT_ROOT / "reward_versions" / "episode_profit_reward_design.md").write_text(design, encoding="utf-8")


def row_metadata(config: dict, reward_version: str, station: str) -> dict:
    cand = PROFIT_REWARD_CANDIDATES[reward_version]
    return {
        "reward_version": reward_version,
        "reward_family": cand.family,
        "cap_name": config["cap"]["cap_name"],
        "season_irrigation_cap": float(config["cap"]["season_irrigation_cap"]),
        "season_n_cap": float(config["cap"]["season_n_cap"]),
        "cv_type": config["training"]["cv_types"].get(station, ""),
    }


def train_profit_policy(config: dict, reward_version: str, station: str, train_year: int):
    from stable_baselines3 import PPO

    seed = int(config["seed"])
    meta = row_metadata(config, reward_version, station)
    smoke_ok, smoke_summary = run_pretrain_smoke_check(config, station, train_year, seed, row_metadata=meta)
    if not smoke_ok:
        raise RuntimeError(f"pretrain smoke check failed: {smoke_summary}")

    policy_name = f"{station}_train{train_year}_{reward_version}_seed{seed}_episode_profit_action_safe"
    env = make_env(config, station, train_year, seed, run_tag=f"{policy_name}_train", evaluation=False, action_safety_enabled=True)
    env = EpisodeProfitRewardWrapper(env, reward_version)
    model_dir = OUTPUT_ROOT / "models" / station
    tensorboard_dir = OUTPUT_ROOT / "tensorboard" / station
    model_dir.mkdir(parents=True, exist_ok=True)
    tensorboard_dir.mkdir(parents=True, exist_ok=True)
    model_path = model_dir / policy_name
    try:
        model = PPO(
            "MlpPolicy",
            env,
            verbose=1,
            seed=seed,
            tensorboard_log=str(tensorboard_dir),
            **ppo_kwargs(config, debug=False),
        )
        model.learn(total_timesteps=int(config["training"]["total_timesteps"]), progress_bar=False)
        model.save(str(model_path))
    finally:
        try:
            env.close()
        except Exception:
            pass
    return model, model_path.with_suffix(".zip"), policy_name


def plot_terminal_breakdown(daily: pd.DataFrame, figure_dir: Path) -> Path:
    figure_dir.mkdir(parents=True, exist_ok=True)
    last = daily.iloc[-1] if len(daily) else pd.Series(dtype=float)
    labels = ["grain_value", "water_cost", "n_cost", "profit_score"]
    values = [
        float(last.get("grain_value_component", 0.0) or 0.0),
        -float(last.get("water_cost_component", 0.0) or 0.0),
        -float(last.get("n_cost_component", 0.0) or 0.0),
        float(last.get("season_profit_score", 0.0) or 0.0),
    ]
    fig, ax = plt.subplots(figsize=(7.5, 4.5))
    ax.bar(labels, values, color=["#4c78a8", "#f58518", "#e45756", "#54a24b"])
    ax.axhline(0, color="black", linewidth=0.8)
    ax.set_title("Terminal profit breakdown")
    ax.set_ylabel("Normalized score")
    fig.tight_layout()
    path = figure_dir / "terminal_profit_breakdown.png"
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def evaluate_profit_model(
    model,
    config: dict,
    reward_version: str,
    station: str,
    train_year: int,
    eval_year: int,
    model_path: Path,
    policy_name: str,
) -> dict:
    seed = int(config["seed"])
    train_info = find_year(config, station, train_year)
    eval_info = find_year(config, station, eval_year)
    env = make_env(config, station, eval_year, seed, run_tag=f"{policy_name}_eval{eval_year}", evaluation=True, action_safety_enabled=True)
    env = EpisodeProfitRewardWrapper(env, reward_version)
    records: list[dict] = []
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        year_info = find_year(config, station, eval_year)
        while not done and step_count < int(config.get("runtime", {}).get("max_steps", 260)):
            latest = latest_observation_dict(env, obs, info)
            dap = int(round(float(latest.get("dap", step_count) or step_count)))
            raw_action, _ = model.predict(obs, deterministic=True)
            raw_action = np.clip(np.asarray(raw_action, dtype=np.float32).flatten(), -1.0, 1.0)
            obs, reward, terminated, truncated, info = env.step(raw_action)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            safety_result = getattr(env.env, "last_safety_result", None)
            if safety_result is None:
                raw_real = {"amir": np.nan, "anfer": np.nan}
                safe_real = {"amir": np.nan, "anfer": np.nan}
                safety_flags = {}
            else:
                raw_real = safety_result.raw_real_action
                safe_real = safety_result.safe_real_action
                safety_flags = {
                    "action_clipped_amir": safety_result.action_clipped_amir,
                    "action_clipped_anfer": safety_result.action_clipped_anfer,
                    "season_irrigation_so_far": safety_result.season_irrigation_so_far,
                    "season_n_so_far": safety_result.season_n_so_far,
                    "safety_rule_triggered": safety_result.safety_rule_triggered,
                }
            norm_action = normalized_action_dict(env, raw_action)
            date = pd.Timestamp(year_info["planting_date"]) + pd.Timedelta(days=max(dap - 1, 0))
            totir_raw = float(latest.get("totir", np.nan)) if latest.get("totir", np.nan) is not None else np.nan
            tofer_raw = float(latest.get("tofer", np.nan)) if latest.get("tofer", np.nan) is not None else np.nan
            row = {
                "station": station,
                "policy_name": policy_name,
                "train_year": train_year,
                "train_year_label": train_info["label"],
                "eval_year": eval_year,
                "eval_year_label": eval_info["label"],
                "seed": seed,
                "date": date.strftime("%Y-%m-%d"),
                "doy": int(date.dayofyear),
                "dap": float(latest.get("dap", dap) or dap),
                "topwt": float(latest.get("topwt", np.nan)) if latest.get("topwt", np.nan) is not None else np.nan,
                "grnwt": float(latest.get("grnwt", np.nan)) if latest.get("grnwt", np.nan) is not None else np.nan,
                "xlai": float(latest.get("xlai", np.nan)) if latest.get("xlai", np.nan) is not None else np.nan,
                "totir_raw": totir_raw,
                "totir": env.total_irrigation if np.isnan(totir_raw) else totir_raw,
                "tofer_raw": tofer_raw,
                "tofer": env.total_n if np.isnan(tofer_raw) else tofer_raw,
                "swfac": float(latest.get("swfac", np.nan)) if latest.get("swfac", np.nan) is not None else np.nan,
                "nstres": float(latest.get("nstres", np.nan)) if latest.get("nstres", np.nan) is not None else np.nan,
                "reward": float(reward),
                "real_action_amir": float(safe_real.get("amir", np.nan)),
                "real_action_anfer": float(safe_real.get("anfer", np.nan)),
                "normalized_action_amir": norm_action.get("amir", np.nan),
                "normalized_action_anfer": norm_action.get("anfer", np.nan),
                "raw_real_action_amir": float(raw_real.get("amir", np.nan)),
                "raw_real_action_anfer": float(raw_real.get("anfer", np.nan)),
                "safe_real_action_amir": float(safe_real.get("amir", np.nan)),
                "safe_real_action_anfer": float(safe_real.get("anfer", np.nan)),
                "done": done,
            }
            row.update(safety_flags)
            row.update(env.last_profit_info)
            row.update(row_metadata(config, reward_version, station))
            records.append(row)
            step_count += 1
    finally:
        try:
            env.close()
        except Exception:
            pass

    daily = pd.DataFrame(records)
    daily_dir = OUTPUT_ROOT / "daily_outputs" / station
    daily_dir.mkdir(parents=True, exist_ok=True)
    daily_csv = daily_dir / f"{station}_train{train_year}_{reward_version}_eval{eval_year}_seed{seed}_daily.csv"
    daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
    fig_dir = OUTPUT_ROOT / "figures" / station / reward_version / f"eval_{eval_year}"
    plot_episode(daily, fig_dir)
    plot_terminal_breakdown(daily, fig_dir)

    episode_completed = bool(records and records[-1]["done"])
    terminal = records[-1] if records else {}
    result = {
        "station": station,
        "reward_version": reward_version,
        "reward_family": PROFIT_REWARD_CANDIDATES[reward_version].family,
        "policy_name": policy_name,
        "train_year": train_year,
        "eval_year": eval_year,
        "seed": seed,
        "season_irrigation_cap": float(config["cap"]["season_irrigation_cap"]),
        "season_n_cap": float(config["cap"]["season_n_cap"]),
        "run_status": "ok" if episode_completed else "failed",
        "episode_completed": episode_completed,
        "episode_length": len(daily),
        "final_grnwt": float(daily["grnwt"].iloc[-1]) if len(daily) else np.nan,
        "final_topwt": float(daily["topwt"].iloc[-1]) if len(daily) else np.nan,
        "total_irrigation": float(daily["real_action_amir"].sum()) if len(daily) else 0.0,
        "total_n_fertilizer": float(daily["real_action_anfer"].sum()) if len(daily) else 0.0,
        "mean_swfac": float(pd.to_numeric(daily["swfac"], errors="coerce").mean()) if len(daily) else np.nan,
        "mean_nstres": float(pd.to_numeric(daily["nstres"], errors="coerce").mean()) if len(daily) else np.nan,
        "mean_reward": float(daily["reward"].mean()) if len(daily) else np.nan,
        "sum_reward": float(daily["reward"].sum()) if len(daily) else 0.0,
        "terminal_reward": float(terminal.get("terminal_reward", 0.0) or 0.0),
        "daily_cost_total": float(terminal.get("daily_cost_total", 0.0) or 0.0),
        "season_profit_score": float(terminal.get("season_profit_score", 0.0) or 0.0),
        "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
        "figure_dir": str(fig_dir.relative_to(PROJECT_ROOT)),
        "model_path": str(model_path.relative_to(PROJECT_ROOT)),
        "notes": "episode_profit_action_safe_eval",
    }
    result["irrigation_saturation_ratio"] = result["total_irrigation"] / result["season_irrigation_cap"]
    result["n_saturation_ratio"] = result["total_n_fertilizer"] / result["season_n_cap"]
    return result


def append_summary(rows: list[dict]) -> pd.DataFrame:
    out = OUTPUT_ROOT / "evaluation" / "episode_profit_reward_evaluation_summary.csv"
    new = pd.DataFrame(rows)
    if out.exists():
        old = pd.read_csv(out)
        df = pd.concat([old, new], ignore_index=True)
        df = df.drop_duplicates(["station", "reward_version", "train_year", "eval_year", "seed"], keep="last")
    else:
        df = new
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return df


def compare_candidates(df: pd.DataFrame) -> pd.DataFrame:
    hla = df[df["station"] == "HLA"].copy()
    hla["quality_gate_pass"] = (hla["run_status"] == "ok") & hla["episode_completed"].astype(str).str.lower().isin(["true", "1"])
    grouped = (
        hla.groupby(["reward_version", "reward_family"], dropna=False)
        .agg(
            eval_count=("eval_year", "count"),
            ok_count=("quality_gate_pass", "sum"),
            mean_yield=("final_grnwt", "mean"),
            std_yield=("final_grnwt", "std"),
            mean_reward=("mean_reward", "mean"),
            mean_irrigation=("total_irrigation", "mean"),
            mean_n=("total_n_fertilizer", "mean"),
            mean_irrigation_saturation_ratio=("irrigation_saturation_ratio", "mean"),
            mean_n_saturation_ratio=("n_saturation_ratio", "mean"),
            profit_score=("season_profit_score", "mean"),
            mean_swfac=("mean_swfac", "mean"),
            mean_nstres=("mean_nstres", "mean"),
        )
        .reset_index()
    )
    base = grouped[grouped["reward_version"] == "P0_current_reward_baseline"]
    base_yield = float(base.iloc[0]["mean_yield"]) if len(base) else float(grouped["mean_yield"].max())
    base_input = 300.0 + 450.0
    grouped["yield_loss_vs_baseline"] = (base_yield - grouped["mean_yield"]) / base_yield
    grouped["input_reduction_vs_baseline"] = 1.0 - (grouped["mean_irrigation"] + grouped["mean_n"]) / base_input
    grouped["strict_pass"] = (
        (grouped["mean_irrigation_saturation_ratio"] < 0.95)
        & (grouped["mean_n_saturation_ratio"] < 0.95)
        & (grouped["yield_loss_vs_baseline"] <= 0.15)
        & (grouped["ok_count"] == grouped["eval_count"])
    )
    grouped["relaxed_pass"] = (
        (
            (grouped["mean_irrigation_saturation_ratio"] < 0.95)
            | (grouped["mean_n_saturation_ratio"] < 0.95)
        )
        & (grouped["yield_loss_vs_baseline"] <= 0.15)
        & (grouped["ok_count"] == grouped["eval_count"])
    )
    grouped["cost_too_strong"] = (
        (grouped["mean_irrigation"] < 1.0)
        & (grouped["mean_n"] < 1.0)
        & (grouped["yield_loss_vs_baseline"] > 0.15)
    )
    out = OUTPUT_ROOT / "evaluation" / "episode_profit_reward_candidate_comparison.csv"
    grouped.to_csv(out, index=False, encoding="utf-8-sig")
    return grouped


def select_candidate(comparison: pd.DataFrame) -> str | None:
    passed = comparison[comparison["strict_pass"]].copy()
    if passed.empty:
        passed = comparison[comparison["relaxed_pass"]].copy()
    if passed.empty:
        return None
    passed = passed.sort_values(["yield_loss_vs_baseline", "input_reduction_vs_baseline"], ascending=[True, False])
    return str(passed.iloc[0]["reward_version"])


def plot_summary(comparison: pd.DataFrame) -> None:
    fig_root = OUTPUT_ROOT / "figures" / "summary"
    fig_root.mkdir(parents=True, exist_ok=True)
    x = comparison["reward_version"]

    plt.figure(figsize=(11, 6))
    plt.scatter(comparison["mean_irrigation"] + comparison["mean_n"], comparison["mean_yield"])
    for _, row in comparison.iterrows():
        plt.annotate(row["reward_version"].split("_")[0], (row["mean_irrigation"] + row["mean_n"], row["mean_yield"]), fontsize=8)
    plt.xlabel("Mean water + nitrogen input")
    plt.ylabel("Mean grain yield")
    plt.tight_layout()
    plt.savefig(fig_root / "episode_profit_yield_vs_input.png", dpi=180)
    plt.close()

    plt.figure(figsize=(12, 6))
    plt.plot(x, comparison["mean_irrigation_saturation_ratio"], marker="o", label="irrigation ratio")
    plt.plot(x, comparison["mean_n_saturation_ratio"], marker="o", label="nitrogen ratio")
    plt.axhline(0.95, color="black", linestyle="--", linewidth=1)
    plt.xticks(rotation=45, ha="right")
    plt.ylabel("Saturation ratio")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_root / "episode_profit_saturation_ratio.png", dpi=180)
    plt.close()

    plt.figure(figsize=(11, 6))
    plt.scatter(comparison["input_reduction_vs_baseline"], comparison["yield_loss_vs_baseline"])
    for _, row in comparison.iterrows():
        plt.annotate(row["reward_version"].split("_")[0], (row["input_reduction_vs_baseline"], row["yield_loss_vs_baseline"]), fontsize=8)
    plt.axhline(0.15, color="black", linestyle="--", linewidth=1)
    plt.xlabel("Input reduction vs baseline")
    plt.ylabel("Yield loss vs baseline")
    plt.tight_layout()
    plt.savefig(fig_root / "episode_profit_yield_loss_vs_input_reduction.png", dpi=180)
    plt.close()

    plt.figure(figsize=(11, 6))
    plt.scatter(comparison["mean_yield"], comparison["profit_score"])
    for _, row in comparison.iterrows():
        plt.annotate(row["reward_version"].split("_")[0], (row["mean_yield"], row["profit_score"]), fontsize=8)
    plt.xlabel("Mean yield")
    plt.ylabel("Mean season profit score")
    plt.tight_layout()
    plt.savefig(fig_root / "episode_profit_profit_score_vs_yield.png", dpi=180)
    plt.close()

    plt.figure(figsize=(12, 6))
    plt.plot(x, comparison["mean_swfac"], marker="o", label="swfac")
    plt.plot(x, comparison["mean_nstres"], marker="o", label="nstres")
    plt.xticks(rotation=45, ha="right")
    plt.ylabel("Mean stress indicators")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_root / "episode_profit_swfac_nstres.png", dpi=180)
    plt.close()


def run_candidate(config: dict, reward_version: str, station: str, train_year: int, eval_years: list[int]) -> list[dict]:
    print(f"[episode-profit] train {station} {train_year} {reward_version}", flush=True)
    model, model_path, policy_name = train_profit_policy(config, reward_version, station, train_year)
    rows = []
    for eval_year in eval_years:
        print(f"[episode-profit] evaluate {station} {reward_version} eval {eval_year}", flush=True)
        rows.append(evaluate_profit_model(model, config, reward_version, station, train_year, int(eval_year), model_path, policy_name))
    return rows


def run_hla_pilot(config: dict) -> tuple[pd.DataFrame, pd.DataFrame, str | None]:
    all_rows: list[dict] = []
    for reward_version in config["reward_candidates"]["versions"]:
        try:
            rows = run_candidate(
                config,
                reward_version,
                config["training"]["hla_pilot_station"],
                int(config["training"]["hla_train_year"]),
                [int(y) for y in config["training"]["hla_eval_years"]],
            )
            all_rows.extend(rows)
            summary = append_summary(all_rows)
            comparison = compare_candidates(summary)
            plot_summary(comparison)
        except Exception as exc:
            error_row = {
                "station": config["training"]["hla_pilot_station"],
                "reward_version": reward_version,
                "reward_family": PROFIT_REWARD_CANDIDATES[reward_version].family,
                "train_year": int(config["training"]["hla_train_year"]),
                "eval_year": np.nan,
                "seed": int(config["seed"]),
                "season_irrigation_cap": float(config["cap"]["season_irrigation_cap"]),
                "season_n_cap": float(config["cap"]["season_n_cap"]),
                "run_status": "failed",
                "episode_completed": False,
                "error_message": f"{type(exc).__name__}: {exc}",
                "traceback": traceback.format_exc()[-2000:],
            }
            all_rows.append(error_row)
            append_summary(all_rows)
            raise
    summary = append_summary(all_rows)
    comparison = compare_candidates(summary)
    plot_summary(comparison)
    return summary, comparison, select_candidate(comparison)


def run_extension(config: dict, selected: str | None) -> pd.DataFrame:
    out = OUTPUT_ROOT / "evaluation" / "cross_site_episode_profit_reward_test.csv"
    if not selected:
        empty = pd.DataFrame([{"selected_reward": "", "run_status": "skipped", "notes": "No HLA candidate passed strict or relaxed criteria."}])
        empty.to_csv(out, index=False, encoding="utf-8-sig")
        return empty
    all_rows: list[dict] = []
    for station in ["SYA", "LCA"]:
        train_year = int(config["training"]["best_train_years"][station])
        eval_years = [int(item["year"]) for item in config["observed_years"][station]]
        all_rows.extend(run_candidate(config, selected, station, train_year, eval_years))
    cross = pd.DataFrame(all_rows)
    cross.to_csv(out, index=False, encoding="utf-8-sig")
    return cross


def write_report(comparison: pd.DataFrame, cross: pd.DataFrame, selected: str | None) -> None:
    report_table = comparison.copy()
    for col in report_table.select_dtypes(include=["float", "int"]).columns:
        report_table[col] = report_table[col].round(4)
    cross_table = cross.copy()
    if len(cross_table):
        for col in cross_table.select_dtypes(include=["float", "int"]).columns:
            cross_table[col] = cross_table[col].round(4)
    tested = report_table["reward_version"].tolist()
    unsaturated = report_table[
        (report_table["mean_irrigation_saturation_ratio"] < 0.95)
        & (report_table["mean_n_saturation_ratio"] < 0.95)
    ]["reward_version"].tolist()
    yield_ok = report_table[report_table["yield_loss_vs_baseline"] <= 0.15]["reward_version"].tolist()
    md = f"""# Episode-level profit reward debug report

Generated at: 2026-06-06

## Goal

This stage redesigns the reward as an episode-level profit objective under action-safe PPO. It does not modify the original site-packages reward, does not train PPO without action safety, and does not overwrite 006_03/006_04 results.

## Why step-wise cost tuning failed

006_03 and 006_04 both changed daily cost terms, but PPO still saturated the 300 mm irrigation / 450 kg ha-1 N cap. That suggests the previous reward signal did not provide a clear season-level trade-off between final yield and total input.

## Terminal reward feasibility

Terminal reward is feasible at the wrapper layer. The original reward callback lacks a done flag, so this stage uses `EpisodeProfitRewardWrapper` around the action-safe env. The wrapper reads final `grnwt` after the environment returns done and accumulates safe `amir`/`anfer` from `SafeActionWrapper`.

## Reward candidates tested

{', '.join(tested)}

## HLA comparison

{df_to_markdown(report_table)}

## Pass summary

- Candidates with both irrigation and N unsaturated: {', '.join(unsaturated) if unsaturated else 'None'}
- Candidates with yield loss <= 15%: {', '.join(yield_ok) if yield_ok else 'None'}
- Recommended candidate: `{selected or 'None'}`

## SYA/LCA extension

{df_to_markdown(cross_table) if len(cross_table) else 'Extension was skipped because HLA did not identify a strict or relaxed passing candidate.'}

## Interpretation

If all P candidates still hit the cap, the remaining problem is probably not just coefficient size. The next likely direction is action-design revision, such as scheduled discrete management events, lower-frequency decisions, or explicit season budget actions, before multi-seed or rainfall-scaling expansion.

## Next step

Proceed to multi-seed only if a candidate is unsaturated and has acceptable yield loss. Proceed to rainfall-scaling budget scenario only after the action/reward behavior is interpretable in the observed-year HLA pilot.
"""
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)
    DOC_MD.write_text(md, encoding="utf-8")
    shutil.copy2(DOC_MD, OUTPUT_ROOT / "reports" / DOC_MD.name)
    write_ppt(comparison, selected)


def write_ppt(comparison: pd.DataFrame, selected: str | None) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        (OUTPUT_ROOT / "reports" / "ppt_generation_skipped_in_container.txt").write_text(
            "python-pptx is not installed in the DSSAT container. Generate the PPT from Windows Python after training.",
            encoding="utf-8",
        )
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_title(slide, title: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True

    def add_body(slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.0), Inches(5.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(15)

    def picture_slide(title: str, file_name: str) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, title)
        path = OUTPUT_ROOT / "figures" / "summary" / file_name
        if path.exists():
            slide.shapes.add_picture(str(path), Inches(0.8), Inches(1.0), width=Inches(11.8))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "006_05 Episode-Level Profit Reward")
    add_body(slide, "目标：把 reward 重构为季节尺度 profit objective。\n约束：action-safe PPO；只做 HLA pilot；不改原始 reward；不覆盖 006_03/006_04。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Terminal Reward Interface")
    add_body(slide, "原始 reward callback 能读 _next_state 和 _history，但没有 done flag。\n本阶段在 gymnasium wrapper 层实现 terminal reward：最后一步读取 final grnwt，并扣除累计灌溉和施氮成本。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Reward Candidates")
    add_body(slide, "P0：当前 reward baseline。\nP1-P3：terminal-only profit，成本由弱到强。\nP4-P6：terminal profit + daily input cost。\n所有系数都是 normalized score，不是真实人民币。")

    picture_slide("Yield vs Input", "episode_profit_yield_vs_input.png")
    picture_slide("Saturation Ratio", "episode_profit_saturation_ratio.png")
    picture_slide("Yield Loss vs Input Reduction", "episode_profit_yield_loss_vs_input_reduction.png")
    picture_slide("Profit Score vs Yield", "episode_profit_profit_score_vs_yield.png")
    picture_slide("SWFAC and NSTRES", "episode_profit_swfac_nstres.png")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Conclusion")
    add_body(slide, f"推荐 candidate：{selected or 'None'}。\n如果仍然打满 cap，下一步优先改 action design 或决策频率，再考虑多 seed 或 rainfall-scaling scenario。")

    DOC_PPT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DOC_PPT)
    shutil.copy2(DOC_PPT, OUTPUT_ROOT / "reports" / DOC_PPT.name)


def main() -> None:
    ensure_dirs()
    config = load_yaml(CONFIG_PATH)
    write_interface_check()
    write_design_files(config)
    summary, comparison, selected = run_hla_pilot(config)
    cross = run_extension(config, selected)
    write_report(comparison, cross, selected)
    print("[episode-profit] done")
    print((OUTPUT_ROOT / "evaluation" / "episode_profit_reward_candidate_comparison.csv").relative_to(PROJECT_ROOT))
    print(f"[episode-profit] selected={selected}")


if __name__ == "__main__":
    main()
