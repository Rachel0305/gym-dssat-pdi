from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))

from ppo_safe_rendering import PROJECT_ROOT, load_yaml


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_ppo_action_safe_site_training.yaml"
REPORT_STEM = f"{date.today().isoformat()}_action_safe_site_level_ppo_training_report"


def fmt(value, digits: int = 2) -> str:
    try:
        return f"{float(value):.{digits}f}"
    except Exception:
        return ""


def load_outputs(config: dict) -> tuple[Path, Path, Path, Path, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    output_root = PROJECT_ROOT / config["paths"]["output_root"]
    plan_path = output_root / "configs" / "action_safe_site_training_plan.csv"
    summary_path = output_root / "evaluation" / "action_safe_site_ppo_evaluation_summary.csv"
    best_path = output_root / "strategy_selection" / "best_policy_by_site.csv"
    ranking_path = output_root / "strategy_selection" / "policy_ranking_by_site.csv"
    plan = pd.read_csv(plan_path) if plan_path.exists() else pd.DataFrame()
    summary = pd.read_csv(summary_path) if summary_path.exists() else pd.DataFrame()
    best = pd.read_csv(best_path) if best_path.exists() else pd.DataFrame()
    return plan_path, summary_path, best_path, ranking_path, plan, summary, best


def write_markdown(config: dict) -> Path:
    plan_path, summary_path, best_path, ranking_path, plan, summary, best = load_outputs(config)
    output_root = PROJECT_ROOT / config["paths"]["output_root"]
    docs_path = PROJECT_ROOT / "docs" / f"{REPORT_STEM}.md"
    lines: list[str] = []
    lines.append("# Action-safe site-level PPO training report")
    lines.append("")
    lines.append(f"Generated at: {date.today().isoformat()}")
    lines.append("")
    lines.append("## Scope")
    lines.append("")
    lines.append("- Only action-safe PPO was trained.")
    lines.append("- Only HLA, SYA, and LCA were included.")
    lines.append("- FQA/YCA, multi-seed training, reward modification, and my_data modification were not performed.")
    lines.append(f"- Timesteps per model: {config.get('training', {}).get('total_timesteps')}")
    lines.append(f"- Output root: `{output_root.relative_to(PROJECT_ROOT)}`")
    lines.append("")
    lines.append("## Quality gates")
    lines.append("")
    lines.append(f"- total_irrigation <= {config['quality_gates']['max_total_irrigation']} mm")
    lines.append(f"- total_n_fertilizer <= {config['quality_gates']['max_total_n_fertilizer']} kg/ha")
    lines.append("- run_status = ok; episode completed; daily CSV and required figures exist")
    lines.append("")
    lines.append("## Generated files")
    lines.append("")
    for path in [plan_path, summary_path, best_path, ranking_path]:
        lines.append(f"- `{path.relative_to(PROJECT_ROOT)}`")
    lines.append("")
    lines.append("## Model-level status")
    lines.append("")
    if summary.empty:
        lines.append("No evaluation summary was generated.")
    else:
        model_status = (
            summary.groupby(["station", "policy_name"], as_index=False)
            .agg(
                evaluations=("eval_year", "count"),
                passed=("quality_gate_pass", lambda s: int(s.astype(str).str.lower().eq("true").sum())),
                mean_yield=("final_grnwt", "mean"),
                mean_irrigation=("total_irrigation", "mean"),
                mean_n=("total_n_fertilizer", "mean"),
                mean_reward=("mean_reward", "mean"),
            )
            .sort_values(["station", "policy_name"])
        )
        lines.append("| station | policy | evaluations | passed | mean_yield | mean_irrigation | mean_n | mean_reward |")
        lines.append("|---|---:|---:|---:|---:|---:|---:|---:|")
        for row in model_status.itertuples():
            lines.append(
                f"| {row.station} | {row.policy_name} | {row.evaluations} | {row.passed} | {fmt(row.mean_yield)} | {fmt(row.mean_irrigation)} | {fmt(row.mean_n)} | {fmt(row.mean_reward)} |"
            )
    lines.append("")
    lines.append("## Best policy by site")
    lines.append("")
    if best.empty:
        lines.append("No best-policy table was generated. This usually means training stopped before any cross-year validation passed.")
    else:
        lines.append("| station | best_policy | validation_years | mean_yield | mean_irrigation | mean_n | stability_score |")
        lines.append("|---|---|---:|---:|---:|---:|---:|")
        for row in best.itertuples():
            lines.append(
                f"| {row.station} | {row.best_policy_name} | {row.validation_years} | {fmt(row.mean_yield)} | {fmt(row.mean_irrigation)} | {fmt(row.mean_n_fertilizer)} | {fmt(row.stability_score, 3)} |"
            )
    lines.append("")
    lines.append("## Interpretation")
    lines.append("")
    lines.append("This stage is a controlled small-batch verification of the observed-year leave-one workflow. The action safety wrapper is a training constraint, not a final paper reward parameter. A policy should only move to longer training, FQA/YCA two-year validation, or multi-seed stability analysis after all quality gates are satisfied.")
    docs_path.write_text("\n".join(lines), encoding="utf-8")
    report_copy = output_root / "reports" / docs_path.name
    report_copy.parent.mkdir(parents=True, exist_ok=True)
    report_copy.write_text(docs_path.read_text(encoding="utf-8"), encoding="utf-8")
    return docs_path


def write_pptx(config: dict, markdown_path: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    _, _, best_path, _, _, summary, best = load_outputs(config)
    ppt_path = PROJECT_ROOT / "docs" / f"{REPORT_STEM}.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_title(slide, title: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.2), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0)

    def add_text(slide, text: str, top: float = 1.0) -> None:
        box = slide.shapes.add_textbox(Inches(0.65), Inches(top), Inches(12.0), Inches(5.8))
        tf = box.text_frame
        tf.word_wrap = True
        for idx, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(17)
            p.font.color.rgb = RGBColor(0, 0, 0)

    def add_table(slide, df: pd.DataFrame, left: float, top: float, width: float, height: float) -> None:
        if df.empty:
            add_text(slide, "No rows available.", top)
            return
        rows = min(len(df), 8) + 1
        cols = len(df.columns)
        table = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height)).table
        for c, col in enumerate(df.columns):
            cell = table.cell(0, c)
            cell.text = str(col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(79, 129, 189)
        for r, (_, row) in enumerate(df.head(8).iterrows(), start=1):
            for c, col in enumerate(df.columns):
                cell = table.cell(r, c)
                cell.text = str(row[col])
                if r % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(230, 236, 247)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment = PP_ALIGN.CENTER
                    for run in paragraph.runs:
                        run.font.name = "Microsoft YaHei"
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(0, 0, 0)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Action-safe site-level PPO training")
    add_text(
        slide,
        "本阶段只训练 HLA、SYA、LCA，并启用 action safety。\n不训练无 safety PPO，不训练 FQA/YCA，不修改 reward，不修改 my_data。\n每个模型训练前运行 null_zero 和 fixed_low_input smoke check，训练后评估训练年份和所有验证年份。",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Quality gates")
    add_text(
        slide,
        f"total_irrigation <= {config['quality_gates']['max_total_irrigation']} mm\n"
        f"total_n_fertilizer <= {config['quality_gates']['max_total_n_fertilizer']} kg/ha\n"
        "run_status 必须为 ok，episode 必须完成，daily CSV 和全部图必须存在。\n任一模型或任一 evaluation 不通过就停止后续站点。",
    )

    if not summary.empty:
        model_status = (
            summary.groupby(["station", "policy_name"], as_index=False)
            .agg(
                evals=("eval_year", "count"),
                passed=("quality_gate_pass", lambda s: int(s.astype(str).str.lower().eq("true").sum())),
                yield_mean=("final_grnwt", "mean"),
                irrig_mean=("total_irrigation", "mean"),
                n_mean=("total_n_fertilizer", "mean"),
            )
        )
        for col in ["yield_mean", "irrig_mean", "n_mean"]:
            model_status[col] = model_status[col].map(lambda v: fmt(v))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, "Model-level result summary")
        add_table(slide, model_status, 0.35, 1.0, 12.6, 5.8)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Best policy by site")
    best_show = best.copy()
    if not best_show.empty:
        cols = ["station", "best_policy_name", "validation_years", "mean_yield", "mean_irrigation", "mean_n_fertilizer", "stability_score"]
        best_show = best_show[[c for c in cols if c in best_show.columns]]
        for col in ["mean_yield", "mean_irrigation", "mean_n_fertilizer", "stability_score"]:
            if col in best_show:
                best_show[col] = best_show[col].map(lambda v: fmt(v, 3 if col == "stability_score" else 2))
    add_table(slide, best_show, 0.35, 1.0, 12.6, 4.8)
    add_text(slide, f"best_policy_by_site.csv: {best_path.relative_to(PROJECT_ROOT)}", top=6.0)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Next decision")
    add_text(
        slide,
        "如果 HLA、SYA、LCA 全部通过质量门，可以进入 FQA/YCA two-year cross validation 或多 seed 稳定性分析。\n"
        "如果某站点失败，应优先查看 quality_gate_reason、daily action、raw vs safe action 和 safety trigger 图，再决定是否调整 action safety 或训练步数。\n"
        f"Markdown report: {markdown_path.relative_to(PROJECT_ROOT)}",
    )

    prs.save(ppt_path)
    report_copy = PROJECT_ROOT / config["paths"]["output_root"] / "reports" / ppt_path.name
    report_copy.write_bytes(ppt_path.read_bytes())
    return ppt_path


def main() -> None:
    config = load_yaml(CONFIG_PATH)
    md = write_markdown(config)
    ppt = write_pptx(config, md)
    print(md.relative_to(PROJECT_ROOT))
    print(ppt.relative_to(PROJECT_ROOT))


if __name__ == "__main__":
    main()
