from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from constrained_ppo_wrappers import ImitationPenaltyRewardWrapper, PriorActionTable, ResidualPriorActionWrapper
from episode_profit_reward import EpisodeProfitRewardWrapper
from ppo_action_safety import ActionSafetyState, apply_action_safety, normalize_action, update_action_safety_state
from ppo_evaluate import latest_observation_dict, make_env, scalar
from ppo_experiment_plan import find_year
from ppo_plot_results import plot_episode
from ppo_safe_rendering import PROJECT_ROOT, load_yaml
from smoke_test_agents import POLICIES


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_multiseed_constrained_ppo_augmented_prior.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "constrained_ppo_multiseed_augmented_prior"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_constrained_ppo_multiseed_augmented_prior_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_constrained_ppo_multiseed_augmented_prior_report.pptx"


def ensure_dirs() -> None:
    for sub in ["configs", "prior_policies", "models", "logs", "tensorboard", "daily_outputs", "evaluation", "figures", "reports"]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)
    for station in ["HLA", "SYA", "LCA"]:
        (OUTPUT_ROOT / "models" / station).mkdir(parents=True, exist_ok=True)
        (OUTPUT_ROOT / "logs" / station).mkdir(parents=True, exist_ok=True)
        (OUTPUT_ROOT / "tensorboard" / station).mkdir(parents=True, exist_ok=True)
        (OUTPUT_ROOT / "daily_outputs" / station).mkdir(parents=True, exist_ok=True)
        (OUTPUT_ROOT / "figures" / station).mkdir(parents=True, exist_ok=True)


def df_to_markdown(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df.empty:
        return ""
    work = df.copy()
    if max_rows is not None:
        work = work.head(max_rows)
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(4)
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    separator = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, separator, *rows])


def profit_score(final_grnwt: float, irrigation: float, nitrogen: float, config: dict) -> float:
    econ = config.get("economics", {})
    return (
        float(econ.get("grain_value_coef", 0.01)) * float(final_grnwt)
        - float(econ.get("water_cost", 0.5)) * float(irrigation)
        - float(econ.get("n_cost", 0.25)) * float(nitrogen)
    )


def write_plan(config: dict) -> None:
    text = """# Multiseed prior and method plan

Generated at: 2026-06-06

## Prior choice

`BC_random_forest_regressor_augmented` is selected because 006_12 recommended it as the best augmented learned prior: it passed all HLA/SYA/LCA evaluations with mean irrigation 0.0, mean N about 81.83 kg/ha, yield loss vs PPO about 5.4%, and profit above the original BC prior.

`BC_mlp_regressor_augmented` is not used because it learned an almost zero-input policy with unacceptable yield loss. `BC_two_stage_classifier_regressor_augmented` is retained only as a backup because it passed the input gate but had lower yield/profit than the random forest prior.

Rainfall-scaling is still premature because 006_12 showed that irrigation labels remain sparse. This stage tests stability of the low-input learned prior before any rainfall stress scenario.

## Methods

- MS0_augmented_RF_prior_replay: pure prior replay, no PPO training.
- MS1_residual_augmented_RF_prior_strict: PPO learns a small residual around the RF prior under 100/200 season guardrails.
- MS2_guardrail_augmented_RF_100_200: PPO uses imitation penalty and strict 100/200 season guardrails.

Failure is defined as failed episodes, mean irrigation >100, mean N >200, yield loss vs PPO >15%, profit below 90% of the RF prior, or any regression toward 300/450 high-input behavior.
"""
    (OUTPUT_ROOT / "evaluation" / "multiseed_prior_and_method_plan.md").write_text(text, encoding="utf-8")


class SimDaySafeActionWrapper:
    def __init__(self, env, config: dict):
        self.env = env
        self.safety_config = {**config.get("action_safety", {}), "enabled": True}
        self.state = ActionSafetyState()
        self.step_count = 0
        self.last_safety_result = None
        self.action_space = env.action_space
        self.observation_space = env.observation_space
        self.metadata = getattr(env, "metadata", {})

    def reset(self, *args, **kwargs):
        self.state = ActionSafetyState()
        self.step_count = 0
        self.last_safety_result = None
        return self.env.reset(*args, **kwargs)

    def step(self, action):
        from ppo_action_safety import denormalize_action

        self.step_count += 1
        raw_real = denormalize_action(self.env.formator.action_names, self.env.formator.action_space_dict, action)
        result = apply_action_safety(raw_real, self.step_count, self.state, self.safety_config)
        safe_norm = normalize_action(self.env.formator.action_names, self.env.formator.action_space_dict, result.safe_real_action)
        obs, reward, terminated, truncated, info = self.env.step(safe_norm)
        update_action_safety_state(self.state, result.safe_real_action, self.step_count)
        self.last_safety_result = result
        return obs, reward, terminated, truncated, info

    def close(self):
        return self.env.close()

    @property
    def unwrapped(self):
        return self.env.unwrapped

    def __getattr__(self, name):
        return getattr(self.env, name)


def prior(config: dict) -> PriorActionTable:
    return PriorActionTable(PROJECT_ROOT / config["prior_policy"]["action_table"])


def base_env(config: dict, station: str, year: int, seed: int, run_tag: str, evaluation: bool):
    env = make_env(config, station, year, seed, run_tag=run_tag, evaluation=evaluation, action_safety_enabled=False)
    return SimDaySafeActionWrapper(env, config)


def wrapped_env(config: dict, station: str, year: int, seed: int, method: dict, run_tag: str, evaluation: bool):
    env = base_env(config, station, year, seed, run_tag, evaluation)
    method_type = method["type"]
    if method_type in {"residual", "imitation_penalty"}:
        env = EpisodeProfitRewardWrapper(env, method.get("reward_version", "P2_terminal_profit_medium_cost"))
    if method_type == "residual":
        env = ResidualPriorActionWrapper(
            env,
            prior(config),
            station,
            irrigation_range=float(method.get("irrigation_residual_range", 10.0)),
            n_range=float(method.get("n_residual_range", 20.0)),
        )
    elif method_type == "imitation_penalty":
        env = ImitationPenaltyRewardWrapper(env, prior(config), station, lambda_bc=float(method.get("lambda_bc", 1.0)))
    return env


def safety_result_from(env) -> Any:
    current = env
    for _ in range(8):
        result = getattr(current, "last_safety_result", None)
        if result is not None:
            return result
        current = getattr(current, "env", None)
        if current is None:
            break
    return None


def prior_info_from(env) -> dict[str, float]:
    info: dict[str, float] = {}
    for attr in ["last_bc_info", "last_residual_info"]:
        current = env
        for _ in range(8):
            value = getattr(current, attr, None)
            if value:
                info.update(value)
            current = getattr(current, "env", None)
            if current is None:
                break
    return info


def normalized_action_dict(env, normalized_action) -> dict[str, float]:
    arr = np.asarray(normalized_action).flatten()
    return {name: float(value) for name, value in zip(env.formator.action_names, arr)}


def evaluate_policy(config: dict, station: str, train_year: int, eval_year: int, seed: int, method: dict, model=None) -> dict[str, Any]:
    method_name = method["name"]
    method_type = method["type"]
    env = wrapped_env(config, station, eval_year, seed, method, f"{method_name}_{station}_seed{seed}_eval{eval_year}", evaluation=True)
    records: list[dict[str, Any]] = []
    prior_table = prior(config)
    cumulative_i = 0.0
    cumulative_n = 0.0
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        planting = pd.Timestamp(find_year(config, station, int(eval_year))["planting_date"])
        while not done and step_count < int(config.get("runtime", {}).get("max_steps", 260)):
            sim_day = step_count + 1
            if method_type == "prior_replay":
                real = prior_table.action(station, sim_day)
                action = normalize_action(env.formator.action_names, env.formator.action_space_dict, real)
            else:
                action, _ = model.predict(obs, deterministic=True)
                action = np.asarray(action, dtype=np.float32).flatten()
                action = np.clip(action, -1.0, 1.0)
            obs, reward, terminated, truncated, info = env.step(action)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            result = safety_result_from(env)
            pinfo = prior_info_from(env)
            if result is not None:
                real_i = float(result.safe_real_action.get("amir", 0.0))
                real_n = float(result.safe_real_action.get("anfer", 0.0))
                raw_i = float(result.raw_real_action.get("amir", 0.0))
                raw_n = float(result.raw_real_action.get("anfer", 0.0))
                safety_rule = result.safety_rule_triggered
            else:
                real_i = real_n = raw_i = raw_n = 0.0
                safety_rule = ""
            cumulative_i += real_i
            cumulative_n += real_n
            date = planting + pd.Timedelta(days=step_count)
            totir_raw = scalar(latest.get("totir"))
            tofer_raw = scalar(latest.get("tofer"))
            prior_action = prior_table.action(station, sim_day)
            bc_dev = ((real_i - prior_action["amir"]) / 100.0) ** 2 + ((real_n - prior_action["anfer"]) / 150.0) ** 2
            records.append(
                {
                    "station": station,
                    "method": method_name,
                    "seed": seed,
                    "prior_policy": config["prior_policy"]["name"],
                    "train_year": int(train_year),
                    "eval_year": int(eval_year),
                    "date": date.strftime("%Y-%m-%d"),
                    "year": int(date.year),
                    "doy": int(date.dayofyear),
                    "sim_day": sim_day,
                    "dap": scalar(latest.get("dap", sim_day)),
                    "topwt": scalar(latest.get("topwt")),
                    "grnwt": scalar(latest.get("grnwt")),
                    "xlai": scalar(latest.get("xlai")),
                    "totir": totir_raw if not np.isnan(totir_raw) else cumulative_i,
                    "tofer": tofer_raw if not np.isnan(tofer_raw) else cumulative_n,
                    "swfac": scalar(latest.get("swfac")),
                    "nstres": scalar(latest.get("nstres")),
                    "reward": float(reward),
                    "real_action_amir": real_i,
                    "real_action_anfer": real_n,
                    "raw_real_action_amir": raw_i,
                    "raw_real_action_anfer": raw_n,
                    "bc_prior_amir": float(prior_action["amir"]),
                    "bc_prior_anfer": float(prior_action["anfer"]),
                    "bc_action_deviation": float(bc_dev**0.5),
                    "safety_rule_triggered": safety_rule,
                    "done": done,
                    "info": json.dumps(info, ensure_ascii=False, default=str),
                    **pinfo,
                }
            )
            step_count += 1
    finally:
        try:
            env.close()
        except Exception:
            pass
    daily = pd.DataFrame(records)
    daily_dir = OUTPUT_ROOT / "daily_outputs" / station
    daily_dir.mkdir(parents=True, exist_ok=True)
    daily_csv = daily_dir / f"{station}_{method_name}_seed{seed}_eval{eval_year}_daily.csv"
    daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
    fig_dir = OUTPUT_ROOT / "figures" / station / method_name / f"seed_{seed}_eval_{eval_year}"
    plot_episode(daily, fig_dir)
    episode_completed = bool(records and records[-1]["done"])
    final_y = float(daily["grnwt"].iloc[-1]) if len(daily) else np.nan
    total_i = float(daily["real_action_amir"].sum()) if len(daily) else 0.0
    total_n = float(daily["real_action_anfer"].sum()) if len(daily) else 0.0
    return {
        "station": station,
        "method": method_name,
        "seed": seed,
        "prior_policy": config["prior_policy"]["name"],
        "train_year": int(train_year),
        "eval_year": int(eval_year),
        "run_status": "ok" if episode_completed else "failed",
        "episode_completed": episode_completed,
        "episode_length": int(len(daily)),
        "final_grnwt": final_y,
        "total_irrigation": total_i,
        "total_n_fertilizer": total_n,
        "profit_score": profit_score(final_y, total_i, total_n, config) if not np.isnan(final_y) else np.nan,
        "mean_swfac": float(pd.to_numeric(daily["swfac"], errors="coerce").mean()) if len(daily) else np.nan,
        "mean_nstres": float(pd.to_numeric(daily["nstres"], errors="coerce").mean()) if len(daily) else np.nan,
        "bc_action_deviation_mean": float(pd.to_numeric(daily["bc_action_deviation"], errors="coerce").mean()) if len(daily) else np.nan,
        "bc_action_deviation_max": float(pd.to_numeric(daily["bc_action_deviation"], errors="coerce").max()) if len(daily) else np.nan,
        "irrigation_saturation_ratio_100_200": total_i / 100.0,
        "n_saturation_ratio_100_200": total_n / 200.0,
        "irrigation_saturation_ratio_300_450": total_i / 300.0,
        "n_saturation_ratio_300_450": total_n / 450.0,
        "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
        "figure_dir": str(fig_dir.relative_to(PROJECT_ROOT)),
        "notes": "simday_safety_action_table_or_constrained_ppo",
    }


def run_smoke_episode(config: dict, station: str, year: int, seed: int, policy_name: str, method: str) -> dict[str, Any]:
    smoke_method = {"name": method, "type": "prior_replay"} if policy_name == "prior_replay" else {"name": method, "type": "fixed_policy"}
    env = base_env(config, station, year, seed, f"smoke_{policy_name}_{station}_{year}_{seed}_{method}", evaluation=True)
    records = []
    prior_table = prior(config)
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        while not done and step_count < int(config.get("runtime", {}).get("max_steps", 260)):
            sim_day = step_count + 1
            if policy_name == "prior_replay":
                real = prior_table.action(station, sim_day)
            else:
                latest = latest_observation_dict(env, obs, info)
                dap = int(round(scalar(latest.get("dap", sim_day))))
                real = POLICIES[policy_name].action_for_dap(dap, env.formator.action_names)
            action = normalize_action(env.formator.action_names, env.formator.action_space_dict, real)
            obs, reward, terminated, truncated, info = env.step(action)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            result = safety_result_from(env)
            safe = result.safe_real_action if result is not None else {"amir": 0.0, "anfer": 0.0}
            records.append(
                {
                    "station": station,
                    "train_year": int(year),
                    "seed": seed,
                    "method": method,
                    "policy_name": policy_name,
                    "sim_day": sim_day,
                    "dap": scalar(latest.get("dap", sim_day)),
                    "grnwt": scalar(latest.get("grnwt")),
                    "real_action_amir": float(safe.get("amir", 0.0)),
                    "real_action_anfer": float(safe.get("anfer", 0.0)),
                    "reward": float(reward),
                    "done": done,
                }
            )
            step_count += 1
    finally:
        try:
            env.close()
        except Exception:
            pass
    daily = pd.DataFrame(records)
    out_dir = OUTPUT_ROOT / "daily_outputs" / station / "smoke"
    out_dir.mkdir(parents=True, exist_ok=True)
    daily_csv = out_dir / f"{station}_{year}_{method}_seed{seed}_{policy_name}_smoke_daily.csv"
    daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
    ok = bool(records and records[-1]["done"])
    return {
        "station": station,
        "train_year": int(year),
        "seed": seed,
        "method": method,
        "policy_name": policy_name,
        "run_status": "ok" if ok else "failed",
        "episode_completed": ok,
        "episode_length": len(daily),
        "final_grnwt": float(daily["grnwt"].iloc[-1]) if len(daily) else np.nan,
        "total_irrigation": float(daily["real_action_amir"].sum()) if len(daily) else 0.0,
        "total_n_fertilizer": float(daily["real_action_anfer"].sum()) if len(daily) else 0.0,
        "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
    }


def run_pretrain_smoke(config: dict, station: str, train_year: int, seed: int, method_name: str) -> bool:
    rows = [run_smoke_episode(config, station, train_year, seed, policy, method_name) for policy in ["null_zero", "prior_replay", "fixed_low_input"]]
    out = OUTPUT_ROOT / "evaluation" / "pretrain_smoke_check_summary.csv"
    new = pd.DataFrame(rows)
    if out.exists():
        old = pd.read_csv(out)
        df = pd.concat([old, new], ignore_index=True)
        df = df.drop_duplicates(["station", "train_year", "seed", "method", "policy_name"], keep="last")
    else:
        df = new
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return all(row["run_status"] == "ok" and bool(row["episode_completed"]) for row in rows)


def ppo_kwargs(config: dict) -> dict:
    allowed = ["learning_rate", "gamma", "n_steps", "batch_size", "ent_coef", "clip_range"]
    return {key: value for key, value in config.get("ppo", {}).items() if key in allowed and value is not None}


def train_one(config: dict, station: str, train_year: int, seed: int, method: dict):
    from stable_baselines3 import PPO

    method_name = method["name"]
    env = wrapped_env(config, station, train_year, seed, method, f"{method_name}_{station}_seed{seed}_train", evaluation=False)
    model_dir = OUTPUT_ROOT / "models" / station
    model_path = model_dir / f"{station}_{method_name}_train{train_year}_seed{seed}"
    try:
        model = PPO(
            "MlpPolicy",
            env,
            verbose=0,
            seed=int(seed),
            tensorboard_log=str(OUTPUT_ROOT / "tensorboard" / station),
            **ppo_kwargs(config),
        )
        model.learn(total_timesteps=int(config["runtime"].get("timesteps", 5000)), progress_bar=False)
        model.save(str(model_path))
    finally:
        try:
            env.close()
        except Exception:
            pass
    return model, model_path.with_suffix(".zip")


def add_deltas(config: dict, summary: pd.DataFrame) -> pd.DataFrame:
    prior_summary = pd.read_csv(PROJECT_ROOT / config["prior_policy"]["baseline_summary"])
    prior_summary = prior_summary[prior_summary["policy_name"].eq(config["prior_policy"]["name"])]
    prior_lookup = prior_summary[["station", "eval_year", "final_grnwt", "total_irrigation", "total_n_fertilizer", "profit_score"]].rename(
        columns={
            "final_grnwt": "augmented_rf_prior_yield",
            "total_irrigation": "augmented_rf_prior_irrigation",
            "total_n_fertilizer": "augmented_rf_prior_n",
            "profit_score": "augmented_rf_prior_profit",
        }
    )
    out = summary.merge(prior_lookup, on=["station", "eval_year"], how="left")
    ref_y = float(config["ppo_reference"]["mean_yield"])
    ref_i = float(config["ppo_reference"]["total_irrigation"])
    ref_n = float(config["ppo_reference"]["total_n"])
    out["yield_loss_vs_augmented_rf_prior"] = (out["augmented_rf_prior_yield"] - out["final_grnwt"]) / out["augmented_rf_prior_yield"]
    out["yield_loss_vs_expert_best"] = np.nan
    out["yield_loss_vs_ppo_baseline"] = (ref_y - out["final_grnwt"]) / ref_y
    out["input_reduction_vs_ppo_baseline"] = 1.0 - ((out["total_irrigation"] + out["total_n_fertilizer"]) / max(1.0, ref_i + ref_n))
    out["cap_regression_300_450"] = (out["total_irrigation"] >= 250.0) | (out["total_n_fertilizer"] >= 400.0)
    out["passes_low_input_gate"] = (
        out["run_status"].eq("ok")
        & (out["total_irrigation"] <= 100.0)
        & (out["total_n_fertilizer"] <= 200.0)
        & (out["yield_loss_vs_ppo_baseline"] <= 0.15)
    )
    return out


def run_all(config: dict) -> None:
    ensure_dirs()
    shutil.copy2(CONFIG_PATH, OUTPUT_ROOT / "configs" / CONFIG_PATH.name)
    prior_src = PROJECT_ROOT / config["prior_policy"]["action_table"]
    shutil.copy2(prior_src, OUTPUT_ROOT / "prior_policies" / prior_src.name)
    write_plan(config)
    rows = []
    train_rows = []
    seeds = [int(seed) for seed in config["runtime"]["seeds"]]
    for method in config["methods"]:
        for station, site in config["sites"].items():
            train_year = int(site["train_year"])
            for seed in seeds:
                print(f"[00613] smoke {method['name']} {station} seed{seed}", flush=True)
                smoke_ok = run_pretrain_smoke(config, station, train_year, seed, method["name"])
                if not smoke_ok:
                    for eval_year in site["eval_years"]:
                        rows.append(
                            {
                                "station": station,
                                "method": method["name"],
                                "seed": seed,
                                "prior_policy": config["prior_policy"]["name"],
                                "train_year": train_year,
                                "eval_year": int(eval_year),
                                "run_status": "failed",
                                "episode_completed": False,
                                "notes": "pretrain_smoke_failed",
                            }
                        )
                    continue
                model = None
                model_path = ""
                if method.get("train", False):
                    print(f"[00613] train {method['name']} {station} seed{seed}", flush=True)
                    model, saved = train_one(config, station, train_year, seed, method)
                    model_path = str(saved.relative_to(PROJECT_ROOT))
                    train_rows.append(
                        {
                            "station": station,
                            "method": method["name"],
                            "seed": seed,
                            "train_year": train_year,
                            "timesteps": int(config["runtime"].get("timesteps", 5000)),
                            "model_path": model_path,
                        }
                    )
                for eval_year in site["eval_years"]:
                    print(f"[00613] eval {method['name']} {station} seed{seed} year{eval_year}", flush=True)
                    row = evaluate_policy(config, station, train_year, int(eval_year), seed, method, model=model)
                    row["model_path"] = model_path
                    rows.append(row)
    summary = add_deltas(config, pd.DataFrame(rows))
    summary.to_csv(OUTPUT_ROOT / "evaluation" / "multiseed_constrained_ppo_summary.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(train_rows).to_csv(OUTPUT_ROOT / "evaluation" / "multiseed_training_summary.csv", index=False, encoding="utf-8-sig")
    make_comparisons(config, summary)


def make_comparisons(config: dict, summary: pd.DataFrame | None = None) -> tuple[pd.DataFrame, pd.DataFrame]:
    if summary is None:
        summary = pd.read_csv(OUTPUT_ROOT / "evaluation" / "multiseed_constrained_ppo_summary.csv")
    ok = summary[summary["run_status"].eq("ok")].copy()
    comp = (
        summary.groupby(["station", "method"], as_index=False)
        .agg(
            eval_count=("eval_year", "count"),
            ok_count=("run_status", lambda s: int((s == "ok").sum())),
            mean_yield=("final_grnwt", "mean"),
            std_yield=("final_grnwt", "std"),
            mean_profit=("profit_score", "mean"),
            std_profit=("profit_score", "std"),
            mean_irrigation=("total_irrigation", "mean"),
            std_irrigation=("total_irrigation", "std"),
            mean_n=("total_n_fertilizer", "mean"),
            std_n=("total_n_fertilizer", "std"),
            mean_yield_loss_vs_ppo=("yield_loss_vs_ppo_baseline", "mean"),
            mean_input_reduction_vs_ppo=("input_reduction_vs_ppo_baseline", "mean"),
            failure_rate=("run_status", lambda s: float((s != "ok").mean())),
            cap_regression_rate=("cap_regression_300_450", "mean"),
        )
    )
    comp["cv_yield"] = comp["std_yield"] / comp["mean_yield"]
    comp.to_csv(OUTPUT_ROOT / "evaluation" / "multiseed_method_comparison.csv", index=False, encoding="utf-8-sig")

    policy_rows = []
    aug = pd.read_csv(PROJECT_ROOT / config["prior_policy"]["baseline_summary"])
    for name in ["best_expert_schedule_replay", "BC_random_forest_regressor_augmented", "BC_two_stage_classifier_regressor_original"]:
        part = aug[aug["policy_name"].eq(name)].copy()
        if not part.empty:
            policy_rows.append(
                {
                    "policy_type": "baseline_or_prior",
                    "policy_name": name,
                    "eval_count": len(part),
                    "mean_yield": float(part["final_grnwt"].mean()),
                    "mean_irrigation": float(part["total_irrigation"].mean()),
                    "mean_n": float(part["total_n_fertilizer"].mean()),
                    "mean_profit": float(part["profit_score"].mean()),
                }
            )
    for method, part in ok.groupby("method"):
        policy_rows.append(
            {
                "policy_type": "multiseed_constrained_ppo",
                "policy_name": method,
                "eval_count": len(part),
                "mean_yield": float(part["final_grnwt"].mean()),
                "mean_irrigation": float(part["total_irrigation"].mean()),
                "mean_n": float(part["total_n_fertilizer"].mean()),
                "mean_profit": float(part["profit_score"].mean()),
            }
        )
    policy_rows.append(
        {
            "policy_type": "ppo_cap_saturated_baseline",
            "policy_name": "old_ppo_cap_saturated_baseline",
            "eval_count": len(summary),
            "mean_yield": float(config["ppo_reference"]["mean_yield"]),
            "mean_irrigation": float(config["ppo_reference"]["total_irrigation"]),
            "mean_n": float(config["ppo_reference"]["total_n"]),
            "mean_profit": profit_score(float(config["ppo_reference"]["mean_yield"]), float(config["ppo_reference"]["total_irrigation"]), float(config["ppo_reference"]["total_n"]), config),
        }
    )
    policy_comp = pd.DataFrame(policy_rows)
    policy_comp.to_csv(OUTPUT_ROOT / "evaluation" / "policy_comparison_multiseed_augmented_prior.csv", index=False, encoding="utf-8-sig")
    recommend(config, comp, policy_comp)
    return comp, policy_comp


def recommend(config: dict, comp: pd.DataFrame, policy_comp: pd.DataFrame) -> pd.DataFrame:
    prior_profit = float(policy_comp[policy_comp["policy_name"].eq(config["prior_policy"]["name"])]["mean_profit"].iloc[0])
    comp = comp.copy()
    comp["site_passes_gate"] = (
        (comp["failure_rate"] == 0.0)
        & (comp["mean_irrigation"] <= 100.0)
        & (comp["mean_n"] <= 200.0)
        & (comp["mean_yield_loss_vs_ppo"] <= 0.15)
        & (comp["mean_profit"] >= 0.9 * prior_profit)
        & (comp["cap_regression_rate"] == 0.0)
    )
    method = (
        comp.groupby("method", as_index=False)
        .agg(
            sites=("station", "nunique"),
            site_pass_count=("site_passes_gate", "sum"),
            mean_yield=("mean_yield", "mean"),
            mean_profit=("mean_profit", "mean"),
            mean_irrigation=("mean_irrigation", "mean"),
            mean_n=("mean_n", "mean"),
            mean_yield_loss_vs_ppo=("mean_yield_loss_vs_ppo", "mean"),
            failure_rate=("failure_rate", "mean"),
            cap_regression_rate=("cap_regression_rate", "mean"),
            cv_yield=("cv_yield", "mean"),
        )
    )
    method["passes_gate"] = (
        (method["sites"] == 3)
        & (method["site_pass_count"] == method["sites"])
        & (method["mean_irrigation"] <= 100.0)
        & (method["mean_n"] <= 200.0)
        & (method["mean_yield_loss_vs_ppo"] <= 0.15)
        & (method["mean_profit"] >= 0.9 * prior_profit)
        & (method["failure_rate"] == 0.0)
        & (method["cap_regression_rate"] == 0.0)
    )
    train_methods = method[method["method"].ne("MS0_augmented_RF_prior_replay") & method["passes_gate"]].copy()
    if len(train_methods):
        best = train_methods.sort_values(["mean_profit", "mean_yield"], ascending=False).head(1).copy()
        best["recommendation"] = "use_finetuned_constrained_ppo_method"
        best["next_step"] = "consider_seed_3_4_extension_before_any_rainfall_scaling"
    else:
        best = method[method["method"].eq("MS0_augmented_RF_prior_replay")].copy()
        best["recommendation"] = "retain_augmented_rf_prior_as_current_best_profit_policy"
        best["next_step"] = "do_not_enter_rainfall_scaling; run_irrigation_responsive_expert_search_and_fix_HLA_yield_gap"
    best.to_csv(OUTPUT_ROOT / "evaluation" / "recommended_multiseed_policy.csv", index=False, encoding="utf-8-sig")
    return best


def make_figures() -> None:
    fig_dir = OUTPUT_ROOT / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    summary = pd.read_csv(OUTPUT_ROOT / "evaluation" / "multiseed_constrained_ppo_summary.csv")
    ok = summary[summary["run_status"].eq("ok")].copy()
    if ok.empty:
        return
    ok.boxplot(column="final_grnwt", by="method", figsize=(10, 5), rot=20)
    plt.suptitle("")
    plt.title("Yield by method")
    plt.tight_layout()
    plt.savefig(fig_dir / "multiseed_yield_boxplot_by_method.png", dpi=180)
    plt.close()

    ok.boxplot(column="profit_score", by="method", figsize=(10, 5), rot=20)
    plt.suptitle("")
    plt.title("Profit by method")
    plt.tight_layout()
    plt.savefig(fig_dir / "multiseed_profit_boxplot_by_method.png", dpi=180)
    plt.close()

    inp = ok.melt(id_vars=["method"], value_vars=["total_irrigation", "total_n_fertilizer"], var_name="input_type", value_name="amount")
    plt.figure(figsize=(10, 5))
    for i, (label, part) in enumerate(inp.groupby("input_type")):
        positions = np.arange(ok["method"].nunique()) + i * 0.25
        grouped = [part[part["method"].eq(m)]["amount"] for m in sorted(ok["method"].unique())]
        plt.boxplot(grouped, positions=positions, widths=0.2)
    plt.xticks(np.arange(ok["method"].nunique()) + 0.125, sorted(ok["method"].unique()), rotation=20, ha="right")
    plt.ylabel("Amount")
    plt.tight_layout()
    plt.savefig(fig_dir / "multiseed_input_boxplot_by_method.png", dpi=180)
    plt.close()

    plt.figure(figsize=(9, 5))
    for method, part in ok.groupby("method"):
        plt.scatter(part["total_irrigation"] + part["total_n_fertilizer"], part["final_grnwt"], label=method, alpha=0.75)
    plt.xlabel("Total input")
    plt.ylabel("Yield")
    plt.legend(fontsize=7)
    plt.tight_layout()
    plt.savefig(fig_dir / "multiseed_yield_vs_input.png", dpi=180)
    plt.close()

    comp = pd.read_csv(OUTPUT_ROOT / "evaluation" / "multiseed_method_comparison.csv")
    site_comp = comp.groupby("method", as_index=False)["cap_regression_rate"].mean()
    plt.figure(figsize=(9, 5))
    plt.barh(site_comp["method"], site_comp["cap_regression_rate"], color="#4F7ECF")
    plt.xlabel("Cap regression rate")
    plt.tight_layout()
    plt.savefig(fig_dir / "cap_regression_rate_by_method.png", dpi=180)
    plt.close()

    site = ok.groupby(["station", "method"], as_index=False)["profit_score"].mean()
    for station, part in site.groupby("station"):
        plt.figure(figsize=(9, 5))
        plt.barh(part["method"], part["profit_score"], color="#4F7ECF")
        plt.xlabel("Mean profit")
        plt.title(station)
        plt.tight_layout()
        plt.savefig(fig_dir / f"site_level_multiseed_policy_comparison_{station}.png", dpi=180)
        plt.close()
    shutil.copy2(fig_dir / "site_level_multiseed_policy_comparison_HLA.png", fig_dir / "site_level_multiseed_policy_comparison.png")

    sample = ok[(ok["method"].ne("MS0_augmented_RF_prior_replay"))].head(1)
    if len(sample):
        daily = pd.read_csv(PROJECT_ROOT / sample["daily_csv_path"].iloc[0])
        plt.figure(figsize=(10, 5))
        plt.plot(daily["sim_day"], daily["real_action_anfer"], label="finetuned N")
        plt.plot(daily["sim_day"], daily["bc_prior_anfer"], label="RF prior N", alpha=0.8)
        plt.plot(daily["sim_day"], daily["real_action_amir"], label="finetuned irrigation")
        plt.plot(daily["sim_day"], daily["bc_prior_amir"], label="RF prior irrigation", alpha=0.8)
        plt.xlabel("Sim day")
        plt.ylabel("Action")
        plt.legend()
        plt.tight_layout()
        plt.savefig(fig_dir / "prior_vs_finetuned_daily_actions.png", dpi=180)
        plt.close()


def write_report(config: dict) -> None:
    summary = pd.read_csv(OUTPUT_ROOT / "evaluation" / "multiseed_constrained_ppo_summary.csv")
    comp = pd.read_csv(OUTPUT_ROOT / "evaluation" / "multiseed_method_comparison.csv")
    policy = pd.read_csv(OUTPUT_ROOT / "evaluation" / "policy_comparison_multiseed_augmented_prior.csv")
    rec = pd.read_csv(OUTPUT_ROOT / "evaluation" / "recommended_multiseed_policy.csv")
    smoke = pd.read_csv(OUTPUT_ROOT / "evaluation" / "pretrain_smoke_check_summary.csv")
    method_overall = comp.groupby("method", as_index=False).agg(
        mean_yield=("mean_yield", "mean"),
        mean_profit=("mean_profit", "mean"),
        mean_irrigation=("mean_irrigation", "mean"),
        mean_n=("mean_n", "mean"),
        mean_yield_loss_vs_ppo=("mean_yield_loss_vs_ppo", "mean"),
        failure_rate=("failure_rate", "mean"),
        cap_regression_rate=("cap_regression_rate", "mean"),
    )
    text = f"""# Constrained PPO multiseed with augmented prior report

Generated at: 2026-06-06

## Scope

This stage used only `BC_random_forest_regressor_augmented` as the prior, only HLA/SYA/LCA, and seeds `{config['runtime']['seeds']}`. It did not train unrestricted PPO, did not train FQA/YCA, and did not enter rainfall-scaling.

## Pretrain smoke checks

{df_to_markdown(smoke.groupby(['method', 'policy_name'], as_index=False).agg(count=('run_status', 'count'), ok=('run_status', lambda s: int((s == 'ok').sum()))))}

## Multiseed method comparison

{df_to_markdown(method_overall)}

## Policy comparison

{df_to_markdown(policy)}

## Recommended policy

{df_to_markdown(rec)}

## Interpretation

If MS1 or MS2 does not clearly improve profit/yield over MS0 and the augmented RF prior, the safer recommendation is to keep the augmented RF prior. Rainfall-scaling remains premature when irrigation actions are still near zero; the next stage should search for irrigation-responsive expert schedules before stress-scenario RL.
"""
    DOC_MD.write_text(text, encoding="utf-8")
    shutil.copy2(DOC_MD, OUTPUT_ROOT / "reports" / DOC_MD.name)
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def title(slide, value: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.45))
        p = box.text_frame.paragraphs[0]
        p.text = value
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0)

    def bullets(slide, lines: list[str]) -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12), Inches(5.8))
        tf = box.text_frame
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(15)
            p.font.color.rgb = RGBColor(0, 0, 0)

    def table(slide, df: pd.DataFrame, max_rows: int = 8) -> None:
        data = df.head(max_rows).copy()
        shape = slide.shapes.add_table(len(data) + 1, len(data.columns), Inches(0.35), Inches(1.0), Inches(12.6), Inches(5.8))
        tbl = shape.table
        for j, col in enumerate(data.columns):
            cell = tbl.cell(0, j)
            cell.text = str(col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(79, 126, 207)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(8)
                p.font.bold = True
                p.font.color.rgb = RGBColor(255, 255, 255)
                p.alignment = PP_ALIGN.CENTER
        for i, (_, row) in enumerate(data.iterrows(), start=1):
            for j, col in enumerate(data.columns):
                value = row[col]
                if isinstance(value, float):
                    value = round(value, 4)
                cell = tbl.cell(i, j)
                cell.text = "" if pd.isna(value) else str(value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(235, 240, 250) if i % 2 == 0 else RGBColor(255, 255, 255)
                for p in cell.text_frame.paragraphs:
                    p.font.name = "Microsoft YaHei"
                    p.font.size = Pt(7)
                    p.font.color.rgb = RGBColor(0, 0, 0)

    slide = prs.slides.add_slide(blank)
    title(slide, "006_13 Multiseed constrained PPO")
    bullets(slide, ["Prior: BC_random_forest_regressor_augmented.", f"Seeds: {config['runtime']['seeds']}; sites: HLA, SYA, LCA.", "Methods: MS0 prior replay, MS1 residual, MS2 strict guardrail."])
    slide = prs.slides.add_slide(blank)
    title(slide, "Method comparison")
    table(slide, method_overall)
    slide = prs.slides.add_slide(blank)
    title(slide, "Policy comparison")
    table(slide, policy)
    slide = prs.slides.add_slide(blank)
    title(slide, "Recommendation")
    table(slide, rec)
    slide = prs.slides.add_slide(blank)
    title(slide, "Next stage")
    bullets(slide, ["Do not enter rainfall-scaling only because PPO is stable.", "If irrigation remains near zero, prioritize irrigation-responsive expert search.", "Extend seeds 3/4 only after reviewing this first batch."])
    prs.save(DOC_PPT)
    shutil.copy2(DOC_PPT, OUTPUT_ROOT / "reports" / DOC_PPT.name)


def report(config: dict) -> None:
    make_comparisons(config)
    make_figures()
    write_report(config)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run", action="store_true")
    parser.add_argument("--report", action="store_true")
    args = parser.parse_args()
    config = load_yaml(CONFIG_PATH)
    ensure_dirs()
    if args.run:
        run_all(config)
    if args.report:
        report(config)
    if not args.run and not args.report:
        run_all(config)
        report(config)


if __name__ == "__main__":
    main()
