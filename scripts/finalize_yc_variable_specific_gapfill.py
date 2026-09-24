#!/usr/bin/env python3
"""Apply YC weather gap-fill gates independently by variable and audit nearby rain stations."""

from __future__ import annotations

import calendar
import csv
import hashlib
import importlib
import io
import json
import math
import shutil
import sys
import subprocess
import time
import urllib.error
import urllib.request
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
import xlrd


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import finalize_yc_weather_and_prepare_cli as previous  # noqa: E402


OUT = ROOT / "results" / "yc_weather_gapfill_finalize"
DOCS = ROOT / "docs"
PREVIOUS = ROOT / "results" / "yc_weather_finalize_and_cli"
RAW = ROOT / "data" / "external" / "yc_chinaflux" / "raw"
STATIONS_URL = "https://www.ncei.noaa.gov/pub/data/ghcn/daily/ghcnd-stations.txt"
STATION_DAILY_URL = "https://www.ncei.noaa.gov/pub/data/ghcn/daily/all/{station}.dly"
SITE_LAT, SITE_LON = 36.830, 116.570
SEARCH_RADIUS_KM = 300.0
MIN_RAIN_PAIRED_DAYS = 3000  # Same minimum paired-day gate used by the prior RAIN validation.
RAIN_EVENT_THRESHOLD = 0.80
RATIO_MIN_COVERAGE = 0.90
TARGET_DATES = [date(2004, 10, day) for day in range(16, 21)]
TRAIN_START, TRAIN_END = date(2004, 1, 1), date(2013, 12, 31)
EXPECTED_DAYS = (TRAIN_END - TRAIN_START).days + 1
METADATA_YEAR_CUTOFF = 2013


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def write_csv(path: Path, fields: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: "" if row.get(field) is None else row.get(field) for field in fields})


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def backup_existing_task_outputs() -> str | None:
    existing = [p for p in OUT.rglob("*") if p.is_file() and "backups" not in p.relative_to(OUT).parts] if OUT.exists() else []
    existing += [p for p in (DOCS / "yc_weather_gapfill_finalize.md", DOCS / "yc_weather_gapfill_finalize.pptx") if p.exists()]
    if not existing:
        return None
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup = OUT / "backups" / stamp
    backup.mkdir(parents=True, exist_ok=False)
    for source in existing:
        if source.is_relative_to(OUT):
            destination = backup / "results" / source.relative_to(OUT)
        else:
            destination = backup / "docs" / source.name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
    return str(backup.relative_to(ROOT))


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def git_start_snapshot() -> dict[str, Any]:
    def git(*args: str) -> str:
        result = subprocess.run(["git", "-C", str(ROOT), *args], check=True, capture_output=True, text=True, encoding="utf-8")
        return result.stdout.strip()
    return {
        "branch": git("branch", "--show-current"),
        "head": git("rev-parse", "HEAD"),
        "tracked_worktree_status": git("status", "--short", "--untracked-files=no").splitlines(),
        "preexisting_untracked_files_observed": True,
        "note": "Only task-specific paths are eligible for staging; pre-existing tracked and untracked files are excluded.",
    }


def variable_gate(validation: dict[str, Any]) -> dict[str, Any]:
    limits = validation["acceptance_limits_declared_in_code"]
    metrics = validation["metrics"]
    checks = validation["gate_checks"]
    if limits.get("rain_event_agreement_minimum") != RAIN_EVENT_THRESHOLD:
        raise RuntimeError("Previous RAIN threshold differs from the task's fixed 80% gate")
    gate_rules = {
        "TMAX": ["TMAX_coverage", "TMAX_corrected_MAE", "TMAX_pearson_r"],
        "TMIN": ["TMIN_coverage", "TMIN_corrected_MAE", "TMIN_pearson_r"],
        "SRAD": ["SRAD_coverage", "SRAD_corrected_MAE", "SRAD_pearson_r"],
    }
    result: dict[str, Any] = {
        "source": str((PREVIOUS / "external_gapfill_validation.json").relative_to(ROOT)),
        "validation_overlap_years": validation["overlap_years"],
        "thresholds_redefined": False,
        "rain_event_agreement_threshold": RAIN_EVENT_THRESHOLD,
        "variables": {},
    }
    for variable, check_names in gate_rules.items():
        block = metrics["temperature_srad"][variable]
        corrected = block["corrected"]
        variable_checks = {key: bool(checks.get(key)) for key in check_names}
        passed = all(variable_checks.values())
        result["variables"][variable] = {
            "gate": "PASS" if passed else "FAIL",
            "checks": variable_checks,
            "overlap_n": corrected["n"],
            "raw_bias_external_minus_official": block["raw"]["bias_external_minus_official"],
            "corrected_mae": corrected["mae"],
            "corrected_rmse": corrected["rmse"],
            "pearson_r": corrected["pearson_r"],
            "bias_correction_official_minus_external": block["bias_correction_parameter_official_minus_external"],
            "correction_formula": validation["bias_correction_formula"],
            "limits": {key: value for key, value in limits.items() if key.startswith(variable + "_") or (variable in ("TMAX", "TMIN") and key == "temperature_minimum_pearson_r")},
        }
    rain = metrics["rain"]
    rain_checks = {key: bool(value) for key, value in checks.items() if key.startswith("RAIN_")}
    event_rate = rain["rain_event_agreement_rate"]
    rain_pass = event_rate >= RAIN_EVENT_THRESHOLD and all(rain_checks.values())
    result["variables"]["RAIN"] = {
        "gate": "PASS" if rain_pass else "BLOCKED",
        "checks": rain_checks,
        "overlap_n": rain["n_paired_days"],
        "event_agreement_days": rain["rain_event_agreement_days"],
        "event_agreement_rate": event_rate,
        "event_agreement_threshold": RAIN_EVENT_THRESHOLD,
        "monthly_ratio_summary": rain["monthly_ratio_summary"],
        "annual_ratio_summary": rain["annual_ratio_summary"],
        "nasa_rain_allowed": False,
        "reason": "降雨事件一致率未达本项目预设80%门槛；NASA RAIN 不用于正式填补。",
    }
    result["reproduction_passed"] = all(result["variables"][v]["gate"] == "PASS" for v in ("TMAX", "TMIN", "SRAD")) and result["variables"]["RAIN"]["gate"] == "BLOCKED"
    return result


def accepted_nasa_rows(validation: dict[str, Any], gates: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[tuple[date, str], float]]:
    proposals = read_csv(PREVIOUS / "external_gapfill_values.csv")
    gaps = read_csv(PREVIOUS / "gaps_before_external_fill.csv")
    gaps_by_key = {(date.fromisoformat(row["date"]), row["variable"]): row for row in gaps}
    accepted: list[dict[str, Any]] = []
    values: dict[tuple[date, str], float] = {}
    for row in proposals:
        variable = row["variable"]
        if variable not in ("TMAX", "TMIN", "SRAD"):
            continue
        if gates["variables"][variable]["gate"] != "PASS":
            continue
        dt = date.fromisoformat(row["date"])
        if not (date(2005, 1, 1) <= dt <= TRAIN_END):
            raise RuntimeError(f"NASA non-rain gap-fill outside 2005-2013: {dt}")
        if row["external_source"] != "NASA_POWER_DAILY_API":
            raise RuntimeError("Unexpected external source in previous gap-fill table")
        original_gap = gaps_by_key.get((dt, variable))
        if original_gap is None:
            raise RuntimeError(f"Refusing to overwrite a non-gap official value: {dt} {variable}")
        original_number = previous.number(original_gap["original_value"])
        if original_number is not None and not (variable == "SRAD" and original_number < 0):
            raise RuntimeError(f"Refusing to overwrite a physically valid official value: {dt} {variable}")
        raw = float(row["external_raw_value"])
        correction = float(row["correction_parameter"])
        filled = float(row["filled_value"])
        expected = float(validation["metrics"]["temperature_srad"][variable]["bias_correction_parameter_official_minus_external"])
        if not math.isfinite(raw) or not math.isfinite(correction) or not math.isfinite(filled) or abs(correction - expected) > 1e-10 or abs((raw + correction) - filled) > 1e-8:
            raise RuntimeError(f"Previous bias-corrected value cannot be reproduced: {dt} {variable}")
        key = (dt, variable)
        if key in values:
            raise RuntimeError(f"Duplicate NASA gap-fill key: {key}")
        values[key] = filled
        accepted.append({
            "date": dt.isoformat(),
            "variable": variable,
            "official_original_value": original_gap["original_value"],
            "nasa_raw_value": raw,
            "bias_correction": correction,
            "accepted_value": filled,
            "source": "nasa_power_train_overlap_bias_corrected",
            "qc_status": "ACCEPTED_VARIABLE_GATE_PASS",
        })
    expected_counts = {"TMAX": 107, "TMIN": 107, "SRAD": 151}
    observed_counts = Counter(row["variable"] for row in accepted)
    if dict(observed_counts) != expected_counts:
        raise RuntimeError(f"NASA variable-specific fill counts differ from audited gaps: {dict(observed_counts)}")
    return accepted, values


def haversine_km(lat: float, lon: float) -> float:
    lat0 = math.radians(SITE_LAT)
    dlat = math.radians(lat) - lat0
    dlon = math.radians(lon) - math.radians(SITE_LON)
    a = math.sin(dlat / 2) ** 2 + math.cos(lat0) * math.cos(math.radians(lat)) * math.sin(dlon / 2) ** 2
    return 6371.0088 * 2 * math.asin(math.sqrt(a))


def fetch_nearby_stations() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    request = urllib.request.Request(STATIONS_URL, headers={"User-Agent": "YC-weather-gapfill-audit/1.0"})
    last_error: str | None = None
    for attempt in range(1, 4):
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                digest = hashlib.sha256()
                stations: list[dict[str, Any]] = []
                for raw_line in response:
                    digest.update(raw_line)
                    line = raw_line.decode("ascii", errors="replace").rstrip("\r\n")
                    try:
                        station_id = line[0:11].strip()
                        lat, lon = float(line[12:20]), float(line[21:30])
                        elevation = float(line[31:37])
                    except (ValueError, IndexError):
                        continue
                    distance = haversine_km(lat, lon)
                    if distance > SEARCH_RADIUS_KM:
                        continue
                    stations.append({
                        "station_id": station_id,
                        "station_name": line[41:71].strip(),
                        "latitude": lat,
                        "longitude": lon,
                        "elevation_m": None if elevation <= -999 else elevation,
                        "distance_km": round(distance, 3),
                    })
                stations.sort(key=lambda item: (item["distance_km"], item["station_id"]))
                return stations, {
                    "url": STATIONS_URL,
                    "retrieved_at_utc": utc_now(),
                    "last_modified": response.headers.get("Last-Modified"),
                    "etag": response.headers.get("ETag"),
                    "bytes_hashed": int(response.headers.get("Content-Length", "0") or 0),
                    "sha256": digest.hexdigest(),
                    "search_radius_km": SEARCH_RADIUS_KM,
                    "candidate_count": len(stations),
                }
        except Exception as exc:  # noqa: BLE001
            last_error = f"{type(exc).__name__}: {exc}"
            if attempt < 3:
                time.sleep(attempt)
    raise RuntimeError(f"Could not retrieve GHCN station inventory: {last_error}")


def fetch_station_prcp(station: dict[str, Any]) -> tuple[dict[date, dict[str, Any]], dict[str, Any]]:
    station_id = station["station_id"]
    url = STATION_DAILY_URL.format(station=station_id)
    request = urllib.request.Request(url, headers={"User-Agent": "YC-weather-gapfill-audit/1.0"})
    last_error: str | None = None
    for attempt in range(1, 4):
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                digest = hashlib.sha256()
                observations: dict[date, dict[str, Any]] = {}
                duplicates = 0
                bytes_read = 0
                for raw_line in response:
                    digest.update(raw_line)
                    bytes_read += len(raw_line)
                    year_bytes = raw_line[11:15]
                    try:
                        year = int(year_bytes)
                    except ValueError:
                        continue
                    if year > METADATA_YEAR_CUTOFF or year < 2004:
                        continue
                    if raw_line[17:21].strip() != b"PRCP":
                        continue
                    try:
                        month = int(raw_line[15:17])
                    except ValueError:
                        continue
                    if year == 2004 and month != 10:
                        continue
                    for day in range(1, calendar.monthrange(year, month)[1] + 1):
                        dt = date(year, month, day)
                        if year == 2004 and dt not in TARGET_DATES:
                            continue
                        offset = 21 + (day - 1) * 8
                        block = raw_line[offset:offset + 8]
                        if len(block) < 5:
                            continue
                        try:
                            value_tenths = int(block[0:5])
                        except ValueError:
                            value_tenths = -9999
                        mflag = block[5:6].decode("ascii", errors="replace").strip()
                        qflag = block[6:7].decode("ascii", errors="replace").strip()
                        sflag = block[7:8].decode("ascii", errors="replace").strip()
                        amount = value_tenths / 10.0 if value_tenths >= 0 else None
                        if value_tenths == -9999:
                            status = "MISSING_SENTINEL"
                        elif value_tenths < 0:
                            status = "NEGATIVE_OR_INVALID_VALUE"
                        elif qflag:
                            status = "QUALITY_FLAGGED"
                        else:
                            status = "VALID"
                        record = {
                            "station_id": station_id,
                            "date": dt.isoformat(),
                            "raw_prcp_tenths_mm": value_tenths if value_tenths != -9999 else None,
                            "station_raw_prcp_mm": amount,
                            "measurement_flag": mflag,
                            "quality_flag": qflag,
                            "source_flag": sflag,
                            "rain_event": bool(amount is not None and amount > 0) or mflag == "T",
                            "qc_status": status,
                        }
                        if dt in observations:
                            duplicates += 1
                        observations[dt] = record
                return observations, {
                    "station_id": station_id,
                    "url": url,
                    "retrieved_at_utc": utc_now(),
                    "last_modified": response.headers.get("Last-Modified"),
                    "etag": response.headers.get("ETag"),
                    "bytes_read": bytes_read,
                    "sha256": digest.hexdigest(),
                    "parsed_weather_periods": "2004-10-16..2004-10-20 and 2005-01-01..2013-12-31 only",
                    "post_2013_rows_hashed_but_values_not_parsed": True,
                    "duplicate_daily_records": duplicates,
                }
        except Exception as exc:  # noqa: BLE001
            last_error = f"{type(exc).__name__}: {exc}"
            if attempt < 3:
                time.sleep(attempt)
    return {}, {
        "station_id": station_id,
        "url": url,
        "retrieved_at_utc": utc_now(),
        "fetch_error": last_error,
        "parsed_weather_periods": "none",
    }


def pearson(x: list[float], y: list[float]) -> float | None:
    if len(x) < 2:
        return None
    mx, my = sum(x) / len(x), sum(y) / len(y)
    dx, dy = [a - mx for a in x], [b - my for b in y]
    den = math.sqrt(sum(a * a for a in dx) * sum(b * b for b in dy))
    return sum(a * b for a, b in zip(dx, dy)) / den if den else None


def summarize_ratios(periods: dict[str, list[tuple[float, float]]], period_kind: str) -> dict[str, Any]:
    ratios: list[float] = []
    details: list[dict[str, Any]] = []
    for period, pairs in sorted(periods.items()):
        if period_kind == "monthly":
            year, month = (int(part) for part in period.split("-"))
            expected_days = calendar.monthrange(year, month)[1]
        else:
            year = int(period)
            expected_days = 366 if calendar.isleap(year) else 365
        coverage = len(pairs) / expected_days
        complete_period = coverage >= RATIO_MIN_COVERAGE
        ref = sum(a for a, _ in pairs)
        station = sum(b for _, b in pairs)
        ratio = station / ref if ref > 0 else None
        details.append({
            "period": period,
            "paired_days": len(pairs),
            "expected_days": expected_days,
            "coverage_fraction": coverage,
            "complete_period": complete_period,
            "official_total_mm": ref,
            "station_total_mm": station,
            "station_to_official_ratio": ratio,
        })
        if complete_period and ratio is not None and math.isfinite(ratio):
            ratios.append(ratio)
    if not ratios:
        summary = {
            "n": 0, "median": None, "minimum": None, "maximum": None,
            "coverage_threshold": RATIO_MIN_COVERAGE,
            "periods_total": len(periods),
            "periods_meeting_coverage": sum(item["complete_period"] for item in details),
            "incomplete_periods_excluded": sum(not item["complete_period"] for item in details),
        }
    else:
        ordered = sorted(ratios)
        middle = len(ordered) // 2
        median = ordered[middle] if len(ordered) % 2 else (ordered[middle - 1] + ordered[middle]) / 2
        summary = {
            "n": len(ordered), "median": median, "minimum": min(ordered), "maximum": max(ordered),
            "coverage_threshold": RATIO_MIN_COVERAGE,
            "periods_total": len(periods),
            "periods_meeting_coverage": sum(item["complete_period"] for item in details),
            "incomplete_periods_excluded": sum(not item["complete_period"] for item in details),
        }
    return {"summary": summary, "periods": details}


def validate_station(station: dict[str, Any], observations: dict[date, dict[str, Any]], official_rain: dict[date, float], download: dict[str, Any]) -> dict[str, Any]:
    valid_overlap = {
        dt: row for dt, row in observations.items()
        if date(2005, 1, 1) <= dt <= TRAIN_END and row["qc_status"] == "VALID"
    }
    paired = [(dt, official_rain[dt], float(row["station_raw_prcp_mm"])) for dt, row in valid_overlap.items() if dt in official_rain]
    official_events = [a > 0 for _, a, _ in paired]
    station_rows = [observations[dt] for dt, _, _ in paired]
    station_events = [bool(row["rain_event"]) for row in station_rows]
    agreement_days = sum(a == b for a, b in zip(official_events, station_events))
    true_positive = sum(a and b for a, b in zip(official_events, station_events))
    false_positive = sum((not a) and b for a, b in zip(official_events, station_events))
    false_negative = sum(a and (not b) for a, b in zip(official_events, station_events))
    official_wet = [(a, b) for _, a, b in paired if a > 0]
    abs_errors = [abs(a - b) for _, a, b in paired]
    wet_abs_errors = [abs(a - b) for a, b in official_wet]
    monthly: dict[str, list[tuple[float, float]]] = defaultdict(list)
    annual: dict[str, list[tuple[float, float]]] = defaultdict(list)
    for dt, ref, site in paired:
        monthly[dt.strftime("%Y-%m")].append((ref, site))
        annual[str(dt.year)].append((ref, site))
    monthly_ratios = summarize_ratios(monthly, "monthly")
    annual_ratios = summarize_ratios(annual, "annual")
    target_records = {dt: observations.get(dt) for dt in TARGET_DATES}
    target_valid_count = sum(row is not None and row["qc_status"] == "VALID" for row in target_records.values())
    station_qc_flags = Counter(row["quality_flag"] for row in observations.values() if row["quality_flag"])
    overlap_source_flags = Counter(row["source_flag"] or "BLANK" for row in station_rows)
    overlap_measurement_flags = Counter(row["measurement_flag"] or "BLANK" for row in station_rows)
    event_rate = agreement_days / len(paired) if paired else None
    enough_overlap = len(paired) >= MIN_RAIN_PAIRED_DAYS
    event_pass = event_rate is not None and event_rate >= RAIN_EVENT_THRESHOLD
    target_complete = target_valid_count == len(TARGET_DATES)
    no_duplicate_days = int(download.get("duplicate_daily_records", 0)) == 0
    passed = bool(enough_overlap and event_pass and target_complete and no_duplicate_days and not download.get("fetch_error"))
    reasons = []
    if not enough_overlap:
        reasons.append(f"有效重叠日不足{MIN_RAIN_PAIRED_DAYS}")
    if not event_pass:
        reasons.append("rain/no-rain事件一致率低于80%")
    if not target_complete:
        reasons.append(f"2004目标五天有效PRCP仅{target_valid_count}/5")
    if not no_duplicate_days:
        reasons.append("GHCN日记录存在重复")
    if download.get("fetch_error"):
        reasons.append("站点文件读取失败")
    return {
        **station,
        "prcp_valid_days_2004_target": target_valid_count,
        "prcp_valid_days_2005_2013": len(valid_overlap),
        "prcp_valid_years_2005_2013": sorted({dt.year for dt in valid_overlap}),
        "paired_days": len(paired),
        "event_agreement_days": agreement_days,
        "event_agreement_rate": event_rate,
        "wet_day_precision": true_positive / (true_positive + false_positive) if true_positive + false_positive else None,
        "wet_day_recall": true_positive / (true_positive + false_negative) if true_positive + false_negative else None,
        "daily_precipitation_mae_mm": sum(abs_errors) / len(abs_errors) if abs_errors else None,
        "wet_day_amount_mae_mm": sum(wet_abs_errors) / len(wet_abs_errors) if wet_abs_errors else None,
        "pearson_r": pearson([a for _, a, _ in paired], [b for _, _, b in paired]),
        "monthly_total_ratio_summary": monthly_ratios["summary"],
        "annual_total_ratio_summary": annual_ratios["summary"],
        "monthly_total_ratios": monthly_ratios["periods"],
        "annual_total_ratios": annual_ratios["periods"],
        "overlap_source_flag_counts": dict(overlap_source_flags),
        "overlap_measurement_flag_counts": dict(overlap_measurement_flags),
        "nonblank_quality_flag_counts": dict(station_qc_flags),
        "duplicate_daily_records": download.get("duplicate_daily_records", 0),
        "raw_file_sha256": download.get("sha256"),
        "data_file_url": download.get("url"),
        "source_last_modified": download.get("last_modified"),
        "source_flag_S_caution_present": bool(overlap_source_flags.get("S", 0)),
        "source_flag_caveat": "NOAA说明S来源由全球同步报文汇总，日降水需谨慎解释；保留来源标志供审查。",
        "unit_handling": "GHCN-Daily PRCP stored in tenths of mm; divided by 10; trace MFLAG=T counts as a wet event with reported amount retained.",
        "processing_integrity_passed": no_duplicate_days and not download.get("fetch_error"),
        "gate_status": "PASS" if passed else "FAIL",
        "failure_reasons": reasons,
    }


def load_station_cache(official_rain: dict[date, float]) -> tuple[list[dict[str, Any]], dict[str, Any], dict[str, dict[date, dict[str, Any]]], list[dict[str, Any]], list[dict[str, Any]]] | None:
    station_csv = OUT / "rain_station_candidates.csv"
    values_csv = OUT / "rain_station_source_values.csv"
    manifest_path = OUT / "run_manifest.json"
    if not (station_csv.exists() and values_csv.exists() and manifest_path.exists()):
        return None
    manifest = load_json(manifest_path)
    cached_sha = manifest.get("scoped_observation_csv_sha256")
    if cached_sha and cached_sha != sha256_file(values_csv):
        return None
    stations_meta = manifest.get("ghcn_station_inventory", {})
    downloads = manifest.get("ghcn_station_files", [])
    stations_rows = read_csv(station_csv)
    source_rows = read_csv(values_csv)
    if not stations_rows or len(downloads) != len(stations_rows) or stations_meta.get("candidate_count") != len(stations_rows):
        return None
    stations = []
    for row in stations_rows:
        stations.append({
            "station_id": row["station_id"],
            "station_name": row["station_name"],
            "latitude": float(row["latitude"]),
            "longitude": float(row["longitude"]),
            "elevation_m": float(row["elevation_m"]) if row.get("elevation_m") else None,
            "distance_km": float(row["distance_km"]),
        })
    observations_by_station: dict[str, dict[date, dict[str, Any]]] = {row["station_id"]: {} for row in stations}
    for row in source_rows:
        dt = date.fromisoformat(row["date"])
        observations_by_station[row["station_id"]][dt] = {
            "station_id": row["station_id"],
            "date": row["date"],
            "raw_prcp_tenths_mm": int(row["raw_prcp_tenths_mm"]) if row.get("raw_prcp_tenths_mm") else None,
            "station_raw_prcp_mm": float(row["station_raw_prcp_mm"]) if row.get("station_raw_prcp_mm") else None,
            "measurement_flag": row.get("measurement_flag", ""),
            "quality_flag": row.get("quality_flag", ""),
            "source_flag": row.get("source_flag", ""),
            "rain_event": row.get("rain_event", "").lower() == "true",
            "qc_status": row["qc_status"],
        }
    download_by_id = {row["station_id"]: row for row in downloads}
    validations = [validate_station(station, observations_by_station[station["station_id"]], official_rain, download_by_id[station["station_id"]]) for station in stations]
    inventory_manifest = dict(stations_meta)
    inventory_manifest["reused_local_scoped_observation_cache"] = True
    return stations, inventory_manifest, observations_by_station, downloads, validations


def local_daily_product_audit(path: Path) -> dict[str, Any]:
    target_values: dict[date, Any] = {}
    valid_count = 0
    total_days = 0
    digest = sha256_file(path)
    with zipfile.ZipFile(path, metadata_encoding="gbk") as archive:
        member = next((item for item in archive.namelist() if "/2004" in item and item.lower().endswith(".xlsx")), None)
        if member is None:
            raise RuntimeError("2004 workbook not found in ChinaFLUX daily archive")
        with archive.open(member) as binary:
            workbook = openpyxl.load_workbook(binary, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                dates = sheet.iter_rows(min_row=3, min_col=1, max_col=3, values_only=True)
                rains = sheet.iter_rows(min_row=3, min_col=25, max_col=25, values_only=True)
                for date_cells, (rain_cell,) in zip(dates, rains):
                    try:
                        dt = date(int(date_cells[0]), int(date_cells[1]), int(date_cells[2]))
                    except (TypeError, ValueError):
                        continue
                    total_days += 1
                    value = previous.number(rain_cell)
                    if value is not None and value >= 0:
                        valid_count += 1
                    if dt in TARGET_DATES:
                        target_values[dt] = rain_cell
            finally:
                workbook.close()
    targets = []
    for dt in TARGET_DATES:
        raw = target_values.get(dt)
        targets.append({"date": dt.isoformat(), "raw_rain_value": raw, "usable_daily_value": previous.number(raw) is not None and previous.number(raw) >= 0 if raw is not None else False})
    return {
        "source": str(path.relative_to(ROOT)),
        "sha256": digest,
        "member": member,
        "year": 2004,
        "daily_rows": total_days,
        "valid_daily_rain_days": valid_count,
        "target_days": targets,
        "target_complete": all(row["usable_daily_value"] for row in targets),
        "decision": "本地逐日产品五天均为缺失码，不能用于恢复；不拆分月/年总量。",
    }


def chinaflux_2004_aggregate_rain(kind: str) -> dict[str, Any]:
    if kind == "monthly":
        archive_path = RAW / "YCA_M_monthly.zip"
        grain_label = "月降水量"
    elif kind == "annual":
        archive_path = RAW / "YCA_M_yearly.zip"
        grain_label = "年降水量"
    else:
        raise ValueError(kind)
    result: dict[str, Any] = {"source": str(archive_path.relative_to(ROOT)), "sha256": sha256_file(archive_path), "year_read": 2004}
    with zipfile.ZipFile(archive_path, metadata_encoding="gbk") as archive:
        member = next((item for item in archive.namelist() if "/2004" in item and item.lower().endswith(".xlsx")), None)
        if member is None:
            raise RuntimeError(f"2004 {kind} ChinaFLUX workbook missing")
        with archive.open(member) as binary:
            workbook = openpyxl.load_workbook(binary, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                headers = [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
                column = next((i for i, value in enumerate(headers) if value == grain_label), None)
                if column is None:
                    raise RuntimeError(f"No {grain_label} column in {member}")
                values = []
                for row in sheet.iter_rows(min_row=3, values_only=True):
                    if not row or row[0] is None or int(row[0]) != 2004:
                        continue
                    amount = previous.number(row[column])
                    if kind == "monthly":
                        values.append({"month": int(row[1]), "precipitation_mm": amount, "valid": amount is not None and amount >= 0})
                    else:
                        values.append({"precipitation_mm": amount, "valid": amount is not None and amount >= 0})
            finally:
                workbook.close()
    result.update({"member": member, "aggregation": kind, "values": values})
    if kind == "annual":
        result["annual_precipitation_mm"] = next((row["precipitation_mm"] for row in values if row["valid"]), None)
    return result


def legacy_local_sources_audit(path: Path) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    with zipfile.ZipFile(path, metadata_encoding="gbk") as archive:
        for member in archive.namelist():
            if not member.lower().endswith(".xls") or not ("5-" in member or "9-" in member):
                continue
            book = xlrd.open_workbook(file_contents=archive.read(member), on_demand=True)
            try:
                sheet = book.sheet_by_index(0)
                title = str(sheet.cell_value(0, 0)) if sheet.nrows else ""
                years = sorted({int(sheet.cell_value(r, 0)) for r in range(2, sheet.nrows) if sheet.ncols and isinstance(sheet.cell_value(r, 0), (int, float))})
                results.append({
                    "archive_member": member,
                    "title": title,
                    "rows": sheet.nrows,
                    "temporal_grain": "monthly aggregate",
                    "years_present": years,
                    "daily_values_for_2004_target": False,
                    "decision": "月统计不能还原指定逐日事件。",
                })
            finally:
                book.release_resources()
    return results


def nasa_rain_support() -> dict[date, float | None]:
    result: dict[date, float | None] = {}
    for row in read_csv(PREVIOUS / "external_gapfill_values.csv"):
        if row["variable"] == "RAIN":
            result[date.fromisoformat(row["date"])] = float(row["filled_value"])
    return result


def write_station_evidence(stations: list[dict[str, Any]], observations_by_station: dict[str, dict[date, dict[str, Any]]], official_rain: dict[date, float], nasa_rain: dict[date, float | None], report_rows: list[dict[str, Any]]) -> None:
    write_csv(OUT / "rain_station_candidates.csv", [
        "station_id", "station_name", "latitude", "longitude", "elevation_m", "distance_km",
        "prcp_valid_days_2004_target", "prcp_valid_days_2005_2013", "prcp_valid_years_2005_2013",
        "paired_days", "event_agreement_days", "event_agreement_rate", "wet_day_precision", "wet_day_recall",
        "daily_precipitation_mae_mm", "wet_day_amount_mae_mm", "monthly_total_ratio_summary", "annual_total_ratio_summary",
        "pearson_r", "overlap_source_flag_counts", "overlap_measurement_flag_counts", "nonblank_quality_flag_counts",
        "source_flag_S_caution_present", "duplicate_daily_records", "gate_status", "failure_reasons", "raw_file_sha256",
    ], [{**row,
          "prcp_valid_years_2005_2013": json.dumps(row["prcp_valid_years_2005_2013"]),
          "monthly_total_ratio_summary": json.dumps(row["monthly_total_ratio_summary"], ensure_ascii=False),
          "annual_total_ratio_summary": json.dumps(row["annual_total_ratio_summary"], ensure_ascii=False),
          "overlap_source_flag_counts": json.dumps(row["overlap_source_flag_counts"], ensure_ascii=False),
          "overlap_measurement_flag_counts": json.dumps(row["overlap_measurement_flag_counts"], ensure_ascii=False),
          "nonblank_quality_flag_counts": json.dumps(row["nonblank_quality_flag_counts"], ensure_ascii=False),
          "failure_reasons": "；".join(row["failure_reasons"])} for row in report_rows])
    source_rows = []
    target_rows = []
    for station in stations:
        sid = station["station_id"]
        observations = observations_by_station.get(sid, {})
        for dt, row in sorted(observations.items()):
            source_rows.append({
                "station_id": sid, "date": dt.isoformat(), "raw_prcp_tenths_mm": row["raw_prcp_tenths_mm"],
                "station_raw_prcp_mm": row["station_raw_prcp_mm"], "measurement_flag": row["measurement_flag"],
                "quality_flag": row["quality_flag"], "source_flag": row["source_flag"], "rain_event": row["rain_event"], "qc_status": row["qc_status"],
            })
            if dt in TARGET_DATES:
                target_rows.append({
                    "station_id": sid, "station_name": station["station_name"], "distance_km": station["distance_km"],
                    "date": dt.isoformat(), "raw_prcp_tenths_mm": row["raw_prcp_tenths_mm"],
                    "station_raw_prcp_mm": row["station_raw_prcp_mm"], "measurement_flag": row["measurement_flag"],
                    "quality_flag": row["quality_flag"], "source_flag": row["source_flag"], "qc_status": row["qc_status"],
                    "nasa_power_prcp_supporting_only": nasa_rain.get(dt),
                })
        for dt in TARGET_DATES:
            if dt not in observations:
                target_rows.append({
                    "station_id": sid, "station_name": station["station_name"], "distance_km": station["distance_km"],
                    "date": dt.isoformat(), "raw_prcp_tenths_mm": None, "station_raw_prcp_mm": None,
                    "measurement_flag": None, "quality_flag": None, "source_flag": None, "qc_status": "NO_RECORD",
                    "nasa_power_prcp_supporting_only": nasa_rain.get(dt),
                })
    write_csv(OUT / "rain_station_source_values.csv", [
        "station_id", "date", "raw_prcp_tenths_mm", "station_raw_prcp_mm", "measurement_flag", "quality_flag", "source_flag", "rain_event", "qc_status",
    ], source_rows)
    write_csv(OUT / "rain_station_target_values.csv", [
        "station_id", "station_name", "distance_km", "date", "raw_prcp_tenths_mm", "station_raw_prcp_mm",
        "measurement_flag", "quality_flag", "source_flag", "qc_status", "nasa_power_prcp_supporting_only",
    ], target_rows)


def resolve_five_rain(stations: list[dict[str, Any]], observations_by_station: dict[str, dict[date, dict[str, Any]]], validations: list[dict[str, Any]], nasa: dict[date, float | None]) -> tuple[list[dict[str, Any]], dict[str, Any] | None]:
    by_id = {row["station_id"]: row for row in validations}
    eligible = [station for station in stations if by_id[station["station_id"]]["gate_status"] == "PASS"]
    eligible.sort(key=lambda st: (st["distance_km"], -(by_id[st["station_id"]]["event_agreement_rate"] or 0), st["station_id"]))
    chosen = eligible[0] if eligible else None
    resolution = []
    for dt in TARGET_DATES:
        selected = observations_by_station.get(chosen["station_id"], {}).get(dt) if chosen else None
        resolution.append({
            "date": dt.isoformat(),
            "selected_rain_mm": selected["station_raw_prcp_mm"] if selected and selected["qc_status"] == "VALID" else None,
            "selected_source": f"ghcn_daily_nearby_station_{chosen['station_id']}" if chosen else None,
            "station_id": chosen["station_id"] if chosen else None,
            "distance_km": chosen["distance_km"] if chosen else None,
            "station_raw_prcp": selected["station_raw_prcp_mm"] if selected else None,
            "nasa_power_prcp": nasa.get(dt),
            "qc_status": "ACCEPTED_GHCN_STATION_GATE_PASS" if chosen else "BLOCKED_NO_STATION_PASSED_GATE",
            "decision_reason": "按80%事件一致率门槛、足够重叠、五日完整性和距离顺序选择。" if chosen else "300 km范围内没有同时满足80%事件一致率、足够重叠及五日有效记录的站点；NASA值仅作辅助。",
        })
    return resolution, chosen


def correlation_and_candidate_rows(candidate: list[dict[str, Any]]) -> dict[str, Any]:
    dates = [date.fromisoformat(row["DATE"]) for row in candidate]
    errors = []
    if len(candidate) != EXPECTED_DAYS or len(set(dates)) != EXPECTED_DAYS:
        errors.append("date_count_or_duplicate_failure")
    if dates != [date.fromordinal(TRAIN_START.toordinal() + i) for i in range(EXPECTED_DAYS)]:
        errors.append("date_sequence_failure")
    sources_missing = 0
    physics_failures = []
    for row in candidate:
        values = {v: (float(row[v]) if row.get(v) not in (None, "") else None) for v in ("TMAX", "TMIN", "SRAD", "RAIN")}
        if any(value is None or not math.isfinite(float(value)) for value in values.values()):
            errors.append(f"missing_or_nonfinite:{row['DATE']}")
            continue
        if values["TMAX"] < values["TMIN"]:
            physics_failures.append(f"tmax_lt_tmin:{row['DATE']}")
        if values["SRAD"] < 0:
            physics_failures.append(f"negative_srad:{row['DATE']}")
        if values["RAIN"] < 0:
            physics_failures.append(f"negative_rain:{row['DATE']}")
        if any(not row.get(f"{v}_source") or row.get(f"{v}_source") == "unresolved" for v in values):
            sources_missing += 1
    if physics_failures:
        errors.extend(physics_failures)
    if sources_missing:
        errors.append(f"provenance_missing_rows:{sources_missing}")
    return {"passed": not errors, "not_run": False, "row_count": len(candidate), "expected_days": EXPECTED_DAYS, "unique_dates": len(set(dates)), "errors": errors, "physics_failure_count": len(physics_failures)}


def load_reusable_candidate(candidate_path: Path, accepted_values: dict[tuple[date, str], float], selected_station: dict[str, Any], selected_observations: dict[date, dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any], list[dict[str, Any]], list[dict[str, Any]]] | None:
    qc_path = OUT / "candidate_qc.json"
    if not candidate_path.exists() or not qc_path.exists():
        return None
    previous_qc = load_json(qc_path)
    if previous_qc.get("passed") is not True or previous_qc.get("row_count") != EXPECTED_DAYS:
        return None
    rows = read_csv(candidate_path)
    for row in rows:
        row["YEAR"], row["DOY"] = int(row["YEAR"]), int(row["DOY"])
        for variable in ("TMAX", "TMIN", "SRAD", "RAIN"):
            row[variable] = float(row[variable])
    by_date = {date.fromisoformat(row["DATE"]): row for row in rows}
    for (dt, variable), expected in accepted_values.items():
        row = by_date.get(dt)
        if row is None or abs(row[variable] - expected) > 1e-8 or row[f"{variable}_source"] != "nasa_power_train_overlap_bias_corrected":
            return None
    selected_id = selected_station["station_id"]
    for dt in TARGET_DATES:
        row, source = by_date.get(dt), selected_observations.get(dt)
        if row is None or source is None or source["qc_status"] != "VALID":
            return None
        if abs(row["RAIN"] - float(source["station_raw_prcp_mm"])) > 1e-8 or row["RAIN_source"] != f"ghcn_daily_nearby_station_{selected_id}":
            return None
    qc = correlation_and_candidate_rows(rows)
    if not qc["passed"]:
        return None
    annual, monthly = monthly_and_annual_summaries(rows)
    return rows, qc, annual, monthly


def monthly_and_annual_summaries(candidate: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    by_year: dict[int, list[dict[str, Any]]] = defaultdict(list)
    by_month: dict[tuple[int, int], list[dict[str, Any]]] = defaultdict(list)
    for row in candidate:
        dt = date.fromisoformat(row["DATE"])
        by_year[dt.year].append(row)
        by_month[(dt.year, dt.month)].append(row)

    def stats(group: list[dict[str, Any]], key: str) -> tuple[float, float, float]:
        values = [float(row[key]) for row in group]
        return sum(values) / len(values), min(values), max(values)

    def rain_stats(group: list[dict[str, Any]]) -> tuple[float, int, float, int]:
        rain = [float(row["RAIN"]) for row in group]
        wet = sum(x > 0 for x in rain)
        longest, current = 0, 0
        for value in rain:
            current = current + 1 if value <= 0 else 0
            longest = max(longest, current)
        return sum(rain), wet, max(rain), longest

    annual, monthly = [], []
    for year, group in sorted(by_year.items()):
        tmax, tmin, srad = stats(group, "TMAX"), stats(group, "TMIN"), stats(group, "SRAD")
        precip, wet, max_rain, dry = rain_stats(group)
        annual.append({"year": year, "days": len(group), "annual_precipitation_mm": precip, "rainy_days": wet, "max_daily_rain_mm": max_rain, "longest_dry_spell_days": dry, "TMAX_mean": tmax[0], "TMAX_min": tmax[1], "TMAX_max": tmax[2], "TMIN_mean": tmin[0], "TMIN_min": tmin[1], "TMIN_max": tmin[2], "SRAD_mean": srad[0], "SRAD_min": srad[1], "SRAD_max": srad[2]})
    for (year, month), group in sorted(by_month.items()):
        tmax, tmin, srad = stats(group, "TMAX"), stats(group, "TMIN"), stats(group, "SRAD")
        precip, wet, max_rain, dry = rain_stats(group)
        monthly.append({"year": year, "month": month, "days": len(group), "monthly_precipitation_mm": precip, "rainy_days": wet, "max_daily_rain_mm": max_rain, "longest_dry_spell_days": dry, "TMAX_mean": tmax[0], "TMAX_min": tmax[1], "TMAX_max": tmax[2], "TMIN_mean": tmin[0], "TMIN_min": tmin[1], "TMIN_max": tmin[2], "SRAD_mean": srad[0], "SRAD_min": srad[1], "SRAD_max": srad[2]})
    return annual, monthly


def backup_summary_csv(path: Path, reason: str) -> None:
    write_csv(path, ["status", "reason"], [{"status": "NOT_RUN_BLOCKED_FIVE_DAY_RAIN", "reason": reason}])


def report_text(summary: dict[str, Any], gates: dict[str, Any], local_audit: dict[str, Any], validations: list[dict[str, Any]]) -> str:
    gate = gates["variables"]
    val_lines = []
    for item in validations:
        def fmt(value: Any, spec: str = ".3f") -> str:
            return format(value, spec) if value is not None else "NA"
        val_lines.append(
            f"| {item['station_id']} | {item['station_name']} | {item['distance_km']:.1f} | {item['paired_days']} | "
            f"{fmt(item.get('event_agreement_rate'), '.2%')} | {fmt(item.get('wet_day_precision'), '.2%')} | {fmt(item.get('wet_day_recall'), '.2%')} | "
            f"{fmt(item.get('daily_precipitation_mae_mm'))} | {fmt(item.get('wet_day_amount_mae_mm'))} | {fmt(item.get('pearson_r'))} | "
            f"{fmt(item.get('monthly_total_ratio_summary', {}).get('median'))} (n={item.get('monthly_total_ratio_summary', {}).get('n', 0)}) | "
            f"{fmt(item.get('annual_total_ratio_summary', {}).get('median'))} (n={item.get('annual_total_ratio_summary', {}).get('n', 0)}) | "
            f"{item['prcp_valid_days_2004_target']}/5 | {item['gate_status']} |"
        )
    nasa_metrics = []
    for variable in ("TMAX", "TMIN", "SRAD"):
        item = gate[variable]
        nasa_metrics.append(f"| {variable} | {item['gate']} | {item['overlap_n']} | {item['corrected_mae']:.3f} | {item['corrected_rmse']:.3f} | {item['pearson_r']:.4f} | {summary['accepted_fill_counts'][variable]} |")
    resolution_lines = []
    for row in summary["five_day_resolution"]:
        value = "未定" if row["selected_rain_mm"] is None else f"{row['selected_rain_mm']:.1f} mm"
        resolution_lines.append(f"| {row['date']} | {value} | {row['station_id'] or '无'} | {row['qc_status']} |")
    candidate_qc = summary["candidate_qc"]
    annual_2004 = summary.get("annual_precipitation_2004_mm")
    annual_text = f"{annual_2004:.1f} mm" if annual_2004 is not None else "未定（五天无通过 Gate 的日值来源）"
    annual_reference = local_audit["chinaflux_annual_rain_2004"].get("annual_precipitation_mm")
    annual_reference_text = f"{annual_reference:.1f} mm" if annual_reference is not None else "无有效年值"
    comparisons = summary.get("chinaflux_2004_candidate_comparison", [])
    annual_comparison = next((row for row in comparisons if row["aggregation"] == "annual"), None)
    monthly_comparisons = [row for row in comparisons if row["aggregation"] == "monthly"]
    monthly_exact = sum(abs(float(row["difference_candidate_minus_chinaflux_mm"])) < 1e-6 for row in monthly_comparisons)
    comparison_text = (
        f"年值差 {float(annual_comparison['difference_candidate_minus_chinaflux_mm']):.1f} mm；"
        f"月值逐项相等 {monthly_exact}/{len(monthly_comparisons)} 月（只作一致性对照，不参与填补决策）。"
        if annual_comparison else "NOT_RUN（candidate 未形成）。"
    )
    if summary["candidate_created"]:
        candidate_qc_text = f"已生成 {summary['candidate_row_count']}/3653 日；日期、变量完整性、物理 Gate 和 provenance QC `{candidate_qc['passed']}`；年/月统计与 ChinaFLUX 聚合比较见对应 CSV。"
    else:
        candidate_qc_text = f"未生成；完整性和物理 QC 未运行（{candidate_qc.get('blocked_by', summary['final_status'])}）。年/月 candidate 统计未运行。"
    fill_ratio_lines = []
    for variable in ("TMAX", "TMIN", "SRAD", "RAIN"):
        count = summary["candidate_gapfill_counts"].get(variable, 0)
        fill_ratio_lines.append(f"| {variable} | {count} | {count / EXPECTED_DAYS:.2%} |")
    selected = summary.get("selected_station")
    if selected:
        selected_validation = next(row for row in validations if row["station_id"] == selected["station_id"])
        partial_2013 = next((row for row in selected_validation["annual_total_ratios"] if row["period"] == "2013"), None)
        complete_month_ratios = [
            row for row in selected_validation["monthly_total_ratios"]
            if row["complete_period"] and row["station_to_official_ratio"] is not None
        ]
        low_month = min(complete_month_ratios, key=lambda row: row["station_to_official_ratio"])
        high_month = max(complete_month_ratios, key=lambda row: row["station_to_official_ratio"])
        flag_counts = selected_validation["overlap_source_flag_counts"]
        target_flags = ",".join(summary.get("selected_target_source_flags", [])) or "NA"
        target_qflags = ",".join(summary.get("selected_target_quality_flags", [])) or "空白"
        partial_2013_text = (
            f"2013 年配对覆盖 {partial_2013['paired_days']}/{partial_2013['expected_days']} 日 "
            f"({partial_2013['coverage_fraction']:.1%})，已从完整年比值摘要中排除。"
            if partial_2013 else ""
        )
        station_detail = (
            f"选中 `{selected['station_id']}`（{selected['station_name']}），距 YC {selected['distance_km']:.2f} km。"
            f"五个目标日 NOAA source flag 为 `{target_flags}`，quality flag 为 `{target_qflags}`。"
            f"湿日 precision {selected_validation['wet_day_precision']:.2%}，recall {selected_validation['wet_day_recall']:.2%}，"
            f"日降水 MAE {selected_validation['daily_precipitation_mae_mm']:.3f} mm，官方湿日雨量 MAE {selected_validation['wet_day_amount_mae_mm']:.3f} mm，"
            f"Pearson r {selected_validation['pearson_r']:.4f}。月/年总量比只汇总覆盖率 ≥{RATIO_MIN_COVERAGE:.0%} 的期间："
            f"完整月 n={selected_validation['monthly_total_ratio_summary']['n']}，中位数/范围 "
            f"{selected_validation['monthly_total_ratio_summary']['median']:.3f} / {selected_validation['monthly_total_ratio_summary']['minimum']:.3f}–{selected_validation['monthly_total_ratio_summary']['maximum']:.3f}；"
            f"完整年 n={selected_validation['annual_total_ratio_summary']['n']}，中位数/范围 "
            f"{selected_validation['annual_total_ratio_summary']['median']:.3f} / {selected_validation['annual_total_ratio_summary']['minimum']:.3f}–{selected_validation['annual_total_ratio_summary']['maximum']:.3f}。"
            f"{partial_2013_text}湿日 precision {selected_validation['wet_day_precision']:.2%}，表明不匹配湿日仍不少；本站结果仅用于本次五个目标日的限定填补，不代表可无条件替代 YC 长期日降雨序列。"
            f"完整月份差异示例：{low_month['period']} 站/官方 {low_month['station_total_mm']:.1f}/{low_month['official_total_mm']:.1f} mm；"
            f"{high_month['period']} {high_month['station_total_mm']:.1f}/{high_month['official_total_mm']:.1f} mm。"
            f"配对来源标志计数：s={flag_counts.get('s', 0)}，S={flag_counts.get('S', 0)}。"
        )
    else:
        station_detail = "没有通过门槛的站点，因此没有选定来源。"
    return f"""# YC 变量级天气缺口补齐审计（003_05_01）

**最终状态：** `{summary['final_status']}`

**范围：** YC/YCA，训练天气 2004–2013；本轮未生成 `.CLI`，未运行 WeatherMan、WGEN、DSSAT 或 PPO。

## 1. 变量级 Gate

上一轮的 `accepted_for_gap_fill` 是全变量联合结果；RAIN 未通过导致所有变量一起被拒绝。本轮保持原阈值和上一轮校正参数，只对各变量自己的验证项独立判定，不修改生产数据。

| 变量 | Gate | overlap n | 校正后 MAE | RMSE | Pearson r | NASA 接受填补天数 |
|---|---:|---:|---:|---:|---:|---:|
{chr(10).join(nasa_metrics)}

RAIN overlap 为 {gate['RAIN']['overlap_n']} 天，事件一致 {gate['RAIN']['event_agreement_days']} 天（{gate['RAIN']['event_agreement_rate']:.2%}），固定门槛 80.00%；因此 NASA RAIN 继续拒绝。TMAX/TMIN/SRAD 沿用 2005–2013 官方值与 NASA 原值计算的加性偏差校正，不覆盖已有有效官方值。

## 2. 本地日降雨源检查

YCA ChinaFLUX 2004 日尺度产品共 {local_audit['daily_rows']} 日，其中 {local_audit['valid_daily_rain_days']} 日有有效降雨值；目标五天原值均为缺失码：

| 日期 | ChinaFLUX 日产品原值 | 可用 |
|---|---:|---:|
{chr(10).join(f"| {item['date']} | {item['raw_rain_value']} | 否 |" for item in local_audit['target_days'])}

项目内自动降水 XLS 与人工气象记录为月统计，不足以还原逐日降水；月/年值不拆分、不反推到五天。

## 3. GHCN-Daily 附近站验证

从 NOAA GHCN-Daily 站点目录按 YC 坐标 36.830°N, 116.570°E 搜索 300 km 内的 {len(validations)} 个站点。PRCP 原始单位为 0.1 mm，按 NOAA 格式除以 10；`-9999` 作为缺测码，非空质量标志不计入有效配对，原始 M/Q/S 标志保留。最低 paired days 采用上一轮 RAIN coverage gate 的 3000 天；事件门槛仍为 80%。月/年总量比只汇总至少 {RATIO_MIN_COVERAGE:.0%} 日覆盖的期间；覆盖不全的期间保留逐期数据但排除在汇总之外，避免把部分月/年误作完整聚合量。

| Station | 名称 | 距离 km | 配对日 | 事件一致率 | Precision | Recall | 日 MAE mm | 湿日 MAE mm | Pearson r | 完整月比中位数 (n) | 完整年比中位数 (n) | 五天 | Gate |
|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|
{chr(10).join(val_lines)}

{station_detail} NOAA 定义 source flag `s` 为中国气象部门来源；`S` 为由全球同步报文汇总，降水需谨慎解释。[GHCN-Daily 官方格式及标志说明](https://www.ncei.noaa.gov/pub/data/ghcn/daily/readme.txt)

### 五天决议

| 日期 | 最终降雨 | 站点 | QC |
|---|---:|---|---|
{chr(10).join(resolution_lines)}

`rain_station_target_values.csv` 保存全部候选站对五天的原始值与 flags；NASA POWER 降雨仅列作辅助，不参与选择。NASA POWER 为网格/模式驱动产品，不称为站点实测。[NASA POWER Daily API 文档](https://power.larc.nasa.gov/docs/services/api/temporal/daily/)

## 4. Candidate、气候统计与 QC

- NASA 接受填补：TMAX {summary['accepted_fill_counts']['TMAX']} 天，TMIN {summary['accepted_fill_counts']['TMIN']} 天，SRAD {summary['accepted_fill_counts']['SRAD']} 天；RAIN 0 天。
- 按 3653 个训练天气日计的来源填补比率：

| 变量 | 已接受 gap-fill 天 | 全训练期占比 |
|---|---:|---:|
{chr(10).join(fill_ratio_lines)}

- 来源层级：2004 非缺失日使用 ChinaFLUX 半小时产品（TMAX/TMIN consistency-based，SRAD 积分，RAIN 完整 48 条求和）；目标五天使用通过验证的邻近 GHCN 日雨量。2005–2013 完整官方 QC 值优先，NASA 只填 TMAX/TMIN/SRAD 的已验证缺口；RAIN 保持官方产品值。
- 2004 candidate 年降雨：{annual_text}。ChinaFLUX 年尺度产品参照为 {annual_reference_text}；仅作独立年值参照，不用于分配目标五天。
- Candidate 与 ChinaFLUX 年/月产品比较：{comparison_text}详见 `chinaflux_2004_candidate_comparison.csv`。
- Candidate QC：`passed={candidate_qc['passed']}`，`not_run={candidate_qc['not_run']}`；原因：{candidate_qc.get('blocked_by', '无') or '无'}。
- `yc_wgen_fitting_weather_2004_2013.csv`：{candidate_qc_text}
- NASA 三变量接受 {sum(summary['accepted_fill_counts'].values())} 个 gap variable-days；五个雨日由地面站 Gate 解决，合计 gap-fill 为 {sum(summary['candidate_gapfill_counts'].values())}/370 个已识别缺口。

## 5. 泄漏与复现

官方源表只解析到 2013；NASA 偏差校正/验证年份为 2005–2013。GHCN `.dly` 只解析 2004-10-16 至 20 及 2005–2013 的 PRCP；2014+ 行仅参与文件 SHA256，未解析天气值。泄漏审计：`{summary['leakage_audit']['passed']}`，最大用于值/校准年份 {summary['leakage_audit']['max_source_weather_years_used_for_values_or_calibration']}。

原始来源 URL、检索时间、HTTP 元数据和文件哈希见 `run_manifest.json`；仅把分析期间的逐日观测子集保存于 `rain_station_source_values.csv`，没有缓存/提交全站历史数据。

## 6. 下一步

当前状态 `{summary['final_status']}`。下一项最小任务是冻结并审阅本轮 candidate，在独立任务中准备 train-only `.CLI` 并开展受控 WGEN pilot；沿用当前单站、训练期范围和资源门槛。本轮未生成 `.CLI`，未运行 WeatherMan/WGEN、DSSAT 或 PPO。
"""


def main() -> int:
    git_start = git_start_snapshot()
    backup = backup_existing_task_outputs()
    OUT.mkdir(parents=True, exist_ok=True)
    previous_validation = load_json(PREVIOUS / "external_gapfill_validation.json")
    previous_leakage = load_json(PREVIOUS / "leakage_audit.json")
    if not previous_leakage.get("passed"):
        raise RuntimeError("Previous leakage audit did not pass")
    gates = variable_gate(previous_validation)
    if not gates["reproduction_passed"]:
        final_status = "BLOCKED_NASA_VARIABLE_GATE_REPRODUCTION"
        accepted_rows: list[dict[str, Any]] = []
        accepted_values: dict[tuple[date, str], float] = {}
    else:
        accepted_rows, accepted_values = accepted_nasa_rows(previous_validation, gates)
        final_status = "BLOCKED_FIVE_DAY_RAIN"
    write_json(OUT / "variable_gate_status.json", gates)
    write_csv(OUT / "accepted_nasa_gapfill.csv", ["date", "variable", "official_original_value", "nasa_raw_value", "bias_correction", "accepted_value", "source", "qc_status"], accepted_rows)

    daily_path = RAW / "YCA_M_daily.zip"
    local_audit = local_daily_product_audit(daily_path)
    local_audit["chinaflux_monthly_rain_2004"] = chinaflux_2004_aggregate_rain("monthly")
    local_audit["chinaflux_annual_rain_2004"] = chinaflux_2004_aggregate_rain("annual")
    legacy_path = RAW / "实体数据及关联说明文档_DP2011_YCA_QXSJ_QXJC.zip"
    local_audit["monthly_source_audit"] = legacy_local_sources_audit(legacy_path)
    write_json(OUT / "local_rain_source_audit.json", local_audit)

    official_met, _ = previous.find_workbooks()
    official = previous.load_official(official_met, "met")
    official_rain = {dt: float(row["RAIN"]) for dt, row in official.items() if 2005 <= dt.year <= 2013 and row.get("RAIN") is not None and row["RAIN"] >= 0}
    cached_station_data = load_station_cache(official_rain)
    if cached_station_data:
        stations, stations_manifest, observations_by_station, downloads, validations = cached_station_data
        print("Reusing scoped NOAA station observations saved in the prior local audit run.")
    else:
        stations, stations_manifest = fetch_nearby_stations()
        observations_by_station = {}
        downloads = []
        validations = []
        for station in stations:
            observations, download = fetch_station_prcp(station)
            observations_by_station[station["station_id"]] = observations
            downloads.append(download)
            validations.append(validate_station(station, observations, official_rain, download))

    nasa_rain = nasa_rain_support()
    write_station_evidence(stations, observations_by_station, official_rain, nasa_rain, validations)
    resolution, selected_station = resolve_five_rain(stations, observations_by_station, validations, nasa_rain)
    write_csv(OUT / "five_day_rain_resolution.csv", ["date", "selected_rain_mm", "selected_source", "station_id", "distance_km", "station_raw_prcp", "nasa_power_prcp", "qc_status", "decision_reason"], resolution)

    valid_candidate = bool(gates["reproduction_passed"] and selected_station)
    candidate_qc: dict[str, Any]
    annual_rows: list[dict[str, Any]] = []
    monthly_rows: list[dict[str, Any]] = []
    chinaflux_comparisons: list[dict[str, Any]] = []
    candidate_path = OUT / "yc_wgen_fitting_weather_2004_2013.csv"
    if valid_candidate:
        selected_id = selected_station["station_id"]
        selected_values = observations_by_station[selected_id]
        cached = load_reusable_candidate(candidate_path, accepted_values, selected_station, selected_values)
        candidate_reused = cached is not None
        if cached:
            output_rows, candidate_qc, annual_rows, monthly_rows = cached
        else:
            chinaflux_daily = previous.load_chinaflux_daily(RAW / "YCA_M_30min.zip", range(2004, 2011))
            _, radiation_path = previous.find_workbooks()
            official_rad = previous.load_official(radiation_path, "rad")
            base = previous.build_base(chinaflux_daily, official, official_rad)
            errors = []
            output_rows = []
            for row in base:
                dt = date.fromisoformat(row["DATE"])
                current = dict(row)
                for variable in ("TMAX", "TMIN", "SRAD"):
                    key = (dt, variable)
                    if key in accepted_values:
                        current[variable] = accepted_values[key]
                        current[f"{variable}_source"] = "nasa_power_train_overlap_bias_corrected"
                        current[f"{variable}_qc"] = "ACCEPTED_VARIABLE_GATE_PASS"
                if dt in TARGET_DATES:
                    rain = selected_values.get(dt)
                    if not rain or rain["qc_status"] != "VALID":
                        errors.append(f"selected_station_target_missing:{dt}")
                    else:
                        current["RAIN"] = rain["station_raw_prcp_mm"]
                        current["RAIN_source"] = f"ghcn_daily_nearby_station_{selected_id}"
                        current["RAIN_qc"] = "ACCEPTED_STATION_OVERLAP_GATE_PASS"
                output_rows.append(current)
            if errors:
                raise RuntimeError("Station passed gate but selected rain values missing: " + ";".join(errors))
            candidate_qc = correlation_and_candidate_rows(output_rows)
        if candidate_qc["passed"]:
            write_csv(candidate_path, ["DATE", "YEAR", "DOY", "SRAD", "TMAX", "TMIN", "RAIN", "SRAD_source", "TMAX_source", "TMIN_source", "RAIN_source", "SRAD_qc", "TMAX_qc", "TMIN_qc", "RAIN_qc"], output_rows)
            if not candidate_reused:
                annual_rows, monthly_rows = monthly_and_annual_summaries(output_rows)
            write_csv(OUT / "candidate_annual_summary.csv", list(annual_rows[0].keys()), annual_rows)
            write_csv(OUT / "candidate_monthly_summary.csv", list(monthly_rows[0].keys()), monthly_rows)
            local_monthly = {int(row["month"]): row["precipitation_mm"] for row in local_audit["chinaflux_monthly_rain_2004"]["values"] if row["valid"]}
            candidate_annual_2004 = next(row["annual_precipitation_mm"] for row in annual_rows if row["year"] == 2004)
            reference_annual_2004 = local_audit["chinaflux_annual_rain_2004"]["annual_precipitation_mm"]
            comparisons = [{"aggregation": "annual", "period": 2004, "candidate_precipitation_mm": candidate_annual_2004, "chinaflux_precipitation_mm": reference_annual_2004, "difference_candidate_minus_chinaflux_mm": candidate_annual_2004 - reference_annual_2004, "days_or_months": 366}]
            for item in (row for row in monthly_rows if row["year"] == 2004 and row["month"] in local_monthly):
                reference = local_monthly[item["month"]]
                comparisons.append({"aggregation": "monthly", "period": f"2004-{item['month']:02d}", "candidate_precipitation_mm": item["monthly_precipitation_mm"], "chinaflux_precipitation_mm": reference, "difference_candidate_minus_chinaflux_mm": item["monthly_precipitation_mm"] - reference, "days_or_months": item["days"]})
            write_csv(OUT / "chinaflux_2004_candidate_comparison.csv", list(comparisons[0].keys()), comparisons)
            chinaflux_comparisons = comparisons
            final_status = "PASS_YC_WEATHER_CANDIDATE"
        else:
            candidate_path.unlink(missing_ok=True)
            final_status = "BLOCKED_CANDIDATE_QC"
            backup_summary_csv(OUT / "candidate_annual_summary.csv", final_status)
            backup_summary_csv(OUT / "candidate_monthly_summary.csv", final_status)
            backup_summary_csv(OUT / "chinaflux_2004_candidate_comparison.csv", final_status)
    else:
        candidate_qc = {"passed": False, "not_run": True, "row_count": None, "expected_days": EXPECTED_DAYS, "blocked_by": "BLOCKED_FIVE_DAY_RAIN" if gates["reproduction_passed"] else "BLOCKED_NASA_VARIABLE_GATE_REPRODUCTION", "errors": ["Candidate construction is conditional on all five rain days passing the station gate."]}
        candidate_path.unlink(missing_ok=True)
        backup_summary_csv(OUT / "candidate_annual_summary.csv", candidate_qc["blocked_by"])
        backup_summary_csv(OUT / "candidate_monthly_summary.csv", candidate_qc["blocked_by"])
        backup_summary_csv(OUT / "chinaflux_2004_candidate_comparison.csv", candidate_qc["blocked_by"])
    write_json(OUT / "candidate_qc.json", candidate_qc)

    fill_counts = dict(Counter(row["variable"] for row in accepted_rows))
    year_fill = Counter((date.fromisoformat(row["date"]).year, row["variable"]) for row in accepted_rows)
    candidate_gapfill_counts = {key: fill_counts.get(key, 0) for key in ("TMAX", "TMIN", "SRAD", "RAIN")}
    if selected_station:
        candidate_gapfill_counts["RAIN"] = len(TARGET_DATES)
        year_fill[(2004, "RAIN")] = len(TARGET_DATES)
    by_year_variable = []
    for year in range(2004, 2014):
        days = 366 if calendar.isleap(year) else 365
        for variable in ("TMAX", "TMIN", "SRAD", "RAIN"):
            count = year_fill[(year, variable)]
            unresolved = len(TARGET_DATES) if year == 2004 and variable == "RAIN" and not selected_station else 0
            status = "GHCN_STATION_GATE_PASS" if year == 2004 and variable == "RAIN" and selected_station else "BLOCKED_FIVE_DAY_RAIN" if unresolved else "NASA_VARIABLE_GATE_PASS" if variable in ("TMAX", "TMIN", "SRAD") and year >= 2005 else "NO_ACCEPTED_GAPFILL"
            by_year_variable.append({"year": year, "variable": variable, "accepted_gapfill_days": count, "unresolved_gap_days": unresolved, "year_days": days, "accepted_fill_fraction_of_year": count / days, "status": status})
    write_csv(OUT / "gap_fill_summary_by_year_variable.csv", list(by_year_variable[0].keys()), by_year_variable)

    source_manifest = {
        "run_time_utc": utc_now(),
        "git_start": git_start,
        "previous_source_files": {
            name: {"path": str((PREVIOUS / name).relative_to(ROOT)), "sha256": sha256_file(PREVIOUS / name)}
            for name in ("external_gapfill_validation.json", "external_gapfill_values.csv", "gaps_before_external_fill.csv", "leakage_audit.json", "run_manifest.json")
        },
        "local_chinaflux_daily": local_audit,
        "ghcn_station_inventory": stations_manifest,
        "ghcn_station_files": downloads,
        "scoped_observation_csv_sha256": sha256_file(OUT / "rain_station_source_values.csv"),
        "official_rain_overlap_years": "2005-2013 only",
        "official_rain_overlap_paired_reference_days": len(official_rain),
        "weather_years_used_for_station_metric_calculation": list(range(2005, 2014)),
        "nasa_calibration_years": "2005-2013 only, reused from previous audit",
        "no_station_full_history_file_saved": True,
    }
    write_json(OUT / "run_manifest.json", source_manifest)
    write_json(OUT / "rain_station_validation.json", {
        "source": "NOAA/NCEI GHCN-Daily station .dly files",
        "source_documentation": "https://www.ncei.noaa.gov/pub/data/ghcn/daily/readme.txt",
        "site": {"latitude": SITE_LAT, "longitude": SITE_LON},
        "search_radius_km": SEARCH_RADIUS_KM,
        "station_count": len(validations),
        "overlap_period": "2005-01-01 through 2013-12-31; valid pair dates only",
        "rain_event_agreement_threshold": RAIN_EVENT_THRESHOLD,
        "minimum_paired_days": MIN_RAIN_PAIRED_DAYS,
        "minimum_paired_days_basis": "Reused previous RAIN_coverage gate (3000 paired days); not a NOAA standard.",
        "aggregate_ratio_summary_min_coverage": RATIO_MIN_COVERAGE,
        "aggregate_ratio_summary_rule": "Only month/year periods with at least 90% of calendar days paired are summarized; incomplete periods remain in period-level records and are excluded from ratio summaries.",
        "rain_event_definition": "official PRCP > 0; station PRCP > 0 or trace MFLAG=T",
        "unit_conversion": "GHCN-Daily PRCP integer tenths of mm divided by 10; no bias or scale adjustment.",
        "quality_rules": {"missing_value": "-9999", "quality_flag": "nonblank QFLAG excluded from paired-day analysis", "source_flag": "retained and reported; NOAA S-source precipitation caution disclosed"},
        "selected_station": selected_station,
        "passing_station_ids": [row["station_id"] for row in validations if row["gate_status"] == "PASS"],
        "stations": validations,
        "final_gate_passed": bool(selected_station),
    })

    official_weather_years = sorted({dt.year for dt in official_rain})
    station_weather_years = sorted({dt.year for obs in observations_by_station.values() for dt in obs})
    source_years_used = set([2004, *official_weather_years, *station_weather_years, *range(2005, 2014)])
    max_used_year = max(source_years_used)
    assert max_used_year <= 2013
    leakage = {
        "passed": max_used_year <= 2013,
        "max_source_weather_years_used_for_values_or_calibration": max_used_year,
        "source_weather_years_used_for_values_or_calibration": {"local_daily_product": [2004], "chinaflux_monthly_and_annual_reference": [2004], "official_daily_validation": official_weather_years, "station_targets": sorted({dt.year for dt in TARGET_DATES}), "station_observations_parsed": station_weather_years, "nasa_bias_correction_reused": list(range(2005, 2014))},
        "post_2013_station_rows_hashed_but_weather_values_not_parsed": True,
        "post_2013_official_weather_cells_read": False,
        "post_2013_overlap_or_calibration_used": False,
        "station_inventory_year_ranges_used_as_weather_values": False,
        "assertion": "assert max(source_weather_years_used_for_values_or_calibration) <= 2013",
    }
    write_json(OUT / "leakage_audit.json", leakage)

    source_inventory = [
        {"source": "ChinaFLUX daily 2004 product", "path": str(daily_path.relative_to(ROOT)), "sha256": local_audit["sha256"], "weather_years_read": "2004 target date and daily PRCP cells only"},
        {"source": "ChinaFLUX monthly 2004 product", "path": "data/external/yc_chinaflux/raw/YCA_M_monthly.zip", "sha256": local_audit["chinaflux_monthly_rain_2004"]["sha256"], "weather_years_read": "2004 monthly precipitation aggregate only"},
        {"source": "ChinaFLUX annual 2004 product", "path": "data/external/yc_chinaflux/raw/YCA_M_yearly.zip", "sha256": local_audit["chinaflux_annual_rain_2004"]["sha256"], "weather_years_read": "2004 annual precipitation aggregate only"},
        {"source": "NOAA GHCN-Daily station inventory", "path": STATIONS_URL, "sha256": stations_manifest["sha256"], "weather_years_read": "station metadata only; no weather values"},
    ]
    for download in downloads:
        source_inventory.append({"source": "NOAA GHCN-Daily station PRCP", "path": download.get("url"), "sha256": download.get("sha256"), "weather_years_read": download.get("parsed_weather_periods")})
    for name in ("external_gapfill_validation.json", "external_gapfill_values.csv", "gaps_before_external_fill.csv", "leakage_audit.json", "run_manifest.json"):
        source_inventory.append({"source": "003_05 audited evidence", "path": str((PREVIOUS / name).relative_to(ROOT)), "sha256": sha256_file(PREVIOUS / name), "weather_years_read": "reuse of already audited 2004-2013 evidence"})
    write_csv(OUT / "source_inventory.csv", ["source", "path", "sha256", "weather_years_read"], source_inventory)

    if selected_station and final_status == "PASS_YC_WEATHER_CANDIDATE":
        candidate_rows = read_csv(candidate_path)
        summary_2004 = next(row for row in annual_rows if row["year"] == 2004)
        annual_2004_precip = float(summary_2004["annual_precipitation_mm"])
    else:
        annual_2004_precip = None
    gate_pass_map = {key: value["gate"] for key, value in gates["variables"].items()}
    summary = {
        "run_time_utc": utc_now(),
        "branch": git_start["branch"],
        "head_at_start": git_start["head"],
        "final_status": final_status,
        "variable_gate_status": gate_pass_map,
        "accepted_fill_counts": {key: fill_counts.get(key, 0) for key in ("TMAX", "TMIN", "SRAD", "RAIN")},
        "candidate_gapfill_counts": candidate_gapfill_counts,
        "accepted_fill_total_variable_days": sum(fill_counts.values()),
        "original_gap_total_variable_days": len(read_csv(PREVIOUS / "gaps_before_external_fill.csv")),
        "gapfill_accepted_count": sum(candidate_gapfill_counts.values()),
        "target_rain_dates": [dt.isoformat() for dt in TARGET_DATES],
        "local_daily_product_audit": local_audit,
        "station_search_radius_km": SEARCH_RADIUS_KM,
        "station_candidates": validations,
        "selected_station": selected_station,
        "selected_target_source_flags": [observations_by_station[selected_station["station_id"]][dt].get("source_flag", "") for dt in TARGET_DATES] if selected_station else [],
        "selected_target_quality_flags": [observations_by_station[selected_station["station_id"]][dt].get("quality_flag", "") for dt in TARGET_DATES] if selected_station else [],
        "five_day_resolution": resolution,
        "candidate_created": candidate_path.exists(),
        "candidate_row_count": candidate_qc.get("row_count"),
        "candidate_qc": candidate_qc,
        "annual_precipitation_2004_mm": annual_2004_precip,
        "annual_reference_chinaflux_2004_mm": local_audit["chinaflux_annual_rain_2004"].get("annual_precipitation_mm"),
        "chinaflux_2004_candidate_comparison": chinaflux_comparisons,
        "leakage_audit": leakage,
        "backup_path": backup,
    }
    write_json(OUT / "audit_summary.json", summary)
    report = report_text(summary, gates, local_audit, validations)
    (DOCS / "yc_weather_gapfill_finalize.md").write_text(report, encoding="utf-8")
    write_experiment_log(summary, gates, local_audit, validations)
    print(json.dumps({
        "final_status": final_status,
        "variable_gate_status": gate_pass_map,
        "accepted_fill_counts": summary["accepted_fill_counts"],
        "station_count": len(validations),
        "passing_stations": [row["station_id"] for row in validations if row["gate_status"] == "PASS"],
        "selected_station": selected_station["station_id"] if selected_station else None,
        "candidate_created": summary["candidate_created"],
        "leakage_passed": leakage["passed"],
    }, ensure_ascii=False, indent=2))
    return 0 if final_status == "PASS_YC_WEATHER_CANDIDATE" else 1


def write_experiment_log(summary: dict[str, Any], gates: dict[str, Any], local_audit: dict[str, Any], validations: list[dict[str, Any]]) -> None:
    lines = [
        "# 实验记录：YC 变量级天气缺口补齐（003_05_01）",
        "",
        f"- 执行时间：{summary['run_time_utc']}",
        f"- 分支：{summary['branch']}；开始 HEAD：{summary['head_at_start']}",
        f"- 最终状态：`{summary['final_status']}`",
        "- 范围：YC/YCA，天气年份 2004–2013；未运行 `.CLI`、WeatherMan、WGEN、DSSAT、PPO。",
        "",
        "## 变量级 Gate",
        "",
    ]
    for variable, gate in gates["variables"].items():
        lines.append(f"- {variable}: `{gate['gate']}`；检查项：`{json.dumps(gate['checks'], ensure_ascii=False)}`")
    lines += [
        f"- RAIN 事件一致率：{gates['variables']['RAIN']['event_agreement_rate']:.6f}；门槛 0.800000，未降低。",
        f"- NASA 接受填补：{summary['accepted_fill_counts']}。只接受 TMAX/TMIN/SRAD，且值与旧表 `raw + correction` 复算一致。",
        "",
        "## 本地雨源检查",
        "",
        f"- ChinaFLUX 日产品有效雨量日：{local_audit['valid_daily_rain_days']}/{local_audit['daily_rows']}。",
        f"- 五天原值：`{json.dumps(local_audit['target_days'], ensure_ascii=False)}`；全部缺测。",
        "- 压缩包中的降水及人工气象记录是月统计，不能代替逐日值。",
        "",
        "## GHCN-Daily 站点验证",
        "",
        f"- 搜索半径：{summary['station_search_radius_km']:.0f} km；站点数：{len(validations)}；有效候选站：`{[r['station_id'] for r in validations if r['gate_status']=='PASS']}`。",
    ]
    for row in validations:
        event_rate = row["event_agreement_rate"]
        event_rate_text = f"{event_rate:.2%}" if event_rate is not None else "NA（无有效配对）"
        lines.append(f"- {row['station_id']} {row['station_name']}，{row['distance_km']:.2f} km，配对 {row['paired_days']} 日，事件一致率 {event_rate_text}，五天 {row['prcp_valid_days_2004_target']}/5，Gate {row['gate_status']}；原因：{row['failure_reasons']}。")
    selected_validation = next(row for row in validations if row["station_id"] == summary["selected_station"]["station_id"])
    partial_2013 = next(row for row in selected_validation["annual_total_ratios"] if row["period"] == "2013")
    lines += [
        "",
        "## 候选数据 QC、泄漏与决策",
        "",
        f"- Candidate 已生成：`{summary['candidate_created']}`；QC：`{summary['candidate_qc']}`。",
        f"- 2004 candidate 年降水：`{summary['annual_precipitation_2004_mm']}`；ChinaFLUX 年统计参照 846.2 mm（不用于拆分）。",
        f"- 济南站湿日 precision：{next(r['wet_day_precision'] for r in validations if r['station_id'] == summary['selected_station']['station_id']):.2%}；月/年总量比仅汇总配对覆盖率至少 {RATIO_MIN_COVERAGE:.0%} 的期间，2013 不完整年度已排除。",
        f"- 济南站 2013 配对覆盖：{partial_2013['paired_days']}/{partial_2013['expected_days']} 日；overlap 来源标志计数：`{selected_validation['overlap_source_flag_counts']}`。",
        f"- 泄漏审计通过：`{summary['leakage_audit']['passed']}`；用于值/校准的最高天气年份：{summary['leakage_audit']['max_source_weather_years_used_for_values_or_calibration']}。",
        "- PPT 结构校验随后运行；无视觉渲染器时仅报告结构验证结果。",
        "- 下一步：冻结并审阅 candidate；独立任务准备 train-only `.CLI` 与受控 WGEN pilot。本轮未运行 WGEN/DSSAT/PPO。",
        "",
    ]
    (OUT / "experiment_log.md").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    raise SystemExit(main())
