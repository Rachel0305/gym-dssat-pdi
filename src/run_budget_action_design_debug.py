from __future__ import annotations

import shutil
import traceback
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from budget_action_wrappers import (
    BudgetActionBaseWrapper,
    ScheduledEventActionWrapper,
    SeasonalBudgetActionWrapper,
    StageBudgetActionWrapper,
)
from episode_profit_reward import EpisodeProfitRewardWrapper, PROFIT_REWARD_CANDIDATES
from ppo_action_safety import SafeActionWrapper
from ppo_evaluate import latest_observation_dict, make_env, normalized_action_dict
from ppo_experiment_plan import find_year
from ppo_plot_results import plot_episode
from ppo_safe_rendering import PROJECT_ROOT, load_yaml
from ppo_train import ppo_kwargs, run_pretrain_smoke_check


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_budget_action_design_debug.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "budget_action_design_debug"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_budget_action_design_debug_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_budget_action_design_debug_report.pptx"


def ensure_dirs() -> None:
    for sub in [
        "configs",
        "wrappers",
        "smoke_checks",
        "rendered_inputs",
        "models/HLA",
        "logs/HLA",
        "tensorboard/HLA",
        "daily_outputs/HLA",
        "evaluation",
        "figures/summary",
        "reports",
    ]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)


def df_to_markdown(df: pd.DataFrame) -> str:
    if df.empty:
        return ""
    work = df.copy()
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    separator = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, separator, *rows])


def design_by_name(config: dict, name: str) -> dict:
    for item in config["action_designs"]:
        if item["name"] == name:
            return dict(item)
    raise KeyError(name)


def row_metadata(config: dict, design: dict, station: str) -> dict:
    reward_version = design["reward_version"]
    return {
        "action_design": design["name"],
        "action_design_type": design["type"],
        "reward_version": reward_version,
        "reward_family": PROFIT_REWARD_CANDIDATES[reward_version].family,
        "cap_name": config["cap"]["cap_name"],
        "season_irrigation_cap": float(config["cap"]["season_irrigation_cap"]),
        "season_n_cap": float(config["cap"]["season_n_cap"]),
        "cv_type": config["training"]["cv_types"].get(station, ""),
    }


def wrap_budget_env(env, config: dict, design: dict):
    cfg = config["budget_action"]
    dtype = design["type"]
    if dtype == "baseline":
        return env
    if dtype == "seasonal_budget":
        return SeasonalBudgetActionWrapper(
            env,
            irrigation_cap=cfg["irrigation_cap"],
            n_cap=cfg["n_cap"],
            irrigation_schedule=cfg["irrigation_schedule"],
            nitrogen_schedule=cfg["nitrogen_schedule"],
        )
    if dtype == "scheduled_event":
        return ScheduledEventActionWrapper(
            env,
            irrigation_events=cfg["irrigation_events"],
            nitrogen_events=cfg["nitrogen_events"],
            single_irrigation_event_max=cfg["single_irrigation_event_max"],
            single_n_event_max=cfg["single_n_event_max"],
            irrigation_cap=cfg["irrigation_cap"],
            n_cap=cfg["n_cap"],
        )
    if dtype == "stage_budget":
        return StageBudgetActionWrapper(env, stages=cfg["stages"], irrigation_cap=cfg["irrigation_cap"], n_cap=cfg["n_cap"])
    raise ValueError(f"Unknown action design type: {dtype}")


def make_budget_env(config: dict, station: str, year: int, seed: int, run_tag: str, evaluation: bool, design: dict):
    base = make_env(config, station, year, seed, run_tag=run_tag, evaluation=evaluation, action_safety_enabled=False)
    budget = wrap_budget_env(base, config, design)
    safe = SafeActionWrapper(budget, {**config.get("action_safety", {}), "enabled": True})
    return EpisodeProfitRewardWrapper(safe, design["reward_version"])


def budget_result_from_env(env) -> Any:
    safe = getattr(env, "env", None)
    inner = getattr(safe, "env", None)
    return getattr(inner, "last_budget_result", None)


def safety_result_from_env(env) -> Any:
    safe = getattr(env, "env", None)
    return getattr(safe, "last_safety_result", None)


def write_review_and_plan(config: dict) -> None:
    previous = PROJECT_ROOT / "Leave_One_experiments" / "action_design_debug" / "evaluation" / "action_design_candidate_comparison.csv"
    prev_text = ""
    if previous.exists():
        prev = pd.read_csv(previous)
        cols = [
            "action_design",
            "mean_irrigation",
            "mean_n",
            "mean_irrigation_saturation_ratio",
            "mean_n_saturation_ratio",
            "mean_design_filtered_days",
            "mean_safety_trigger_days",
            "mean_decision_days",
        ]
        prev_text = "\n\n## 006_06 comparison\n\n" + df_to_markdown(prev[[c for c in cols if c in prev.columns]])
    review = """# A/B action design failure review

Generated at: 2026-06-06

The 006_06 decision-interval and window-gated wrappers reduced action opportunities, but did not change the total-budget decision. PPO could still request large positive actions on the remaining allowed days, and SafeActionWrapper still spent the entire 300/450 seasonal cap.

The high safety-trigger-day counts show that action safety remained the real cap controller. A7/A10/A15 and B window gating filtered daily actions but did not make PPO choose a smaller seasonal budget.

Therefore 006_07 moves from daily action filtering to explicit budget or event semantics: PPO selects a season budget, scheduled event amounts, or stage budgets; wrapper rules execute those decisions.
""" + prev_text
    (OUTPUT_ROOT / "evaluation" / "action_design_failure_review.md").write_text(review, encoding="utf-8")

    plan = pd.DataFrame(
        [
            {
                "design": item["name"],
                "type": item["type"],
                "reward_version": item["reward_version"],
                "description": {
                    "baseline": "daily continuous action with action safety",
                    "seasonal_budget": "PPO selects season budget, wrapper releases fixed split at DAP 25/50/75 and 1/30/60",
                    "scheduled_event": "PPO selects amounts only at fixed event DAPs",
                    "stage_budget": "PPO selects stage budget at stage starts",
                }.get(item["type"], ""),
            }
            for item in config["action_designs"]
        ]
    )
    plan.to_csv(OUTPUT_ROOT / "configs" / "budget_action_design_plan.csv", index=False, encoding="utf-8-sig")
    (OUTPUT_ROOT / "evaluation" / "budget_action_design_plan.md").write_text("# Budget/action design plan\n\n" + df_to_markdown(plan), encoding="utf-8")


def write_design_files(config: dict) -> None:
    shutil.copy2(CONFIG_PATH, OUTPUT_ROOT / "configs" / CONFIG_PATH.name)
    shutil.copy2(PROJECT_ROOT / "src" / "budget_action_wrappers.py", OUTPUT_ROOT / "wrappers" / "budget_action_wrappers.py")
    write_review_and_plan(config)


def train_budget_policy(config: dict, design: dict, station: str, train_year: int):
    from stable_baselines3 import PPO

    seed = int(config["seed"])
    meta = row_metadata(config, design, station)
    smoke_ok, smoke_summary = run_pretrain_smoke_check(config, station, train_year, seed, row_metadata=meta)
    if not smoke_ok:
        raise RuntimeError(f"pretrain smoke check failed: {smoke_summary}")

    policy_name = f"{station}_train{train_year}_{design['name']}_seed{seed}_budget_action"
    env = make_budget_env(config, station, train_year, seed, f"{policy_name}_train", False, design)
    model_dir = OUTPUT_ROOT / "models" / station
    tensorboard_dir = OUTPUT_ROOT / "tensorboard" / station
    model_dir.mkdir(parents=True, exist_ok=True)
    tensorboard_dir.mkdir(parents=True, exist_ok=True)
    model_path = model_dir / policy_name
    try:
        model = PPO(
            "MlpPolicy",
            env,
            verbose=1,
            seed=seed,
            tensorboard_log=str(tensorboard_dir),
            **ppo_kwargs(config, debug=False),
        )
        model.learn(total_timesteps=int(config["training"]["total_timesteps"]), progress_bar=False)
        model.save(str(model_path))
    finally:
        try:
            env.close()
        except Exception:
            pass
    return model, model_path.with_suffix(".zip"), policy_name


def plot_budget_episode(daily: pd.DataFrame, figure_dir: Path) -> None:
    plot_episode(daily, figure_dir)
    figure_dir.mkdir(parents=True, exist_ok=True)
    x = pd.to_numeric(daily["dap"], errors="coerce")

    fig, axes = plt.subplots(2, 1, figsize=(8.5, 6.0), sharex=True)
    for col in ["scheduled_event_irrigation", "executed_irrigation", "real_action_amir"]:
        if col in daily.columns:
            axes[0].plot(x, pd.to_numeric(daily[col], errors="coerce"), label=col)
    axes[0].set_ylabel("Irrigation")
    axes[0].grid(alpha=0.25)
    axes[0].legend(frameon=False)
    for col in ["scheduled_event_n", "executed_n", "real_action_anfer"]:
        if col in daily.columns:
            axes[1].plot(x, pd.to_numeric(daily[col], errors="coerce"), label=col)
    axes[1].set_ylabel("N fertilizer")
    axes[1].set_xlabel("DAP")
    axes[1].grid(alpha=0.25)
    axes[1].legend(frameon=False)
    fig.suptitle("Budget/event actions")
    fig.tight_layout()
    fig.savefig(figure_dir / "budget_or_event_actions.png", dpi=160)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8.5, 4.5))
    for col in ["budget_remaining_irrigation", "budget_remaining_n"]:
        if col in daily.columns:
            ax.plot(x, pd.to_numeric(daily[col], errors="coerce"), label=col)
    ax.set_title("Budget remaining")
    ax.set_xlabel("DAP")
    ax.grid(alpha=0.25)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(figure_dir / "budget_remaining_timeseries.png", dpi=160)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8.5, 4.5))
    if "safety_rule_triggered" in daily.columns:
        triggered = daily["safety_rule_triggered"].fillna("").astype(str).str.len().gt(0).astype(int)
        ax.plot(x, triggered, label="safe_rule_triggered")
    if "is_scheduled_event_day" in daily.columns:
        ax.plot(x, pd.to_numeric(daily["is_scheduled_event_day"], errors="coerce"), label="event_day")
    ax.set_ylim(-0.05, 1.05)
    ax.set_title("Safe trigger and event days")
    ax.set_xlabel("DAP")
    ax.grid(alpha=0.25)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(figure_dir / "safe_trigger_timeseries.png", dpi=160)
    plt.close(fig)


def evaluate_budget_model(model, config: dict, design: dict, station: str, train_year: int, eval_year: int, model_path: Path, policy_name: str) -> dict:
    seed = int(config["seed"])
    train_info = find_year(config, station, train_year)
    eval_info = find_year(config, station, eval_year)
    env = make_budget_env(config, station, eval_year, seed, f"{policy_name}_eval{eval_year}", True, design)
    records: list[dict] = []
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        year_info = find_year(config, station, eval_year)
        while not done and step_count < int(config.get("runtime", {}).get("max_steps", 260)):
            latest = latest_observation_dict(env, obs, info)
            dap = int(round(float(latest.get("dap", step_count) or step_count)))
            raw_action, _ = model.predict(obs, deterministic=True)
            raw_action = np.clip(np.asarray(raw_action, dtype=np.float32).flatten(), -1.0, 1.0)
            obs, reward, terminated, truncated, info = env.step(raw_action)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            budget_result = budget_result_from_env(env)
            safety_result = safety_result_from_env(env)
            if budget_result is None:
                budget_fields = {
                    "raw_policy_action": ",".join(f"{x:.6f}" for x in raw_action.tolist()),
                    "budget_action_irrigation": np.nan,
                    "budget_action_n": np.nan,
                    "scheduled_event_irrigation": np.nan,
                    "scheduled_event_n": np.nan,
                    "executed_irrigation": np.nan,
                    "executed_n": np.nan,
                    "budget_remaining_irrigation": np.nan,
                    "budget_remaining_n": np.nan,
                    "stage_name": "",
                    "event_name": "",
                    "is_budget_decision_day": False,
                    "is_scheduled_event_day": False,
                    "budget_rule_triggered": "",
                }
            else:
                budget_fields = budget_result.__dict__.copy()
                budget_fields["raw_policy_action"] = ",".join(f"{x:.6f}" for x in budget_result.raw_policy_action)
            if safety_result is None:
                safe_real = {"amir": np.nan, "anfer": np.nan}
                safety_flags = {}
            else:
                safe_real = safety_result.safe_real_action
                safety_flags = {
                    "action_clipped_amir": safety_result.action_clipped_amir,
                    "action_clipped_anfer": safety_result.action_clipped_anfer,
                    "safety_rule_triggered": safety_result.safety_rule_triggered,
                }
            norm_action = normalized_action_dict(env, raw_action)
            date = pd.Timestamp(year_info["planting_date"]) + pd.Timedelta(days=max(dap - 1, 0))
            row = {
                "station": station,
                "policy_name": policy_name,
                "action_design": design["name"],
                "action_design_type": design["type"],
                "train_year": train_year,
                "train_year_label": train_info["label"],
                "eval_year": eval_year,
                "eval_year_label": eval_info["label"],
                "seed": seed,
                "date": date.strftime("%Y-%m-%d"),
                "doy": int(date.dayofyear),
                "dap": float(latest.get("dap", dap) or dap),
                "topwt": float(latest.get("topwt", np.nan)) if latest.get("topwt", np.nan) is not None else np.nan,
                "grnwt": float(latest.get("grnwt", np.nan)) if latest.get("grnwt", np.nan) is not None else np.nan,
                "xlai": float(latest.get("xlai", np.nan)) if latest.get("xlai", np.nan) is not None else np.nan,
                "totir": env.total_irrigation,
                "tofer": env.total_n,
                "swfac": float(latest.get("swfac", np.nan)) if latest.get("swfac", np.nan) is not None else np.nan,
                "nstres": float(latest.get("nstres", np.nan)) if latest.get("nstres", np.nan) is not None else np.nan,
                "reward": float(reward),
                "real_action_amir": float(safe_real.get("amir", np.nan)),
                "real_action_anfer": float(safe_real.get("anfer", np.nan)),
                "normalized_action_amir": norm_action.get("amir", np.nan),
                "normalized_action_anfer": norm_action.get("anfer", np.nan),
                "done": done,
            }
            row.update(budget_fields)
            row.update(safety_flags)
            row.update(env.last_profit_info)
            row.update(row_metadata(config, design, station))
            records.append(row)
            step_count += 1
    finally:
        try:
            env.close()
        except Exception:
            pass

    daily = pd.DataFrame(records)
    daily_dir = OUTPUT_ROOT / "daily_outputs" / station
    daily_dir.mkdir(parents=True, exist_ok=True)
    daily_csv = daily_dir / f"{station}_train{train_year}_{design['name']}_eval{eval_year}_seed{seed}_daily.csv"
    daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
    fig_dir = OUTPUT_ROOT / "figures" / station / design["name"] / f"eval_{eval_year}"
    plot_budget_episode(daily, fig_dir)
    episode_completed = bool(records and records[-1]["done"])
    result = {
        "station": station,
        "action_design": design["name"],
        "action_design_type": design["type"],
        "reward_version": design["reward_version"],
        "policy_name": policy_name,
        "train_year": train_year,
        "eval_year": eval_year,
        "seed": seed,
        "season_irrigation_cap": float(config["cap"]["season_irrigation_cap"]),
        "season_n_cap": float(config["cap"]["season_n_cap"]),
        "run_status": "ok" if episode_completed else "failed",
        "episode_completed": episode_completed,
        "episode_length": len(daily),
        "final_grnwt": float(daily["grnwt"].iloc[-1]) if len(daily) else np.nan,
        "total_irrigation": float(daily["real_action_amir"].sum()) if len(daily) else 0.0,
        "total_n_fertilizer": float(daily["real_action_anfer"].sum()) if len(daily) else 0.0,
        "mean_swfac": float(pd.to_numeric(daily["swfac"], errors="coerce").mean()) if len(daily) else np.nan,
        "mean_nstres": float(pd.to_numeric(daily["nstres"], errors="coerce").mean()) if len(daily) else np.nan,
        "mean_reward": float(daily["reward"].mean()) if len(daily) else np.nan,
        "sum_reward": float(daily["reward"].sum()) if len(daily) else 0.0,
        "num_budget_decision_days": int(pd.to_numeric(daily.get("is_budget_decision_day", pd.Series(dtype=float)), errors="coerce").sum()) if len(daily) else 0,
        "num_scheduled_event_days": int(pd.to_numeric(daily.get("is_scheduled_event_day", pd.Series(dtype=float)), errors="coerce").sum()) if len(daily) else 0,
        "num_safe_trigger_days": int(daily.get("safety_rule_triggered", pd.Series(dtype=str)).fillna("").astype(str).str.len().gt(0).sum()) if len(daily) else 0,
        "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
        "figure_dir": str(fig_dir.relative_to(PROJECT_ROOT)),
        "model_path": str(model_path.relative_to(PROJECT_ROOT)),
        "notes": "budget_action_eval",
    }
    result["irrigation_saturation_ratio"] = result["total_irrigation"] / result["season_irrigation_cap"]
    result["n_saturation_ratio"] = result["total_n_fertilizer"] / result["season_n_cap"]
    return result


def append_summary(rows: list[dict]) -> pd.DataFrame:
    out = OUTPUT_ROOT / "evaluation" / "budget_action_design_evaluation_summary.csv"
    new = pd.DataFrame(rows)
    if out.exists():
        old = pd.read_csv(out)
        df = pd.concat([old, new], ignore_index=True)
        df = df.drop_duplicates(["station", "action_design", "train_year", "eval_year", "seed"], keep="last")
    else:
        df = new
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return df


def compare_designs(df: pd.DataFrame) -> pd.DataFrame:
    hla = df[df["station"] == "HLA"].copy()
    hla["quality_gate_pass"] = (hla["run_status"] == "ok") & hla["episode_completed"].astype(str).str.lower().isin(["true", "1"])
    grouped = (
        hla.groupby(["action_design", "action_design_type", "reward_version"], dropna=False)
        .agg(
            eval_count=("eval_year", "count"),
            ok_count=("quality_gate_pass", "sum"),
            mean_yield=("final_grnwt", "mean"),
            std_yield=("final_grnwt", "std"),
            mean_reward=("mean_reward", "mean"),
            mean_irrigation=("total_irrigation", "mean"),
            mean_n=("total_n_fertilizer", "mean"),
            mean_irrigation_saturation_ratio=("irrigation_saturation_ratio", "mean"),
            mean_n_saturation_ratio=("n_saturation_ratio", "mean"),
            mean_swfac=("mean_swfac", "mean"),
            mean_nstres=("mean_nstres", "mean"),
            mean_budget_decision_days=("num_budget_decision_days", "mean"),
            mean_scheduled_event_days=("num_scheduled_event_days", "mean"),
            mean_safe_trigger_days=("num_safe_trigger_days", "mean"),
        )
        .reset_index()
    )
    base = grouped[grouped["action_design"] == "baseline_daily_action_current"]
    base_yield = float(base.iloc[0]["mean_yield"]) if len(base) else float(grouped["mean_yield"].max())
    grouped["yield_loss_vs_baseline"] = (base_yield - grouped["mean_yield"]) / base_yield
    grouped["input_reduction_vs_baseline"] = 1.0 - (grouped["mean_irrigation"] + grouped["mean_n"]) / (300.0 + 450.0)
    grouped["strict_pass"] = (
        (grouped["mean_irrigation_saturation_ratio"] < 0.95)
        & (grouped["mean_n_saturation_ratio"] < 0.95)
        & (grouped["yield_loss_vs_baseline"] <= 0.15)
        & (grouped["ok_count"] == grouped["eval_count"])
    )
    grouped["relaxed_pass"] = (
        (
            (grouped["mean_irrigation_saturation_ratio"] < 0.95)
            | (grouped["mean_n_saturation_ratio"] < 0.95)
        )
        & (grouped["yield_loss_vs_baseline"] <= 0.15)
        & (grouped["ok_count"] == grouped["eval_count"])
    )
    grouped["action_too_restrictive"] = (
        (grouped["mean_irrigation"] < 1.0)
        & (grouped["mean_n"] < 1.0)
        & (grouped["yield_loss_vs_baseline"] > 0.15)
    )
    out = OUTPUT_ROOT / "evaluation" / "budget_action_design_candidate_comparison.csv"
    grouped.to_csv(out, index=False, encoding="utf-8-sig")
    return grouped


def select_design(comparison: pd.DataFrame) -> str | None:
    passed = comparison[comparison["strict_pass"]].copy()
    if passed.empty:
        passed = comparison[comparison["relaxed_pass"]].copy()
    if passed.empty:
        return None
    passed = passed.sort_values(["yield_loss_vs_baseline", "input_reduction_vs_baseline"], ascending=[True, False])
    return str(passed.iloc[0]["action_design"])


def plot_summary(comparison: pd.DataFrame) -> None:
    fig_root = OUTPUT_ROOT / "figures" / "summary"
    fig_root.mkdir(parents=True, exist_ok=True)
    labels = comparison["action_design"]
    plt.figure(figsize=(11, 6))
    plt.scatter(comparison["mean_irrigation"] + comparison["mean_n"], comparison["mean_yield"])
    for _, row in comparison.iterrows():
        plt.annotate(row["action_design"].replace("_", "\n"), (row["mean_irrigation"] + row["mean_n"], row["mean_yield"]), fontsize=7)
    plt.xlabel("Mean water + nitrogen input")
    plt.ylabel("Mean grain yield")
    plt.tight_layout()
    plt.savefig(fig_root / "budget_action_yield_vs_input.png", dpi=180)
    plt.close()

    plt.figure(figsize=(12, 6))
    plt.plot(labels, comparison["mean_irrigation_saturation_ratio"], marker="o", label="irrigation ratio")
    plt.plot(labels, comparison["mean_n_saturation_ratio"], marker="o", label="nitrogen ratio")
    plt.axhline(0.95, color="black", linestyle="--", linewidth=1)
    plt.xticks(rotation=45, ha="right")
    plt.ylabel("Saturation ratio")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_root / "budget_action_saturation_ratio.png", dpi=180)
    plt.close()

    plt.figure(figsize=(11, 6))
    plt.scatter(comparison["input_reduction_vs_baseline"], comparison["yield_loss_vs_baseline"])
    for _, row in comparison.iterrows():
        plt.annotate(row["action_design"].replace("_", "\n"), (row["input_reduction_vs_baseline"], row["yield_loss_vs_baseline"]), fontsize=7)
    plt.axhline(0.15, color="black", linestyle="--", linewidth=1)
    plt.xlabel("Input reduction vs baseline")
    plt.ylabel("Yield loss vs baseline")
    plt.tight_layout()
    plt.savefig(fig_root / "budget_action_yield_loss_vs_input_reduction.png", dpi=180)
    plt.close()

    plt.figure(figsize=(12, 6))
    plt.bar(labels, comparison["mean_safe_trigger_days"])
    plt.xticks(rotation=45, ha="right")
    plt.ylabel("Mean safe trigger days")
    plt.tight_layout()
    plt.savefig(fig_root / "budget_action_safe_trigger_days.png", dpi=180)
    plt.close()

    plt.figure(figsize=(12, 6))
    plt.plot(labels, comparison["mean_swfac"], marker="o", label="swfac")
    plt.plot(labels, comparison["mean_nstres"], marker="o", label="nstres")
    plt.xticks(rotation=45, ha="right")
    plt.ylabel("Mean stress indicators")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_root / "budget_action_swfac_nstres.png", dpi=180)
    plt.close()


def run_design(config: dict, design: dict, station: str, train_year: int, eval_years: list[int]) -> list[dict]:
    print(f"[budget-action] train {station} {train_year} {design['name']}", flush=True)
    model, model_path, policy_name = train_budget_policy(config, design, station, train_year)
    rows = []
    for eval_year in eval_years:
        print(f"[budget-action] evaluate {station} {design['name']} eval {eval_year}", flush=True)
        rows.append(evaluate_budget_model(model, config, design, station, train_year, int(eval_year), model_path, policy_name))
    return rows


def run_hla_pilot(config: dict) -> tuple[pd.DataFrame, pd.DataFrame, str | None]:
    all_rows: list[dict] = []
    for design in config["action_designs"]:
        try:
            rows = run_design(
                config,
                design,
                config["training"]["hla_pilot_station"],
                int(config["training"]["hla_train_year"]),
                [int(y) for y in config["training"]["hla_eval_years"]],
            )
            all_rows.extend(rows)
            summary = append_summary(all_rows)
            comparison = compare_designs(summary)
            plot_summary(comparison)
        except Exception as exc:
            all_rows.append(
                {
                    "station": config["training"]["hla_pilot_station"],
                    "action_design": design["name"],
                    "action_design_type": design["type"],
                    "reward_version": design["reward_version"],
                    "train_year": int(config["training"]["hla_train_year"]),
                    "eval_year": np.nan,
                    "seed": int(config["seed"]),
                    "run_status": "failed",
                    "episode_completed": False,
                    "error_message": f"{type(exc).__name__}: {exc}",
                    "traceback": traceback.format_exc()[-2000:],
                }
            )
            append_summary(all_rows)
            raise
    summary = append_summary(all_rows)
    comparison = compare_designs(summary)
    plot_summary(comparison)
    return summary, comparison, select_design(comparison)


def write_blocker_report() -> None:
    text = """# Budget action design blocker report

Generated at: 2026-06-06

D/E/F explicit action semantics did not identify a strict or relaxed HLA pilot candidate. The next step should shift away from PPO-only online learning:

1. Rule-based expert policy plus imitation learning.
2. Offline search over fixed management schedules.
3. Bayesian optimization over water/N budgets and event amounts.
4. DSSAT scenario ensemble rather than PPO.
"""
    (OUTPUT_ROOT / "evaluation" / "budget_action_design_blocker_report.md").write_text(text, encoding="utf-8")


def run_extension(config: dict, selected: str | None) -> pd.DataFrame:
    out = OUTPUT_ROOT / "evaluation" / "cross_site_budget_action_design_test.csv"
    if not selected:
        empty = pd.DataFrame([{"selected_action_design": "", "run_status": "skipped", "notes": "No HLA budget/event design passed strict or relaxed criteria."}])
        empty.to_csv(out, index=False, encoding="utf-8-sig")
        write_blocker_report()
        return empty
    design = design_by_name(config, selected)
    all_rows: list[dict] = []
    for station in ["SYA", "LCA"]:
        train_year = int(config["training"]["best_train_years"][station])
        eval_years = [int(item["year"]) for item in config["observed_years"][station]]
        all_rows.extend(run_design(config, design, station, train_year, eval_years))
    cross = pd.DataFrame(all_rows)
    cross.to_csv(out, index=False, encoding="utf-8-sig")
    return cross


def write_report(comparison: pd.DataFrame, cross: pd.DataFrame, selected: str | None) -> None:
    table = comparison.copy()
    for col in table.select_dtypes(include=["float", "int"]).columns:
        table[col] = table[col].round(4)
    cross_table = cross.copy()
    if len(cross_table):
        for col in cross_table.select_dtypes(include=["float", "int"]).columns:
            cross_table[col] = cross_table[col].round(4)
    unsaturated = table[(table["mean_irrigation_saturation_ratio"] < 0.95) & (table["mean_n_saturation_ratio"] < 0.95)]["action_design"].tolist()
    yield_ok = table[table["yield_loss_vs_baseline"] <= 0.15]["action_design"].tolist()
    md = f"""# Budget/action design debug report

Generated at: 2026-06-06

## Goal

This stage moves beyond daily filtering to explicit seasonal budget, scheduled event, and stage budget action semantics. It does not modify my_data or site-packages reward files and does not overwrite 006_03/006_04/006_05/006_06.

## Implemented wrappers

- `SeasonalBudgetActionWrapper`: PPO chooses season water/N budget on DAP 1; wrapper releases fixed splits.
- `ScheduledEventActionWrapper`: PPO chooses event amounts only at fixed DAP events.
- `StageBudgetActionWrapper`: PPO chooses stage budget at stage starts.

## Wrapper order

`GymDssatWrapper -> BudgetActionWrapper -> SafeActionWrapper -> EpisodeProfitRewardWrapper`

## HLA comparison

{df_to_markdown(table)}

## Pass summary

- Designs with both irrigation and N unsaturated: {', '.join(unsaturated) if unsaturated else 'None'}
- Designs with yield loss <= 15%: {', '.join(yield_ok) if yield_ok else 'None'}
- Recommended design: `{selected or 'None'}`

## SYA/LCA extension

{df_to_markdown(cross_table) if len(cross_table) else 'Extension was skipped because HLA did not identify a strict or relaxed passing budget/event design.'}

## Interpretation

If D/E/F still select cap-level inputs, PPO is choosing maximum budget even when action semantics are explicit. If D/E/F reduce input but lose too much yield, the next step is offline schedule search or Bayesian optimization to build a stronger prior before RL.

## Next step

Do not enter multi-seed or rainfall-scaling until a budget/event design is interpretable. If this stage fails, shift to imitation learning, offline schedule search, Bayesian optimization, or DSSAT scenario ensemble.
"""
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)
    DOC_MD.write_text(md, encoding="utf-8")
    shutil.copy2(DOC_MD, OUTPUT_ROOT / "reports" / DOC_MD.name)
    write_ppt(comparison, selected)


def write_ppt(comparison: pd.DataFrame, selected: str | None) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        (OUTPUT_ROOT / "reports" / "ppt_generation_skipped_in_container.txt").write_text(
            "python-pptx is not installed in the DSSAT container. Generate the PPT from Windows Python after training.",
            encoding="utf-8",
        )
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_title(slide, title: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True

    def add_body(slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.0), Inches(5.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(15)

    def picture_slide(title: str, file_name: str) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, title)
        path = OUTPUT_ROOT / "figures" / "summary" / file_name
        if path.exists():
            slide.shapes.add_picture(str(path), Inches(0.8), Inches(1.0), width=Inches(11.8))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "006_07 Budget/Event Action Design")
    add_body(slide, "目标：从每日动作过滤转向显式季节预算、固定事件和阶段预算动作。\n约束：HLA 2011 pilot；action safety；不改原始 reward/my_data。")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Why D/E/F")
    add_body(slide, "A7/A10/A15 和窗口门控仍然打满 300/450。\n这说明减少动作机会不够，需要让 PPO 直接选择总预算或关键事件量。")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Implemented Designs")
    add_body(slide, "D：第一天选择季节水氮预算，按固定比例释放。\nE：只在 DAP 1/25/30/50/60/75 固定事件日决策。\nF：在阶段开始日选择阶段预算。\nD/E plus profit 使用 episode-level profit reward。")
    picture_slide("Yield vs Input", "budget_action_yield_vs_input.png")
    picture_slide("Saturation Ratio", "budget_action_saturation_ratio.png")
    picture_slide("Yield Loss vs Input Reduction", "budget_action_yield_loss_vs_input_reduction.png")
    picture_slide("Safe Trigger Days", "budget_action_safe_trigger_days.png")
    picture_slide("SWFAC and NSTRES", "budget_action_swfac_nstres.png")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Conclusion")
    add_body(slide, f"推荐 design：{selected or 'None'}。\n如果 D/E/F 仍失败，下一步应转向 imitation learning、offline schedule search、Bayesian optimization 或 DSSAT scenario ensemble。")
    DOC_PPT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DOC_PPT)
    shutil.copy2(DOC_PPT, OUTPUT_ROOT / "reports" / DOC_PPT.name)


def main() -> None:
    ensure_dirs()
    config = load_yaml(CONFIG_PATH)
    write_design_files(config)
    summary, comparison, selected = run_hla_pilot(config)
    cross = run_extension(config, selected)
    write_report(comparison, cross, selected)
    print("[budget-action] done")
    print((OUTPUT_ROOT / "evaluation" / "budget_action_design_candidate_comparison.csv").relative_to(PROJECT_ROOT))
    print(f"[budget-action] selected={selected}")


if __name__ == "__main__":
    main()
