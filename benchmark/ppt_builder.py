"""Generate and structurally QA the 021_00 benchmark framework PPTX."""

from __future__ import annotations

import logging
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


LOGGER = logging.getLogger(__name__)
BLACK = RGBColor(25, 25, 25)
GRAY = RGBColor(95, 95, 95)
LIGHT = RGBColor(235, 238, 241)
BLUE = RGBColor(54, 91, 134)
GREEN = RGBColor(39, 111, 78)


def _items(value: Any, fallback: str) -> list[str]:
    if value is None:
        return [fallback]
    if isinstance(value, str):
        return [value]
    if isinstance(value, Mapping):
        return [f"{key}: {item}" for key, item in value.items()]
    if isinstance(value, Sequence):
        return [str(item) for item in value] or [fallback]
    return [str(value)]


def _add_text(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: float = 18,
    bold: bool = False,
    color: RGBColor = BLACK,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> Any:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _title(slide: Any, text: str) -> None:
    _add_text(slide, text, 0.65, 0.35, 12.0, 0.55, size=27, bold=True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.0), Inches(12.0), Inches(0.02))
    rule.fill.solid()
    rule.fill.fore_color.rgb = BLACK
    rule.line.fill.background()


def _bullets(slide: Any, items: list[str], top: float = 1.35, size: float = 17) -> None:
    box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.4), Inches(5.3))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items[:7]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = str(item)
        paragraph.level = 0
        paragraph.space_after = Pt(12)
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = BLACK
        paragraph.text = f"• {paragraph.text}"


def _footer(slide: Any, text: str = "021_00 Benchmark Framework Refactor") -> None:
    _add_text(slide, text, 0.7, 7.12, 12.0, 0.18, size=7.5, color=GRAY)


def _notes(slide: Any, text: str) -> None:
    try:
        frame = slide.notes_slide.notes_text_frame
        frame.text = text
    except (AttributeError, NotImplementedError):
        LOGGER.debug("Speaker notes are not supported by this python-pptx build")


def _workflow_slide(slide: Any) -> None:
    _title(slide, "统一流程把配置、计算与证据串成可追溯链条")
    labels = [
        "config",
        "validation",
        "result lookup",
        "train / resume",
        "evaluate",
        "summarize",
        "plot",
        "MD / PPT",
        "Git backup",
    ]
    x0, y, width, height, gap = 0.45, 2.35, 1.18, 0.9, 0.22
    for index, label in enumerate(labels):
        left = x0 + index * (width + gap)
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(y), Inches(width), Inches(height)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT if index not in (0, 8) else RGBColor(218, 228, 238)
        box.line.color.rgb = BLACK
        frame = box.text_frame
        frame.clear()
        frame.vertical_anchor = 1
        p = frame.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = index in (0, 8)
        if index < len(labels) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(left + width + 0.02),
                Inches(y + 0.31),
                Inches(gap - 0.01),
                Inches(0.28),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GRAY
            arrow.line.fill.background()
    _add_text(
        slide,
        "旧脚本由 adapter 接入；任何复用、失败或部分完成都写入 manifest。",
        1.0,
        4.2,
        11.2,
        0.6,
        size=18,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    _notes(slide, "流程图展示本次重构的控制面；底层已验证 DQN/reward/IC 不在本任务中修改。")


def _example_manifest(manifest: Mapping[str, Any]) -> list[str]:
    return [
        f"experiment_id: {manifest.get('experiment_id', 'pending')}",
        f"config_hash: {str(manifest.get('config_hash', 'pending'))[:16]}...",
        f"status: {manifest.get('status', 'pending')}",
        f"station/year/seed: {manifest.get('station_code', 'NA')} / {manifest.get('year', 'NA')} / {manifest.get('seed', 'NA')}",
        "resume_scope: checkpoint + replay buffer + RNG; DSSAT process is restarted at episode boundary",
    ]


def build_pptx_report(
    output_dir: str | Path,
    *,
    config: Mapping[str, Any] | None = None,
    manifest: Mapping[str, Any] | None = None,
    tables: Mapping[str, Any] | None = None,
    figures: Mapping[str, Any] | None = None,
    context: Mapping[str, Any] | None = None,
    output_path: str | Path | None = None,
) -> Path:
    """Build the actual methods-style Chinese PPTX and run structural QA."""

    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    path = Path(output_path) if output_path else destination / "benchmark_report.pptx"
    path.parent.mkdir(parents=True, exist_ok=True)
    cfg = dict(config or {})
    run_manifest = dict(manifest or {})
    ctx = dict(context or {})
    experiment = dict(cfg.get("experiment", {}))

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    _add_text(slide, "DQN Benchmark Framework Refactor", 0.8, 1.4, 11.8, 0.8, size=32, bold=True)
    _add_text(slide, "论文实验前的统一配置、复用、续跑与报告基础工程", 0.82, 2.35, 11.2, 0.45, size=19)
    _add_text(slide, f"2026-07-11  |  {experiment.get('experiment_id', '021_00')}", 0.82, 6.55, 11.0, 0.3, size=10, color=GRAY)
    _notes(slide, "本次只重构工程控制面，不修改冻结 reward、IC 或已验证 DQN 主线。")

    slide = prs.slides.add_slide(blank)
    _title(slide, "脚本复制已经成为论文实验的主要工程风险")
    _bullets(slide, _items(ctx.get("background"), "五站点证据已经形成，但入口和输出分散。")[:4])
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "审计确认了复用价值，也暴露了 resume 与输入身份边界")
    _bullets(slide, _items(ctx.get("current_problems"), "当前问题未注入。")[:6], size=16)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _workflow_slide(slide)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "YAML 把站点差异和冻结算法参数分离")
    configuration = _items(ctx.get("configuration"), "配置摘要未注入。")[:5]
    _bullets(slide, configuration, size=16)
    _add_text(slide, "科学配置进入 config_hash；dry-run、resume 等运行控制不改变科学身份。", 0.9, 5.85, 11.3, 0.5, size=16, bold=True, color=BLUE)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "复用等级避免把“可重画图”误写成“可精确续训”")
    _bullets(slide, _items(ctx.get("reuse"), "复用结果未注入。")[:6], size=16)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "统一目录让每个输出都能回溯到配置和 manifest")
    _bullets(slide, [
        "configs / logs / checkpoints / evaluations",
        "summaries / figures / tables / reports / manifests",
        "CSV 与图一一对应；PNG 300 dpi 与 SVG 成对输出",
        "模型与 replay buffer 保存在本地，不纳入 Git 大文件备份",
    ], size=17)
    _footer(slide)

    smoke = ctx.get("smoke_test", {})
    slide = prs.slides.add_slide(blank)
    _title(slide, "Smoke test 只验证管线，不作为算法性能证据")
    _bullets(slide, _items(smoke.get("checks") if isinstance(smoke, Mapping) else smoke, "smoke test 尚未运行。")[:7], size=15)
    _add_text(slide, f"状态：{smoke.get('status', '未验证') if isinstance(smoke, Mapping) else '未验证'}", 0.9, 6.1, 4.0, 0.4, size=18, bold=True, color=GREEN)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "一个命令对应一个不可混淆的实验配置")
    command = _items(ctx.get("commands"), "python -m benchmark.benchmark_runner --config <yaml> --dry-run")[0]
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.5))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(247, 247, 247); box.line.color.rgb = BLACK
    _add_text(slide, command, 1.15, 2.55, 11.0, 0.65, size=16, color=BLACK)
    _add_text(slide, "先 dry-run；确认后再 train / evaluate / report。", 1.0, 4.5, 11.0, 0.5, size=18, bold=True, align=PP_ALIGN.CENTER)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "Manifest 记录身份、状态、来源和失败证据")
    _bullets(slide, _example_manifest(run_manifest), size=15)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "兼容式重构优先保护既有证据")
    _bullets(slide, _items(ctx.get("risks"), "风险与保护措施未注入。")[:7], size=15)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "下一轮实验只启用已通过 smoke 的能力")
    _bullets(slide, _items(ctx.get("ready"), "尚无已验证能力可宣布。")[:6], size=16)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "后续模板已预留，但默认全部 dry-run")
    _bullets(slide, _items(ctx.get("next_steps"), "等待 smoke test 后启动 021_01–021_08。")[:7], size=15)
    _footer(slide)

    slide = prs.slides.add_slide(blank)
    _title(slide, "Methods Source 与 References")
    references = _items(ctx.get("references"), "参考文献未注入。")[:8]
    _bullets(slide, references, top=1.25, size=11.5)
    _footer(slide)

    prs.save(path)
    qa_path = destination / "qa_report.md"
    qa = audit_pptx(path)
    qa_path.write_text(qa, encoding="utf-8")
    LOGGER.info("PPTX written to %s", path)
    return path


def audit_pptx(path: str | Path) -> str:
    """Reopen a PPTX and report bounds, density, notes, and placeholders."""

    presentation = Presentation(path)
    width, height = presentation.slide_width, presentation.slide_height
    defects: list[tuple[str, int, str]] = []
    notes_count = 0
    for slide_index, slide in enumerate(presentation.slides, 1):
        slide_chars = 0
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width or shape.top + shape.height > height:
                defects.append(("high", slide_index, "shape exceeds slide bounds"))
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                slide_chars += len(text)
                if any(token in text.lower() for token in ("lorem", "xxxx", "todo placeholder")):
                    defects.append(("high", slide_index, "unreplaced placeholder text"))
                if len(max(text.split(), key=len, default="")) > 80:
                    defects.append(("medium", slide_index, "long unbroken token may overflow"))
        if slide_chars > 900:
            defects.append(("medium", slide_index, f"text density is high ({slide_chars} characters)"))
        try:
            note_text = slide.notes_slide.notes_text_frame.text.strip()
            if note_text:
                notes_count += 1
        except (AttributeError, NotImplementedError):
            pass
    high = [item for item in defects if item[0] == "high"]
    medium = [item for item in defects if item[0] == "medium"]
    lines = [
        "# PPTX QA report",
        "",
        f"- file: `{Path(path)}`",
        f"- slide_count: {len(presentation.slides)}",
        f"- notes_slides: {notes_count}",
        f"- high_severity_defects: {len(high)}",
        f"- medium_severity_defects: {len(medium)}",
        "- rendered_preview: unavailable; used python-pptx structural review",
        "",
        "## Defects",
        "",
    ]
    if defects:
        lines.extend(f"- {severity} — slide {slide}: {message}" for severity, slide, message in defects)
    else:
        lines.append("- No structural bounds, placeholder, or density defect detected.")
    lines.extend([
        "",
        "## Corrective review",
        "",
        "- The deck uses a process-wide workflow, text-led safeguards, and manifest evidence rather than repeated decorative cards.",
        "- All generated shapes were re-opened and checked against the slide canvas.",
        "- A manual PowerPoint rendering check is still recommended before an advisor presentation.",
    ])
    return "\n".join(lines) + "\n"


# Backwards-friendly alias.
build_ppt = build_pptx_report

