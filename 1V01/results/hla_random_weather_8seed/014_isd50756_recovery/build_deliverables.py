#!/usr/bin/env python
"""Build the Chinese 014 report and five-slide evidence deck."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
DOCS = ROOT / "docs"
GATE = json.loads((OUT / "final_gate.json").read_text(encoding="utf-8"))
IDENTITY = json.loads((OUT / "station_identity.json").read_text(encoding="utf-8"))
OCC = json.loads((OUT / "isd_occurrence_summary.json").read_text(encoding="utf-8"))
CLIMATE = json.loads((OUT / "2013_climate_sanity.json").read_text(encoding="utf-8"))
PROV = json.loads((OUT / "input_provenance.json").read_text(encoding="utf-8"))
GH = GATE["same_station_ghcn_qa"]
HL = GATE["hla_d222_qa"]
GS = GATE["gsod_crosscheck"]
REPORT = DOCS / "hla_isd50756_recovery_014.md"
DECK = DOCS / "hla_isd50756_recovery_014.pptx"

NAVY = RGBColor(28, 54, 69)
TEAL = RGBColor(20, 112, 126)
PALE = RGBColor(242, 247, 247)
INK = RGBColor(36, 47, 54)
MUTED = RGBColor(91, 108, 117)
RED = RGBColor(165, 61, 52)
WHITE = RGBColor(255, 255, 255)
FONT = "Microsoft YaHei"


def build_report():
    d222 = CLIMATE["selected_source_subtotals"].get("HLA_LOCAL_D222_CHAIN", {})
    ghcn = CLIMATE["selected_source_subtotals"].get("NOAA_GHCN_DAILY", {})
    secondary = CLIMATE["selected_source_subtotals"].get("NOAA_GHCN_DAILY_SECONDARY", {})
    event_rows = []
    for e in CLIMATE["isd_vs_top_selected_events"][:5]:
        near = "; ".join(
            f"{x['gsod_date_utc']} {x['rain_mm']:.3f} mm"
            for x in e["gsod_nearby_positive_rain"]
        ) or "无可用 A-G 记录"
        event_rows.append(
            f"| {e['date']} | {e['selected_rain_mm']:.1f} | {e['selected_source']} | {near} |"
        )
    source_rows = []
    for x in PROV["retrieved_files"]:
        source_rows.append(
            f"| {Path(x['path']).name} | {x['size_bytes']:,} | {x['sha256']} | "
            f"{x['retrieved_at_local']} | [{x['url']}]({x['url']}) |"
        )
    july30 = next((x for x in CLIMATE["isd_vs_top_selected_events"] if x["date"] == "2013-07-30"), None)
    july30_gsod = "; ".join(
        f"{x['gsod_date_utc']}={x['rain_mm']:.3f} mm"
        for x in (july30["gsod_nearby_positive_rain"] if july30 else [])
    )
    report = f"""# HLA 同站 NOAA ISD 降水恢复复审（014）

审计日期：2026-10-05（Asia/Shanghai）

## 结论

NOAA ISD 站点身份核对通过：507560-99999 的 2013 SYNOP 报文明确包含 WMO 50756，站名为 HAILUN，观测记录的主位置与 GHCN CHM00050756 目录位置一致。但 2013 ISD 中的 AA1–AA4 降水 occurrence 没有一条能成为质量合格、可量化且可唯一分配的降水 interval。491 条记录由 374 个未闭合的累计期开始标记、50 个 trace、67 个缺失 condition code 构成；另有 2 条质量码为 suspect。

ISD 能识别部分缺口日期存在原始降水字段/窗口标记，却不能恢复可信日降水值。013 的 111 个缺口日中，16 日有 AA 时间窗/标记覆盖，95 日没有 AA 时间窗信息；有效完整日值 0 天、恢复 0 天、仍缺 111 天。ISD 回退未接受，雨量 provenance 仍为 PARTIAL，READY_TO_RESUME_010 = FALSE。

| Gate | 结果 |
|---|---|
| ISD 与 HLA/50756 同站身份 | {GATE['isd_same_station_identity']} |
| 2013 可接受 ISD 完整本地日值 | {GATE['isd_daily_valid_days_2013']} 天 |
| 013 缺口恢复 | {GATE['isd_recovered_unresolved_days']} / {GATE['starting_unresolved_days']} 天 |
| 014 后剩余缺口 | {GATE['remaining_unresolved_days']} 天 |
| ISD 同站回退 | {GATE['isd_same_station_fallback']} |
| 2013 雨量来源闭合 | {GATE['hla_rain_provenance']} |
| 可恢复 010 | {str(GATE['ready_to_resume_010']).upper()} |

## 1. 站点身份

| 核对项 | NOAA ISD | NOAA GHCN-Daily |
|---|---|---|
| 标识 | USAF 507560 / WBAN 99999 | CHM00050756 |
| 站名 / 国家 | HAILUN / CH | HAILUN / CH |
| WMO | 2013 报文 remarks 中 50756 共 {IDENTITY['isd_synop_explicit_wmo_50756_records']} 条 | 站点目录 WMO 字段 50756 |
| 2013 记录主位置 | {IDENTITY['isd_2013_modal_observation_location']['latitude']:.3f}°N, {IDENTITY['isd_2013_modal_observation_location']['longitude']:.3f}°E, {IDENTITY['isd_2013_modal_observation_location']['elevation_m']:.0f} m | {IDENTITY['ghcn_station_catalog']['latitude']:.3f}°N, {IDENTITY['ghcn_station_catalog']['longitude']:.3f}°E, {IDENTITY['ghcn_station_catalog']['elevation_m']:.0f} m |

ISD 站史静态目录列为 47.450°N、126.867°E、248 m；它与 ISD 2013 报文众数位置及 GHCN 站点目录（47.433°N、126.967°E、240 m）不同。名称、国家、WMO 报文编号与 USAF/WBAN 标识相符，且 2013 报文位置与 GHCN 完全吻合，故身份判为 PASS；静态站史坐标差异保留供复核。GSOD 文件使用的坐标与静态 ISD 站史一致。

## 2. 原始记录与来源

官方原始文件 507560-99999-2013.gz 共 {OCC['raw_records']:,} 条记录，UTC 日期覆盖 {OCC['date_first_utc']} 至 {OCC['date_last_utc']}；其中 {OCC['aa_occurrence_count']} 个 AA1–AA4 occurrence，{OCC['records_without_AA1_AA4']:,} 条记录不含这些降水字段。ISD 记录按 UTC 标时；本地日按 Asia/Shanghai（UTC+08:00）界定。

文件、大小、SHA-256 和获取时间如下。原始 NOAA 文件保留在项目 data/external/noaa_isd_hailun_50756/raw/，不纳入 Git。

| 文件 | 字节 | SHA-256 | 获取时间（本地） | NOAA 来源 |
|---|---:|---|---|---|
{chr(10).join(source_rows)}

冻结输入包括 013 后 source selection、013 unresolved 清单、013 年度 sanity JSON 和项目已有的 GHCN CHM00050756.dly。014 没有修改这些输入。

## 3. AA1–AA4 解析和 interval 选择

[NOAA 官方 ISD 格式文档]({PROV['retrieved_files'][3]['url']})定义 AA1–AA4 为重复的液态降水 period、depth、condition、quality 字段。period 以小时计，99 缺测；depth 单位为毫米、缩放因子 10（原始整数除以 10），9999 缺测。condition 和 quality 原码均保存在 occurrence 表。

| 字段 | 观察到的值 |
|---|---|
| AA slot | AA1 {OCC['slot_counts'].get('AA1', 0)} 条；AA2 {OCC['slot_counts'].get('AA2', 0)} 条；AA3/AA4 0 条 |
| period | 6 h：{OCC['period_hours_counts'].get('06', 0)}；24 h：{OCC['period_hours_counts'].get('24', 0)}；99 缺测：{OCC['period_99_missing_count']} |
| depth | 9999 缺测：{OCC['depth_9999_missing_count']}；原始 0000：{OCC['raw_zero_depth_count']}；原始正数：{OCC['raw_positive_depth_count']} |
| condition | 2 trace：{OCC['condition_code_counts'].get('2', 0)}；3 累计期开始：{OCC['condition_code_counts'].get('3', 0)}；9 condition 缺失：{OCC['condition_code_counts'].get('9', 0)} |
| quality | 1 全部 QC 通过：{OCC['quality_code_counts'].get('1', 0)}；2 suspect：{OCC['quality_code_counts'].get('2', 0)} |
| 累计期闭合 | 未配对开始标记 {OCC['unpaired_accumulation_begin_count']}；配对结束标记 {OCC['paired_accumulation_end_count']} |
| 可用 interval | {OCC['valid_numeric_intervals']} 条 |

condition 3 表示累计期刚开始，降水量要到累计期结束才有；本文件中没有可配对的 condition 4 结束 occurrence，因此不将开始标记的 depth 当作日雨量。condition 2 是 trace，保留标记而不臆造毫米数，也不将它当成精确干日。condition 9 表示 condition 缺失，无法判断数值是否属于有效测量、trace 或累计期边界。quality 2 为 suspect，不用于正式日值。

每个 AA occurrence 都保留时窗、UTC/local 时间、原行号、原始字段和行尾文本。condition 3 按观测时刻作为累计期开始；其他 occurrence 按观测时刻作为报告窗口结束。日值要求 24 小时本地日被质量合格、无重叠的有效 interval 完整覆盖。完整方案优先；多种方案时按时长平方和最小选择较短独立时窗；若同等方案产生不同总量则标为 AMBIGUOUS_OVERLAP。跨午夜窗口不整段归入某一日。没有把重叠 AA 窗口直接求和。

2013 本地日状态：OTHER_UNRESOLVED 211 日，NO_PRECIP_RECORD 152 日，QUALITY_FAIL 2 日；VALID_COMPLETE_DAY 和 VALID_DRY_DAY 均为 0。UTC 日诊断另行保存，仅供比较。

## 4. GSOD、GHCN 50756 与 HLA/D222 验证

### GSOD 同站派生日值参考

同站 GSOD 年文件可访问，覆盖 365 日；其中 159 日有 A–G 降水汇总值。NOAA GSOD readme 规定 PRCP 以 0.01 英寸计，本审计乘 25.4 换算毫米；H 表示不完整（可能有 trace），I 表示当天没有降水观测，99.99 为缺测，这些都未被转为干日。GSOD 与 ISD 属于同一 NOAA 观测链，并非独立证据；GSOD 日按 UTC 汇总，并非中国本地 00:00–24:00 日。

ISD 完整本地日与 GSOD 可比日数为 0，UTC 诊断可比日数也为 0。因此 exact/tolerance agreement、MAE、RMSE、wet/dry agreement 和系统偏移均不可计算。GSOD 只用于查看事件附近是否存在官方 UTC 日降水报告，不用于补入 source hierarchy。

### GHCN-Daily CHM00050756

GHCN 2013 有 {GH['valid_ghcn_2013_days']} 个质量可用 PRCP 日；由于 ISD 没有完整日值，同日可比数为 0，−1/0/+1 日诊断也都没有可比值。Wet/dry、MAE、RMSE、Pearson、Spearman、exact-zero 和累计差均不可计算。DAY_BOUNDARY_QA = {GATE['day_boundary_qa']}，当前证据不足以确认边界，也不足以接受回退。

### HLA/D222 明确数值日

冻结 source table 中 2004–2013 共 {HL['local_D222_numeric_truth_days_2004_2013']} 个 D222 明确数字值，2013 年有 {HL['local_D222_numeric_truth_days_2013']} 天；空白没有视为 0。ISD 可比数为 0，所有误差和 wet/dry 指标不可计算。不能用“同站身份通过”替代实际重叠验证。

## 5. 013 的 111 个缺口日

| 项目 | 天数 |
|---|---:|
| 起始 unresolved | 111 |
| 有原始 AA 时窗/标记覆盖 | {GATE['unresolved_days_with_any_raw_aa_window_information']} |
| 没有 AA 时窗记录 | {111 - GATE['unresolved_days_with_any_raw_aa_window_information']} |
| 可形成 VALID_COMPLETE_DAY / VALID_DRY_DAY | 0 |
| 正式恢复 | 0 |
| 仍 unresolved | 111 |
| NASA 参考为 0 / 大于 0 | {GATE['unresolved_days_nasa_reference_zero']} / {GATE['unresolved_days_nasa_reference_positive']} |

NASA 仅作参考，没有用于补值。触及 AA 窗口的缺口日只有未闭合累计期开始、trace 或缺少 condition 码等信息，不能建立可信数值日雨量。

## 6. 2013 雨量 sanity

| 来源 | 已选日数 | 已选小计 |
|---|---:|---:|
| HLA_LOCAL_D222_CHAIN | {d222.get('days', 0)} | {d222.get('subtotal_mm', 0):.1f} mm |
| NOAA_GHCN_DAILY（50756） | {ghcn.get('days', 0)} | {ghcn.get('subtotal_mm', 0):.1f} mm |
| NOAA_GHCN_DAILY_SECONDARY | {secondary.get('days', 0)} | {secondary.get('subtotal_mm', 0):.1f} mm |
| 013 已选观测日小计 | {CLIMATE['selected_days']} | {CLIMATE['selected_observed_subtotal_mm']:.1f} mm |
| NOAA ISD 可接受本地日值 | {CLIMATE['isd_local_daily_valid_days']} | {CLIMATE['isd_local_daily_available_subtotal_mm']:.1f} mm |

2013 仍缺 {CLIMATE['unresolved_days']} 天；1,149.8 mm 是已选观测日小计，不是全年总量。2004–2012 完整年范围为 {CLIMATE['2004_2012_complete_year_total_min_mm']:.1f}–{CLIMATE['2004_2012_complete_year_total_max_mm']:.1f} mm，NASA 2013 年参考为 {CLIMATE['nasa_2013_total_mm']:.2f} mm；GHCN 50756 的 154 个有效日小计为 {CLIMATE['ghcn50756_2013_available_subtotal_mm']:.1f} mm。2013 sanity 保持 {CLIMATE['status']}。由于全年不完整，不报告年度总量、完整年 wet-day count 或完整年最大日雨量。

已选 2013 中有 {CLIMATE['selected_days_at_or_above_20mm']} 个 ≥20 mm 日期，合计 {CLIMATE['subtotal_of_days_at_or_above_20mm']:.1f} mm（占已选小计 {CLIMATE['top_15_selected_days_share_of_observed_subtotal_pct']:.1f}%）；前 5 日合计 {CLIMATE['top_5_selected_days_sum_mm']:.1f} mm（{CLIMATE['top_5_selected_days_share_of_observed_subtotal_pct']:.1f}%）。最大已选日是 2013-07-30 的 178.2 mm，来源为 HLA/D222 链。

15 个 ≥20 mm 事件日期都在同日或相邻 UTC 日出现至少一个正的 GSOD A–G 值。这只能表示附近有降水报告，不能验证同一中国本地日的数值。尤其 178.2 mm 的 7 月 30 日值，附近 GSOD G 日值为 {july30_gsod}；日报周期不同且量级明显不一致，需保留为未解决的事件幅度差异。

| 日期 | 已选雨量 (mm) | 已选来源 | 附近 GSOD UTC 日参考 |
|---|---:|---|---|
{chr(10).join(event_rows)}

源表按日期唯一，三个 source 日数之和为 254；没有跨源重复累加日值。未因高值偏大而删除或裁剪降雨。ISD 没有可用重建值，无法支持或否定这些已选本地日量；GSOD 日界和 ISD 小时 interval QA 不足以替代逐时闭合。

## 7. 最终 Gate 与执行边界

    ISD_STATION_IDENTITY = PASS
    HLA_ISD_50756_FALLBACK = NOT_ACCEPTED
    HLA_RAIN_PROVENANCE = PARTIAL
    READY_TO_RESUME_010 = FALSE
    starting unresolved = 111
    ISD recovered = 0
    remaining unresolved = 111
    HLA_2013_CLIMATE_SANITY = SUSPICIOUS

未运行 fitting weather、CLI、WGEN、DSSAT 或 PPO；未处理 FQA；未使用 2014 年及之后天气；未改变 012/013 source hierarchy。daily_source_selection_after_014.csv 保留 013 选值，缺口列于 unresolved_dates_after_014.csv。

下一步需要为 111 日取得可追溯的海伦站逐日/逐时官方数值源，或另有严格验证过的同站观测链；当前 ISD 文件只提供标记和不闭合信息。

## 结果文件

- results/hla_random_weather_8seed/014_isd50756_recovery/isd_precip_occurrences_2013.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/isd_interval_selection_2013.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/isd_daily_reconstruction_2013.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/isd_utc_day_diagnostics_2013.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/isd_vs_ghcn50756_daily_qa.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/isd_vs_hla_truth_qa.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/gsod_crosscheck.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/isd_recovery_on_111_days.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/daily_source_selection_after_014.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/unresolved_dates_after_014.csv
- results/hla_random_weather_8seed/014_isd50756_recovery/final_gate.json
"""
    DOCS.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(report, encoding="utf-8")
    (OUT / "findings.md").write_text(report, encoding="utf-8")


def rect(slide, x, y, w, h, color, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    return shape


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = content
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def header(slide, title, subtitle, page):
    rect(slide, 0, 0, 13.333, 0.56, NAVY)
    text(slide, 0.52, 0.08, 10.6, 0.36, title, 22, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    text(slide, 11.15, 0.08, 1.65, 0.34, f"014  |  {page}/5", 10, WHITE, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)
    text(slide, 0.56, 0.69, 12.2, 0.28, subtitle, 10, MUTED)


def card(slide, x, y, w, h, number, label, detail, number_color=TEAL):
    rect(slide, x, y, w, h, PALE, True)
    text(slide, x + 0.15, y + 0.12, w - 0.3, 0.55, number, 29, number_color, True)
    text(slide, x + 0.15, y + 0.72, w - 0.3, 0.35, label, 15, NAVY, True)
    text(slide, x + 0.15, y + 1.1, w - 0.3, h - 1.17, detail, 10, MUTED)


def bullet(slide, x, y, w, content, size=13):
    rect(slide, x, y + 0.1, 0.09, 0.09, TEAL, True)
    text(slide, x + 0.18, y, w - 0.18, 0.55, content, size, INK)


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    header(s, "海伦 2013 降水缺口：同站 ISD 复审", "014 · 只检查 NOAA 50756 同站小时/天气报降水", 1)
    card(s, 0.6, 1.28, 3.85, 1.9, "111", "待恢复日期", "013 的剩余日期全部位于 2013")
    card(s, 4.75, 1.28, 3.85, 1.9, "50756", "同站身份通过", "HAILUN；WMO 50756；小时记录坐标与 GHCN 一致")
    card(s, 8.9, 1.28, 3.85, 1.9, "0 / 111", "本次恢复", "没有质量通过且完整的本地日值", RED)
    rect(s, 0.6, 3.55, 12.15, 2.55, PALE, True)
    text(s, 0.9, 3.82, 5.6, 0.4, "为什么转向同站 ISD", 20, NAVY, True)
    bullet(s, 0.95, 4.35, 5.45, "013 中 8 个远处 GHCN 候选均未通过冻结的 80% 事件一致率门槛。")
    bullet(s, 0.95, 5.0, 5.45, "014 只审计同一 Hailun / WMO 50756 的 NOAA ISD 2013 文件。")
    text(s, 6.85, 4.02, 5.45, 0.4, "本轮门禁", 15, TEAL, True)
    text(s, 6.85, 4.52, 5.45, 1.1, "ISD fallback  NOT_ACCEPTED\nRain provenance  PARTIAL\nREADY_TO_RESUME_010  FALSE", 17, RED, True)
    text(s, 0.62, 6.42, 12.05, 0.38, "没有运行 fitting weather、CLI、WGEN、DSSAT、PPO 或 FQA。", 10, MUTED)

    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    header(s, "AA 降水字段：标记多，数值 interval 不成立", "occurrence 留存 condition、quality、原始值、时窗和行号", 2)
    card(s, 0.6, 1.3, 2.85, 1.55, "491", "AA occurrences", "2,914 条 ISD 报文中")
    card(s, 3.65, 1.3, 2.85, 1.55, "374", "累计期开始", "没有配对的结束码 4")
    card(s, 6.7, 1.3, 2.85, 1.55, "50 + 67", "trace + condition 缺失", "trace 未换算为任意毫米数")
    card(s, 9.75, 1.3, 2.85, 1.55, "0", "有效数值 interval", "另有 2 条 quality=2 suspect", RED)
    rect(s, 0.65, 3.2, 5.95, 2.45, PALE, True)
    text(s, 0.95, 3.48, 5.3, 0.38, "官方字段语义", 19, NAVY, True)
    bullet(s, 1.0, 4.0, 5.2, "depth 原始值 ÷ 10 = mm；period=99、depth=9999 缺测。")
    bullet(s, 1.0, 4.58, 5.2, "condition=3 是累计期开始；结束报告之前雨量不可用。")
    bullet(s, 1.0, 5.16, 5.2, "condition=2 为 trace；condition=9 是 condition 缺失。")
    rect(s, 6.95, 3.2, 5.7, 2.45, PALE, True)
    text(s, 7.25, 3.48, 5.05, 0.38, "不重复累计规则", 19, NAVY, True)
    bullet(s, 7.3, 4.0, 4.95, "本地日必须由完整、质量通过的 interval 链覆盖。")
    bullet(s, 7.3, 4.58, 4.95, "区间只选不重叠集合；完整覆盖优先，时长较短优先。")
    bullet(s, 7.3, 5.16, 4.95, "跨午夜、未闭合标记或多解总量冲突不自动补值。")
    text(s, 0.68, 6.05, 12.0, 0.45, "没有 AA 时窗被直接相加；VALID_COMPLETE_DAY 和 VALID_DRY_DAY 均为 0。", 14, RED, True)

    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    header(s, "同站身份通过，日值一致性无法计算", "对照资料存在，但 ISD 没有完整本地日供配对", 3)
    rect(s, 0.65, 1.35, 5.6, 4.65, PALE, True)
    text(s, 0.95, 1.65, 5.0, 0.4, "身份核验", 20, NAVY, True)
    bullet(s, 1.0, 2.2, 4.95, "ISD 507560-99999：HAILUN / China；959 条 SYNOP 报文含 WMO 50756。")
    bullet(s, 1.0, 3.0, 4.95, "GHCN CHM00050756：WMO 50756；47.433°N, 126.967°E, 240 m。")
    bullet(s, 1.0, 3.8, 4.95, "2013 ISD 记录主位置与 GHCN 目录完全吻合。")
    text(s, 1.0, 4.72, 4.9, 0.8, "ISD 静态站史为 47.450°N, 126.867°E, 248 m；该元数据差异留档。", 12, MUTED)
    rect(s, 6.65, 1.35, 6.0, 4.65, PALE, True)
    text(s, 6.95, 1.65, 5.4, 0.4, "重叠验证", 20, NAVY, True)
    text(s, 7.0, 2.25, 5.25, 0.5, f"GHCN：{GH['valid_ghcn_2013_days']} 天有效；同日 ISD pair = 0", 14, INK)
    text(s, 7.0, 2.95, 5.25, 0.5, f"D222：{HL['local_D222_numeric_truth_days_2013']} 个数值日；ISD pair = 0", 14, INK)
    text(s, 7.0, 3.65, 5.25, 0.5, "GSOD：159 天可用 A–G；ISD pair = 0", 14, INK)
    text(s, 7.0, 4.45, 5.15, 1.05, "wet/dry、MAE、RMSE、相关系数和 −1/0/+1 日诊断均不可计算。\nDAY_BOUNDARY_QA = NOT_ASSESSABLE", 15, RED, True)
    text(s, 0.68, 6.27, 12.0, 0.4, "GSOD 来源于同一观测链且按 UTC 日汇总，仅作附近日期参考。", 11, MUTED)

    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    header(s, "111 个缺口日：没有可安全恢复的数值", "16 日触及 AA 窗口/标记，但都不是合格日值", 4)
    card(s, 0.65, 1.32, 3.75, 1.7, "16", "有 AA 窗口信息", "只含未闭合标记、trace 或缺失条件")
    card(s, 4.8, 1.32, 3.75, 1.7, "0", "有效完整本地日", "覆盖、质量和日界都须通过", RED)
    card(s, 8.95, 1.32, 3.75, 1.7, "111", "仍未解决", "hierarchy 没有新增值", RED)
    rect(s, 0.75, 3.45, 11.85, 2.0, PALE, True)
    text(s, 1.05, 3.72, 5.0, 0.4, "缺口分类", 18, NAVY, True)
    text(s, 1.1, 4.25, 4.9, 0.62, "95 日：无 AA 窗口记录\n16 日：只有不完整字段或标记", 15, INK)
    text(s, 6.65, 3.72, 5.2, 0.4, "NASA 仅作参考", 18, NAVY, True)
    text(s, 6.7, 4.25, 5.0, 0.62, "35 日为 0\n76 日大于 0\n没有 NASA 补值", 15, INK)
    text(s, 0.82, 5.95, 11.6, 0.64, "同日 GHCN、D222、GSOD 配对数均为 0；门禁维持 NOT_ACCEPTED。", 15, RED, True)

    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    header(s, "2013 小计仍可疑，010 暂不恢复", "ISD 没有验证或替换高值；保留原值和未解决状态", 5)
    rect(s, 0.65, 1.3, 6.0, 2.15, PALE, True)
    text(s, 0.95, 1.58, 5.45, 0.46, "观测日小计  1,149.8 mm", 24, NAVY, True)
    text(s, 0.98, 2.18, 5.4, 0.86, "254 / 365 天已选；111 天仍缺\n2004–2012 完整年范围 350.0–677.3 mm\nNASA 2013 参考 746.07 mm", 14, INK)
    rect(s, 6.95, 1.3, 5.7, 2.15, PALE, True)
    text(s, 7.25, 1.58, 5.0, 0.46, "高值保留原样", 21, NAVY, True)
    text(s, 7.28, 2.18, 5.0, 0.86, "15 日 ≥20 mm：641.6 mm\n前 5 日：388.9 mm（33.8%）\n最大 178.2 mm：7 月 30 日，D222 链", 14, INK)
    rect(s, 0.65, 3.78, 6.0, 2.05, PALE, True)
    text(s, 0.95, 4.05, 5.4, 0.38, "GSOD 附近事件参考", 18, NAVY, True)
    text(s, 0.98, 4.55, 5.35, 0.95, "高值日期附近有正的 UTC 日报告，但日界不同且来自同一观测链。7/30 的 178.2 mm 与附近 GSOD 日值幅度差异很大，不能据此校正。", 12, INK)
    rect(s, 6.95, 3.78, 5.7, 2.05, NAVY, True)
    text(s, 7.25, 4.05, 5.0, 0.38, "最终状态", 18, WHITE, True)
    text(s, 7.28, 4.55, 5.0, 1.0, "ISD fallback    NOT_ACCEPTED\nRain provenance    PARTIAL\nREADY_TO_RESUME_010    FALSE", 16, WHITE, True)
    text(s, 0.72, 6.25, 11.9, 0.4, "下一步：寻找 111 日可追溯的海伦站官方逐日/逐时数值记录；不降 QA、不删高值。", 11, MUTED)

    DECK.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DECK)


if __name__ == "__main__":
    build_report()
    make_deck()
    print(json.dumps({"report": str(REPORT), "pptx": str(DECK), "slides": 5}, ensure_ascii=False))
