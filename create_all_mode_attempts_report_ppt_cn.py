from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path.cwd()
OUT_DIR = ROOT / "figures_hl"
ASSET_DIR = OUT_DIR / "all_mode_attempts_report_assets"
OUT_DIR.mkdir(exist_ok=True)
ASSET_DIR.mkdir(exist_ok=True)
PPT_PATH = OUT_DIR / "all_mode_water_nitrogen_attempts_report.pptx"

FONT = "Microsoft YaHei"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(90, 90, 90)
BLUE = RGBColor(68, 114, 196)
LIGHT_BLUE = RGBColor(221, 235, 247)
WHITE = RGBColor(255, 255, 255)


def read_csv(relative_path: str) -> pd.DataFrame:
    df = pd.read_csv(ROOT / relative_path, keep_default_na=False)
    if "agent" in df.columns:
        df["agent"] = df["agent"].replace("", "null")
    return df


penalty_df = read_csv("output_hl/rain_scaling_train/HL_scale0p6_penalty_scan_50k/penalty_scan_summary.csv")
safe_df = read_csv("output_hl/rain_scaling_train/HL_scale0p6_safe_penalty_scan_50k/safe_penalty_scan_summary.csv")
budget_df = read_csv("output_hl/rain_scaling_train/HL_scale0p6_budget_scan_50k/budget_scan_summary.csv")
budget_100_df = read_csv("output_hl/rain_scaling_train/HL_scale0p6_budget220_100_100k/diagnose_ppo/all_water_stress_summary.csv")


def select_row(df: pd.DataFrame, combo: str | None = None, agent: str = "ppo") -> pd.Series:
    sub = df[df["agent"].astype(str).eq(agent)]
    if combo is not None and "combo" in sub.columns:
        sub = sub[sub["combo"].eq(combo)]
    if sub.empty:
        raise ValueError(f"Missing row: combo={combo}, agent={agent}")
    return sub.iloc[0]


expert = select_row(budget_df, "budget220_100_daily20_5_pen20", "expert")
null = select_row(budget_df, "budget220_100_daily20_5_pen20", "null")
budget_50 = select_row(budget_df, "budget220_100_daily20_5_pen20", "ppo")
budget_100 = select_row(budget_100_df, None, "ppo")


def save_attempt_chart(path: Path, df: pd.DataFrame, title: str, combos: list[tuple[str, str]]) -> None:
    rows = []
    for label, combo in combos:
        r = select_row(df, combo, "ppo")
        rows.append((label, float(r["max_grnwt"]), float(r["total_anfer"]), float(r["total_amir"])))

    labels = [r[0] for r in rows]
    yields = [r[1] for r in rows]
    ns = [r[2] for r in rows]
    waters = [r[3] for r in rows]

    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    fig, ax1 = plt.subplots(figsize=(9, 4.6), dpi=160)
    ax1.bar(labels, yields, color=["#4472C4", "#5B9BD5", "#70AD47", "#A9D18E"][: len(labels)], width=0.55)
    ax1.set_ylabel("产量 max_grnwt (kg/ha)")
    ax1.set_title(title, color="black")
    ax1.grid(axis="y", alpha=0.25)
    ax1.tick_params(axis="x", rotation=12)
    for i, value in enumerate(yields):
        ax1.text(i, value + max(yields) * 0.02, f"{value:.0f}", ha="center", fontsize=8)

    ax2 = ax1.twinx()
    ax2.plot(labels, ns, marker="o", color="#C00000", label="总施氮 kg/ha")
    ax2.plot(labels, waters, marker="s", color="#00A2E8", label="总灌水 mm")
    ax2.set_ylabel("投入总量")
    ax2.legend(loc="upper right", fontsize=8)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_final_chart(path: Path) -> None:
    labels = ["Null", "Expert", "PPO 50k", "PPO 100k"]
    vals = [
        (float(null["max_grnwt"]), float(null["total_anfer"]), float(null["total_amir"])),
        (float(expert["max_grnwt"]), float(expert["total_anfer"]), float(expert["total_amir"])),
        (float(budget_50["max_grnwt"]), float(budget_50["total_anfer"]), float(budget_50["total_amir"])),
        (float(budget_100["max_grnwt"]), float(budget_100["total_anfer"]), float(budget_100["total_amir"])),
    ]
    yields = [v[0] for v in vals]
    ns = [v[1] for v in vals]
    waters = [v[2] for v in vals]

    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    fig, ax1 = plt.subplots(figsize=(9, 4.6), dpi=160)
    ax1.bar(labels, yields, color=["#A5A5A5", "#4472C4", "#70AD47", "#2F5597"], width=0.55)
    ax1.set_ylabel("产量 max_grnwt (kg/ha)")
    ax1.set_title("最终对比：HL scale0.6", color="black")
    ax1.grid(axis="y", alpha=0.25)
    for i, value in enumerate(yields):
        ax1.text(i, value + max(yields) * 0.02, f"{value:.0f}", ha="center", fontsize=8)

    ax2 = ax1.twinx()
    ax2.plot(labels, ns, marker="o", color="#C00000", label="总施氮 kg/ha")
    ax2.plot(labels, waters, marker="s", color="#00A2E8", label="总灌水 mm")
    ax2.set_ylabel("投入总量")
    ax2.legend(loc="upper left", fontsize=8)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


penalty_chart = ASSET_DIR / "penalty_scan_chart.png"
safe_chart = ASSET_DIR / "safe_scan_chart.png"
budget_chart = ASSET_DIR / "budget_scan_chart.png"
final_chart = ASSET_DIR / "final_comparison_chart.png"

save_attempt_chart(
    penalty_chart,
    penalty_df,
    "尝试1：只加 reward 惩罚",
    [
        ("pen15 ns20", "pen15_icost15_ns20_nlim300_nex5_ilim120_iex10"),
        ("pen15 ns50", "pen15_icost15_ns50_nlim250_nex10_ilim80_iex20"),
        ("pen20 ns50", "pen20_icost20_ns50_nlim220_nex15_ilim80_iex30"),
    ],
)
save_attempt_chart(
    safe_chart,
    safe_df,
    "尝试2：安全日动作上限",
    [
        ("safe60_20", "safe60_20_pen15_icost15_ns20_nlim300_nex5_ilim120_iex10"),
        ("safe40_15", "safe40_15_pen15_icost15_ns50_nlim250_nex10_ilim80_iex20"),
        ("safe30_10", "safe30_10_pen20_icost20_ns50_nlim220_nex15_ilim80_iex30"),
        ("safe20_5", "safe20_5_pen20_icost20_ns80_nlim180_nex20_ilim60_iex40"),
    ],
)
save_attempt_chart(
    budget_chart,
    budget_df,
    "尝试3：日上限 + 季节总预算",
    [
        ("budget180_60", "budget180_60_daily20_5_pen20"),
        ("budget220_100", "budget220_100_daily20_5_pen20"),
        ("budget300_150", "budget300_150_daily30_10_pen15"),
    ],
)
save_final_chart(final_chart)


def set_run_font(run, size: int | None = None, bold: bool = False, color: RGBColor = BLACK) -> None:
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        child = rpr.find(qn(tag))
        if child is None:
            child = etree.Element(qn(tag))
            rpr.append(child)
        child.set("typeface", FONT)


def set_paragraph_text(paragraph, text: str, size: int, bold: bool = False, color: RGBColor = BLACK) -> None:
    paragraph.text = ""
    run = paragraph.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color)


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.55))
    set_paragraph_text(box.text_frame.paragraphs[0], title, 24, True, BLACK)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.48), Inches(0.82), Inches(12.3), Inches(0.35))
        set_paragraph_text(sub.text_frame.paragraphs[0], subtitle, 11, False, GRAY)


def add_bullets(slide, bullets: list[str | tuple[str, int]], x=0.65, y=1.25, w=12.0, h=5.8, size=15) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in bullets:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(5)
        set_paragraph_text(p, text, size if level == 0 else size - 2, False, BLACK)


def add_table(slide, data: list[list[str]], x, y, w, h, size=9) -> None:
    rows = len(data)
    cols = len(data[0])
    table = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(data[r][c])
            set_run_font(run, size=size, bold=(r == 0), color=WHITE if r == 0 else BLACK)
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (LIGHT_BLUE if r % 2 == 0 else WHITE)


def add_chart(slide, path: Path, x=6.4, y=1.25, w=6.1, h=4.3) -> None:
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_footer(slide, number: int) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(7.13), Inches(12.4), Inches(0.2))
    p = box.text_frame.paragraphs[0]
    set_paragraph_text(p, f"水氮联合优化 all 模式实验记录 | {number}", 8, False, GRAY)
    p.alignment = PP_ALIGN.RIGHT


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def new_slide():
    return prs.slides.add_slide(prs.slide_layouts[6])


slides = []

s = new_slide(); slides.append(s)
box = s.shapes.add_textbox(Inches(0.75), Inches(1.45), Inches(11.8), Inches(0.9))
set_paragraph_text(box.text_frame.paragraphs[0], "水氮联合优化 all 模式新尝试汇报", 34, True, BLACK)
sub = s.shapes.add_textbox(Inches(0.78), Inches(2.45), Inches(11.8), Inches(0.7))
set_paragraph_text(sub.text_frame.paragraphs[0], "从 reward 惩罚、日动作上限，到季节总预算约束的完整实验记录", 18, False, BLACK)
add_table(
    s,
    [
        ["实验场景", "核心结论"],
        ["海伦站 2007 玉米；历史天气作为完美预报；降雨缩放到 60% 构造干旱调试场景。", "单纯 reward 惩罚不能解决过量动作；日上限改善动作尺度；季节预算首次得到可解释策略。"],
    ],
    0.8, 4.25, 11.8, 1.25, 11,
)

s = new_slide(); slides.append(s); add_title(s, "汇报内容与实验线索")
add_bullets(s, [
    "导师要求：最终模型需要同时考虑产量、氮回收/氮素利用、施肥量和灌水量。",
    "前期发现：五站点原始年份多数为雨养或弱灌溉需求，直接训练灌溉优化空间有限。",
    "为走通天气预报驱动的水氮管理优化，先用海伦站降雨缩放 0.6 构造干旱调试场景。",
    "之后依次尝试：reward 惩罚增强、日安全动作上限、季节总预算 wrapper。",
    "每一轮都保留 Null / Expert / PPO 诊断，指标包括 yield、total_anfer、total_amir、swfac/nstres 胁迫天数。",
])

s = new_slide(); slides.append(s); add_title(s, "背景：为什么要从施肥扩展到水氮联合优化")
add_bullets(s, [
    "原始 fertilization_reward 只处理施肥：reward = trnu × coef - penality × anfer。",
    "导师要求不是靠经验指定 coef/penality，而是通过参数扫描寻找更优组合。",
    "后续论文目标更大：能够基于天气预报进行水氮联合管理优化。",
    "当前先使用历史天气作为完美预报来走通方法框架。",
    "因此 all mode 不只是加一个灌溉动作，而是要保证 PPO 在有限水氮资源下学会何时施肥、何时灌水。",
])

s = new_slide(); slides.append(s); add_title(s, "诊断指标：如何判断 PPO 策略是否合理")
add_table(s, [
    ["指标", "含义", "判断方向"],
    ["max_grnwt", "最终/最高籽粒产量 proxy", "越高越好"],
    ["total_anfer", "全季总施氮量 kg/ha", "越少越好，但不能牺牲产量"],
    ["total_amir", "全季总灌水量 mm", "越少越好，但不能造成水分胁迫"],
    ["swfac", "水分胁迫指标；gym 中已做 1 - 原值", "越低越好，0 近似无胁迫"],
    ["nstres", "氮胁迫指标；gym 中已做 1 - 原值", "越低越好，0 近似无胁迫"],
    ["reward", "当前 reward 函数下累计得分", "辅助判断，不能单独替代农学指标"],
], 0.6, 1.25, 12.1, 3.3, 10)
add_bullets(s, ["注意：maize 的 swfac/nstres 在 gym 后处理中做了 1 - 原值转换，因此诊断表中数值越大表示胁迫越强。"], y=5.25, h=0.8, size=14)

s = new_slide(); slides.append(s); add_title(s, "调试场景：海伦站降雨缩放 0.6")
add_bullets(s, [
    "原始五站点年份多为降雨充足或接近平衡，灌溉优化空间不明显。",
    "为了测试水氮联合优化能力，构造海伦站 HL scale0.6：把降雨缩放到 60%。",
    "该情景 PRCP = 207.3 mm，ETCP = 496.0 mm，水分收支差约 -288.7 mm。",
    "Null 模式产量极低，说明该情景确实需要管理干预。",
    "该情景不是最终实证站点结论，而是用于调试模型是否会因缺水而灌水的方法场景。",
], w=6.4)
add_table(s, [
    ["Agent", "Yield", "N kg/ha", "Irr. mm", "swfac days", "nstres days"],
    ["Null", "469", "0", "0", "0", "114"],
    ["Expert", "6276", "165", "30", "18", "34"],
], 7.1, 1.45, 5.3, 1.55, 10)

s = new_slide(); slides.append(s); add_title(s, "所有新尝试的逻辑路线")
add_table(s, [
    ["阶段", "做法", "中间讨论结论", "结果"],
    ["1", "增强 reward 惩罚", "只调 reward 可能不足以约束 PPO 动作空间", "失败：极端过量水肥"],
    ["2", "日安全动作上限", "把 PPO 动作映射到合理日施用范围", "改善：过量下降，但仍每天少量累积"],
    ["3", "季节总预算", "预算用完后动作强制为 0，逼 PPO 学 timing", "成功：出现可解释策略"],
    ["4", "100k 确认", "验证最佳预算策略是否稳定", "稳定：接近 50k 结果"],
], 0.55, 1.25, 12.2, 3.0, 10)
add_bullets(s, ["核心判断：PPO 的问题不是完全不会优化，而是原环境动作空间太自由。先把农学可行性写进动作层，再让 reward 做偏好排序，效果更稳定。"], y=4.8, h=1.1, size=15)

s = new_slide(); slides.append(s); add_title(s, "尝试 1：只增加 reward 惩罚")
add_bullets(s, [
    "保留原动作空间，不限制 PPO 每天施多少氮/灌多少水。",
    "新增总施肥/总灌水超量惩罚：超过设定 limit 后按 cost 扣分。",
    "新增无水分胁迫灌水惩罚：如果 swfac 低于阈值仍灌水，则额外扣分。",
    "参数例：pen20_icost20_ns50_nlim220_nex15_ilim80_iex30。",
    ("pen20：施肥惩罚系数 20；icost20：灌水基础成本 20；ns50：无水分胁迫灌水额外成本 50。", 1),
    ("nlim220/nex15：总施氮超过 220 后按 15 扣分；ilim80/iex30：总灌水超过 80 后按 30 扣分。", 1),
], size=14)

s = new_slide(); slides.append(s); add_title(s, "尝试 1 结果：reward 惩罚不足以防止过量动作")
add_bullets(s, [
    "三组 50k 扫描全部失败。",
    "PPO 总施氮约 13,600 kg/ha，总灌水约 3,500 mm。",
    "产量约 4,866-4,868 kg/ha，低于 Expert 的 6,276 kg/ha。",
    "解释：PPO 初始探索可以产生极端大动作；长期累计惩罚不能及时阻止动作空间灾难。",
    "中间结论：必须先约束动作可行域，不能只靠 reward 扣分。",
], w=5.7, size=14)
add_chart(s, penalty_chart)

s = new_slide(); slides.append(s); add_title(s, "尝试 2：安全日动作上限 Safe Action Wrapper")
add_bullets(s, [
    "新增 SafeActionGymDssatWrapper，不修改原始 DSSAT 环境。",
    "关键修正：不是先按原动作范围反归一化再 clip，而是直接把 PPO 的 [-1,1] 映射到 [0, safe_cap]。",
    "命名解释：safe20_5 表示每天最多施氮 20 kg/ha、每天最多灌水 5 mm。",
    "safe60_20 表示每天最多施氮 60 kg/ha、每天最多灌水 20 mm。",
    "局限：日上限只能限制每天最多多少，不能限制全季总共多少。",
], size=15)

s = new_slide(); slides.append(s); add_title(s, "尝试 2 结果：日上限显著改善，但仍然过量")
add_bullets(s, [
    "safe20_5 是日上限扫描中最好的一组。",
    "50k: yield ≈ 7456，total_anfer ≈ 1355，total_amir ≈ 352。",
    "100k follow-up: yield ≈ 7457，total_anfer ≈ 1350，total_amir ≈ 351。",
    "产量高，但水氮总量仍远高于 Expert。",
    "中间结论：步数增加不是关键，必须限制季节总预算。",
], w=5.7, size=14)
add_chart(s, safe_chart)

s = new_slide(); slides.append(s); add_title(s, "尝试 3：季节总预算 Budgeted Safe Action Wrapper")
add_bullets(s, [
    "新增 BudgetedSafeActionGymDssatWrapper：日动作上限 + 季节总预算。",
    "每次 reset 时清零累计水氮用量。",
    "每次 step 先按日上限映射动作，再按剩余季节预算二次裁剪。",
    "预算用完后，对应动作自动变成 0。",
    "这会迫使 PPO 学会什么时候使用有限水氮资源，而不是每天少量累积。",
], w=7.0, size=15)
add_table(s, [
    ["命名", "含义"],
    ["budget220_100", "全季施氮预算 220 kg/ha；全季灌水预算 100 mm"],
    ["daily20_5", "每天施氮上限 20 kg/ha；每天灌水上限 5 mm"],
    ["pen20", "施肥 reward 中对 anfer 的惩罚系数为 20"],
], 7.45, 1.4, 5.0, 2.1, 10)

s = new_slide(); slides.append(s); add_title(s, "尝试 3 结果：季节预算首次得到可解释策略")
add_bullets(s, [
    "budget180_60：最保守，水氮接近 Expert，但产量略低。",
    "budget220_100：当前最佳折中，产量高于 Expert，水氮增加仍在可解释范围内。",
    "budget300_150：预算更宽，水氮更多，但产量并未继续明显提高。",
    "结论：季节预算不是越大越好，过大预算会带来额外投入但收益有限。",
], w=5.7, size=14)
add_chart(s, budget_chart)

s = new_slide(); slides.append(s); add_title(s, "预算扫描详细结果（HL scale0.6, 50k）")
add_table(s, [
    ["策略", "产量 kg/ha", "总氮 kg/ha", "总灌水 mm", "swfac天数", "nstres天数", "解释"],
    ["Expert", "6276", "165", "30", "18", "34", "固定专家策略"],
    ["budget180_60", "6165", "166", "51", "21", "32", "保守，接近专家用氮"],
    ["budget220_100", "6729", "206", "91", "17", "23", "最佳折中，产量提升"],
    ["budget300_150", "6695", "278", "133", "18", "16", "投入增加但收益不增"],
], 0.45, 1.25, 12.45, 2.65, 10)
add_bullets(s, [
    "为什么 budget220_100 最合适：相对 Expert，产量提高约 453 kg/ha，施氮增加约 41 kg/ha，灌水增加约 61 mm。",
    "为什么 budget180_60 也有价值：几乎保持专家施氮水平，灌水略增，产量只低约 111 kg/ha，可作为保守节水节肥对照。",
], y=4.55, h=1.5, size=14)

s = new_slide(); slides.append(s); add_title(s, "100k 确认：budget220_100 结果稳定")
add_bullets(s, [
    "对 50k 最优的 budget220_100 单独扩展到 100k。",
    "100k PPO：yield ≈ 6669，total_anfer ≈ 202，total_amir ≈ 92。",
    "50k PPO：yield ≈ 6729，total_anfer ≈ 206，total_amir ≈ 91。",
    "两次结果接近，说明当前策略不是短步数偶然结果。",
    "该配置可以作为后续论文方法调试的主配置。",
], w=5.7, size=14)
add_chart(s, final_chart)

s = new_slide(); slides.append(s); add_title(s, "三类尝试的对比总结")
add_table(s, [
    ["尝试", "代表配置", "PPO产量", "PPO总氮", "PPO总灌水", "结论"],
    ["Reward惩罚", "pen20 ns50", "4866", "13650", "3524", "失败，动作空间灾难"],
    ["日上限", "safe20_5", "7456", "1355", "352", "改善但仍过量"],
    ["季节预算", "budget220_100", "6729", "206", "91", "当前最佳"],
    ["100k确认", "budget220_100", "6669", "202", "92", "稳定复现"],
], 0.55, 1.25, 12.2, 2.7, 11)
add_bullets(s, ["中间讨论形成的关键判断：PPO 并不是完全不会学习，而是原始动作空间不符合农学管理约束。先把可行域约束清楚，再让 reward 做偏好排序，效果明显更好。"], y=4.55, h=1.1, size=15)

s = new_slide(); slides.append(s); add_title(s, "当前形成的方法框架")
add_bullets(s, [
    "1. 用 DSSAT 构建作物-土壤-天气响应环境。",
    "2. 当前先使用历史天气作为完美天气预报，未来可替换为真实天气预报。",
    "3. PPO 观察作物状态、水分状态、氮胁迫等指标。",
    "4. PPO 输出水氮动作。",
    "5. BudgetedSafeActionWrapper 将动作限制在日上限和季节预算内。",
    "6. DSSAT 返回产量、水分胁迫、氮胁迫、trnu 等状态，用于训练与诊断。",
], size=15)

s = new_slide(); slides.append(s); add_title(s, "代码与结果文件记录")
add_table(s, [
    ["文件/目录", "作用"],
    ["sb3_safe_action_wrapper.py", "SafeAction 与 BudgetedSafeAction wrapper"],
    ["train_hl_all_multisite_safe.py", "日动作上限训练脚本"],
    ["train_hl_all_multisite_budget.py", "日上限 + 季节预算训练脚本"],
    ["run_hl_scale0p6_safe_penalty_scan.py", "safe 动作上限扫描"],
    ["run_hl_scale0p6_budget_scan.py", "季节预算扫描"],
    ["diagnose_all_water_stress_sites.py", "支持 PPO + safe/budget wrapper 的诊断"],
    ["output_hl/rain_scaling_train/...", "训练、诊断、summary 和 findings 输出"],
], 0.45, 1.2, 12.45, 3.55, 10)
add_bullets(s, ["GitHub 记录：bd155ff Add safe action all-mode scan；65da954 Add seasonal budget all-mode scan。"], y=5.3, h=0.8, size=14)

s = new_slide(); slides.append(s); add_title(s, "当前结果的限制与需要说明的点")
add_bullets(s, [
    "HL scale0.6 是降雨缩放构造的干旱调试场景，不等同于真实干旱年份验证。",
    "目前主要是单站点、单年份、单随机种子的流程验证，后续需要多 seed 和多站点。",
    "budget220_100 的预算值目前是调试得到的工程约束，需要结合农学经验或文献进一步论证。",
    "turfac 尚未安全暴露，当前主要使用 swfac/nstres 作为决策诊断指标。",
    "最终论文中应把方法可行性和实证推广验证分开叙述。",
], size=15)

s = new_slide(); slides.append(s); add_title(s, "下一步建议")
add_bullets(s, [
    "1. 对 budget220_100 输出逐日 action trace 图，检查施肥/灌水发生在哪些 DAP 和作物阶段。",
    "2. 做 3-5 个随机种子重复，确认结果均值和方差。",
    "3. 在真实干旱年份或用户搜集的新天气数据上测试，而不只依赖 rain scale0.6。",
    "4. 将预算约束推广到其他站点，比较不同水分需求情景下 PPO 是否自动调整灌水。",
    "5. 为论文写法准备：历史天气 = 完美预报的方法说明，并预留真实天气预报接口。",
], size=15)

s = new_slide(); slides.append(s); add_title(s, "阶段性结论")
add_bullets(s, [
    "单纯 reward 惩罚不能解决 all mode 的过量动作问题。",
    "日动作上限能显著改善动作尺度，但仍无法控制季节总量。",
    "季节总预算 wrapper 是关键突破：它让 PPO 学习有限水氮资源下的调度策略。",
    "当前最佳配置：budget220_100_daily20_5_pen20。",
    "在 HL scale0.6 情景下，PPO 100k 产量约 6669 kg/ha，高于 Expert 的 6276 kg/ha；总氮约 202 kg/ha，总灌水约 92 mm。",
    "这套流程已经可以作为水氮联合优化方法的原型，后续重点转向多 seed、多站点和真实干旱年份验证。",
], size=16)

for i, slide in enumerate(slides, start=1):
    add_footer(slide, i)

prs.save(PPT_PATH)
print(PPT_PATH)
