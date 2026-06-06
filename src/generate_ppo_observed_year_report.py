from __future__ import annotations

from datetime import datetime
from pathlib import Path
import shutil

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ppo_safe_rendering import PROJECT_ROOT, load_yaml


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_ppo_observed_years.yaml"
DOCS = PROJECT_ROOT / "docs"
PPO_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "ppo_observed_years"
REPORT_MD = DOCS / "2026-06-05_ppo_observed_year_training_framework.md"
REPORT_PPTX = DOCS / "2026-06-05_ppo_observed_year_training_framework.pptx"
BLUE = RGBColor(68, 114, 196)
LIGHT_BLUE = RGBColor(232, 238, 249)


def read_csv(path: Path) -> pd.DataFrame:
    if path.exists():
        return pd.read_csv(path)
    return pd.DataFrame()


def rel(path: Path) -> str:
    return str(path.relative_to(PROJECT_ROOT)).replace("\\", "/")


def md_table(df: pd.DataFrame, max_rows: int = 30) -> str:
    if df.empty:
        return "_No data._"
    return df.head(max_rows).to_markdown(index=False)


def build_render_check() -> Path:
    rows = []
    for template in sorted((PPO_ROOT / "rendered_inputs").glob("**/*.jinja2")):
        text = template.read_text(encoding="utf-8", errors="replace")
        parts = template.parts
        try:
            idx = parts.index("rendered_inputs")
            station = parts[idx + 1]
            year = parts[idx + 2]
            run_tag = parts[idx + 3]
        except Exception:
            station, year, run_tag = "", "", ""
        checks = {
            "template_exists": template.exists(),
            "irrigation_section_present": "*IRRIGATION AND WATER MANAGEMENT" in text,
            "fertilizer_section_present": "*FERTILIZERS (INORGANIC)" in text,
            "simulation_controls_present": "*SIMULATION CONTROLS" in text,
            "planting_section_present": "*PLANTING DETAILS" in text,
            "mi_mf_enabled": " 1 1 0 0 1 1 1" in text or " 1  1  0  0  1  1  1" in text,
        }
        for item, ok in checks.items():
            rows.append(
                {
                    "station": station,
                    "year": year,
                    "run_tag": run_tag,
                    "file_path": rel(template),
                    "check_item": item,
                    "status": "pass" if ok else "fail",
                }
            )
    out = PPO_ROOT / "evaluation" / "rendered_input_check.csv"
    out.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out


def write_markdown() -> None:
    config = load_yaml(CONFIG_PATH)
    plan = read_csv(PPO_ROOT / "configs" / "ppo_observed_year_experiment_plan.csv")
    smoke = read_csv(PPO_ROOT / "smoke_checks" / "pretrain_smoke_check_summary.csv")
    evaluation = read_csv(PPO_ROOT / "evaluation" / "ppo_evaluation_summary.csv")
    selection = read_csv(PPO_ROOT / "strategy_selection" / "best_policy_by_site.csv")
    render_check_path = build_render_check()
    render_check = read_csv(render_check_path)
    model_files = sorted((PPO_ROOT / "models").glob("**/*.zip"))
    daily_files = sorted((PPO_ROOT / "daily_outputs").glob("**/*.csv"))
    figure_files = sorted((PPO_ROOT / "figures").glob("**/*.png"))
    scripts = [
        "src/ppo_experiment_plan.py",
        "src/ppo_safe_rendering.py",
        "src/ppo_train.py",
        "src/ppo_evaluate.py",
        "src/ppo_plot_results.py",
        "src/ppo_strategy_selection.py",
        "experiments/ppo_observed_years/config_ppo_observed_years.yaml",
        "experiments/ppo_observed_years/generate_experiment_plan.py",
        "experiments/ppo_observed_years/train_one_policy.py",
        "experiments/ppo_observed_years/evaluate_one_policy.py",
        "experiments/ppo_observed_years/run_debug_hla_2007.py",
        "experiments/ppo_observed_years/run_all_trainings_DISABLED_BY_DEFAULT.py",
        "experiments/ppo_observed_years/run_all_evaluations_DISABLED_BY_DEFAULT.py",
    ]
    debug = config["debug"]
    final_hparams = {k: v for k, v in config["ppo"].items() if k != "total_timesteps_full"}
    need_confirm = [k for k, v in final_hparams.items() if v is None]
    status_text = "success" if not evaluation.empty and (evaluation["run_status"] == "ok").all() else "incomplete"
    lines = [
        "# PPO observed-year leave-one training framework",
        "",
        f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "## 1. 本阶段目标",
        "基于已经跑通的 observed-year smoke test，生成可逐站点调试的 PPO 训练、验证、绘图和策略选择框架。本阶段不批量训练全部 PPO，只运行 HLA 2007 seed0 小步数 debug PPO。",
        "",
        "## 2. 为什么现在可以进入 PPO 框架生成",
        "上一阶段 smoke test timeout 已排查完成，29 个原失败案例全部重跑成功。LCA/SYA/YCA 的问题来自临时 rendered DSSAT management section 和跨年份静态管理日期，已在安全渲染逻辑中修复。",
        "",
        "## 3. 为什么不直接批量训练全部站点",
        "完整计划共有 14 个 observed-year PPO 模型。当前阶段只验证框架和 HLA 2007 单模型流程，避免长时间训练时难以定位环境、渲染、动作记录或 OOM 问题。",
        "",
        "## 4. 实测年份和训练-验证组合",
        f"组合表：`{rel(PPO_ROOT / 'configs' / 'ppo_observed_year_experiment_plan.csv')}`",
        "",
        md_table(plan),
        "",
        "## 5. 安全 rendered input 逻辑",
        "PPO 训练和评估使用 `src/ppo_safe_rendering.py`，继承 smoke test debug 的修复：不覆盖 my_data 原始模板；每个站点-年份生成临时 rendered input；启用 MI/MF；保证 irrigation/fertilizer sections 存在；替换跨年份残留的静态灌溉/施肥事件；复制对应 QC WTH 到临时目录。",
        "",
        f"rendered input 检查表：`{rel(render_check_path)}`",
        "",
        md_table(render_check.groupby(['check_item', 'status']).size().reset_index(name='count') if not render_check.empty else render_check),
        "",
        "## 6. 生成的代码和配置",
        "\n".join(f"- `{item}`" for item in scripts),
        "",
        "## 7. HLA 2007 pretrain smoke check",
        f"pretrain smoke check 表：`{rel(PPO_ROOT / 'smoke_checks' / 'pretrain_smoke_check_summary.csv')}`",
        "",
        md_table(smoke),
        "",
        "## 8. HLA 2007 debug PPO",
        f"- station: `{debug['station']}`",
        f"- train_year: `{debug['train_year']}`",
        f"- seed: `{config['seed']}`",
        f"- configured debug total_timesteps: `{debug['total_timesteps']}`",
        "- note: SB3 PPO 按 n_steps rollout 成批更新，因此实际日志显示 total_timesteps 可略高于配置值。",
        f"- status: `{status_text}`",
        "",
        "模型文件：",
        "\n".join(f"- `{rel(path)}`" for path in model_files) if model_files else "- no model file",
        "",
        "## 9. train/eval daily output 和图",
        f"- daily CSV 数：{len(daily_files)}",
        f"- 响应图 PNG 数：{len(figure_files)}",
        "",
        md_table(evaluation),
        "",
        "## 10. 策略选择脚本输出",
        f"策略选择输出：`{rel(PPO_ROOT / 'strategy_selection' / 'best_policy_by_site.csv')}`",
        "",
        md_table(selection),
        "",
        "## 11. 仍需人工确认的 PPO 超参数",
        "最终批量训练超参数仍保留为 null，避免把 debug 参数误当最终参数。需要确认："
        + (", ".join(need_confirm) if need_confirm else "无。"),
        "",
        "## 12. 下一步建议",
        "可以进入下一步 HLA/SYA/LCA 的逐站点批量训练准备，但建议顺序仍然是：每次只启动一个模型，先跑对应 pretrain smoke check，再训练，再评估，确认 daily CSV 和图完整后再进入下一个模型。",
    ]
    REPORT_MD.parent.mkdir(parents=True, exist_ok=True)
    REPORT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    report_dir = PPO_ROOT / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REPORT_MD, report_dir / REPORT_MD.name)


def set_font(paragraph, size=15, bold=False, color=RGBColor(0, 0, 0)) -> None:
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.5))
    p = box.text_frame.paragraphs[0]
    p.text = title
    set_font(p, 23, True)


def add_bullets(slide, bullets: list[str], top=1.05, size=16) -> None:
    box = slide.shapes.add_textbox(Inches(0.65), Inches(top), Inches(12), Inches(5.9))
    tf = box.text_frame
    tf.clear()
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.space_after = Pt(6)
        set_font(p, size)


def add_table(slide, df: pd.DataFrame, max_rows=10, top=1.1, font_size=9) -> None:
    if df.empty:
        add_bullets(slide, ["No data."], top=top)
        return
    data = df.head(max_rows)
    rows, cols = len(data) + 1, len(data.columns)
    table = slide.shapes.add_table(rows, cols, Inches(0.45), Inches(top), Inches(12.4), Inches(5.8)).table
    for j, col in enumerate(data.columns):
        cell = table.cell(0, j)
        cell.text = str(col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            set_font(p, font_size, True, RGBColor(255, 255, 255))
    for i, (_, row) in enumerate(data.iterrows(), start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = "" if pd.isna(value) else str(value)
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_BLUE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                set_font(p, font_size)


def write_ppt() -> None:
    config = load_yaml(CONFIG_PATH)
    plan = read_csv(PPO_ROOT / "configs" / "ppo_observed_year_experiment_plan.csv")
    smoke = read_csv(PPO_ROOT / "smoke_checks" / "pretrain_smoke_check_summary.csv")
    evaluation = read_csv(PPO_ROOT / "evaluation" / "ppo_evaluation_summary.csv")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "PPO observed-year 框架生成")
    add_bullets(slide, ["目标：生成 leave-one PPO 训练/验证/绘图框架。", "本阶段只跑 HLA 2007 seed0 小步数 debug PPO。", "不批量训练全部站点，不修改 reward，不覆盖 my_data 原始数据。"])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "前置 smoke test 修复结果")
    add_bullets(slide, ["上一阶段 29 个 timeout 案例已全部重跑成功。", "修复点：MI/MF 管理层级、irrigation/fertilizer sections、跨年份静态管理日期。", "PPO 必须复用同一套 safe rendered input 逻辑。"])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "PPO 实验设计")
    add_table(slide, plan[["station", "train_year", "train_year_label", "validation_years", "cross_validation_type"]], max_rows=14, font_size=8)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "安全 rendered input 机制")
    add_bullets(slide, ["每次训练/评估生成临时 rendered input。", "不修改 my_data 原始模板、SOL、WTH、CUL。", "保证 MI/MF、灌溉 section、施肥 section 可用。", "替换 simulation start 前的历史灌溉/施肥事件。"])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "pretrain smoke check 机制")
    add_table(slide, smoke, max_rows=4, font_size=10)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "PPO daily output 设计")
    add_bullets(slide, ["训练后评估训练年和 validation years。", "保存 dap/topwt/grnwt/xlai/totir/tofer/swfac/nstres/reward。", "同时保存 real_action_amir/anfer 和 normalized_action_amir/anfer。", "每次评估生成 5 张 daily response 图。"])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "HLA 2007 debug 试训结果")
    add_table(slide, evaluation[["station", "train_year", "eval_year", "run_status", "episode_length", "final_grnwt", "total_irrigation", "total_n_fertilizer"]], max_rows=5, font_size=9)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "后续完整训练计划")
    add_bullets(slide, ["下一步建议先逐个运行 HLA、SYA、LCA。", "每个模型先 pretrain smoke check，再训练，再评估。", "FQA/YCA 只有 2 年，报告中应标记为 limited two-year cross validation。"])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "风险和注意事项")
    add_bullets(slide, ["最终 PPO 超参数仍需确认，配置中保留 null。", "debug PPO 只用于验证流程，不代表最终策略效果。", "模型文件不要上传大批量版本到 GitHub。", "后续训练仍需关注 OOM 和 DSSAT 子进程清理。"])

    REPORT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(REPORT_PPTX)
    report_dir = PPO_ROOT / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REPORT_PPTX, report_dir / REPORT_PPTX.name)


def main() -> None:
    write_markdown()
    write_ppt()
    print(rel(REPORT_MD))
    print(rel(REPORT_PPTX))


if __name__ == "__main__":
    main()
