"""Create the 020 temporary-freeze records and selective backup manifest.

This script only writes new closeout files under docs/ and results/.
"""
from __future__ import annotations

import hashlib
import json
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs'
OUT=ROOT/'results/sy_lc_weather_augmentation_closeout'
FREEZE=DOCS/'sy_lc_weather_augmentation_temporary_freeze.md'
MANIFEST_MD=DOCS/'sy_lc_weather_augmentation_backup_manifest.md'
RECORD_MD=DOCS/'sy_lc_weather_augmentation_closeout_020.md'
PPT=DOCS/'sy_lc_weather_augmentation_closeout_020.pptx'
MANIFEST_JSON=OUT/'backup_manifest.json'
FREEZE_JSON=OUT/'freeze_status.json'

REPORTS=[
 'docs/sy_lc_random_weather_source_audit_015.md',
 'docs/provenance_resolution_report.md',
 'docs/sy_lc_wgen_readiness_review.md',
 'docs/lca_cli_parameter_audit_016.md',
 'docs/lca_wgen_window_review.md',
 'docs/sya_train_only_fill_sensitivity_review.md',
 'docs/sya_cli_parameter_audit_018.md',
 'docs/sya_wgen_readiness_review_019.md',
]
SCRIPTS=[
 'src/audit_sy_lc_random_weather_source_015.py',
 'src/verify_sy_lc_weather_audit_015.py',
 'src/revise_sy_lc_weather_provenance_015.py',
 'src/trace_sy_lc_nonrain_provenance_015.py',
 'src/review_sy_lc_wgen_readiness_015.py',
 'src/analyze_sy_lc_wgen_train_only_sensitivity_015.py',
 'src/audit_lca_cli_parameters_016.py',
 'src/review_lca_wgen_windows.py',
 'src/review_sya_train_only_fill_017.py',
 'src/audit_sya_cli_parameters_018.py',
 'src/review_sya_wgen_readiness_019.py',
 'src/archive_sy_lc_weather_augmentation_020.py',
 'scripts/build_dssat_cli.py',
 'src/weather_preprocess.py',
]
RESULT_DIRS=[
 'results/sy_lc_random_weather_015',
 'results/lca_cli_parameter_audit_016',
 'results/lca_wgen_window_review',
 'results/sya_train_only_fill_sensitivity',
 'results/sya_cli_parameter_audit_018',
 'results/sya_wgen_readiness_review_019',
]
PROMPTS=[f'prompt_02/{n}' for n in (
 '015_sya_lca_random_weather_source_audit.md',
 '016_lca_cli_parameter_audit.md',
 '017_sya_train_only_fill_sensitivity.md',
 '018_sya_cli_parameter_audit.md',
 '019_sya_wgen_readiness_review.md',
 '020_sy_lc_weather_augmentation_temporary_freeze_and_backup.md',
)]


def sha(path):
    h=hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):
            h.update(block)
    return h.hexdigest()


def git(*args):
    p=subprocess.run(['git',*args],cwd=ROOT,text=True,capture_output=True)
    return p.returncode,p.stdout.strip()


def slide(prs,title,body,source=None):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    bg=s.background.fill;bg.solid();bg.fore_color.rgb=RGBColor(248,250,252)
    tb=s.shapes.add_textbox(Inches(.65),Inches(.48),Inches(11.9),Inches(.8))
    tf=tb.text_frame;tf.clear()
    p=tf.paragraphs[0];p.text=title;p.font.name='Microsoft YaHei';p.font.size=Pt(28);p.font.bold=True;p.font.color.rgb=RGBColor(20,43,67)
    box=s.shapes.add_textbox(Inches(.78),Inches(1.55),Inches(11.85),Inches(4.95))
    t=box.text_frame;t.clear();t.word_wrap=True
    for i,line in enumerate(body):
        p=t.paragraphs[0] if i==0 else t.add_paragraph()
        p.text=line;p.font.name='Microsoft YaHei';p.font.size=Pt(20);p.font.color.rgb=RGBColor(37,57,75)
        p.space_after=Pt(19)
    if source:
        foot=s.shapes.add_textbox(Inches(.78),Inches(6.88),Inches(11.7),Inches(.34))
        p=foot.text_frame.paragraphs[0];p.text=source;p.font.name='Microsoft YaHei';p.font.size=Pt(10);p.font.color.rgb=RGBColor(92,110,128)
    return s


def make_ppt():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    slide(prs,'SYA / LCA 天气增强阶段性收尾',[
        '当前实验对两站随机天气增强实行临时冻结。',
        'historical-weather PPO 冻结结果继续用于本轮研究。',
        '审查范围：来源、填补、CLI 参数和 WGEN 输入适宜性。'
    ],'020 closeout · 2026-10-06')
    slide(prs,'SYA：2005 年缺段削弱日际变率',[
        '462 个非降水缺口中，424 个位于 2005 年。',
        '训练期独立均值填补解决了验证期信息进入填补均值的问题。',
        '2005 年 1–4 月 SRAD、TMAX、TMIN 月内标准差均为 0。',
        '019 gate：BLOCKED_BY_WEATHER_STATISTICS。'
    ],'依据：017 敏感性、018 CLI 审计、019 变率复审')
    slide(prs,'LCA：1 月降水参数未稳定',[
        '2005–2013 的 1 月仅有 4 个湿日。',
        '2000–2013 增至 26 个，其中 14 个来自 2000 年。',
        '候选窗口的 1 月 Gamma ALPHA 均达到 0.998 截断值。',
        '当前 gate：BLOCKED_BY_JANUARY_GAMMA_SAMPLE。'
    ],'依据：016 CLI 审计、长历史窗口审查')
    slide(prs,'历史天气 PPO 结果继续保留',[
        '本轮阻塞针对 WGEN 随机天气的参数拟合输入。',
        '已有 historical-weather PPO 结果按原冻结实验口径保留。',
        '不因 WGEN 阻塞重训，也不据此判定历史天气结果无效。'
    ],'范围说明：本次未复算 PPO 或 DSSAT')
    slide(prs,'两站本轮均临时退出 WGEN',[
        'SYA：EXCLUDED_FROM_WGEN_FOR_CURRENT_EXPERIMENT。',
        'LCA：EXCLUDED_FROM_WGEN_FOR_CURRENT_EXPERIMENT。',
        '状态：TEMPORARILY_FROZEN_FOR_CURRENT_WEATHER_AUGMENTATION_EXPERIMENT。',
        '未来获得更好数据或方法后可重新开启。'
    ],'机器状态：results/sy_lc_weather_augmentation_closeout/freeze_status.json')
    slide(prs,'Git 归档范围',[
        '归档 015–019 报告、审查脚本、关键 CSV、gate/QC JSON 和候选 CLI。',
        '保留原始 WTH 与既有 PPO 结果，不把模型文件加入本次提交。',
        '逐文件 SHA256 与 Git 状态记录在 backup manifest。'
    ],'清单：docs/sy_lc_weather_augmentation_backup_manifest.md')
    slide(prs,'重新开启的条件',[
        'SYA：可靠补测或独立重建 2005 年长缺段，并重新审查变率。',
        'LCA：更长、连续且质量合格的降水记录，或重新验证冬季参数化。',
        '若采用新的天气生成方法，需要独立完成方法和输出 QC。'
    ],'本轮停止 WGEN、DSSAT、PPO')
    prs.save(PPT)


def main():
    for p in (FREEZE,MANIFEST_MD,RECORD_MD,PPT,MANIFEST_JSON,FREEZE_JSON):
        if p.exists():raise SystemExit(f'Refusing to overwrite {p}')
    for p in REPORTS+SCRIPTS+PROMPTS:
        if not (ROOT/p).is_file():raise FileNotFoundError(p)
    for d in RESULT_DIRS:
        if not (ROOT/d).is_dir():raise FileNotFoundError(d)
    OUT.mkdir(parents=True)
    status={
      'scope':'current_weather_augmentation_experiment',
      'freeze_type':'temporary',
      'status':'TEMPORARILY_FROZEN_FOR_CURRENT_WEATHER_AUGMENTATION_EXPERIMENT',
      'historical_weather_ppo_results':{
        'SYA':'TEMPORARILY_FROZEN_ACCEPTED_FOR_CURRENT_EXPERIMENT',
        'LCA':'TEMPORARILY_FROZEN_ACCEPTED_FOR_CURRENT_EXPERIMENT'},
      'weather_augmentation':{
        'SYA':{'status':'EXCLUDED_FROM_WGEN_FOR_CURRENT_EXPERIMENT','reason':'BLOCKED_BY_WEATHER_STATISTICS','reopen_allowed':True},
        'LCA':{'status':'EXCLUDED_FROM_WGEN_FOR_CURRENT_EXPERIMENT','reason':'BLOCKED_BY_JANUARY_GAMMA_SAMPLE','reopen_allowed':True}},
      'WGEN_run':False,'synthetic_weather_generated':False,'DSSAT_run':False,'PPO_run':False,
      'historical_ppo_revalidated_in_closeout':False,
      'evidence':{
        'SYA':'results/sya_wgen_readiness_review_019/wgen_readiness_gate.json',
        'LCA':'results/lca_wgen_window_review/window_summary.json'}}
    FREEZE_JSON.write_text(json.dumps(status,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    FREEZE.write_text('''# SYA / LCA 天气增强实验临时冻结说明

日期：2026-10-06。目的：在现有天气资料和 WGEN 方法约束下，结束本轮 SYA/LCA 随机天气增强推进，同时保存可复核的实验链。统一状态为 `TEMPORARILY_FROZEN_FOR_CURRENT_WEATHER_AUGMENTATION_EXPERIMENT`。该状态可在数据或方法改善后重新开启。

| 站点 | 本轮天气增强 | 当前限制 | historical-weather PPO |
|---|---|---|---|
| SYA | `EXCLUDED_FROM_WGEN_FOR_CURRENT_EXPERIMENT` | `BLOCKED_BY_WEATHER_STATISTICS` | 保留当前冻结结果 |
| LCA | `EXCLUDED_FROM_WGEN_FOR_CURRENT_EXPERIMENT` | `BLOCKED_BY_JANUARY_GAMMA_SAMPLE` | 保留当前冻结结果 |

## 已排除的来源问题

RAIN 原始 XLS 空白表示 0 mm 降水。LCA cleaned CSV 与最终 WTH 的 1,113 处差异来自历史下载补充流程，差异本身不再列为来源异常。SYA 的 462 项非降水缺口确认采用旧站点×月份均值填补；LCA 残余 22 项数值与旧填补一致，但逐日外部补测链仍不完整。方法审查允许带记录的缺测填补进入候选参数拟合，不要求所有值都是原始观测。原 015 provenance gate 作为历史审查状态保留，后续方法 gate 记录具体 WGEN 限制。

## 本轮真正的 WGEN 限制

SYA train-only 候选修正了旧均值填补对 2014–2023 验证期的依赖；018 CLI 的 168 项参数和格式通过。但 2005 年 1–4 月 SRAD、TMAX、TMIN 各自整月恒定，最长连续受填补影响 139 天。均值填补不能恢复日际变率，因此 019 判 `BLOCKED_BY_WEATHER_STATISTICS`。

LCA 2005–2013 CLI 参数齐全，但 1 月只有 4 个湿日，Gamma ALPHA 达 0.998 截断上限。扩至 2000–2013 后 1 月有 26 个湿日，其中 14 个集中在 2000 年，原始 ALPHA 仍超过截断范围；1995–2013 无连续合格序列。本轮维持 `BLOCKED_BY_JANUARY_GAMMA_SAMPLE`。

## 历史天气 PPO 与论文表述

已冻结的 historical-weather PPO 结果继续按原实验范围用于当前研究。本轮仅审查随机天气增强输入，没有重新计算或重新验证 PPO 结果，也不据此判定历史天气模拟或 PPO 结果无效。论文宜写为：两站的历史天气策略结果按既有条件报告；SYA 因 2005 年连续均值填补导致月内日际变率缺失，LCA 因冬季湿日稀少且 Gamma 参数稳定性不足，未纳入本轮 WGEN 增强比较。不要把缺少随机天气对照解读为两站策略失败。

## 重新开启条件

SYA 需要可靠的 2005 年外部补测，或独立验证的天气重建方案，并重新审查月内和湿/干条件变率。LCA 需要更长、连续且质量合格的降水记录，或独立验证的冬季降水参数化方案。采用新的天气生成方法时，也需重新完成输入来源、参数和输出 QC。本轮实验到此收尾；未运行 WGEN、DSSAT 或 PPO，未修改原 WTH 或冻结 PPO 结果。
''',encoding='utf-8')
    RECORD_MD.write_text('''# 020 SYA / LCA 天气增强实验收尾记录

日期：2026-10-06。任务范围为 015–019 审计链的临时冻结、归档清单和 Git 备份。总判定见 `sy_lc_weather_augmentation_temporary_freeze.md` 与 `results/sy_lc_weather_augmentation_closeout/freeze_status.json`；本记录有对应的 `.pptx` 汇报版。

SYA：017 训练期独立填补与 018 CLI 计算完成，019 发现 2005 年 1–4 月三个非降水变量月内标准差为 0，故本轮不进入 WGEN。LCA：016 CLI 与长窗口审查显示 1 月降水 Gamma 形状仍受稀少且集中湿日限制，本轮不进入 WGEN。两站的 historical-weather PPO 冻结结果继续保留；本任务未重新训练或改动这些结果。

归档采用逐文件清单，保留报告、脚本、机器 gate、参数 QC、关键 CSV 与候选 CLI。原天气输入、模型 checkpoint 和无关实验文件不纳入新增提交。具体路径、SHA256、大小和 Git 状态见 `sy_lc_weather_augmentation_backup_manifest.md` / `results/sy_lc_weather_augmentation_closeout/backup_manifest.json`。

本轮没有运行 WGEN、DSSAT、PPO，没有生成随机天气，也没有更改 WTH。Git 分支、提交与推送结果以最终执行记录为准；不得把临时冻结解释为永久终止。
''',encoding='utf-8')
    make_ppt()
    result_files=[]
    for d in RESULT_DIRS:
        result_files.extend((ROOT/d).rglob('*'))
    result_files=[p for p in result_files if p.is_file() and p.suffix.lower() in ('.json','.csv','.cli','.md','.txt')]
    selected={*(ROOT/p for p in REPORTS+SCRIPTS+PROMPTS),*result_files,FREEZE,FREEZE_JSON,RECORD_MD,PPT,MANIFEST_MD,MANIFEST_JSON}
    refs=paths=[ROOT/'DSSAT_auto_validation/multisite_new_cultivar_inputs_013/SY'/f'CNSY{y%100:02d}01.WTH' for y in range(2005,2014)]
    refs += [ROOT/'DSSAT_auto_validation/multisite_new_cultivar_inputs_013_lowIC_manual/LC'/f'CNLC{y%100:02d}01.WTH' for y in range(2000,2014)]
    records=[]
    for p in sorted(selected|set(refs),key=lambda x:x.as_posix()):
        rel=p.relative_to(ROOT).as_posix()
        if p in refs:category='protected_source_WTH';include=False;task='source'
        elif p in result_files:category='machine_evidence';include=True;task='015-019'
        elif p in (FREEZE,FREEZE_JSON,RECORD_MD,PPT,MANIFEST_MD,MANIFEST_JSON):category='020_closeout';include=True;task='020'
        elif rel in REPORTS:category='audit_report';include=True;task='015-019'
        elif rel in SCRIPTS:category='audit_or_validation_code';include=True;task='015-020'
        else:category='task_prompt';include=True;task='015-020'
        tracked=git('ls-files','--error-unmatch','--',rel)[0]==0
        ignored=git('check-ignore','--',rel)[0]==0
        size=p.stat().st_size if p.exists() else None
        records.append({'path':rel,'category':category,'include_in_git':include,'sha256':sha(p) if p.exists() and p not in (MANIFEST_MD,MANIFEST_JSON) else None,
            'sha256_note':'self-referential manifest; hash omitted' if p in (MANIFEST_MD,MANIFEST_JSON) else None,
            'task':task,'key_evidence':category in ('machine_evidence','audit_report','020_closeout'),
            'large_or_excluded':not include,'size_bytes':size,
            'git_state_before_020':'tracked' if tracked else 'ignored_untracked' if ignored else 'untracked',
            'force_add_required':include and ignored and not tracked})
    manifest={'scope':'SYA_LCA_weather_augmentation_015_020','branch':git('branch','--show-current')[1],
       'remote':git('remote','get-url','origin')[1],
       'files':records,'manifest_hash_note':'Manifest files omit their own SHA256 to avoid recursive self-reference.',
       'force_add_policy':'Only explicitly listed small ignored evidence files may use git add -f.'}
    MANIFEST_JSON.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    totals={'included':sum(r['include_in_git'] for r in records),'source_references':sum(not r['include_in_git'] for r in records),
       'included_bytes':sum(r['size_bytes'] or 0 for r in records if r['include_in_git'])}
    lines=['# SYA / LCA 天气增强实验备份清单','',
      f"分支：`{manifest['branch']}`；remote：`{manifest['remote']}`。计划提交 {totals['included']} 个文件，已计大小约 {totals['included_bytes']/1024:.1f} KiB；另列 {totals['source_references']} 个受保护 WTH 输入作为哈希参照，不新增提交。",'',
      'SHA256、生成任务、是否关键、文件大小、Git 原有 tracked/untracked/ignored 状态及 force-add 标志以 `backup_manifest.json` 为准。清单自身因自引用无法写入自身 SHA256。','',
      '| 类别 | 文件数 | Git 处理 |','|---|---:|---|']
    cats=sorted(set(r['category'] for r in records))
    for c in cats:
        subset=[r for r in records if r['category']==c]
        lines.append(f"| {c} | {len(subset)} | {'只读参照，不新增提交' if c=='protected_source_WTH' else '逐项暂存'} |")
    lines += (['','关键报告：']+[f'- `{p}`' for p in REPORTS]+['','机器 gate / QC：']+
        [f"- `{r['path']}`" for r in records if r['include_in_git'] and r['path'].endswith(('/final_gate.json','/methodology_gate.json','/parameter_qc.json','/wgen_readiness_gate.json','/freeze_status.json'))])
    lines+=['','候选 CLI：','- `results/lca_cli_parameter_audit_016/CNLC_2005_2013_candidate.CLI`','- `results/sya_cli_parameter_audit_018/CNSY_2005_2013_train_only_candidate.CLI`','',
      '源 WTH、PPO checkpoint、原始 XLS 以及无关实验结果不因本次归档而添加。忽略规则下的 CSV/CLI 等，只对本 JSON 明确 `include_in_git=true` 且 `force_add_required=true` 的小型文件逐项使用 `git add -f`。','']
    MANIFEST_MD.write_text('\n'.join(lines),encoding='utf-8')
    print(json.dumps(totals,ensure_ascii=False))


if __name__=='__main__':main()
