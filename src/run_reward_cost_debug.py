from __future__ import annotations

import inspect
import math
import shutil
import sys
import traceback
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd

from ppo_safe_rendering import PROJECT_ROOT, load_yaml
from ppo_train import train_one_policy
from reward_candidates import REWARD_CANDIDATES, candidate_table, patch_reward_candidate, unpatch_reward_candidate


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_reward_cost_debug.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "reward_cost_debug"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_reward_cost_revision_and_debug_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_reward_cost_revision_and_debug_report.pptx"


def ensure_dirs() -> None:
    for sub in [
        "configs",
        "reward_versions",
        "smoke_checks",
        "rendered_inputs",
        "models",
        "logs",
        "tensorboard",
        "daily_outputs",
        "evaluation",
        "figures/summary",
        "reports",
    ]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)


def df_to_markdown(df: pd.DataFrame) -> str:
    if df.empty:
        return ""
    work = df.copy()
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    separator = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, separator, *rows])


def write_candidate_files() -> None:
    import gym_dssat_pdi.envs.configs.rewards as rewards

    review_path = OUTPUT_ROOT / "reward_versions" / "current_reward_review.md"
    source_path = inspect.getsourcefile(rewards) or "unknown"
    all_reward_source = inspect.getsource(rewards.all_reward)
    fert_source = inspect.getsource(rewards.fertilization_reward)
    irrig_source = inspect.getsource(rewards.irrigation_reward)
    review = f"""# Current reward review

Generated at: 2026-06-06

## Reward source

- File path: `{source_path}`
- Mode used by PPO: `all`
- Functions: `all_reward`, `fertilization_reward`, `irrigation_reward`

## Code snippets

```python
{fert_source}

{irrig_source}

{all_reward_source}
```

## Inputs and state variables

- Inputs: previous state, next state, history, cultivar.
- Crop state variables used directly: `trnu`, `topwt`, `swfac`.
- Action variables used directly: `amir`, `anfer`.
- Cumulative actions are available through `history['action']`.

## Cost structure

- Daily irrigation cost: yes, through `irrigation_reward = delta_topwt - 15 * amir`.
- Daily nitrogen cost: yes, through `fertilization_reward = trnu * coef - penality * anfer`.
- Season irrigation cost: optional env-var hook only, default 0.
- Season nitrogen cost: optional env-var hook only, default 0.
- Nonlinear excess penalty: no, only optional linear excess hooks with default 0.
- Terminal grain-yield benefit: no explicit terminal reward.

## Main problem

The current reward gives positive growth/nitrogen-recovery incentives every step while action safety caps total water and nitrogen externally. In the 006_02 season-cap sensitivity analysis, PPO filled every cap level from 100/150 to 300/450. This indicates that the current reward does not make PPO internalize seasonal water and nitrogen scarcity; the wrapper prevents unsafe totals, but the reward still rewards using the entire allowed budget.
"""
    review_path.write_text(review, encoding="utf-8")

    design_path = OUTPUT_ROOT / "reward_versions" / "reward_candidate_design.md"
    rows = pd.DataFrame(candidate_table())
    rows.to_csv(OUTPUT_ROOT / "reward_versions" / "reward_candidate_table.csv", index=False, encoding="utf-8-sig")
    design = """# Reward candidate design

Generated at: 2026-06-06

This stage does not overwrite the original gym-DSSAT reward file. Candidate rewards are defined in `src/reward_candidates.py` and are injected at runtime by patching the imported `gym_dssat_pdi.envs.configs.rewards.all_reward` function before environment creation.

## Candidate families

- `current_reward_baseline`: original `all_reward`.
- `candidate_A_linear_cost`: original reward minus extra linear daily irrigation and nitrogen costs.
- `candidate_B_linear_plus_excess_penalty`: original reward minus linear costs and a quadratic penalty after cumulative irrigation exceeds 200 mm or cumulative nitrogen exceeds 300 kg/ha.
- `candidate_C_incremental_yield_minus_total_cost`: incremental grain-yield proxy minus water and nitrogen action costs. The reward callback has no explicit `done` flag, so terminal yield is approximated by positive daily `grnwt` increments.

## Candidate table

""" + df_to_markdown(rows)
    design_path.write_text(design, encoding="utf-8")


def build_plan(config: dict) -> pd.DataFrame:
    cap = config["cap"]
    rows = []
    order = 1
    for reward_version in config["reward_candidates"]:
        cand = REWARD_CANDIDATES[reward_version]
        rows.append(
            {
                "run_order": order,
                "stage": "HLA_pilot",
                "station": "HLA",
                "train_year": 2011,
                "eval_years": "2011,2007,2009",
                "reward_version": reward_version,
                "reward_family": cand.family,
                "cap_name": cap["cap_name"],
                "season_irrigation_cap": cap["season_irrigation_cap"],
                "season_n_cap": cap["season_n_cap"],
                "seed": config["seed"],
                "total_timesteps": config["training"]["total_timesteps"],
            }
        )
        order += 1
    plan = pd.DataFrame(rows)
    out = OUTPUT_ROOT / "configs" / "reward_cost_debug_plan.csv"
    plan.to_csv(out, index=False, encoding="utf-8-sig")
    shutil.copy2(CONFIG_PATH, OUTPUT_ROOT / "configs" / CONFIG_PATH.name)
    return plan


def row_metadata(config: dict, reward_version: str, station: str) -> dict:
    cand = REWARD_CANDIDATES[reward_version]
    cap = config["cap"]
    return {
        "reward_version": reward_version,
        "reward_family": cand.family,
        "cap_name": cap["cap_name"],
        "season_irrigation_cap": float(cap["season_irrigation_cap"]),
        "season_n_cap": float(cap["season_n_cap"]),
        "cv_type": config["training"]["cv_types"].get(station, ""),
    }


def enrich_evaluation_summary() -> pd.DataFrame:
    src = OUTPUT_ROOT / "evaluation" / "ppo_evaluation_summary.csv"
    if not src.exists():
        return pd.DataFrame()
    df = pd.read_csv(src)
    for col in ["total_irrigation", "total_n_fertilizer", "season_irrigation_cap", "season_n_cap", "final_grnwt"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    df["irrigation_saturation_ratio"] = df["total_irrigation"] / df["season_irrigation_cap"]
    df["n_saturation_ratio"] = df["total_n_fertilizer"] / df["season_n_cap"]
    df["quality_gate_pass"] = (
        (df["run_status"] == "ok")
        & (df["episode_completed"].astype(str).str.lower().isin(["true", "1"]))
        & (df["irrigation_saturation_ratio"] <= 1.000001)
        & (df["n_saturation_ratio"] <= 1.000001)
    )
    out = OUTPUT_ROOT / "evaluation" / "reward_candidate_evaluation_summary.csv"
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return df


def compare_candidates(df: pd.DataFrame) -> pd.DataFrame:
    pilot = df[df["station"] == "HLA"].copy()
    grouped = (
        pilot.groupby(["reward_version", "reward_family"], dropna=False)
        .agg(
            eval_count=("eval_year", "count"),
            ok_count=("quality_gate_pass", "sum"),
            mean_yield=("final_grnwt", "mean"),
            std_yield=("final_grnwt", "std"),
            mean_reward=("mean_reward", "mean"),
            mean_irrigation=("total_irrigation", "mean"),
            mean_n=("total_n_fertilizer", "mean"),
            mean_irrigation_saturation_ratio=("irrigation_saturation_ratio", "mean"),
            mean_n_saturation_ratio=("n_saturation_ratio", "mean"),
        )
        .reset_index()
    )
    baseline = grouped[grouped["reward_version"] == "current_reward_baseline"]
    if len(baseline):
        b = baseline.iloc[0]
        grouped["yield_loss_vs_current_reward"] = (b["mean_yield"] - grouped["mean_yield"]) / b["mean_yield"]
        grouped["input_reduction_vs_current_reward"] = 1.0 - (
            grouped["mean_irrigation"] + grouped["mean_n"]
        ) / (b["mean_irrigation"] + b["mean_n"])
    else:
        grouped["yield_loss_vs_current_reward"] = math.nan
        grouped["input_reduction_vs_current_reward"] = math.nan
    grouped["yield_per_100mm_irrigation"] = grouped["mean_yield"] / (grouped["mean_irrigation"] / 100.0)
    grouped["yield_per_100kg_n"] = grouped["mean_yield"] / (grouped["mean_n"] / 100.0)
    grouped["recommendation_pass"] = (
        (grouped["mean_irrigation_saturation_ratio"] < 0.95)
        & (grouped["mean_n_saturation_ratio"] < 0.95)
        & (grouped["yield_loss_vs_current_reward"] <= 0.10)
        & (grouped["ok_count"] == grouped["eval_count"])
    )
    out = OUTPUT_ROOT / "evaluation" / "reward_candidate_comparison.csv"
    grouped.to_csv(out, index=False, encoding="utf-8-sig")
    return grouped


def best_candidate(comparison: pd.DataFrame) -> str | None:
    passed = comparison[comparison["recommendation_pass"]].copy()
    if passed.empty:
        return None
    passed = passed.sort_values(
        ["input_reduction_vs_current_reward", "yield_loss_vs_current_reward", "mean_yield"],
        ascending=[False, True, False],
    )
    return str(passed.iloc[0]["reward_version"])


def copy_candidate_figures(df: pd.DataFrame) -> None:
    for _, row in df.iterrows():
        figure_dir = row.get("figure_dir")
        reward_version = row.get("reward_version")
        eval_year = row.get("eval_year")
        station = row.get("station")
        if not figure_dir or not reward_version or pd.isna(eval_year):
            continue
        src = PROJECT_ROOT / str(figure_dir)
        dst = OUTPUT_ROOT / "figures" / str(station) / str(reward_version) / f"eval_{int(eval_year)}"
        if not src.exists() or dst.exists():
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copytree(src, dst)


def plot_summary(df: pd.DataFrame, comparison: pd.DataFrame) -> None:
    fig_root = OUTPUT_ROOT / "figures" / "summary"
    fig_root.mkdir(parents=True, exist_ok=True)
    pilot = df[df["station"] == "HLA"].copy()
    plt.rcParams["font.family"] = ["DejaVu Sans"]

    plt.figure(figsize=(10, 6))
    for reward_version, sub in pilot.groupby("reward_version"):
        plt.scatter(sub["total_irrigation"] + sub["total_n_fertilizer"], sub["final_grnwt"], label=reward_version)
    plt.xlabel("Total irrigation + nitrogen input")
    plt.ylabel("Final grain yield")
    plt.legend(fontsize=8)
    plt.tight_layout()
    plt.savefig(fig_root / "reward_candidates_yield_vs_input.png", dpi=180)
    plt.close()

    x = comparison["reward_version"]
    plt.figure(figsize=(11, 6))
    plt.plot(x, comparison["mean_irrigation"], marker="o", label="irrigation")
    plt.plot(x, comparison["mean_n"], marker="o", label="nitrogen")
    plt.xticks(rotation=45, ha="right")
    plt.ylabel("Mean seasonal input")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_root / "reward_candidates_irrigation_n_ratio.png", dpi=180)
    plt.close()

    plt.figure(figsize=(10, 6))
    plt.scatter(comparison["mean_yield"], comparison["mean_reward"])
    for _, row in comparison.iterrows():
        plt.annotate(row["reward_version"], (row["mean_yield"], row["mean_reward"]), fontsize=8)
    plt.xlabel("Mean final grain yield")
    plt.ylabel("Mean reward")
    plt.tight_layout()
    plt.savefig(fig_root / "reward_candidates_reward_vs_yield.png", dpi=180)
    plt.close()

    plt.figure(figsize=(11, 6))
    plt.plot(x, comparison["mean_irrigation_saturation_ratio"], marker="o", label="irrigation ratio")
    plt.plot(x, comparison["mean_n_saturation_ratio"], marker="o", label="nitrogen ratio")
    plt.axhline(0.95, color="black", linestyle="--", linewidth=1)
    plt.xticks(rotation=45, ha="right")
    plt.ylabel("Mean saturation ratio")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_root / "reward_candidates_saturation_ratio.png", dpi=180)
    plt.close()

    plt.figure(figsize=(11, 6))
    for reward_version, sub in pilot.groupby("reward_version"):
        sub = sub.sort_values("eval_year")
        plt.plot(sub["eval_year"], sub["final_grnwt"], marker="o", label=reward_version)
    plt.xlabel("Eval year")
    plt.ylabel("Final grain yield")
    plt.legend(fontsize=7, ncol=2)
    plt.tight_layout()
    plt.savefig(fig_root / "reward_candidates_cross_year_yield_stability.png", dpi=180)
    plt.close()


def run_training(config: dict, reward_version: str, station: str, train_year: int) -> None:
    patch_reward_candidate(reward_version)
    meta = row_metadata(config, reward_version, station)
    try:
        train_one_policy(
            station=station,
            train_year=train_year,
            seed=int(config["seed"]),
            total_timesteps=int(config["training"]["total_timesteps"]),
            config_path=CONFIG_PATH,
            debug=False,
            policy_tag=reward_version,
            row_metadata=meta,
        )
    finally:
        unpatch_reward_candidate()


def run_hla_pilot(config: dict) -> tuple[pd.DataFrame, pd.DataFrame, str | None]:
    for reward_version in config["reward_candidates"]:
        print(f"[reward-cost-debug] HLA pilot: {reward_version}", flush=True)
        run_training(config, reward_version, "HLA", 2011)
    df = enrich_evaluation_summary()
    copy_candidate_figures(df)
    comparison = compare_candidates(df)
    plot_summary(df, comparison)
    return df, comparison, best_candidate(comparison)


def run_extension(config: dict, selected_reward: str | None) -> pd.DataFrame:
    if not selected_reward:
        out = OUTPUT_ROOT / "evaluation" / "cross_site_reward_candidate_test.csv"
        empty = pd.DataFrame(
            [
                {
                    "selected_reward": "",
                    "run_status": "skipped",
                    "notes": "No HLA candidate satisfied saturation and yield-loss criteria.",
                }
            ]
        )
        empty.to_csv(out, index=False, encoding="utf-8-sig")
        return empty

    for station in ["SYA", "LCA"]:
        train_year = int(config["training"]["best_train_years"][station])
        print(f"[reward-cost-debug] Extension: {station} {train_year} {selected_reward}", flush=True)
        run_training(config, selected_reward, station, train_year)
    df = enrich_evaluation_summary()
    copy_candidate_figures(df)
    cross = df[df["station"].isin(["SYA", "LCA"]) & (df["reward_version"] == selected_reward)].copy()
    out = OUTPUT_ROOT / "evaluation" / "cross_site_reward_candidate_test.csv"
    cross.to_csv(out, index=False, encoding="utf-8-sig")
    return cross


def write_report(comparison: pd.DataFrame, cross: pd.DataFrame, selected_reward: str | None) -> None:
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)
    table = comparison.copy()
    numeric_cols = table.select_dtypes(include=["float", "int"]).columns
    table[numeric_cols] = table[numeric_cols].round(4)
    cross_table = cross.copy()
    if len(cross_table):
        numeric_cols = cross_table.select_dtypes(include=["float", "int"]).columns
        cross_table[numeric_cols] = cross_table[numeric_cols].round(4)
    selected_text = selected_reward or "None"
    passed = comparison[comparison["recommendation_pass"]]["reward_version"].tolist()
    md = f"""# Reward cost revision and debug report

Generated at: 2026-06-06

## Why revise reward

The 006_02 season-cap sensitivity analysis showed that action-safe PPO filled every tested cap from 100/150 to 300/450. This means the safety wrapper prevents unlimited actions, but the reward itself still encourages using all allowed water and nitrogen.

## Current reward issue

- Current `all_reward` combines nitrogen recovery reward and biomass-growth irrigation reward.
- It includes daily action costs, but no default nonlinear seasonal penalty and no explicit terminal grain-yield economic return.
- PPO therefore treats the season cap as the effective budget and learns to exhaust it.

## Candidate design

- A candidates add linear irrigation and nitrogen costs.
- B candidates add linear costs plus quadratic penalties after 200 mm irrigation and 300 kg/ha nitrogen.
- C1 approximates an economic reward by using positive daily grain-yield increments minus input costs.

## HLA pilot setup

- Station: HLA
- Train year: 2011
- Eval years: 2011, 2007, 2009
- Action safety cap: 300 mm irrigation / 450 kg ha-1 N
- Timesteps: 5000
- Seed: 0

## HLA reward candidate comparison

{df_to_markdown(table)}

## Candidates passing initial filter

{', '.join(passed) if passed else 'No candidate passed the initial filter.'}

Initial filter:

- mean irrigation saturation ratio < 0.95
- mean nitrogen saturation ratio < 0.95
- yield loss vs current reward <= 10%
- all HLA eval episodes passed quality gate

## Recommended candidate

Recommended reward for next debug stage: `{selected_text}`.

## SYA/LCA extension

{df_to_markdown(cross_table) if len(cross_table) else 'Extension was skipped because HLA did not identify a passing candidate.'}

## Interpretation

Action safety cap and reward cost are different mechanisms. The cap is a safety valve. Reward cost is the signal that should make PPO internalize water and nitrogen scarcity. If a candidate no longer fills 300/450 while keeping yield loss below 10%, it is a better debug reward for later multi-seed tests. This is still not the final paper reward; the next version should use interpretable water, nitrogen, and grain price parameters.

## Next step

Proceed to multi-seed only if the selected candidate remains below cap on HLA and the SYA/LCA extension is acceptable. If the candidate still saturates or yield collapses, tune the cost coefficients before running rainfall-scaling budget scenarios.
"""
    DOC_MD.write_text(md, encoding="utf-8")
    shutil.copy2(DOC_MD, OUTPUT_ROOT / "reports" / DOC_MD.name)
    write_ppt(comparison, selected_reward)


def write_ppt(comparison: pd.DataFrame, selected_reward: str | None) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        note = OUTPUT_ROOT / "reports" / "ppt_generation_skipped_in_container.txt"
        note.write_text("python-pptx is not installed in the DSSAT container. Generate the PPT from Windows Python after training.", encoding="utf-8")
        return

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_title(slide, title: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.5))
        tf = box.text_frame
        tf.text = title
        p = tf.paragraphs[0]
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True

    def add_body(slide, text: str, top: float = 0.95) -> None:
        box = slide.shapes.add_textbox(Inches(0.65), Inches(top), Inches(12.0), Inches(5.8))
        tf = box.text_frame
        tf.word_wrap = True
        for idx, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(16)

    def add_picture_slide(title: str, path: Path) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, title)
        if path.exists():
            slide.shapes.add_picture(str(path), Inches(0.8), Inches(1.0), width=Inches(11.8))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Reward Cost Revision Debug")
    add_body(
        slide,
        "目标：解释为什么 current reward 会让 PPO 打满 cap，并测试带水氮成本项的 reward candidate。\n"
        "约束：只做 action-safe PPO；不改原始 reward；不做多 seed；不覆盖旧结果。\n"
        "HLA pilot：train 2011，eval 2011/2007/2009，cap=300 mm irrigation / 450 kg ha-1 N。",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Current Reward Problem")
    add_body(
        slide,
        "上一阶段 cap sensitivity 中，五站点所有 cap 均被打满。\n"
        "这说明安全上限只是外部阀门，reward 本身仍鼓励使用全部可用水氮。\n"
        "当前 reward 缺少默认季节总量成本、非线性超量惩罚和显式终局产量收益。",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Reward Candidate Families")
    add_body(
        slide,
        "A：current reward + 线性灌溉/施氮成本。\n"
        "B：A + 超过 200 mm / 300 kg ha-1 后的二次惩罚。\n"
        "C：用日增籽粒产量代理 terminal yield benefit，再扣除水氮成本。\n"
        "注意：C1 是 debug 近似，因为 gym reward 回调没有显式 done 标志。",
    )

    summary_dir = OUTPUT_ROOT / "figures" / "summary"
    add_picture_slide("Yield vs Input", summary_dir / "reward_candidates_yield_vs_input.png")
    add_picture_slide("Saturation Ratio", summary_dir / "reward_candidates_saturation_ratio.png")
    add_picture_slide("Reward vs Yield", summary_dir / "reward_candidates_reward_vs_yield.png")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Recommendation")
    passed = comparison[comparison["recommendation_pass"]]["reward_version"].tolist()
    add_body(
        slide,
        f"通过初筛的 candidate：{', '.join(passed) if passed else '无'}\n"
        f"推荐进入下一阶段：{selected_reward or '暂无'}\n"
        "如果推荐项在 SYA/LCA 也不打满 300/450 且产量损失可控，可以进入多 seed；否则继续调成本系数。",
    )

    DOC_PPT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DOC_PPT)
    shutil.copy2(DOC_PPT, OUTPUT_ROOT / "reports" / DOC_PPT.name)


def main() -> None:
    ensure_dirs()
    config = load_yaml(CONFIG_PATH)
    write_candidate_files()
    build_plan(config)
    try:
        df, comparison, selected = run_hla_pilot(config)
        cross = run_extension(config, selected)
        comparison = compare_candidates(enrich_evaluation_summary())
        plot_summary(enrich_evaluation_summary(), comparison)
        write_report(comparison, cross, selected)
    except Exception:
        err = OUTPUT_ROOT / "reports" / "reward_cost_debug_error.log"
        err.write_text(traceback.format_exc(), encoding="utf-8")
        print(traceback.format_exc(), file=sys.stderr)
        raise


if __name__ == "__main__":
    main()
