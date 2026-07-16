from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PPT_PATH = next(
    path for path in ROOT.glob("*codex.pptx")
    if not path.name.startswith("~$") and "-building" not in path.stem
)
OUTPUT_PATH = ROOT / "杜雨秋-开题-codex-全稿美化.pptx"
MEDIA = ROOT / "ppt_work" / "extracted_media"
WORK = ROOT / "ppt_work"

NAVY = RGBColor(25, 74, 150)
BLUE = RGBColor(0, 112, 192)
LIGHT_BLUE = RGBColor(229, 240, 250)
PALE_BLUE = RGBColor(242, 247, 252)
GREEN = RGBColor(51, 139, 87)
LIGHT_GREEN = RGBColor(232, 244, 231)
ORANGE = RGBColor(224, 145, 48)
LIGHT_ORANGE = RGBColor(252, 242, 226)
RED = RGBColor(190, 72, 60)
DARK = RGBColor(31, 31, 31)
MID = RGBColor(90, 100, 112)
LINE = RGBColor(205, 214, 224)
WHITE = RGBColor(255, 255, 255)


def clear_slide(slide):
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)


def set_background(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def set_text_style(run, size=16, bold=False, color=DARK, font="微软雅黑"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=16,
    bold=False,
    color=DARK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
    font="微软雅黑",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    run = p.add_run()
    run.text = text
    set_text_style(run, size, bold, color, font)
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=15, color=DARK, bullet=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if bullet else "") + line
        p.space_after = Pt(7)
        p.line_spacing = 1.12
        for run in p.runs:
            set_text_style(run, size, False, color)
    return box


def add_header(slide, title, section=None):
    if section:
        add_text(slide, section, 0.45, 0.16, 0.66, 0.34, 12, True, WHITE,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(0.14), Inches(0.72), Inches(0.38))
        badge.fill.solid(); badge.fill.fore_color.rgb = NAVY
        badge.line.fill.background()
        slide.shapes._spTree.remove(slide.shapes._spTree[-2])
        slide.shapes._spTree.insert(2, badge._element)
    title_size = 18.5 if len(title) > 28 else (20.5 if len(title) > 22 else 23)
    add_text(slide, title, 1.22 if section else 0.48, 0.12, 8.25 if section else 9.0, 0.48,
             title_size, True, NAVY, valign=MSO_ANCHOR.MIDDLE, margin=0)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.72), Inches(9.1), Inches(0.035))
    line.fill.solid(); line.fill.fore_color.rgb = NAVY; line.line.fill.background()


def add_box(slide, x, y, w, h, title, body=None, fill=PALE_BLUE, line=NAVY, title_color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line; shape.line.width = Pt(1.1)
    add_text(slide, title, x + 0.12, y + 0.10, w - 0.24, 0.34, 16, True, title_color,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    if body:
        add_text(slide, body, x + 0.18, y + 0.53, w - 0.36, h - 0.63, 13.5, False, DARK,
                 PP_ALIGN.LEFT, MSO_ANCHOR.TOP, margin=0)
    return shape


def add_arrow(slide, x, y, w=0.45, h=0.34, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_picture_contain(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                    Inches(pw), Inches(ph))


def add_source(slide, text):
    add_text(slide, text, 0.48, 7.12, 9.05, 0.22, 7.5, False, MID,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def build_outline_slide(slide, active_index):
    clear_slide(slide)
    set_background(slide, WHITE)
    sections = [
        ("01", "研究背景及意义"),
        ("02", "国内外研究现状"),
        ("03", "研究目标与内容"),
        ("04", "总体思路与技术方法"),
        ("05", "研究基础及条件"),
        ("06", "研究进度安排"),
    ]
    add_text(slide, "汇报大纲", 0.58, 0.48, 3.2, 0.72, 30, True, NAVY,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "RESEARCH PROPOSAL", 6.75, 0.63, 2.60, 0.28, 10, True, MID,
             PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE, margin=0)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(1.30), Inches(8.78), Inches(0.035))
    line.fill.solid(); line.fill.fore_color.rgb = NAVY; line.line.fill.background()
    for i, (num, title) in enumerate(sections):
        row, col = divmod(i, 2)
        x = 0.66 + col * 4.55
        y = 1.75 + row * 1.56
        active = i == active_index
        fill = NAVY if active else PALE_BLUE
        border = NAVY if active else LINE
        color = WHITE if active else DARK
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(4.05), Inches(1.04))
        box.fill.solid(); box.fill.fore_color.rgb = fill
        box.line.color.rgb = border; box.line.width = Pt(1.0)
        add_text(slide, num, x + 0.20, y + 0.21, 0.58, 0.54, 16, True,
                 WHITE if active else NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x + 0.92, y + 0.18, 2.90, 0.62, 18, active, color,
                 PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
        if active:
            marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.10), Inches(1.04))
            marker.fill.solid(); marker.fill.fore_color.rgb = ORANGE; marker.line.fill.background()
    add_text(slide, "当前部分", 0.66, 6.64, 1.02, 0.30, 11, True, MID,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, sections[active_index][1], 1.60, 6.56, 3.8, 0.42, 15, True, NAVY,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)


def make_mitigation_crops():
    paths = [
        MEDIA / "doc_image9_bb9c28cb.png",
        MEDIA / "doc_image10_cdf40117.png",
    ]
    outputs = []
    for path, name in zip(paths, ["maize_mitigation_selected.png", "wheat_mitigation_selected.png"]):
        out = WORK / name
        with Image.open(path) as im:
            # Retain the global bar, mitigation map, labels and legend.
            crop = im.crop((0, int(im.height * 0.37), im.width, int(im.height * 0.67)))
            crop.save(out)
        outputs.append(out)
    return outputs


def build_deck():
    prs = Presentation(str(PPT_PATH))
    slides = prs.slides
    maize_crop, wheat_crop = make_mitigation_crops()

    # Slides 1-18: rebuild the opening half with the same visual system.
    slide = slides[0]; clear_slide(slide); set_background(slide, WHITE)
    cover_mark = MEDIA / "ppt_s01_p01_7e0bc7c8.png"
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(1.15), Inches(7.5))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.15), Inches(0), Inches(0.10), Inches(7.5))
    accent.fill.solid(); accent.fill.fore_color.rgb = GREEN; accent.line.fill.background()
    add_picture_contain(slide, cover_mark, 1.55, 0.52, 2.45, 1.00)
    add_text(slide, "博士论文开题答辩", 1.62, 1.72, 3.0, 0.46, 17, True, GREEN,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "基于数据驱动与强化学习相结合的\n农田水氮智能管理研究",
             1.60, 2.30, 7.65, 1.62, 30, True, NAVY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.62), Inches(4.18), Inches(1.36), Inches(0.07))
    line.fill.solid(); line.fill.fore_color.rgb = ORANGE; line.line.fill.background()
    add_text(slide, "汇报人：杜雨秋", 1.62, 4.66, 2.35, 0.38, 16, False, DARK,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "指导教师：陶福禄", 4.24, 4.66, 2.75, 0.38, 16, False, DARK,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "自然地理教研室（陆表）", 1.62, 5.24, 3.40, 0.36, 14, False, MID,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "2026年7月17日", 1.62, 6.36, 2.15, 0.38, 14, True, NAVY,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "中国科学院地理科学与资源研究所", 5.70, 6.36, 3.40, 0.38, 12, False, MID,
             PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE, margin=0)
    set_notes(slide, "尊敬的各位老师，上午好，我是2023级博士生杜雨秋，研究方向为全球变化生态学，导师为陶福禄研究员。我的开题题目是基于数据驱动与强化学习相结合的农田水氮智能管理研究。")

    build_outline_slide(slides[1], 0)

    # Slide 3: nitrogen supports food production but is inefficiently used.
    slide = slides[2]; clear_slide(slide); add_header(slide, "氮肥支撑粮食生产，但超过一半未被作物有效利用", "1.1")
    cycle_img = MEDIA / "ppt_s03_p03_e6b31f49.png"
    add_picture_contain(slide, cycle_img, 0.42, 1.12, 6.12, 4.90)
    add_box(slide, 6.78, 1.20, 2.65, 1.55, "粮食生产", "合成氮肥是现代农业增产的重要基础", LIGHT_GREEN, GREEN, GREEN)
    metric = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.78), Inches(3.08), Inches(2.65), Inches(1.58))
    metric.fill.solid(); metric.fill.fore_color.rgb = NAVY; metric.line.fill.background()
    add_text(slide, "> 50%", 6.95, 3.22, 2.30, 0.72, 28, True, WHITE,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "肥料氮未进入作物或土壤稳定库", 6.98, 4.00, 2.24, 0.42, 12.5, False, WHITE,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    add_box(slide, 6.78, 4.98, 2.65, 1.10, "环境代价", "气态排放 + 水文流失", LIGHT_ORANGE, ORANGE, ORANGE)
    add_text(slide, "核心矛盾：投入不断增加，但作物吸收与环境容量有限",
             0.72, 6.37, 8.60, 0.46, 15.5, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_source(slide, "Source: Gruber & Galloway, Nature 451, 293–296 (2008). https://doi.org/10.1038/nature06592")
    set_notes(slide, "氮肥对粮食生产不可替代，但氮投入与作物需求之间存在明显错配。全球农业系统中超过一半的肥料氮没有被作物有效利用，而是进入环境。")

    # Slide 4: environmental consequences.
    slide = slides[3]; clear_slide(slide); add_header(slide, "农业氮损失同时影响气候、空气、水体与生物多样性", "1.1")
    impacts = [
        ("N₂O", "气候变化", "强温室气体\n农业是人为排放的重要来源", RGBColor(249, 235, 234), RED),
        ("NH₃ / NOx", "空气污染", "形成颗粒物与臭氧\n影响人体健康", LIGHT_ORANGE, ORANGE),
        ("淋失 / 径流", "水体富营养化", "硝态氮进入地下水、\n河流与沿海水体", LIGHT_BLUE, BLUE),
        ("Nr累积", "生态系统退化", "土壤酸化、群落变化\n与生物多样性下降", LIGHT_GREEN, GREEN),
    ]
    for i, (chem, title, body, fill, accent) in enumerate(impacts):
        x = 0.42 + i * 2.40
        add_box(slide, x, 1.42, 2.10, 3.92, title, body, fill, accent, accent)
        add_text(slide, chem, x + 0.30, 1.74, 1.50, 0.64, 20, True, accent,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.40), Inches(4.82), Inches(1.30), Inches(0.06))
        bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    add_text(slide, "同一单位氮损失可沿多条路径产生跨介质、跨尺度环境效应",
             0.70, 6.13, 8.60, 0.55, 16, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "氮损失不是单一温室气体问题。气态路径影响气候和空气质量，水文路径导致地下水和地表水污染，长期活性氮累积还会破坏生态系统稳定性。")

    # Slide 5: F-Nr concept.
    slide = slides[4]; clear_slide(slide); add_header(slide, "F-Nr统一表征单位氮投入经不同路径损失的比例", "1.1")
    add_text(slide, "F-Nr", 0.72, 1.25, 2.25, 0.88, 32, True, NAVY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, "Fraction of Nitrogen Released", 0.72, 2.08, 2.25, 0.42, 12, False, MID,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    formula = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.10), Inches(1.24), Inches(6.12), Inches(1.38))
    formula.fill.solid(); formula.fill.fore_color.rgb = PALE_BLUE; formula.line.color.rgb = NAVY
    add_text(slide, "F-Nrᵣ = 特定路径 r 的氮损失量 / 氮投入量 × 100%",
             3.34, 1.55, 5.64, 0.70, 21, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, "统一核算五类主要路径", 0.70, 3.18, 2.55, 0.42, 16, True, DARK,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    pathways = [("N₂O", RED), ("NH₃", ORANGE), ("NO", GREEN), ("N leaching", BLUE), ("N runoff", NAVY)]
    for i, (label, color) in enumerate(pathways):
        x = 0.62 + i * 1.83
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(3.92), Inches(1.25), Inches(1.25))
        c.fill.solid(); c.fill.fore_color.rgb = color; c.line.fill.background()
        add_text(slide, label, x, 3.92, 1.25, 1.25, 14, True, WHITE,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "优势：不同地区、不同作物和不同损失路径可以在统一尺度下比较",
             0.72, 5.92, 8.55, 0.56, 15.5, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "本研究使用F-Nr统一表示单位氮投入通过某一具体路径损失到环境中的比例，并分别对五种主要损失路径开展建模和制图。")

    # Slide 6: spatial uncertainty.
    slide = slides[5]; clear_slide(slide); add_header(slide, "固定排放因子难以解释氮损失的强烈时空异质性", "1.1")
    sutton_img = MEDIA / "ppt_s06_p03_b5de6ab2.png"
    add_picture_contain(slide, sutton_img, 4.52, 1.12, 4.92, 4.92)
    drivers = [
        ("气候", "温度、降水与极端事件", GREEN),
        ("土壤", "质地、pH、有机碳与水分", ORANGE),
        ("管理", "氮投入、肥料类型、耕作与灌溉", BLUE),
        ("水文", "淋失、径流与间接N₂O形成", NAVY),
    ]
    for i, (title, body, color) in enumerate(drivers):
        y = 1.20 + i * 1.22
        add_box(slide, 0.50, y, 3.58, 0.92, title, body, WHITE, color, color)
    add_text(slide, "固定全球因子 → 难以识别热点与区域差异 → 减排措施缺乏针对性",
             0.68, 6.30, 8.64, 0.48, 15.5, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_source(slide, "Source: Sutton et al., Our Nutrient World (2013), nitrogen cascade framework.")
    set_notes(slide, "氮损失由气候、土壤、管理和水文过程共同控制，尤其是淋失与径流具有高度时空异质性。固定排放因子很难反映这些差异。")

    # Slide 7: maize and wheat.
    slide = slides[6]; clear_slide(slide); add_header(slide, "玉米和小麦是全球氮投入与损失控制的关键作物", "1.1")
    nw_img = MEDIA / "ppt_s07_p02_d13e5ff7.png"
    add_picture_contain(slide, nw_img, 0.40, 1.08, 6.45, 5.52)
    add_box(slide, 7.02, 1.22, 2.38, 1.45, "氮肥投入", "两种作物合计消耗全球作物氮肥约38%", LIGHT_GREEN, GREEN, GREEN)
    add_box(slide, 7.02, 3.02, 2.38, 1.45, "氮损失", "覆盖气态排放、淋失与径流等多条路径", LIGHT_BLUE, BLUE, BLUE)
    add_box(slide, 7.02, 4.82, 2.38, 1.45, "研究价值", "种植广、投入高，适合作物特异的全球比较", LIGHT_ORANGE, ORANGE, ORANGE)
    add_source(slide, "Source: Wang et al., Nature Geoscience 17, 1008–1015 (2024). https://doi.org/10.1038/s41561-024-01542-x")
    set_notes(slide, "玉米和小麦既是粮食安全的核心作物，也是全球氮肥投入和活性氮损失的重要来源，因此本研究以这两类旱地谷物为对象开展作物特异建模。")

    # Slide 8: method trade-offs.
    slide = slides[7]; clear_slide(slide); add_header(slide, "现有方法在观测真实性、空间覆盖与过程表达之间存在权衡", "1.1")
    field_img = MEDIA / "ppt_s08_p04_958097c7.png"
    dssat_img = MEDIA / "ppt_s08_p03_a051dc72.png"
    add_picture_contain(slide, field_img, 0.50, 1.28, 2.65, 2.20)
    add_picture_contain(slide, dssat_img, 6.82, 1.28, 2.65, 2.20)
    add_box(slide, 0.50, 3.72, 2.65, 1.85, "田间观测", "测量可靠、机制直接\n但成本高、空间覆盖有限", LIGHT_GREEN, GREEN, GREEN)
    add_box(slide, 3.68, 1.28, 2.64, 4.29, "统计与回归", "计算效率较高\n\n依赖有限变量与样本\n难以刻画复杂非线性\n和空间异质性", PALE_BLUE, NAVY, NAVY)
    add_box(slide, 6.82, 3.72, 2.65, 1.85, "过程模型", "可表达水氮过程\n但数据、参数与计算成本较高", LIGHT_ORANGE, ORANGE, ORANGE)
    add_text(slide, "需要一种能够整合多源数据、适用于全球空间外推且可解释的方法",
             0.72, 6.26, 8.55, 0.50, 15.5, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "田间观测真实但覆盖有限，经验模型简单但难以表达复杂关系，过程模型机制清晰但全球应用成本高。因此需要数据驱动方法补充全球尺度诊断能力。")

    # Slide 9: machine learning opportunity.
    slide = slides[8]; clear_slide(slide); add_header(slide, "机器学习能够融合多源数据并捕捉非线性空间差异", "1.1")
    rf_img = MEDIA / "ppt_s09_p04_9ce4f809.png"
    add_picture_contain(slide, rf_img, 4.48, 1.18, 4.85, 3.70)
    datasets = [("气候", GREEN), ("土壤", ORANGE), ("地形", NAVY), ("植被", BLUE), ("管理", RED)]
    for i, (label, color) in enumerate(datasets):
        y = 1.22 + i * 0.87
        add_box(slide, 0.55, y, 2.12, 0.66, label, None, WHITE, color, color)
        add_arrow(slide, 2.84, y + 0.17, 0.46, 0.28, color)
    add_box(slide, 3.36, 2.00, 0.92, 1.98, "RF", None, LIGHT_BLUE, NAVY, NAVY)
    add_box(slide, 0.55, 5.30, 2.52, 1.00, "空间预测", "高分辨率F-Nr地图", LIGHT_GREEN, GREEN, GREEN)
    add_box(slide, 3.72, 5.30, 2.52, 1.00, "模型解释", "变量贡献与响应方向", LIGHT_BLUE, BLUE, BLUE)
    add_box(slide, 6.89, 5.30, 2.52, 1.00, "决策支撑", "热点与区域优先级", LIGHT_ORANGE, ORANGE, ORANGE)
    set_notes(slide, "随机森林适合处理高维、多源和非线性数据。它不仅可以生成全球高分辨率预测，还可以结合解释方法识别关键驱动因素。")

    # Slide 10: the remaining decision gap.
    slide = slides[9]; clear_slide(slide); add_header(slide, "全球制图能够诊断问题，但不能直接生成逐日管理动作", "1.1")
    add_box(slide, 0.55, 1.34, 3.25, 3.85, "全球尺度诊断", "氮损失发生在哪里？\n\n哪些气候、土壤和管理因素\n控制不同损失路径？\n\n不同措施可减排多少？", LIGHT_GREEN, GREEN, GREEN)
    add_arrow(slide, 4.08, 3.05, 0.72, 0.50, NAVY)
    add_box(slide, 4.97, 1.34, 4.45, 3.85, "田间尺度决策", "今天是否需要灌溉或施氮？\n\n在多少投入下兼顾产量、收益、\nWUE、NUE和环境效益？\n\n面对天气变化如何动态调整？", LIGHT_BLUE, BLUE, BLUE)
    add_text(slide, "研究转折：从“识别氮损失”进一步走向“动态优化水氮管理”",
             0.70, 5.83, 8.60, 0.66, 18, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "全球制图解决的是诊断问题，但农业管理需要在每日尺度回答何时投入、投入多少以及如何根据作物和天气状态调整，因此需要进一步引入动态决策方法。")

    # Slide 11: water-nitrogen coupling.
    slide = slides[10]; clear_slide(slide); add_header(slide, "水分调节氮吸收、迁移与损失，水氮管理必须协同优化", "1.1")
    center = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.72), Inches(2.02), Inches(2.56), Inches(2.56))
    center.fill.solid(); center.fill.fore_color.rgb = NAVY; center.line.fill.background()
    add_text(slide, "作物—土壤\n水氮耦合", 3.72, 2.02, 2.56, 2.56, 20, True, WHITE,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    nodes = [
        (0.50, 1.20, "作物吸收", "根系活性与氮需求", LIGHT_GREEN, GREEN),
        (0.50, 4.47, "水分供应", "降水、灌溉与蒸散", LIGHT_BLUE, BLUE),
        (7.04, 1.20, "氮素转化", "矿化、硝化与反硝化", LIGHT_ORANGE, ORANGE),
        (7.04, 4.47, "环境损失", "气态排放、淋失与径流", RGBColor(249, 235, 234), RED),
    ]
    for x, y, title, body, fill, color in nodes:
        add_box(slide, x, y, 2.45, 1.45, title, body, fill, color, color)
    add_text(slide, "灌溉会改变土壤水分、氮迁移和作物需求；施氮也会影响产量与环境风险",
             0.74, 6.28, 8.52, 0.50, 15, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "水分和氮素并不是两个独立管理变量。水分影响根系吸收、微生物过程和氮迁移，因此仅优化施氮或仅优化灌溉都可能得到次优方案。")

    # Slide 12: reinforcement learning.
    slide = slides[11]; clear_slide(slide); add_header(slide, "强化学习适合处理连续反馈与长期收益驱动的农田决策", "1.1")
    add_box(slide, 0.60, 1.58, 2.35, 2.42, "智能体 Agent", "观测作物、土壤与天气\n选择灌溉和施氮动作", LIGHT_BLUE, BLUE, BLUE)
    add_arrow(slide, 3.16, 2.48, 0.68, 0.45, NAVY)
    add_box(slide, 4.05, 1.58, 2.35, 2.42, "环境 Environment", "DSSAT推进作物—土壤过程\n返回下一状态与产出", LIGHT_GREEN, GREEN, GREEN)
    add_arrow(slide, 6.60, 2.48, 0.68, 0.45, NAVY)
    add_box(slide, 7.50, 1.58, 1.90, 2.42, "奖励 Reward", "产量/收益\n资源成本\n环境损失", LIGHT_ORANGE, ORANGE, ORANGE)
    back = slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(2.25), Inches(4.55), Inches(5.55), Inches(0.46))
    back.fill.solid(); back.fill.fore_color.rgb = MID; back.line.fill.background()
    add_text(slide, "根据长期累计奖励不断更新策略", 2.50, 4.55, 5.05, 0.46, 14, True, WHITE,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "优势：无需预先枚举全部管理组合，可根据状态变化学习动态策略",
             0.72, 5.98, 8.56, 0.56, 16, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "强化学习通过智能体与环境交互学习策略，目标是最大化整个生育期的累计收益，因此适合具有动态变化、延迟反馈和多目标权衡的农田管理问题。")

    build_outline_slide(slides[12], 1)

    # Slide 14: evolution of nitrogen-loss estimation.
    slide = slides[13]; clear_slide(slide); add_header(slide, "氮损失估算正从固定因子走向空间显式与数据驱动", "2.1")
    stages = [
        ("Tier 1", "全球固定因子", "简单可推广\n但不确定性高", MID),
        ("Tier 2", "国家/区域因子", "区域适配增强\n仍依赖有限观测", GREEN),
        ("Tier 3", "过程模型", "机制表达充分\n数据与参数成本高", ORANGE),
        ("机器学习", "空间显式预测", "融合多源变量\n捕捉非线性异质性", BLUE),
    ]
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.90), Inches(3.18), Inches(8.20), Inches(0.06))
    line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background()
    for i, (tag, title, body, color) in enumerate(stages):
        x = 0.52 + i * 2.40
        add_box(slide, x, 1.32, 2.05, 1.44, tag, title, WHITE, color, color)
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.80), Inches(2.93), Inches(0.44), Inches(0.44))
        node.fill.solid(); node.fill.fore_color.rgb = color; node.line.fill.background()
        add_box(slide, x, 3.72, 2.05, 1.60, "特点", body, PALE_BLUE, color, color)
    add_text(slide, "发展方向：在保持空间预测能力的同时，提高解释性并支撑管理决策",
             0.70, 6.15, 8.60, 0.54, 15.5, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "氮损失估算从全球固定排放因子逐步发展到国家因子、过程模型和机器学习。当前关键不是单纯追求更复杂模型，而是同时兼顾空间预测、解释与管理价值。")

    # Slide 15: ML literature status and gap.
    slide = slides[14]; clear_slide(slide); add_header(slide, "机器学习提升了空间预测能力，但机制解释与管理转化仍不足", "2.2")
    add_box(slide, 0.55, 1.28, 4.05, 4.62, "已有进展", "• 广泛应用于作物产量和氮状态预测\n\n• 可融合遥感、气候与土壤等高维数据\n\n• RF、GBM和深度学习能够表达非线性关系\n\n• 区域与国家尺度空间预测不断增加", LIGHT_GREEN, GREEN, GREEN)
    add_box(slide, 5.00, 1.28, 4.45, 4.62, "仍待解决", "• 全球作物特异氮损失研究仍有限\n\n• 动态植被、极端气候与水文变量考虑不足\n\n• 预测贡献与因果机制容易混淆\n\n• 空间识别结果难以直接转化为逐日管理动作", LIGHT_ORANGE, ORANGE, ORANGE)
    add_text(slide, "本研究的切入点：作物特异全球制图 + SHAP/SEM机制解析 + 管理情景评估",
             0.68, 6.25, 8.64, 0.52, 15.5, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "已有机器学习农业研究很多，但真正针对全球氮损失、多路径、作物特异和机制解释的工作仍有限，并且预测结果通常没有进一步连接到动态管理。")

    # Slide 16: crop models.
    slide = slides[15]; clear_slide(slide); add_header(slide, "作物模型能够描述水氮过程，但传统情景遍历难以动态决策", "2.3")
    add_picture_contain(slide, dssat_img, 0.46, 1.16, 4.15, 4.98)
    add_box(slide, 4.90, 1.25, 4.40, 1.62, "过程优势", "模拟作物生长、土壤水分平衡与氮循环\n可评价灌溉和施肥对产量及环境的影响", LIGHT_GREEN, GREEN, GREEN)
    add_box(slide, 4.90, 3.16, 4.40, 1.62, "优化局限", "依赖预设情景和大量组合模拟\n环境变化后通常需要重新设计方案", LIGHT_ORANGE, ORANGE, ORANGE)
    add_box(slide, 4.90, 5.07, 4.40, 1.12, "需要突破", "从离线情景比较转向基于状态的动态决策", LIGHT_BLUE, BLUE, BLUE)
    add_source(slide, "Source: DSSAT framework adapted from Jones et al., European Journal of Agronomy 18, 235–265 (2003).")
    set_notes(slide, "DSSAT等作物模型能够提供可信的作物和土壤过程，但传统优化通常靠预设管理组合反复模拟，无法自然地根据每日状态和天气信息调整动作。")

    # Slide 17: crop models + RL.
    slide = slides[16]; clear_slide(slide); add_header(slide, "作物模型与强化学习开始融合，但水氮协同和泛化验证仍薄弱", "2.4")
    platforms = [("CropGym", "作物氮管理环境", GREEN), ("CyclesGym", "长期农艺策略学习", ORANGE), ("gym-DSSAT", "高保真作物模型接口", BLUE)]
    for i, (name, desc, color) in enumerate(platforms):
        x = 0.55 + i * 3.16
        add_box(slide, x, 1.36, 2.85, 1.66, name, desc, WHITE, color, color)
        if i < 2:
            add_arrow(slide, x + 2.92, 2.02, 0.22, 0.28, NAVY)
    add_box(slide, 1.30, 3.50, 2.10, 1.18, "作物模型", "提供状态转移与产量响应", LIGHT_GREEN, GREEN, GREEN)
    add_arrow(slide, 3.60, 3.90, 0.54, 0.34, NAVY)
    add_box(slide, 4.34, 3.50, 2.10, 1.18, "强化学习", "基于长期奖励更新管理策略", LIGHT_BLUE, BLUE, BLUE)
    add_arrow(slide, 6.64, 3.90, 0.54, 0.34, NAVY)
    add_box(slide, 7.38, 3.50, 1.72, 1.18, "动态动作", "灌溉 + 施氮", LIGHT_ORANGE, ORANGE, ORANGE)
    add_text(slide, "现有不足", 0.65, 5.18, 1.20, 0.34, 15, True, RED,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    add_text(slide, "单一作物/单一措施较多  ·  多站点跨年份泛化不足  ·  天气预报信息价值缺少系统评估",
             1.70, 5.08, 7.65, 0.56, 14.5, True, DARK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    add_text(slide, "研究空缺：融合DSSAT、天气预报与DQN的多站点水氮联合决策",
             0.70, 6.05, 8.60, 0.50, 16, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_source(slide, "Platforms: Overweg et al. (2021), CropGym; Turchetta et al. (2022), CyclesGym; Gautron et al. (2022), gym-DSSAT.")
    set_notes(slide, "CropGym、CyclesGym和gym-DSSAT表明作物模型可以与强化学习结合，但当前研究多集中于单一作物或单一管理措施，对天气信息、水氮协同和跨站点泛化关注不足。")

    build_outline_slide(slides[17], 2)
    build_outline_slide(slides[21], 3)
    build_outline_slide(slides[28], 4)

    # Slide 19: objectives.
    slide = slides[18]; clear_slide(slide); add_header(slide, "研究目标", "3.1")
    add_text(slide, "构建“全球氮损失诊断—区域水氮决策—跨环境验证”的农田智能管理体系",
             0.75, 1.02, 8.5, 0.72, 21, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    objectives = [
        ("01", "识别全球氮损失", "量化玉米和小麦各路径氮损失，解析驱动机制与减排潜力", LIGHT_GREEN, GREEN),
        ("02", "构建数字决策环境", "耦合DSSAT与gym-DSSAT，引入未来天气信息形成可交互环境", LIGHT_BLUE, BLUE),
        ("03", "优化并验证水氮策略", "利用DQN联合优化灌溉与施氮，检验多站点、多年份泛化能力", LIGHT_ORANGE, ORANGE),
    ]
    for i, (num, title, body, fill, accent) in enumerate(objectives):
        x = 0.55 + i * 3.18
        add_box(slide, x, 2.12, 2.86, 3.55, title, body, fill, accent, accent)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 1.08), Inches(1.78), Inches(0.70), Inches(0.70))
        circle.fill.solid(); circle.fill.fore_color.rgb = accent; circle.line.fill.background()
        add_text(slide, num, x + 1.08, 1.78, 0.70, 0.70, 15, True, WHITE,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "核心落点：在保障产量的同时，提高WUE与NUE并降低水氮投入和氮损失风险",
             0.75, 6.18, 8.5, 0.54, 15, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "本研究不是把机器学习、作物模型和强化学习简单叠加，而是围绕三个递进目标展开：先识别全球氮损失问题，再构建区域尺度可交互环境，最终形成并验证水氮联合决策策略。")

    # Slide 20: research content 1.
    slide = slides[19]; clear_slide(slide); add_header(slide, "研究内容一：全球氮损失格局、驱动机制与减排潜力", "3.2")
    steps = [
        ("多源数据", "气候·土壤·地形\n植被·管理·观测", LIGHT_GREEN, GREEN),
        ("机器学习制图", "作物特异RF模型\n5 arcmin全球预测", LIGHT_BLUE, BLUE),
        ("机制解析", "SHAP贡献归因\nSEM因果路径", LIGHT_ORANGE, ORANGE),
        ("缓解情景", "提高NUE\n采用EEFs", RGBColor(249, 235, 234), RED),
    ]
    for i, (title, body, fill, accent) in enumerate(steps):
        x = 0.45 + i * 2.40
        add_box(slide, x, 1.34, 2.03, 2.05, title, body, fill, accent, accent)
        if i < 3: add_arrow(slide, x + 2.08, 2.15, 0.28, 0.30, NAVY)
    add_text(slide, "五类活性氮损失路径", 0.65, 4.03, 2.25, 0.38, 16, True, NAVY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    pathways = [("N₂O", RED), ("NH₃", ORANGE), ("NO", GREEN), ("淋失", BLUE), ("径流", NAVY)]
    for i, (label, color) in enumerate(pathways):
        x = 0.55 + i * 1.02
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(4.58), Inches(0.82), Inches(0.82))
        c.fill.solid(); c.fill.fore_color.rgb = color; c.line.fill.background()
        add_text(slide, label, x, 4.58, 0.82, 0.82, 14, True, WHITE,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    add_box(slide, 6.05, 4.02, 3.30, 1.85, "输出", "全球空间格局与热点\n关键控制因素及作用路径\n区域差异化减排潜力", PALE_BLUE, NAVY, NAVY)
    add_text(slide, "回答：氮损失发生在哪里、为何发生，以及可减缓多少？",
             0.65, 6.28, 8.7, 0.52, 16, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "第一部分已经形成论文成果。汇报时重点强调三个科学问题：空间格局、控制机制和缓解潜力，而不是展开每个数据集的细节。")

    # Slide 21: research content 2.
    slide = slides[20]; clear_slide(slide); add_header(slide, "研究内容二：融合天气预报的农田水氮联合优化", "3.2")
    add_box(slide, 0.48, 1.26, 2.28, 2.30, "环境状态", "作物生育阶段\n土壤水分与无机氮\n历史气象与未来天气窗口", LIGHT_GREEN, GREEN, GREEN)
    add_arrow(slide, 2.88, 2.18, 0.45, 0.35, NAVY)
    add_box(slide, 3.38, 1.26, 2.28, 2.30, "DQN智能体", "离散水氮联合动作\n经验回放与目标网络\n长期累计收益优化", LIGHT_BLUE, BLUE, BLUE)
    add_arrow(slide, 5.78, 2.18, 0.45, 0.35, NAVY)
    add_box(slide, 6.28, 1.26, 2.95, 2.30, "DSSAT环境响应", "逐日更新作物—土壤过程\n输出产量、WUE、NUE\n以及投入和氮损失", LIGHT_ORANGE, ORANGE, ORANGE)
    add_text(slide, "多站点 × 多年份 × 多信息情景", 0.68, 4.18, 8.65, 0.45, 18, True, NAVY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    labels = ["东北", "黄淮海", "西北", "西南", "华南"]
    for i, label in enumerate(labels):
        x = 0.67 + i * 1.78
        add_box(slide, x, 4.88, 1.47, 0.82, label, None, WHITE, LINE, NAVY)
    add_text(slide, "比较无预报、完美预报及不同预报窗口下策略表现，评估信息价值与跨环境泛化能力",
             0.7, 6.22, 8.6, 0.50, 15, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "第二部分把天气预报作为状态信息的一部分。历史逐日天气在回放实验中代表完美预报，用于估计未来气象信息能够带来的决策收益上界。")

    # Slide 23: preserve the existing overall roadmap and add a title.
    slide = slides[22]
    for shape in list(slide.shapes):
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            slide.shapes._spTree.remove(shape._element)
    picture = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    picture.left = Inches(0)
    picture.top = Inches(0.84)
    picture.width = Inches(10)
    picture.height = Inches(6.22)
    add_text(slide, "4.1 总体技术路线", 0.42, 0.10, 3.10, 0.46, 24, True, NAVY,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=0)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(0.64), Inches(9.1), Inches(0.03))
    line.fill.solid(); line.fill.fore_color.rgb = NAVY; line.line.fill.background()
    set_notes(slide, "总体技术路线包含两个尺度：全球尺度完成氮损失识别、机制解析与缓解评估；区域尺度以五个典型站点构建DSSAT环境，开展强化学习水氮联合优化。")

    # Slide 24: data-driven method.
    slide = slides[23]; clear_slide(slide); add_header(slide, "全球氮损失：从多源观测到可解释空间预测", "4.2")
    flow = [
        ("观测与环境数据", "田间F-Nr观测\n气候·土壤·管理", LIGHT_GREEN, GREEN),
        ("特征工程", "生育期统计\n极端气候与水文变量", PALE_BLUE, BLUE),
        ("作物特异RF", "玉米/小麦分别建模\n重复交叉验证与Bootstrap", LIGHT_BLUE, NAVY),
        ("解释与情景", "SHAP + SEM\nNUE与EEFs情景", LIGHT_ORANGE, ORANGE),
    ]
    for i, item in enumerate(flow):
        x = 0.38 + i * 2.42
        add_box(slide, x, 1.36, 2.08, 2.06, item[0], item[1], item[2], item[3], item[3])
        if i < 3: add_arrow(slide, x + 2.10, 2.18, 0.28, 0.30, NAVY)
    add_text(slide, "模型评价", 0.62, 4.15, 1.55, 0.38, 16, True, NAVY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    for i, metric in enumerate(["R²", "RMSE", "MAE", "空间外推", "不确定性"]):
        x = 0.48 + i * 1.16
        add_box(slide, x, 4.70, 0.98, 0.72, metric, None, WHITE, LINE, NAVY)
    add_box(slide, 6.55, 4.05, 2.78, 1.65, "最终产品", "5 arcmin全球F-Nr地图\n驱动因素空间排序\n缓解潜力与优先区域", LIGHT_GREEN, GREEN, GREEN)
    add_text(slide, "数据驱动模型负责全球诊断，区域决策优化由DSSAT与强化学习进一步承接",
             0.65, 6.25, 8.7, 0.50, 15, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "这一页只讲方法链条。随机森林用于空间预测，SHAP用于变量贡献解释，结构方程模型用于分析直接和间接作用路径，最终通过NUE和EEFs情景评价缓解潜力。")

    # Slide 25: DSSAT/gym environment.
    slide = slides[24]; clear_slide(slide); add_header(slide, "gym-DSSAT将作物过程模型转化为可交互决策环境", "4.3")
    dssat_img = MEDIA / "ppt_s08_p03_a051dc72.png"
    add_picture_contain(slide, dssat_img, 0.42, 1.18, 3.25, 4.98)
    add_arrow(slide, 3.72, 3.20, 0.55, 0.40, NAVY)
    add_box(slide, 4.34, 1.28, 2.10, 1.35, "输入", "气象·土壤·品种\n初始水氮·管理规则", LIGHT_GREEN, GREEN, GREEN)
    add_box(slide, 4.34, 3.00, 2.10, 1.35, "逐日过程", "作物生长\n土壤水分与氮循环", LIGHT_BLUE, BLUE, BLUE)
    add_box(slide, 4.34, 4.72, 2.10, 1.35, "输出状态", "LAI·生物量·物候\n土壤水氮·环境损失", LIGHT_ORANGE, ORANGE, ORANGE)
    add_arrow(slide, 6.55, 3.20, 0.55, 0.40, NAVY)
    add_box(slide, 7.18, 1.55, 2.25, 3.95, "gym交互接口", "观察状态 sₜ\n\n执行动作 aₜ\n灌溉量 + 施氮量\n\n获得奖励 rₜ\n更新到 sₜ₊₁", PALE_BLUE, NAVY, NAVY)
    add_text(slide, "环境角色：提供可信的作物—土壤状态转移，而非替代真实田间验证",
             0.70, 6.38, 8.6, 0.43, 14.5, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_source(slide, "Source: DSSAT framework adapted from Jones et al., Eur. J. Agron. (2003); gym-DSSAT: Gautron et al. (2022), arXiv:2207.03270.")
    set_notes(slide, "DSSAT负责逐日模拟作物生长和土壤水氮过程，gym-DSSAT负责把这些过程封装为强化学习可调用的观察、动作和奖励接口。")

    # Slide 26: DQN decision model.
    slide = slides[25]; clear_slide(slide); add_header(slide, "DQN在离散动作空间中学习长期水氮管理策略", "4.4")
    dqn_img = MEDIA / "ppt_s31_p01_8115580f.png"
    add_picture_contain(slide, dqn_img, 0.65, 1.15, 5.90, 2.75)
    add_box(slide, 6.82, 1.28, 2.45, 1.23, "联合动作", "aₜ =（灌溉等级，施氮等级）", LIGHT_GREEN, GREEN, GREEN)
    add_box(slide, 6.82, 2.76, 2.45, 1.23, "Q值输出", "选择长期累计回报最大的动作", LIGHT_BLUE, BLUE, BLUE)
    add_box(slide, 0.68, 4.55, 2.48, 1.42, "经验回放", "随机采样历史转移\n降低连续样本相关性", PALE_BLUE, NAVY, NAVY)
    add_box(slide, 3.78, 4.55, 2.48, 1.42, "目标网络", "延迟更新目标Q值\n提高训练稳定性", LIGHT_ORANGE, ORANGE, ORANGE)
    add_box(slide, 6.88, 4.55, 2.48, 1.42, "奖励函数", "产量/收益 − 水氮成本\n− 氮损失与违规惩罚", LIGHT_GREEN, GREEN, GREEN)
    add_text(slide, "优化目标不是单季最高产，而是产量、投入效率与环境效益的综合平衡",
             0.75, 6.35, 8.5, 0.42, 15, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_source(slide, "Algorithm basis: Mnih et al., Nature 518, 529–533 (2015); Double DQN: van Hasselt et al. (2015).")
    set_notes(slide, "选择DQN是因为当前水氮动作被离散化。经验回放和目标网络用于提高训练稳定性；奖励函数同时考虑产量、经济收益、水氮投入和环境损失。")

    # Slide 27: forecast information.
    slide = slides[26]; clear_slide(slide); add_header(slide, "历史天气回放构造完美预报，量化未来信息的决策价值", "4.5")
    add_text(slide, "决策日 t", 0.68, 1.22, 1.15, 0.36, 16, True, NAVY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    # Timeline.
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(2.03), Inches(8.05), Inches(0.05))
    line.fill.solid(); line.fill.fore_color.rgb = MID; line.line.fill.background()
    for i in range(9):
        x = 0.90 + i * 1.02
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(1.88), Inches(0.34), Inches(0.34))
        c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 2 else (BLUE if i > 2 else MID)
        c.line.fill.background()
        label = ["t−2", "t−1", "t", "t+1", "t+2", "…", "t+H", "", ""][i]
        if label: add_text(slide, label, x - 0.18, 2.25, 0.70, 0.30, 11, i == 2, DARK,
                           PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    add_box(slide, 0.48, 2.85, 2.65, 2.20, "当前可观测状态", "作物生育阶段与生长状态\n土壤水分和无机氮\n截至t日的历史气象", PALE_BLUE, NAVY, NAVY)
    add_arrow(slide, 3.27, 3.70, 0.45, 0.36, NAVY)
    add_box(slide, 3.82, 2.85, 2.62, 2.20, "完美预报窗口", "以历史年份真实天气作为\nt+1至t+H日预报\n降水·温度·太阳辐射", LIGHT_GREEN, GREEN, GREEN)
    add_arrow(slide, 6.58, 3.70, 0.45, 0.36, NAVY)
    add_box(slide, 7.13, 2.85, 2.40, 2.20, "DQN决策", "根据状态与未来天气\n选择当日灌溉/施氮动作\nDSSAT推进至下一日", LIGHT_ORANGE, ORANGE, ORANGE)
    add_text(slide, "sₜ = [作物状态，土壤水氮，历史天气，未来H天天气]",
             1.02, 5.55, 7.95, 0.50, 17, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, "定位：完美预报是理想信息上界；后续通过无预报、不同H及预报扰动检验稳健性",
             0.68, 6.30, 8.65, 0.48, 14.5, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "这里的完美预报并不是声称现实预报没有误差，而是把历史年份已经发生的天气序列作为未来已知信息，建立决策收益的理论上界。后续可以通过缩短预报窗口或加入扰动评价稳健性。")

    # Slide 28: validation design.
    slide = slides[27]; clear_slide(slide); add_header(slide, "多站点、多年份和多基准共同检验策略泛化能力", "4.6")
    add_text(slide, "训练与验证设计", 0.55, 1.08, 2.0, 0.42, 17, True, NAVY,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    cols = [
        ("空间维度", "五个典型农业生态站点\n不同气候、土壤与管理背景", LIGHT_GREEN, GREEN),
        ("时间维度", "多年份训练与独立年份测试\n覆盖干旱、湿润与正常年份", PALE_BLUE, BLUE),
        ("信息维度", "无预报 vs 完美预报\n不同预报窗口与扰动情景", LIGHT_ORANGE, ORANGE),
    ]
    for i, (title, body, fill, accent) in enumerate(cols):
        add_box(slide, 0.48 + i * 3.14, 1.62, 2.85, 1.68, title, body, fill, accent, accent)
    add_text(slide, "对照策略", 0.55, 3.76, 1.50, 0.38, 16, True, NAVY,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    policies = ["零投入", "农民管理", "专家规则", "DSSAT自动管理", "DQN策略"]
    for i, name in enumerate(policies):
        add_box(slide, 0.48 + i * 1.89, 4.24, 1.62, 0.78, name, None,
                LIGHT_BLUE if i == 4 else WHITE, NAVY if i == 4 else LINE, NAVY)
    add_text(slide, "评价指标", 0.55, 5.45, 1.50, 0.38, 16, True, NAVY,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    metrics = ["产量/收益", "WUE", "NUE", "灌溉与施氮", "氮损失", "稳定性"]
    for i, name in enumerate(metrics):
        add_text(slide, name, 0.60 + i * 1.48, 5.95, 1.25, 0.46, 13.5, True, DARK,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.74 + i * 1.48), Inches(6.48), Inches(0.98), Inches(0.06))
        bar.fill.solid(); bar.fill.fore_color.rgb = [GREEN, BLUE, NAVY, ORANGE, RED, MID][i]
        bar.line.fill.background()
    set_notes(slide, "验证不只比较最终收益，还要看不同站点、不同年份和不同天气信息条件下的稳定性。基准策略包括零投入、农民管理、专家规则和DSSAT自动管理。")

    # Slide 30: published paper and real figure.
    slide = slides[29]; clear_slide(slide); add_header(slide, "已发表论文完成全球氮损失识别与机制分析", "5.1")
    map_img = MEDIA / "doc_image8_b474f7fd.png"
    add_picture_contain(slide, map_img, 0.40, 1.00, 6.45, 5.86)
    add_box(slide, 6.95, 1.10, 2.55, 1.33, "论文成果", "Journal of Cleaner Production\n559 (2026) 148256", LIGHT_BLUE, NAVY, NAVY)
    add_box(slide, 6.95, 2.70, 2.55, 2.46, "核心发现", "• 水文路径贡献超过55%\n• 极端降水显著影响淋失与径流\n• 气候和土壤质地控制气态损失", LIGHT_GREEN, GREEN, GREEN)
    add_text(slide, "已形成全球5 arcmin作物特异氮损失数据与分析流程",
             6.98, 5.56, 2.50, 0.82, 15, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_source(slide, "Source: Du et al., Journal of Cleaner Production 559 (2026) 148256. https://doi.org/10.1016/j.jclepro.2026.148256")
    set_notes(slide, "这是本人已发表论文的核心结果图。该研究建立了玉米和小麦五类氮损失路径的全球高分辨率制图与机制分析基础，也是本开题中数据驱动部分的研究基础。")

    # Slide 31: mitigation results from the real paper.
    slide = slides[30]; clear_slide(slide); add_header(slide, "联合提高NUE与采用EEFs具有最大的氮损失缓解潜力", "5.2")
    add_text(slide, "玉米", 0.55, 0.92, 4.25, 0.35, 16, True, NAVY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, "小麦", 5.18, 0.92, 4.25, 0.35, 16, True, NAVY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_picture_contain(slide, maize_crop, 0.48, 1.28, 4.38, 2.78)
    add_picture_contain(slide, wheat_crop, 5.14, 1.28, 4.38, 2.78)
    add_box(slide, 0.75, 4.48, 3.95, 1.34, "玉米", "联合情景减排约 23.9%", LIGHT_ORANGE, ORANGE, ORANGE)
    add_box(slide, 5.30, 4.48, 3.95, 1.34, "小麦", "联合情景减排 80.6%", LIGHT_GREEN, GREEN, GREEN)
    add_text(slide, "减排策略需要体现作物与区域差异，不能使用单一全球方案",
             0.75, 6.20, 8.5, 0.46, 15.5, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_source(slide, "Source: selected panels from Du et al., Journal of Cleaner Production 559 (2026) 148256.")
    set_notes(slide, "情景分析表明，提高氮利用效率并采用增效氮肥的联合策略效果最好，但玉米和小麦的减排幅度差异显著，因此需要作物和区域差异化管理。")

    # Slide 32: RL foundation and feasibility.
    slide = slides[31]; clear_slide(slide); add_header(slide, "数据、模型与计算条件支撑后续强化学习研究", "5.3")
    foundations = [
        ("站点与数据", "已整理典型站点气象、土壤、品种和管理资料\n具备多年份DSSAT输入基础", LIGHT_GREEN, GREEN),
        ("模型与环境", "已完成DSSAT与gym-DSSAT环境配置和运行测试\n已设计状态、动作与奖励框架", LIGHT_BLUE, BLUE),
        ("算法与验证", "已由连续动作探索转向离散DQN框架\n拟开展多基准、多站点与跨年份验证", LIGHT_ORANGE, ORANGE),
    ]
    for i, (title, body, fill, accent) in enumerate(foundations):
        add_box(slide, 0.48 + i * 3.16, 1.25, 2.88, 2.42, title, body, fill, accent, accent)
    add_text(slide, "已具备的技术链", 0.60, 4.28, 1.65, 0.36, 16, True, NAVY,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    chain = ["数据预处理", "DSSAT模拟", "gym交互", "DQN训练", "结果评价"]
    for i, label in enumerate(chain):
        x = 0.50 + i * 1.88
        add_box(slide, x, 4.82, 1.55, 0.80, label, None, WHITE, NAVY, NAVY)
        if i < 4: add_arrow(slide, x + 1.59, 5.05, 0.22, 0.26, NAVY)
    add_text(slide, "风险控制：先开展小规模烟雾测试，再进行长周期训练；完整保存日志、模型与评价结果",
             0.65, 6.28, 8.7, 0.52, 14.5, True, DARK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "本研究的数据、模型和软件环境已经具备。后续重点是完善天气信息编码、DQN训练和泛化验证，并通过小规模测试控制计算风险。")

    # Slide 33: schedule as a clean timeline.
    slide = slides[32]; clear_slide(slide); add_header(slide, "研究进度安排", "6")
    periods = [
        ("2026.07–12", "全球氮损失", "完善评估、驱动与情景分析", GREEN),
        ("2027.01–06", "数字环境", "构建gym-DSSAT水氮环境", BLUE),
        ("2027.07–12", "策略优化", "开展DQN训练与泛化验证", ORANGE),
        ("2028.01–03", "综合分析", "整合结果并完成论文撰写", NAVY),
        ("2028.04–06", "论文答辩", "修改论文与准备答辩", RED),
    ]
    axis = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.92), Inches(3.24), Inches(8.18), Inches(0.08))
    axis.fill.solid(); axis.fill.fore_color.rgb = LINE; axis.line.fill.background()
    for i, (period, title, body, color) in enumerate(periods):
        x = 0.48 + i * 1.86
        y = 1.34 if i % 2 == 0 else 3.78
        add_box(slide, x, y, 1.66, 1.58, title, body, WHITE, color, color)
        add_text(slide, period, x, y + (1.72 if i % 2 == 0 else -0.48), 1.66, 0.36, 11.5, True, color,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.64), Inches(3.02), Inches(0.38), Inches(0.38))
        node.fill.solid(); node.fill.fore_color.rgb = color; node.line.fill.background()
    add_text(slide, "阶段成果逐步衔接：全球诊断 → 环境构建 → 策略优化 → 综合与答辩",
             0.72, 6.36, 8.56, 0.48, 15.5, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    set_notes(slide, "研究计划分为五个阶段。前期完善已发表工作的扩展分析，随后集中构建强化学习环境并开展训练验证，最后完成综合分析、论文撰写和答辩准备。")

    # Slide 34: closing.
    slide = slides[33]; clear_slide(slide); set_background(slide, WHITE)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.02))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
    add_text(slide, "感谢各位老师指导", 0.80, 2.05, 8.40, 0.86, 32, True, NAVY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "请批评指正", 0.80, 3.13, 8.40, 0.58, 20, False, GREEN,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.28), Inches(4.08), Inches(1.44), Inches(0.06))
    line.fill.solid(); line.fill.fore_color.rgb = ORANGE; line.line.fill.background()
    add_text(slide, "杜雨秋  ·  2026年7月17日", 0.80, 4.55, 8.40, 0.45, 14, True, DARK,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "中国科学院地理科学与资源研究所", 0.80, 5.18, 8.40, 0.38, 12, False, MID,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, margin=0)
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.92), Inches(10), Inches(0.58))
    footer.fill.solid(); footer.fill.fore_color.rgb = LIGHT_GREEN; footer.line.fill.background()

    temp_path = OUTPUT_PATH.with_name(OUTPUT_PATH.stem + "-building.pptx")
    prs.save(str(temp_path))
    temp_path.replace(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build_deck()
