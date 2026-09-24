#!/usr/bin/env python3
"""Create the Chinese 003_05 YC weather finalization gate presentation."""

from __future__ import annotations

import json
import hashlib
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "yc_weather_finalize_and_cli"
OUTPUT = ROOT / "docs" / "yc_weather_finalize_and_cli.pptx"
SUMMARY = RESULTS / "audit_summary.json"

INK = RGBColor(34, 50, 55)
MUTED = RGBColor(101, 117, 119)
TEAL = RGBColor(26, 116, 112)
PALE_TEAL = RGBColor(227, 241, 238)
ORANGE = RGBColor(211, 111, 67)
PALE_ORANGE = RGBColor(250, 238, 229)
GOLD = RGBColor(208, 164, 72)
PALE_GOLD = RGBColor(249, 245, 226)
LINE = RGBColor(220, 229, 226)
PAPER = RGBColor(250, 252, 251)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(81, 127, 99)


def text(slide, x, y, w, h, value, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(value)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    obj = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = fill
    obj.line.color.rgb = line or fill
    return obj


def base_slide(prs, title, section, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    rect(slide, 0, 0, 13.333, 0.10, TEAL)
    text(slide, 0.62, 0.34, 5.7, 0.28, section.upper(), 9, TEAL, True)
    text(slide, 0.62, 0.72, 12.1, 0.58, title, 25, INK, True)
    rect(slide, 0.62, 7.12, 12.08, 0.012, LINE)
    text(slide, 0.64, 7.20, 8.5, 0.22, "YC TRAIN WEATHER  |  003_05", 8, MUTED, True)
    text(slide, 12.0, 7.18, 0.7, 0.24, f"{page:02d} / 06", 8, MUTED, False, PP_ALIGN.RIGHT)
    return slide


def bullet(slide, x, y, w, value, color=INK, size=15, marker_color=TEAL):
    rect(slide, x, y + 0.08, 0.09, 0.09, marker_color, shape=MSO_SHAPE.OVAL)
    text(slide, x + 0.23, y, w - 0.23, 0.52, value, size, color)


def metric(slide, x, y, w, label, value, detail, accent=TEAL):
    rect(slide, x, y, w, 1.35, WHITE, LINE)
    rect(slide, x, y, 0.07, 1.35, accent)
    text(slide, x + 0.20, y + 0.13, w - 0.35, 0.28, label, 10, MUTED, True)
    text(slide, x + 0.20, y + 0.43, w - 0.35, 0.48, value, 25, accent, True)
    text(slide, x + 0.20, y + 0.98, w - 0.35, 0.26, detail, 9, MUTED)


def cell_text(cell, value, size=11, color=INK, bold=False):
    cell.text = str(value)
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.035)
    cell.margin_bottom = Inches(0.035)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in cell.text_frame.paragraphs:
        p.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.bold = bold


def table(slide, x, y, w, h, headers, rows, widths=None, font_size=10):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if widths:
        for col, width in zip(tbl.columns, widths):
            col.width = Inches(width)
    for ci, header in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK
        cell_text(cell, header, font_size, WHITE, True)
    for ri, row in enumerate(rows, start=1):
        for ci, value in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else PALE_TEAL
            cell_text(cell, value, font_size, INK)
    return shape


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def build():
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    validation = summary.get("external_validation", {})
    metrics = validation.get("metrics", {})
    rain = metrics.get("rain", {})
    annual = summary.get("rain_source_comparison_2005_2010", [])
    gaps = summary.get("gap_counts", {})
    failed = [k for k, value in validation.get("gate_checks", {}).items() if not value]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = base_slide(prs, "YC 训练天气定稿与 `.CLI` 准备", "决策摘要", 1)
    rect(slide, 0.62, 1.60, 7.05, 4.88, PALE_ORANGE, PALE_ORANGE)
    text(slide, 0.96, 1.92, 6.3, 0.35, "最终 Gate", 12, ORANGE, True)
    text(slide, 0.96, 2.34, 6.25, 0.64, summary.get("final_status", "BLOCKED_EXTERNAL_GAPFILL"), 22, INK, True)
    text(slide, 0.96, 3.20, 6.12, 0.82,
         "NASA POWER 的温度与辐射 overlap 通过；降雨事件一致率未达到预设门槛。", 17, INK)
    rect(slide, 0.96, 4.32, 5.92, 0.015, RGBColor(224, 204, 190))
    bullet(slide, 0.98, 4.62, 5.95, "不生成 fitting candidate，不估算年/月气候统计。", size=14, marker_color=ORANGE)
    bullet(slide, 0.98, 5.29, 5.95, "不生成 `.CLI`，不运行 WGEN、DSSAT 或 PPO。", size=14, marker_color=ORANGE)
    metric(slide, 8.02, 1.66, 2.18, "原始变量级 gaps", str(summary.get("gap_total_variable_days", 0)), "仅限 2004–2013", ORANGE)
    metric(slide, 10.42, 1.66, 2.18, "降雨事件一致率", f"{rain.get('rain_event_agreement_rate', 0):.1%}", "门槛 80.0%", ORANGE)
    metric(slide, 8.02, 3.24, 2.18, "接受填补", str(summary.get("gapfill_accepted_count", 0)), "candidate 未形成", MUTED)
    metric(slide, 10.42, 3.24, 2.18, "候选完整日数", "0 / 3,653", "QC 未运行", MUTED)
    text(slide, 8.08, 5.15, 4.45, 0.92, "停在数据质量 Gate 是预期行为：保留证据，不用猜值换取通过。", 15, INK, True)

    slide = base_slide(prs, "旧阻塞放宽了什么，哪些仍然不能猜", "决策原则", 2)
    text(slide, 0.74, 1.52, 5.35, 0.36, "旧审计阻塞", 13, MUTED, True)
    text(slide, 7.17, 1.52, 5.35, 0.36, "本轮采用的准则", 13, TEAL, True)
    rect(slide, 0.74, 1.98, 5.60, 3.44, WHITE, LINE)
    rect(slide, 7.00, 1.98, 5.60, 3.44, PALE_TEAL, PALE_TEAL)
    bullet(slide, 1.05, 2.32, 4.96, "温度列必须逐一证明对应传感器高度。", size=14, marker_color=ORANGE)
    bullet(slide, 1.05, 3.16, 4.96, "官方产品每个单元格都要有实测/插补标志。", size=14, marker_color=ORANGE)
    bullet(slide, 1.05, 4.00, 4.96, "QC 产品中的 0 mm 雨量曾被视为未决。", size=14, marker_color=ORANGE)
    bullet(slide, 7.32, 2.32, 4.95, "温度按重叠期一致性选“近地面”，不称为已确认 2 m。", size=14)
    bullet(slide, 7.32, 3.16, 4.95, "正式接受 2005–2013 官方 QC 日产品，并如实标注产品值。", size=14)
    bullet(slide, 7.32, 4.00, 4.95, "官方发布的 0 mm 直接接受；只填真缺测/非法值。", size=14)
    rect(slide, 0.74, 5.78, 11.86, 0.82, PALE_GOLD, PALE_GOLD)
    text(slide, 1.02, 5.98, 11.2, 0.40, "未放宽：训练期只到 2013；外部值只填 gap；所有来源和 QC 均保留。", 16, INK, True)

    slide = base_slide(prs, "数据源层级与缺口规模", "来源审计", 3)
    rect(slide, 0.72, 1.58, 5.30, 4.95, WHITE, LINE)
    text(slide, 1.00, 1.86, 4.75, 0.35, "2004 · ChinaFLUX 30 min", 16, TEAL, True)
    bullet(slide, 1.02, 2.40, 4.65, "温度：近地面空气温度日极值；按重叠期一致性选择，不声称已确认 2 m。", size=12)
    bullet(slide, 1.02, 3.20, 4.65, "SRAD：48 个半小时辐射积分；RAIN：48 个半小时求和。", size=12)
    bullet(slide, 1.02, 4.00, 4.65, "RAIN 仅 5 日不完整：2004-10-16 至 2004-10-20。", size=12, marker_color=ORANGE)
    text(slide, 1.00, 5.17, 4.68, 0.35, "2005–2013 · 官方 QC 日产品", 16, TEAL, True)
    bullet(slide, 1.02, 5.66, 4.65, "TMAX / TMIN / RAIN / SRAD；官方 0 mm 保留。", size=12)
    year_rows = []
    for year in range(2004, 2014):
        fields = gaps.get(str(year), {})
        count = sum(fields.values())
        components = ", ".join(f"{var} {n}" for var, n in sorted(fields.items())) or "无"
        year_rows.append([year, count, components])
    table(slide, 6.38, 1.62, 6.20, 4.84, ["年份", "缺口项", "按变量拆分"], year_rows, [1.15, 1.25, 3.8], 10)
    text(slide, 6.44, 6.57, 6.08, 0.26, "合计 370 个 date-variable gaps；每条日期、原值与原因见 CSV。", 10, MUTED)

    slide = base_slide(prs, "NASA POWER overlap：三项通过，降雨未通过", "外部源验证", 4)
    rows = []
    for variable in ("TMAX", "TMIN", "SRAD"):
        item = metrics.get("temperature_srad", {}).get(variable, {})
        corrected = item.get("corrected", {})
        rows.append([variable, corrected.get("n", 0), f"{corrected.get('mae', 0):.2f}", f"{corrected.get('rmse', 0):.2f}", f"{corrected.get('pearson_r', 0):.3f}", "通过"])
    table(slide, 0.72, 1.55, 7.10, 2.45, ["变量", "配对日", "校正 MAE", "校正 RMSE", "Pearson r", "结论"], rows,
          [1.10, 1.25, 1.25, 1.45, 1.25, 0.80], 10)
    rect(slide, 8.20, 1.55, 4.40, 2.45, PALE_ORANGE, PALE_ORANGE)
    text(slide, 8.54, 1.84, 3.68, 0.35, "RAIN · Gate Fail", 16, ORANGE, True)
    text(slide, 8.54, 2.30, 3.64, 0.47, f"{rain.get('rain_event_agreement_rate', 0):.1%}", 29, ORANGE, True)
    text(slide, 8.54, 2.90, 3.60, 0.78, f"事件量 MAE {rain.get('event_amount_errors_on_official_wet_days', {}).get('mae', 0):.2f} mm\n月总量比中位数 {rain.get('monthly_ratio_summary', {}).get('median', 0):.2f}", 12, INK)
    rect(slide, 0.72, 4.32, 11.88, 1.75, WHITE, LINE)
    text(slide, 1.02, 4.60, 11.2, 0.40, "结论：外部降雨不够可靠，不能据此补 2004 的五个缺口日。", 18, INK, True)
    text(slide, 1.02, 5.17, 11.1, 0.55, "事件一致率预设门槛 80%；NASA POWER 实测对照为 71.98%。年/月总量比虽接近 1，不能替代逐日事件可信度。", 13, MUTED)

    slide = base_slide(prs, "Candidate 未落盘；2014+ 保持隔离", "候选与泄漏", 5)
    metric(slide, 0.78, 1.58, 2.72, "候选天气", "未生成", "外部 overlap Gate fail", ORANGE)
    metric(slide, 3.68, 1.58, 2.72, "Candidate QC", "未运行", "不把 blocked 当作 pass", ORANGE)
    metric(slide, 6.58, 1.58, 2.72, "2004 年雨量", "未定稿", "5 日缺口未获准填补", MUTED)
    metric(slide, 9.48, 1.58, 2.72, "2014+ 泄漏", "未发现", "cell values 未读取", GREEN)
    rect(slide, 0.78, 3.38, 11.42, 2.54, WHITE, LINE)
    text(slide, 1.08, 3.66, 10.7, 0.35, "泄漏审计边界", 15, TEAL, True)
    bullet(slide, 1.10, 4.20, 10.45, "ChinaFLUX 日值只读 2004–2010；官方气象值到 2013，2014 仅读年份标记并停止。", size=13)
    bullet(slide, 1.10, 4.82, 10.45, "NASA 请求截止 2013-12-31；偏差参数只用 2005–2013 重叠日。", size=13)
    bullet(slide, 1.10, 5.44, 10.45, "`assert max(source_weather_years_used_for_values_or_calibration) <= 2013`。", size=13)
    text(slide, 0.82, 6.28, 11.4, 0.42, "2005–2010 source totals 不一致仍如实列出；2005/2006 官方日值高于 ChinaFLUX 51.2 / 23.4 mm。", 12, MUTED)

    slide = base_slide(prs, "CLI Gate 顺延：先解决每日降雨缺口", "下一步", 6)
    rect(slide, 0.78, 1.60, 5.54, 4.72, PALE_ORANGE, PALE_ORANGE)
    text(slide, 1.10, 1.92, 4.80, 0.36, "本轮停止点", 13, ORANGE, True)
    text(slide, 1.10, 2.44, 4.82, 0.74, "BLOCKED_EXTERNAL_GAPFILL", 19, INK, True)
    bullet(slide, 1.12, 3.48, 4.68, "未生成候选 `.CLI` 输入文件。", size=14, marker_color=ORANGE)
    bullet(slide, 1.12, 4.31, 4.68, "未伪造 WGEN 参数；未触碰生产天气文件。", size=14, marker_color=ORANGE)
    bullet(slide, 1.12, 5.14, 4.68, "WGEN pilot、DSSAT smoke、PPO 全部未启动。", size=14, marker_color=ORANGE)
    text(slide, 6.92, 1.82, 5.28, 0.36, "最小后续任务", 13, TEAL, True)
    bullet(slide, 6.94, 2.34, 5.14, "为 2004-10-16 至 20 找到可信逐日雨量，或引入更合适的公开源并做同样的 2005–2013 overlap。", size=14)
    bullet(slide, 6.94, 3.65, 5.14, "外部降雨 Gate 通过后再构建 3,653 日 candidate 并跑完整 QC。", size=14)
    bullet(slide, 6.94, 4.63, 5.14, "候选通过后使用 WeatherMan 官方 `Generate > Calculate Parameters` GUI 流程；不生成随机天气。", size=14)
    text(slide, 6.95, 5.82, 5.14, 0.52, "WeatherMan 参考：DSSAT User's Guide Vol. 3。", 11, MUTED)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        backup_dir = ROOT / "backups" / "yc_weather_finalize_and_cli"
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup_path = backup_dir / f"yc_weather_finalize_and_cli_{datetime.now(timezone.utc):%Y%m%dT%H%M%SZ}.pptx"
        shutil.copy2(OUTPUT, backup_path)
    prs.core_properties.title = "YC 2004-2013 天气定稿与 CLI 准备"
    prs.core_properties.subject = "003_05 overlap validation and weather candidate gate"
    prs.core_properties.author = "OpenAI Codex"
    prs.save(OUTPUT)
    check = Presentation(OUTPUT)
    max_width, max_height = check.slide_width, check.slide_height
    out_of_bounds = []
    for slide_number, slide in enumerate(check.slides, start=1):
        for shape_number, shape in enumerate(slide.shapes, start=1):
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > max_width or shape.top + shape.height > max_height:
                out_of_bounds.append({"slide": slide_number, "shape": shape_number, "name": shape.name})
    with zipfile.ZipFile(OUTPUT) as archive:
        zip_error = archive.testzip()
    validation = {
        "path": str(OUTPUT.relative_to(ROOT)), "slide_count": len(check.slides),
        "expected_slide_count": 6, "all_slides_have_shapes": all(len(s.shapes) > 0 for s in check.slides),
        "final_status": summary.get("final_status"), "validation_time_utc": datetime.now(timezone.utc).isoformat(),
        "sha256": sha256_file(OUTPUT), "zip_package_integrity": zip_error is None,
        "shapes_within_slide_bounds": not out_of_bounds, "out_of_bounds_shapes": out_of_bounds,
        "visual_renderer_available": False,
    }
    (RESULTS / "pptx_validation_summary.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    log_path = RESULTS / "experiment_log.md"
    if log_path.exists():
        log = log_path.read_text(encoding="utf-8")
        marker = "\n## PPTX 生成与检查\n"
        if marker in log:
            log = log.split(marker, 1)[0]
        log += marker + f"\n- 文件：`docs/yc_weather_finalize_and_cli.pptx`，6 页；SHA256 `{validation['sha256']}`。\n- 结构检查：ZIP 完整={validation['zip_package_integrity']}；所有形状在画布内={validation['shapes_within_slide_bounds']}。\n- 未执行图像级渲染；当前只确认结构与边界，视觉渲染状态记为 unavailable。\n"
        log_path.write_text(log, encoding="utf-8")
    print(json.dumps(validation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
