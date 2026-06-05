from __future__ import annotations

import math
import re
from datetime import datetime
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = PROJECT_ROOT / "data"
WEATHER_QC_DIR = PROJECT_ROOT / "weather_clean_qc"
YEAR_DIR = PROJECT_ROOT / "Leave_One_experiments" / "year_classification"
FIG_DIR = (
    PROJECT_ROOT
    / "Leave_One_experiments"
    / "figures"
    / "observed_phenology_rainfall_diagnosis"
)
DOCS_DIR = PROJECT_ROOT / "docs"

OBSERVED_INPUT = DATA_DIR / "observed_phenology_dates.csv"
OBSERVED_FALLBACK = YEAR_DIR / "station_phenology_records.csv"
STANDARDIZED_OUTPUT = DATA_DIR / "observed_phenology_dates_standardized.csv"
DIAG_OUTPUT = YEAR_DIR / "observed_phenology_rainfall_diagnosis.csv"
RANK_OUTPUT = YEAR_DIR / "observed_phenology_rainfall_rank_by_station.csv"
REPORT_OUTPUT = DOCS_DIR / "phase2c_observed_phenology_rainfall_diagnosis.md"
PPT_OUTPUT = DOCS_DIR / "phase2c_observed_phenology_rainfall_diagnosis.pptx"
PHASE2B_SELECTED = YEAR_DIR / "selected_years_qc.csv"
PHASE2B_REPORT = (
    PROJECT_ROOT
    / "Leave_One_experiments"
    / "reports"
    / "phase2b_rain_qc_and_reclassification_report.md"
)

DATE_COL_CANDIDATES = ["DATE", "date", "Date", "YYYYMMDD", "year_doy"]
RAIN_COL_CANDIDATES = ["RAIN", "rain", "Rain", "PREC", "PRCP"]
STATIONS = ["HLA", "SYA", "LCA", "YCA", "FQA"]
BLUE = RGBColor(31, 78, 121)
LIGHT_BLUE = RGBColor(221, 235, 247)


def ensure_dirs() -> None:
    for directory in [DATA_DIR, YEAR_DIR, FIG_DIR, DOCS_DIR]:
        directory.mkdir(parents=True, exist_ok=True)


def copy_observed_input_if_needed() -> None:
    if OBSERVED_INPUT.exists():
        return
    if not OBSERVED_FALLBACK.exists():
        raise FileNotFoundError(
            "Missing data/observed_phenology_dates.csv and fallback "
            "Leave_One_experiments/year_classification/station_phenology_records.csv"
        )
    df = pd.read_csv(OBSERVED_FALLBACK)
    needed = [
        "station_code",
        "station_name",
        "year",
        "planting_date",
        "silking_date",
        "mature_date",
        "harvest_date",
        "source",
        "note",
    ]
    df[needed].to_csv(OBSERVED_INPUT, index=False, encoding="utf-8-sig")


def parse_date(value: object) -> pd.Timestamp:
    ts = pd.to_datetime(value, errors="coerce")
    if pd.isna(ts):
        raise ValueError(f"Invalid date: {value!r}")
    return pd.Timestamp(ts).normalize()


def to_yyddd(ts: pd.Timestamp) -> str:
    return f"{ts.year % 100:02d}{ts.dayofyear:03d}"


def qc_phenology(row: pd.Series) -> str:
    flags: list[str] = []
    if not row["planting_date_dt"] < row["silking_date_dt"]:
        flags.append("invalid_planting_silking_order")
    if not row["silking_date_dt"] < row["mature_date_dt"]:
        flags.append("invalid_silking_mature_order")
    if not row["mature_date_dt"] <= row["harvest_date_dt"]:
        flags.append("invalid_mature_after_harvest")
    if not row["harvest_date_dt"] > row["planting_date_dt"]:
        flags.append("invalid_harvest_before_planting")
    if not row["mature_date_dt"] > row["planting_date_dt"]:
        flags.append("invalid_mature_before_planting")
    if row["planting_to_silking_days"] < 35 or row["planting_to_silking_days"] > 110:
        flags.append("suspect_silking_duration")
    if row["planting_to_mature_days"] < 80 or row["planting_to_mature_days"] > 190:
        flags.append("suspect_mature_duration")
    if row["planting_to_harvest_days"] < 80 or row["planting_to_harvest_days"] > 220:
        flags.append("suspect_harvest_duration")
    if row["mature_to_harvest_days"] < 0:
        flags.append("invalid_mature_after_harvest")
    if row["mature_to_harvest_days"] > 30:
        flags.append("long_mature_to_harvest_gap")
    return "pass" if not flags else ";".join(dict.fromkeys(flags))


def standardize_phenology() -> pd.DataFrame:
    copy_observed_input_if_needed()
    df = pd.read_csv(OBSERVED_INPUT)
    required = [
        "station_code",
        "station_name",
        "year",
        "planting_date",
        "silking_date",
        "mature_date",
        "harvest_date",
        "source",
        "note",
    ]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"Observed phenology file missing columns: {missing}")

    for col in ["planting_date", "silking_date", "mature_date", "harvest_date"]:
        df[f"{col}_dt"] = df[col].map(parse_date)
        df[col] = df[f"{col}_dt"].dt.strftime("%Y-%m-%d")
        df[col.replace("_date", "_yyddd")] = df[f"{col}_dt"].map(to_yyddd)

    df["planting_to_silking_days"] = (
        df["silking_date_dt"] - df["planting_date_dt"]
    ).dt.days
    df["silking_to_mature_days"] = (
        df["mature_date_dt"] - df["silking_date_dt"]
    ).dt.days
    df["mature_to_harvest_days"] = (
        df["harvest_date_dt"] - df["mature_date_dt"]
    ).dt.days
    df["planting_to_mature_days"] = (
        df["mature_date_dt"] - df["planting_date_dt"]
    ).dt.days
    df["planting_to_harvest_days"] = (
        df["harvest_date_dt"] - df["planting_date_dt"]
    ).dt.days
    df["phenology_qc_flag"] = df.apply(qc_phenology, axis=1)

    out_cols = [
        "station_code",
        "station_name",
        "year",
        "planting_date",
        "silking_date",
        "mature_date",
        "harvest_date",
        "planting_yyddd",
        "silking_yyddd",
        "mature_yyddd",
        "harvest_yyddd",
        "planting_to_silking_days",
        "silking_to_mature_days",
        "mature_to_harvest_days",
        "planting_to_mature_days",
        "planting_to_harvest_days",
        "phenology_qc_flag",
        "source",
        "note",
    ]
    df[out_cols].to_csv(STANDARDIZED_OUTPUT, index=False, encoding="utf-8-sig")
    return df


def detect_column(columns: list[str], candidates: list[str], label: str) -> str:
    for candidate in candidates:
        if candidate in columns:
            return candidate
    raise ValueError(f"Could not detect {label} column from: {columns}")


def load_weather() -> tuple[dict[str, pd.DataFrame], dict[str, dict[str, str]]]:
    if not WEATHER_QC_DIR.exists():
        raise FileNotFoundError("weather_clean_qc directory is required for Phase 2c.")
    weather: dict[str, pd.DataFrame] = {}
    fields: dict[str, dict[str, str]] = {}
    for station in STATIONS:
        path = WEATHER_QC_DIR / f"{station}_weather_cleaned_qc.csv"
        if not path.exists():
            raise FileNotFoundError(f"Missing QC weather file: {path}")
        df = pd.read_csv(path)
        date_col = detect_column(list(df.columns), DATE_COL_CANDIDATES, "date")
        rain_col = detect_column(list(df.columns), RAIN_COL_CANDIDATES, "RAIN")
        df = df.copy()
        df["DATE_STD"] = pd.to_datetime(df[date_col], errors="coerce")
        df["RAIN_STD"] = pd.to_numeric(df[rain_col], errors="coerce")
        df = df.sort_values("DATE_STD").reset_index(drop=True)
        weather[station] = df
        fields[station] = {"date_col": date_col, "rain_col": rain_col}
    return weather, fields


def window_stats(
    df: pd.DataFrame,
    start: pd.Timestamp,
    end: pd.Timestamp,
) -> dict[str, float | int | str]:
    expected_days = int((end - start).days) + 1
    window = df[(df["DATE_STD"] >= start) & (df["DATE_STD"] <= end)].copy()
    available = int(window["DATE_STD"].nunique())
    missing_weather_days = expected_days - available
    missing_rain_days = int(window["RAIN_STD"].isna().sum())
    negative_rain_days = int((window["RAIN_STD"] < 0).sum())
    extreme_rain_days = int((window["RAIN_STD"] > 250).sum())
    zero_rain_days = int((window["RAIN_STD"] == 0).sum())
    rain_sum = float(window["RAIN_STD"].sum(skipna=True))

    flags: list[str] = []
    if missing_weather_days:
        flags.append("missing_weather_days")
    if missing_rain_days:
        flags.append("missing_rain_days")
    if negative_rain_days:
        flags.append("negative_rain_days")
    if extreme_rain_days:
        flags.append("extreme_rain_days_gt_250mm")
    if available > 0 and zero_rain_days == available:
        flags.append("all_zero_rain_in_window")

    return {
        "days_expected": expected_days,
        "days_available": available,
        "missing_weather_days": missing_weather_days,
        "missing_rain_days": missing_rain_days,
        "negative_rain_days": negative_rain_days,
        "zero_rain_days": zero_rain_days,
        "extreme_rain_days": extreme_rain_days,
        "rain_mm": round(rain_sum, 3),
        "qc_flag": "pass" if not flags else ";".join(flags),
    }


def diagnose_rainfall(phenology: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, dict[str, str]]]:
    weather, detected_fields = load_weather()
    rows: list[dict[str, object]] = []

    for _, row in phenology.iterrows():
        station = row["station_code"]
        start = row["planting_date_dt"]
        mature = row["mature_date_dt"]
        harvest = row["harvest_date_dt"]
        station_weather = weather[station]
        harvest_stats = window_stats(station_weather, start, harvest)
        mature_stats = window_stats(station_weather, start, mature)

        post_days = int((harvest - mature).days)
        post_rain = round(
            harvest_stats["rain_mm"] - mature_stats["rain_mm"],
            3,
        )
        rain_diff = post_rain
        ratio = (
            math.nan
            if harvest_stats["rain_mm"] == 0
            else round(rain_diff / harvest_stats["rain_mm"], 4)
        )

        rain_flags = []
        for flag in [harvest_stats["qc_flag"], mature_stats["qc_flag"]]:
            if flag != "pass":
                rain_flags.extend(flag.split(";"))
        if harvest_stats["rain_mm"] == 0:
            rain_flags.append("zero_harvest_window_rain")

        rows.append(
            {
                "station_code": station,
                "station_name": row["station_name"],
                "year": int(row["year"]),
                "planting_date": row["planting_date"],
                "mature_date": row["mature_date"],
                "harvest_date": row["harvest_date"],
                "harvest_window_start": row["planting_date"],
                "harvest_window_end": row["harvest_date"],
                "harvest_window_days": harvest_stats["days_expected"],
                "harvest_window_rain_mm": harvest_stats["rain_mm"],
                "mature_window_start": row["planting_date"],
                "mature_window_end": row["mature_date"],
                "mature_window_days": mature_stats["days_expected"],
                "mature_window_rain_mm": mature_stats["rain_mm"],
                "post_mature_window_days": post_days,
                "post_mature_window_rain_mm": post_rain,
                "rain_difference_harvest_minus_mature": rain_diff,
                "rain_difference_ratio": ratio,
                "weather_days_expected": harvest_stats["days_expected"],
                "weather_days_available": harvest_stats["days_available"],
                "missing_weather_days": harvest_stats["missing_weather_days"],
                "missing_rain_days": harvest_stats["missing_rain_days"],
                "negative_rain_days": harvest_stats["negative_rain_days"],
                "zero_rain_days": harvest_stats["zero_rain_days"],
                "phenology_qc_flag": row["phenology_qc_flag"],
                "rain_qc_flag": (
                    "pass" if not rain_flags else ";".join(dict.fromkeys(rain_flags))
                ),
                "observed_rain_rank": "",
                "observed_rain_label": "",
                "recommended_for_smoke_test": "",
                "note": row["note"],
            }
        )

    diagnosis = pd.DataFrame(rows)
    diagnosis = add_rank_labels(diagnosis)
    diagnosis.to_csv(DIAG_OUTPUT, index=False, encoding="utf-8-sig")
    return diagnosis, detected_fields


def add_rank_labels(diagnosis: pd.DataFrame) -> pd.DataFrame:
    out = diagnosis.copy()
    for station, group in out.groupby("station_code"):
        ordered = group.sort_values("harvest_window_rain_mm")
        n = len(ordered)
        median_rain = ordered["harvest_window_rain_mm"].median()
        mid_idx = (
            (ordered["harvest_window_rain_mm"] - median_rain).abs().sort_values().index[0]
            if n >= 3
            else None
        )
        for pos, idx in enumerate(ordered.index, start=1):
            if n == 2:
                label = "observed_lower_rain_year" if pos == 1 else "observed_higher_rain_year"
                rank = "rank_1_lower" if pos == 1 else "rank_2_higher"
            elif n == 3:
                label = [
                    "observed_low_rain_year",
                    "observed_mid_rain_year",
                    "observed_high_rain_year",
                ][pos - 1]
                rank = f"rank_{pos}"
            else:
                rank = f"rank_{pos}"
                if pos == 1:
                    label = "observed_low_rain_year"
                elif idx == mid_idx:
                    label = "observed_mid_rain_year"
                elif pos == n:
                    label = "observed_high_rain_year"
                else:
                    label = f"observed_intermediate_rain_year_{pos}"
            if out.loc[idx, "phenology_qc_flag"] == "pass" and out.loc[idx, "rain_qc_flag"] == "pass":
                recommended = "yes"
            else:
                recommended = "review_before_use"
            out.loc[idx, "observed_rain_rank"] = rank
            out.loc[idx, "observed_rain_label"] = label
            out.loc[idx, "recommended_for_smoke_test"] = recommended
    return out.sort_values(["station_code", "harvest_window_rain_mm"]).reset_index(drop=True)


def create_rank_table(diagnosis: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for station, group in diagnosis.groupby("station_code"):
        sorted_group = group.sort_values("harvest_window_rain_mm")
        rain_range = float(sorted_group["harvest_window_rain_mm"].max() - sorted_group["harvest_window_rain_mm"].min())
        n = len(sorted_group)
        enough_gradient = n >= 3 and rain_range >= 100
        for _, row in sorted_group.iterrows():
            rows.append(
                {
                    "station_code": station,
                    "station_name": row["station_name"],
                    "year": int(row["year"]),
                    "harvest_window_rain_mm": row["harvest_window_rain_mm"],
                    "mature_window_rain_mm": row["mature_window_rain_mm"],
                    "observed_rain_rank": row["observed_rain_rank"],
                    "observed_rain_label": row["observed_rain_label"],
                    "station_observed_year_count": n,
                    "station_observed_rain_range_mm": round(rain_range, 3),
                    "can_use_relative_low_mid_high": "yes" if enough_gradient else "limited",
                    "dry_normal_wet_candidate_advice": (
                        "relative_candidates_only" if enough_gradient else "do_not_call_dry_normal_wet"
                    ),
                    "recommended_next_step": (
                        "observed_year_smoke_test"
                        if enough_gradient
                        else "historical_weather_substitution_scenario;rainfall_scaling_sensitivity_scenario"
                    ),
                }
            )
    rank = pd.DataFrame(rows)
    rank.to_csv(RANK_OUTPUT, index=False, encoding="utf-8-sig")
    return rank


def load_phase2b_selected() -> pd.DataFrame:
    if not PHASE2B_SELECTED.exists():
        return pd.DataFrame()
    return pd.read_csv(PHASE2B_SELECTED)


def setup_plot_style() -> None:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 150


def save_figures(diagnosis: pd.DataFrame, rank: pd.DataFrame, phase2b: pd.DataFrame) -> list[Path]:
    setup_plot_style()
    paths: list[Path] = []

    fig, ax = plt.subplots(figsize=(9.5, 5.2))
    data = diagnosis.sort_values(["station_code", "harvest_window_rain_mm"])
    labels = data["station_code"] + "\n" + data["year"].astype(str)
    ax.bar(labels, data["harvest_window_rain_mm"], color="#2F75B5")
    ax.set_title("Phase 2c observed years: harvest-window RAIN")
    ax.set_ylabel("RAIN (mm)")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    path = FIG_DIR / "phase2c_harvest_window_rain_by_station_year.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    paths.append(path)

    fig, ax = plt.subplots(figsize=(9.5, 5.2))
    x = range(len(data))
    width = 0.38
    ax.bar([i - width / 2 for i in x], data["mature_window_rain_mm"], width, label="planting -> mature", color="#70AD47")
    ax.bar([i + width / 2 for i in x], data["harvest_window_rain_mm"], width, label="planting -> harvest", color="#2F75B5")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels)
    ax.set_title("Mature-window vs harvest-window RAIN")
    ax.set_ylabel("RAIN (mm)")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False)
    fig.tight_layout()
    path = FIG_DIR / "phase2c_mature_vs_harvest_window_rain.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    paths.append(path)

    fig, ax = plt.subplots(figsize=(9.5, 5.2))
    ax.bar(labels, data["harvest_window_days"], color="#5B9BD5")
    ax.set_title("Observed planting-to-harvest duration")
    ax.set_ylabel("Days")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    path = FIG_DIR / "phase2c_planting_to_harvest_days.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    paths.append(path)

    if not phase2b.empty:
        rows = []
        for _, row in phase2b.iterrows():
            station = row["station"]
            for label, ycol, rcol in [
                ("phase2b_dry", "dry_year", "dry_year_rain"),
                ("phase2b_normal", "normal_year", "normal_year_rain"),
                ("phase2b_wet", "wet_year", "wet_year_rain"),
            ]:
                rows.append({"station_code": station, "scenario": label, "year": int(row[ycol]), "rain_mm": float(row[rcol])})
        p2b_long = pd.DataFrame(rows)
        observed_station = diagnosis.groupby("station_code", as_index=False)["harvest_window_rain_mm"].agg(["min", "median", "max"]).reset_index()
        observed_station = observed_station.melt(id_vars="station_code", var_name="scenario", value_name="rain_mm")
        observed_station["scenario"] = "phase2c_observed_" + observed_station["scenario"].astype(str)
        combo = pd.concat([p2b_long[["station_code", "scenario", "rain_mm"]], observed_station], ignore_index=True)
        pivot = combo.pivot_table(index="station_code", columns="scenario", values="rain_mm", aggfunc="first")
        fig, ax = plt.subplots(figsize=(9.5, 5.2))
        pivot[[c for c in pivot.columns if "phase2b" in c]].plot(kind="bar", ax=ax, width=0.8, color=["#A9D18E", "#5B9BD5", "#ED7D31"])
        ax.set_title("Phase 2b fallback representative RAIN by station")
        ax.set_ylabel("RAIN (mm)")
        ax.grid(axis="y", alpha=0.25)
        ax.legend(frameon=False, fontsize=8)
        fig.tight_layout()
        path = FIG_DIR / "phase2b_fallback_representative_years.png"
        fig.savefig(path, bbox_inches="tight")
        plt.close(fig)
        paths.append(path)
    return paths


def md_table(df: pd.DataFrame, cols: list[str]) -> str:
    if df.empty:
        return "无。\n"
    return df[cols].to_markdown(index=False)


def make_report(
    diagnosis: pd.DataFrame,
    rank: pd.DataFrame,
    detected_fields: dict[str, dict[str, str]],
    phase2b: pd.DataFrame,
    figure_paths: list[Path],
) -> None:
    phenology_issues = diagnosis[diagnosis["phenology_qc_flag"] != "pass"]
    rain_issues = diagnosis[diagnosis["rain_qc_flag"] != "pass"]
    field_lines = [
        f"- {st}: date=`{fields['date_col']}`, RAIN=`{fields['rain_col']}`"
        for st, fields in sorted(detected_fields.items())
    ]
    overlap_lines = []
    if not phase2b.empty:
        observed = diagnosis.groupby("station_code")["year"].apply(lambda x: set(map(int, x))).to_dict()
        for _, row in phase2b.iterrows():
            station = row["station"]
            p2b_years = {int(row["dry_year"]), int(row["normal_year"]), int(row["wet_year"])}
            overlap = sorted(observed.get(station, set()) & p2b_years)
            overlap_lines.append(
                f"- {station}: Phase 2b={sorted(p2b_years)}, Phase 2c observed={sorted(observed.get(station, set()))}, overlap={overlap}"
            )

    enough = rank.groupby("station_code")["can_use_relative_low_mid_high"].first().to_dict()
    smoke = diagnosis[diagnosis["recommended_for_smoke_test"] == "yes"]

    lines = [
        "# Phase 2c 实测生育期降雨诊断",
        "",
        "## 1. 任务背景",
        "Phase 2b 已完成 QC 后天气与代表年份初选，但由于 jinja2 模板缺少明确收获期，当时只能使用 `planting + 150 days` 作为临时生育期窗口。本阶段改用实测 `planting_date -> harvest_date`，只诊断已有真实试验年份，不训练 PPO，不修改 reward，也不生成新的情景天气。",
        "",
        "## 2. 输入数据",
        f"- 实测物候输入：`{OBSERVED_INPUT.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- QC 后天气目录：`{WEATHER_QC_DIR.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- Phase 2b 对比文件：`{PHASE2B_SELECTED.relative_to(PROJECT_ROOT).as_posix()}`",
        "",
        "## 3. 日期 QC 结果",
        "所有日期已标准化为 `YYYY-MM-DD`，并生成 DSSAT 常用 `YYDDD` 字段。检查规则包括播种-吐丝-成熟-收获顺序，以及几个阶段持续天数阈值。",
        "",
        md_table(
            diagnosis,
            [
                "station_code",
                "year",
                "planting_date",
                "mature_date",
                "harvest_date",
                "harvest_window_days",
                "phenology_qc_flag",
            ],
        ),
        "",
        "异常日期记录：",
        md_table(phenology_issues, ["station_code", "year", "phenology_qc_flag", "note"]),
        "",
        "## 4. 天气数据读取和 RAIN 字段识别",
        *field_lines,
        "",
        "RAIN QC 检查包括窗口内天气日期完整性、RAIN 缺失、负值、极端值和全 0 情况。",
        "",
        "异常 RAIN 记录：",
        md_table(rain_issues, ["station_code", "year", "rain_qc_flag"]),
        "",
        "## 5. 每个站点-年份的完整生育期降雨",
        md_table(
            diagnosis,
            [
                "station_code",
                "year",
                "harvest_window_days",
                "harvest_window_rain_mm",
                "mature_window_rain_mm",
                "post_mature_window_rain_mm",
                "rain_difference_ratio",
            ],
        ),
        "",
        "## 6. 成熟到收获期间降雨影响",
        "成熟窗口和收获窗口的差值即成熟后到收获期的降雨。若差值很小，使用成熟期或收获期作为窗口终点对降雨排序影响有限；若差值较大，后续分类应优先使用收获窗口。",
        "",
        "## 7. 真实年份降雨排序",
        md_table(
            rank,
            [
                "station_code",
                "year",
                "harvest_window_rain_mm",
                "observed_rain_rank",
                "observed_rain_label",
                "can_use_relative_low_mid_high",
                "recommended_next_step",
            ],
        ),
        "",
        "## 8. 是否可以称为相对低雨/中雨/高雨",
        "本阶段默认不使用 dry / normal / wet 命名。只有至少 3 个实测年份且站内降雨跨度较明显时，才建议作为 observed experimental years 内部的相对低雨/中雨/高雨候选，不代表 2000-2022 长期气候分位数。",
        "",
        "\n".join(f"- {st}: {status}" for st, status in sorted(enough.items())),
        "",
        "## 9. 与 Phase 2b 代表年份对比",
        "Phase 2b 使用 QC 后天气和 `planting + 150 days` fallback 选择长期历史代表年；Phase 2c 只使用已有实测物候年份。因此 Phase 2c 不直接继承 Phase 2b 代表年，也不把 2-4 年真实试验记录强行扩展到 2000-2022 全部历史年份。",
        "",
        *overlap_lines,
        "",
        "## 10. 推荐下一步 smoke test 年份",
        "建议先对通过 QC 的实测年份做 NullAgent、fixed_low_N、fixed_medium_N、fixed_high_N smoke test，不直接进入 PPO。",
        "",
        md_table(
            smoke,
            [
                "station_code",
                "year",
                "observed_rain_label",
                "harvest_window_rain_mm",
                "recommended_for_smoke_test",
            ],
        ),
        "",
        "## 11. 后续情景方案建议",
        "- 对真实年份梯度不足或只有 2 年记录的站点，建议设计 `historical_weather_substitution_scenario`：固定真实试验年的品种、土壤初始条件和管理方案，替换为同站点历史低雨/中雨/高雨年份天气，并明确标记为虚拟天气替换情景。",
        "- 如需机制敏感性分析，建议设计 `rainfall_scaling_sensitivity_scenario`：以真实试验年份天气为 baseline，构造 RAIN x 0.6、0.8、1.0、1.2、1.4，只作为降雨敏感性分析。",
        "",
        "## 12. 本阶段生成的文件",
        f"- `{OBSERVED_INPUT.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- `{STANDARDIZED_OUTPUT.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- `{DIAG_OUTPUT.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- `{RANK_OUTPUT.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- `{REPORT_OUTPUT.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- `{PPT_OUTPUT.relative_to(PROJECT_ROOT).as_posix()}`",
        *[f"- `{p.relative_to(PROJECT_ROOT).as_posix()}`" for p in figure_paths],
        "",
        "## 13. 尚未解决的问题和风险",
        "- 部分站点只有 2 年真实试验记录，不能构成完整低雨/中雨/高雨梯度。",
        "- 实测年份排序是站内 observed years 的相对排序，不等同长期气候 dry/normal/wet 分类。",
        "- 后续 smoke test 应先使用固定策略验证环境、天气、管理文件是否能稳定跑通，再进入 PPO。",
    ]
    REPORT_OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def add_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.45))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 0, 0)


def add_bullets(slide, bullets: list[str], left=0.65, top=1.05, width=12.0, height=5.8, size=18) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(0, 0, 0)
        p.space_after = Pt(8)


def add_table(slide, df: pd.DataFrame, cols: list[str], left=0.45, top=1.05, width=12.4, height=5.8) -> None:
    data = df[cols].copy()
    max_rows = min(len(data), 12)
    rows = max_rows + 1
    table = slide.shapes.add_table(rows, len(cols), Inches(left), Inches(top), Inches(width), Inches(height)).table
    for j, col in enumerate(cols):
        cell = table.cell(0, j)
        cell.text = col
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER
    for i in range(max_rows):
        for j, col in enumerate(cols):
            val = data.iloc[i][col]
            text = "" if pd.isna(val) else str(round(val, 1) if isinstance(val, float) else val)
            cell = table.cell(i + 1, j)
            cell.text = text
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_BLUE
            for p in cell.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(9)
                p.font.color.rgb = RGBColor(0, 0, 0)


def make_ppt(diagnosis: pd.DataFrame, rank: pd.DataFrame, figure_paths: list[Path]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slides = [
        ("Phase 2c 实测生育期降雨诊断", ["目标：使用实测 planting_date -> harvest_date 替代 Phase 2b fallback。", "范围：只诊断真实试验年份，不训练 PPO，不修改 reward。"]),
        ("背景", ["Phase 2b 的 planting + 150 days 是临时窗口。", "Phase 2c 使用实测播种、吐丝、成熟、收获日期。", "本阶段判断真实试验年份是否足以代表降雨梯度。"]),
        ("输入数据", [f"实测物候：{OBSERVED_INPUT.name}", "天气：weather_clean_qc/*_weather_cleaned_qc.csv", "对比：selected_years_qc.csv 与 Phase 2b 报告。"]),
        ("方法", ["完整窗口：planting_date -> harvest_date。", "成熟窗口：planting_date -> mature_date。", "差异：harvest_window_rain - mature_window_rain。"]),
    ]
    for title, bullets in slides:
        s = prs.slides.add_slide(blank)
        add_title(s, title)
        add_bullets(s, bullets)

    s = prs.slides.add_slide(blank)
    add_title(s, "日期 QC 结果")
    add_table(s, diagnosis, ["station_code", "year", "harvest_window_days", "phenology_qc_flag"])

    s = prs.slides.add_slide(blank)
    add_title(s, "天气 RAIN QC 结果")
    add_table(s, diagnosis, ["station_code", "year", "weather_days_expected", "weather_days_available", "rain_qc_flag"])

    for title, image_name in [
        ("Harvest-window RAIN 对比", "phase2c_harvest_window_rain_by_station_year.png"),
        ("Mature-window 与 Harvest-window 对比", "phase2c_mature_vs_harvest_window_rain.png"),
        ("生育期长度", "phase2c_planting_to_harvest_days.png"),
        ("Phase 2b 代表年份对比", "phase2b_fallback_representative_years.png"),
    ]:
        path = FIG_DIR / image_name
        if path.exists():
            s = prs.slides.add_slide(blank)
            add_title(s, title)
            s.shapes.add_picture(str(path), Inches(0.65), Inches(1.05), width=Inches(12.0))

    s = prs.slides.add_slide(blank)
    add_title(s, "真实年份降雨排序")
    add_table(s, rank, ["station_code", "year", "harvest_window_rain_mm", "observed_rain_label", "can_use_relative_low_mid_high"])

    s = prs.slides.add_slide(blank)
    add_title(s, "是否足以代表低雨/中雨/高雨")
    add_bullets(s, ["至少 3 个真实年份且降雨跨度明显的站点，可作为 observed years 内部相对候选。", "只有 2 年记录的站点不能强行命名为中雨年。", "这些候选不代表 2000-2022 长期气候分位数。"])

    s = prs.slides.add_slide(blank)
    add_title(s, "推荐下一步 smoke test 年份")
    add_table(s, diagnosis, ["station_code", "year", "observed_rain_label", "recommended_for_smoke_test"])

    s = prs.slides.add_slide(blank)
    add_title(s, "后续情景方案建议")
    add_bullets(s, ["observed_year_smoke_test：NullAgent 与固定低/中/高氮策略。", "historical_weather_substitution_scenario：固定真实试验管理，替换历史代表年天气。", "rainfall_scaling_sensitivity_scenario：RAIN x 0.6/0.8/1.0/1.2/1.4。"])

    s = prs.slides.add_slide(blank)
    add_title(s, "文件输出和 GitHub 备份说明")
    add_bullets(s, ["本阶段文件已本地生成。", "是否提交/推送 GitHub 需要用户确认，当前未自动 push。", "关键文件位于 data/、docs/、scripts/、Leave_One_experiments/year_classification/ 和 figures/。"], size=17)

    prs.save(PPT_OUTPUT)


def main() -> None:
    ensure_dirs()
    phenology = standardize_phenology()
    diagnosis, detected_fields = diagnose_rainfall(phenology)
    rank = create_rank_table(diagnosis)
    phase2b = load_phase2b_selected()
    figure_paths = save_figures(diagnosis, rank, phase2b)
    make_report(diagnosis, rank, detected_fields, phase2b, figure_paths)
    make_ppt(diagnosis, rank, figure_paths)

    outputs = [
        OBSERVED_INPUT,
        STANDARDIZED_OUTPUT,
        DIAG_OUTPUT,
        RANK_OUTPUT,
        REPORT_OUTPUT,
        PPT_OUTPUT,
        *figure_paths,
    ]
    print("Phase 2c observed phenology rainfall diagnosis complete.")
    for path in outputs:
        print(path.relative_to(PROJECT_ROOT).as_posix())


if __name__ == "__main__":
    main()
