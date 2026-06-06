from __future__ import annotations

import itertools
import json
import math
import shutil
import traceback
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from offline_schedule_policy import DeterministicSchedulePolicy, ScheduleEvent
from ppo_action_safety import ActionSafetyState, apply_action_safety, normalize_action, update_action_safety_state
from ppo_evaluate import latest_observation_dict, make_env, scalar
from ppo_experiment_plan import find_year
from ppo_plot_results import plot_episode
from ppo_safe_rendering import PROJECT_ROOT, load_yaml


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_offline_schedule_search.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "offline_schedule_search"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_offline_schedule_search_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_offline_schedule_search_report.pptx"


def ensure_dirs() -> None:
    for sub in [
        "configs",
        "rendered_inputs",
        "candidate_schedules",
        "daily_outputs/HLA",
        "evaluation",
        "figures/HLA",
        "figures/summary",
        "expert_policy",
        "reports",
    ]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)


def df_to_markdown(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df.empty:
        return ""
    work = df.copy()
    if max_rows is not None:
        work = work.head(max_rows)
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(4)
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    separator = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, separator, *rows])


def write_pause_note() -> None:
    text = """# Why pause PPO and search offline

Generated at: 2026-06-06

006_03 to 006_07 showed a repeated pattern: step-wise cost rewards, stronger cost coefficients, episode-level profit rewards, low-frequency/window action wrappers, and explicit budget/event action wrappers all still drove HLA PPO policies to the seasonal action-safety cap of 300 mm irrigation and 450 kg ha-1 nitrogen.

Reward cost tuning was insufficient because the learned policy still treated available seasonal input as beneficial. Terminal profit reward was insufficient because the policy continued to select cap-level actions before the terminal score could create an interpretable low-input behavior. Low-frequency, window, budget, and event wrappers changed when actions could happen, but did not provide an external prior for what a reasonable management schedule should look like.

The next methodological step is therefore non-RL schedule search. A deterministic DSSAT scenario ensemble can expose the yield-water-nitrogen trade-off, identify Pareto-efficient schedules, and construct expert schedules. These expert schedules can later support imitation learning, behavior cloning, reward calibration, or constrained PPO.
"""
    (OUTPUT_ROOT / "evaluation" / "why_pause_ppo_and_search_offline.md").write_text(text, encoding="utf-8")


def copy_config() -> None:
    shutil.copy2(CONFIG_PATH, OUTPUT_ROOT / "configs" / CONFIG_PATH.name)


def generate_site_schedules(config: dict, station: str, train_year: int) -> pd.DataFrame:
    search = config["offline_search"]
    n_values = [float(v) for v in search["grid"]["nitrogen_amounts"]]
    i_values = [float(v) for v in search["grid"]["irrigation_amounts"]]
    n_cap = float(search["grid"]["conservative_n_cap"])
    i_cap = float(search["grid"]["conservative_irrigation_cap"])
    rows: list[dict[str, Any]] = []
    sid = 0
    for n_combo in itertools.product(n_values, repeat=3):
        if sum(n_combo) > n_cap:
            continue
        for i_combo in itertools.product(i_values, repeat=3):
            if sum(i_combo) > i_cap:
                continue
            sid += 1
            row: dict[str, Any] = {
                "schedule_id": f"{station}{train_year}_S{sid:04d}",
                "station": station,
                "train_year": int(train_year),
                "total_N": float(sum(n_combo)),
                "total_irrigation": float(sum(i_combo)),
            }
            for event, amount in zip(search["nitrogen_events"], n_combo):
                row[f"{event['name']}_DAP"] = int(event["dap"])
                row[f"{event['name']}_amount"] = float(amount)
            for event, amount in zip(search["irrigation_events"], i_combo):
                row[f"{event['name']}_DAP"] = int(event["dap"])
                row[f"{event['name']}_amount"] = float(amount)
            rows.append(row)
    schedules = pd.DataFrame(rows)
    max_n = config["runtime"].get("max_hla_schedules")
    if max_n:
        schedules = schedules.head(int(max_n)).copy()
    out = OUTPUT_ROOT / "candidate_schedules" / f"{station}_{train_year}_coarse_grid_schedules.csv"
    schedules.to_csv(out, index=False, encoding="utf-8-sig")
    return schedules


def generate_hla_schedules(config: dict) -> pd.DataFrame:
    search = config["offline_search"]
    return generate_site_schedules(config, str(search["station"]), int(search["train_year"]))


def policy_from_schedule(config: dict, schedule: pd.Series) -> DeterministicSchedulePolicy:
    events: list[ScheduleEvent] = []
    for event in config["offline_search"]["nitrogen_events"]:
        name = str(event["name"])
        events.append(ScheduleEvent(name=name, dap=int(schedule[f"{name}_DAP"]), nitrogen=float(schedule[f"{name}_amount"])))
    for event in config["offline_search"]["irrigation_events"]:
        name = str(event["name"])
        events.append(ScheduleEvent(name=name, dap=int(schedule[f"{name}_DAP"]), irrigation=float(schedule[f"{name}_amount"])))
    return DeterministicSchedulePolicy(events)


def clip_to_action_space(env, real_action: dict[str, float]) -> tuple[dict[str, float], str]:
    spaces = getattr(env.formator.action_space_dict, "spaces", env.formator.action_space_dict)
    clipped: dict[str, float] = {}
    notes: list[str] = []
    for name in env.formator.action_names:
        space = spaces[name]
        low = float(np.asarray(space.low).flatten()[0])
        high = float(np.asarray(space.high).flatten()[0])
        raw = float(real_action.get(name, 0.0))
        value = min(max(raw, low), high)
        clipped[name] = value
        if abs(value - raw) > 1e-9:
            notes.append(f"{name}_env_bound_clip_{raw:g}_to_{value:g}")
    return clipped, ";".join(notes)


def profit_score(final_grnwt: float, irrigation: float, nitrogen: float, config: dict) -> float:
    econ = config["offline_search"]["economics"]
    return (
        float(econ["grain_value_coef"]) * float(final_grnwt)
        - float(econ["water_cost"]) * float(irrigation)
        - float(econ["n_cost"]) * float(nitrogen)
    )


def evaluate_schedule(config: dict, schedule: pd.Series, station: str, eval_year: int, save_daily: bool = True) -> dict:
    seed = int(config.get("seed", 0))
    schedule_id = str(schedule["schedule_id"])
    run_tag = f"{schedule_id}_eval{eval_year}"
    env = make_env(config, station, int(eval_year), seed, run_tag=run_tag, evaluation=True, action_safety_enabled=False)
    policy = policy_from_schedule(config, schedule)
    safety_state = ActionSafetyState()
    safety_config = {**config.get("action_safety", {}), "enabled": True}
    records: list[dict[str, Any]] = []
    cumulative_irrigation = 0.0
    cumulative_n = 0.0
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        year_info = find_year(config, station, int(eval_year))
        while not done and step_count < int(config.get("runtime", {}).get("max_steps", 260)):
            latest = latest_observation_dict(env, obs, info)
            dap_raw = scalar(latest.get("dap", np.nan))
            dap = int(round(dap_raw)) if not np.isnan(dap_raw) and dap_raw > 0 else step_count + 1
            planned = policy.action_for_dap(dap)
            bounded, env_clip_note = clip_to_action_space(env, planned)
            safety_result = apply_action_safety(bounded, dap, safety_state, safety_config)
            safe_action = safety_result.safe_real_action
            normalized_action = normalize_action(env.formator.action_names, env.formator.action_space_dict, safe_action)
            obs, reward, terminated, truncated, info = env.step(normalized_action)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            update_action_safety_state(safety_state, safe_action, dap)
            real_amir = float(safe_action.get("amir", 0.0))
            real_anfer = float(safe_action.get("anfer", 0.0))
            cumulative_irrigation += real_amir
            cumulative_n += real_anfer
            date = pd.Timestamp(year_info["planting_date"]) + pd.Timedelta(days=max(dap - 1, 0))
            totir_raw = scalar(latest.get("totir"))
            tofer_raw = scalar(latest.get("tofer"))
            records.append(
                {
                    "station": station,
                    "schedule_id": schedule_id,
                    "eval_year": int(eval_year),
                    "date": date.strftime("%Y-%m-%d"),
                    "year": int(date.year),
                    "doy": int(date.dayofyear),
                    "dap": scalar(latest.get("dap", dap)),
                    "planned_event_name": policy.event_name_for_dap(dap),
                    "is_schedule_event_day": policy.is_event_day(dap),
                    "planned_irrigation": planned.get("amir", 0.0),
                    "planned_n": planned.get("anfer", 0.0),
                    "env_bounded_irrigation": bounded.get("amir", 0.0),
                    "env_bounded_n": bounded.get("anfer", 0.0),
                    "real_action_amir": real_amir,
                    "real_action_anfer": real_anfer,
                    "normalized_action_amir": float(normalized_action[env.formator.action_names.index("amir")]) if "amir" in env.formator.action_names else np.nan,
                    "normalized_action_anfer": float(normalized_action[env.formator.action_names.index("anfer")]) if "anfer" in env.formator.action_names else np.nan,
                    "raw_real_action_amir": planned.get("amir", 0.0),
                    "raw_real_action_anfer": planned.get("anfer", 0.0),
                    "safe_real_action_amir": real_amir,
                    "safe_real_action_anfer": real_anfer,
                    "env_clip_note": env_clip_note,
                    "safety_rule_triggered": safety_result.safety_rule_triggered,
                    "topwt": scalar(latest.get("topwt")),
                    "grnwt": scalar(latest.get("grnwt")),
                    "xlai": scalar(latest.get("xlai")),
                    "totir_raw": totir_raw,
                    "totir": totir_raw if not np.isnan(totir_raw) else cumulative_irrigation,
                    "tofer_raw": tofer_raw,
                    "tofer": tofer_raw if not np.isnan(tofer_raw) else cumulative_n,
                    "swfac": scalar(latest.get("swfac")),
                    "nstres": scalar(latest.get("nstres")),
                    "reward": float(reward),
                    "done": done,
                    "info": json.dumps(info, ensure_ascii=False, default=str),
                }
            )
            step_count += 1
    finally:
        try:
            env.close()
        except Exception:
            pass
    daily = pd.DataFrame(records)
    daily_csv_path = ""
    if save_daily:
        daily_dir = OUTPUT_ROOT / "daily_outputs" / station
        daily_dir.mkdir(parents=True, exist_ok=True)
        daily_csv = daily_dir / f"{station}_{schedule_id}_eval{eval_year}_daily.csv"
        daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
        daily_csv_path = str(daily_csv.relative_to(PROJECT_ROOT))
    episode_completed = bool(records and records[-1]["done"])
    final_grnwt = float(daily["grnwt"].iloc[-1]) if len(daily) else np.nan
    total_irrig = float(daily["real_action_amir"].sum()) if len(daily) else 0.0
    total_n = float(daily["real_action_anfer"].sum()) if len(daily) else 0.0
    return {
        "schedule_id": schedule_id,
        "station": station,
        "eval_year": int(eval_year),
        "run_status": "ok" if episode_completed else "failed",
        "episode_completed": episode_completed,
        "episode_length": int(len(daily)),
        "final_grnwt": final_grnwt,
        "final_topwt": float(daily["topwt"].iloc[-1]) if len(daily) else np.nan,
        "final_xlai": float(daily["xlai"].iloc[-1]) if len(daily) else np.nan,
        "total_irrigation": total_irrig,
        "total_n_fertilizer": total_n,
        "planned_total_irrigation": float(schedule.get("total_irrigation", 0.0)),
        "planned_total_n": float(schedule.get("total_N", 0.0)),
        "mean_swfac": float(pd.to_numeric(daily["swfac"], errors="coerce").mean()) if len(daily) else np.nan,
        "mean_nstres": float(pd.to_numeric(daily["nstres"], errors="coerce").mean()) if len(daily) else np.nan,
        "yield_per_100mm_irrigation": final_grnwt / (total_irrig / 100.0) if total_irrig > 0 else np.nan,
        "yield_per_100kg_n": final_grnwt / (total_n / 100.0) if total_n > 0 else np.nan,
        "profit_score": profit_score(final_grnwt, total_irrig, total_n, config) if not np.isnan(final_grnwt) else np.nan,
        "daily_csv_path": daily_csv_path,
        "notes": "",
    }


def run_site_coarse_grid(config: dict, station: str, train_year: int, schedules: pd.DataFrame) -> pd.DataFrame:
    out = OUTPUT_ROOT / "evaluation" / f"{station}_{train_year}_coarse_grid_summary.csv"
    done_ids: set[str] = set()
    existing = pd.DataFrame()
    if out.exists():
        existing = pd.read_csv(out)
        done_ids = set(existing.loc[existing["run_status"].eq("ok"), "schedule_id"].astype(str))
    rows: list[dict] = []
    for idx, schedule in schedules.iterrows():
        sid = str(schedule["schedule_id"])
        if sid in done_ids:
            continue
        if idx % 25 == 0:
            print(f"[offline-search] {station} {train_year} coarse grid {idx + 1}/{len(schedules)} {sid}", flush=True)
        try:
            rows.append(evaluate_schedule(config, schedule, station, train_year, save_daily=True))
        except Exception as exc:
            rows.append(
                {
                    "schedule_id": sid,
                    "station": station,
                    "eval_year": int(train_year),
                    "run_status": "failed",
                    "episode_completed": False,
                    "error_message": f"{type(exc).__name__}: {exc}",
                    "traceback": traceback.format_exc()[-2000:],
                }
            )
        if rows and len(rows) % 25 == 0:
            merged = pd.concat([existing, pd.DataFrame(rows)], ignore_index=True) if len(existing) else pd.DataFrame(rows)
            merged = merged.drop_duplicates(["schedule_id", "station", "eval_year"], keep="last")
            merged.to_csv(out, index=False, encoding="utf-8-sig")
    summary = pd.concat([existing, pd.DataFrame(rows)], ignore_index=True) if len(existing) else pd.DataFrame(rows)
    summary = summary.drop_duplicates(["schedule_id", "station", "eval_year"], keep="last")
    summary.to_csv(out, index=False, encoding="utf-8-sig")
    return summary


def run_hla_coarse_grid(config: dict, schedules: pd.DataFrame) -> pd.DataFrame:
    return run_site_coarse_grid(config, "HLA", 2011, schedules)


def pareto_frontier(df: pd.DataFrame) -> pd.Series:
    work = df.copy()
    y = pd.to_numeric(work["final_grnwt"], errors="coerce").fillna(-np.inf).to_numpy()
    water = pd.to_numeric(work["total_irrigation"], errors="coerce").fillna(np.inf).to_numpy()
    n = pd.to_numeric(work["total_n_fertilizer"], errors="coerce").fillna(np.inf).to_numpy()
    is_pareto = np.ones(len(work), dtype=bool)
    for i in range(len(work)):
        if not is_pareto[i]:
            continue
        dominates = (y >= y[i]) & (water <= water[i]) & (n <= n[i]) & ((y > y[i]) | (water < water[i]) | (n < n[i]))
        if dominates.any():
            is_pareto[i] = False
    return pd.Series(is_pareto, index=df.index)


def select_site_candidates(config: dict, station: str, train_year: int, summary: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    ok = summary[summary["run_status"].eq("ok") & summary["episode_completed"].astype(str).str.lower().isin(["true", "1"])].copy()
    ref_yield = float(config["offline_search"]["ppo_reference"]["mean_yield"])
    ref_i = float(config["offline_search"]["ppo_reference"]["total_irrigation"])
    ref_n = float(config["offline_search"]["ppo_reference"]["total_n"])
    ok["yield_loss_vs_ppo_baseline"] = (ref_yield - ok["final_grnwt"]) / ref_yield
    ok["irrigation_reduction_vs_ppo_baseline"] = 1.0 - ok["total_irrigation"] / ref_i
    ok["n_reduction_vs_ppo_baseline"] = 1.0 - ok["total_n_fertilizer"] / ref_n
    ok["is_pareto"] = pareto_frontier(ok)
    pareto = ok[ok["is_pareto"]].copy().sort_values(["final_grnwt", "profit_score"], ascending=[False, False])
    pareto.to_csv(OUTPUT_ROOT / "evaluation" / f"{station}_{train_year}_pareto_frontier.csv", index=False, encoding="utf-8-sig")

    labels: list[pd.DataFrame] = []
    top_n = int(config["offline_search"]["top_selection"]["top_n_each"])
    for label, part in [
        ("top_yield", ok.sort_values("final_grnwt", ascending=False).head(top_n)),
        ("top_profit", ok.sort_values("profit_score", ascending=False).head(top_n)),
        ("pareto_balanced", pareto.sort_values(["yield_loss_vs_ppo_baseline", "profit_score"], ascending=[True, False]).head(top_n)),
        (
            "low_input_within_10pct_yield_loss",
            ok[ok["yield_loss_vs_ppo_baseline"] <= float(config["offline_search"]["top_selection"]["low_input_yield_loss_threshold"])]
            .assign(total_input=lambda x: x["total_irrigation"] + x["total_n_fertilizer"])
            .sort_values(["total_input", "profit_score"], ascending=[True, False])
            .head(top_n),
        ),
    ]:
        if len(part):
            temp = part.copy()
            temp["expert_candidate_type"] = label
            labels.append(temp)
    candidates = pd.concat(labels, ignore_index=True) if labels else pd.DataFrame()
    if len(candidates):
        candidates = candidates.sort_values(["schedule_id", "expert_candidate_type"]).drop_duplicates("schedule_id", keep="first")
        candidates = candidates.sort_values(["yield_loss_vs_ppo_baseline", "profit_score"], ascending=[True, False])
        max_unique = int(config["offline_search"]["top_selection"]["max_unique_schedules"])
        candidates = candidates.head(max_unique)
    candidates.to_csv(OUTPUT_ROOT / "evaluation" / f"{station}_{train_year}_expert_schedule_candidates.csv", index=False, encoding="utf-8-sig")
    return pareto, candidates


def select_candidates(config: dict, summary: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    return select_site_candidates(config, "HLA", 2011, summary)


def schedule_lookup(schedules: pd.DataFrame) -> dict[str, pd.Series]:
    return {str(row["schedule_id"]): row for _, row in schedules.iterrows()}


def run_site_cross_year(config: dict, station: str, schedules: pd.DataFrame, candidates: pd.DataFrame, eval_years: list[int]) -> pd.DataFrame:
    out = OUTPUT_ROOT / "evaluation" / f"{station}_top_schedule_cross_year_summary.csv"
    if candidates.empty:
        empty = pd.DataFrame()
        empty.to_csv(out, index=False, encoding="utf-8-sig")
        return empty
    lookup = schedule_lookup(schedules)
    rows: list[dict] = []
    for sid in candidates["schedule_id"].astype(str).tolist():
        schedule = lookup[sid]
        for eval_year in eval_years:
            print(f"[offline-search] cross-year {station} {sid} eval {eval_year}", flush=True)
            rows.append(evaluate_schedule(config, schedule, station, int(eval_year), save_daily=True))
    cross = pd.DataFrame(rows)
    cross.to_csv(out, index=False, encoding="utf-8-sig")
    return cross


def run_cross_year(config: dict, schedules: pd.DataFrame, candidates: pd.DataFrame) -> pd.DataFrame:
    return run_site_cross_year(config, "HLA", schedules, candidates, [int(y) for y in config["offline_search"]["eval_years"]])


def rank_site_expert_schedules(config: dict, station: str, cross: pd.DataFrame) -> pd.DataFrame:
    out = OUTPUT_ROOT / "expert_policy" / f"{station}_expert_schedule_ranking.csv"
    if cross.empty:
        pd.DataFrame().to_csv(out, index=False, encoding="utf-8-sig")
        return pd.DataFrame()
    ref_yield = float(config["offline_search"]["ppo_reference"]["mean_yield"])
    grouped = (
        cross.groupby("schedule_id")
        .agg(
            eval_count=("eval_year", "count"),
            ok_count=("run_status", lambda s: int((s == "ok").sum())),
            mean_yield=("final_grnwt", "mean"),
            std_yield=("final_grnwt", "std"),
            mean_profit=("profit_score", "mean"),
            mean_irrigation=("total_irrigation", "mean"),
            mean_n=("total_n_fertilizer", "mean"),
            mean_swfac=("mean_swfac", "mean"),
            mean_nstres=("mean_nstres", "mean"),
        )
        .reset_index()
    )
    grouped["mean_yield_loss_vs_ppo_baseline"] = (ref_yield - grouped["mean_yield"]) / ref_yield

    def minmax(series: pd.Series, invert: bool = False) -> pd.Series:
        values = pd.to_numeric(series, errors="coerce")
        if values.max() == values.min():
            score = pd.Series(1.0, index=series.index)
        else:
            score = (values - values.min()) / (values.max() - values.min())
        return 1.0 - score if invert else score

    grouped["yield_stability_score"] = 0.65 * minmax(grouped["mean_yield"]) + 0.35 * minmax(grouped["std_yield"].fillna(0.0), invert=True)
    grouped["profit_stability_score"] = minmax(grouped["mean_profit"])
    grouped["overall_score"] = (
        0.40 * minmax(grouped["mean_yield"])
        + 0.25 * minmax(grouped["mean_profit"])
        + 0.15 * minmax(grouped["std_yield"].fillna(0.0), invert=True)
        + 0.10 * minmax(grouped["mean_irrigation"], invert=True)
        + 0.10 * minmax(grouped["mean_n"], invert=True)
    )
    grouped = grouped.sort_values("overall_score", ascending=False)
    grouped.to_csv(out, index=False, encoding="utf-8-sig")
    return grouped


def rank_expert_schedules(config: dict, cross: pd.DataFrame) -> pd.DataFrame:
    return rank_site_expert_schedules(config, "HLA", cross)


def generate_imitation_dataset(ranking: pd.DataFrame, cross: pd.DataFrame) -> Path | None:
    if ranking.empty:
        return None
    best_id = str(ranking.iloc[0]["schedule_id"])
    rows = []
    for path in cross.loc[cross["schedule_id"].astype(str).eq(best_id), "daily_csv_path"].dropna().astype(str):
        daily_path = PROJECT_ROOT / path
        if not daily_path.exists():
            continue
        daily = pd.read_csv(daily_path)
        state_cols = [c for c in ["dap", "topwt", "grnwt", "xlai", "swfac", "nstres", "totir", "tofer"] if c in daily.columns]
        for _, row in daily.iterrows():
            rows.append(
                {
                    "station": row.get("station", ""),
                    "year": row.get("eval_year", row.get("year", "")),
                    "date": row.get("date", ""),
                    "dap": row.get("dap", ""),
                    "state_variables": json.dumps({col: row.get(col, np.nan) for col in state_cols}, ensure_ascii=False, default=str),
                    "expert_action_irrigation": row.get("real_action_amir", 0.0),
                    "expert_action_n": row.get("real_action_anfer", 0.0),
                    "schedule_id": best_id,
                    "expert_policy_type": "offline_schedule_search_best_overall",
                }
            )
    if not rows:
        return None
    out = OUTPUT_ROOT / "expert_policy" / "imitation_dataset.csv"
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")
    return out


def plot_outputs(summary: pd.DataFrame, pareto: pd.DataFrame, cross: pd.DataFrame, ranking: pd.DataFrame) -> None:
    fig_dir = OUTPUT_ROOT / "figures" / "HLA"
    fig_dir.mkdir(parents=True, exist_ok=True)
    ok = summary[summary["run_status"].eq("ok")].copy()
    if len(ok):
        plt.figure(figsize=(8.5, 5.0))
        plt.scatter(ok["total_irrigation"], ok["final_grnwt"], s=22, alpha=0.55, label="all schedules")
        if len(pareto):
            plt.scatter(pareto["total_irrigation"], pareto["final_grnwt"], s=32, label="pareto", color="#d95f02")
        plt.xlabel("Total irrigation (mm)")
        plt.ylabel("Grain yield")
        plt.grid(alpha=0.25)
        plt.legend(frameon=False)
        plt.tight_layout()
        plt.savefig(fig_dir / "HLA_yield_vs_irrigation.png", dpi=180)
        plt.close()

        plt.figure(figsize=(8.5, 5.0))
        plt.scatter(ok["total_n_fertilizer"], ok["final_grnwt"], s=22, alpha=0.55)
        plt.xlabel("Total nitrogen (kg/ha)")
        plt.ylabel("Grain yield")
        plt.grid(alpha=0.25)
        plt.tight_layout()
        plt.savefig(fig_dir / "HLA_yield_vs_nitrogen.png", dpi=180)
        plt.close()

        plt.figure(figsize=(8.5, 5.0))
        plt.scatter(ok["total_irrigation"] + ok["total_n_fertilizer"], ok["profit_score"], s=22, alpha=0.55)
        plt.xlabel("Total water + nitrogen input")
        plt.ylabel("Profit score")
        plt.grid(alpha=0.25)
        plt.tight_layout()
        plt.savefig(fig_dir / "HLA_profit_vs_input.png", dpi=180)
        plt.close()

        fig = plt.figure(figsize=(8.5, 6.0))
        ax = fig.add_subplot(111, projection="3d")
        ax.scatter(ok["total_irrigation"], ok["total_n_fertilizer"], ok["final_grnwt"], s=12, alpha=0.35)
        if len(pareto):
            ax.scatter(pareto["total_irrigation"], pareto["total_n_fertilizer"], pareto["final_grnwt"], s=28, color="#d95f02")
        ax.set_xlabel("Irrigation")
        ax.set_ylabel("N")
        ax.set_zlabel("Yield")
        fig.tight_layout()
        fig.savefig(fig_dir / "HLA_pareto_frontier_yield_water_n.png", dpi=180)
        plt.close(fig)

    if len(cross):
        plt.figure(figsize=(10, 5.5))
        top_ids = ranking.head(8)["schedule_id"].astype(str).tolist() if len(ranking) else cross["schedule_id"].astype(str).unique()[:8].tolist()
        subset = cross[cross["schedule_id"].astype(str).isin(top_ids)]
        for sid, part in subset.groupby("schedule_id"):
            plt.plot(part["eval_year"].astype(str), part["final_grnwt"], marker="o", label=sid)
        plt.xlabel("Eval year")
        plt.ylabel("Grain yield")
        plt.grid(alpha=0.25)
        plt.legend(frameon=False, fontsize=7)
        plt.tight_layout()
        plt.savefig(fig_dir / "HLA_top_schedules_cross_year_yield.png", dpi=180)
        plt.close()

        plt.figure(figsize=(10, 5.5))
        for sid, part in subset.groupby("schedule_id"):
            plt.plot(part["eval_year"].astype(str), part["profit_score"], marker="o", label=sid)
        plt.xlabel("Eval year")
        plt.ylabel("Profit score")
        plt.grid(alpha=0.25)
        plt.legend(frameon=False, fontsize=7)
        plt.tight_layout()
        plt.savefig(fig_dir / "HLA_top_schedules_cross_year_profit.png", dpi=180)
        plt.close()

    if len(ranking):
        best = str(ranking.iloc[0]["schedule_id"])
        daily_paths = cross.loc[cross["schedule_id"].astype(str).eq(best), "daily_csv_path"].dropna().astype(str).tolist()
        frames = []
        for p in daily_paths:
            path = PROJECT_ROOT / p
            if path.exists():
                frames.append(pd.read_csv(path))
        if frames:
            daily = pd.concat(frames, ignore_index=True)
            plt.figure(figsize=(10, 5.5))
            for year, part in daily.groupby("eval_year"):
                plt.plot(part["dap"], part["real_action_amir"], label=f"{year} irrigation")
                plt.plot(part["dap"], part["real_action_anfer"], linestyle="--", label=f"{year} N")
            plt.xlabel("DAP")
            plt.ylabel("Action amount")
            plt.grid(alpha=0.25)
            plt.legend(frameon=False, fontsize=7)
            plt.tight_layout()
            plt.savefig(fig_dir / "HLA_expert_schedule_actions.png", dpi=180)
            plt.close()

            plt.figure(figsize=(10, 5.5))
            for year, part in daily.groupby("eval_year"):
                plt.plot(part["dap"], part["swfac"], label=f"{year} swfac")
                plt.plot(part["dap"], part["nstres"], linestyle="--", label=f"{year} nstres")
            plt.xlabel("DAP")
            plt.ylabel("Stress indicators")
            plt.grid(alpha=0.25)
            plt.legend(frameon=False, fontsize=7)
            plt.tight_layout()
            plt.savefig(fig_dir / "HLA_swfac_nstres_for_top_schedules.png", dpi=180)
            plt.close()

    summary_fig = OUTPUT_ROOT / "figures" / "summary" / "all_available_sites_expert_schedule_summary.png"
    if len(ranking):
        top = ranking.head(10).copy()
        plt.figure(figsize=(10, 5.5))
        plt.bar(top["schedule_id"], top["overall_score"])
        plt.xticks(rotation=45, ha="right")
        plt.ylabel("Overall score")
        plt.tight_layout()
        plt.savefig(summary_fig, dpi=180)
        plt.close()


def plot_site_outputs(station: str, summary: pd.DataFrame, pareto: pd.DataFrame, cross: pd.DataFrame, ranking: pd.DataFrame) -> None:
    if station == "HLA":
        plot_outputs(summary, pareto, cross, ranking)
        return
    fig_dir = OUTPUT_ROOT / "figures" / station
    fig_dir.mkdir(parents=True, exist_ok=True)
    ok = summary[summary["run_status"].eq("ok")].copy() if len(summary) else pd.DataFrame()
    if len(ok):
        plt.figure(figsize=(8.5, 5.0))
        plt.scatter(ok["total_irrigation"], ok["final_grnwt"], s=22, alpha=0.55, label="all schedules")
        if len(pareto):
            plt.scatter(pareto["total_irrigation"], pareto["final_grnwt"], s=32, color="#d95f02", label="pareto")
        plt.xlabel("Total irrigation (mm)")
        plt.ylabel("Grain yield")
        plt.grid(alpha=0.25)
        plt.legend(frameon=False)
        plt.tight_layout()
        plt.savefig(fig_dir / f"{station}_yield_vs_irrigation.png", dpi=180)
        plt.close()

        plt.figure(figsize=(8.5, 5.0))
        plt.scatter(ok["total_n_fertilizer"], ok["final_grnwt"], s=22, alpha=0.55)
        plt.xlabel("Total nitrogen (kg/ha)")
        plt.ylabel("Grain yield")
        plt.grid(alpha=0.25)
        plt.tight_layout()
        plt.savefig(fig_dir / f"{station}_yield_vs_nitrogen.png", dpi=180)
        plt.close()

        plt.figure(figsize=(8.5, 5.0))
        plt.scatter(ok["total_irrigation"] + ok["total_n_fertilizer"], ok["profit_score"], s=22, alpha=0.55)
        plt.xlabel("Total water + nitrogen input")
        plt.ylabel("Profit score")
        plt.grid(alpha=0.25)
        plt.tight_layout()
        plt.savefig(fig_dir / f"{station}_profit_vs_input.png", dpi=180)
        plt.close()

    if len(cross):
        top_ids = ranking.head(8)["schedule_id"].astype(str).tolist() if len(ranking) else cross["schedule_id"].astype(str).unique()[:8].tolist()
        subset = cross[cross["schedule_id"].astype(str).isin(top_ids)]
        plt.figure(figsize=(10, 5.5))
        for sid, part in subset.groupby("schedule_id"):
            plt.plot(part["eval_year"].astype(str), part["final_grnwt"], marker="o", label=sid)
        plt.xlabel("Eval year")
        plt.ylabel("Grain yield")
        plt.grid(alpha=0.25)
        plt.legend(frameon=False, fontsize=7)
        plt.tight_layout()
        plt.savefig(fig_dir / f"{station}_top_schedules_cross_year_yield.png", dpi=180)
        plt.close()

        plt.figure(figsize=(10, 5.5))
        for sid, part in subset.groupby("schedule_id"):
            plt.plot(part["eval_year"].astype(str), part["profit_score"], marker="o", label=sid)
        plt.xlabel("Eval year")
        plt.ylabel("Profit score")
        plt.grid(alpha=0.25)
        plt.legend(frameon=False, fontsize=7)
        plt.tight_layout()
        plt.savefig(fig_dir / f"{station}_top_schedules_cross_year_profit.png", dpi=180)
        plt.close()


def extension_decision(config: dict, ranking: pd.DataFrame) -> tuple[bool, str]:
    if ranking.empty:
        return False, "No HLA ranking was generated."
    best = ranking.iloc[0]
    gate = config["offline_search"]["extension_gate"]
    ok = (
        float(best["mean_yield_loss_vs_ppo_baseline"]) <= float(gate["mean_yield_loss_vs_ppo_baseline"])
        and float(best["mean_irrigation"]) <= float(gate["mean_irrigation_max"])
        and float(best["mean_n"]) <= float(gate["mean_n_max"])
        and int(best["ok_count"]) == int(best["eval_count"])
    )
    if ok:
        return True, f"HLA best schedule {best['schedule_id']} passed the low-input <=10% yield-loss gate."
    return False, f"HLA best schedule {best['schedule_id']} did not pass the <=10% yield-loss and low-input gate."


def run_extension_sites(config: dict, should_extend: bool) -> dict[str, dict[str, Any]]:
    results: dict[str, dict[str, Any]] = {}
    if not should_extend:
        for station in ["SYA", "LCA"]:
            out = OUTPUT_ROOT / "expert_policy" / f"{station}_expert_schedule_ranking.csv"
            pd.DataFrame([{"station": station, "run_status": "skipped", "notes": "HLA did not pass extension gate."}]).to_csv(out, index=False, encoding="utf-8-sig")
            results[station] = {"run_status": "skipped", "ranking_rows": 0}
        return results
    for station, site_cfg in config.get("extension_sites", {}).items():
        train_year = int(site_cfg["train_year"])
        eval_years = [int(y) for y in site_cfg["eval_years"]]
        print(f"[offline-search] extension {station} train_year {train_year}", flush=True)
        schedules = generate_site_schedules(config, station, train_year)
        summary = run_site_coarse_grid(config, station, train_year, schedules)
        pareto, candidates = select_site_candidates(config, station, train_year, summary)
        cross = run_site_cross_year(config, station, schedules, candidates, eval_years)
        ranking = rank_site_expert_schedules(config, station, cross)
        plot_site_outputs(station, summary, pareto, cross, ranking)
        results[station] = {
            "run_status": "ok",
            "schedule_count": int(len(schedules)),
            "coarse_ok": int((summary["run_status"] == "ok").sum()) if len(summary) else 0,
            "candidate_count": int(len(candidates)),
            "cross_year_count": int(len(cross)),
            "ranking_rows": int(len(ranking)),
            "best_schedule_id": str(ranking.iloc[0]["schedule_id"]) if len(ranking) else "",
            "best_mean_yield": float(ranking.iloc[0]["mean_yield"]) if len(ranking) else np.nan,
            "best_mean_irrigation": float(ranking.iloc[0]["mean_irrigation"]) if len(ranking) else np.nan,
            "best_mean_n": float(ranking.iloc[0]["mean_n"]) if len(ranking) else np.nan,
        }
    return results


def write_reports(config: dict, schedules: pd.DataFrame, summary: pd.DataFrame, pareto: pd.DataFrame, candidates: pd.DataFrame, cross: pd.DataFrame, ranking: pd.DataFrame, imitation_path: Path | None, extend: bool, extension_note: str) -> None:
    best = ranking.head(1).copy()
    top_table_cols = [
        "schedule_id",
        "mean_yield",
        "std_yield",
        "mean_profit",
        "mean_irrigation",
        "mean_n",
        "mean_yield_loss_vs_ppo_baseline",
        "overall_score",
    ]
    coarse_cols = ["schedule_id", "final_grnwt", "total_irrigation", "total_n_fertilizer", "profit_score", "yield_per_100mm_irrigation", "yield_per_100kg_n"]
    ok_summary = summary[summary["run_status"].eq("ok")].copy()
    low10 = candidates[candidates.get("yield_loss_vs_ppo_baseline", pd.Series(dtype=float)) <= 0.10] if len(candidates) else pd.DataFrame()
    extension_rows = []
    for station in ["SYA", "LCA"]:
        p = OUTPUT_ROOT / "expert_policy" / f"{station}_expert_schedule_ranking.csv"
        if p.exists():
            site = pd.read_csv(p)
            if "overall_score" in site.columns and len(site):
                first = site.iloc[0]
                extension_rows.append(
                    {
                        "station": station,
                        "run_status": "ok",
                        "best_schedule_id": first.get("schedule_id", ""),
                        "mean_yield": first.get("mean_yield", np.nan),
                        "mean_irrigation": first.get("mean_irrigation", np.nan),
                        "mean_n": first.get("mean_n", np.nan),
                        "overall_score": first.get("overall_score", np.nan),
                    }
                )
            elif len(site):
                extension_rows.append(site.iloc[0].to_dict())
    extension_table = pd.DataFrame(extension_rows)

    md = f"""# Offline schedule search report

Generated at: 2026-06-06

## Goal

This stage pauses PPO tuning and uses deterministic DSSAT schedule search to identify an expert prior for water and nitrogen management. It does not train PPO, does not enter rainfall scaling, and does not modify my_data or previous 006_03-006_07 outputs.

## Search Space

- Station/year: HLA 2011 coarse grid.
- Nitrogen events: DAP 1, 30, 60.
- Irrigation events: DAP 25, 50, 75.
- N event amounts: {config['offline_search']['grid']['nitrogen_amounts']}.
- Irrigation event amounts: {config['offline_search']['grid']['irrigation_amounts']}.
- Conservative caps: total N <= {config['offline_search']['grid']['conservative_n_cap']} kg/ha, total irrigation <= {config['offline_search']['grid']['conservative_irrigation_cap']} mm.
- Profit score: 0.01 * grain yield - 0.5 * irrigation - 0.25 * nitrogen.

## Why PPO Is Paused

006_03-006_07 repeatedly saturated 300/450 under cost rewards, terminal profit, low-frequency/window action designs, and explicit budget/event wrappers. The offline search tests whether reasonable non-RL schedules exist before returning to imitation learning or constrained PPO.

## HLA Coarse Grid Result

- Candidate schedules generated: {len(schedules)}
- Completed HLA 2011 evaluations: {int((summary['run_status'] == 'ok').sum()) if len(summary) else 0}
- Pareto schedules: {len(pareto)}
- Expert candidates selected for cross-year validation: {len(candidates)}
- Low-input candidates within 10% of PPO baseline yield: {len(low10)}

## Top HLA 2011 Schedules

{df_to_markdown(ok_summary.sort_values('profit_score', ascending=False)[[c for c in coarse_cols if c in ok_summary.columns]].head(10))}

## Cross-Year Ranking

{df_to_markdown(ranking[[c for c in top_table_cols if c in ranking.columns]].head(10))}

## Best Schedule

{df_to_markdown(best[[c for c in top_table_cols if c in best.columns]]) if len(best) else 'No best schedule was available.'}

## Extension Decision

- Extend SYA/LCA: {extend}
- Reason: {extension_note}
- Imitation dataset: {str(imitation_path.relative_to(PROJECT_ROOT)) if imitation_path else 'not generated'}

## SYA/LCA Extension

{df_to_markdown(extension_table) if len(extension_table) else 'Extension was not executed.'}

## Interpretation

If HLA identifies schedules with lower water/N and acceptable yield loss, these schedules should be treated as expert priors, not final RL policies. They can be used for imitation learning, constrained PPO target shaping, or reward calibration. Rainfall-scaling budget scenarios should remain a separate stress-test path.
"""
    DOC_MD.parent.mkdir(parents=True, exist_ok=True)
    DOC_MD.write_text(md, encoding="utf-8")
    shutil.copy2(DOC_MD, OUTPUT_ROOT / "reports" / DOC_MD.name)
    write_ppt(schedules, summary, pareto, candidates, ranking, imitation_path, extend, extension_note)


def write_ppt(schedules: pd.DataFrame, summary: pd.DataFrame, pareto: pd.DataFrame, candidates: pd.DataFrame, ranking: pd.DataFrame, imitation_path: Path | None, extend: bool, extension_note: str) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        (OUTPUT_ROOT / "reports" / "ppt_generation_skipped_in_container.txt").write_text(
            "python-pptx is not installed in the DSSAT container. Generate the PPT from Windows Python after search.",
            encoding="utf-8",
        )
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_title(slide, text: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.45))
        p = box.text_frame.paragraphs[0]
        p.text = text
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.bold = True

    def add_body(slide, text: str, size: int = 15) -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.0), Inches(5.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(size)

    def add_picture(title: str, path: Path) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, title)
        if path.exists():
            slide.shapes.add_picture(str(path), Inches(0.8), Inches(1.0), width=Inches(11.7))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "006_08 Offline Schedule Search")
    add_body(slide, "目的：暂停 PPO 小修小补，先用 deterministic schedule search 找水氮管理 expert prior。\n范围：HLA 2011 coarse grid；固定 N DAP 1/30/60，灌溉 DAP 25/50/75。\n约束：不训练 PPO，不改 my_data，不覆盖 006_03 到 006_07。")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Why Pause PPO")
    add_body(slide, "006_03 到 006_07 均显示 PPO 倾向打满 300/450。\nreward cost、terminal profit、低频动作、窗口门控、预算/事件动作都没有产生主动节约。\n因此先用 DSSAT scenario ensemble 找可解释的 expert schedule。")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Search Space and Outputs")
    add_body(slide, f"候选 schedules：{len(schedules)}。\n完成 HLA 2011 evaluations：{int((summary['run_status'] == 'ok').sum()) if len(summary) else 0}。\nPareto schedules：{len(pareto)}。\n跨年验证 candidates：{len(candidates)}。")
    for title, filename in [
        ("Yield vs Irrigation", "HLA_yield_vs_irrigation.png"),
        ("Yield vs Nitrogen", "HLA_yield_vs_nitrogen.png"),
        ("Profit vs Input", "HLA_profit_vs_input.png"),
        ("Pareto Frontier", "HLA_pareto_frontier_yield_water_n.png"),
        ("Cross-Year Yield", "HLA_top_schedules_cross_year_yield.png"),
        ("Cross-Year Profit", "HLA_top_schedules_cross_year_profit.png"),
        ("Expert Schedule Actions", "HLA_expert_schedule_actions.png"),
        ("SWFAC and NSTRES", "HLA_swfac_nstres_for_top_schedules.png"),
    ]:
        add_picture(title, OUTPUT_ROOT / "figures" / "HLA" / filename)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Conclusion")
    best = ranking.iloc[0].to_dict() if len(ranking) else {}
    ext_lines = []
    for station in ["SYA", "LCA"]:
        p = OUTPUT_ROOT / "expert_policy" / f"{station}_expert_schedule_ranking.csv"
        if p.exists():
            site = pd.read_csv(p)
            if len(site) and "overall_score" in site.columns:
                row = site.iloc[0]
                ext_lines.append(f"{station}: {row['schedule_id']}, yield={float(row['mean_yield']):.1f}, irrigation={float(row['mean_irrigation']):.1f}, N={float(row['mean_n']):.1f}")
            elif len(site):
                ext_lines.append(f"{station}: {site.iloc[0].get('run_status', 'unknown')}")
    add_body(
        slide,
        f"Best schedule：{best.get('schedule_id', 'None')}。\n"
        f"Mean yield：{best.get('mean_yield', float('nan')):.2f}；Mean irrigation：{best.get('mean_irrigation', float('nan')):.2f}；Mean N：{best.get('mean_n', float('nan')):.2f}。\n"
        f"Expand SYA/LCA：{extend}。\nReason：{extension_note}\n"
        f"Imitation dataset：{str(imitation_path.relative_to(PROJECT_ROOT)) if imitation_path else 'not generated'}\n"
        + ("\n".join(ext_lines) if ext_lines else ""),
        size=14,
    )
    DOC_PPT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DOC_PPT)
    shutil.copy2(DOC_PPT, OUTPUT_ROOT / "reports" / DOC_PPT.name)


def main() -> None:
    ensure_dirs()
    config = load_yaml(CONFIG_PATH)
    copy_config()
    write_pause_note()
    schedules = generate_hla_schedules(config)
    print(f"[offline-search] generated schedules: {len(schedules)}", flush=True)
    summary = run_hla_coarse_grid(config, schedules)
    pareto, candidates = select_candidates(config, summary)
    cross = run_cross_year(config, schedules, candidates)
    ranking = rank_expert_schedules(config, cross)
    imitation_path = generate_imitation_dataset(ranking, cross)
    plot_outputs(summary, pareto, cross, ranking)
    extend, extension_note = extension_decision(config, ranking)
    run_extension_sites(config, extend)
    write_reports(config, schedules, summary, pareto, candidates, cross, ranking, imitation_path, extend, extension_note)
    print("[offline-search] done", flush=True)
    print((OUTPUT_ROOT / "evaluation" / "HLA_2011_coarse_grid_summary.csv").relative_to(PROJECT_ROOT), flush=True)
    print((OUTPUT_ROOT / "expert_policy" / "HLA_expert_schedule_ranking.csv").relative_to(PROJECT_ROOT), flush=True)


if __name__ == "__main__":
    main()
