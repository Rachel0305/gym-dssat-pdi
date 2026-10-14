from pathlib import Path
import json,pandas as pd
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
R=Path.cwd();O=R/'results/hla_random_weather_8seed/011_ghcn50756_fallback';gate=json.loads((O/'final_gate.json').read_text());met=json.loads((O/'overlap_daily_metrics.json').read_text())['metrics']
report='''# 011 HLA GHCN 50756 降水补源审计

执行日期：2026-10-04。仅使用 2004–2013 拟合期证据。

## 结论

`HLA_2004_GHCN_FALLBACK = ACCEPTED`，限 2004-01-01 至 2004-08-31。

`HLA_RAIN_PROVENANCE = FAIL`，`READY_TO_RESUME_010 = FALSE`。

011 文档将目标窗口补源通过与完整拟合期通过联系起来，但其前提“其余日期已有认可的来源方案”并未在 009/010 成立。009 的 MIXED 空白语义仍然有效。此次补源解除 244 天目标缺口，余下 2,447 个本地空白日期仍需确定来源，不能擅自填零，也不能扩大此次 GHCN 授权窗口。因此不生成名义完整的 3,653 天候选。

## 原始 provenance、空间及 QC

复用 008 官方 NOAA 缓存，SHA-256 与 008 存档完全一致：

`3925709643D003BCD53B53AD37D0F7A4816F10E60EDC73E83D63EBF214139232`

原路径：`results/hl_fq_random_weather_8seed/source_audit/ghcn_daily_raw/CHM00050756.dly`。

HAILUN：47.433°N、126.967°E、240 m。HLA 冻结 WTH 表头：47.450°N、126.900°E、234 m。大圆距离 5.3817 km，高差 6 m。距离基于表头精度；没有独立地形调查，不推断站点类型。结合事件和量级证据，空间代表性为 ACCEPTABLE。

依据 NOAA GHCN-Daily 官方 README：PRCP 原值除以10换算 mm，-9999为缺测，非空 QFLAG 排除，MFLAG/SFLAG 原样保留。MFLAG T 为微量，数值0不另造正降水；SFLAG s 和 S 不能混同。未修改原始记录。

2004 年 366/366 天可用，目标 244/244 天可用；无重复、负降水或非空质量标记。目标累计278.9 mm、71个数值>0湿日、最大33.7 mm；全年357.3 mm。2004–2013共3,442/3,653天可用，211个缺测均在2013，不能把2013的820.9 mm称为完整年总量。

## 日界与事件

实际非空 D222 数值才作为真值；全空白不进入 wet/dry truth。旧表复核得到相同结果，不视为独立真值。

| GHCN 日期 shift | 可比日 | 事件一致率 | MAE mm | Pearson |
|---|---:|---:|---:|---:|
| -1 | 912 | 53.07% | 6.174 | 0.115 |
| 0 | 935 | 85.13% | 1.913 | 0.802 |
| +1 | 913 | 53.23% | 6.001 | 0.181 |

0日明显最合理，未移动正式序列。0日RMSE=6.813 mm，Spearman=0.871，事件命中率93.40%，false wet=89、false dry=50，累计差=-115.5 mm。真值样本有177个明确零值，但仍偏向有记录日，不能据此推断所有未知空白都是干日。

分别保留2004年9–12月与2005–2013窗口。最大单日分别出现在本地2013-07-30和GHCN 2013-08-12，不能把不同极值日期当成统一日界偏移。

## 月、年与极端

2005–2012 GHCN 可用累计/本地实际数值小计为0.987–1.146。可比日的月累计比例中位数1.054，范围0–4.042。低雨量月份相对误差较大，未发现贯穿年份的×10或÷10量级错误。所有本地月年累计明确标记为观测小计，未知日不赋零；2013 GHCN覆盖不足单列。

GHCN全拟合期≥25 mm事件42次，≥50 mm事件8次，≥100 mm事件1次。最大100.8 mm（2013-08-12），同日D222=66.9 mm、NASA=42.39 mm；2008-08-10 GHCN=80.8、D222=80.7 mm。2006-07-28 GHCN=67.5、D222=115.2 mm，保留空间/观测差异，不裁剪。2004目标窗口最大33.7 mm，无孤立巨大值。

2004 NASA前八月319.57 mm、全年406.55 mm，与GHCN分别278.9/357.3 mm有差异，但主要夏季降雨时段一致。NESDC原字段仅作原样参照：9月两字段为0，与D22244.5、GHCN42.5、NASA47.16 mm冲突；“最高”55/93行大于“合计”，不能擅自调换字段或把它作为月累计真值。NASA未晋升为来源。

## 输出与下一步

已生成目标244天 `accepted_target_window_fallback.csv`，含原flags、source、固定source_reason和距离。全部必需审计CSV/JSON均在本目录，另有 `unresolved_local_dates_after_target.csv` 列出2,447个未闭合日期。

最小恢复条件：为2004-09-01至2013-12-31其余本地空白日提供获准的逐日来源方案或足够的空白语义证据，重点包括此前111个2013 blocker。不得将本次接受范围自动扩大。待3653天RAIN级QC通过后才能恢复010 Phase B。

未运行CLI、WGEN、DSSAT或PPO。未处理FQA，未改动008/009/010证据。010整体Phase A尚未通过，因此未追加“已解除”说明。

## Git 与执行记录

建议提交：011 prompt、run_audit.py、build_deliverables.py、小型审计CSV/JSON、findings.md、本报告与PPT。原始天气和原始Excel不提交。未commit或push。

建议 commit：`hla: audit GHCN 50756 fallback for 2004 rainfall`。

执行中首次解析未兼容Excel文本数字，修正为显式数值转换后重跑；一次临时PowerShell引号错误和一次UTF-8读取错误均已修正，未影响原始数据。原始缓存哈希经比对通过。

官方格式参考：https://www.ncei.noaa.gov/pub/data/ghcn/daily/readme.txt （2026-10-04核对）。
'''
# Populate event counts from evidence rather than hand-entered totals.
ec=pd.read_csv(O/'extreme_counts_by_year.csv');report=report.replace('事件42次',f'事件{int(ec.ge25.sum())}次').replace('事件8次',f'事件{int(ec.ge50.sum())}次')
(O/'findings.md').write_text(report,encoding='utf-8');(R/'docs/hla_ghcn50756_rain_fallback_audit_011.md').write_text(report,encoding='utf-8')
slides=[('011 HLA 2004 降水补源审计',['目标244天：GHCN 50756补源可接受','完整3653天：尚不能恢复010','其余2447个本地空白日来源未闭合']),('官方缓存与站点代表性',['HAILUN：47.433°N，126.967°E，240 m','距HLA 5.3817 km，高差6 m','2004目标窗口244/244天可用，无非空QFLAG','累计278.9 mm，最大33.7 mm']),('0日对齐通过事件门槛',['-1日：53.07%；0日：85.13%；+1日：53.23%','0日935个实际非空可比日，冻结门槛80%','MAE 1.913 mm，RMSE 6.813 mm','全空白日不作为干日真值']),('2004多源月尺度证据',['GHCN前八月278.9 mm，NASA 319.57 mm','9月：D222 44.5，GHCN 42.5，NASA 47.16 mm','NESDC 9月为0，原字段语义仍未闭合','NASA和NESDC仅作交叉参照']),('目标补源通过，完整RAIN仍受阻',['仅接受2004-01-01至2004-08-31的GHCN补源','009仍将其余2447个全空白日标为语义未明确','READY_TO_RESUME_010 = FALSE','待其余日值来源闭合后重做3653天RAIN级QC'])]
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
for i,(title,lines) in enumerate(slides):
 s=p.slides.add_slide(p.slide_layouts[6]);tb=s.shapes.add_textbox(Inches(.65),Inches(.6),Inches(12),Inches(1));r=tb.text_frame.paragraphs[0].add_run();r.text=title;r.font.size=Pt(32);r.font.name='Microsoft YaHei';r.font.bold=True
 tb=s.shapes.add_textbox(Inches(.8),Inches(2),Inches(11.6),Inches(4.4));tf=tb.text_frame;tf.word_wrap=True
 for j,line in enumerate(lines):
  pp=tf.paragraphs[0] if j==0 else tf.add_paragraph();pp.text=line;pp.font.name='Microsoft YaHei';pp.font.size=Pt(24);pp.space_after=Pt(24)
 tb=s.shapes.add_textbox(Inches(.8),Inches(7),Inches(11.5),Inches(.25));tb.text_frame.text=f'2026-10-04  |  正式证据：011 CSV / JSON / MD   {i+1}/5';tb.text_frame.paragraphs[0].font.size=Pt(11)
p.save(R/'docs/hla_ghcn50756_rain_fallback_audit_011.pptx')
print('Report and 5-slide PPT saved. Counts:',ec.ge25.sum(),ec.ge50.sum(),ec.ge100.sum())
