"""Build the lightweight 010 experiment-record deck after the Phase A stop."""

from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
BASE = ROOT / "results" / "hla_random_weather_8seed" / "010"
DOCS = ROOT / "docs"
PLOTS = BASE / "plots"
PLOTS.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(25, 50, 74)
RED = RGBColor(181, 55, 55)
GREEN = RGBColor(44, 122, 94)
GRAY = RGBColor(92, 101, 110)
LIGHT = RGBColor(239, 243, 246)


def set_text(run, size=20, bold=False, color=NAVY):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.2), Inches(0.8))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_text(run, 26, True)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(1.02), Inches(12.0), Inches(0.4))
        p2 = sub.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        set_text(r2, 11, False, GRAY)


def bullets(slide, items, left=0.8, top=1.55, width=11.8, height=5.1, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(12)
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = NAVY


def status_card(slide, x, y, w, h, label, value, color):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = color
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    set_text(r, 13, True, GRAY)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = value
    set_text(r2, 22, True, color)


def build_chart():
    plt.rcParams["font.family"] = "Microsoft YaHei"
    plt.rcParams["axes.unicode_minus"] = False
    path = BASE / "source_audit" / "hla_2004_monthly_crosscheck.csv"
    df = pd.read_csv(path)
    fig, ax = plt.subplots(figsize=(10.8, 4.4))
    months = df["month"]
    ax.plot(months, df["raw_最高"], marker="o", linewidth=2, label="NESDC 原字段‘最高’（语义未闭合）")
    ax.plot(months, df["d222_20_20_mm"], marker="o", linewidth=2, label="D222 月累计")
    ax.plot(months, df["nasa_rain_mm"], marker="o", linewidth=2, label="NASA 月累计")
    ax.plot(months, df["ghcn_rain_mm"], marker="o", linewidth=2, label="GHCN 月累计")
    ax.set_xticks(range(1, 13))
    ax.set_xlabel("2004 月份")
    ax.set_ylabel("降水（mm）")
    ax.set_title("HLA 2004 月尺度交叉核对")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=2, fontsize=8, frameon=False)
    fig.tight_layout()
    out = PLOTS / "hla_2004_monthly_crosscheck.png"
    fig.savefig(out, dpi=170, bbox_inches="tight")
    plt.close(fig)
    return out


def main():
    chart = build_chart()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    title(slide, "010 HLA 随机天气恢复任务", "Phase A 降水 provenance 审计 · 2026-10-04")
    status_card(slide, 0.9, 1.75, 5.5, 2.1, "HLA_RAIN_PROVENANCE", "FAIL", RED)
    status_card(slide, 6.9, 1.75, 5.5, 2.1, "READY_FOR_HLA_WGEN", "FALSE", RED)
    bullets(slide, ["任务按冻结门禁停止；未运行 WGEN、DSSAT 或 PPO。"], top=4.5, size=21)

    slide = prs.slides.add_slide(blank)
    title(slide, "新数据的真实粒度", "NESDC 2000–2008 年黑龙江海伦站农田气象监测数据")
    bullets(
        slide,
        [
            "产品说明明确：时间分辨率为月。",
            "降水表字段仅为 年 / 月 / 合计 / 最高；没有日字段。",
            "降水实体表实际覆盖 2001–2008，共 93 个月记录。",
            "缺测编码、质量控制码和 blank 语义未提供。",
        ],
        size=20,
    )

    slide = prs.slides.add_slide(blank)
    title(slide, "2004 年无法形成逐日来源", "2004 年只有 1–9 月九条月记录；10–12 月缺失")
    slide.shapes.add_picture(str(chart), Inches(0.8), Inches(1.45), width=Inches(11.8))
    note = slide.shapes.add_textbox(Inches(0.85), Inches(6.55), Inches(11.6), Inches(0.45))
    p = note.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "9 月 NESDC 两个原始字段均为 0；D222 / 项目旧表为 44.5 mm，NASA / GHCN 为 47.16 / 42.5 mm。"
    set_text(r, 12, True, RED)

    slide = prs.slides.add_slide(blank)
    title(slide, "Phase A 阻断原因")
    bullets(
        slide,
        [
            "BLOCKED_BY_NESDC_MONTHLY_NOT_DAILY",
            "BLOCKED_BY_2004_DAILY_PRECIPITATION_PROVENANCE",
            "BLOCKED_BY_NESDC_D222_2004_09_CONFLICT",
            "BLOCKED_BY_NESDC_FIELD_SEMANTICS",
        ],
        size=18,
    )

    slide = prs.slides.add_slide(blank)
    title(slide, "停止状态与最小恢复条件")
    status_card(slide, 0.8, 1.55, 2.35, 1.45, "Fitting weather", "NOT RUN", GRAY)
    status_card(slide, 3.45, 1.55, 2.35, 1.45, "CLI / WGEN", "NOT RUN", GRAY)
    status_card(slide, 6.1, 1.55, 2.35, 1.45, "DSSAT smoke", "NOT RUN", GRAY)
    status_card(slide, 8.75, 1.55, 2.35, 1.45, "PPO seeds", "0 / 8", GRAY)
    bullets(
        slide,
        [
            "最小恢复证据：HLA 2004-01-01 至 2004-08-31 的权威逐日降水。",
            "必须含日期、单位、缺测/0/blank 定义、QC，并与 D222 有足够重叠核对日界。",
            "获得证据后从 Phase A 重跑；FQA 继续冻结。",
        ],
        top=3.55,
        size=18,
    )

    out = DOCS / "hla_random_weather_8seed_010_record.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
