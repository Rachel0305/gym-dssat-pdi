"""Read-only comparison of the downloaded HLA automatic-gauge monthly dataset."""
from __future__ import annotations

import csv
import hashlib
import json
from collections import defaultdict
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ZIP = ROOT / "data/external/CERN_sciencedb821.zip"
PDF = ROOT / "data/external/CERN_sciencedb821_doc.pdf"
D222_SELECTION = ROOT / "results/hla_random_weather_8seed/014_isd50756_recovery/daily_source_selection_after_014.csv"


def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def main():
    if (OUT / "final_gate.json").exists():
        raise FileExistsError("040 final gate exists; refusing overwrite")
    with ZipFile(ZIP) as archive:
        members = [m for m in archive.infolist() if not m.is_dir()]
        if len(members) != 1 or not members[0].filename.endswith(".xlsx"):
            raise ValueError("Expected one xlsx member in ScienceDB 821 archive")
        xlsx_bytes = archive.read(members[0])
    workbook = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    if "降水" not in workbook.sheetnames:
        raise ValueError("Precipitation sheet missing")
    sheet = workbook["降水"]
    header = [cell.value for cell in sheet[1]]
    expected = ["年份", "月份", "月合计值（mm）", "月小时降水极大值（mm）", "极大值日期"]
    if header != expected:
        raise ValueError(f"Unexpected precipitation columns: {header}")
    auto = {}
    for year, month, total, hourly_max, max_day in sheet.iter_rows(min_row=2, values_only=True):
        key = int(year), int(month)
        if key in auto:
            raise ValueError(f"Duplicate auto precipitation month: {key}")
        auto[key] = {"total": float(total), "hourly_max": float(hourly_max), "max_day": int(max_day)}
    workbook.close()
    if len(auto) != 120 or any((year, month) not in auto for year in range(2009, 2019) for month in range(1, 13)):
        raise ValueError("Automatic monthly series incomplete")

    d222 = defaultdict(list)
    selected = defaultdict(list)
    with D222_SELECTION.open(newline="", encoding="utf-8-sig") as stream:
        for row in csv.DictReader(stream):
            key = int(row["year"]), int(row["month"])
            if row["local_d222_rain_mm"].strip():
                d222[key].append(float(row["local_d222_rain_mm"]))
            if row["selected_rain_mm"].strip():
                selected[key].append(float(row["selected_rain_mm"]))
    monthly = []
    for year in range(2009, 2014):
        for month in range(1, 13):
            key = year, month
            a = auto[key]
            manual_subtotal = round(sum(d222[key]), 3)
            selected_subtotal = round(sum(selected[key]), 3)
            monthly.append({
                "year": year, "month": month, "auto_published_monthly_total_mm": a["total"],
                "auto_hourly_max_mm": a["hourly_max"], "auto_hourly_max_day_of_month": a["max_day"],
                "d222_numeric_days": len(d222[key]), "d222_numeric_subtotal_mm": manual_subtotal,
                "selected_days": len(selected[key]), "selected_subtotal_mm": selected_subtotal,
                "d222_subtotal_minus_auto_published_monthly_mm": round(manual_subtotal - a["total"], 3),
                "comparison_note": "D222 is incomplete-day subtotal; automatic monthly total may sum available hours on partly missing days; different instruments",
            })
    with (OUT / "monthly_comparison.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(monthly[0]))
        writer.writeheader()
        writer.writerows(monthly)
    annual = []
    for year in range(2009, 2019):
        auto_total = round(sum(auto[year, month]["total"] for month in range(1, 13)), 3)
        manual_total = round(sum(sum(d222[year, month]) for month in range(1, 13)), 3) if year <= 2013 else None
        annual.append({"year": year, "auto_12_month_total_mm": auto_total,
                       "d222_numeric_day_subtotal_mm": manual_total,
                       "d222_numeric_day_count": sum(len(d222[year, month]) for month in range(1, 13)) if year <= 2013 else None})
    with (OUT / "annual_context.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(annual[0]))
        writer.writeheader()
        writer.writerows(annual)
    payload = {
        "task": "040", "site": "HLA", "status": "MONTHLY_REFERENCE_AUDITED_DAILY_EVENT_UNRESOLVED",
        "zip_sha256": sha(ZIP), "paper_pdf_sha256": sha(PDF),
        "contained_xlsx_sha256": hashlib.sha256(xlsx_bytes).hexdigest().upper(),
        "d222_selection_sha256": sha(D222_SELECTION), "auto_months": len(auto),
        "auto_2013_annual_mm": next(row["auto_12_month_total_mm"] for row in annual if row["year"] == 2013),
        "auto_2013_rank_descending_2009_2018": 1 + sum(row["auto_12_month_total_mm"] > 714.8 for row in annual),
        "d222_2013_numeric_day_subtotal_mm": next(row["d222_numeric_day_subtotal_mm"] for row in annual if row["year"] == 2013),
        "auto_2013_july_mm": auto[2013, 7]["total"], "auto_2013_july_hourly_max_mm": auto[2013, 7]["hourly_max"],
        "auto_2013_july_hourly_max_day": auto[2013, 7]["max_day"],
        "daily_2013_07_30_total_available": False,
        "d222_178_2_mm_confirmed_by_automatic_data": False,
        "d222_178_2_mm_refuted_by_automatic_data": False,
        "weather_modified": False, "wgen_executed": False, "dssat_executed": False, "ppo_executed": False,
        "formal_training_ready": False,
    }
    (OUT / "final_gate.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
