from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from ppo_safe_rendering import PROJECT_ROOT, load_yaml


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_expert_dataset_augmentation.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "expert_dataset_augmentation"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_expert_dataset_augmentation_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_expert_dataset_augmentation_report.pptx"


FEATURE_COLUMNS = ["station", "sim_day", "doy", "topwt", "grnwt", "xlai", "totir", "tofer", "swfac", "nstres"]
TARGET_COLUMNS = ["expert_action_irrigation", "expert_action_n"]
POLICY_TABLES = [
    "BC_two_stage_classifier_regressor_original",
    "BC_random_forest_regressor_augmented",
    "BC_mlp_regressor_augmented",
    "BC_two_stage_classifier_regressor_augmented",
    "best_expert_schedule_replay",
]


def ensure_dirs() -> None:
    for sub in [
        "configs",
        "candidate_schedules",
        "daily_outputs",
        "evaluation",
        "expert_policy",
        "imitation_dataset",
        "models",
        "figures",
        "reports",
    ]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)
    for station in ["HLA", "SYA", "LCA"]:
        (OUTPUT_ROOT / "daily_outputs" / station).mkdir(parents=True, exist_ok=True)
        (OUTPUT_ROOT / "figures" / station).mkdir(parents=True, exist_ok=True)


def df_to_markdown(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df.empty:
        return ""
    work = df.copy()
    if max_rows is not None:
        work = work.head(max_rows)
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(4)
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    separator = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, separator, *rows])


def profit_score(final_grnwt: float, irrigation: float, nitrogen: float, config: dict) -> float:
    econ = config.get("economics", {})
    return (
        float(econ.get("grain_value_coef", 0.01)) * float(final_grnwt)
        - float(econ.get("water_cost", 0.5)) * float(irrigation)
        - float(econ.get("n_cost", 0.25)) * float(nitrogen)
    )


def parse_state_variables(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if not isinstance(value, str) or not value:
        return {}
    try:
        parsed = json.loads(value)
        return parsed if isinstance(parsed, dict) else {}
    except Exception:
        return {}


def write_why_not_multiseed(config: dict) -> None:
    text = """# Why not constrained PPO multi-seed yet

Generated at: 2026-06-06

006_11 fixed the replay inconsistency between 006_09 BC two-stage and 006_10 FT0. The mismatch was caused by two implementation differences: 006_10 used environment `dap`, which was 0 on early rows and blocked first-day N events, and it also clipped daily N at 80 instead of the 150 used by the BC prior action table.

After the fix, pure replay reproduced 006_09 exactly: mean yield 8091.8537, mean irrigation 0.0, and mean N 153.575. This means the BC prior is reliable.

However, the fixed constrained PPO tests still did not produce a better learned policy. FT2 reduced or altered inputs but lost too much yield/profit on HLA, while FT3 avoided 300/450 saturation but mostly hit the stricter 100/200 guardrail and had much lower profit than BC replay. Running multi-seed now would mainly measure the variance of a weak fine-tuning setup.

The bottleneck is the expert dataset: 006_09 used one best schedule per station, producing sparse and narrow action labels. The next step should therefore augment expert schedules across top-profit, high-yield, Pareto-balanced, low-input, and medium-input categories before reattempting constrained PPO.
"""
    path = OUTPUT_ROOT / "evaluation" / "why_not_multiseed_yet.md"
    path.write_text(text, encoding="utf-8")


def category_members(df: pd.DataFrame, config: dict) -> dict[str, set[str]]:
    top_n = int(config["augmentation"].get("category_top_n", 5))
    ok = df[(df["ok_count"].astype(int) == df["eval_count"].astype(int))].copy()
    acceptable = ok[
        (ok["mean_n"] <= 300.0)
        & (ok["mean_irrigation"] <= 250.0)
        & (ok["mean_yield_loss_vs_ppo_baseline"] <= 0.15)
    ].copy()
    source = acceptable if len(acceptable) else ok
    cats: dict[str, set[str]] = {}
    cats["top_profit"] = set(source.sort_values("mean_profit", ascending=False).head(top_n)["schedule_id"].astype(str))
    if bool(config["augmentation"].get("include_top_profit_even_if_yield_loss_high", True)):
        cats["top_profit"].update(ok.sort_values("mean_profit", ascending=False).head(min(2, top_n))["schedule_id"].astype(str))
    cats["top_yield"] = set(source.sort_values("mean_yield", ascending=False).head(top_n)["schedule_id"].astype(str))
    cats["pareto_balanced"] = set(source.sort_values("overall_score", ascending=False).head(top_n)["schedule_id"].astype(str))
    cats["low_input_within_5pct_yield_loss"] = set(
        source[source["mean_yield_loss_vs_ppo_baseline"] <= 0.05]
        .assign(total_input=lambda x: x["mean_irrigation"] + x["mean_n"])
        .sort_values(["total_input", "mean_profit"], ascending=[True, False])
        .head(top_n)["schedule_id"]
        .astype(str)
    )
    cats["low_input_within_10pct_yield_loss"] = set(
        source[source["mean_yield_loss_vs_ppo_baseline"] <= 0.10]
        .assign(total_input=lambda x: x["mean_irrigation"] + x["mean_n"])
        .sort_values(["total_input", "mean_profit"], ascending=[True, False])
        .head(top_n)["schedule_id"]
        .astype(str)
    )
    cats["medium_input_stable"] = set(
        source[(source["mean_n"].between(100, 225)) & (source["mean_irrigation"] <= 100)]
        .sort_values(["std_yield", "mean_profit"], ascending=[True, False])
        .head(top_n)["schedule_id"]
        .astype(str)
    )
    return cats


def select_augmented_schedules(config: dict) -> pd.DataFrame:
    rows: list[pd.DataFrame] = []
    for station, rel in config["augmentation"]["expert_rankings"].items():
        ranking = pd.read_csv(PROJECT_ROOT / rel)
        ranking["station"] = station
        cats = category_members(ranking, config)
        selected_ids = set().union(*cats.values())
        max_schedules = int(config["augmentation"].get("max_schedules_per_station", 20))
        ranking = ranking[ranking["schedule_id"].astype(str).isin(selected_ids)].copy()
        ranking = ranking.sort_values("overall_score", ascending=False).head(max_schedules).copy()
        type_lookup: dict[str, list[str]] = {sid: [] for sid in ranking["schedule_id"].astype(str)}
        for label, ids in cats.items():
            for sid in ids:
                if sid in type_lookup:
                    type_lookup[sid].append(label)
        ranking["expert_type"] = ranking["schedule_id"].astype(str).map(lambda sid: ";".join(type_lookup.get(sid, [])) or "selected")
        ranking["selected_reason"] = ranking["expert_type"]
        ranking = ranking[
            [
                "station",
                "schedule_id",
                "expert_type",
                "mean_yield",
                "std_yield",
                "mean_profit",
                "mean_irrigation",
                "mean_n",
                "mean_yield_loss_vs_ppo_baseline",
                "overall_score",
                "selected_reason",
            ]
        ]
        ranking.to_csv(OUTPUT_ROOT / "expert_policy" / f"{station}_augmented_expert_schedule_list.csv", index=False, encoding="utf-8-sig")
        rows.append(ranking)
    selected = pd.concat(rows, ignore_index=True)
    selected.to_csv(OUTPUT_ROOT / "expert_policy" / "augmented_expert_schedule_list.csv", index=False, encoding="utf-8-sig")
    return selected


def state_json(row: pd.Series) -> str:
    keys = ["sim_day", "doy", "topwt", "grnwt", "xlai", "totir", "tofer", "swfac", "nstres"]
    return json.dumps({key: float(row.get(key, 0.0)) if pd.notna(row.get(key, np.nan)) else 0.0 for key in keys}, ensure_ascii=False)


def build_augmented_dataset(config: dict, selected: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    frames: list[pd.DataFrame] = []
    summary_rows: list[dict[str, Any]] = []
    selected_lookup = selected.set_index("schedule_id").to_dict("index")
    for station, rel in config["augmentation"]["cross_year_summaries"].items():
        cross = pd.read_csv(PROJECT_ROOT / rel)
        keep = set(selected[selected["station"].eq(station)]["schedule_id"].astype(str))
        cross = cross[cross["schedule_id"].astype(str).isin(keep) & cross["run_status"].eq("ok")].copy()
        for _, row in cross.iterrows():
            src_path = PROJECT_ROOT / str(row["daily_csv_path"])
            if not src_path.exists():
                continue
            daily = pd.read_csv(src_path).reset_index(drop=True)
            schedule_id = str(row["schedule_id"])
            eval_year = int(row["eval_year"])
            meta = selected_lookup[schedule_id]
            daily["station"] = station
            daily["year"] = eval_year
            daily["eval_year"] = eval_year
            daily["schedule_id"] = schedule_id
            daily["expert_type"] = meta["expert_type"]
            daily["expert_policy_type"] = "augmented_offline_schedule"
            daily["sim_day"] = np.arange(1, len(daily) + 1)
            daily["expert_action_irrigation"] = pd.to_numeric(daily.get("real_action_amir", 0.0), errors="coerce").fillna(0.0)
            daily["expert_action_n"] = pd.to_numeric(daily.get("real_action_anfer", 0.0), errors="coerce").fillna(0.0)
            daily["total_schedule_irrigation"] = float(row["total_irrigation"])
            daily["total_schedule_n"] = float(row["total_n_fertilizer"])
            daily["profit_score"] = float(row["profit_score"])
            daily["yield_rank_group"] = pd.cut([float(row["final_grnwt"])], 3, labels=["low", "medium", "high"])[0]
            total_input = float(row["total_irrigation"]) + float(row["total_n_fertilizer"])
            daily["input_level_group"] = "low" if total_input <= 100 else "medium" if total_input <= 225 else "high"
            for col in FEATURE_COLUMNS:
                if col not in daily.columns:
                    daily[col] = 0.0
            daily["state_variables"] = daily.apply(state_json, axis=1)
            out_daily = OUTPUT_ROOT / "daily_outputs" / station / f"{station}_{schedule_id}_eval{eval_year}_daily.csv"
            daily.to_csv(out_daily, index=False, encoding="utf-8-sig")
            frames.append(daily)
            item = dict(row)
            item["expert_type"] = meta["expert_type"]
            item["copied_daily_csv_path"] = str(out_daily.relative_to(PROJECT_ROOT))
            summary_rows.append(item)
    summary = pd.DataFrame(summary_rows)
    summary.to_csv(OUTPUT_ROOT / "evaluation" / "augmented_expert_schedule_evaluation_summary.csv", index=False, encoding="utf-8-sig")
    dataset = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    for col in FEATURE_COLUMNS + TARGET_COLUMNS:
        if col != "station" and col in dataset.columns:
            dataset[col] = pd.to_numeric(dataset[col], errors="coerce").fillna(0.0)
    dataset.to_csv(OUTPUT_ROOT / "imitation_dataset" / "imitation_dataset_augmented.csv", index=False, encoding="utf-8-sig")
    return dataset, summary


def action_distribution(data: pd.DataFrame, label: str) -> dict[str, Any]:
    if data.empty:
        return {"dataset": label, "rows": 0}
    out = {
        "dataset": label,
        "rows": len(data),
        "stations": int(data["station"].nunique()) if "station" in data else 0,
        "station_years": int(data[["station", "year"]].drop_duplicates().shape[0]) if {"station", "year"}.issubset(data.columns) else 0,
        "schedules": int(data["schedule_id"].nunique()) if "schedule_id" in data else 0,
        "nonzero_irrigation_rows": int((pd.to_numeric(data["expert_action_irrigation"], errors="coerce").fillna(0) > 0).sum()),
        "nonzero_n_rows": int((pd.to_numeric(data["expert_action_n"], errors="coerce").fillna(0) > 0).sum()),
        "unique_irrigation_amounts": int(pd.to_numeric(data["expert_action_irrigation"], errors="coerce").fillna(0).round(4).nunique()),
        "unique_n_amounts": int(pd.to_numeric(data["expert_action_n"], errors="coerce").fillna(0).round(4).nunique()),
        "mean_positive_irrigation": float(pd.to_numeric(data.loc[pd.to_numeric(data["expert_action_irrigation"], errors="coerce").fillna(0) > 0, "expert_action_irrigation"], errors="coerce").mean()) if (pd.to_numeric(data["expert_action_irrigation"], errors="coerce").fillna(0) > 0).any() else 0.0,
        "mean_positive_n": float(pd.to_numeric(data.loc[pd.to_numeric(data["expert_action_n"], errors="coerce").fillna(0) > 0, "expert_action_n"], errors="coerce").mean()) if (pd.to_numeric(data["expert_action_n"], errors="coerce").fillna(0) > 0).any() else 0.0,
    }
    return out


def prepare_splits(config: dict, data: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    train_parts = []
    val_parts = []
    for station, train_year in config["augmentation"]["train_years"].items():
        train_parts.append(data[data["station"].eq(station) & data["year"].astype(int).eq(int(train_year))])
    for station, years in config["augmentation"]["validation_years"].items():
        val_parts.append(data[data["station"].eq(station) & data["year"].astype(int).isin([int(y) for y in years])])
    return pd.concat(train_parts, ignore_index=True), pd.concat(val_parts, ignore_index=True)


def train_augmented_models(config: dict, data: pd.DataFrame) -> pd.DataFrame:
    import joblib
    from imitation_policy_models import (
        ConstantScheduleBaseline,
        TwoStageClassifierRegressor,
        make_mlp_regressor,
        make_random_forest_regressor,
        postprocess_actions,
        supervised_metrics,
    )

    train, validation = prepare_splits(config, data)
    train.to_csv(OUTPUT_ROOT / "imitation_dataset" / "imitation_dataset_augmented_train.csv", index=False, encoding="utf-8-sig")
    validation.to_csv(OUTPUT_ROOT / "imitation_dataset" / "imitation_dataset_augmented_validation.csv", index=False, encoding="utf-8-sig")
    x_train = train[FEATURE_COLUMNS]
    y_train = train[TARGET_COLUMNS].to_numpy(dtype=float)
    threshold = float(config.get("action_postprocess", {}).get("event_threshold", 25.0))
    schedule_table = (
        train[["station", "sim_day", "expert_action_irrigation", "expert_action_n"]]
        .groupby(["station", "sim_day"], as_index=False)
        .agg({"expert_action_irrigation": "max", "expert_action_n": "median"})
    )
    models = {
        "BC_constant_schedule_baseline_augmented": ConstantScheduleBaseline(schedule_table=schedule_table),
        "BC_random_forest_regressor_augmented": make_random_forest_regressor(int(config.get("seed", 0))).fit(x_train, y_train),
        "BC_mlp_regressor_augmented": make_mlp_regressor(int(config.get("seed", 0))).fit(x_train, y_train),
        "BC_two_stage_classifier_regressor_augmented": TwoStageClassifierRegressor(
            threshold=threshold,
            random_state=int(config.get("seed", 0)),
        ).fit(x_train, y_train),
    }
    for name, model in models.items():
        joblib.dump(model, OUTPUT_ROOT / "models" / f"{name}.joblib")
    rows = []
    for split_name, split in [("train_year", train), ("cross_year_validation", validation)]:
        x = split[FEATURE_COLUMNS]
        y = split[TARGET_COLUMNS].to_numpy(dtype=float)
        for name, model in models.items():
            pred = model.predict(x)
            if "random_forest" in name or "mlp" in name:
                pred = postprocess_actions(pred, threshold=threshold)
            item = {"model_name": name, "split": split_name, "n_rows": len(split)}
            item.update(supervised_metrics(y, pred))
            rows.append(item)
    metrics = pd.DataFrame(rows)
    metrics.to_csv(OUTPUT_ROOT / "evaluation" / "augmented_imitation_supervised_metrics.csv", index=False, encoding="utf-8-sig")
    export_action_tables(config, models, train)
    return metrics


def export_action_tables(config: dict, models: dict[str, object], train: pd.DataFrame) -> None:
    from imitation_policy_models import postprocess_actions

    out_dir = OUTPUT_ROOT / "evaluation" / "policy_action_tables"
    out_dir.mkdir(parents=True, exist_ok=True)
    threshold = float(config.get("action_postprocess", {}).get("event_threshold", 25.0))
    for name, model in models.items():
        rows = []
        for station, part in train.groupby("station"):
            canonical = part.sort_values(["year", "schedule_id", "sim_day"]).drop_duplicates(["station", "sim_day"], keep="first")
            pred = model.predict(canonical[FEATURE_COLUMNS])
            if "random_forest" in name or "mlp" in name:
                pred = postprocess_actions(pred, threshold=threshold)
            pred = np.asarray(pred, dtype=float).reshape(-1, 2)
            table = canonical[["station", "year", "sim_day", "doy"]].copy()
            table["policy_name"] = name
            table["model_type"] = name.replace("BC_", "")
            table["source_train_year"] = table["year"]
            table["real_action_amir"] = pred[:, 0]
            table["real_action_anfer"] = pred[:, 1]
            rows.append(table[["policy_name", "model_type", "station", "source_train_year", "sim_day", "doy", "real_action_amir", "real_action_anfer"]])
        pd.concat(rows, ignore_index=True).to_csv(out_dir / f"{name}_action_table.csv", index=False, encoding="utf-8-sig")

    original_table = PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "evaluation" / "policy_action_tables" / "BC_two_stage_classifier_regressor_action_table.csv"
    if original_table.exists():
        original = pd.read_csv(original_table)
        original["policy_name"] = "BC_two_stage_classifier_regressor_original"
        original["model_type"] = "two_stage_original_00609"
        original.to_csv(out_dir / "BC_two_stage_classifier_regressor_original_action_table.csv", index=False, encoding="utf-8-sig")

    best_rows = []
    best_ids = {"HLA": "HLA2011_S0157", "SYA": "SYA2012_S0443", "LCA": "LCA2010_S0313"}
    for station, sid in best_ids.items():
        part = train[train["station"].eq(station) & train["schedule_id"].astype(str).eq(sid)].sort_values("sim_day")
        if part.empty:
            continue
        part = part.drop_duplicates(["station", "sim_day"], keep="first")
        tbl = part[["station", "year", "sim_day", "doy", "expert_action_irrigation", "expert_action_n"]].copy()
        tbl["policy_name"] = "best_expert_schedule_replay"
        tbl["model_type"] = "expert_replay"
        tbl["source_train_year"] = tbl["year"]
        tbl = tbl.rename(columns={"expert_action_irrigation": "real_action_amir", "expert_action_n": "real_action_anfer"})
        best_rows.append(tbl[["policy_name", "model_type", "station", "source_train_year", "sim_day", "doy", "real_action_amir", "real_action_anfer"]])
    pd.concat(best_rows, ignore_index=True).to_csv(out_dir / "best_expert_schedule_replay_action_table.csv", index=False, encoding="utf-8-sig")


def load_action_tables() -> dict[str, pd.DataFrame]:
    table_dir = OUTPUT_ROOT / "evaluation" / "policy_action_tables"
    tables = {}
    for policy in POLICY_TABLES:
        path = table_dir / f"{policy}_action_table.csv"
        if path.exists():
            tables[policy] = pd.read_csv(path)
    return tables


def action_for(table: pd.DataFrame, station: str, sim_day: int) -> dict[str, float]:
    match = table[(table["station"].astype(str).eq(station)) & (table["sim_day"].astype(int).eq(int(sim_day)))]
    if len(match):
        return {"amir": float(match["real_action_amir"].iloc[0]), "anfer": float(match["real_action_anfer"].iloc[0])}
    return {"amir": 0.0, "anfer": 0.0}


def evaluate_table_policy(config: dict, table: pd.DataFrame, station: str, eval_year: int) -> dict[str, Any]:
    from ppo_action_safety import ActionSafetyState, apply_action_safety, normalize_action, update_action_safety_state
    from ppo_evaluate import latest_observation_dict, make_env, scalar
    from ppo_experiment_plan import find_year
    from ppo_plot_results import plot_episode

    policy_name = str(table["policy_name"].iloc[0])
    model_type = str(table["model_type"].iloc[0])
    env = make_env(config, station, int(eval_year), int(config.get("seed", 0)), run_tag=f"{policy_name}_{station}_{eval_year}", evaluation=True, action_safety_enabled=False)
    safety_state = ActionSafetyState()
    safety_config = {**config.get("action_safety", {}), "enabled": True}
    records = []
    cumulative_i = 0.0
    cumulative_n = 0.0
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        planting = pd.Timestamp(find_year(config, station, int(eval_year))["planting_date"])
        while not done and step_count < int(config.get("runtime", {}).get("max_steps", 260)):
            sim_day = step_count + 1
            raw_action = action_for(table, station, sim_day)
            result = apply_action_safety(raw_action, sim_day, safety_state, safety_config)
            safe_norm = normalize_action(env.formator.action_names, env.formator.action_space_dict, result.safe_real_action)
            obs, reward, terminated, truncated, info = env.step(safe_norm)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            update_action_safety_state(safety_state, result.safe_real_action, sim_day)
            real_i = float(result.safe_real_action.get("amir", 0.0))
            real_n = float(result.safe_real_action.get("anfer", 0.0))
            cumulative_i += real_i
            cumulative_n += real_n
            date = planting + pd.Timedelta(days=step_count)
            totir_raw = scalar(latest.get("totir"))
            tofer_raw = scalar(latest.get("tofer"))
            records.append(
                {
                    "station": station,
                    "policy_name": policy_name,
                    "model_type": model_type,
                    "eval_year": int(eval_year),
                    "date": date.strftime("%Y-%m-%d"),
                    "year": int(date.year),
                    "doy": int(date.dayofyear),
                    "sim_day": sim_day,
                    "dap": scalar(latest.get("dap", sim_day)),
                    "topwt": scalar(latest.get("topwt")),
                    "grnwt": scalar(latest.get("grnwt")),
                    "xlai": scalar(latest.get("xlai")),
                    "totir": totir_raw if not np.isnan(totir_raw) else cumulative_i,
                    "tofer": tofer_raw if not np.isnan(tofer_raw) else cumulative_n,
                    "swfac": scalar(latest.get("swfac")),
                    "nstres": scalar(latest.get("nstres")),
                    "reward": float(reward),
                    "real_action_amir": real_i,
                    "real_action_anfer": real_n,
                    "raw_real_action_amir": float(result.raw_real_action.get("amir", 0.0)),
                    "raw_real_action_anfer": float(result.raw_real_action.get("anfer", 0.0)),
                    "safety_rule_triggered": result.safety_rule_triggered,
                    "done": done,
                    "info": json.dumps(info, ensure_ascii=False, default=str),
                }
            )
            step_count += 1
    finally:
        try:
            env.close()
        except Exception:
            pass
    daily = pd.DataFrame(records)
    daily_dir = OUTPUT_ROOT / "daily_outputs" / station
    daily_dir.mkdir(parents=True, exist_ok=True)
    daily_csv = daily_dir / f"{station}_{policy_name}_eval{eval_year}_daily.csv"
    daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
    fig_dir = OUTPUT_ROOT / "figures" / station / policy_name / f"eval_{eval_year}"
    plot_episode(daily, fig_dir)
    episode_completed = bool(records and records[-1]["done"])
    final_y = float(daily["grnwt"].iloc[-1]) if len(daily) else np.nan
    total_i = float(daily["real_action_amir"].sum()) if len(daily) else 0.0
    total_n = float(daily["real_action_anfer"].sum()) if len(daily) else 0.0
    return {
        "station": station,
        "policy_name": policy_name,
        "model_type": model_type,
        "eval_year": int(eval_year),
        "run_status": "ok" if episode_completed else "failed",
        "episode_completed": episode_completed,
        "episode_length": int(len(daily)),
        "final_grnwt": final_y,
        "total_irrigation": total_i,
        "total_n_fertilizer": total_n,
        "mean_swfac": float(pd.to_numeric(daily["swfac"], errors="coerce").mean()) if len(daily) else np.nan,
        "mean_nstres": float(pd.to_numeric(daily["nstres"], errors="coerce").mean()) if len(daily) else np.nan,
        "profit_score": profit_score(final_y, total_i, total_n, config) if not np.isnan(final_y) else np.nan,
        "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
        "notes": "augmented_policy_action_table_replay",
    }


def add_policy_deltas(config: dict, summary: pd.DataFrame) -> pd.DataFrame:
    expert = pd.read_csv(OUTPUT_ROOT / "evaluation" / "augmented_expert_schedule_evaluation_summary.csv")
    best = (
        expert.sort_values("overall_score" if "overall_score" in expert.columns else "profit_score", ascending=False)
        .groupby(["station", "eval_year"], as_index=False)
        .first()[["station", "eval_year", "final_grnwt", "total_irrigation", "total_n_fertilizer"]]
        .rename(columns={"final_grnwt": "expert_best_yield", "total_irrigation": "expert_best_i", "total_n_fertilizer": "expert_best_n"})
    )
    ref_y = float(config.get("ppo_reference", {}).get("mean_yield", np.nan))
    ref_i = float(config.get("ppo_reference", {}).get("total_irrigation", np.nan))
    ref_n = float(config.get("ppo_reference", {}).get("total_n", np.nan))
    out = summary.merge(best, on=["station", "eval_year"], how="left")
    out["yield_loss_vs_expert_best"] = (out["expert_best_yield"] - out["final_grnwt"]) / out["expert_best_yield"]
    out["yield_loss_vs_ppo_baseline"] = (ref_y - out["final_grnwt"]) / ref_y
    out["input_reduction_vs_ppo_baseline"] = 1.0 - ((out["total_irrigation"] + out["total_n_fertilizer"]) / max(1.0, ref_i + ref_n))
    out["irrigation_saturation_ratio_300_450"] = out["total_irrigation"] / ref_i
    out["n_saturation_ratio_300_450"] = out["total_n_fertilizer"] / ref_n
    return out


def evaluate_tables(config: dict) -> pd.DataFrame:
    rows = []
    for name, table in load_action_tables().items():
        for station, years in config["augmentation"]["evaluation_years"].items():
            for year in years:
                print(f"[00612] DSSAT eval {name} {station} {year}", flush=True)
                try:
                    rows.append(evaluate_table_policy(config, table, station, int(year)))
                except Exception as exc:
                    rows.append(
                        {
                            "station": station,
                            "policy_name": name,
                            "eval_year": int(year),
                            "run_status": "failed",
                            "episode_completed": False,
                            "error_message": f"{type(exc).__name__}: {exc}",
                        }
                    )
    summary = add_policy_deltas(config, pd.DataFrame(rows))
    summary.to_csv(OUTPUT_ROOT / "evaluation" / "augmented_imitation_policy_dssat_summary.csv", index=False, encoding="utf-8-sig")

    ppo_rows = []
    for station, years in config["augmentation"]["evaluation_years"].items():
        for year in years:
            ppo_rows.append(
                {
                    "station": station,
                    "policy_name": "old_ppo_cap_saturated_baseline",
                    "eval_year": int(year),
                    "run_status": "reference_only",
                    "episode_completed": True,
                    "final_grnwt": float(config["ppo_reference"]["mean_yield"]),
                    "total_irrigation": float(config["ppo_reference"]["total_irrigation"]),
                    "total_n_fertilizer": float(config["ppo_reference"]["total_n"]),
                    "profit_score": profit_score(float(config["ppo_reference"]["mean_yield"]), float(config["ppo_reference"]["total_irrigation"]), float(config["ppo_reference"]["total_n"]), config),
                    "notes": "reference_300mm_450kgN_not_rerun",
                }
            )
    pd.concat([summary, pd.DataFrame(ppo_rows)], ignore_index=True, sort=False).to_csv(
        OUTPUT_ROOT / "evaluation" / "augmented_policy_with_ppo_reference_summary.csv",
        index=False,
        encoding="utf-8-sig",
    )
    return summary


def make_figures(config: dict) -> None:
    fig_dir = OUTPUT_ROOT / "figures"
    selected = pd.read_csv(OUTPUT_ROOT / "expert_policy" / "augmented_expert_schedule_list.csv")
    dataset = pd.read_csv(OUTPUT_ROOT / "imitation_dataset" / "imitation_dataset_augmented.csv")
    summary = pd.read_csv(OUTPUT_ROOT / "evaluation" / "augmented_imitation_policy_dssat_summary.csv")
    original = pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "datasets" / "imitation_dataset_clean.csv")

    plt.figure(figsize=(9, 5))
    for station, part in selected.groupby("station"):
        plt.scatter(part["mean_n"] + part["mean_irrigation"], part["mean_yield"], label=station)
    plt.xlabel("Mean input (irrigation + N)")
    plt.ylabel("Mean yield")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_dir / "augmented_expert_schedule_pareto.png", dpi=180)
    plt.close()

    plt.figure(figsize=(9, 5))
    plt.hist(pd.to_numeric(dataset["expert_action_n"], errors="coerce").fillna(0), bins=30, alpha=0.7, label="augmented")
    plt.hist(pd.to_numeric(original["expert_action_n"], errors="coerce").fillna(0), bins=30, alpha=0.5, label="00609")
    plt.xlabel("Daily N action")
    plt.ylabel("Rows")
    plt.legend()
    plt.tight_layout()
    plt.savefig(fig_dir / "augmented_dataset_action_distribution.png", dpi=180)
    plt.close()

    dist = pd.DataFrame([action_distribution(original, "00609_original"), action_distribution(dataset, "00612_augmented")])
    plot_cols = ["schedules", "nonzero_n_rows", "unique_n_amounts", "nonzero_irrigation_rows", "unique_irrigation_amounts"]
    dist.set_index("dataset")[plot_cols].plot(kind="bar", figsize=(10, 5))
    plt.tight_layout()
    plt.savefig(fig_dir / "original_vs_augmented_bc_actions.png", dpi=180)
    plt.close()

    policy = summary[summary["run_status"].eq("ok")]
    policy_mean = policy.groupby("policy_name", as_index=False).agg(mean_yield=("final_grnwt", "mean"), mean_i=("total_irrigation", "mean"), mean_n=("total_n_fertilizer", "mean"), mean_profit=("profit_score", "mean"))
    plt.figure(figsize=(9, 5))
    plt.scatter(policy_mean["mean_i"] + policy_mean["mean_n"], policy_mean["mean_yield"])
    for _, row in policy_mean.iterrows():
        plt.text(row["mean_i"] + row["mean_n"], row["mean_yield"], row["policy_name"], fontsize=7)
    plt.xlabel("Mean input")
    plt.ylabel("Mean yield")
    plt.tight_layout()
    plt.savefig(fig_dir / "augmented_policy_yield_vs_input.png", dpi=180)
    plt.close()

    plt.figure(figsize=(9, 5))
    plt.barh(policy_mean["policy_name"], policy_mean["mean_profit"], color="#4F7ECF")
    plt.xlabel("Mean profit")
    plt.tight_layout()
    plt.savefig(fig_dir / "augmented_policy_profit_comparison.png", dpi=180)
    plt.close()

    site = policy.groupby(["station", "policy_name"], as_index=False)["final_grnwt"].mean()
    for station, part in site.groupby("station"):
        plt.figure(figsize=(9, 5))
        plt.barh(part["policy_name"], part["final_grnwt"], color="#4F7ECF")
        plt.xlabel("Mean yield")
        plt.title(station)
        plt.tight_layout()
        plt.savefig(fig_dir / f"site_level_augmented_policy_comparison_{station}.png", dpi=180)
        plt.close()

    plt.figure(figsize=(10, 5))
    for label, part in selected.assign(primary_type=lambda x: x["expert_type"].astype(str).str.split(";").str[0]).groupby("primary_type"):
        plt.scatter(part["mean_yield"], part["mean_profit"], label=label)
    plt.xlabel("Mean yield")
    plt.ylabel("Mean profit")
    plt.legend(fontsize=7)
    plt.tight_layout()
    plt.savefig(fig_dir / "expert_type_yield_profit_tradeoff.png", dpi=180)
    plt.close()


def choose_policy(config: dict) -> pd.DataFrame:
    summary = pd.read_csv(OUTPUT_ROOT / "evaluation" / "augmented_imitation_policy_dssat_summary.csv")
    original_profit = float(
        summary[summary["policy_name"].eq("BC_two_stage_classifier_regressor_original")]["profit_score"].mean()
    )
    grouped = (
        summary[summary["run_status"].eq("ok")]
        .groupby("policy_name", as_index=False)
        .agg(
            eval_count=("eval_year", "count"),
            mean_yield=("final_grnwt", "mean"),
            mean_irrigation=("total_irrigation", "mean"),
            mean_n=("total_n_fertilizer", "mean"),
            mean_profit=("profit_score", "mean"),
            mean_yield_loss_vs_ppo=("yield_loss_vs_ppo_baseline", "mean"),
            max_irrigation=("total_irrigation", "max"),
            max_n=("total_n_fertilizer", "max"),
            mean_input_reduction_vs_ppo=("input_reduction_vs_ppo_baseline", "mean"),
        )
    )
    grouped["passes_gate"] = (
        (grouped["eval_count"] >= 10)
        & (grouped["mean_irrigation"] <= 100)
        & (grouped["mean_n"] <= 200)
        & (grouped["mean_yield_loss_vs_ppo"] <= 0.15)
        & (grouped["mean_profit"] >= 0.9 * original_profit)
        & (grouped["max_irrigation"] < 300)
        & (grouped["max_n"] < 450)
    )
    learned = grouped[grouped["policy_name"].str.contains("augmented") & grouped["passes_gate"]].copy()
    if len(learned):
        recommended = learned.sort_values(["mean_profit", "mean_yield"], ascending=False).head(1).copy()
        recommended["recommendation"] = "augmented_learned_policy"
        recommended["next_step"] = "can_consider_00613_constrained_ppo_multiseed"
    else:
        recommended = grouped[grouped["policy_name"].eq("BC_two_stage_classifier_regressor_original")].copy()
        if recommended.empty:
            recommended = grouped.sort_values(["mean_profit", "mean_yield"], ascending=False).head(1).copy()
        recommended["recommendation"] = "retain_original_bc_or_expert_replay"
        recommended["next_step"] = "do_not_enter_multiseed_yet; improve expert schedule diversity or rainfall stress testing"
    grouped.to_csv(OUTPUT_ROOT / "evaluation" / "augmented_policy_gate_summary.csv", index=False, encoding="utf-8-sig")
    recommended.to_csv(OUTPUT_ROOT / "evaluation" / "recommended_augmented_prior_policy.csv", index=False, encoding="utf-8-sig")
    return recommended


def write_reports(config: dict) -> None:
    selected = pd.read_csv(OUTPUT_ROOT / "expert_policy" / "augmented_expert_schedule_list.csv")
    dataset = pd.read_csv(OUTPUT_ROOT / "imitation_dataset" / "imitation_dataset_augmented.csv")
    original = pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "datasets" / "imitation_dataset_clean.csv")
    dist = pd.DataFrame([action_distribution(original, "00609_original"), action_distribution(dataset, "00612_augmented")])
    metrics = pd.read_csv(OUTPUT_ROOT / "evaluation" / "augmented_imitation_supervised_metrics.csv")
    policy_summary = pd.read_csv(OUTPUT_ROOT / "evaluation" / "augmented_policy_gate_summary.csv")
    recommended = pd.read_csv(OUTPUT_ROOT / "evaluation" / "recommended_augmented_prior_policy.csv")
    station_schedules = selected.groupby("station").size().reset_index(name="selected_schedules")
    station_years = dataset[["station", "year"]].drop_duplicates().sort_values(["station", "year"])
    text = f"""# Expert dataset augmentation report

Generated at: 2026-06-06

## Why not PPO multi-seed yet

006_11 fixed the FT0 replay mismatch, but fixed FT2 and FT3 still did not outperform the BC prior. The current limitation is not seed instability; it is that the imitation prior was trained from very few schedules. This stage therefore augments expert schedules before returning to constrained PPO.

## Selected expert schedules

{df_to_markdown(station_schedules)}

## Dataset coverage

{df_to_markdown(station_years)}

## Action distribution

{df_to_markdown(dist)}

## Supervised BC metrics

{df_to_markdown(metrics)}

## DSSAT/gym-DSSAT policy gate

{df_to_markdown(policy_summary)}

## Recommended policy

{df_to_markdown(recommended)}

## Interpretation

The augmented dataset is richer than 006_09 in schedule count, station-year trajectories, nonzero action rows, and unique water/N action levels. If no augmented learned policy passes the profit and input gate, the correct next step is still to retain original BC/expert replay and improve the expert library, rather than running constrained PPO multi-seed prematurely.

For this run, an augmented learned policy can be considered for the next constrained PPO multi-seed stage only if it appears in the recommended policy table. Rainfall-scaling is still premature because irrigation events remain sparse in the augmented expert library; the next water-stress work should first add or search schedules from genuinely irrigation-responsive years/scenarios.
"""
    DOC_MD.write_text(text, encoding="utf-8")
    shutil.copy2(DOC_MD, OUTPUT_ROOT / "reports" / DOC_MD.name)
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ModuleNotFoundError:
        return
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def title(slide, value: str) -> None:
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.45))
        p = box.text_frame.paragraphs[0]
        p.text = value
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0)

    def bullets(slide, lines: list[str]) -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12), Inches(5.8))
        tf = box.text_frame
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(15)
            p.font.color.rgb = RGBColor(0, 0, 0)

    def table(slide, df: pd.DataFrame, max_rows: int = 8) -> None:
        data = df.head(max_rows).copy()
        shape = slide.shapes.add_table(len(data) + 1, len(data.columns), Inches(0.35), Inches(1.0), Inches(12.6), Inches(5.8))
        tbl = shape.table
        for j, col in enumerate(data.columns):
            cell = tbl.cell(0, j)
            cell.text = str(col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(79, 126, 207)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(8)
                p.font.bold = True
                p.font.color.rgb = RGBColor(255, 255, 255)
                p.alignment = PP_ALIGN.CENTER
        for i, (_, row) in enumerate(data.iterrows(), start=1):
            for j, col in enumerate(data.columns):
                value = row[col]
                if isinstance(value, float):
                    value = round(value, 4)
                cell = tbl.cell(i, j)
                cell.text = "" if pd.isna(value) else str(value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(235, 240, 250) if i % 2 == 0 else RGBColor(255, 255, 255)
                for p in cell.text_frame.paragraphs:
                    p.font.name = "Microsoft YaHei"
                    p.font.size = Pt(7)
                    p.font.color.rgb = RGBColor(0, 0, 0)

    slide = prs.slides.add_slide(blank)
    title(slide, "006_12 Expert dataset augmentation")
    bullets(slide, ["目标：先扩充 expert schedules，再决定是否回到 constrained PPO。", "不做 PPO multi-seed，不进入 rainfall-scaling，不覆盖 006_08 到 006_11。"])
    slide = prs.slides.add_slide(blank)
    title(slide, "Selected schedules")
    table(slide, station_schedules)
    slide = prs.slides.add_slide(blank)
    title(slide, "Action distribution")
    table(slide, dist)
    slide = prs.slides.add_slide(blank)
    title(slide, "Policy gate")
    table(slide, policy_summary[["policy_name", "eval_count", "mean_yield", "mean_irrigation", "mean_n", "mean_profit", "passes_gate"]])
    slide = prs.slides.add_slide(blank)
    title(slide, "Recommendation")
    table(slide, recommended[["policy_name", "recommendation", "next_step"]])
    slide = prs.slides.add_slide(blank)
    title(slide, "Rainfall-scaling decision")
    bullets(slide, ["Constrained PPO multi-seed can be considered only with the recommended augmented prior.", "Rainfall-scaling remains premature because augmented irrigation labels are still sparse.", "Next water-stress work should add irrigation-responsive expert schedules before stress-scenario RL."])
    prs.save(DOC_PPT)
    shutil.copy2(DOC_PPT, OUTPUT_ROOT / "reports" / DOC_PPT.name)


def prepare(config: dict) -> None:
    ensure_dirs()
    shutil.copy2(CONFIG_PATH, OUTPUT_ROOT / "configs" / CONFIG_PATH.name)
    write_why_not_multiseed(config)
    selected = select_augmented_schedules(config)
    dataset, _ = build_augmented_dataset(config, selected)
    original = pd.read_csv(PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "datasets" / "imitation_dataset_clean.csv")
    pd.DataFrame([action_distribution(original, "00609_original"), action_distribution(dataset, "00612_augmented")]).to_csv(
        OUTPUT_ROOT / "evaluation" / "augmented_dataset_action_distribution_summary.csv",
        index=False,
        encoding="utf-8-sig",
    )
    train_augmented_models(config, dataset)


def report(config: dict) -> None:
    make_figures(config)
    choose_policy(config)
    write_reports(config)


def main() -> None:
    parser = argparse.ArgumentParser(description="Augment expert schedules before constrained PPO.")
    parser.add_argument("--prepare", action="store_true", help="Select schedules, build dataset, train augmented BC models, export action tables.")
    parser.add_argument("--evaluate-tables", action="store_true", help="Evaluate exported action tables in DSSAT/gym-DSSAT.")
    parser.add_argument("--report", action="store_true", help="Generate figures, Markdown, and PPT from existing CSV outputs.")
    args = parser.parse_args()
    config = load_yaml(CONFIG_PATH)
    ensure_dirs()
    if args.prepare:
        prepare(config)
    if args.evaluate_tables:
        evaluate_tables(config)
    if args.report:
        report(config)
    if not (args.prepare or args.evaluate_tables or args.report):
        prepare(config)
        evaluate_tables(config)
        report(config)


if __name__ == "__main__":
    main()
