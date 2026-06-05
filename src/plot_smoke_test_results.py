from __future__ import annotations

import shutil
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SMOKE_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "smoke_tests"
SUMMARY_PATH = SMOKE_ROOT / "evaluation" / "smoke_test_summary.csv"
MINIMAL_PATH = SMOKE_ROOT / "evaluation" / "smoke_test_minimal_summary.csv"
REPORT_PATH = PROJECT_ROOT / "docs" / "2026-06-05_observed_year_smoke_test_report.md"
PPT_PATH = PROJECT_ROOT / "docs" / "2026-06-05_observed_year_smoke_test_report.pptx"
REPORT_COPY = SMOKE_ROOT / "reports" / REPORT_PATH.name
PPT_COPY = SMOKE_ROOT / "reports" / PPT_PATH.name
BLUE = RGBColor(31, 78, 121)
LIGHT_BLUE = RGBColor(221, 235, 247)


def setup_style() -> None:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False


def save_station_comparison_figures(summary: pd.DataFrame) -> list[Path]:
    setup_style()
    paths: list[Path] = []
    metrics = [
        ("final_grnwt", "final_grnwt_by_year_policy", "Final grain weight"),
        ("total_irrigation", "total_irrigation_by_year_policy", "Total irrigation"),
        ("total_n_fertilizer", "total_n_by_year_policy", "Total nitrogen fertilizer"),
        ("mean_swfac", "mean_swfac_by_year_policy", "Mean SWFAC"),
        ("mean_nstres", "mean_nstres_by_year_policy", "Mean NSTRES"),
    ]
    ok = summary[summary["run_status"] == "ok"].copy()
    for station in sorted(summary["station"].dropna().unique()):
        group = ok[ok["station"] == station]
        out_dir = SMOKE_ROOT / "figures" / station / "station_policy_comparison"
        out_dir.mkdir(parents=True, exist_ok=True)
        for metric, filename, title in metrics:
            fig, ax = plt.subplots(figsize=(8.5, 4.5))
            if group.empty:
                ax.text(0.5, 0.5, "No successful episodes for this station", ha="center", va="center", fontsize=13)
                ax.set_axis_off()
            else:
                pivot = group.pivot_table(index="year", columns="policy_name", values=metric, aggfunc="first")
                pivot.plot(kind="bar", ax=ax)
                ax.set_ylabel(metric)
                ax.grid(axis="y", alpha=0.25)
                ax.legend(frameon=False, fontsize=8)
            ax.set_title(f"{station} {title}")
            fig.tight_layout()
            path = out_dir / f"{station}_{filename}.png"
            fig.savefig(path, dpi=150)
            plt.close(fig)
            paths.append(path)
    return paths


def md_table(df: pd.DataFrame, columns: list[str], max_rows: int | None = None) -> str:
    if df.empty:
        return "无。\n"
    data = df[columns].copy()
    if max_rows is not None:
        data = data.head(max_rows)
    return data.to_markdown(index=False)


def make_report(summary: pd.DataFrame, minimal: pd.DataFrame, comparison_figures: list[Path]) -> None:
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    SMOKE_ROOT.joinpath("reports").mkdir(parents=True, exist_ok=True)
    status_counts = summary["run_status"].value_counts().reset_index()
    status_counts.columns = ["run_status", "count"]
    by_station = pd.crosstab(summary["station"], summary["run_status"]).reset_index()
    failures = summary[summary["run_status"] != "ok"].copy()
    ok = summary[summary["run_status"] == "ok"].copy()
    policy_means = ok.groupby(["station", "policy_name"], as_index=False).agg(
        final_grnwt_mean=("final_grnwt", "mean"),
        total_irrigation_mean=("total_irrigation", "mean"),
        total_n_mean=("total_n_fertilizer", "mean"),
        mean_swfac=("mean_swfac", "mean"),
        mean_nstres=("mean_nstres", "mean"),
    )
    daily_count = len(list((SMOKE_ROOT / "daily_outputs").glob("**/*_daily.csv")))
    episode_fig_count = len([p for p in (SMOKE_ROOT / "figures").glob("**/*.png") if "station_policy_comparison" not in str(p)])
    comparison_count = len(comparison_figures)
    can_enter_ppo = "否。YCA 全部超时，LCA/SYA 固定水氮策略大量超时；需要先排查这些环境输入或动作导致的卡顿。"
    lines = [
        "# 2026-06-05 observed year smoke test report",
        "",
        "## 1. 本阶段目的",
        "本阶段是环境体检，不是优化。目标是在进入 PPO 前验证 Phase 2c 推荐的真实试验年份是否能在 gym-DSSAT 中稳定跑完，并保存 daily CSV、响应图和汇总表。",
        "",
        "## 2. 使用的站点和年份",
        "测试年份来自 Phase 2c 的实测物候年份，共 14 个 station-year：FQA 2008/2010，HLA 2007/2011/2009，LCA 2010/2011/2008/2009，SYA 2014/2015/2012，YCA 2014/2008。",
        "",
        "## 3. 固定策略",
        "- null_zero: amir=0, anfer=0。",
        "- fixed_low_input: DAP 1/30 施氮 30/20 kg ha-1。",
        "- fixed_medium_input: DAP 1/30/60 施氮 50/50/50 kg ha-1，DAP 30/60 灌水 30/30 mm。",
        "- fixed_high_input: DAP 1/30/60 施氮 80/80/90 kg ha-1，DAP 30/60/90 灌水 40/40/40 mm。",
        "",
        "## 4. 最小 smoke test 结果",
        md_table(minimal, ["station", "year", "policy_name", "run_status", "episode_completed", "episode_length", "daily_csv_path"]),
        "",
        "## 5. 56 次 episode 批量运行状态",
        md_table(status_counts, ["run_status", "count"]),
        "",
        "按站点统计：",
        md_table(by_station, list(by_station.columns)),
        "",
        "## 6. 失败案例和错误信息",
        md_table(failures, ["station", "year", "policy_name", "run_status", "error_message"], max_rows=40),
        "",
        "## 7. daily output 和图表完整性",
        f"- 成功 episode 数：{len(ok)}。",
        f"- daily CSV 数：{daily_count}。",
        f"- episode 响应图数：{episode_fig_count}，按 5 张/成功 episode 计算应为 {len(ok) * 5}。",
        f"- 站点级对比图数：{comparison_count}。",
        "",
        "## 8. 固定策略响应对比",
        md_table(policy_means, ["station", "policy_name", "final_grnwt_mean", "total_irrigation_mean", "total_n_mean", "mean_swfac", "mean_nstres"]),
        "",
        "## 9. 变量缺失或异常",
        "成功 episode 的 daily CSV 均包含 prompt 要求的核心列：dap、topwt、grnwt、xlai、totir、tofer、swfac、nstres、reward、real_action_amir、real_action_anfer、normalized_action_amir、normalized_action_anfer。若环境没有 tofer，则脚本使用施氮动作累计值补充。",
        "",
        "## 10. 是否可以进入 PPO 训练脚本生成",
        can_enter_ppo,
        "",
        "## 11. 需要单独排错的站点或年份",
        "- LCA: null_zero 全部成功，但 fixed_low/medium/high 全部 timeout，优先检查固定动作是否导致 DSSAT 交互进程等待或输出异常。",
        "- SYA: null_zero 全部成功，但 fixed_low/medium/high 全部 timeout。",
        "- YCA: null_zero 和全部固定策略均 timeout，优先检查 YCA 临时模板、WTH 命名、土壤/品种/管理组合和 DSSAT 日志。",
        "",
        "## 12. 生成文件",
        f"- `{SUMMARY_PATH.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- `{MINIMAL_PATH.relative_to(PROJECT_ROOT).as_posix()}`",
        "- `Leave_One_experiments/smoke_tests/daily_outputs/`",
        "- `Leave_One_experiments/smoke_tests/figures/`",
        "- `Leave_One_experiments/smoke_tests/rendered_inputs/`",
        f"- `{REPORT_PATH.relative_to(PROJECT_ROOT).as_posix()}`",
        f"- `{PPT_PATH.relative_to(PROJECT_ROOT).as_posix()}`",
    ]
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    shutil.copyfile(REPORT_PATH, REPORT_COPY)


def add_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.5))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 0, 0)


def add_bullets(slide, bullets: list[str], top=1.05, size=18) -> None:
    box = slide.shapes.add_textbox(Inches(0.65), Inches(top), Inches(12.0), Inches(5.8))
    tf = box.text_frame
    tf.clear()
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(0, 0, 0)
        p.space_after = Pt(8)


def add_table(slide, df: pd.DataFrame, cols: list[str], max_rows=10) -> None:
    data = df[cols].head(max_rows).copy()
    table = slide.shapes.add_table(len(data) + 1, len(cols), Inches(0.45), Inches(1.05), Inches(12.4), Inches(5.8)).table
    for j, col in enumerate(cols):
        cell = table.cell(0, j)
        cell.text = col
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER
    for i in range(len(data)):
        for j, col in enumerate(cols):
            val = data.iloc[i][col]
            cell = table.cell(i + 1, j)
            if isinstance(val, float):
                text = f"{val:.1f}"
            else:
                text = "" if pd.isna(val) else str(val)
            cell.text = text
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_BLUE
            for p in cell.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(9)
                p.font.color.rgb = RGBColor(0, 0, 0)


def make_ppt(summary: pd.DataFrame, minimal: pd.DataFrame) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    ok = summary[summary["run_status"] == "ok"].copy()
    failures = summary[summary["run_status"] != "ok"].copy()
    status_counts = summary["run_status"].value_counts().reset_index()
    status_counts.columns = ["run_status", "count"]
    by_station = pd.crosstab(summary["station"], summary["run_status"]).reset_index()
    policy_means = ok.groupby(["station", "policy_name"], as_index=False).agg(
        final_grnwt_mean=("final_grnwt", "mean"),
        total_irrigation_mean=("total_irrigation", "mean"),
        total_n_mean=("total_n_fertilizer", "mean"),
    )
    slides = [
        ("任务目标", ["对 Phase 2c 推荐的 14 个真实 station-year 做环境体检。", "策略：NullAgent 与固定低/中/高水氮投入。", "不训练 PPO，不修改 reward，不覆盖 my_data 原始数据。"]),
        ("输入数据和站点年份", ["FQA 2008/2010；HLA 2007/2011/2009。", "LCA 2010/2011/2008/2009；SYA 2014/2015/2012。", "YCA 2014/2008。"]),
        ("固定策略设计", ["null_zero: 不灌水、不施肥。", "fixed_low: 总氮 50 kg/ha。", "fixed_medium: 总氮 150 kg/ha，总灌水 60 mm。", "fixed_high: 总氮 250 kg/ha，总灌水 120 mm。"]),
    ]
    for title, bullets in slides:
        s = prs.slides.add_slide(blank)
        add_title(s, title)
        add_bullets(s, bullets)
    s = prs.slides.add_slide(blank)
    add_title(s, "最小 smoke test 结果")
    add_table(s, minimal, ["station", "year", "policy_name", "run_status", "episode_completed", "episode_length"])
    s = prs.slides.add_slide(blank)
    add_title(s, "56 次 episode 批量运行状态")
    add_table(s, status_counts, ["run_status", "count"])
    s = prs.slides.add_slide(blank)
    add_title(s, "按站点运行状态")
    add_table(s, by_station, list(by_station.columns))
    s = prs.slides.add_slide(blank)
    add_title(s, "固定策略产量对比")
    add_table(s, policy_means, ["station", "policy_name", "final_grnwt_mean", "total_irrigation_mean", "total_n_mean"], max_rows=14)
    example = SMOKE_ROOT / "figures" / "HLA" / "2007" / "null_zero" / "crop_growth_timeseries.png"
    if example.exists():
        s = prs.slides.add_slide(blank)
        add_title(s, "示例 daily response 图")
        s.shapes.add_picture(str(example), Inches(0.8), Inches(1.1), width=Inches(11.5))
    s = prs.slides.add_slide(blank)
    add_title(s, "失败或异常案例")
    add_table(s, failures, ["station", "year", "policy_name", "run_status", "error_message"], max_rows=12)
    s = prs.slides.add_slide(blank)
    add_title(s, "下一步建议")
    add_bullets(s, ["暂不建议直接进入 PPO。", "先排查 YCA 全部 timeout，以及 LCA/SYA 固定策略 timeout。", "优先检查临时模板、WTH 文件名、土壤文件、动作上限和 DSSAT 交互等待。"])
    prs.save(PPT_PATH)
    shutil.copyfile(PPT_PATH, PPT_COPY)


def main() -> None:
    summary = pd.read_csv(SUMMARY_PATH)
    minimal = pd.read_csv(MINIMAL_PATH)
    figs = save_station_comparison_figures(summary)
    make_report(summary, minimal, figs)
    make_ppt(summary, minimal)
    print(REPORT_PATH.relative_to(PROJECT_ROOT))
    print(PPT_PATH.relative_to(PROJECT_ROOT))
    print(f"station_comparison_figures={len(figs)}")


if __name__ == "__main__":
    main()
