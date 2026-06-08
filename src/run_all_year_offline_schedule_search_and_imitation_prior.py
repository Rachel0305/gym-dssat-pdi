from __future__ import annotations

import argparse
import itertools
import json
import pickle
import shutil
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import yaml

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except ModuleNotFoundError:
    Presentation = None
    RGBColor = None
    Inches = None
    Pt = None

from offline_schedule_policy import DeterministicSchedulePolicy, ScheduleEvent
from ppo_action_safety import ActionSafetyState, apply_action_safety, normalize_action, update_action_safety_state
from ppo_evaluate import latest_observation_dict, make_env, scalar
from ppo_experiment_plan import find_year
from ppo_plot_results import plot_episode
from ppo_safe_rendering import PROJECT_ROOT, SITE_INFO


CONFIG_PATH = PROJECT_ROOT / "experiments" / "ppo_observed_years" / "config_all_year_offline_schedule_search.yaml"
OUTPUT_ROOT = PROJECT_ROOT / "Leave_One_experiments" / "all_year_offline_schedule_search"
DOC_MD = PROJECT_ROOT / "docs" / "2026-06-06_all_year_offline_schedule_search_and_imitation_prior_report.md"
DOC_PPT = PROJECT_ROOT / "docs" / "2026-06-06_all_year_offline_schedule_search_and_imitation_prior_report.pptx"

FEATURE_COLUMNS = [
    "station",
    "sim_day",
    "doy",
    "topwt",
    "grnwt",
    "xlai",
    "totir",
    "tofer",
    "swfac",
    "nstres",
]
TARGET_COLUMNS = ["expert_action_irrigation", "expert_action_n"]


@dataclass(frozen=True)
class ScheduleSpec:
    schedule_id: str
    station_code: str
    year: int
    schedule_family: str
    n_by_dap: dict[int, float]
    i_by_dap: dict[int, float]

    @property
    def total_n(self) -> float:
        return float(sum(self.n_by_dap.values()))

    @property
    def total_irrigation(self) -> float:
        return float(sum(self.i_by_dap.values()))


def ensure_dirs() -> None:
    for sub in [
        "configs",
        "candidate_schedules",
        "daily_outputs",
        "evaluation",
        "expert_policy",
        "imitation_dataset",
        "models",
        "figures",
        "reports",
        "rendered_inputs",
        "logs",
    ]:
        (OUTPUT_ROOT / sub).mkdir(parents=True, exist_ok=True)


def load_yaml(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def df_to_markdown(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df.empty:
        return ""
    work = df.copy()
    work.columns = [str(col) for col in work.columns]
    if max_rows is not None:
        work = work.head(max_rows)
    for col in work.select_dtypes(include=["float", "int"]).columns:
        work[col] = pd.to_numeric(work[col], errors="coerce").round(4)
    for col in work.columns:
        work[col] = work[col].map(lambda value: "" if pd.isna(value) else str(value))
    header = "| " + " | ".join(work.columns) + " |"
    sep = "| " + " | ".join(["---"] * len(work.columns)) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in work.astype(str).values.tolist()]
    return "\n".join([header, sep, *rows])


def read_csv(path: str | Path) -> pd.DataFrame:
    return pd.read_csv(PROJECT_ROOT / path if not Path(path).is_absolute() else path)


def date_from_doy(year: int, doy: int) -> pd.Timestamp:
    return pd.Timestamp(year=int(year), month=1, day=1) + pd.Timedelta(days=int(doy) - 1)


def representative_phenology() -> dict[str, dict[str, int]]:
    path = PROJECT_ROOT / "data" / "observed_phenology_dates_standardized.csv"
    pheno = pd.read_csv(path, parse_dates=["planting_date", "harvest_date"])
    rows: dict[str, dict[str, int]] = {}
    for station, group in pheno.groupby("station_code"):
        planting_doy = int(round(group["planting_date"].dt.dayofyear.median()))
        harvest_days = int(round((group["harvest_date"] - group["planting_date"]).dt.days.median() + 1))
        rows[str(station)] = {"planting_doy": planting_doy, "harvest_days": harvest_days}
    return rows


def build_env_config(config: dict, selection: pd.DataFrame | None = None) -> dict:
    base = load_yaml(PROJECT_ROOT / config["paths"]["base_env_config"])
    base["paths"]["output_root"] = config["paths"]["output_root"]
    base["paths"]["rendered_input_root"] = str(Path(config["paths"]["output_root"]) / "rendered_inputs")
    base["seed"] = int(config.get("seed", 0))
    base["runtime"]["max_steps"] = int(config["runtime"].get("max_steps", 260))
    base["action_safety"] = config["action_safety"]
    base["economics"] = config["economics"]
    reps = representative_phenology()
    if selection is None:
        inventory = read_csv(config["paths"]["weather_inventory_csv"])
        selection = inventory[inventory["weather_qc_status"].eq("ok")][["station_code", "year", "weather_file"]].copy()
    observed_years: dict[str, list[dict[str, Any]]] = {}
    for row in selection.drop_duplicates(["station_code", "year"]).itertuples(index=False):
        station = str(row.station_code)
        year = int(row.year)
        rep = reps[station]
        planting = date_from_doy(year, rep["planting_doy"])
        observed_years.setdefault(station, []).append(
            {
                "year": year,
                "planting_date": planting.strftime("%Y-%m-%d"),
                "weather_file": str(row.weather_file),
                "source": "all_year_weather_pool_representative_phenology",
            }
        )
    base["observed_years"] = observed_years
    return base


def scenario_group(value: str) -> str:
    text = str(value)
    if "irrigation_responsive_year" in text:
        return "irrigation_responsive"
    if "water_stress_year" in text:
        return "water_stress"
    if "dry_year" in text:
        return "dry"
    if "wet_year" in text:
        return "wet"
    return "normal"


def select_search_years(config: dict, pool: pd.DataFrame, plan: pd.DataFrame) -> pd.DataFrame:
    pool = pool.copy()
    pool["scenario_group"] = pool["scenario_type"].map(scenario_group)
    pool["selected_for_search"] = False
    reasons: dict[int, list[str]] = {idx: [] for idx in pool.index}

    priority = pool["has_water_stress"].astype(bool) | pool["irrigation_responsive"].astype(bool)
    pool.loc[priority, "selected_for_search"] = True
    for idx in pool.index[priority]:
        if bool(pool.at[idx, "has_water_stress"]):
            reasons[idx].append("water_stress_year")
        if bool(pool.at[idx, "irrigation_responsive"]):
            reasons[idx].append("irrigation_responsive_year")

    if config["selection"].get("always_include_2020_2023", True):
        calib_years = set(zip(plan["station_code"].astype(str), plan["year"].astype(int)))
        is_2020_2023 = pool.apply(lambda r: (str(r["station_code"]), int(r["year"])) in calib_years, axis=1)
        pool.loc[is_2020_2023, "selected_for_search"] = True
        for idx in pool.index[is_2020_2023]:
            reasons[idx].append("calibration_validation_candidate_2020_2023")

    for station, group in pool.groupby("station_code"):
        for group_name, limit_key in [
            ("dry", "dry_years_per_station"),
            ("normal", "normal_years_per_station"),
            ("wet", "wet_years_per_station"),
        ]:
            limit = int(config["selection"].get(limit_key, 0))
            if limit <= 0:
                continue
            sub = group[group["scenario_group"].eq(group_name)].copy()
            sub = sub.sort_values(["recommended_for_ppo_eval", "recommended_for_ppo_train", "growing_season_rain"], ascending=[False, False, True])
            selected_idx = sub.head(limit).index
            pool.loc[selected_idx, "selected_for_search"] = True
            for idx in selected_idx:
                reasons[idx].append(f"stratified_{group_name}")

    selected = pool[pool["selected_for_search"]].copy()
    max_years = int(config["runtime"].get("max_station_years", 0))
    if max_years and len(selected) > max_years:
        selected["priority_score"] = (
            selected["irrigation_responsive"].astype(int) * 100
            + selected["has_water_stress"].astype(int) * 50
            + selected["recommended_for_validation"].astype(int) * 10
            + selected["recommended_for_calibration"].astype(int) * 8
            + selected["recommended_for_ppo_eval"].astype(int) * 4
            + selected["recommended_for_ppo_train"].astype(int)
        )
        selected = selected.sort_values(["priority_score", "station_code", "year"], ascending=[False, True, True]).head(max_years)
        pool["selected_for_search"] = pool.index.isin(selected.index)

    pool["selected_for_train_pool"] = pool["selected_for_search"] & pool["recommended_for_ppo_train"].astype(bool)
    pool["selected_for_eval_pool"] = pool["selected_for_search"] & pool["recommended_for_ppo_eval"].astype(bool)
    pool["selection_reason"] = [";".join(dict.fromkeys(reasons[idx])) for idx in pool.index]
    cols = [
        "station_code",
        "station_name",
        "year",
        "weather_file",
        "scenario_type",
        "scenario_group",
        "has_water_stress",
        "has_nitrogen_stress",
        "irrigation_responsive",
        "recommended_for_calibration",
        "recommended_for_validation",
        "recommended_for_ppo_train",
        "recommended_for_ppo_eval",
        "selected_for_search",
        "selected_for_train_pool",
        "selected_for_eval_pool",
        "selection_reason",
        "growing_season_rain",
        "max_swfac",
        "nstres_days_gt_0p05",
        "yield_gain_from_irrigation_at_same_N",
    ]
    out = pool.loc[:, cols].sort_values(["selected_for_search", "station_code", "year"], ascending=[False, True, True])
    out_path = OUTPUT_ROOT / "configs" / "all_year_search_year_selection.csv"
    out.to_csv(out_path, index=False, encoding="utf-8-sig")
    selected_out = out[out["selected_for_search"]].copy()
    md = [
        "# All-Year Search Year Selection",
        "",
        f"- Selected station-years: {len(selected_out)}",
        f"- Water-stress selected: {int(selected_out['has_water_stress'].sum())}",
        f"- Irrigation-responsive selected: {int(selected_out['irrigation_responsive'].sum())}",
        f"- 2020-2023 selected: {int(selected_out['year'].isin([2020, 2021, 2022, 2023]).sum())}",
        "",
        df_to_markdown(selected_out[["station_code", "year", "scenario_group", "selection_reason", "growing_season_rain", "max_swfac"]], max_rows=120),
    ]
    (OUTPUT_ROOT / "configs" / "all_year_search_year_selection.md").write_text("\n".join(md), encoding="utf-8")
    return out


def schedule_from_amounts(station: str, year: int, family: str, n_amounts: tuple[float, ...], i_events: list[int], i_amounts: tuple[float, ...], sid: int) -> ScheduleSpec:
    n_by_dap = {dap: float(amount) for dap, amount in zip([1, 30, 60], n_amounts) if float(amount) > 0}
    i_by_dap = {dap: float(amount) for dap, amount in zip(i_events, i_amounts) if float(amount) > 0}
    schedule_id = f"{station}{year}_{family}_S{sid:04d}"
    return ScheduleSpec(schedule_id, station, year, family, n_by_dap, i_by_dap)


def generate_schedules_for_year(config: dict, row: pd.Series) -> pd.DataFrame:
    station = str(row["station_code"])
    year = int(row["year"])
    stress_like = bool(row["has_water_stress"]) or bool(row["irrigation_responsive"])
    search = config["offline_search"]
    n_events = [1, 30, 60]
    n_amounts = [float(v) for v in search["nitrogen_amounts"]]
    n_combos = [combo for combo in itertools.product(n_amounts, repeat=3) if sum(combo) <= float(search["total_n_cap"])]
    if stress_like:
        i_events = [int(v) for v in search["irrigation_events_full"]]
        i_amounts = [float(v) for v in search["irrigation_amounts_full"]]
        i_combos = [combo for combo in itertools.product(i_amounts, repeat=len(i_events)) if sum(combo) <= float(search["total_irrigation_cap"])]
        max_schedules = int(search["max_schedules_stress_year"])
    else:
        i_events = [int(v) for v in search["irrigation_events_simple"]]
        i_amounts = [float(v) for v in search["irrigation_amounts_simple"]]
        i_combos = [combo for combo in itertools.product(i_amounts, repeat=len(i_events)) if sum(combo) <= float(search["total_irrigation_cap"])]
        max_schedules = int(search["max_schedules_normal_year"])

    curated_n = [(0, 0, 0), (50, 50, 50), (75, 75, 75), (100, 50, 50), (50, 100, 50), (50, 50, 100), (100, 75, 75)]
    curated_i = [tuple([0] * len(i_events))]
    for total in ([20, 40, 60, 80, 120, 160] if stress_like else [30, 60, 120]):
        amount = total / len(i_events)
        rounded = min(i_amounts, key=lambda v: abs(v - amount))
        curated_i.append(tuple([rounded] * len(i_events)))
    if stress_like:
        for idx in range(len(i_events)):
            for amount in [20, 40, 60]:
                combo = [0.0] * len(i_events)
                combo[idx] = amount
                curated_i.append(tuple(combo))

    rows: list[dict[str, Any]] = []
    specs: list[ScheduleSpec] = []
    sid = 0
    seen: set[tuple] = set()
    for n_combo in curated_n + n_combos:
        if sum(n_combo) > float(search["total_n_cap"]):
            continue
        for i_combo in curated_i + i_combos:
            if sum(i_combo) > float(search["total_irrigation_cap"]):
                continue
            key = tuple(n_combo) + tuple(i_combo)
            if key in seen:
                continue
            seen.add(key)
            sid += 1
            family = "stress_grid" if stress_like else "normal_grid"
            specs.append(schedule_from_amounts(station, year, family, tuple(map(float, n_combo)), i_events, tuple(map(float, i_combo)), sid))
            if sid >= max_schedules:
                break
        if sid >= max_schedules:
            break
    for spec in specs:
        row_dict: dict[str, Any] = {
            "station_code": station,
            "year": year,
            "schedule_id": spec.schedule_id,
            "schedule_family": spec.schedule_family,
            "total_n": spec.total_n,
            "total_irrigation": spec.total_irrigation,
        }
        for dap in n_events:
            row_dict[f"N_DAP{dap}"] = spec.n_by_dap.get(dap, 0.0)
        for dap in i_events:
            row_dict[f"I_DAP{dap}"] = spec.i_by_dap.get(dap, 0.0)
        rows.append(row_dict)
    return pd.DataFrame(rows)


def policy_from_schedule(schedule: pd.Series) -> DeterministicSchedulePolicy:
    events: list[ScheduleEvent] = []
    for col in schedule.index:
        if str(col).startswith("N_DAP") and float(schedule[col]) > 0:
            dap = int(str(col).replace("N_DAP", ""))
            events.append(ScheduleEvent(name=f"N{dap}", dap=dap, nitrogen=float(schedule[col])))
        if str(col).startswith("I_DAP") and float(schedule[col]) > 0:
            dap = int(str(col).replace("I_DAP", ""))
            events.append(ScheduleEvent(name=f"I{dap}", dap=dap, irrigation=float(schedule[col])))
    return DeterministicSchedulePolicy(events)


def clip_to_action_space(env, real_action: dict[str, float]) -> tuple[dict[str, float], str]:
    spaces = getattr(env.formator.action_space_dict, "spaces", env.formator.action_space_dict)
    clipped: dict[str, float] = {}
    notes: list[str] = []
    for name in env.formator.action_names:
        space = spaces[name]
        low = float(np.asarray(space.low).flatten()[0])
        high = float(np.asarray(space.high).flatten()[0])
        raw = float(real_action.get(name, 0.0))
        value = min(max(raw, low), high)
        clipped[name] = value
        if abs(value - raw) > 1e-9:
            notes.append(f"{name}_env_bound_clip_{raw:g}_to_{value:g}")
    return clipped, ";".join(notes)


def profit(final_grnwt: float, irrigation: float, nitrogen: float, config: dict, low_water: bool = False) -> float:
    econ = config["economics"]
    water_cost = float(econ["low_water_cost"] if low_water else econ["water_cost"])
    return float(econ["grain_value_coef"]) * float(final_grnwt) - water_cost * float(irrigation) - float(econ["n_cost"]) * float(nitrogen)


def summarize_daily(daily: pd.DataFrame, config: dict) -> dict[str, Any]:
    final_grnwt = float(pd.to_numeric(daily["grnwt"], errors="coerce").iloc[-1]) if len(daily) else np.nan
    final_topwt = float(pd.to_numeric(daily["topwt"], errors="coerce").iloc[-1]) if len(daily) else np.nan
    final_xlai = float(pd.to_numeric(daily["xlai"], errors="coerce").iloc[-1]) if len(daily) else np.nan
    total_i = float(pd.to_numeric(daily["real_action_amir"], errors="coerce").fillna(0).sum()) if len(daily) else 0.0
    total_n = float(pd.to_numeric(daily["real_action_anfer"], errors="coerce").fillna(0).sum()) if len(daily) else 0.0
    swfac = pd.to_numeric(daily.get("swfac", pd.Series(dtype=float)), errors="coerce")
    nstres = pd.to_numeric(daily.get("nstres", pd.Series(dtype=float)), errors="coerce")
    return {
        "episode_completed": bool(len(daily) and bool(daily["done"].iloc[-1])),
        "episode_length": int(len(daily)),
        "final_grnwt": final_grnwt,
        "final_topwt": final_topwt,
        "final_xlai": final_xlai,
        "total_irrigation": total_i,
        "total_n": total_n,
        "profit_default": profit(final_grnwt, total_i, total_n, config, low_water=False) if not np.isnan(final_grnwt) else np.nan,
        "profit_low_water_cost": profit(final_grnwt, total_i, total_n, config, low_water=True) if not np.isnan(final_grnwt) else np.nan,
        "mean_swfac": float(swfac.mean()) if len(swfac) else np.nan,
        "max_swfac": float(swfac.max()) if len(swfac) else np.nan,
        "swfac_stress_days_gt_0p05": int((swfac > 0.05).sum()) if len(swfac) else 0,
        "mean_nstres": float(nstres.mean()) if len(nstres) else np.nan,
        "max_nstres": float(nstres.max()) if len(nstres) else np.nan,
        "nstres_days_gt_0p05": int((nstres > 0.05).sum()) if len(nstres) else 0,
        "irrigation_event_count": int((pd.to_numeric(daily["real_action_amir"], errors="coerce").fillna(0) > 0).sum()) if len(daily) else 0,
        "n_event_count": int((pd.to_numeric(daily["real_action_anfer"], errors="coerce").fillna(0) > 0).sum()) if len(daily) else 0,
    }


def evaluate_schedule(config: dict, env_config: dict, schedule: pd.Series, scenario: pd.Series) -> dict[str, Any]:
    station = str(schedule["station_code"])
    year = int(schedule["year"])
    schedule_id = str(schedule["schedule_id"])
    daily_dir = OUTPUT_ROOT / "daily_outputs" / station
    daily_dir.mkdir(parents=True, exist_ok=True)
    daily_csv = daily_dir / f"{station}_{year}_{schedule_id}_daily.csv"
    if daily_csv.exists():
        daily = pd.read_csv(daily_csv)
        summary = summarize_daily(daily, config)
        return {
            **base_schedule_summary(schedule, scenario),
            "run_status": "ok_cached" if summary["episode_completed"] else "failed_cached",
            "error_message": "",
            **summary,
            "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
        }

    env = make_env(env_config, station, year, int(config.get("seed", 0)), run_tag=f"{schedule_id}", evaluation=True, action_safety_enabled=False)
    policy = policy_from_schedule(schedule)
    safety_state = ActionSafetyState()
    safety_config = {**config["action_safety"], "enabled": True}
    records: list[dict[str, Any]] = []
    try:
        obs, info = env.reset()
        done = False
        step_count = 0
        year_info = find_year(env_config, station, year)
        planting = pd.Timestamp(year_info["planting_date"])
        while not done and step_count < int(config["runtime"].get("max_steps", 260)):
            latest = latest_observation_dict(env, obs, info)
            dap_raw = scalar(latest.get("dap", np.nan))
            dap = int(round(dap_raw)) if not np.isnan(dap_raw) and dap_raw > 0 else step_count + 1
            planned = policy.action_for_dap(dap)
            bounded, env_clip_note = clip_to_action_space(env, planned)
            safety_result = apply_action_safety(bounded, dap, safety_state, safety_config)
            safe_action = safety_result.safe_real_action
            normalized_action = normalize_action(env.formator.action_names, env.formator.action_space_dict, safe_action)
            obs, reward, terminated, truncated, info = env.step(normalized_action)
            done = bool(terminated or truncated)
            latest = latest_observation_dict(env, obs, info)
            update_action_safety_state(safety_state, safe_action, dap)
            date = planting + pd.Timedelta(days=max(dap - 1, 0))
            real_i = float(safe_action.get("amir", 0.0))
            real_n = float(safe_action.get("anfer", 0.0))
            records.append(
                {
                    "station": station,
                    "year": year,
                    "schedule_id": schedule_id,
                    "date": date.strftime("%Y-%m-%d"),
                    "doy": int(date.dayofyear),
                    "sim_day": int(step_count + 1),
                    "dap": dap,
                    "topwt": scalar(latest.get("topwt")),
                    "grnwt": scalar(latest.get("grnwt")),
                    "xlai": scalar(latest.get("xlai")),
                    "totir": scalar(latest.get("totir")),
                    "tofer": scalar(latest.get("tofer")),
                    "swfac": scalar(latest.get("swfac")),
                    "nstres": scalar(latest.get("nstres")),
                    "reward": float(reward),
                    "planned_event_name": policy.event_name_for_dap(dap),
                    "is_schedule_event_day": bool(policy.is_event_day(dap)),
                    "planned_irrigation": float(planned.get("amir", 0.0)),
                    "planned_n": float(planned.get("anfer", 0.0)),
                    "real_action_amir": real_i,
                    "real_action_anfer": real_n,
                    "expert_action_irrigation": real_i,
                    "expert_action_n": real_n,
                    "env_clip_note": env_clip_note,
                    "safety_rule_triggered": safety_result.safety_rule_triggered,
                    "done": done,
                }
            )
            step_count += 1
    except Exception as exc:
        return {**base_schedule_summary(schedule, scenario), "run_status": "failed", "error_message": traceback.format_exc(), "daily_csv_path": ""}
    finally:
        try:
            env.close()
        except Exception:
            pass
    daily = pd.DataFrame(records)
    daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
    if config["runtime"].get("plot_daily_episodes", False):
        plot_episode(daily, OUTPUT_ROOT / "figures" / "daily" / station / schedule_id)
    summary = summarize_daily(daily, config)
    return {
        **base_schedule_summary(schedule, scenario),
        "run_status": "ok" if summary["episode_completed"] else "failed",
        "error_message": "",
        **summary,
        "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
    }


def base_schedule_summary(schedule: pd.Series, scenario: pd.Series) -> dict[str, Any]:
    return {
        "station_code": str(schedule["station_code"]),
        "station_name": str(scenario.get("station_name", schedule["station_code"])),
        "year": int(schedule["year"]),
        "weather_file": str(scenario.get("weather_file", "")),
        "scenario_type": str(scenario.get("scenario_type", "")),
        "scenario_group": str(scenario.get("scenario_group", "")),
        "has_water_stress": bool(scenario.get("has_water_stress", False)),
        "has_nitrogen_stress": bool(scenario.get("has_nitrogen_stress", False)),
        "irrigation_responsive": bool(scenario.get("irrigation_responsive", False)),
        "schedule_id": str(schedule["schedule_id"]),
        "schedule_family": str(schedule["schedule_family"]),
        "planned_total_irrigation": float(schedule["total_irrigation"]),
        "planned_total_n": float(schedule["total_n"]),
    }


def run_schedule_search(config: dict, env_config: dict, selection: pd.DataFrame) -> pd.DataFrame:
    selected = selection[selection["selected_for_search"].astype(bool)].copy()
    schedule_frames = []
    for _, scenario in selected.iterrows():
        schedules = generate_schedules_for_year(config, scenario)
        schedule_frames.append(schedules)
        out = OUTPUT_ROOT / "candidate_schedules" / f"{scenario.station_code}_{int(scenario.year)}_candidate_schedules.csv"
        schedules.to_csv(out, index=False, encoding="utf-8-sig")
    all_schedules = pd.concat(schedule_frames, ignore_index=True) if schedule_frames else pd.DataFrame()
    all_schedules.to_csv(OUTPUT_ROOT / "candidate_schedules" / "all_year_candidate_schedules.csv", index=False, encoding="utf-8-sig")

    summary_path = OUTPUT_ROOT / "evaluation" / "all_year_schedule_search_summary.csv"
    existing = pd.read_csv(summary_path) if summary_path.exists() else pd.DataFrame()
    done_keys = set(zip(existing.get("station_code", []), existing.get("year", []), existing.get("schedule_id", []), existing.get("run_status", [])))
    rows: list[dict[str, Any]] = []
    scenario_lookup = selected.set_index(["station_code", "year"])
    for idx, schedule in all_schedules.iterrows():
        key_ok = (schedule["station_code"], int(schedule["year"]), schedule["schedule_id"], "ok")
        key_cached = (schedule["station_code"], int(schedule["year"]), schedule["schedule_id"], "ok_cached")
        if key_ok in done_keys or key_cached in done_keys:
            old = existing[
                existing["station_code"].astype(str).eq(str(schedule["station_code"]))
                & existing["year"].astype(int).eq(int(schedule["year"]))
                & existing["schedule_id"].astype(str).eq(str(schedule["schedule_id"]))
            ].iloc[0].to_dict()
            rows.append(old)
            continue
        scenario = scenario_lookup.loc[(str(schedule["station_code"]), int(schedule["year"]))]
        print(f"[00616-search] {schedule['station_code']} {int(schedule['year'])} {schedule['schedule_id']}", flush=True)
        rows.append(evaluate_schedule(config, env_config, schedule, scenario))
        if len(rows) % 25 == 0:
            pd.DataFrame(rows).to_csv(summary_path, index=False, encoding="utf-8-sig")
    summary = pd.DataFrame(rows)
    summary.to_csv(summary_path, index=False, encoding="utf-8-sig")
    return summary


def pareto_frontier(df: pd.DataFrame) -> pd.Series:
    work = df[["final_grnwt", "profit_low_water_cost", "total_irrigation", "total_n"]].copy()
    values = work.to_numpy(dtype=float)
    is_pareto = np.ones(len(values), dtype=bool)
    for i, row in enumerate(values):
        if not is_pareto[i]:
            continue
        better_or_equal = (values[:, 0] >= row[0]) & (values[:, 1] >= row[1]) & (values[:, 2] <= row[2]) & (values[:, 3] <= row[3])
        strictly_better = (values[:, 0] > row[0]) | (values[:, 1] > row[1]) | (values[:, 2] < row[2]) | (values[:, 3] < row[3])
        dominated_by_other = better_or_equal & strictly_better
        dominated_by_other[i] = False
        if dominated_by_other.any():
            is_pareto[i] = False
    return pd.Series(is_pareto, index=df.index)


def select_expert_library(config: dict, summary: pd.DataFrame) -> pd.DataFrame:
    ok = summary[summary["run_status"].astype(str).str.startswith("ok")].copy()
    if ok.empty:
        return ok
    selected_rows: list[pd.DataFrame] = []
    top_n = int(config["offline_search"].get("top_schedules_per_station_year", 5))
    for (station, year), group in ok.groupby(["station_code", "year"]):
        group = group.copy()
        best_yield = group["final_grnwt"].max()
        group["yield_loss_vs_best_in_year"] = (best_yield - group["final_grnwt"]) / best_yield
        group["is_pareto"] = pareto_frontier(group)
        categories = [
            ("top_profit_low_water", group.sort_values("profit_low_water_cost", ascending=False).head(1)),
            ("top_yield", group.sort_values("final_grnwt", ascending=False).head(1)),
            ("pareto_balanced", group[group["is_pareto"]].sort_values(["yield_loss_vs_best_in_year", "profit_low_water_cost"], ascending=[True, False]).head(2)),
            (
                "low_input_within_10pct_yield",
                group[group["yield_loss_vs_best_in_year"] <= 0.10].sort_values(["total_irrigation", "total_n", "profit_low_water_cost"], ascending=[True, True, False]).head(1),
            ),
            (
                "irrigation_positive_candidate",
                group[group["total_irrigation"] > 0].sort_values(["profit_low_water_cost", "final_grnwt"], ascending=[False, False]).head(1),
            ),
        ]
        for label, frame in categories:
            if frame.empty:
                continue
            temp = frame.copy()
            temp["expert_type"] = label
            selected_rows.append(temp)
    expert = pd.concat(selected_rows, ignore_index=True) if selected_rows else pd.DataFrame()
    expert = expert.sort_values(["station_code", "year", "expert_type", "schedule_id"]).drop_duplicates(["station_code", "year", "schedule_id"])
    expert.to_csv(OUTPUT_ROOT / "expert_policy" / "all_year_expert_schedule_library.csv", index=False, encoding="utf-8-sig")
    return expert


def build_imitation_dataset(expert: pd.DataFrame) -> pd.DataFrame:
    frames = []
    for row in expert.itertuples(index=False):
        path = PROJECT_ROOT / str(row.daily_csv_path)
        if not path.exists():
            continue
        daily = pd.read_csv(path)
        daily["station"] = str(row.station_code)
        daily["station_code"] = str(row.station_code)
        daily["station_name"] = str(row.station_name)
        daily["year"] = int(row.year)
        daily["weather_file"] = str(row.weather_file)
        daily["scenario_type"] = str(row.scenario_type)
        daily["scenario_group"] = str(row.scenario_group)
        daily["schedule_id"] = str(row.schedule_id)
        daily["expert_type"] = str(row.expert_type)
        daily["total_schedule_irrigation"] = float(row.total_irrigation)
        daily["total_schedule_n"] = float(row.total_n)
        daily["profit_default"] = float(row.profit_default)
        daily["profit_low_water_cost"] = float(row.profit_low_water_cost)
        daily["source_daily_csv"] = str(row.daily_csv_path)
        for col in ["topwt", "grnwt", "xlai", "totir", "tofer", "swfac", "nstres"]:
            if col not in daily.columns:
                daily[col] = 0.0
        daily["state_variables"] = daily[["sim_day", "doy", "topwt", "grnwt", "xlai", "totir", "tofer", "swfac", "nstres"]].apply(
            lambda s: json.dumps(s.to_dict(), ensure_ascii=False), axis=1
        )
        daily["expert_action_irrigation"] = pd.to_numeric(daily["expert_action_irrigation"], errors="coerce").fillna(0.0)
        daily["expert_action_n"] = pd.to_numeric(daily["expert_action_n"], errors="coerce").fillna(0.0)
        frames.append(daily)
    data = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    keep = [
        "station_code",
        "station",
        "station_name",
        "year",
        "weather_file",
        "scenario_type",
        "scenario_group",
        "schedule_id",
        "expert_type",
        "date",
        "doy",
        "sim_day",
        "dap",
        "state_variables",
        "swfac",
        "nstres",
        "topwt",
        "grnwt",
        "xlai",
        "totir",
        "tofer",
        "expert_action_irrigation",
        "expert_action_n",
        "total_schedule_irrigation",
        "total_schedule_n",
        "profit_default",
        "profit_low_water_cost",
        "source_daily_csv",
    ]
    data = data[[col for col in keep if col in data.columns]].copy()
    out = OUTPUT_ROOT / "imitation_dataset" / "all_year_imitation_dataset.csv"
    data.to_csv(out, index=False, encoding="utf-8-sig")
    return data


def prior_dataset_counts() -> pd.DataFrame:
    candidates = [
        ("00609_original", PROJECT_ROOT / "Leave_One_experiments" / "imitation_learning_prior" / "datasets" / "imitation_dataset_clean.csv"),
        ("00609_original_alt", PROJECT_ROOT / "Leave_One_experiments" / "offline_schedule_search" / "expert_policy" / "imitation_dataset.csv"),
        ("00612_augmented", PROJECT_ROOT / "Leave_One_experiments" / "expert_dataset_augmentation" / "imitation_dataset" / "imitation_dataset_augmented.csv"),
    ]
    rows = []
    for label, path in candidates:
        if not path.exists():
            continue
        data = pd.read_csv(path)
        i_col = "expert_action_irrigation" if "expert_action_irrigation" in data.columns else "real_action_amir"
        n_col = "expert_action_n" if "expert_action_n" in data.columns else "real_action_anfer"
        rows.append(
            {
                "dataset": label,
                "path": str(path.relative_to(PROJECT_ROOT)),
                "rows": len(data),
                "station_years": data[["station", "year"]].drop_duplicates().shape[0] if {"station", "year"}.issubset(data.columns) else np.nan,
                "nonzero_irrigation_rows": int((pd.to_numeric(data.get(i_col, 0), errors="coerce").fillna(0) > 0).sum()),
                "nonzero_n_rows": int((pd.to_numeric(data.get(n_col, 0), errors="coerce").fillna(0) > 0).sum()),
            }
        )
    return pd.DataFrame(rows)


def dataset_action_distribution(data: pd.DataFrame, prior_counts: pd.DataFrame) -> pd.DataFrame:
    row = {
        "dataset": "00616_all_year",
        "path": "Leave_One_experiments/all_year_offline_schedule_search/imitation_dataset/all_year_imitation_dataset.csv",
        "rows": len(data),
        "station_years": data[["station", "year"]].drop_duplicates().shape[0] if len(data) else 0,
        "nonzero_irrigation_rows": int((data["expert_action_irrigation"] > 0).sum()) if len(data) else 0,
        "nonzero_n_rows": int((data["expert_action_n"] > 0).sum()) if len(data) else 0,
    }
    comp = pd.concat([prior_counts, pd.DataFrame([row])], ignore_index=True)
    comp.to_csv(OUTPUT_ROOT / "evaluation" / "all_year_imitation_dataset_action_distribution.csv", index=False, encoding="utf-8-sig")
    md = [
        "# All-Year Imitation Dataset Check",
        "",
        df_to_markdown(comp),
        "",
        "The all-year dataset is expected to increase nonzero irrigation rows by adding water-stress and irrigation-responsive station-years from the 006_15 scenario pool.",
    ]
    (OUTPUT_ROOT / "evaluation" / "all_year_imitation_dataset_check.md").write_text("\n".join(md), encoding="utf-8")
    return comp


def prepare_features(data: pd.DataFrame) -> pd.DataFrame:
    out = data.copy()
    out["station"] = out["station"].astype(str)
    for col in FEATURE_COLUMNS:
        if col not in out.columns:
            out[col] = 0.0
    for col in [c for c in FEATURE_COLUMNS if c != "station"] + TARGET_COLUMNS:
        out[col] = pd.to_numeric(out[col], errors="coerce").fillna(0.0)
    return out


def train_imitation_models(config: dict, data: pd.DataFrame, selection: pd.DataFrame) -> tuple[dict[str, Any], pd.DataFrame, pd.DataFrame]:
    from imitation_policy_models import (
        TwoStageClassifierRegressor,
        make_mlp_regressor,
        make_random_forest_regressor,
        postprocess_actions,
        supervised_metrics,
    )

    data = prepare_features(data)
    train_years = set(tuple(x) for x in selection[selection["selected_for_train_pool"].astype(bool)][["station_code", "year"]].itertuples(index=False, name=None))
    eval_years = set(tuple(x) for x in selection[selection["selected_for_eval_pool"].astype(bool)][["station_code", "year"]].itertuples(index=False, name=None))
    data["_key"] = list(zip(data["station"].astype(str), data["year"].astype(int)))
    train = data[data["_key"].isin(train_years)].copy()
    validation = data[data["_key"].isin(eval_years)].copy()
    if validation.empty:
        validation = data.sample(frac=0.25, random_state=int(config.get("seed", 0))).copy()
    if train.empty:
        train = data.drop(validation.index).copy()

    x_train, y_train = train[FEATURE_COLUMNS], train[TARGET_COLUMNS]
    x_val, y_val = validation[FEATURE_COLUMNS], validation[TARGET_COLUMNS]
    models: dict[str, Any] = {}
    specs = {
        "BC_random_forest_regressor_all_year": make_random_forest_regressor(random_state=int(config.get("seed", 0))),
        "BC_two_stage_classifier_regressor_all_year": TwoStageClassifierRegressor(threshold=float(config["imitation"]["event_threshold"]), random_state=int(config.get("seed", 0))),
        "BC_mlp_regressor_all_year": make_mlp_regressor(random_state=int(config.get("seed", 0))),
    }
    metrics_rows = []
    for name, model in specs.items():
        print(f"[00616-train] {name}", flush=True)
        model.fit(x_train, y_train)
        pred = model.predict(x_val)
        pred = postprocess_actions(pred, threshold=float(config["imitation"]["event_threshold"]))
        metrics = supervised_metrics(y_val, pred)
        water_mask = validation["scenario_type"].astype(str).str.contains("water_stress_year", na=False)
        irr_mask = validation["scenario_type"].astype(str).str.contains("irrigation_responsive_year", na=False)
        for mask, label in [(water_mask, "water_stress"), (irr_mask, "irrigation_responsive")]:
            if mask.any():
                sub_pred = model.predict(x_val.loc[mask])
                sub_pred = postprocess_actions(sub_pred, threshold=float(config["imitation"]["event_threshold"]))
                sub_true = y_val.loc[mask].to_numpy(dtype=float)
                true_event = sub_true[:, 0] > 0
                pred_event = sub_pred[:, 0] > 0
                if true_event.any():
                    metrics[f"{label}_irrigation_event_recall"] = float(((true_event & pred_event).sum()) / true_event.sum())
                else:
                    metrics[f"{label}_irrigation_event_recall"] = np.nan
            else:
                metrics[f"{label}_irrigation_event_recall"] = np.nan
        metrics_rows.append({"policy_name": name, "train_rows": len(train), "validation_rows": len(validation), **metrics})
        with (OUTPUT_ROOT / "models" / f"{name}.pkl").open("wb") as f:
            pickle.dump(model, f)
        models[name] = model
    metrics_df = pd.DataFrame(metrics_rows)
    metrics_df.to_csv(OUTPUT_ROOT / "evaluation" / "all_year_imitation_supervised_metrics.csv", index=False, encoding="utf-8-sig")
    return models, metrics_df, validation.drop(columns=["_key"], errors="ignore")


def predict_action(model, station: str, sim_day: int, doy: int, latest: dict, threshold: float) -> dict[str, float]:
    from imitation_policy_models import postprocess_actions

    row = {
        "station": station,
        "sim_day": int(sim_day),
        "doy": int(doy),
        "topwt": scalar(latest.get("topwt", 0.0)),
        "grnwt": scalar(latest.get("grnwt", 0.0)),
        "xlai": scalar(latest.get("xlai", 0.0)),
        "totir": scalar(latest.get("totir", 0.0)),
        "tofer": scalar(latest.get("tofer", 0.0)),
        "swfac": scalar(latest.get("swfac", 0.0)),
        "nstres": scalar(latest.get("nstres", 0.0)),
    }
    features = pd.DataFrame([row], columns=FEATURE_COLUMNS)
    pred = postprocess_actions(model.predict(features), threshold=threshold)
    return {"amir": float(pred.reshape(-1, 2)[0, 0]), "anfer": float(pred.reshape(-1, 2)[0, 1])}


def evaluate_model_policy(config: dict, env_config: dict, model, policy_name: str, scenario: pd.Series, expert_lookup: pd.DataFrame) -> dict[str, Any]:
    station = str(scenario["station_code"])
    year = int(scenario["year"])
    daily_dir = OUTPUT_ROOT / "daily_outputs" / station
    daily_dir.mkdir(parents=True, exist_ok=True)
    daily_csv = daily_dir / f"{station}_{year}_{policy_name}_daily.csv"
    if daily_csv.exists():
        daily = pd.read_csv(daily_csv)
    else:
        env = make_env(env_config, station, year, int(config.get("seed", 0)), run_tag=f"{policy_name}_{station}_{year}", evaluation=True, action_safety_enabled=False)
        safety_state = ActionSafetyState()
        safety_config = {**config["action_safety"], "enabled": True}
        records = []
        try:
            obs, info = env.reset()
            done = False
            step_count = 0
            year_info = find_year(env_config, station, year)
            planting = pd.Timestamp(year_info["planting_date"])
            while not done and step_count < int(config["runtime"].get("max_steps", 260)):
                latest = latest_observation_dict(env, obs, info)
                sim_day = step_count + 1
                date = planting + pd.Timedelta(days=step_count)
                raw = predict_action(model, station, sim_day, int(date.dayofyear), latest, float(config["imitation"]["event_threshold"]))
                safety_result = apply_action_safety(raw, sim_day, safety_state, safety_config)
                safe_norm = normalize_action(env.formator.action_names, env.formator.action_space_dict, safety_result.safe_real_action)
                obs, reward, terminated, truncated, info = env.step(safe_norm)
                done = bool(terminated or truncated)
                latest = latest_observation_dict(env, obs, info)
                update_action_safety_state(safety_state, safety_result.safe_real_action, sim_day)
                records.append(
                    {
                        "station": station,
                        "year": year,
                        "policy_name": policy_name,
                        "date": date.strftime("%Y-%m-%d"),
                        "doy": int(date.dayofyear),
                        "sim_day": sim_day,
                        "dap": scalar(latest.get("dap", sim_day)),
                        "topwt": scalar(latest.get("topwt")),
                        "grnwt": scalar(latest.get("grnwt")),
                        "xlai": scalar(latest.get("xlai")),
                        "totir": scalar(latest.get("totir")),
                        "tofer": scalar(latest.get("tofer")),
                        "swfac": scalar(latest.get("swfac")),
                        "nstres": scalar(latest.get("nstres")),
                        "reward": float(reward),
                        "real_action_amir": float(safety_result.safe_real_action.get("amir", 0.0)),
                        "real_action_anfer": float(safety_result.safe_real_action.get("anfer", 0.0)),
                        "raw_real_action_amir": float(safety_result.raw_real_action.get("amir", 0.0)),
                        "raw_real_action_anfer": float(safety_result.raw_real_action.get("anfer", 0.0)),
                        "safety_rule_triggered": safety_result.safety_rule_triggered,
                        "done": done,
                    }
                )
                step_count += 1
        finally:
            try:
                env.close()
            except Exception:
                pass
        daily = pd.DataFrame(records)
        daily.to_csv(daily_csv, index=False, encoding="utf-8-sig")
    summary = summarize_daily(daily, config)
    expert = expert_lookup[(expert_lookup["station_code"].astype(str).eq(station)) & (expert_lookup["year"].astype(int).eq(year))]
    best_yield = float(expert["final_grnwt"].max()) if len(expert) else np.nan
    best_profit = float(expert["profit_low_water_cost"].max()) if len(expert) else np.nan
    return {
        "station_code": station,
        "station_name": str(scenario.get("station_name", station)),
        "year": year,
        "scenario_type": str(scenario.get("scenario_type", "")),
        "scenario_group": str(scenario.get("scenario_group", "")),
        "policy_name": policy_name,
        "run_status": "ok" if summary["episode_completed"] else "failed",
        **summary,
        "yield_loss_vs_best_expert": (best_yield - summary["final_grnwt"]) / best_yield if best_yield and not np.isnan(best_yield) else np.nan,
        "profit_gap_vs_best_expert": best_profit - summary["profit_low_water_cost"] if not np.isnan(best_profit) else np.nan,
        "daily_csv_path": str(daily_csv.relative_to(PROJECT_ROOT)),
        "notes": "action_safe_all_year_imitation_prior_eval",
    }


def evaluate_imitation_policies(config: dict, env_config: dict, models: dict[str, Any], selection: pd.DataFrame, expert: pd.DataFrame) -> pd.DataFrame:
    if not config["runtime"].get("evaluate_imitation_policy", True):
        return pd.DataFrame()
    eval_pool = selection[selection["selected_for_eval_pool"].astype(bool)].copy()
    if len(eval_pool) > int(config["runtime"].get("max_eval_station_years", 24)):
        eval_pool["priority_score"] = eval_pool["irrigation_responsive"].astype(int) * 100 + eval_pool["has_water_stress"].astype(int) * 50 + eval_pool["year"].isin([2020, 2021, 2022, 2023]).astype(int) * 10
        eval_pool = eval_pool.sort_values(["priority_score", "station_code", "year"], ascending=[False, True, True]).head(int(config["runtime"]["max_eval_station_years"]))
    rows = []
    for name, model in models.items():
        for _, scenario in eval_pool.iterrows():
            print(f"[00616-eval] {name} {scenario.station_code} {int(scenario.year)}", flush=True)
            try:
                rows.append(evaluate_model_policy(config, env_config, model, name, scenario, expert))
            except Exception:
                rows.append(
                    {
                        "station_code": str(scenario.station_code),
                        "station_name": str(scenario.station_name),
                        "year": int(scenario.year),
                        "scenario_type": str(scenario.scenario_type),
                        "scenario_group": str(scenario.scenario_group),
                        "policy_name": name,
                        "run_status": "failed",
                        "error_message": traceback.format_exc(),
                    }
                )
    result = pd.DataFrame(rows)
    result.to_csv(OUTPUT_ROOT / "evaluation" / "all_year_imitation_policy_dssat_summary.csv", index=False, encoding="utf-8-sig")
    return result


def recommend_prior(config: dict, policy_eval: pd.DataFrame, metrics: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for name, group in policy_eval.groupby("policy_name"):
        ok = group[group["run_status"].eq("ok")]
        if ok.empty:
            rows.append({"policy_name": name, "recommended": False, "reason": "no_ok_dssat_evaluation"})
            continue
        water = ok[ok["scenario_type"].astype(str).str.contains("water_stress_year", na=False)]
        mean_i = ok["total_irrigation"].mean()
        mean_n = ok["total_n"].mean()
        mean_loss = ok["yield_loss_vs_best_expert"].mean()
        water_nonzero = float((water["total_irrigation"] > 0).mean()) if len(water) else 0.0
        metric_row = metrics[metrics["policy_name"].eq(name)]
        irr_recall = float(metric_row["irrigation_event_recall"].iloc[0]) if len(metric_row) and "irrigation_event_recall" in metric_row else np.nan
        passed = bool(mean_i <= 160 and mean_n <= 250 and mean_loss <= 0.15 and water_nonzero > 0)
        rows.append(
            {
                "policy_name": name,
                "recommended": passed,
                "mean_irrigation": mean_i,
                "mean_n": mean_n,
                "mean_yield_loss_vs_best_expert": mean_loss,
                "water_stress_nonzero_irrigation_rate": water_nonzero,
                "supervised_irrigation_event_recall": irr_recall,
                "reason": "passes_basic_all_year_prior_gate" if passed else "does_not_pass_basic_all_year_prior_gate",
            }
        )
    rec = pd.DataFrame(rows).sort_values(["recommended", "water_stress_nonzero_irrigation_rate", "mean_yield_loss_vs_best_expert"], ascending=[False, False, True])
    rec.to_csv(OUTPUT_ROOT / "evaluation" / "recommended_all_year_prior_policy.csv", index=False, encoding="utf-8-sig")
    return rec


def plot_outputs(selection: pd.DataFrame, summary: pd.DataFrame, expert: pd.DataFrame, dataset_dist: pd.DataFrame, metrics: pd.DataFrame, policy_eval: pd.DataFrame) -> None:
    fig_dir = OUTPUT_ROOT / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    ok = summary[summary["run_status"].astype(str).str.startswith("ok")].copy()
    if not ok.empty:
        plt.figure(figsize=(8, 5))
        plt.scatter(ok["total_irrigation"], ok["final_grnwt"], s=10, alpha=0.35, label="searched")
        if not expert.empty:
            plt.scatter(expert["total_irrigation"], expert["final_grnwt"], s=24, alpha=0.8, label="expert")
        plt.xlabel("Total irrigation (mm)")
        plt.ylabel("Yield grnwt")
        plt.legend()
        plt.tight_layout()
        plt.savefig(fig_dir / "all_year_schedule_search_pareto_by_station.png", dpi=180)
        plt.close()

    if not expert.empty:
        coverage = expert.groupby(["scenario_group", "expert_type"]).size().reset_index(name="count")
        pivot = coverage.pivot_table(index="scenario_group", columns="expert_type", values="count", fill_value=0)
        ax = pivot.plot(kind="bar", stacked=True, figsize=(9, 4.5))
        ax.set_ylabel("Expert schedules")
        ax.grid(axis="y", alpha=0.25)
        ax.get_figure().tight_layout()
        ax.get_figure().savefig(fig_dir / "all_year_expert_library_scenario_coverage.png", dpi=180)
        plt.close(ax.get_figure())

    if not dataset_dist.empty:
        ax = dataset_dist.set_index("dataset")[["nonzero_irrigation_rows", "nonzero_n_rows"]].plot(kind="bar", figsize=(9, 4.5))
        ax.set_ylabel("Nonzero action rows")
        ax.grid(axis="y", alpha=0.25)
        ax.get_figure().tight_layout()
        ax.get_figure().savefig(fig_dir / "all_year_imitation_action_distribution.png", dpi=180)
        plt.close(ax.get_figure())

    dataset_path = OUTPUT_ROOT / "imitation_dataset" / "all_year_imitation_dataset.csv"
    if dataset_path.exists():
        data = pd.read_csv(dataset_path)
        irr = data[pd.to_numeric(data["expert_action_irrigation"], errors="coerce").fillna(0) > 0]
        if not irr.empty:
            ax = irr["dap"].plot(kind="hist", bins=30, figsize=(8, 4), color="#4F81BD")
            ax.set_xlabel("DAP")
            ax.set_ylabel("Irrigation event rows")
            ax.get_figure().tight_layout()
            ax.get_figure().savefig(fig_dir / "irrigation_event_dap_distribution.png", dpi=180)
            plt.close(ax.get_figure())

    if not metrics.empty:
        cols = [c for c in ["irrigation_event_recall", "nitrogen_event_recall", "irrigation_mae", "nitrogen_mae"] if c in metrics.columns]
        ax = metrics.set_index("policy_name")[cols].plot(kind="bar", figsize=(10, 4.5))
        ax.grid(axis="y", alpha=0.25)
        ax.get_figure().tight_layout()
        ax.get_figure().savefig(fig_dir / "all_year_bc_supervised_metrics.png", dpi=180)
        plt.close(ax.get_figure())

    if not policy_eval.empty:
        ok_eval = policy_eval[policy_eval["run_status"].eq("ok")].copy()
        if not ok_eval.empty:
            plt.figure(figsize=(8, 5))
            for name, group in ok_eval.groupby("policy_name"):
                plt.scatter(group["total_irrigation"], group["final_grnwt"], label=name, alpha=0.75)
            plt.xlabel("Total irrigation (mm)")
            plt.ylabel("Yield grnwt")
            plt.legend(fontsize=7)
            plt.tight_layout()
            plt.savefig(fig_dir / "all_year_policy_yield_vs_input.png", dpi=180)
            plt.close()

            ax = ok_eval.groupby("policy_name")["profit_low_water_cost"].mean().sort_values().plot(kind="barh", figsize=(8, 4))
            ax.set_xlabel("Mean low-water-cost profit")
            ax.get_figure().tight_layout()
            ax.get_figure().savefig(fig_dir / "all_year_policy_profit_comparison.png", dpi=180)
            plt.close(ax.get_figure())

            water = ok_eval[ok_eval["scenario_type"].astype(str).str.contains("water_stress_year", na=False)]
            if not water.empty:
                ax = water.groupby("policy_name")["total_irrigation"].mean().sort_values().plot(kind="barh", figsize=(8, 4))
                ax.set_xlabel("Mean irrigation on water-stress years")
                ax.get_figure().tight_layout()
                ax.get_figure().savefig(fig_dir / "water_stress_year_policy_comparison.png", dpi=180)
                plt.close(ax.get_figure())

            irr_resp = ok_eval[ok_eval["scenario_type"].astype(str).str.contains("irrigation_responsive_year", na=False)]
            if not irr_resp.empty:
                ax = irr_resp.groupby("policy_name")["total_irrigation"].mean().sort_values().plot(kind="barh", figsize=(8, 4))
                ax.set_xlabel("Mean irrigation on irrigation-responsive years")
                ax.get_figure().tight_layout()
                ax.get_figure().savefig(fig_dir / "irrigation_responsive_year_policy_comparison.png", dpi=180)
                plt.close(ax.get_figure())


def write_report(config: dict, selection: pd.DataFrame, summary: pd.DataFrame, expert: pd.DataFrame, data: pd.DataFrame, dataset_dist: pd.DataFrame, metrics: pd.DataFrame, policy_eval: pd.DataFrame, rec: pd.DataFrame) -> None:
    selected = selection[selection["selected_for_search"].astype(bool)].copy()
    ok = summary[summary["run_status"].astype(str).str.startswith("ok")]
    next_step = "006_17_all_year_constrained_ppo_with_all_year_prior" if (not rec.empty and bool(rec["recommended"].any())) else "improve_all_year_expert_library_before_constrained_ppo"
    lines = [
        "# All-Year Offline Schedule Search and Imitation Prior Report",
        "",
        "Generated at: 2026-06-06",
        "",
        "## Scope",
        "",
        "This stage uses the 006_15 all-year weather scenario pool to run deterministic schedule search and train imitation priors. It does not train PPO, does not enter rainfall scaling, and does not modify my_data.",
        "",
        "## Station-Year Selection",
        "",
        f"- Selected station-years: {len(selected)}",
        f"- Selected water-stress years: {int(selected['has_water_stress'].sum())}",
        f"- Selected irrigation-responsive years: {int(selected['irrigation_responsive'].sum())}",
        f"- Selected 2020-2023 calibration/validation candidates: {int(selected['year'].isin([2020, 2021, 2022, 2023]).sum())}",
        "",
        df_to_markdown(selected.groupby("station_code").agg(station_years=("year", "count"), water_stress=("has_water_stress", "sum"), irrigation_responsive=("irrigation_responsive", "sum")).reset_index()),
        "",
        "## Schedule Search",
        "",
        f"- Schedule evaluations completed or cached: {len(ok)} / {len(summary)}",
        f"- Expert schedules retained: {len(expert)}",
        f"- Expert schedules with nonzero irrigation: {int((expert['total_irrigation'] > 0).sum()) if len(expert) else 0}",
        "",
        "The expert library intentionally keeps top yield, top low-water-cost profit, Pareto-balanced, low-input-within-10%-yield-loss, and irrigation-positive schedules. It therefore does not collapse to profit-only or irrigation=0-only schedules.",
        "",
        "## Imitation Dataset",
        "",
        df_to_markdown(dataset_dist),
        "",
        "## Supervised Imitation Metrics",
        "",
        df_to_markdown(metrics),
        "",
        "## DSSAT/gym-DSSAT Prior Evaluation",
        "",
        df_to_markdown(policy_eval.groupby("policy_name").agg(ok=("run_status", lambda s: int((s == "ok").sum())), mean_yield=("final_grnwt", "mean"), mean_irrigation=("total_irrigation", "mean"), mean_n=("total_n", "mean"), mean_yield_loss=("yield_loss_vs_best_expert", "mean")).reset_index() if not policy_eval.empty else pd.DataFrame()),
        "",
        "## Recommended Prior",
        "",
        df_to_markdown(rec),
        "",
        "## Next Step",
        "",
        f"Recommended next route: `{next_step}`.",
    ]
    DOC_MD.write_text("\n".join(lines), encoding="utf-8")
    reports_dir = OUTPUT_ROOT / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(DOC_MD, reports_dir / DOC_MD.name)
    if Presentation is not None:
        write_ppt(selection, summary, expert, dataset_dist, metrics, policy_eval, rec)
        shutil.copyfile(DOC_PPT, reports_dir / DOC_PPT.name)


def add_title(slide, text: str):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.45))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_body(slide, text: str, top: float = 0.9):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.0), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(0, 0, 0)


def add_picture_if_exists(slide, path: Path, top: float = 1.2):
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(0.7), Inches(top), width=Inches(11.8))


def write_ppt(selection: pd.DataFrame, summary: pd.DataFrame, expert: pd.DataFrame, dataset_dist: pd.DataFrame, metrics: pd.DataFrame, policy_eval: pd.DataFrame, rec: pd.DataFrame) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    selected = selection[selection["selected_for_search"].astype(bool)]
    slides = [
        ("006_16 all-year offline schedule search", f"Selected station-years: {len(selected)}\nWater-stress years: {int(selected['has_water_stress'].sum())}\nIrrigation-responsive years: {int(selected['irrigation_responsive'].sum())}\nNo PPO training; no rainfall scaling; no my_data modification."),
        ("Why this stage", "006_15 showed that all-year weather increases swfac samples compared with observed-year-only.\nThis stage builds expert schedules first, then trains imitation priors.\nThe goal is an interpretable prior for later constrained PPO, not a final RL result."),
        ("Schedule search design", "N events: DAP 1, 30, 60; total N cap 250 kg/ha.\nIrrigation events: DAP 20, 35, 50, 65, 80, 95 for stress/response years; simplified events for low-response years.\nExpert schedules keep profit, yield, Pareto-balanced, low-input, and irrigation-positive categories."),
        ("Key counts", f"Schedule evaluations: {len(summary)}\nExpert schedules retained: {len(expert)}\nExpert schedules with irrigation > 0: {int((expert['total_irrigation'] > 0).sum()) if len(expert) else 0}\nImitation rows: {int(dataset_dist.loc[dataset_dist['dataset'].eq('00616_all_year'), 'rows'].iloc[0]) if len(dataset_dist) and (dataset_dist['dataset'].eq('00616_all_year')).any() else 0}"),
    ]
    for title, body in slides:
        slide = prs.slides.add_slide(blank)
        add_title(slide, title)
        add_body(slide, body)
    for title, fig in [
        ("Expert library coverage", "all_year_expert_library_scenario_coverage.png"),
        ("Action distribution", "all_year_imitation_action_distribution.png"),
        ("BC supervised metrics", "all_year_bc_supervised_metrics.png"),
        ("Policy yield vs input", "all_year_policy_yield_vs_input.png"),
        ("Water-stress policy comparison", "water_stress_year_policy_comparison.png"),
    ]:
        slide = prs.slides.add_slide(blank)
        add_title(slide, title)
        add_picture_if_exists(slide, OUTPUT_ROOT / "figures" / fig)
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Recommendation")
    best = rec.iloc[0].to_dict() if len(rec) else {}
    add_body(slide, "\n".join([f"{k}: {v}" for k, v in best.items()]) if best else "No recommended learned prior yet.")
    prs.save(DOC_PPT)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default=str(CONFIG_PATH))
    parser.add_argument("--selection-only", action="store_true")
    parser.add_argument("--search-only", action="store_true")
    parser.add_argument("--report-only", action="store_true")
    args = parser.parse_args()
    config = load_yaml(Path(args.config))
    ensure_dirs()
    shutil.copyfile(Path(args.config), OUTPUT_ROOT / "configs" / Path(args.config).name)
    pool = read_csv(config["paths"]["scenario_pool_csv"])
    plan = read_csv(config["paths"]["calibration_plan_csv"])
    selection_path = OUTPUT_ROOT / "configs" / "all_year_search_year_selection.csv"
    if args.report_only and selection_path.exists():
        selection = pd.read_csv(selection_path)
    else:
        selection = select_search_years(config, pool, plan)
    if args.selection_only:
        print(selection_path.relative_to(PROJECT_ROOT))
        return
    env_config = build_env_config(config, selection[selection["selected_for_search"].astype(bool)])
    summary_path = OUTPUT_ROOT / "evaluation" / "all_year_schedule_search_summary.csv"
    expert_path = OUTPUT_ROOT / "expert_policy" / "all_year_expert_schedule_library.csv"
    data_path = OUTPUT_ROOT / "imitation_dataset" / "all_year_imitation_dataset.csv"
    metrics_path = OUTPUT_ROOT / "evaluation" / "all_year_imitation_supervised_metrics.csv"
    eval_path = OUTPUT_ROOT / "evaluation" / "all_year_imitation_policy_dssat_summary.csv"
    rec_path = OUTPUT_ROOT / "evaluation" / "recommended_all_year_prior_policy.csv"
    if args.report_only and summary_path.exists() and expert_path.exists() and data_path.exists() and metrics_path.exists():
        summary = pd.read_csv(summary_path)
        expert = pd.read_csv(expert_path)
        data = pd.read_csv(data_path)
        metrics = pd.read_csv(metrics_path)
        policy_eval = pd.read_csv(eval_path) if eval_path.exists() else pd.DataFrame()
        rec = pd.read_csv(rec_path) if rec_path.exists() else pd.DataFrame()
        dataset_dist = pd.read_csv(OUTPUT_ROOT / "evaluation" / "all_year_imitation_dataset_action_distribution.csv")
    else:
        summary = run_schedule_search(config, env_config, selection)
        expert = select_expert_library(config, summary)
        data = build_imitation_dataset(expert)
        prior_counts = prior_dataset_counts()
        dataset_dist = dataset_action_distribution(data, prior_counts)
        if args.search_only:
            plot_outputs(selection, summary, expert, dataset_dist, pd.DataFrame(), pd.DataFrame())
            print((OUTPUT_ROOT / "configs" / "all_year_search_year_selection.csv").relative_to(PROJECT_ROOT))
            print(summary_path.relative_to(PROJECT_ROOT))
            print(expert_path.relative_to(PROJECT_ROOT))
            print(data_path.relative_to(PROJECT_ROOT))
            return
        models, metrics, _validation = train_imitation_models(config, data, selection)
        policy_eval = evaluate_imitation_policies(config, env_config, models, selection, expert)
        rec = recommend_prior(config, policy_eval, metrics)
    plot_outputs(selection, summary, expert, dataset_dist, metrics, policy_eval)
    write_report(config, selection, summary, expert, data, dataset_dist, metrics, policy_eval, rec)
    print((OUTPUT_ROOT / "configs" / "all_year_search_year_selection.csv").relative_to(PROJECT_ROOT))
    print(summary_path.relative_to(PROJECT_ROOT))
    print(expert_path.relative_to(PROJECT_ROOT))
    print(data_path.relative_to(PROJECT_ROOT))
    print(metrics_path.relative_to(PROJECT_ROOT))
    print(DOC_MD.relative_to(PROJECT_ROOT))
    print(DOC_PPT.relative_to(PROJECT_ROOT))


if __name__ == "__main__":
    main()
